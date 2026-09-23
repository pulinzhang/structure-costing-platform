"""写出芜湖 T13J · 2026年07月材料变更再变样表（变回去）。不 import 业务包。

必须先成功导入 2026年06月那一本。本本变更前 = 库里最新 DC03，变更后回到 DC01。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

OUT = Path(__file__).resolve().parent / "t13j-2026-07-material-change.xlsx"

SHEET_HELP = "材料变更填写说明"
SHEET_ORDER = "工程变更单（材质）"

HEADERS = (
    "基地",
    "项目",
    "变更版本",
    "生效年月",
    "冲压工厂",
    "零件号",
    "零件名称",
    "变更前材料",
    "变更前料厚",
    "变更后材料",
    "变更后料厚",
    "互换性",
)

ROWS = (
    (
        "芜湖",
        "T13J",
        "MAT-202607",
        "2026年07月",
        "安徽振华汽车部件有限责任公司",
        "403014115AA",
        "左B柱气囊导向支架",
        "DC03",
        1.5,
        "DC01",
        1.5,
        "不影响互换性",
    ),
    (
        "芜湖",
        "T13J",
        "MAT-202607",
        "2026年07月",
        "安徽振华汽车部件有限责任公司",
        "403014116AA",
        "右B柱气囊导向支架",
        "DC03",
        1.5,
        "DC01",
        1.5,
        "不影响互换性",
    ),
)

HELP_LINES = (
    "材料变更填写说明（再变回去，测试配套）",
    "",
    "先导入 t13j-2026-06-material-change.xlsx（MAT-202606，DC01→DC03）。",
    "再本本：上传页必须选 芜湖 / T13J / 2026 / 07。",
    "变更前必须是库里最新材料 DC03，不能写工艺投产列 DC01。",
    "变更后回到 DC01（再变回去）。工艺行材料始终仍是 DC01。",
    "变更前料厚、变更后料厚必须是数字，大于 0，最多两位小数，不要写成文字。变更前必须等于工艺料厚。变更后相对变更前不能超过百分之二十。本两件都是 1.5。",
    "变更版本 MAT-202607 配 2026年07月。只准两张表。禁止合并、公式、修空格。",
    "本样两行：403014115AA / 403014116AA，DC03→DC01，料厚 1.5→1.5。",
)

_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
_DATA_FONT = Font(name="微软雅黑", size=9)
_TITLE_FONT = Font(name="微软雅黑", size=14, bold=True, color="1F4E79")
_HELP_FONT = Font(name="微软雅黑", size=10)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_WIDTHS = (10, 10, 14, 14, 36, 16, 22, 14, 12, 14, 12, 16)


def _write_help(sheet: Worksheet) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = HELP_LINES[0]
    sheet["A1"].font = _TITLE_FONT
    for index, line in enumerate(HELP_LINES[1:], start=2):
        sheet.cell(index, 1, line).font = _HELP_FONT
        sheet.cell(index, 1).alignment = _LEFT
    sheet.column_dimensions["A"].width = 110
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"


def _write_order(sheet: Worksheet) -> None:
    for col, title in enumerate(HEADERS, start=1):
        cell = sheet.cell(1, col, title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN
        sheet.column_dimensions[get_column_letter(col)].width = _WIDTHS[col - 1]
    for row_index, values in enumerate(ROWS, start=2):
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, col, value)
            cell.font = _DATA_FONT
            cell.alignment = _CENTER
            cell.border = _THIN
            if col in (9, 11):
                cell.number_format = "0.00"

        sheet.row_dimensions[row_index].height = 22
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{1 + len(ROWS)}"


def main() -> None:
    workbook = Workbook()
    help_sheet = workbook.active
    assert help_sheet is not None
    help_sheet.title = SHEET_HELP
    _write_help(help_sheet)
    _write_order(workbook.create_sheet(SHEET_ORDER))
    workbook.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
