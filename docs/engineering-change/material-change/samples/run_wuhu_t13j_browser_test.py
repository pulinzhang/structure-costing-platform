"""芜湖 T13J 材料变更网页测试。

先开开发服务（http://127.0.0.1:8000），在仓库根目录执行：

    python -X utf8 docs/engineering-change/material-change/samples/run_wuhu_t13j_browser_test.py

样表就在本目录。负例由脚本改格子，写到本目录 `_cases/`，跑完删掉。
库里若已有 MAT-202607 / MAT-202606，脚本会先回滚再测。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import Page, sync_playwright

BASE = "http://127.0.0.1:8000"
SAMPLE_DIR = Path(__file__).resolve().parent
SAMPLE_06 = SAMPLE_DIR / "t13j-2026-06-material-change.xlsx"
SAMPLE_07 = SAMPLE_DIR / "t13j-2026-07-material-change.xlsx"
CASE_DIR = SAMPLE_DIR / "_cases"
SHEET_ORDER = "工程变更单（材质）"
LOCATION = "芜湖"
VEHICLE = "T13J"


def login(page: Page) -> None:
    page.goto(f"{BASE}/login", wait_until="domcontentloaded")
    page.fill("#loginUsername", "admin")
    page.fill("#loginPassword", "admin123")
    page.click("#btnLogin")
    page.wait_for_url("**/app**", timeout=20000)


def set_select(page: Page, select_id: str, value: str) -> None:
    page.evaluate(
        """([selectId, value]) => {
            const select = document.getElementById(selectId);
            const opt = [...select.options].find(
                (item) => item.value === value || (item.textContent || '').trim() === value
            );
            if (!opt) throw new Error(selectId + ' missing ' + value);
            select.value = opt.value;
            select.dispatchEvent(new Event('change', { bubbles: true }));
        }""",
        [select_id, value],
    )


def save_case(name: str, source: Path, editor) -> Path:
    CASE_DIR.mkdir(exist_ok=True)
    workbook = load_workbook(source)
    editor(workbook)
    path = CASE_DIR / name
    workbook.save(path)
    workbook.close()
    return path


def upload(page: Page, path: Path, year: int, month: int) -> str:
    page.goto(f"{BASE}/app#/engineering-change-material-upload", wait_until="domcontentloaded")
    page.wait_for_function(
        "() => document.querySelectorAll('#ecmUploadLocation option').length > 1",
        timeout=20000,
    )
    set_select(page, "ecmUploadLocation", LOCATION)
    page.wait_for_function(
        "() => [...document.querySelectorAll('#ecmUploadVehicleModel option')].some(o => o.value === 'T13J')",
        timeout=20000,
    )
    set_select(page, "ecmUploadVehicleModel", VEHICLE)
    set_select(page, "ecmUploadYear", str(year))
    set_select(page, "ecmUploadMonth", str(month))
    page.set_input_files("#ecmUploadFileInput", str(path.resolve()))
    page.wait_for_function(
        "() => { const btn = document.getElementById('ecmUploadSubmitBtn'); return btn && !btn.disabled; }",
        timeout=10000,
    )
    page.click("#ecmUploadSubmitBtn")
    page.wait_for_function(
        """() => {
            const title = document.getElementById('ecmProgressTitle');
            const err = document.getElementById('ecmUploadErrorPanel');
            const text = title ? title.textContent || '' : '';
            const errVisible = err && !err.classList.contains('hidden');
            return text.includes('解析成功') || text.includes('解析失败') || errVisible;
        }""",
        timeout=120000,
    )
    panel = page.locator("#ecmUploadErrorPanel")
    if panel.count() and panel.first.is_visible():
        return page.locator("#ecmUploadErrorBody").inner_text().strip()
    title = page.locator("#ecmProgressTitle").inner_text().strip()
    if "失败" in title:
        return page.locator("#ecmUploadErrorBody").inner_text().strip()
    return "ok"


def expect_error(page: Page, path: Path, year: int, month: int, phrase: str, label: str) -> None:
    message = upload(page, path, year, month)
    print(f"{label}: {message.splitlines()[0][:120]}")
    if message == "ok" or phrase not in message:
        raise SystemExit(f"{label} 未报「{phrase}」\n{message[:400]}")


def open_filtered_list(page: Page, status: str) -> None:
    page.goto(f"{BASE}/app#/engineering-change-material-list", wait_until="domcontentloaded")
    page.wait_for_function(
        "() => document.querySelectorAll('#ecmListLocationFilter option').length > 1",
        timeout=20000,
    )
    set_select(page, "ecmListLocationFilter", LOCATION)
    page.wait_for_function(
        "() => [...document.querySelectorAll('#ecmListVehicleFilter option')].some(o => o.value === 'T13J')",
        timeout=20000,
    )
    set_select(page, "ecmListVehicleFilter", VEHICLE)
    page.wait_for_function(
        """() => {
            const loading = document.getElementById('ecmLoading');
            return !loading || loading.classList.contains('hidden');
        }""",
        timeout=20000,
    )
    with page.expect_response(
        lambda response: "/api/engineering-change/material/uploads" in response.url
        and f"status={status}" in response.url
        and response.status == 200,
        timeout=20000,
    ):
        set_select(page, "ecmListStatusFilter", status)
    page.wait_for_function(
        """() => {
            const loading = document.getElementById('ecmLoading');
            const body = document.getElementById('ecmTableBody');
            const empty = document.getElementById('ecmEmpty');
            const hidden = !loading || loading.classList.contains('hidden');
            const ready = (body && body.querySelector('tr')) || (empty && !empty.classList.contains('hidden'));
            return hidden && ready;
        }""",
        timeout=20000,
    )


def confirm_delete(page: Page) -> None:
    page.locator("#appConfirmOkBtn").click()


def close_delete_error(page: Page) -> None:
    dialog = page.locator("#appDeleteErrorDialog")
    if dialog.count() and dialog.first.is_visible():
        page.click("#appDeleteErrorOkBtn")


def purge_failed(page: Page) -> None:
    for _ in range(30):
        open_filtered_list(page, "failed")
        button = page.locator("#ecmTableBody button[data-action='delete']")
        if button.count() == 0:
            return
        button.first.click()
        confirm_delete(page)
        page.wait_for_timeout(500)
    raise SystemExit("失败头删不完")


def rollback_version(page: Page, version: str) -> str:
    """成功列表里删这一本。没有这本返回 missing。被更晚刀拦住时返回提示全文。"""
    close_delete_error(page)
    open_filtered_list(page, "success")
    row = page.locator("#ecmTableBody tr").filter(has_text=version)
    if row.count() == 0:
        return "missing"
    row.first.locator("button[data-action='delete']").click()
    confirm_delete(page)
    for _ in range(40):
        dialog = page.locator("#appDeleteErrorDialog")
        if dialog.count() and dialog.first.is_visible():
            message = page.locator("#appDeleteErrorBody").inner_text().strip()
            page.click("#appDeleteErrorOkBtn")
            return message
        body = page.locator("#ecmTableBody")
        if body.count() and version not in body.inner_text():
            return "ok"
        page.wait_for_timeout(250)
    raise SystemExit(f"回滚 {version} 没有结果")


def clear_existing_books(page: Page) -> None:
    for version in ("MAT-202607", "MAT-202606"):
        result = rollback_version(page, version)
        print(f"清旧本 {version}: {result.splitlines()[0][:120]}")
        if result not in ("ok", "missing"):
            raise SystemExit(f"清不掉 {version}")


def drawer_text(page: Page) -> str:
    page.goto(
        f"{BASE}/app#/stamping-parts?location=%E8%8A%9C%E6%B9%96&vehicle_model=T13J&part_number=403014115AA",
        wait_until="domcontentloaded",
    )
    page.wait_for_function(
        """() => {
            const loading = document.getElementById('partsLoading');
            const body = document.getElementById('partsTableBody');
            const hidden = !loading || loading.classList.contains('hidden');
            return hidden && body && body.innerText.includes('403014115AA');
        }""",
        timeout=20000,
    )
    page.locator("#partsTableBody button[data-action='detail']").first.click()
    page.wait_for_function(
        "() => (document.getElementById('partDrawerContent')?.innerText || '').includes('材料变更')",
        timeout=20000,
    )
    return page.locator("#partDrawerContent").inner_text()


def both_rows(sheet, column: int, value) -> None:
    sheet.cell(2, column, value)
    sheet.cell(3, column, value)


def build_cases() -> list[tuple[Path, int, int, str, str]]:
    def sheet_name(workbook) -> None:
        workbook[SHEET_ORDER].title = "工程变更单"

    def header(workbook) -> None:
        sheet = workbook[SHEET_ORDER]
        sheet.insert_rows(1)
        sheet.cell(1, 1, "标题")

    def bad_month(workbook) -> None:
        both_rows(workbook[SHEET_ORDER], 4, "2026年13月")

    def bad_version(workbook) -> None:
        both_rows(workbook[SHEET_ORDER], 3, "MAT-202607")

    def same_material(workbook) -> None:
        workbook[SHEET_ORDER].cell(2, 10, "DC01")

    def thick_before(workbook) -> None:
        workbook[SHEET_ORDER].cell(2, 9, 1.6)

    def thick_delta(workbook) -> None:
        workbook[SHEET_ORDER].cell(2, 11, 2.0)

    def thick_format(workbook) -> None:
        workbook[SHEET_ORDER].cell(2, 11, 1.555)

    def space(workbook) -> None:
        workbook[SHEET_ORDER].cell(2, 12, "不影响互换性 ")

    def numeric_part(workbook) -> None:
        workbook[SHEET_ORDER].cell(2, 6, 403014115)

    def duplicate_part(workbook) -> None:
        sheet = workbook[SHEET_ORDER]
        sheet.cell(3, 6, "403014115AA")
        sheet.cell(3, 7, "左B柱气囊导向支架")

    return [
        (save_case("sheet.xlsx", SAMPLE_06, sheet_name), 2026, 6, "请只留两张表", "表名不对"),
        (save_case("header.xlsx", SAMPLE_06, header), 2026, 6, "第 1 行", "表头不在第1行"),
        (save_case("month.xlsx", SAMPLE_06, bad_month), 2026, 6, "2026年6月", "生效月13"),
        (
            save_case("version.xlsx", SAMPLE_06, bad_version),
            2026,
            6,
            "MAT-202606",
            "版本和月份不符",
        ),
        (
            save_case("same.xlsx", SAMPLE_06, same_material),
            2026,
            6,
            "变更前和变更后材料写成一样",
            "前后材料相同",
        ),
        (
            save_case("thick-before.xlsx", SAMPLE_06, thick_before),
            2026,
            6,
            "变更前料厚",
            "变更前料厚不对",
        ),
        (
            save_case("thick-delta.xlsx", SAMPLE_06, thick_delta),
            2026,
            6,
            "超过两成",
            "料厚差超过两成",
        ),
        (
            save_case("thick-format.xlsx", SAMPLE_06, thick_format),
            2026,
            6,
            "最多两位小数",
            "料厚三位小数",
        ),
        (save_case("space.xlsx", SAMPLE_06, space), 2026, 6, "前后空格", "互换性多空格"),
        (save_case("part.xlsx", SAMPLE_06, numeric_part), 2026, 6, "改成文字", "零件号是数字"),
        (
            save_case("dup-row.xlsx", SAMPLE_06, duplicate_part),
            2026,
            6,
            "写了不止一行",
            "同一零件写两行",
        ),
        (SAMPLE_07, 2026, 7, "正在用的牌号", "没传06就传07"),
    ]


def main() -> None:
    if not SAMPLE_06.is_file() or not SAMPLE_07.is_file():
        raise SystemExit(f"样表不在 {SAMPLE_DIR}")
    cases = build_cases()

    def wrong_before(workbook) -> None:
        both_rows(workbook[SHEET_ORDER], 8, "DC01")
        both_rows(workbook[SHEET_ORDER], 10, "DC03")

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1440, "height": 900})
        page.set_default_timeout(20000)
        login(page)
        clear_existing_books(page)

        for path, year, month, phrase, label in cases:
            expect_error(page, path, year, month, phrase, label)
        purge_failed(page)
        print("失败头已清掉")

        if upload(page, SAMPLE_06, 2026, 6) != "ok":
            raise SystemExit("06 应该成功")
        print("06 成功")
        expect_error(page, SAMPLE_06, 2026, 6, "这个版本已经传过了", "重复上传 MAT-202606")

        expect_error(
            page,
            save_case("07-before.xlsx", SAMPLE_07, wrong_before),
            2026,
            7,
            "正在用的牌号",
            "07变更前仍写DC01",
        )
        purge_failed(page)

        if upload(page, SAMPLE_07, 2026, 7) != "ok":
            raise SystemExit("07 应该成功")
        print("07 成功")
        expect_error(page, SAMPLE_07, 2026, 7, "这个版本已经传过了", "重复上传 MAT-202607")
        purge_failed(page)

        blocked = rollback_version(page, "MAT-202606")
        print("06 在 07 还在时:", blocked.splitlines()[0][:120])
        if "更晚" not in blocked:
            raise SystemExit("06 被错误回滚了")

        drawer = drawer_text(page)
        if "MAT-202606" not in drawer or "MAT-202607" not in drawer:
            raise SystemExit("抽屉两刀不在")
        if "2025年01月" not in drawer or "DC01" not in drawer.split("材料变更")[0]:
            raise SystemExit("工艺材料或生失效被改了")
        print("抽屉两刀在，工艺仍是 DC01 / 2025年01月")

        july_rollback = rollback_version(page, "MAT-202607")
        print("07 回滚:", july_rollback.splitlines()[0][:200])
        if july_rollback != "ok":
            raise SystemExit("07 回滚失败")
        if rollback_version(page, "MAT-202606") != "ok":
            raise SystemExit("06 回滚失败")
        drawer = drawer_text(page)
        if "暂无材料变更" not in drawer or "MAT-2026" in drawer:
            raise SystemExit("回滚后抽屉没清空")
        print("两本回滚后抽屉已空")
        browser.close()

    if CASE_DIR.is_dir():
        for path in CASE_DIR.glob("*.xlsx"):
            path.unlink()
        CASE_DIR.rmdir()
    print("完成")


if __name__ == "__main__":
    main()
