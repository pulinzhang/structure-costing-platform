"""青岛 T28 · 2025年07月 工程变更导入测试样本：一本 Excel、封皮五行。"""

from __future__ import annotations

import sqlite3
import sys
from copy import copy
from dataclasses import dataclass, field, replace
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

REPO = Path(__file__).resolve().parents[3]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

from app.parsers.bom.meta_parser import parse_bom_meta  # noqa: E402
from app.parsers.bom.node_parser import parse_bom_nodes  # noqa: E402
from app.services.bom.bom_template_builder import (  # noqa: E402
    build_bom_material_template,
)
from app.services.stamping.die.template_builder import (  # noqa: E402
    COL_COUNT as DIE_COL_COUNT,
)
from app.services.stamping.die.template_builder import (  # noqa: E402
    build_die_layout_template,
)
from app.services.stamping.upload.template_builder import (  # noqa: E402
    build_stamping_process_template,
)

SAMPLES = Path(__file__).resolve().parent
OUT = SAMPLES / "t28-2025-07-import-bundle.xlsx"
DATABASE_PATH = REPO / "data" / "stamping_costing.db"

LOCATION = "青岛"
VEHICLE_MODEL = "T28"
CUSTOMER_NAME = "奇瑞公司"
AS_OF_YEAR = 2025
AS_OF_MONTH = 7
PERIOD_LABEL = "2025年07月"
CHANGE_VERSION = "EC-202507"
STAMPING_SUPPLIER = "烟台海众汽车零部件有限公司"
SLASH = "/"
NOT_ASSEMBLED = "此件不装"
ADD_PART = "增加零件"
ADD_PART_NUMBER = "501007040AA"
ADD_PART_NAME = "中通道后纵梁加强板"
ADD_PART_INTERNAL = "ACA613002"
_ADD_PART_FILL = PatternFill("solid", fgColor="E2EFDA")

_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_DATA_FONT = Font(name="微软雅黑", size=9)
_NOT_ASSEMBLED_FILL = PatternFill("solid", fgColor="FFE4E4")
_DIE_DATA_FONT = Font(name="宋体", size=12)
_DIE_NOTE_TITLE_FONT = Font(name="宋体", size=12, bold=True)
_DIE_NOTE_FONT = Font(name="宋体", size=10)
_DIE_NOTE_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
_DIE_PART_COLS = (1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 23, 25, 26, 27, 28, 29, 30)
_DIE_ALLOC_NOTES = (
    "【模具费分摊说明】上传只读上方表格，不读本段。",
    "1. 先选采购基地并填写项目（新项目在本页填写）。模具方式或生产方式含「双槽 / 合并 / 一出二」→ 各序默认分摊 0.50；否则 1.00（单干）。左右件各记一半，避免同一副模记两次。",
    "2. 共模件若第一序单独开模（常见落料），上传后到模具清单把该序比例改成 1，后面共模序保持 0.50。",
    "3. 级进模写在生产方式。模具费不打 0.80（那是加工费折扣）。级进且共模仍是 0.50。",
    "4. 模具编号可空，系统按 DP-40-项目-流水自动补号；已填写的编号保持不变。",
    "5. 预估重量（T）= 长×宽×高×掏空×7.85÷10⁹。工序费 = 重量 × 单价 × 分摊比例 ÷ 分摊万台。",
)
_DIE_THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)


@dataclass
class BomNode:
    node_id: int
    level: int
    seq_no: int | None
    internal_number: str | None
    part_number: str
    part_name: str | None
    qty: Decimal | None
    qty_text: str | None
    weld_info: str | None
    material: str | None
    thickness_text: str | None
    material_mill: str | None
    part_weight: Decimal | None
    gauge_developed: str | None
    remark: str | None
    children: list[BomNode] = field(default_factory=list)


@dataclass(frozen=True)
class ProcessOp:
    process_seq: str | None
    process_code: str | None
    die_sets: int | None
    die_number: str | None
    tonnage: int | None


@dataclass(frozen=True)
class ProcessPart:
    part_number: str
    part_name: str | None
    material: str | None
    thickness: float | Decimal | None
    part_weight: float | Decimal | None
    sheet_length: float | Decimal | None
    sheet_width: float | Decimal | None
    parts_per_sheet: int | None
    sheet_type: str | None
    die_type: str | None
    operations: list[ProcessOp]


OUT = SAMPLES / "t28-2025-07-import-bundle.xlsx"


@dataclass(frozen=True)
class CoverRow:
    change_type: str
    before_assembly: str
    after_assembly: str | None
    assembly_name: str
    part_level: str
    before_part: str | None
    after_part: str | None
    part_name: str | None


@dataclass(frozen=True)
class Scenario:
    cover: CoverRow
    root_part_number: str
    bump_root: bool
    extra_old: str | None = None
    extra_new: str | None = None
    not_assembled_part_number: str | None = None
    process_old_part_number: str | None = None
    added_part_number: str | None = None
    added_part_name: str | None = None
    added_internal_number: str | None = None
    added_material: str | None = None
    added_thickness_text: str | None = None
    added_part_weight: Decimal | None = None


SCENARIOS = (
    Scenario(
        cover=CoverRow(
            "完成停止",
            "501009533AA",
            None,
            "发盖左铰链安装支架总成",
            "总成",
            None,
            None,
            None,
        ),
        root_part_number="501009533AA",
        bump_root=False,
    ),
    Scenario(
        cover=CoverRow(
            "改图",
            "501006873AA",
            "501006873AB",
            "流水槽加强梁总成",
            "一级",
            "501006875AA",
            "501006875AB",
            "雨刮电机安装支架Ⅰ",
        ),
        root_part_number="501006873AA",
        bump_root=True,
        extra_old="501006875AA",
        extra_new="501006875AB",
        process_old_part_number="501006875AA",
    ),
    Scenario(
        cover=CoverRow(
            "改图",
            "501009144AA",
            "501009144AB",
            "后防撞梁总成",
            "二级",
            "501009122AA",
            "501009122AB",
            "后保下安装支架",
        ),
        root_part_number="501009144AA",
        bump_root=True,
        extra_old="501009122AA",
        extra_new="501009122AB",
        process_old_part_number="501009122AA",
    ),
    Scenario(
        cover=CoverRow(
            NOT_ASSEMBLED,
            "501008897AA",
            "501008897AB",
            "后围总成",
            "一级",
            "504000341AA",
            None,
            "左D柱下结构增强胶块",
        ),
        root_part_number="501008897AA",
        bump_root=True,
        not_assembled_part_number="504000341AA",
    ),
    Scenario(
        cover=CoverRow(
            ADD_PART,
            "501007029AA",
            "501007029AB",
            "中通道左后纵梁总成",
            "一级",
            None,
            ADD_PART_NUMBER,
            ADD_PART_NAME,
        ),
        root_part_number="501007029AA",
        bump_root=True,
        added_part_number=ADD_PART_NUMBER,
        added_part_name=ADD_PART_NAME,
        added_internal_number=ADD_PART_INTERNAL,
        added_material="HC340/590DP",
        added_thickness_text="T=1.0",
        added_part_weight=Decimal("0.235"),
        process_old_part_number="501007031AA",
    ),
)


def _copy_sheet(source: Worksheet, dest_wb: Workbook, title: str) -> Worksheet:
    dest = dest_wb.create_sheet(title)
    for row in source.iter_rows():
        for cell in row:
            cloned = dest.cell(cell.row, cell.column, cell.value)
            if cell.has_style:
                cloned.font = copy(cell.font)
                cloned.border = copy(cell.border)
                cloned.fill = copy(cell.fill)
                cloned.number_format = cell.number_format
                cloned.protection = copy(cell.protection)
                cloned.alignment = copy(cell.alignment)
    for merged in source.merged_cells.ranges:
        dest.merge_cells(str(merged))
    for letter, dim in source.column_dimensions.items():
        dest.column_dimensions[letter].width = dim.width
    for index, dim in source.row_dimensions.items():
        dest.row_dimensions[index].height = dim.height
    dest.freeze_panes = source.freeze_panes
    dest.sheet_view.zoomScale = source.sheet_view.zoomScale or 100
    dest.print_title_rows = source.print_title_rows
    dest.print_area = source.print_area
    return dest


def _load_template(content: bytes):
    return load_workbook(BytesIO(content))


def _unmerge_in_rows(sheet: Worksheet, min_row: int, max_row: int) -> None:
    ranges = [
        str(item)
        for item in sheet.merged_cells.ranges
        if item.min_row >= min_row and item.max_row <= max_row
    ]
    for item in ranges:
        sheet.unmerge_cells(item)


def _clear_row_range(sheet: Worksheet, start_row: int, end_row: int, max_col: int) -> None:
    _unmerge_in_rows(sheet, start_row, end_row)
    for row in range(start_row, end_row + 1):
        for col in range(1, max_col + 1):
            sheet.cell(row, col).value = None


def _merge_vertical(sheet: Worksheet, col: int, start_row: int, end_row: int) -> None:
    if end_row > start_row:
        sheet.merge_cells(start_row=start_row, start_column=col, end_row=end_row, end_column=col)


def _as_decimal(value: object) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value))


def _as_int(value: object) -> int | None:
    if value is None or value == "":
        return None
    return int(value)


def _excel_number(value: Decimal | float | int | None) -> int | float | None:
    if value is None:
        return None
    number = float(value)
    if number.is_integer():
        return int(number)
    return number


def _hits_month(
    effective_year: int,
    effective_month: int,
    expiry_year: int | None,
    expiry_month: int | None,
) -> bool:
    as_of = AS_OF_YEAR * 12 + AS_OF_MONTH
    start = effective_year * 12 + effective_month
    if start > as_of:
        return False
    if expiry_year is None or expiry_month is None:
        return True
    return as_of <= expiry_year * 12 + expiry_month


def _rename_part(part_number: str, old: str, new: str) -> str:
    if part_number == old:
        return new
    if part_number.startswith(old) and part_number[len(old) :].startswith("-"):
        return new + part_number[len(old) :]
    return part_number


def _bumped_root(root_part_number: str) -> str:
    if not root_part_number.endswith("AA"):
        raise ValueError(f"树根不是 AA 后缀: {root_part_number}")
    return root_part_number[:-2] + "AB"


def _is_listed_change(part_number: str, scenario: Scenario) -> bool:
    if scenario.not_assembled_part_number == part_number:
        return True
    if scenario.extra_old:
        if part_number == scenario.extra_old:
            return True
        suffix = part_number[len(scenario.extra_old) :]
        if part_number.startswith(scenario.extra_old) and suffix.startswith("-"):
            return True
    return False


def _prune_unchanged(node: BomNode, scenario: Scenario) -> BomNode | None:
    kept_children = [
        pruned
        for child in node.children
        if (pruned := _prune_unchanged(child, scenario)) is not None
    ]
    keep = bool(kept_children) or _is_listed_change(node.part_number, scenario)
    if node.level == 0 and scenario.bump_root:
        keep = True
    if not keep:
        return None
    return BomNode(
        node_id=node.node_id,
        level=node.level,
        seq_no=node.seq_no,
        internal_number=node.internal_number,
        part_number=node.part_number,
        part_name=node.part_name,
        qty=node.qty,
        qty_text=node.qty_text,
        weld_info=node.weld_info,
        material=node.material,
        thickness_text=node.thickness_text,
        material_mill=node.material_mill,
        part_weight=node.part_weight,
        gauge_developed=node.gauge_developed,
        remark=node.remark,
        children=kept_children,
    )


def _apply_renames(part_number: str, scenario: Scenario) -> str:
    renamed = part_number
    if scenario.extra_old and scenario.extra_new:
        renamed = _rename_part(renamed, scenario.extra_old, scenario.extra_new)
    if scenario.bump_root:
        renamed = _rename_part(
            renamed, scenario.root_part_number, _bumped_root(scenario.root_part_number)
        )
    return renamed


def _added_part_node(scenario: Scenario) -> BomNode:
    if not scenario.added_part_number:
        raise ValueError("增加零件场景没有新零件号")
    return BomNode(
        node_id=0,
        level=1,
        seq_no=None,
        internal_number=scenario.added_internal_number,
        part_number=scenario.added_part_number,
        part_name=scenario.added_part_name,
        qty=Decimal("1"),
        qty_text="1",
        weld_info=None,
        material=scenario.added_material,
        thickness_text=scenario.added_thickness_text,
        material_mill=None,
        part_weight=scenario.added_part_weight,
        gauge_developed=None,
        remark=None,
        children=[],
    )


def _attach_added_part(tree: BomNode, scenario: Scenario) -> BomNode:
    if not scenario.added_part_number:
        return tree
    return BomNode(
        node_id=tree.node_id,
        level=tree.level,
        seq_no=tree.seq_no,
        internal_number=tree.internal_number,
        part_number=tree.part_number,
        part_name=tree.part_name,
        qty=tree.qty,
        qty_text=tree.qty_text,
        weld_info=tree.weld_info,
        material=tree.material,
        thickness_text=tree.thickness_text,
        material_mill=tree.material_mill,
        part_weight=tree.part_weight,
        gauge_developed=tree.gauge_developed,
        remark=tree.remark,
        children=[_added_part_node(scenario)],
    )


def _qty_cell(node: BomNode, *, is_not_assembled: bool) -> object:
    if is_not_assembled:
        return 0
    if node.qty_text:
        return node.qty_text
    return _excel_number(node.qty) if node.qty is not None else 1


def _connect() -> sqlite3.Connection:
    if not DATABASE_PATH.exists():
        raise FileNotFoundError(f"找不到现网库: {DATABASE_PATH}")
    conn = sqlite3.connect(f"file:{DATABASE_PATH}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def _load_children(cur: sqlite3.Cursor, parent_id: int) -> list[BomNode]:
    rows = list(
        cur.execute(
            "SELECT id, level, seq_no, internal_number, part_number, part_name, qty, "
            "qty_text, weld_info, material, thickness_text, material_mill, part_weight, "
            "gauge_developed, remark FROM bom_nodes "
            "WHERE parent_id = ? ORDER BY sort_no, id",
            (parent_id,),
        )
    )
    nodes: list[BomNode] = []
    for row in rows:
        node = BomNode(
            node_id=row["id"],
            level=row["level"],
            seq_no=_as_int(row["seq_no"]),
            internal_number=row["internal_number"],
            part_number=row["part_number"],
            part_name=row["part_name"],
            qty=_as_decimal(row["qty"]),
            qty_text=row["qty_text"],
            weld_info=row["weld_info"],
            material=row["material"],
            thickness_text=row["thickness_text"],
            material_mill=row["material_mill"],
            part_weight=_as_decimal(row["part_weight"]),
            gauge_developed=row["gauge_developed"],
            remark=row["remark"],
            children=_load_children(cur, row["id"]),
        )
        nodes.append(node)
    return nodes


def load_tree(conn: sqlite3.Connection, root_part_number: str) -> BomNode:
    cur = conn.cursor()
    row = cur.execute(
        "SELECT id, level, seq_no, internal_number, part_number, part_name, qty, "
        "qty_text, weld_info, material, thickness_text, material_mill, part_weight, "
        "gauge_developed, remark, effective_year, effective_month, expiry_year, "
        "expiry_month FROM bom_nodes "
        "WHERE location = ? AND vehicle_model = ? AND level = 0 AND part_number = ?",
        (LOCATION, VEHICLE_MODEL, root_part_number),
    ).fetchone()
    if row is None:
        raise ValueError(f"库里没有 {LOCATION} {VEHICLE_MODEL} 树根 {root_part_number}")
    if not _hits_month(
        row["effective_year"],
        row["effective_month"],
        row["expiry_year"],
        row["expiry_month"],
    ):
        raise ValueError(f"{root_part_number} 在 {PERIOD_LABEL} 未命中")
    return BomNode(
        node_id=row["id"],
        level=row["level"],
        seq_no=_as_int(row["seq_no"]),
        internal_number=row["internal_number"],
        part_number=row["part_number"],
        part_name=row["part_name"],
        qty=_as_decimal(row["qty"]),
        qty_text=row["qty_text"],
        weld_info=row["weld_info"],
        material=row["material"],
        thickness_text=row["thickness_text"],
        material_mill=row["material_mill"],
        part_weight=_as_decimal(row["part_weight"]),
        gauge_developed=row["gauge_developed"],
        remark=row["remark"],
        children=_load_children(cur, row["id"]),
    )


def load_process_part(conn: sqlite3.Connection, part_number: str) -> ProcessPart:
    cur = conn.cursor()
    row = cur.execute(
        "SELECT id, part_number, part_name, material, thickness, part_weight, "
        "sheet_length, sheet_width, parts_per_sheet, sheet_type, die_type "
        "FROM stamping_parts WHERE location = ? AND vehicle_model = ? AND part_number = ?",
        (LOCATION, VEHICLE_MODEL, part_number),
    ).fetchone()
    if row is None:
        raise ValueError(f"库里没有工艺行 {part_number}")
    ops = [
        ProcessOp(
            process_seq=item["process_seq"],
            process_code=item["process_code"],
            die_sets=_as_int(item["die_sets"]),
            die_number=item["die_number"],
            tonnage=_as_int(item["tonnage"]),
        )
        for item in cur.execute(
            "SELECT process_seq, process_code, die_sets, die_number, tonnage "
            "FROM stamping_processes WHERE part_id = ? ORDER BY process_order, id",
            (row["id"],),
        )
    ]
    return ProcessPart(
        part_number=row["part_number"],
        part_name=row["part_name"],
        material=row["material"],
        thickness=row["thickness"],
        part_weight=row["part_weight"],
        sheet_length=row["sheet_length"],
        sheet_width=row["sheet_width"],
        parts_per_sheet=_as_int(row["parts_per_sheet"]),
        sheet_type=row["sheet_type"],
        die_type=row["die_type"],
        operations=ops,
    )


def _write_cover(sheet: Worksheet) -> None:
    title = Font(name="微软雅黑", size=14, bold=True)
    label = Font(name="微软雅黑", size=10)
    head = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4472C4")
    sheet["A1"] = "工程变更（封皮）"
    sheet["A1"].font = title
    sheet.merge_cells("A1:J1")
    sheet["A2"] = (
        "一本 Excel、五种改法。后面三张表和日常导入模板相同。"
        f"基地 {LOCATION}，顾客 {CUSTOMER_NAME}，投产清单版本 D/00。"
        "先写变更类型和原来的交货总成，再写这一次动到的零件。"
        "总成名称、零件名称变更前后相同，各只一列。零件层级=总成时零件名称空。"
        "材料清单只写变化的行，没变的零件不要抄。"
        "完成停止：变更后总成空，材料清单不写这个总成。"
        f"{NOT_ASSEMBLED}：总成还交货，变更前总成/变更后总成 AA→AB，变更后零件空；该行用量写 0、备注含「{NOT_ASSEMBLED}」。"
        f"{ADD_PART}：加在原来的总成上，不是新开一棵总成。变更前总成/变更后总成 AA→AB，变更前零件空、变更后零件写新号，用量 > 0。"
    )
    sheet["A2"].font = Font(name="微软雅黑", size=9, color="666666")
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.merge_cells("A2:J2")
    sheet.row_dimensions[2].height = 108
    headers = (
        "项目",
        "生效年月",
        "变更类型",
        "变更前总成",
        "变更后总成",
        "总成名称",
        "零件层级",
        "变更前零件",
        "变更后零件",
        "零件名称",
    )
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(4, col, header)
        cell.font = head
        cell.fill = fill
        cell.border = _THIN
        cell.alignment = _CENTER
    for offset, scenario in enumerate(SCENARIOS):
        excel_row = 5 + offset
        cover = scenario.cover
        values = (
            VEHICLE_MODEL,
            PERIOD_LABEL,
            cover.change_type,
            cover.before_assembly,
            cover.after_assembly or "",
            cover.assembly_name,
            cover.part_level,
            cover.before_part or "",
            cover.after_part or "",
            cover.part_name or "",
        )
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(excel_row, col, value)
            cell.font = label
            cell.border = _THIN
            cell.alignment = _CENTER
            if cover.change_type == NOT_ASSEMBLED and col >= 3:
                cell.fill = _NOT_ASSEMBLED_FILL
            if cover.change_type == ADD_PART and col >= 3:
                cell.fill = _ADD_PART_FILL
        sheet.row_dimensions[excel_row].height = 28
    last_cover_row = 4 + len(SCENARIOS)
    sheet.merge_cells(f"A5:A{last_cover_row}")
    sheet.merge_cells(f"B5:B{last_cover_row}")
    widths = (12, 14, 12, 22, 22, 28, 10, 22, 22, 28)
    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + col)].width = width
    sheet.freeze_panes = "A5"


def _has_not_assembled_remark(remark: str | None) -> bool:
    return bool(remark) and NOT_ASSEMBLED in remark


def _style_data_row(
    sheet: Worksheet,
    row: int,
    *,
    is_not_assembled: bool = False,
    is_added_part: bool = False,
) -> None:
    fill = None
    if is_not_assembled:
        fill = _NOT_ASSEMBLED_FILL
    elif is_added_part:
        fill = _ADD_PART_FILL
    for col in range(1, 24):
        cell = sheet.cell(row, col)
        cell.font = _DATA_FONT
        cell.border = _THIN
        cell.alignment = _CENTER
        if fill is not None:
            cell.fill = fill
    sheet.row_dimensions[row].height = 22


def _write_leaf_attrs(sheet: Worksheet, row: int, node: BomNode, *, remark: str | None) -> None:
    sheet.cell(row, 18, node.material or None)
    sheet.cell(row, 19, node.thickness_text or None)
    sheet.cell(row, 20, node.material_mill or None)
    sheet.cell(row, 21, _excel_number(node.part_weight))
    sheet.cell(row, 22, node.gauge_developed or SLASH)
    sheet.cell(row, 23, remark)


def _write_l2_identity(sheet: Worksheet, row: int, node: BomNode | None) -> None:
    if node is None:
        sheet.cell(row, 13, SLASH)
        sheet.cell(row, 14, SLASH)
        sheet.cell(row, 15, SLASH)
        sheet.cell(row, 17, SLASH)
        return
    sheet.cell(row, 13, node.internal_number or SLASH)
    sheet.cell(row, 14, node.part_number)
    sheet.cell(row, 15, node.part_name)
    sheet.cell(row, 17, _qty_cell(node, is_not_assembled=False))


def _apply_node_renames(node: BomNode, scenario: Scenario) -> BomNode:
    remark = node.remark
    qty = node.qty
    qty_text = node.qty_text
    is_not_assembled = scenario.not_assembled_part_number == node.part_number
    if is_not_assembled:
        qty = Decimal("0")
        qty_text = "0"
        remark = NOT_ASSEMBLED
    return BomNode(
        node_id=node.node_id,
        level=node.level,
        seq_no=node.seq_no,
        internal_number=node.internal_number,
        part_number=_apply_renames(node.part_number, scenario),
        part_name=node.part_name,
        qty=qty,
        qty_text=qty_text,
        weld_info=node.weld_info,
        material=node.material,
        thickness_text=node.thickness_text,
        material_mill=node.material_mill,
        part_weight=node.part_weight,
        gauge_developed=node.gauge_developed,
        remark=remark,
        children=[_apply_node_renames(child, scenario) for child in node.children],
    )


def _write_bom_header(sheet: Worksheet, doc_no: str) -> None:
    sheet["A1"] = f"{VEHICLE_MODEL}项目产品材料清单"
    sheet["A2"] = (
        "编制部门： 技术中心                    "
        f"编制日期：2025-07-01  版本： {CHANGE_VERSION}  编号：{doc_no}"
    )
    sheet["A3"] = f"项目代号：  {VEHICLE_MODEL}"
    sheet["I3"] = "产品类别：  车身件"
    sheet["N3"] = f"顾客名称：  {CUSTOMER_NAME}"
    sheet["P3"] = "阶段标识： D"
    sheet["R3"] = "■批量生产清单"
    sheet["V3"] = "修订日期：2025-07-01"


def _append_bom_tree(sheet: Worksheet, root: BomNode, start_row: int) -> int:
    excel_row = start_row
    l1_spans: list[tuple[int, int]] = []
    for l1 in root.children:
        group_start = excel_row
        leaves = l1.children
        if not leaves:
            _style_data_row(
                sheet,
                excel_row,
                is_not_assembled=_has_not_assembled_remark(l1.remark),
                is_added_part=l1.part_number == ADD_PART_NUMBER,
            )
            sheet.cell(excel_row, 8, l1.internal_number)
            sheet.cell(excel_row, 9, l1.part_number)
            sheet.cell(excel_row, 10, l1.part_name)
            sheet.cell(excel_row, 11, _qty_cell(l1, is_not_assembled=l1.qty == Decimal("0")))
            sheet.cell(excel_row, 12, l1.weld_info or SLASH)
            _write_l2_identity(sheet, excel_row, None)
            _write_leaf_attrs(sheet, excel_row, l1, remark=l1.remark)
            excel_row += 1
        else:
            for index, l2 in enumerate(leaves):
                _style_data_row(
                    sheet,
                    excel_row,
                    is_not_assembled=_has_not_assembled_remark(l2.remark),
                    is_added_part=l2.part_number == ADD_PART_NUMBER,
                )
                if index == 0:
                    sheet.cell(excel_row, 8, l1.internal_number)
                    sheet.cell(excel_row, 9, l1.part_number)
                    sheet.cell(excel_row, 10, l1.part_name)
                    sheet.cell(excel_row, 11, _qty_cell(l1, is_not_assembled=False))
                    sheet.cell(excel_row, 12, l1.weld_info or SLASH)
                _write_l2_identity(sheet, excel_row, l2)
                _write_leaf_attrs(sheet, excel_row, l2, remark=l2.remark)
                excel_row += 1
        l1_spans.append((group_start, excel_row - 1))

    last_row = excel_row - 1
    if last_row < start_row:
        return start_row
    for col in range(1, 8):
        _merge_vertical(sheet, col, start_row, last_row)
    sheet.cell(start_row, 1, root.seq_no)
    sheet.cell(start_row, 2, root.internal_number)
    sheet.cell(start_row, 3, root.part_number)
    sheet.cell(start_row, 4, root.part_name)
    sheet.cell(start_row, 6, _qty_cell(root, is_not_assembled=False))
    sheet.cell(start_row, 7, root.weld_info or SLASH)
    for start, end in l1_spans:
        for col in range(8, 13):
            _merge_vertical(sheet, col, start, end)
    return excel_row


def _strip_process_sample(sheet: Worksheet) -> None:
    sheet["D2"] = f"{VEHICLE_MODEL}项目冲压工艺确认表-工程变更"
    sheet["D3"] = STAMPING_SUPPLIER
    sheet["S3"] = CHANGE_VERSION
    max_row = sheet.max_row or 6
    _clear_row_range(sheet, 7, max(max_row, 7), 20)


def _fill_process_part(
    sheet: Worksheet,
    part: ProcessPart,
    new_part_number: str,
    *,
    start_row: int,
    seq_no: int,
) -> int:
    operations = part.operations or [ProcessOp(None, None, None, None, None)]
    end_row = start_row + len(operations) - 1
    for row in range(start_row, end_row + 1):
        for col in range(1, 20):
            cell = sheet.cell(row, col)
            cell.border = _THIN
            cell.font = _DATA_FONT
            cell.alignment = _CENTER
        sheet.row_dimensions[row].height = 22
    common_cols = (1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 19)
    for col in common_cols:
        _merge_vertical(sheet, col, start_row, end_row)
    sheet.cell(start_row, 1, seq_no)
    sheet.cell(start_row, 2, new_part_number)
    sheet.cell(start_row, 3, part.part_name)
    sheet.cell(start_row, 5, part.material)
    sheet.cell(start_row, 6, _excel_number(_as_decimal(part.thickness)))
    sheet.cell(start_row, 7, _excel_number(_as_decimal(part.part_weight)))
    sheet.cell(start_row, 8, _excel_number(_as_decimal(part.sheet_length)))
    sheet.cell(start_row, 9, _excel_number(_as_decimal(part.sheet_width)))
    sheet.cell(start_row, 10, part.parts_per_sheet)
    sheet.cell(start_row, 13, part.sheet_type)
    sheet.cell(start_row, 19, part.die_type)
    for offset, operation in enumerate(operations):
        row = start_row + offset
        sheet.cell(row, 14, operation.process_seq)
        sheet.cell(row, 15, operation.process_code)
        sheet.cell(row, 16, operation.die_sets)
        sheet.cell(row, 17, operation.die_number)
        if operation.tonnage is not None:
            sheet.cell(row, 18, operation.tonnage)
    return end_row + 1


def _strip_die_below_header(sheet: Worksheet) -> None:
    max_row = max(sheet.max_row or 6, 40)
    _clear_row_range(sheet, 6, max_row, DIE_COL_COUNT)


def _style_die_data_row(sheet: Worksheet, row: int) -> None:
    for col in range(1, DIE_COL_COUNT + 1):
        cell = sheet.cell(row, col)
        cell.border = _DIE_THIN
        cell.font = _DIE_DATA_FONT
        cell.alignment = _CENTER
    sheet.row_dimensions[row].height = 18.75


def _fill_die_part(
    sheet: Worksheet,
    part: ProcessPart,
    new_part_number: str,
    *,
    start_row: int,
    seq_no: int,
) -> int:
    operations = part.operations or [ProcessOp(None, None, None, None, None)]
    end_row = start_row + len(operations) - 1
    for row in range(start_row, end_row + 1):
        _style_die_data_row(sheet, row)
    for col in _DIE_PART_COLS:
        _merge_vertical(sheet, col, start_row, end_row)
    sheet.cell(start_row, 1, seq_no)
    sheet.cell(start_row, 2, new_part_number)
    sheet.cell(start_row, 3, part.part_name)
    sheet.cell(start_row, 5, part.material)
    sheet.cell(start_row, 6, _excel_number(_as_decimal(part.thickness)))
    sheet.cell(start_row, 7, _excel_number(_as_decimal(part.part_weight)))
    sheet.cell(start_row, 8, _excel_number(_as_decimal(part.sheet_length)))
    sheet.cell(start_row, 9, _excel_number(_as_decimal(part.sheet_width)))
    sheet.cell(start_row, 10, part.sheet_type)
    sheet.cell(start_row, 11, part.parts_per_sheet)
    sheet.cell(start_row, 25, part.die_type)
    for offset, operation in enumerate(operations):
        row = start_row + offset
        sheet.cell(row, 14, operation.process_seq)
        sheet.cell(row, 15, operation.process_code)
        sheet.cell(row, 16, operation.die_sets)
        sheet.cell(row, 17, operation.die_number)
        if operation.tonnage is not None:
            sheet.cell(row, 24, operation.tonnage)
    return end_row + 1


def _write_die_alloc_notes(sheet: Worksheet, data_end_row: int) -> None:
    start_row = data_end_row + 2
    for offset, line in enumerate(_DIE_ALLOC_NOTES):
        excel_row = start_row + offset
        sheet.merge_cells(
            start_row=excel_row,
            start_column=1,
            end_row=excel_row,
            end_column=DIE_COL_COUNT,
        )
        cell = sheet.cell(excel_row, 1, line)
        cell.font = _DIE_NOTE_TITLE_FONT if offset == 0 else _DIE_NOTE_FONT
        cell.alignment = _DIE_NOTE_ALIGN
        sheet.row_dimensions[excel_row].height = 22


def _write_die_layout(
    sheet: Worksheet,
    process_parts: list[tuple[ProcessPart, str]],
) -> None:
    sheet.cell(1, 1, f"{VEHICLE_MODEL}项目冲压模具布点工艺清单")
    sheet.cell(1, 30, CHANGE_VERSION)
    _strip_die_below_header(sheet)
    excel_row = 6
    for seq_no, (part, new_part_number) in enumerate(process_parts, start=1):
        excel_row = _fill_die_part(
            sheet,
            part,
            new_part_number,
            start_row=excel_row,
            seq_no=seq_no,
        )
    last_data_row = excel_row - 1
    _write_die_alloc_notes(sheet, last_data_row if last_data_row >= 6 else 6)


def _save(workbook: Workbook, path: Path) -> Path:
    temp_out = path.with_suffix(".building.xlsx")
    workbook.save(temp_out)
    try:
        temp_out.replace(path)
        return path
    except OSError:
        print(f"目标文件被占用，已写出: {temp_out}")
        return temp_out


def build_workbook(
    trees: list[BomNode],
    process_parts: list[tuple[ProcessPart, str]],
) -> Workbook:
    bom_wb = _load_template(build_bom_material_template()[0])
    process_wb = _load_template(build_stamping_process_template()[0])
    die_wb = _load_template(build_die_layout_template()[0])

    combined = Workbook()
    cover = combined.active
    if cover is None:
        cover = combined.create_sheet("封皮")
    cover.title = "封皮"
    _write_cover(cover)

    bom_sheet = _copy_sheet(bom_wb["车身件材料清单"], combined, "车身件材料清单")
    _write_bom_header(bom_sheet, "EC-T28-202507")
    _clear_row_range(bom_sheet, 5, 200, 23)
    excel_row = 5
    for tree in trees:
        excel_row = _append_bom_tree(bom_sheet, tree, excel_row)
    _copy_sheet(bom_wb["填写说明"], combined, "材料清单填写说明")

    process_sheet = _copy_sheet(process_wb["冲压工艺确认表"], combined, "冲压工艺确认表")
    _strip_process_sample(process_sheet)
    process_row = 7
    for seq_no, (part, new_part_number) in enumerate(process_parts, start=1):
        process_row = _fill_process_part(
            process_sheet,
            part,
            new_part_number,
            start_row=process_row,
            seq_no=seq_no,
        )
    _copy_sheet(process_wb["填写说明"], combined, "工艺确认填写说明")

    die_sheet = _copy_sheet(die_wb["零件汇总"], combined, "零件汇总")
    _write_die_layout(die_sheet, process_parts)
    return combined


def _verify(path: Path) -> str:
    workbook = load_workbook(path, data_only=True)
    cover = workbook["封皮"]
    cover_rows = [
        (
            cover.cell(row, 3).value,
            cover.cell(row, 4).value,
            cover.cell(row, 5).value or "（空）",
            cover.cell(row, 6).value,
            cover.cell(row, 7).value,
            cover.cell(row, 8).value or "（空）",
            cover.cell(row, 9).value or "（空）",
            cover.cell(row, 10).value or "（空）",
        )
        for row in range(5, 4 + len(SCENARIOS) + 1)
    ]
    bom = workbook["车身件材料清单"]
    meta = parse_bom_meta(bom)
    nodes, _, _ = parse_bom_nodes(bom)
    l0 = [node.part_number for node in nodes if node.level == 0]
    listed = [(node.level, node.part_number) for node in nodes]
    process_sheet = workbook["冲压工艺确认表"]
    process_numbers = [
        process_sheet.cell(row, 2).value for row in range(7, 16) if process_sheet.cell(row, 2).value
    ]
    die = workbook["零件汇总"]
    die_listed = [
        (die.cell(row, 2).value, die.cell(row, 17).value or "（空）")
        for row in range(6, 12)
        if die.cell(row, 2).value
    ]
    glue = next(
        (node for node in nodes if node.part_number == "504000341AA"),
        None,
    )
    add_cover = cover_rows[-1]
    if add_cover[1] != "501007029AA" or add_cover[2] != "501007029AB":
        raise ValueError(f"增加零件总成两列必须是原来的总成 AA→AB，得到 {add_cover}")
    if add_cover[5] != "（空）" or add_cover[6] != ADD_PART_NUMBER:
        raise ValueError(f"增加零件零件两列必须是空→{ADD_PART_NUMBER}，得到 {add_cover}")
    if add_cover[3] != "中通道左后纵梁总成" or add_cover[7] != ADD_PART_NAME:
        raise ValueError(f"增加零件名称列必须写出总成名和零件名，得到 {add_cover}")
    lines = [
        path.name,
        f"  封皮 {cover_rows}",
        f"  材料清单版本={meta.bom_version} L0={l0} 节点={listed}",
        f"  工艺新号={process_numbers}",
        f"  布点新号={die_listed}",
    ]
    if glue is not None:
        glue_fill = None
        for row in range(5, 40):
            if bom.cell(row, 9).value == "504000341AA":
                glue_fill = bom.cell(row, 1).fill.fgColor.rgb
                break
        lines.append(
            f"  胶块 {glue.part_number} 用量={glue.qty} 备注={glue.remark} 底色={glue_fill}"
        )
    workbook.close()
    return "\n".join(lines)


def main() -> None:
    conn = _connect()
    try:
        trees: list[BomNode] = []
        process_parts: list[tuple[ProcessPart, str]] = []
        for scenario in SCENARIOS:
            if scenario.cover.change_type == "完成停止":
                continue
            tree = load_tree(conn, scenario.root_part_number)
            pruned = _prune_unchanged(tree, scenario)
            if pruned is None:
                raise ValueError(f"{scenario.root_part_number} 剪完后没有可写的行")
            pruned = _attach_added_part(pruned, scenario)
            trees.append(_apply_node_renames(pruned, scenario))
            if scenario.process_old_part_number and scenario.cover.after_part:
                process_part = load_process_part(conn, scenario.process_old_part_number)
                if scenario.added_part_name:
                    process_part = replace(process_part, part_name=scenario.added_part_name)
                process_parts.append(
                    (
                        process_part,
                        scenario.cover.after_part,
                    )
                )
        trees.sort(key=lambda node: node.seq_no or 0)
        workbook = build_workbook(trees, process_parts)
        out = _save(workbook, OUT)
        print(_verify(out))
    finally:
        conn.close()


if __name__ == "__main__":
    main()
