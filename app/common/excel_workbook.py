"""Excel Workbook 无业务辅助。"""

from __future__ import annotations

from decimal import Decimal
from typing import cast

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

# openpyxl 新建 Workbook 时默认首表名；仅此类「尚未占用」的空表可改名复用
_DEFAULT_SHEET_TITLES = frozenset({"Sheet", "工作表1"})

# 写入 Worksheet.cell 之前的业务值。None 由 excel_write_value 收成空串。
ExcelCellValue = str | int | float | Decimal | bool | None


def excel_write_value(value: object) -> str | int | float | Decimal | bool:
    """把业务导出值收成 openpyxl Worksheet.cell 可写的类型。"""
    if value is None:
        return ""
    if isinstance(value, bool | int | float | Decimal | str):
        return value
    return str(value)


def resolve_worksheet(wb: Workbook, sheet_title: str) -> Worksheet:
    """空默认表则复用改名；否则 create_sheet（重名加 _2/_3…）。"""
    default_sheet = wb.worksheets[0] if wb.worksheets else None
    if (
        default_sheet is not None
        and len(wb.worksheets) == 1
        and default_sheet.title in _DEFAULT_SHEET_TITLES
        and default_sheet.max_row == 1
        and default_sheet.max_column == 1
        and default_sheet["A1"].value is None
    ):
        default_sheet.title = sheet_title
        return default_sheet
    title = sheet_title
    existing = {sheet.title for sheet in wb.worksheets}
    if title in existing:
        suffix = 2
        while f"{sheet_title}_{suffix}" in existing:
            suffix += 1
        title = f"{sheet_title}_{suffix}"
    return cast(Worksheet, wb.create_sheet(title))
