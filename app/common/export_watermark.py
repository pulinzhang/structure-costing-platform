"""导出 Excel 内控水印：屏上浅色铺底 + 打印页眉页脚。"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Protocol, cast

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image, ImageDraw, ImageFont

from app.common.audit import display_name
from app.common.timezone import now
from app.schemas.vo.user_vo import UserVO

_WATERMARK_PREFIX = "导出人："
_TILE_WIDTH_PX = 720
_TILE_HEIGHT_PX = 400
_TILE_ROW_STEP = 22
_TILE_COL_STEP = 10
_MAX_TILE_ROWS = 8
_MAX_TILE_COLS = 4

_FONT_CANDIDATES = (
    Path(r"C:\Windows\Fonts\msyh.ttc"),
    Path(r"C:\Windows\Fonts\simhei.ttf"),
    Path(r"C:\Windows\Fonts\simsun.ttc"),
    Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
    Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
    Path("/System/Library/Fonts/PingFang.ttc"),
    Path("/System/Library/Fonts/STHeiti Light.ttc"),
)


class _HeaderFooterPart(Protocol):
    text: str


class _HeaderFooterItem(Protocol):
    left: _HeaderFooterPart
    center: _HeaderFooterPart
    right: _HeaderFooterPart


def format_export_watermark_text(
    exporter: UserVO,
    *,
    exported_at: datetime | None = None,
) -> str:
    """导出人：部门-姓名，时间：YYYY-MM-DD HH:MM。"""
    person = display_name(exporter)
    department = str(getattr(exporter, "department", None) or "").strip()
    who = f"{department}-{person}" if department else person
    stamp = exported_at if exported_at is not None else now()
    when = stamp.strftime("%Y-%m-%d %H:%M")
    return f"{_WATERMARK_PREFIX}{who}，时间：{when}"


def workbook_to_xlsx_bytes(
    workbook: Workbook,
    exporter: UserVO,
    *,
    exported_at: datetime | None = None,
) -> bytes:
    """盖水印后写入 xlsx 字节。业务导出落盘统一走这里。"""
    apply_workbook_watermark(
        workbook, format_export_watermark_text(exporter, exported_at=exported_at)
    )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def apply_workbook_watermark(workbook: Workbook, text: str) -> None:
    """对每个工作表追加打印水印，并铺浅色斜排图（截图可见）。"""
    if not text.strip():
        return
    png_bytes = _build_watermark_png(text)
    for sheet in workbook.worksheets:
        _append_print_watermark(sheet, text)
        _tile_watermark_images(sheet, png_bytes)


def _append_print_watermark(worksheet: Worksheet, text: str) -> None:
    header = cast(_HeaderFooterItem | None, worksheet.oddHeader)
    footer = cast(_HeaderFooterItem | None, worksheet.oddFooter)
    _append_part(header.center if header is not None else None, text)
    if footer is None:
        return
    center = (footer.center.text or "").strip()
    if not center:
        footer.center.text = text
        return
    if text in center or text in (footer.left.text or "") or text in (footer.right.text or ""):
        return
    _append_part(footer.left, text)


def _append_part(part: _HeaderFooterPart | None, text: str) -> None:
    if part is None:
        return
    existing = (part.text or "").strip()
    if not existing:
        part.text = text
        return
    if text in existing:
        return
    part.text = f"{existing}  {text}"


def _tile_watermark_images(worksheet: Worksheet, png_bytes: bytes) -> None:
    used_rows = max(int(worksheet.max_row or 1), 1)
    used_cols = max(int(worksheet.max_column or 1), 1)
    row_tiles = min(_MAX_TILE_ROWS, max(1, (used_rows + _TILE_ROW_STEP - 1) // _TILE_ROW_STEP))
    col_tiles = min(_MAX_TILE_COLS, max(1, (used_cols + _TILE_COL_STEP - 1) // _TILE_COL_STEP))
    for row_index in range(row_tiles):
        for col_index in range(col_tiles):
            row = 1 + row_index * _TILE_ROW_STEP
            col = 1 + col_index * _TILE_COL_STEP
            anchor = f"{get_column_letter(col)}{row}"
            image = ExcelImage(_png_buffer(png_bytes))
            image.width = _TILE_WIDTH_PX
            image.height = _TILE_HEIGHT_PX
            worksheet.add_image(image, anchor)


def _png_buffer(png_bytes: bytes) -> BytesIO:
    buffer = BytesIO(png_bytes)
    buffer.name = "export_watermark.png"
    return buffer


def _build_watermark_png(text: str) -> bytes:
    image = Image.new("RGBA", (_TILE_WIDTH_PX, _TILE_HEIGHT_PX), (255, 255, 255, 0))
    layer = Image.new("RGBA", (_TILE_WIDTH_PX, _TILE_HEIGHT_PX), (0, 0, 0, 0))
    draw = ImageDraw.Draw(layer)
    font = _load_font(22)
    draw.text((36, 160), text, font=font, fill=(150, 150, 150, 72))
    rotated = layer.rotate(28, resample=Image.Resampling.BICUBIC, expand=False)
    composed = Image.alpha_composite(image, rotated)
    output = BytesIO()
    composed.save(output, format="PNG")
    return output.getvalue()


def _load_font(size: int) -> ImageFont.ImageFont | ImageFont.FreeTypeFont:
    for path in _FONT_CANDIDATES:
        if not path.is_file():
            continue
        try:
            return ImageFont.truetype(str(path), size=size)
        except OSError:
            continue
    return ImageFont.load_default()
