"""项目：导入模板与一览 Excel。"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.common.excel_workbook import resolve_worksheet
from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.timezone import now
from app.exceptions.base import BusinessException
from app.parsers.masterdata.vehicle_model_parser import (
    HEADER_CURRENT_SUPPLIED_UNITS,
    HEADER_EFFECTIVE_MONTH,
    HEADER_EFFECTIVE_YEAR,
    HEADER_EXPIRY_MONTH,
    HEADER_EXPIRY_YEAR,
    HEADER_LOCATION,
    HEADER_VEHICLE_MODEL,
)
from app.repositories.masterdata.vehicle_model import VehicleModelRepository
from app.schemas.vo.user_vo import UserVO
from app.services.masterdata.location import LocationService

IMPORT_HEADERS = (
    HEADER_LOCATION,
    HEADER_VEHICLE_MODEL,
    HEADER_EFFECTIVE_YEAR,
    HEADER_EFFECTIVE_MONTH,
    HEADER_EXPIRY_YEAR,
    HEADER_EXPIRY_MONTH,
    HEADER_CURRENT_SUPPLIED_UNITS,
)


def new_import_sheet() -> tuple[Workbook, Worksheet, Font, Border]:
    workbook = Workbook()
    sheet = resolve_worksheet(workbook, "项目")
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    data_font = Font(name="微软雅黑", size=9)
    thin_border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    for col, title in enumerate(IMPORT_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 16
    for col_letter in ("C", "D", "E", "F", "G"):
        sheet.column_dimensions[col_letter].width = 14
    sheet.column_dimensions["G"].width = 16
    return workbook, sheet, data_font, thin_border


def write_import_row(
    sheet: Worksheet,
    row_idx: int,
    values: tuple[str | int | None, ...],
    data_font: Font,
    thin_border: Border,
) -> None:
    for col, value in enumerate(values, start=1):
        cell = sheet.cell(row=row_idx, column=col, value=value)
        cell.font = data_font
        cell.border = thin_border


def build_template() -> tuple[bytes, str]:
    workbook, sheet, data_font, thin_border = new_import_sheet()
    write_import_row(
        sheet,
        2,
        ("青岛", "T13J", 2026, 1, None, None, 0),
        data_font,
        thin_border,
    )
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue(), "项目导入模板.xlsx"


def export_excel(
    repository: VehicleModelRepository,
    location_service: LocationService,
    db: Session,
    current_user: UserVO,
    location: str | None = None,
) -> tuple[bytes, str]:
    name = (location or "").strip()
    if name:
        canonical = location_service.resolve_canonical(db, name)
        if canonical is None:
            raise BusinessException("没有这座基地")
        records = repository.list_by_location(db, canonical)
    else:
        records = repository.list_all(db)

    workbook, sheet, data_font, thin_border = new_import_sheet()
    for row_idx, record in enumerate(records, start=2):
        write_import_row(
            sheet,
            row_idx,
            (
                record.location,
                record.vehicle_model,
                record.effective_year,
                record.effective_month,
                record.expiry_year,
                record.expiry_month,
                record.current_supplied_units,
            ),
            data_font,
            thin_border,
        )
    stamp = now().strftime("%Y%m%d_%H%M%S")
    return workbook_to_xlsx_bytes(workbook, current_user), f"项目一览_{stamp}.xlsx"
