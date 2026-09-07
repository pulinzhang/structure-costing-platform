"""项目主数据 Service。"""

from __future__ import annotations

import logging
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.common.audit import audit_create, audit_update
from app.common.excel_workbook import resolve_worksheet
from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.timezone import now
from app.common.vehicle_model import normalize_vehicle_model
from app.exceptions.base import BusinessException, NotFoundException
from app.models.masterdata.vehicle_model import VehicleModel
from app.parsers.masterdata.vehicle_model_parser import (
    HEADER_EFFECTIVE_MONTH,
    HEADER_EFFECTIVE_YEAR,
    HEADER_EXPIRY_MONTH,
    HEADER_EXPIRY_YEAR,
    HEADER_LOCATION,
    HEADER_VEHICLE_MODEL,
    parse_vehicle_model_excel,
)
from app.repositories.masterdata.vehicle_model import VehicleModelRepository
from app.schemas.dto.vehicle_model_dto import (
    VehicleModelCreateDTO,
    VehicleModelInsert,
    VehicleModelPatch,
    VehicleModelUpdateDTO,
)
from app.schemas.vo.user_vo import UserVO
from app.schemas.vo.vehicle_model_vo import (
    PERIOD_EXPIRED,
    PERIOD_INTRODUCIBLE,
    PERIOD_PENDING,
    LocationCardVO,
    VehicleModelListVO,
    VehicleModelUploadResultVO,
    VehicleModelVO,
)
from app.services.masterdata.location import LocationService

logger = logging.getLogger("service.vehicle_model")

_PLACEHOLDER_CODES = frozenset({"UNKNOWN"})
_PAGE_HINT = "请到项目一览表手工增加或上传（#/vehicle-model-list）"
_FILE_MISMATCH = "文件项目为 {parsed}，与填写的项目 {selected} 不一致"
_IMPORT_HEADERS = (
    HEADER_LOCATION,
    HEADER_VEHICLE_MODEL,
    HEADER_EFFECTIVE_YEAR,
    HEADER_EFFECTIVE_MONTH,
    HEADER_EXPIRY_YEAR,
    HEADER_EXPIRY_MONTH,
)


def _period(year: int, month: int) -> tuple[int, int]:
    return (year, month)


def _new_import_sheet() -> tuple[Workbook, Worksheet, Font, Border]:
    workbook = Workbook()
    sheet = resolve_worksheet(workbook, "项目")
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    data_font = Font(name="微软雅黑", size=9)
    thin_border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    for col, title in enumerate(_IMPORT_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 16
    for col_letter in ("C", "D", "E", "F"):
        sheet.column_dimensions[col_letter].width = 12
    return workbook, sheet, data_font, thin_border


def _write_import_row(
    sheet: Worksheet,
    row_idx: int,
    values: tuple[str | int | None, ...],
    data_font: Font,
    thin_border: Border,
) -> None:
    for col, value in enumerate(values, start=1):
        cell = sheet.cell(row=row_idx, column=col, value=value)
        cell.font = data_font
        cell.border = thin_border


class VehicleModelService:
    """项目 CRUD、上传与解析契约。"""

    def __init__(self) -> None:
        self.repository = VehicleModelRepository()
        self.location_service = LocationService()

    def normalize(self, value: str | None) -> str | None:
        code = normalize_vehicle_model(value)
        if not code or code in _PLACEHOLDER_CODES:
            return None
        return code

    def require_registered(
        self,
        db: Session,
        location: str,
        vehicle_model: str,
    ) -> str:
        code = self.normalize(vehicle_model)
        if not code:
            raise BusinessException(f"没有该项目，{_PAGE_HINT}")
        record = self.repository.get(db, location, code)
        if record is None:
            raise BusinessException(f"没有该项目，{_PAGE_HINT}")
        return record.vehicle_model

    def require_introducible(
        self,
        db: Session,
        location: str,
        vehicle_model: str,
        year: int,
        month: int,
    ) -> tuple[str, str]:
        canonical = self.location_service.require_enabled(db, location)
        code = self.normalize(vehicle_model)
        if not code:
            raise BusinessException(f"没有该项目，{_PAGE_HINT}")
        record = self.repository.get(db, canonical, code)
        if record is None:
            raise BusinessException(f"没有该项目，{_PAGE_HINT}")
        as_of = _period(year, month)
        if as_of < _period(record.effective_year, record.effective_month):
            raise BusinessException(f"该项目尚未生效，{_PAGE_HINT}")
        if not self._is_not_expired(record, year, month):
            raise BusinessException(f"该项目已失效，{_PAGE_HINT}")
        return (canonical, record.vehicle_model)

    def require_overlaps_year(
        self,
        db: Session,
        location: str,
        vehicle_model: str,
        year: int,
    ) -> tuple[str, str]:
        canonical = self.location_service.require_enabled(db, location)
        code = self.normalize(vehicle_model)
        if not code:
            raise BusinessException(f"没有该项目，{_PAGE_HINT}")
        record = self.repository.get(db, canonical, code)
        if record is None:
            raise BusinessException(f"没有该项目，{_PAGE_HINT}")
        if not self._overlaps_year(record, year):
            raise BusinessException(f"该项目与 {year} 年无交集，{_PAGE_HINT}")
        return (canonical, record.vehicle_model)

    def assert_matches_selected(self, selected: str | None, parsed: str | None) -> None:
        selected_code = self.normalize(selected)
        parsed_code = self.normalize(parsed)
        if not selected_code or not parsed_code or selected_code != parsed_code:
            raise BusinessException(
                _FILE_MISMATCH.format(
                    parsed=parsed_code or (parsed or "").strip() or "空",
                    selected=selected_code or (selected or "").strip() or "空",
                )
            )

    def list_introducible(
        self,
        db: Session,
        location: str,
        year: int,
        month: int,
    ) -> list[str]:
        canonical = self.location_service.resolve_canonical(db, location)
        if canonical is None:
            return []
        try:
            self.location_service.require_enabled(db, canonical)
        except BusinessException:
            return []
        return [
            row.vehicle_model
            for row in self.repository.list_by_location(db, canonical)
            if self._is_in_period(row, year, month)
        ]

    def list_registered(self, db: Session, location: str) -> list[str]:
        canonical = self.location_service.resolve_canonical(db, location)
        if canonical is None:
            return []
        return [row.vehicle_model for row in self.repository.list_by_location(db, canonical)]

    def list_filter_codes(self, db: Session, location: str | None = None) -> list[str]:
        """列表筛：有基地则该基地已登记；无则各基地代号并集。"""
        raw = (location or "").strip()
        if raw:
            return self.list_registered(db, raw)
        return self.repository.list_distinct_codes(db)

    def list_overlaps_year(self, db: Session, location: str, year: int) -> list[str]:
        canonical = self.location_service.resolve_canonical(db, location)
        if canonical is None:
            return []
        try:
            self.location_service.require_enabled(db, canonical)
        except BusinessException:
            return []
        return [
            row.vehicle_model
            for row in self.repository.list_by_location(db, canonical)
            if self._overlaps_year(row, year)
        ]

    def list_page(self, db: Session) -> VehicleModelListVO:
        clock = now()
        as_of_year = clock.year
        as_of_month = clock.month
        location_rows = self.location_service.list_all(db)
        items_by_location: dict[str, list[VehicleModel]] = {}
        for row in self.repository.list_all(db):
            items_by_location.setdefault(row.location, []).append(row)

        cards: list[LocationCardVO] = []
        items: list[VehicleModelVO] = []
        for location_vo in location_rows:
            rows = items_by_location.get(location_vo.location, [])
            pending = 0
            introducible = 0
            expired = 0
            for row in rows:
                status = self._period_status(row, as_of_year, as_of_month)
                items.append(self._to_vo(row, status))
                if status == PERIOD_PENDING:
                    pending += 1
                elif status == PERIOD_INTRODUCIBLE:
                    introducible += 1
                else:
                    expired += 1
            cards.append(
                LocationCardVO(
                    location=location_vo.location,
                    is_enabled=location_vo.is_enabled,
                    project_count=len(rows),
                    introducible_count=introducible if location_vo.is_enabled else 0,
                    pending_count=pending,
                    expired_count=expired,
                )
            )
        return VehicleModelListVO(
            as_of_year=as_of_year,
            as_of_month=as_of_month,
            locations=cards,
            items=items,
        )

    def create(
        self,
        db: Session,
        dto: VehicleModelCreateDTO,
        current_user: UserVO,
    ) -> VehicleModelVO:
        canonical = self.location_service.require_enabled(db, dto.location)
        code = self.normalize(dto.vehicle_model)
        if not code:
            raise BusinessException("项目代号无效")
        self._validate_period(
            dto.effective_year,
            dto.effective_month,
            dto.expiry_year,
            dto.expiry_month,
        )
        if self.repository.get(db, canonical, code) is not None:
            raise BusinessException("该基地已有此项目")
        audit = audit_create(current_user)
        try:
            record = self.repository.create(
                db,
                VehicleModelInsert(
                    location=canonical,
                    vehicle_model=code,
                    effective_year=dto.effective_year,
                    effective_month=dto.effective_month,
                    expiry_year=dto.expiry_year,
                    expiry_month=dto.expiry_month,
                    created_by=audit["created_by"],
                    updated_by=audit["updated_by"],
                ),
            )
            db.commit()
            db.refresh(record)
        except IntegrityError:
            db.rollback()
            raise BusinessException("该基地已有此项目") from None
        except BusinessException:
            db.rollback()
            raise
        except Exception:
            db.rollback()
            logger.exception("创建项目失败")
            raise
        clock = now()
        return self._to_vo(record, self._period_status(record, clock.year, clock.month))

    def update(
        self,
        db: Session,
        location: str,
        vehicle_model: str,
        dto: VehicleModelUpdateDTO,
        current_user: UserVO,
    ) -> VehicleModelVO:
        record = self._require_row(db, location, vehicle_model)
        self._validate_period(
            dto.effective_year,
            dto.effective_month,
            dto.expiry_year,
            dto.expiry_month,
        )
        audit = audit_update(current_user)
        patch: VehicleModelPatch = {
            "effective_year": dto.effective_year,
            "effective_month": dto.effective_month,
            "expiry_year": dto.expiry_year,
            "expiry_month": dto.expiry_month,
            "updated_by": audit["updated_by"],
        }
        try:
            self.repository.apply_patch(db, record, patch)
            db.commit()
        except IntegrityError:
            db.rollback()
            raise BusinessException("仍被引用，不能修改") from None
        except BusinessException:
            db.rollback()
            raise
        except Exception:
            db.rollback()
            logger.exception("更新项目失败")
            raise
        refreshed = self.repository.get(db, record.location, record.vehicle_model)
        if refreshed is None:
            raise NotFoundException("项目不存在")
        clock = now()
        return self._to_vo(refreshed, self._period_status(refreshed, clock.year, clock.month))

    def delete(self, db: Session, location: str, vehicle_model: str) -> None:
        record = self._require_row(db, location, vehicle_model)
        try:
            self.repository.delete(db, record)
            db.commit()
        except IntegrityError:
            db.rollback()
            raise BusinessException("仍被引用，不能删除") from None
        except BusinessException:
            db.rollback()
            raise
        except Exception:
            db.rollback()
            logger.exception("删除项目失败")
            raise

    def process_upload(
        self,
        db: Session,
        file_content: bytes,
        current_user: UserVO,
    ) -> VehicleModelUploadResultVO:
        parsed = parse_vehicle_model_excel(file_content)
        if parsed.errors:
            raise BusinessException("；".join(parsed.errors))
        if not parsed.rows:
            raise BusinessException("没有可导入的项目行")

        pending: list[VehicleModelInsert] = []
        seen: dict[tuple[str, str], tuple[int, int, int | None, int | None]] = {}
        skipped_count = 0
        errors: list[str] = []
        audit = audit_create(current_user)

        for row in parsed.rows:
            try:
                canonical = self.location_service.require_enabled(db, row.location)
            except BusinessException as exc:
                errors.append(f"第 {row.row_number} 行：{exc.message}")
                continue
            code = self.normalize(row.vehicle_model)
            if not code:
                errors.append(f"第 {row.row_number} 行：项目代号无效")
                continue
            try:
                self._validate_period(
                    row.effective_year,
                    row.effective_month,
                    row.expiry_year,
                    row.expiry_month,
                )
            except BusinessException as exc:
                errors.append(f"第 {row.row_number} 行：{exc.message}")
                continue
            period = (
                row.effective_year,
                row.effective_month,
                row.expiry_year,
                row.expiry_month,
            )
            key = (canonical, code)
            if key in seen:
                if seen[key] == period:
                    skipped_count += 1
                    continue
                errors.append(
                    f"第 {row.row_number} 行：与文件内已有「{canonical} / {code}」生效失效不一致"
                )
                continue
            seen[key] = period
            if self.repository.get(db, canonical, code) is not None:
                skipped_count += 1
                continue
            pending.append(
                VehicleModelInsert(
                    location=canonical,
                    vehicle_model=code,
                    effective_year=row.effective_year,
                    effective_month=row.effective_month,
                    expiry_year=row.expiry_year,
                    expiry_month=row.expiry_month,
                    created_by=audit["created_by"],
                    updated_by=audit["updated_by"],
                )
            )

        if errors:
            raise BusinessException("；".join(errors))

        try:
            for payload in pending:
                self.repository.create(db, payload)
            db.commit()
        except IntegrityError:
            db.rollback()
            raise BusinessException("该基地已有此项目") from None
        except BusinessException:
            db.rollback()
            raise
        except Exception:
            db.rollback()
            logger.exception("导入项目失败")
            raise
        return VehicleModelUploadResultVO(
            created_count=len(pending),
            skipped_count=skipped_count,
        )

    def build_template(self) -> tuple[bytes, str]:
        workbook, sheet, data_font, thin_border = _new_import_sheet()
        _write_import_row(
            sheet,
            2,
            ("青岛", "T13J", 2026, 1, None, None),
            data_font,
            thin_border,
        )
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        return buffer.getvalue(), "项目导入模板.xlsx"

    def export_excel(
        self,
        db: Session,
        current_user: UserVO,
        location: str | None = None,
    ) -> tuple[bytes, str]:
        name = (location or "").strip()
        if name:
            canonical = self.location_service.resolve_canonical(db, name)
            if canonical is None:
                raise BusinessException("没有这座基地")
            records = self.repository.list_by_location(db, canonical)
        else:
            records = self.repository.list_all(db)

        workbook, sheet, data_font, thin_border = _new_import_sheet()
        for row_idx, record in enumerate(records, start=2):
            _write_import_row(
                sheet,
                row_idx,
                (
                    record.location,
                    record.vehicle_model,
                    record.effective_year,
                    record.effective_month,
                    record.expiry_year,
                    record.expiry_month,
                ),
                data_font,
                thin_border,
            )
        stamp = now().strftime("%Y%m%d_%H%M%S")
        return workbook_to_xlsx_bytes(workbook, current_user), f"项目一览_{stamp}.xlsx"

    def _require_row(
        self,
        db: Session,
        location: str,
        vehicle_model: str,
    ) -> VehicleModel:
        name = (location or "").strip()
        code = self.normalize(vehicle_model)
        if not name or not code:
            raise NotFoundException("项目不存在")
        record = self.repository.get(db, name, code)
        if record is None:
            raise NotFoundException("项目不存在")
        return record

    def _validate_period(
        self,
        effective_year: int,
        effective_month: int,
        expiry_year: int | None,
        expiry_month: int | None,
    ) -> None:
        if effective_month < 1 or effective_month > 12:
            raise BusinessException("生效月必须为 1–12")
        has_year = expiry_year is not None
        has_month = expiry_month is not None
        if has_year != has_month:
            raise BusinessException("失效年与失效月必须同空或同有")
        if has_month and (expiry_month is None or expiry_month < 1 or expiry_month > 12):
            raise BusinessException("失效月必须为 1–12")
        if (
            has_year
            and has_month
            and expiry_year is not None
            and expiry_month is not None
            and _period(expiry_year, expiry_month) < _period(effective_year, effective_month)
        ):
            raise BusinessException("失效不得早于生效")

    def _is_not_expired(self, record: VehicleModel, year: int, month: int) -> bool:
        if record.expiry_year is None or record.expiry_month is None:
            return True
        return _period(year, month) <= _period(record.expiry_year, record.expiry_month)

    def _is_in_period(self, record: VehicleModel, year: int, month: int) -> bool:
        as_of = _period(year, month)
        if as_of < _period(record.effective_year, record.effective_month):
            return False
        return self._is_not_expired(record, year, month)

    def _overlaps_year(self, record: VehicleModel, year: int) -> bool:
        if _period(record.effective_year, record.effective_month) > _period(year, 12):
            return False
        if record.expiry_year is None or record.expiry_month is None:
            return True
        return _period(record.expiry_year, record.expiry_month) >= _period(year, 1)

    def _period_status(
        self,
        record: VehicleModel,
        year: int,
        month: int,
    ) -> str:
        as_of = _period(year, month)
        if as_of < _period(record.effective_year, record.effective_month):
            return PERIOD_PENDING
        if not self._is_not_expired(record, year, month):
            return PERIOD_EXPIRED
        return PERIOD_INTRODUCIBLE

    def _to_vo(self, record: VehicleModel, period_status: str) -> VehicleModelVO:
        return VehicleModelVO(
            location=record.location,
            vehicle_model=record.vehicle_model,
            effective_year=record.effective_year,
            effective_month=record.effective_month,
            expiry_year=record.expiry_year,
            expiry_month=record.expiry_month,
            period_status=period_status,
        )
