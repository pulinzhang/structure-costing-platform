"""牌号分类：导入模板与一览 Excel。

核心列与 Parser 同契约：大类 / 牌号 / 备注。
列表导出另附大类序号与审计列，导入时忽略。
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.common.excel_workbook import resolve_worksheet
from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.timezone import now
from app.parsers.masterdata.steel_grade_category_parser import (
    HEADER_CATEGORY_NAME,
    HEADER_GRADE,
    HEADER_REMARK,
)
from app.repositories.masterdata.steel_grade_category import SteelGradeCategoryRepository
from app.schemas.vo.user_vo import UserVO

IMPORT_HEADERS = (
    HEADER_CATEGORY_NAME,
    HEADER_GRADE,
    HEADER_REMARK,
)

EXPORT_EXTRA_HEADERS = (
    "大类序号",
    "创建人",
    "创建时间",
    "更新人",
    "更新时间",
)

EXPORT_HEADERS = IMPORT_HEADERS + EXPORT_EXTRA_HEADERS


def _dt(value: datetime | None) -> str:
    if value is None:
        return ""
    return f"{value:%Y-%m-%d %H:%M:%S}"


def catalog_sheet_styles() -> tuple[Font, Font, PatternFill, Alignment, Border]:
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    data_font = Font(name="微软雅黑", size=9)
    thin_border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    return header_font, data_font, header_fill, header_alignment, thin_border


def attach_import_sheet(workbook: Workbook) -> tuple[Worksheet, Font, Border]:
    sheet = resolve_worksheet(workbook, "牌号分类")
    header_font, data_font, header_fill, header_alignment, thin_border = catalog_sheet_styles()
    for col, title in enumerate(IMPORT_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    sheet.column_dimensions["A"].width = 10
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 24
    return sheet, data_font, thin_border


def write_import_row(
    sheet: Worksheet,
    row_idx: int,
    values: tuple[str | None, ...],
    data_font: Font,
    thin_border: Border,
) -> None:
    for col, value in enumerate(values, start=1):
        cell = sheet.cell(row=row_idx, column=col, value=value)
        cell.font = data_font
        cell.border = thin_border


def build_template() -> tuple[bytes, str]:
    workbook = Workbook()
    sheet, data_font, thin_border = attach_import_sheet(workbook)
    write_import_row(
        sheet,
        2,
        ("冷轧", "DC01", "样例行，可删"),
        data_font,
        thin_border,
    )
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue(), "牌号分类导入模板.xlsx"


def export_excel(
    repository: SteelGradeCategoryRepository,
    db: Session,
    current_user: UserVO,
    *,
    grade: str | None = None,
    category_name: str | None = None,
) -> tuple[bytes, str]:
    records = repository.list_for_export(
        db,
        grade=grade,
        category_name=category_name,
    )
    workbook = Workbook()
    sheet = resolve_worksheet(workbook, "牌号分类")
    header_font, data_font, header_fill, header_alignment, thin_border = catalog_sheet_styles()
    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")

    for col, title in enumerate(EXPORT_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, record in enumerate(records, start=2):
        values: list[str] = [
            record.category_name,
            record.grade,
            record.remark or "",
            str(record.category_seq),
            record.created_by or "",
            _dt(record.created_at),
            record.updated_by or "",
            _dt(record.updated_at),
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_idx, column=col, value=value)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = align_center if col == 4 else align_left

    widths = [10, 24, 24, 10, 12, 20, 12, 20]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "A2"
    last_row = max(1, len(records) + 1)
    sheet.auto_filter.ref = f"A1:H{last_row}"

    stamp = now().strftime("%Y%m%d_%H%M%S")
    return workbook_to_xlsx_bytes(workbook, current_user), f"牌号分类_{stamp}.xlsx"
