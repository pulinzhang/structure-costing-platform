"""牌号分类档案 Excel 导出。"""

from __future__ import annotations

import logging
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.common.excel_workbook import resolve_worksheet
from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.timezone import now
from app.repositories.masterdata.steel_grade_category import SteelGradeCategoryRepository
from app.schemas.dto.steel_grade_category_dto import SteelGradeCategoryExportQuery
from app.schemas.vo.user_vo import UserVO

logger = logging.getLogger("service.steel_grade_category.export")


def _dt(value: datetime | None) -> str:
    if value is None:
        return ""
    return f"{value:%Y-%m-%d %H:%M:%S}"


class SteelGradeCategoryExportService:
    """按当前筛选导出牌号分类档案。"""

    def __init__(self) -> None:
        self.repository = SteelGradeCategoryRepository()

    def export(
        self,
        db: Session,
        query: SteelGradeCategoryExportQuery,
        current_user: UserVO,
    ) -> tuple[bytes, str]:
        records = self.repository.list_for_export(
            db,
            grade=query.grade,
            category_name=query.category_name,
        )

        wb = Workbook()
        ws = resolve_worksheet(wb, "牌号分类")

        header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        data_font = Font(name="微软雅黑", size=9)
        align_left = Alignment(horizontal="left", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )

        headers = [
            "大类序号",
            "大类",
            "牌号",
            "备注",
            "创建人",
            "创建时间",
            "更新人",
            "更新时间",
        ]
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for row_idx, record in enumerate(records, start=2):
            values: list[str] = [
                str(record.category_seq),
                record.category_name,
                record.grade,
                record.remark or "",
                record.created_by or "",
                _dt(record.created_at),
                record.updated_by or "",
                _dt(record.updated_at),
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.font = data_font
                cell.border = thin_border
                if col == 1:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        widths = [10, 10, 24, 24, 12, 20, 12, 20]
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        ws.freeze_panes = "A2"
        last_row = max(1, len(records) + 1)
        ws.auto_filter.ref = f"A1:H{last_row}"

        content = workbook_to_xlsx_bytes(wb, current_user)

        stamp = now().strftime("%Y%m%d_%H%M%S")
        filename = f"牌号分类_{stamp}.xlsx"
        logger.info("导出牌号分类 %s 条 -> %s", len(records), filename)
        return content, filename
