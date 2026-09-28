"""供应商：导入模板与一览 Excel。"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.common.excel_workbook import resolve_worksheet
from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.timezone import now
from app.parsers.masterdata.supplier_parser import (
    HEADER_ABBREVIATION,
    HEADER_ALIASES,
    HEADER_ERP_CODE,
    HEADER_FULL_NAME,
    HEADER_STATUS,
)
from app.repositories.masterdata.supplier.supplier_repository import SupplierRepository
from app.schemas.vo.user_vo import UserVO

IMPORT_HEADERS = (
    HEADER_ABBREVIATION,
    HEADER_FULL_NAME,
    HEADER_ERP_CODE,
    HEADER_ALIASES,
)

EXPORT_HEADERS = (*IMPORT_HEADERS, HEADER_STATUS)


def catalog_sheet_styles() -> tuple[Font, Font, PatternFill, Alignment, Border]:
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A1A22", end_color="1A1A22", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    data_font = Font(name="微软雅黑", size=9)
    thin_border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    return header_font, data_font, header_fill, header_alignment, thin_border


def attach_catalog_sheet(workbook: Workbook) -> tuple[Worksheet, Font, Border]:
    sheet = resolve_worksheet(workbook, "供应商")
    header_font, data_font, header_fill, header_alignment, thin_border = catalog_sheet_styles()
    for col, title in enumerate(IMPORT_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 40
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 36
    return sheet, data_font, thin_border


def attach_export_sheet(workbook: Workbook) -> tuple[Worksheet, Font, Border]:
    sheet = resolve_worksheet(workbook, "供应商")
    header_font, data_font, header_fill, header_alignment, thin_border = catalog_sheet_styles()
    for col, title in enumerate(EXPORT_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 40
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 36
    sheet.column_dimensions["E"].width = 10
    return sheet, data_font, thin_border


def write_import_row(
    sheet: Worksheet,
    row_idx: int,
    values: tuple[str | None, ...],
    data_font: Font,
    thin_border: Border,
) -> None:
    for col, value in enumerate(values, start=1):
        cell = sheet.cell(row=row_idx, column=col, value=value)
        cell.font = data_font
        cell.border = thin_border


def write_catalog_to_workbook(
    repository: SupplierRepository,
    db: Session,
    workbook: Workbook,
    *,
    keyword: str | None = None,
    is_enabled: bool | None = True,
) -> None:
    records = repository.list_for_export(
        db,
        keyword=keyword,
        is_enabled=is_enabled,
    )
    sheet, data_font, thin_border = attach_export_sheet(workbook)
    for row_idx, record in enumerate(records, start=2):
        aliases = "、".join(item.alias for item in record.aliases)
        write_import_row(
            sheet,
            row_idx,
            (
                record.abbreviation,
                record.full_name,
                record.erp_code,
                aliases or None,
                "已启用" if record.is_enabled else "已禁用",
            ),
            data_font,
            thin_border,
        )


def build_template() -> tuple[bytes, str]:
    workbook = Workbook()
    sheet, data_font, thin_border = attach_catalog_sheet(workbook)
    write_import_row(
        sheet,
        2,
        ("常庆", "安庆常庆汽车部件有限公司", "CQ001", "常庆汽配、安庆常庆"),
        data_font,
        thin_border,
    )
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue(), "供应商导入模板.xlsx"


def export_excel(
    repository: SupplierRepository,
    db: Session,
    current_user: UserVO,
    *,
    keyword: str | None = None,
    is_enabled: bool | None = None,
) -> tuple[bytes, str]:
    workbook = Workbook()
    write_catalog_to_workbook(
        repository,
        db,
        workbook,
        keyword=keyword,
        is_enabled=is_enabled,
    )
    stamp = now().strftime("%Y%m%d_%H%M%S")
    return workbook_to_xlsx_bytes(workbook, current_user), f"供应商列表_{stamp}.xlsx"
