"""openpyxl 色板、字体、金额单元格。"""

from decimal import Decimal

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.common.money import Q4, excel_money_cell

FONT_NAME = "微软雅黑"
MONEY_FORMAT = "0.0000"


def fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def font(
    *,
    size: int = 10,
    bold: bool = False,
    color: str = "000000",
) -> Font:
    return Font(name=FONT_NAME, size=size, bold=bold, color=color)


def thin_border() -> Border:
    side = Side(style="thin", color="BFBFBF")
    return Border(left=side, right=side, top=side, bottom=side)


def money_cell(value: Decimal | None) -> Decimal | str:
    return excel_money_cell(value, Q4, allow_negative=True)


def money_cell_blank_zero(value: Decimal | None) -> Decimal | str:
    number = money_cell(value)
    if number == "" or number == 0:
        return ""
    return number


def alignments() -> tuple[Alignment, Alignment, Alignment, Alignment]:
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    indent = Alignment(horizontal="left", vertical="center", indent=1)
    return center, left, right, indent
