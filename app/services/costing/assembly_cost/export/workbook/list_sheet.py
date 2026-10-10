"""成本一览 / 成品核价一览。"""

from decimal import Decimal

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.common.audit import audit_status_label
from app.common.excel_workbook import ExcelCellValue
from app.common.money import Q4, format_money, quantize_money
from app.models.costing.assembly_cost import AssemblyCost
from app.services.costing.assembly_cost.codes import (
    part_short_code,
    procurement_part_short_code,
)
from app.services.costing.assembly_cost.export.workbook.columns import (
    FINANCE_LIST_SHEET,
    LIST_CENTER_COLS,
    LIST_HEADERS,
    LIST_KPI_COL,
    LIST_MONEY_COLS,
    LIST_WIDTHS,
    PROCUREMENT_LIST_CENTER_COLS,
    PROCUREMENT_LIST_HEADERS,
    PROCUREMENT_LIST_KPI_COL,
    PROCUREMENT_LIST_MONEY_COLS,
    PROCUREMENT_LIST_SHEET,
    PROCUREMENT_LIST_WIDTHS,
)
from app.services.costing.assembly_cost.export.workbook.labels import (
    calc_status_label,
    confirm_status_label,
    object_kind_label,
    tree_level_label,
)
from app.services.costing.assembly_cost.export.workbook.layout import (
    SheetLayout,
    apply_freeze_and_filter,
)
from app.services.costing.assembly_cost.export.workbook.styles import (
    MONEY_FORMAT,
    fill,
    font,
    money_cell,
    money_cell_blank_zero,
)


def write_list_sheet(
    layout: SheetLayout,
    ws: Worksheet,
    records: list[AssemblyCost],
    *,
    filter_text: str,
    view: str,
    confirm_map: dict[int, str],
    seq_no_by_source_node_id: dict[int, int | None] | None = None,
    missing_l0_note: str = "",
) -> None:
    is_procurement = view == "procurement"
    headers = PROCUREMENT_LIST_HEADERS if is_procurement else LIST_HEADERS
    widths = PROCUREMENT_LIST_WIDTHS if is_procurement else LIST_WIDTHS
    money_cols = PROCUREMENT_LIST_MONEY_COLS if is_procurement else LIST_MONEY_COLS
    kpi_col = PROCUREMENT_LIST_KPI_COL if is_procurement else LIST_KPI_COL
    center_cols = PROCUREMENT_LIST_CENTER_COLS if is_procurement else LIST_CENTER_COLS
    ws.title = PROCUREMENT_LIST_SHEET if is_procurement else FINANCE_LIST_SHEET
    col_count = len(headers)
    if not is_procurement:
        base_note = (
            "说明：金额为系统核价快照（材料 / 焊接 / 月总账 / 间接 / 总费用）；"
            "表头与末行为本表合计（当前导出范围内的总成）；"
            "末行公式不进筛选、不随筛选变化；"
            "明细 Sheet 仅作佐证，不以 Excel 公式回算；"
            "序号与材料清单总成序号一致。"
        )
        note = f"{missing_l0_note} {base_note}" if missing_l0_note else base_note
    else:
        note = (
            "说明：本表为核算参考，不参与采购询价/合同业务。"
            "金额列为材料成本、焊接费/电泳黑漆等、总成本（关注列）。"
            "表头与末行为本册对象合计；末行公式不进筛选、不随筛选变化。"
        )
    header_row = layout.write_banner(
        ws,
        title="零件成本" if is_procurement else "成品核价一览表",
        filter_text=filter_text,
        count_label=_list_count_label(records, is_procurement=is_procurement),
        note=note,
        col_count=col_count,
    )
    if missing_l0_note and not is_procurement:
        ws.row_dimensions[3].height = 36
        ws.cell(row=3, column=1).alignment = Alignment(
            horizontal="left",
            vertical="center",
            indent=1,
            wrap_text=True,
        )
    ws.row_dimensions[2].height = 36
    ws.cell(row=2, column=1).alignment = Alignment(
        horizontal="left",
        vertical="center",
        indent=1,
        wrap_text=True,
    )
    layout.write_simple_header(
        ws,
        header_row,
        headers,
        money_cols=money_cols,
        kpi_col=kpi_col,
        header_height=32 if is_procurement else 22,
    )

    l0_seq_map = seq_no_by_source_node_id or {}
    for row_offset, record in enumerate(records, start=1):
        row_idx = header_row + row_offset
        row_data: list[ExcelCellValue]
        if is_procurement:
            confirm_status = ""
            if record.procurement_object_id is not None:
                raw = confirm_map.get(int(record.procurement_object_id))
                confirm_status = confirm_status_label(raw)
            row_data = [
                row_offset,
                procurement_part_short_code(part_number=record.part_number),
                record.part_number or "",
                record.part_name or "",
                object_kind_label(record.object_kind),
                tree_level_label(record.level),
                money_cell(record.material_amount),
                money_cell_blank_zero(record.welding_amount),
                money_cell(record.total_amount),
                confirm_status,
                calc_status_label(record.calc_status),
                audit_status_label(record.audit_status),
            ]
        else:
            source_node_id = getattr(record, "source_node_id", None)
            mapped_seq = l0_seq_map.get(int(source_node_id)) if source_node_id is not None else None
            row_data = [
                mapped_seq if mapped_seq is not None else "",
                record.year,
                record.month,
                record.location or "",
                record.part_number or "",
                part_short_code(part_number=record.part_number, level=record.level),
                record.part_name or "",
                record.vehicle_model or "",
                record.bom_version or "",
                money_cell(record.material_amount),
                money_cell(record.welding_amount),
                money_cell(record.total_amount),
                money_cell(record.indirect_amount),
                money_cell(record.grand_total),
                calc_status_label(record.calc_status),
                audit_status_label(record.audit_status),
            ]
        layout.append_styled_row(
            ws,
            row_idx,
            row_data,
            money_cols=money_cols,
            seq=row_offset,
            center_cols=center_cols,
            kpi_col=kpi_col,
        )

    layout.apply_column_widths(ws, widths)
    _write_list_total_row(
        layout,
        ws,
        header_row=header_row,
        record_count=len(records),
        col_count=col_count,
        money_cols=money_cols,
        kpi_col=kpi_col,
    )
    apply_freeze_and_filter(
        ws,
        header_row=header_row,
        col_count=col_count,
        row_count=len(records),
        filter_empty_placeholder=False,
    )
    if is_procurement:
        layout.polish_sheet(ws, header_row=header_row)


def _sum_amount_field(records: list[AssemblyCost], field_name: str) -> Decimal:
    total = Decimal("0")
    for record in records:
        value = getattr(record, field_name, None)
        if value is None or value == "":
            continue
        total += quantize_money(value, Q4, allow_negative=True)
    return total.quantize(Q4)


def _money_text(records: list[AssemblyCost], field_name: str) -> str:
    return format_money(_sum_amount_field(records, field_name), Q4, 4)


def _list_count_label(records: list[AssemblyCost], *, is_procurement: bool) -> str:
    total = len(records)
    ok_count = sum(1 for record in records if record.calc_status == "ok")
    material = _money_text(records, "material_amount")
    welding = _money_text(records, "welding_amount")
    if is_procurement:
        cost_total = _money_text(records, "total_amount")
        return (
            f"共 {total} 个对象（核算正常 {ok_count}） · "
            f"本册对象合计　材料 {material}　焊接 {welding}　总成本 {cost_total}"
        )
    month_total = _money_text(records, "total_amount")
    indirect = _money_text(records, "indirect_amount")
    grand_total = _money_text(records, "grand_total")
    return (
        f"共 {total} 条（核算正常 {ok_count}） · "
        f"本表合计　材料 {material}　焊接 {welding}　"
        f"月总账 {month_total}　间接 {indirect}　总费用 {grand_total}"
    )


def _write_list_total_row(
    layout: SheetLayout,
    ws: Worksheet,
    *,
    header_row: int,
    record_count: int,
    col_count: int,
    money_cols: set[int],
    kpi_col: int,
) -> None:
    total_row = header_row + record_count + 1
    data_start = header_row + 1
    data_end = header_row + record_count
    for col in range(1, col_count + 1):
        cell = ws.cell(row=total_row, column=col, value=None)
        cell.font = font(size=10, bold=True)
        cell.border = layout.border
        cell.fill = fill("ECEFF1")
        if col in money_cols:
            cell.alignment = layout.right
            cell.number_format = MONEY_FORMAT
        else:
            cell.alignment = layout.center
    label = ws.cell(row=total_row, column=1, value="合计")
    label.font = font(size=10, bold=True)
    label.fill = fill("ECEFF1")
    label.alignment = layout.center
    kpi_fill = fill("FFCDD2")
    for col in sorted(money_cols):
        letter = get_column_letter(col)
        cell = ws.cell(row=total_row, column=col)
        if record_count > 0:
            cell.value = f"=SUM({letter}{data_start}:{letter}{data_end})"
        else:
            cell.value = Decimal("0.0000")
        cell.font = font(size=10, bold=True, color="B71C1C")
        cell.fill = kpi_fill if col == kpi_col else fill("FFF3E0")
        cell.alignment = layout.right
        cell.number_format = MONEY_FORMAT
    ws.row_dimensions[total_row].height = 22
