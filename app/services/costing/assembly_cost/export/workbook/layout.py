"""表头横幅、普通表头、数据行、打印与合并。"""

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.common.excel_workbook import ExcelCellValue
from app.common.timezone import now
from app.services.costing.assembly_cost.export.workbook.styles import (
    MONEY_FORMAT,
    alignments,
    fill,
    font,
    thin_border,
)


class SheetLayout:
    def __init__(self) -> None:
        self.center, self.left, self.right, self.indent = alignments()
        self.border = thin_border()

    def write_banner(
        self,
        ws: Worksheet,
        *,
        title: str,
        filter_text: str,
        count_label: str,
        note: str,
        col_count: int,
    ) -> int:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
        c1 = ws.cell(row=1, column=1, value=title)
        c1.font = font(size=15, bold=True, color="FFFFFF")
        c1.fill = fill("1A237E")
        c1.alignment = self.indent
        ws.row_dimensions[1].height = 28

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
        c2 = ws.cell(
            row=2,
            column=1,
            value=(f"{filter_text} · {count_label}" f"　　导出 {now().strftime('%Y-%m-%d %H:%M')}"),
        )
        c2.font = font(size=9, color="455A64")
        c2.fill = fill("E8EAF6")
        c2.alignment = self.indent
        ws.row_dimensions[2].height = 20

        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=col_count)
        c3 = ws.cell(row=3, column=1, value=note)
        c3.font = font(size=9, color="B71C1C")
        c3.fill = fill("FFF3E0")
        c3.alignment = self.indent
        ws.row_dimensions[3].height = 20
        return 4

    def write_simple_header(
        self,
        ws: Worksheet,
        header_row: int,
        headers: list[str],
        *,
        money_cols: set[int],
        kpi_col: int | None = None,
        header_height: int = 22,
    ) -> None:
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col, value=title)
            cell.font = font(size=10, bold=True, color="FFFFFF")
            cell.alignment = self.center
            cell.border = self.border
            if kpi_col is not None and col == kpi_col:
                cell.fill = fill("C62828")
            elif col in money_cols:
                cell.fill = fill("EF6C00")
            else:
                cell.fill = fill("546E7A")
        ws.row_dimensions[header_row].height = header_height

    def append_styled_row(
        self,
        ws: Worksheet,
        row_idx: int,
        row_data: list[ExcelCellValue],
        *,
        money_cols: set[int],
        seq: int,
        center_cols: set[int] | None = None,
        kpi_col: int | None = None,
    ) -> None:
        stripe = "F7F7F7" if seq % 2 == 0 else "FFFFFF"
        kpi_fill = "FFEBEE" if seq % 2 == 0 else "FFCDD2"
        aligned_center = center_cols or {1}
        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.font = font(size=9)
            cell.border = self.border
            if kpi_col is not None and col == kpi_col:
                cell.fill = fill(kpi_fill)
                cell.font = font(size=10, bold=True, color="B71C1C")
                cell.alignment = self.right
                cell.number_format = MONEY_FORMAT
            elif col in money_cols:
                cell.fill = fill(stripe)
                cell.alignment = self.right
                if value != "":
                    cell.number_format = MONEY_FORMAT
            elif col in aligned_center:
                cell.fill = fill(stripe)
                cell.alignment = self.center
            else:
                cell.fill = fill(stripe)
                cell.alignment = self.left
        ws.row_dimensions[row_idx].height = 20

    def apply_column_widths(self, ws: Worksheet, widths: list[int]) -> None:
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width

    def merge_item_group(
        self,
        ws: Worksheet,
        *,
        start_row: int,
        end_row: int,
        columns: range,
        left_cols: set[int] | None = None,
    ) -> None:
        if end_row <= start_row:
            return
        left_aligned = left_cols or set()
        for col in columns:
            ws.merge_cells(
                start_row=start_row,
                start_column=col,
                end_row=end_row,
                end_column=col,
            )
            cell = ws.cell(row=start_row, column=col)
            cell.alignment = self.left if col in left_aligned else self.center

    def polish_sheet(self, ws: Worksheet, *, header_row: int) -> None:
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 110
        ws.print_title_rows = f"1:{header_row}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.35
        ws.page_margins.right = 0.35
        ws.page_margins.top = 0.4
        ws.page_margins.bottom = 0.4


def apply_freeze_and_filter(
    ws: Worksheet,
    *,
    header_row: int,
    col_count: int,
    row_count: int,
    filter_empty_placeholder: bool = True,
) -> None:
    ws.freeze_panes = f"A{header_row + 1}"
    last_col = get_column_letter(col_count)
    if row_count > 0:
        last_row = header_row + row_count
    elif filter_empty_placeholder:
        last_row = header_row + 1
    else:
        last_row = header_row
    ws.auto_filter.ref = f"A{header_row}:{last_col}{last_row}"
