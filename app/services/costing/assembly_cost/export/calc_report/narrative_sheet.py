"""过程说明 Sheet。"""

from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.vo.assembly_cost_vo import AssemblyCostCalcLogVO
from app.services.costing.assembly_cost.export.calc_report.columns import NARRATIVE_COL_WIDTH
from app.services.costing.assembly_cost.export.calc_report.layout import apply_report_page
from app.services.costing.assembly_cost.export.calc_report.styles import (
    PAPER,
    STONE,
    alignments,
    fill,
    font,
)


def write_narrative_sheet(ws: Worksheet, log: AssemblyCostCalcLogVO) -> None:
    _center, _left, _right, indent = alignments()
    apply_report_page(ws, tab_color=STONE, fit_height=0, landscape=False)
    ws.column_dimensions["A"].width = NARRATIVE_COL_WIDTH
    ws.merge_cells("A1:A1")
    title = ws["A1"]
    title.value = "计算过程说明"
    title.font = font(size=15, bold=True, color="FFFFFF")
    title.fill = fill(STONE)
    title.alignment = indent
    ws.row_dimensions[1].height = 28

    text = (log.detail_text or "").strip() or "（无过程说明）"
    ws.merge_cells("A3:A3")
    body = ws["A3"]
    body.value = text
    body.font = Font(name="Consolas", size=9, color=STONE)
    body.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    body.fill = fill(PAPER)
    line_count = text.count("\n") + 1
    ws.row_dimensions[3].height = min(400, max(80, line_count * 14))
