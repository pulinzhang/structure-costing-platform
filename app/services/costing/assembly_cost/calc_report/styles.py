"""计算报告浅色纸面色板（Brabus light）与单元格样式。"""

from decimal import Decimal

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.common.money import Q4, excel_money_cell

FONT_NAME = "微软雅黑"
MONEY_FORMAT = "0.0000"
ACCENT = "D32F2F"
INK = "18181B"
MUTED = "52525B"
PAPER = "FFFFFF"
CREAM = "FAFAFA"
STONE = "3F3F46"
MONEY_HEADER = "D32F2F"
SUMMARY_FILL = "FAFAFA"


def fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def font(*, size: int = 10, bold: bool = False, color: str = "000000") -> Font:
    return Font(name=FONT_NAME, size=size, bold=bold, color=color)


def thin_border() -> Border:
    side = Side(style="thin", color="E4E4E7")
    return Border(left=side, right=side, top=side, bottom=side)


def money_cell(value: object) -> Decimal | str:
    return excel_money_cell(value, Q4, allow_negative=True)


def alignments() -> tuple[Alignment, Alignment, Alignment, Alignment]:
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")
    indent = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    return center, left, right, indent
