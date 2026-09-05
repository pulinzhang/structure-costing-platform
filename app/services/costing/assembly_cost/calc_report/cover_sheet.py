"""封面摘要 Sheet。"""

from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet

from app.common.timezone import now
from app.schemas.vo.assembly_cost_vo import AssemblyCostCalcLogVO
from app.services.costing.assembly_cost.calc_report.columns import COVER_COL_WIDTHS
from app.services.costing.assembly_cost.calc_report.labels import (
    format_when,
    status_label,
)
from app.services.costing.assembly_cost.calc_report.layout import apply_report_page
from app.services.costing.assembly_cost.calc_report.styles import (
    ACCENT,
    CREAM,
    INK,
    MUTED,
    PAPER,
    SUMMARY_FILL,
    alignments,
    fill,
    font,
)


def write_cover_sheet(ws: Worksheet, log: AssemblyCostCalcLogVO, filter_text: str) -> None:
    _center, left, _right, indent = alignments()
    apply_report_page(ws, tab_color=ACCENT, fit_height=1)
    ws.column_dimensions["A"].width = COVER_COL_WIDTHS[0]
    ws.column_dimensions["B"].width = COVER_COL_WIDTHS[1]
    ws.column_dimensions["C"].width = COVER_COL_WIDTHS[2]
    ws.column_dimensions["D"].width = COVER_COL_WIDTHS[3]

    ws.merge_cells("A1:D1")
    title = ws["A1"]
    title.value = "总成核价计算报告"
    title.font = font(size=20, bold=True, color="FFFFFF")
    title.fill = fill(ACCENT)
    title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:D2")
    sub = ws["A2"]
    sub.value = (
        f"Assembly Costing Process Report  ·  {filter_text}"
        f"  ·  导出 {now().strftime('%Y-%m-%d %H:%M')}"
    )
    sub.font = font(size=9, color=MUTED)
    sub.fill = fill(CREAM)
    sub.alignment = indent
    ws.row_dimensions[2].height = 22

    period = f"{log.year}-{log.month:02d}"
    rows: list[tuple[str, str, str, str]] = [
        ("范围", f"{log.location} / {log.vehicle_model} / {log.bom_version}", "核价期", period),
        (
            "状态",
            status_label(log.status),
            "操作人",
            log.operator_name or "—",
        ),
        ("开始", format_when(log.started_at), "结束", format_when(log.finished_at)),
        (
            "核价卡",
            str(log.subject_total if log.subject_total is not None else log.l0_total or 0),
            "批次",
            log.calc_batch_id,
        ),
        (
            "成功",
            str(log.ok_count),
            "部分成功",
            str(log.partial_count),
        ),
        (
            "失败",
            str(log.failed_count),
            "跳过",
            str(log.skipped_approved_count),
        ),
    ]
    start = 4
    for offset, (label_a, value_a, label_b, value_b) in enumerate(rows):
        row = start + offset
        for col, (label, value) in enumerate(
            ((label_a, value_a), (label_b, value_b)),
            start=1,
        ):
            label_cell = ws.cell(row=row, column=col * 2 - 1, value=label)
            label_cell.font = font(size=9, color=MUTED)
            label_cell.fill = fill(CREAM)
            label_cell.alignment = indent
            value_cell = ws.cell(row=row, column=col * 2, value=value)
            value_cell.font = font(size=11, bold=True, color=INK)
            value_cell.fill = fill(PAPER)
            value_cell.alignment = left
        ws.row_dimensions[row].height = 22

    summary_row = start + len(rows) + 1
    ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=4)
    summary = ws.cell(row=summary_row, column=1, value=log.message or "（无汇总说明）")
    summary.font = font(size=10, color=MUTED)
    summary.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
    summary.fill = fill(SUMMARY_FILL)
    ws.row_dimensions[summary_row].height = 48

    note_row = summary_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=4)
    note = ws.cell(
        row=note_row,
        column=1,
        value="金额为本次计算快照，与报告页一致；明细 Sheet 不作公式回算。",
    )
    note.font = font(size=9, color=ACCENT)
    note.alignment = indent
