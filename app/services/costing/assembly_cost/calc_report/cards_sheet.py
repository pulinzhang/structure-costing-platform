"""核价卡明细 Sheet。"""

from decimal import Decimal

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.vo.assembly_cost_vo import (
    AssemblyCostCalcCardTraceVO,
    AssemblyCostCalcLogVO,
)
from app.services.costing.assembly_cost.calc_report.columns import (
    CARD_CENTER_COLS,
    CARD_HEADERS,
    CARD_KPI_COL,
    CARD_MONEY_COLS,
    CARD_STATUS_COL,
    CARD_WIDTHS,
)
from app.services.costing.assembly_cost.calc_report.labels import (
    STATUS_COLOR,
    status_label,
)
from app.services.costing.assembly_cost.calc_report.layout import apply_report_page
from app.services.costing.assembly_cost.calc_report.styles import (
    ACCENT,
    CREAM,
    INK,
    MONEY_FORMAT,
    MONEY_HEADER,
    MUTED,
    PAPER,
    STONE,
    alignments,
    fill,
    font,
    money_cell,
    thin_border,
)
from app.services.costing.assembly_cost.codes import part_short_code


def write_cards_sheet(
    ws: Worksheet,
    log: AssemblyCostCalcLogVO,
    cards: list[AssemblyCostCalcCardTraceVO],
    filter_text: str,
) -> None:
    center, left, right, indent = alignments()
    border = thin_border()
    apply_report_page(ws, tab_color=ACCENT, fit_height=0, show_grid=None)
    col_count = len(CARD_HEADERS)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    c1 = ws.cell(row=1, column=1, value="核价卡明细")
    c1.font = font(size=15, bold=True, color="FFFFFF")
    c1.fill = fill(ACCENT)
    c1.alignment = indent
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    c2 = ws.cell(
        row=2,
        column=1,
        value=(
            f"{log.location} / {log.vehicle_model} / {log.bom_version}"
            f"  ·  {log.year}-{log.month:02d}  ·  {filter_text}"
        ),
    )
    c2.font = font(size=9, color=MUTED)
    c2.fill = fill(CREAM)
    c2.alignment = indent
    ws.row_dimensions[2].height = 20

    header_row = 4
    for col, title in enumerate(CARD_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = font(size=9, bold=True, color="FFFFFF")
        cell.alignment = center
        cell.border = border
        if col == CARD_KPI_COL:
            cell.fill = fill(ACCENT)
        elif col in CARD_MONEY_COLS:
            cell.fill = fill(MONEY_HEADER)
        else:
            cell.fill = fill(STONE)
    ws.row_dimensions[header_row].height = 22

    if not cards:
        empty_row = header_row + 1
        ws.merge_cells(
            start_row=empty_row,
            start_column=1,
            end_row=empty_row,
            end_column=col_count,
        )
        empty = ws.cell(row=empty_row, column=1, value="没有符合筛选的核价卡。")
        empty.font = font(size=10, color=MUTED)
        empty.alignment = Alignment(horizontal="center", vertical="center")
        empty.fill = fill(PAPER)
        ws.row_dimensions[empty_row].height = 28

    for seq, card in enumerate(cards, start=1):
        row_idx = header_row + seq
        stripe = "FAF7F2" if seq % 2 == 0 else "FFFFFF"
        kpi_fill = "FFF1F2" if seq % 2 == 0 else "FFE4E6"
        values: list[str | int | Decimal] = [
            seq,
            f"L{card.level}",
            "树根" if card.is_tree_root else "",
            card.part_number,
            part_short_code(part_number=card.part_number, level=card.level),
            card.part_name or "",
            card.object_kind or "",
            status_label(card.status),
            money_cell(card.material_amount),
            money_cell(card.welding_amount),
            money_cell(card.should_cost),
            money_cell(card.grand_total),
            card.missing_leaf_count,
            card.message or "",
            "；".join(card.missing_parts),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = border
            if col == CARD_KPI_COL:
                cell.fill = fill(kpi_fill)
                cell.font = font(size=10, bold=True, color=ACCENT)
                cell.alignment = right
                cell.number_format = MONEY_FORMAT
            elif col in CARD_MONEY_COLS:
                cell.fill = fill(stripe)
                cell.font = font(size=9)
                cell.alignment = right
                cell.number_format = MONEY_FORMAT
            elif col == CARD_STATUS_COL:
                cell.fill = fill(stripe)
                cell.font = font(
                    size=9,
                    bold=True,
                    color=STATUS_COLOR.get(card.status, INK),
                )
                cell.alignment = center
            elif col in CARD_CENTER_COLS:
                cell.fill = fill(stripe)
                cell.font = font(size=9)
                cell.alignment = center
            else:
                cell.fill = fill(stripe)
                cell.font = font(size=9, color=INK)
                cell.alignment = left
        ws.row_dimensions[row_idx].height = 22

    for col, width in enumerate(CARD_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = f"A{header_row + 1}"
    last_data_row = header_row + max(len(cards), 1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(col_count)}{last_data_row}"
