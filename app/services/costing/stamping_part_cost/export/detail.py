"""明细口径 Excel 导出。"""

from __future__ import annotations

from typing import Any


import logging
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.common.audit import audit_status_label
from app.repositories.stamping.stamping_process_fee_repository import (
    StampingProcessFeeRepository,
)
from app.repositories.yishiyiyi.scrap_price.scrap_price_repository import (
    StampingScrapPriceRepository,
)
from app.repositories.costing.steel_price.steel_price_repository import SteelPriceRepository
from app.services.costing.stamping_part_cost.export.formatters import format_datetime
from app.services.costing.stamping_part_cost.export.styles import (
    THIN_BORDER,
    WEAK_DATA_FONT,
    WEAK_HEADER_FILL,
    WEAK_HEADER_FONT,
)
from app.services.costing.stamping_part_cost.export.trace_lookups import load_trace_maps
from app.services.costing.stamping_part_cost.helpers import (
    format_number as _num,
    format_process_fee_trace,
    format_scrap_price_trace,
    format_steel_price_excel_trace,
    part_excel_row_col,
)

logger = logging.getLogger("service.stamping_part_cost_export")


def export_detail(
    db: Session,
    rows: list[Any],
    *,
    steel_price_repo: SteelPriceRepository,
    scrap_price_repo: StampingScrapPriceRepository,
    process_fee_repo: StampingProcessFeeRepository,
) -> tuple[bytes, str]:
    """明细导出：业务字段 + 末尾弱化溯源/关联 ID + 审核元数据。"""
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "零件价格核算"

    headers = [
        "年度",
        "月份",
        "基地",
        "冲压工厂",
        "车型",
        "版本",
        "零件号",
        "零件名称",
        "材质",
        "料厚",
        "单重kg",
        "料片长",
        "料片宽",
        "可制件数",
        "密度ρ_steel",
        "材料损耗率λ",
        "废料回收率η",
        "材料消耗M kg",
        "净重H kg",
        "废料重量O kg",
        "材料利用率%",
        "钢厂",
        "钢材吨价",
        "废料吨价",
        "①钢材采购 M×P×(1+λ)",
        "②损耗废料扣除 M×λ×η×Q",
        "③边角废料扣除 O×η×Q",
        "废料扣除合计 ②+③",
        "原材料成本R ①−②−③",
        "加工费",
        "单件成本",
        "计算状态",
        "失败原因",
        "审核状态",
        "计算批次",
        "公式版本",
        "计算时间",
        # 溯源（弱化）
        "来源Excel",
        "工作表",
        "行/列",
        "钢材价溯源",
        "废料价溯源",
        "加工费溯源",
        "钢材价ID",
        "废料价ID",
        "加工费ID",
    ]
    weak_from_col = 38  # 来源Excel 起
    header_fill = PatternFill("solid", fgColor="1E1E1E")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    for col, title in enumerate(headers, start=1):
        cell = sheet.cell(1, col, title)
        if col >= weak_from_col:
            cell.fill = WEAK_HEADER_FILL
            cell.font = WEAK_HEADER_FONT
        else:
            cell.fill = header_fill
            cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = THIN_BORDER

    steel_by_id, scrap_by_id, fee_by_id = load_trace_maps(
        db,
        rows,
        steel_price_repo=steel_price_repo,
        scrap_price_repo=scrap_price_repo,
        process_fee_repo=process_fee_repo,
    )

    for row_index, row in enumerate(rows, start=2):
        values = [
            row.year,
            row.month,
            row.location,
            row.stamping_supplier,
            row.vehicle_model,
            row.version,
            row.part_number,
            row.part_name or "",
            row.material or "",
            _num(row.thickness, 4),
            _num(row.part_weight, 4),
            _num(row.sheet_length, 2),
            _num(row.sheet_width, 2),
            row.parts_per_sheet if row.parts_per_sheet is not None else "",
            _num(row.density, 4),
            _num(getattr(row, "material_loss_rate", None), 6),
            _num(getattr(row, "scrap_recovery_rate", None), 6),
            _num(row.blank_weight_kg, 10),
            _num(row.finished_total_kg, 10),
            _num(row.scrap_weight_kg, 10),
            _num(row.material_utilization_pct, 2),
            row.steel_mill or "",
            _num(row.steel_ton_price, 2),
            _num(row.scrap_ton_price, 2),
            _num(row.material_amount_sheet, 4),
            _num(getattr(row, "loss_scrap_deduction", None), 4),
            _num(getattr(row, "process_scrap_deduction", None), 4),
            _num(row.scrap_amount_sheet, 4),
            _num(row.material_net_per_part, 4),
            _num(row.process_fee_amount, 3),
            _num(row.unit_cost, 4),
            row.calc_status or "",
            row.calc_fail_reason or "",
            audit_status_label(row.audit_status),
            row.calc_batch_id or "",
            row.formula_version or "",
            format_datetime(row.calculated_at),
            row.excel_filename or "",
            row.excel_sheet_name or "",
            part_excel_row_col(row),
            format_steel_price_excel_trace(
                steel_by_id.get(row.steel_price_id)
                if row.steel_price_id is not None
                else None
            ),
            format_scrap_price_trace(
                scrap_by_id.get(row.scrap_price_id)
                if row.scrap_price_id is not None
                else None
            ),
            format_process_fee_trace(
                fee_by_id.get(row.process_fee_id)
                if row.process_fee_id is not None
                else None
            ),
            row.steel_price_id if row.steel_price_id is not None else "",
            row.scrap_price_id if row.scrap_price_id is not None else "",
            row.process_fee_id if row.process_fee_id is not None else "",
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, col, value)
            cell.border = THIN_BORDER
            if col >= weak_from_col:
                cell.font = WEAK_DATA_FONT
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=col in {41, 42, 43},
                )
            else:
                cell.alignment = Alignment(vertical="center")

    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 32
    buffer = BytesIO()
    workbook.save(buffer)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"冲压零件价格核算明细_{stamp}.xlsx"
    logger.info("导出明细格式 %s 条 → %s", len(rows), filename)
    return buffer.getvalue(), filename
