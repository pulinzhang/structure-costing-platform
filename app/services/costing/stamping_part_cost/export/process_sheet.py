"""工艺清单口径 Excel 导出。"""

from __future__ import annotations

import logging
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from app.common.bom.manufacturing_category import short_part_number_digits
from app.common.excel_workbook import resolve_worksheet
from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.swing_shear import steel_kg_price_with_swing_shear
from app.common.timezone import now
from app.repositories.bom.bom_repository import BomNodeRepository, BomSnapshotRepository
from app.repositories.costing.steel_price.steel_price_repository import SteelPriceRepository
from app.repositories.stamping.stamping_process_fee_repository import (
    StampingProcessFeeRepository,
)
from app.repositories.yishiyiyi.scrap_price.scrap_price_repository import (
    StampingScrapPriceRepository,
)
from app.schemas.vo.user_vo import UserVO
from app.services.costing.stamping_part_cost.export.bom_resolver import (
    resolve_assembly_supply_no,
)
from app.services.costing.stamping_part_cost.export.formatters import formula_number
from app.services.costing.stamping_part_cost.export.tailor_weld_display import (
    blank_term_has_mass_inputs,
    build_blank_row_m_formula,
    build_steel_price_index,
    build_tailor_group_n_formula,
    build_tailor_group_o_formula,
    build_tailor_group_r_formula,
    load_blanks_by_part_id,
    resolve_part_cost_display,
)
from app.services.costing.stamping_part_cost.export.trace_lookups import load_trace_maps
from app.services.costing.stamping_part_cost.helpers import (
    DEFAULT_MATERIAL_LOSS_RATE,
    DEFAULT_SCRAP_RECOVERY_RATE,
    DEFAULT_STEEL_DENSITY,
    format_process_fee_trace,
    format_scrap_price_trace,
    format_steel_price_excel_trace,
    part_excel_row_col,
    to_decimal,
)
from app.services.costing.stamping_part_cost.helpers import (
    format_number as _num,
)
from app.services.settings.system.settings_service import SettingsService

logger = logging.getLogger("service.stamping_part_cost_export")

# 拼焊总成共用列：跨子片行合并（料厚/长/宽/消耗/单价等按片分行）
_TAILOR_MERGED_COLS = frozenset(
    {
        "A",
        "B",
        "C",
        "D",
        "E",
        "F",
        "G",
        "I",
        "O",
        "P",
        "R",
        "S",
        "T",
        "U",
        "V",
        "W",
        "X",
        "Y",
        "Z",
        "AA",
        "AB",
        "AC",
        "AD",
        "AE",
        "AG",
        "AH",
        "AJ",
        "AK",
        "AL",
        "AM",
        "AN",
        "AO",
        "AP",
        "AQ",
    }
)
_NUMBER_COLS = frozenset(
    {
        "H",
        "I",
        "K",
        "L",
        "M",
        "N",
        "O",
        "P",
        "Q",
        "R",
        "S",
        "T",
        "U",
        "V",
        "W",
        "X",
        "Y",
        "Z",
        "AA",
        "AB",
        "AQ",
    }
)
_TRACE_COLS = frozenset({"AC", "AD", "AE", "AF", "AG", "AH", "AI", "AJ", "AK", "AL", "AM"})
_ID_COLS = frozenset({"AI", "AJ", "AK", "AL", "AM"})
_RESULT_COLS = frozenset({"S"})
_FINAL_COLS = frozenset({"AB"})
_FEE_RATE_COLS = frozenset({"X", "Y", "Z", "AA"})
_TEXT_LEFT_COLS = frozenset({"B", "C", "D", "E", "F", "G", "J"})
PROCESS_SHEET_HEADERS: tuple[tuple[str, str], ...] = (
    ("A", "序号"),
    ("B", "总成供货号"),
    ("C", "零件号"),
    ("D", "短码"),
    ("E", "零件名称"),
    ("F", "零件简图"),
    ("G", "材质"),
    ("H", "料厚"),
    ("I", "净重"),
    ("J", "零件尺寸"),
    ("K", "长"),
    ("L", "宽(步距)"),
    ("M", "可制件数"),
    ("N", "预估单件材料消耗(KG)"),
    ("O", "材料利用率"),
    ("P", "废料重量"),
    ("Q", "钢板单价(元/KG)"),
    ("R", "废料单价(元/KG)"),
    ("S", "原材料成本"),
    ("T", "加工费"),
    ("U", "模具费"),
    ("V", "检具费"),
    ("W", "直接成本"),
    ("X", "仓储+包装"),
    ("Y", "管理费"),
    ("Z", "财务费"),
    ("AA", "税前利润"),
    ("AB", "最终价"),
    ("AC", "来源Excel"),
    ("AD", "工作表"),
    ("AE", "行/列"),
    ("AF", "钢材价溯源"),
    ("AG", "废料价溯源"),
    ("AH", "加工费溯源"),
    ("AI", "钢材价ID"),
    ("AJ", "废料价ID"),
    ("AK", "加工费ID"),
    ("AL", "模具费ID"),
    ("AM", "检具费ID"),
    ("AN", "是否利废"),
    ("AO", "利用哪类废料"),
    ("AP", "是否摆剪"),
    ("AQ", "摆剪加价(元/吨)"),
)
_PROCESS_COL_LETTERS = tuple(letter for letter, _title in PROCESS_SHEET_HEADERS)
_FILENAME_ILLEGAL = '<>:"/\\|?*'


def _steel_kg_for_material(row: object) -> Decimal | None:
    """钢板单价(元/kg) = (联动吨价 + 本月摆剪加价) / 1000。"""
    steel_ton = to_decimal(getattr(row, "steel_ton_price", None))
    if steel_ton is None:
        return None
    surcharge = to_decimal(getattr(row, "swing_shear_fee_per_ton", None)) or Decimal("0")
    return steel_kg_price_with_swing_shear(steel_ton, surcharge)


def _swing_shear_fee_cell(row: object) -> Decimal | str:
    if not getattr(row, "has_swing_shear", False):
        return ""
    number = to_decimal(getattr(row, "swing_shear_fee_per_ton", None))
    if number is None:
        return ""
    return _num(number, 2)


def _part_identity_cells(
    row: object,
    *,
    seq: object,
    assembly_supply_no: str,
) -> dict[str, object]:
    return {
        "A": seq,
        "B": assembly_supply_no,
        "C": getattr(row, "part_number", None) or "",
        "D": short_part_number_digits(getattr(row, "part_number", None)) or "",
        "E": getattr(row, "part_name", None) or "",
        "F": "",
        "G": getattr(row, "material", None) or "",
    }


def _offal_and_swing_cells(row: object) -> dict[str, object]:
    return {
        "AN": "是" if getattr(row, "is_offal_nested", False) else "否",
        "AO": getattr(row, "offal_source", None) or "",
        "AP": "是" if getattr(row, "has_swing_shear", False) else "否",
        "AQ": _swing_shear_fee_cell(row),
    }


def _fee_amount_cell(value: object, digits: int) -> Decimal | str:
    """加工费 / 模具费 / 检具费：None 或 0 导出为空，不写 0。"""
    number = to_decimal(value)
    if number is None or number == 0:
        return ""
    return _num(number, digits)


def _finished_fee_formula_cells(
    excel_row: int,
    *,
    warehouse_packaging_rate: Decimal,
    management_rate: Decimal,
    finance_rate: Decimal,
    pretax_profit_rate: Decimal,
) -> dict[str, object]:
    """直接成本 + 四项费用（固定顺序）+ 最终价，均为 Excel 公式。

    顺序：仓储+包装 → 管理费 → 财务费 → 税前利润
    费率须为系统配置 Decimal（禁止用费用÷直接成本反推，会出脏小数）。
    直接成本为 0 时单元格显示空（与加工费/模具费/检具费一致）。
    """
    v = f"W{excel_row}"
    direct = f"S{excel_row}+T{excel_row}+U{excel_row}+V{excel_row}"
    return {
        "W": f'=IF(({direct})=0,"",{direct})',
        "X": f'=IF({v}="","",{v}*{formula_number(warehouse_packaging_rate)})',
        "Y": f'=IF({v}="","",{v}*{formula_number(management_rate)})',
        "Z": f'=IF({v}="","",{v}*{formula_number(finance_rate)})',
        "AA": f'=IF({v}="","",{v}*{formula_number(pretax_profit_rate)})',
        "AB": (f'=IF({v}="","",{v}+X{excel_row}+Y{excel_row}' f"+Z{excel_row}+AA{excel_row})"),
    }


def _filename_token(value: object) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    for char in _FILENAME_ILLEGAL:
        text = text.replace(char, "")
    text = " ".join(text.split())
    return text or None


def _unique_row_value(rows: list[Any], attr: str) -> object | None:
    values = {getattr(row, attr, None) for row in rows}
    values.discard(None)
    values.discard("")
    if len(values) != 1:
        return None
    return next(iter(values))


def _unique_row_int(rows: list[Any], attr: str) -> int | None:
    raw = _unique_row_value(rows, attr)
    if isinstance(raw, bool):
        return None
    if isinstance(raw, int):
        return raw
    if isinstance(raw, str):
        text = raw.strip()
        if text.isdigit():
            return int(text)
    return None


def build_process_sheet_filename(
    rows: list[Any],
    *,
    year: int | None = None,
    month: int | None = None,
    location: str | None = None,
    stamping_supplier: str | None = None,
    vehicle_model: str | None = None,
    version: str | None = None,
) -> str:
    """冲压零件核价结果导出文件名：年月 + 基地 + 工厂 + 车型版本（有则带上）。"""
    year_value = year if year is not None else _unique_row_int(rows, "year")
    month_value = month if month is not None else _unique_row_int(rows, "month")
    location_value = location or _unique_row_value(rows, "location")
    supplier_value = stamping_supplier or _unique_row_value(rows, "stamping_supplier")
    vehicle_value = vehicle_model or _unique_row_value(rows, "vehicle_model")
    version_value = version or _unique_row_value(rows, "version")

    parts = ["冲压零件核价结果"]
    if year_value is not None and month_value is not None:
        parts.append(f"{year_value}年{month_value}月")
    elif year_value is not None:
        parts.append(f"{year_value}年")
    for token in (
        _filename_token(location_value),
        _filename_token(supplier_value),
        _filename_token(vehicle_value),
        _filename_token(version_value),
    ):
        if token:
            parts.append(token)
    stamp = now().strftime("%Y%m%d_%H%M%S")
    return "_".join(parts) + f"_{stamp}.xlsx"


def _part_size_formula(material_row: int, thickness_row: int) -> str:
    """零件尺寸 = 材质-料厚；整数料厚用 TEXT 保留 1 位小数。"""
    return f'=G{material_row}&"-"&TEXT(H{thickness_row},"0.0###")'


def write_process_sheet_to_workbook(
    workbook: Workbook,
    db: Session,
    rows: list[Any],
    *,
    bom_snapshot_repo: BomSnapshotRepository,
    bom_node_repo: BomNodeRepository,
    steel_price_repo: SteelPriceRepository,
    scrap_price_repo: StampingScrapPriceRepository,
    process_fee_repo: StampingProcessFeeRepository,
    year: int | None = None,
    month: int | None = None,
    location: str | None = None,
    stamping_supplier: str | None = None,
    vehicle_model: str | None = None,
    version: str | None = None,
    sheet_title: str = "冲压核价结果",
) -> tuple[str, int]:
    """
    将工艺清单口径核价结果写入已有 Workbook（不 save）。
    返回 (主表 sheet 名, 主表零件行数近似用 rows 长度)。
    """
    sheet = resolve_worksheet(workbook, sheet_title)

    last_col = len(_PROCESS_COL_LETTERS)
    last_col_letter = _PROCESS_COL_LETTERS[-1]
    trace_cols = _TRACE_COLS
    id_cols = _ID_COLS
    fee_rate_cols = _FEE_RATE_COLS

    settings_service = SettingsService()
    warehouse_packaging_rate = settings_service.get_finished_warehouse_packaging_rate(db)
    management_rate = settings_service.get_finished_management_rate(db)
    finance_rate = settings_service.get_finished_finance_rate(db)
    pretax_profit_rate = settings_service.get_finished_pretax_profit_rate(db)

    # —— 样式 ——
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    title_font = Font(name="微软雅黑", size=14, bold=True, color="1F1F1F")
    meta_font = Font(name="微软雅黑", size=9, color="666666")
    group_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    header_font = Font(name="微软雅黑", size=9, bold=True, color="1F1F1F")
    weak_header_font = Font(name="微软雅黑", size=8, bold=True, color="757575")
    data_font = Font(name="微软雅黑", size=9, color="000000")
    weak_data_font = Font(name="微软雅黑", size=8, color="9E9E9E")
    fee_rate_font = Font(name="微软雅黑", size=8, color="757575")
    accent_font = Font(name="微软雅黑", size=9, bold=True, color="C62828")

    fill_title = PatternFill("solid", fgColor="F5F5F5")
    fill_group_id = PatternFill("solid", fgColor="424242")
    fill_group_product = PatternFill("solid", fgColor="455A64")
    fill_group_process = PatternFill("solid", fgColor="1565C0")
    fill_group_fee = PatternFill("solid", fgColor="2E7D32")
    fill_group_fee_rate = PatternFill("solid", fgColor="607D8B")
    fill_group_final = PatternFill("solid", fgColor="C62828")
    fill_group_trace = PatternFill("solid", fgColor="78909C")
    fill_header = PatternFill("solid", fgColor="ECEFF1")
    fill_header_trace = PatternFill("solid", fgColor="ECEFF1")
    fill_header_fee_rate = PatternFill("solid", fgColor="ECEFF1")
    fill_zebra = PatternFill("solid", fgColor="FAFAFA")
    fill_zebra_trace = PatternFill("solid", fgColor="F5F5F5")
    fill_fee_rate = PatternFill("solid", fgColor="F7F7F7")
    fill_fee_rate_zebra = PatternFill("solid", fgColor="F0F0F0")
    fill_result = PatternFill("solid", fgColor="FFF3E0")
    fill_final = PatternFill("solid", fgColor="FFEBEE")

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")

    # 范围摘要 + 核价口径（本批快照参数）
    meta_parts: list[str] = []
    density_text = "7.85"
    loss_rate_text = "0.003"
    loss_permille_text = "3"
    recovery_text = "0.97"
    recovery_percent_text = "97"
    if rows:
        sample = rows[0]
        if sample.year and sample.month:
            meta_parts.append(f"{sample.year}年{sample.month}月")
        if sample.location:
            meta_parts.append(str(sample.location))
        if sample.stamping_supplier:
            meta_parts.append(str(sample.stamping_supplier))
        density = to_decimal(getattr(sample, "density", None))
        loss_rate = to_decimal(getattr(sample, "material_loss_rate", None))
        recovery = to_decimal(getattr(sample, "scrap_recovery_rate", None))
        if density is not None:
            density_text = formula_number(density)
        if loss_rate is not None:
            loss_rate_text = formula_number(loss_rate)
            loss_permille_text = formula_number(loss_rate * Decimal("1000"))
        if recovery is not None:
            recovery_text = formula_number(recovery)
            recovery_percent_text = formula_number(recovery * Decimal("100"))
    meta_text = " · ".join(meta_parts) if meta_parts else "按筛选条件导出"
    formula_hint = (
        f"核价口径："
        f"N=L×W×T×ρ_steel÷1,000,000÷n；"
        f"P=N−I；"
        f"S=N×Q×(1+λ)−N×λ×η×R−P×η×R；"
        f"直接成本=S+T+U+V；"
        f"最终价=直接成本+仓储包装费+管理费+财务费+税前利润。"
        f"拼焊件按子片分行：H/K/L/Q 各片一格；"
        f"子片 N=K×L×H×ρ/1e6/M；"
        f"总成 P=SUM(N)−I，S=Σ(Ni×Qi×(1+λ))−SUM(N)×λ×η×R−P×η×R。"
        f"其中 ρ_steel={density_text} g/cm³，"
        f"λ={loss_rate_text}（{loss_permille_text}‰），"
        f"η={recovery_text}（{recovery_percent_text}%），"
        f"Q/R 为元/kg（吨价÷1000；Q 含本月摆剪加价），"
        f"÷1,000,000 为 mm·g/cm³ → kg 量纲换算。"
    )

    # Row 1 标题
    sheet.merge_cells(f"A1:{last_col_letter}1")
    title_cell = sheet["A1"]
    title_cell.value = "冲压零件核价结果（材料费 + 加工费 + 模具费 + 检具费）"
    title_cell.font = title_font
    title_cell.fill = fill_title
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    for col in range(1, last_col + 1):
        sheet.cell(1, col).fill = fill_title
        sheet.cell(1, col).border = thin
    sheet.row_dimensions[1].height = 28

    # Row 2 说明
    sheet.merge_cells(f"A2:{last_col_letter}2")
    meta_cell = sheet["A2"]
    meta_cell.value = (
        f"{meta_text}　|　共 {len(rows)} 件　|　{formula_hint}"
        f"　|　AC~AH 人可读溯源可核对/补填；AI~AM 为关联 ID"
    )
    meta_cell.font = meta_font
    meta_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for col in range(1, last_col + 1):
        sheet.cell(2, col).fill = fill_title
        sheet.cell(2, col).border = thin
    sheet.row_dimensions[2].height = 48

    # Row 3 分组表头
    group_fills = {
        range(1, 3): fill_group_id,
        range(3, 11): fill_group_product,
        range(11, 20): fill_group_process,
        range(20, 21): fill_group_fee,
        range(21, 23): fill_group_fee,
        range(23, 24): fill_group_fee,
        range(24, 28): fill_group_fee_rate,
        range(28, 29): fill_group_final,
        range(29, 35): fill_group_trace,
        range(35, 40): fill_group_trace,
        range(40, 42): fill_group_product,
        range(42, 44): fill_group_id,
    }
    for merge_range, label in (
        ("A3:B3", "标识"),
        ("C3:J3", "产品信息"),
        ("K3:S3", "材料费"),
        ("T3:T3", "加工费"),
        ("U3:U3", "模具费"),
        ("V3:V3", "检具费"),
        ("W3:W3", "直接成本"),
        ("X3:AA3", "成品费率费用"),
        ("AB3:AB3", "最终价"),
        ("AC3:AH3", "溯源"),
        ("AI3:AM3", "关联ID"),
        ("AN3:AO3", "利废"),
        ("AP3:AQ3", "摆剪"),
    ):
        sheet.merge_cells(merge_range)
        cell = sheet[merge_range.split(":")[0]]
        cell.value = label
        cell.font = group_font
        cell.alignment = center
    for col in range(1, last_col + 1):
        cell = sheet.cell(3, col)
        cell.border = thin
        cell.font = group_font
        cell.alignment = center
        for col_range, fill in group_fills.items():
            if col in col_range:
                cell.fill = fill
                break
    sheet.row_dimensions[3].height = 22

    # Row 4 列名
    for col_letter, title in PROCESS_SHEET_HEADERS:
        cell = sheet[f"{col_letter}4"]
        cell.value = title
        is_trace = col_letter in trace_cols
        is_fee_rate = col_letter in fee_rate_cols
        if is_trace or is_fee_rate:
            cell.font = weak_header_font
        else:
            cell.font = header_font
        if is_trace:
            cell.fill = fill_header_trace
        elif is_fee_rate:
            cell.fill = fill_header_fee_rate
        else:
            cell.fill = fill_header
        cell.alignment = center
        cell.border = thin
    sheet.row_dimensions[4].height = 40

    # 数据行
    data_start = 5
    number_cols = _NUMBER_COLS
    result_cols = _RESULT_COLS
    final_cols = _FINAL_COLS
    text_left_cols = _TEXT_LEFT_COLS
    snapshot_l0_cache: dict[tuple[str, str, str], dict[str, str]] = {}
    fill_plain = PatternFill("solid", fgColor="FFFFFF")

    steel_by_id, scrap_by_id, fee_by_id = load_trace_maps(
        db,
        rows,
        steel_price_repo=steel_price_repo,
        scrap_price_repo=scrap_price_repo,
        process_fee_repo=process_fee_repo,
    )
    blanks_by_part = load_blanks_by_part_id(db, rows)
    steel_index = build_steel_price_index(db, rows, steel_repo=steel_price_repo)

    def paint_data_row(excel_row: int, values: dict[str, object], zebra: bool) -> None:
        for col_letter in _PROCESS_COL_LETTERS:
            cell = sheet[f"{col_letter}{excel_row}"]
            if col_letter in values:
                cell.value = values[col_letter]
            is_trace = col_letter in trace_cols
            is_fee_rate = col_letter in fee_rate_cols
            if is_trace:
                cell.font = weak_data_font
            elif is_fee_rate:
                cell.font = fee_rate_font
            elif col_letter in final_cols:
                cell.font = accent_font
            else:
                cell.font = data_font
            cell.border = thin
            if col_letter in final_cols:
                cell.fill = fill_final
            elif col_letter in result_cols:
                cell.fill = fill_result
            elif is_fee_rate:
                cell.fill = fill_fee_rate_zebra if zebra else fill_fee_rate
            elif is_trace:
                cell.fill = fill_zebra_trace if zebra else fill_plain
            elif zebra:
                cell.fill = fill_zebra
            else:
                cell.fill = fill_plain

            if col_letter in {"AF", "AG", "AH"}:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif col_letter in id_cols:
                cell.alignment = center
            elif col_letter in text_left_cols or col_letter == "AA":
                cell.alignment = left
            elif col_letter in number_cols or col_letter == "A":
                cell.alignment = right if col_letter != "A" else center
            else:
                cell.alignment = center

            if col_letter == "H":
                cell.number_format = "0.0###"
            elif col_letter in {"I", "Q", "R", "T", "U", "V"}:
                cell.number_format = "0.####"
            elif col_letter in {"K", "L"}:
                cell.number_format = "0.##"
            elif col_letter == "O":
                cell.number_format = "0.00%"
            elif col_letter in {"N", "P", "S", "W", "X", "Y", "Z", "AA", "AB"}:
                cell.number_format = "0.0000"
            elif col_letter == "M" or col_letter in id_cols:
                cell.number_format = "0"

        multi_line = any(
            isinstance(values.get(col), str) and "\n" in str(values.get(col))
            for col in ("AF", "AG", "AH")
        )
        sheet.row_dimensions[excel_row].height = 52 if multi_line else 18

    def merge_tailor_group(first_row: int, last_row: int) -> None:
        if last_row <= first_row:
            return
        for col_letter in _TAILOR_MERGED_COLS:
            sheet.merge_cells(f"{col_letter}{first_row}:{col_letter}{last_row}")

    excel_row = data_start
    for offset, row in enumerate(rows):
        density = to_decimal(getattr(row, "density", None)) or DEFAULT_STEEL_DENSITY
        loss_rate = (
            to_decimal(getattr(row, "material_loss_rate", None)) or DEFAULT_MATERIAL_LOSS_RATE
        )
        recovery = (
            to_decimal(getattr(row, "scrap_recovery_rate", None)) or DEFAULT_SCRAP_RECOVERY_RATE
        )
        display = resolve_part_cost_display(
            row,
            blanks_by_part.get(int(row.source_part_id or 0), []),
            steel_index,
        )
        scrap_ton = to_decimal(row.scrap_ton_price)
        scrap_kg = (scrap_ton / Decimal("1000")) if scrap_ton is not None else None
        assembly_supply_no = resolve_assembly_supply_no(
            db,
            bom_snapshot_repo=bom_snapshot_repo,
            bom_node_repo=bom_node_repo,
            location=row.location,
            vehicle_model=row.vehicle_model,
            version=row.version,
            part_number=row.part_number,
            snapshot_l0_cache=snapshot_l0_cache,
        )
        seq = row.seq_no if row.seq_no is not None else offset + 1
        zebra = offset % 2 == 1
        scrap_trace = format_scrap_price_trace(
            scrap_by_id.get(row.scrap_price_id) if row.scrap_price_id is not None else None
        )
        fee_trace = format_process_fee_trace(
            fee_by_id.get(row.process_fee_id) if row.process_fee_id is not None else None
        )
        process_fee_cell = _fee_amount_cell(row.process_fee_amount, 3)
        die_fee_cell = _fee_amount_cell(getattr(row, "die_fee_amount", None), 4)
        gauge_fee_cell = _fee_amount_cell(getattr(row, "gauge_fee_amount", None), 4)

        if display.is_tailor_welded and len(display.blank_terms) >= 2:
            terms = display.blank_terms
            first_row = excel_row
            last_row = excel_row + len(terms) - 1
            priced_rows = [
                first_row + index
                for index, term in enumerate(terms)
                if blank_term_has_mass_inputs(term) and term.steel_kg_price is not None
            ]
            can_formula = len(priced_rows) == len(terms)
            is_offal_nested = bool(getattr(row, "is_offal_nested", False))
            if is_offal_nested:
                group_n = ""
                group_o = ""
            else:
                group_n = build_tailor_group_n_formula(first_row, last_row)
                group_o = build_tailor_group_o_formula(first_row, last_row)
            if can_formula:
                group_r: Decimal | str = build_tailor_group_r_formula(
                    priced_rows,
                    first_row=first_row,
                    last_row=last_row,
                    loss_rate=loss_rate,
                    recovery=recovery,
                )
            else:
                net_per_part = to_decimal(row.material_net_per_part)
                group_r = net_per_part if net_per_part is not None else ""

            for index, term in enumerate(terms):
                blank_row = first_row + index
                blank_values: dict[str, object] = {
                    "H": (_num(term.thickness, 4) if term.thickness is not None else ""),
                    "J": _part_size_formula(first_row, blank_row),
                    "K": (_num(term.length, 2) if term.length is not None else ""),
                    "L": (_num(term.width, 2) if term.width is not None else ""),
                    "M": (term.parts_per_sheet if term.parts_per_sheet is not None else ""),
                    "N": (
                        build_blank_row_m_formula(blank_row, density)
                        if blank_term_has_mass_inputs(term)
                        else ""
                    ),
                    "Q": (float(term.steel_kg_price) if term.steel_kg_price is not None else ""),
                    "AF": format_steel_price_excel_trace(term.steel),
                    "AI": term.steel.id if term.steel is not None else "",
                }
                if index == 0:
                    blank_values.update(
                        {
                            **_part_identity_cells(
                                row,
                                seq=seq,
                                assembly_supply_no=assembly_supply_no,
                            ),
                            "I": (_num(row.part_weight, 4) if row.part_weight is not None else ""),
                            "O": group_n,
                            "P": group_o,
                            "R": float(scrap_kg) if scrap_kg is not None else "",
                            "S": group_r,
                            "T": process_fee_cell,
                            "U": die_fee_cell,
                            "V": gauge_fee_cell,
                            **_finished_fee_formula_cells(
                                first_row,
                                warehouse_packaging_rate=warehouse_packaging_rate,
                                management_rate=management_rate,
                                finance_rate=finance_rate,
                                pretax_profit_rate=pretax_profit_rate,
                            ),
                            "AC": row.excel_filename or "",
                            "AD": row.excel_sheet_name or "",
                            "AE": part_excel_row_col(row),
                            "AG": scrap_trace,
                            "AH": fee_trace,
                            "AJ": (row.scrap_price_id if row.scrap_price_id is not None else ""),
                            "AK": (row.process_fee_id if row.process_fee_id is not None else ""),
                            "AL": (
                                row.die_fee_id
                                if getattr(row, "die_fee_id", None) is not None
                                else ""
                            ),
                            "AM": (
                                row.gauge_fee_id
                                if getattr(row, "gauge_fee_id", None) is not None
                                else ""
                            ),
                            **_offal_and_swing_cells(row),
                        }
                    )
                paint_data_row(blank_row, blank_values, zebra)
            merge_tailor_group(first_row, last_row)
            excel_row = last_row + 1
            continue

        steel_kg = _steel_kg_for_material(row)
        is_offal_nested = bool(getattr(row, "is_offal_nested", False))
        if is_offal_nested:
            # 利废件没有新板边角废料，也没有材料利用率。
            # O 列是百分比格式：写 100 会显示成 10000%；写 1 才是 100%。此处留空，不展示。
            material_n = f"=I{excel_row}" if row.part_weight is not None else ""
            material_o = ""
            material_p = ""
            steel_q: float | str = ""
            material_s = f"=I{excel_row}*R{excel_row}" if row.part_weight is not None else ""
            steel_trace = "利废：成品净重×废料单价（从废料里捡出的成品）"
        else:
            material_n = (
                f"=K{excel_row}*L{excel_row}*H{excel_row}"
                f"*{formula_number(density)}/1000000/M{excel_row}"
            )
            material_o = f"=I{excel_row}/N{excel_row}"
            material_p = f"=N{excel_row}-I{excel_row}"
            steel_q = float(steel_kg) if steel_kg is not None else ""
            material_s = (
                f"=N{excel_row}*Q{excel_row}*(1+{formula_number(loss_rate)})"
                f"-N{excel_row}*{formula_number(loss_rate)}"
                f"*{formula_number(recovery)}*R{excel_row}"
                f"-P{excel_row}*R{excel_row}*{formula_number(recovery)}"
            )
            steel_trace = format_steel_price_excel_trace(
                steel_by_id.get(row.steel_price_id) if row.steel_price_id is not None else None
            )
        values: dict[str, object] = {
            **_part_identity_cells(
                row,
                seq=seq,
                assembly_supply_no=assembly_supply_no,
            ),
            "H": _num(row.thickness, 4) if row.thickness is not None else "",
            "I": _num(row.part_weight, 4) if row.part_weight is not None else "",
            "J": _part_size_formula(excel_row, excel_row),
            "K": _num(row.sheet_length, 2) if row.sheet_length is not None else "",
            "L": _num(row.sheet_width, 2) if row.sheet_width is not None else "",
            "M": row.parts_per_sheet if row.parts_per_sheet is not None else "",
            "N": material_n,
            "O": material_o,
            "P": material_p,
            "Q": steel_q,
            "R": float(scrap_kg) if scrap_kg is not None else "",
            "S": material_s,
            "T": process_fee_cell,
            "U": die_fee_cell,
            "V": gauge_fee_cell,
            **_finished_fee_formula_cells(
                excel_row,
                warehouse_packaging_rate=warehouse_packaging_rate,
                management_rate=management_rate,
                finance_rate=finance_rate,
                pretax_profit_rate=pretax_profit_rate,
            ),
            "AC": row.excel_filename or "",
            "AD": row.excel_sheet_name or "",
            "AE": part_excel_row_col(row),
            "AF": steel_trace,
            "AG": scrap_trace,
            "AH": fee_trace,
            "AI": row.steel_price_id if row.steel_price_id is not None else "",
            "AJ": row.scrap_price_id if row.scrap_price_id is not None else "",
            "AK": row.process_fee_id if row.process_fee_id is not None else "",
            "AL": row.die_fee_id if getattr(row, "die_fee_id", None) is not None else "",
            "AM": row.gauge_fee_id if getattr(row, "gauge_fee_id", None) is not None else "",
            **_offal_and_swing_cells(row),
        }
        paint_data_row(excel_row, values, zebra)
        excel_row += 1

    last_data_row = excel_row - 1 if rows else 4

    widths = {
        "A": 6,
        "B": 22,
        "C": 14,
        "D": 8,
        "E": 20,
        "F": 10,
        "G": 14,
        "H": 12,
        "I": 9,
        "J": 18,
        "K": 14,
        "L": 14,
        "M": 10,
        "N": 14,
        "O": 11,
        "P": 10,
        "Q": 16,
        "R": 13,
        "S": 12,
        "T": 10,
        "U": 10,
        "V": 10,
        "W": 11,
        "X": 10,
        "Y": 10,
        "Z": 12,
        "AA": 10,
        "AB": 11,
        "AC": 28,
        "AD": 14,
        "AE": 8,
        "AF": 42,
        "AG": 28,
        "AH": 36,
        "AI": 10,
        "AJ": 10,
        "AK": 10,
        "AL": 10,
        "AM": 10,
        "AN": 10,
        "AO": 14,
        "AP": 10,
        "AQ": 16,
    }
    for col_letter, width in widths.items():
        sheet.column_dimensions[col_letter].width = width

    sheet.freeze_panes = "E5"  # 冻结至短码列（A~D：序号/总成供货号/零件号/短码）
    if rows:
        sheet.auto_filter.ref = f"A4:{last_col_letter}{last_data_row}"
    sheet.print_title_rows = "1:4"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    # 附录：核价口径定义（次于「单件」工作表，打开文件默认见单件）
    note = resolve_worksheet(workbook, "冲压核价口径")
    note_font_title = Font(name="微软雅黑", size=12, bold=True, color="1F1F1F")
    note_font = Font(name="微软雅黑", size=10, color="333333")
    note_header = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    note_header_fill = PatternFill("solid", fgColor="455A64")
    note["A1"] = "冲压零件核价口径定义"
    note["A1"].font = note_font_title
    note.merge_cells("A1:C1")
    note.row_dimensions[1].height = 24
    # 用户看到的是纯中文+符号，没有下划线变量名
    note["A2"] = (
        f"本批次参数：钢材密度 ρ（柔 Rho）={density_text} g/cm³；"
        f"材料损耗率 λ（拉姆达 Lambda）={loss_rate_text}（{loss_permille_text}‰）；"
        f"废料回收率 η（伊塔 Eta）={recovery_text}（{recovery_percent_text}%）"
    )
    note["A2"].font = meta_font
    note.merge_cells("A2:C2")

    note_headers = ("符号", "定义", "计算关系 / 单位")
    for col, text in enumerate(note_headers, start=1):
        cell = note.cell(4, col, text)
        cell.font = note_header
        cell.fill = note_header_fill
        cell.alignment = center
        cell.border = thin

    note_rows = [
        ("B", "总成供货号", "同基地 BOM 沿 parent 上溯至 L0；多总成逗号分隔"),
        ("C", "零件号", "本行核算零件号"),
        ("D", "短码", "零件号数字体末四位"),
        ("L, W, T", "板料长、宽(步距)、料厚", "mm；拼焊件每个子片单独一行"),
        ("n", "可制件数", "件/板；拼焊件按子片行填写"),
        ("ρ_steel", "钢材密度", f"{density_text} g/cm³"),
        ("1,000,000", "量纲换算系数", "mm³·(g/cm³) → kg"),
        (
            "N",
            "预估单件材料消耗",
            "普通件 / 拼焊子片行：=K×L×H×ρ/1e6/M；" "总成消耗 = SUM(各子片 N)",
        ),
        ("I", "零件净重", "kg；拼焊件为焊后总成净重，跨子片行合并"),
        ("P", "废料重量", "普通件 P=N−I；拼焊件 P=SUM(N)−I（合并格）"),
        (
            "Q",
            "钢板单价",
            "元/kg（吨价÷1000，含本月摆剪加价）；拼焊件各子片一行各自单价",
        ),
        ("R", "废料单价", "元/kg（吨价÷1000）"),
        ("λ", "材料损耗率", f"{loss_rate_text}（{loss_permille_text}‰）"),
        ("η", "废料回收率", f"{recovery_text}（{recovery_percent_text}%）"),
        ("①", "钢材采购额", "普通件 N×Q×(1+λ)；拼焊件 Σ(Ni×Qi×(1+λ))"),
        ("②", "损耗废料扣除", "N×λ×η×R（拼焊件与 Σ Ni×λ×η×R 等价）"),
        ("③", "边角废料扣除", "P × η × R（总成一次）"),
        (
            "S",
            "原材料成本",
            "普通件：=N×Q×(1+λ)−N×λ×η×R−P×η×R；"
            "拼焊件：=Σ(Ni×Qi×(1+λ))−SUM(N)×λ×η×R−P×η×R（引用各子片行）",
        ),
        ("T", "加工费", "元/件；取自加工费核算结果（拼焊父件一行）；0 或空则单元格留空"),
        ("U", "模具费", "元/件；取自年度模具费最终价；无档案或 0 则单元格留空"),
        ("V", "检具费", "元/件；检具费（万元）÷ 分摊万台；无生效档案或 0 则单元格留空"),
        ("W", "直接成本", "S + T + U + V；合计为 0 则单元格留空"),
        ("X", "仓储+包装", "直接成本 × 仓储+包装管理费用率"),
        ("Y", "管理费", "直接成本 × 管理费率"),
        ("Z", "财务费", "直接成本 × 财务费率"),
        ("AA", "税前利润", "直接成本 × 税前利润率"),
        (
            "AB",
            "最终价",
            "直接成本 + 仓储包装费 + 管理费 + 财务费 + 税前利润"
            "（计算时按系统设置费率落库；导出列为公式）",
        ),
        (
            "AC~AH",
            "溯源（弱化）",
            "按文件/表/行定位核对；钢材含基价·厚度·宽度·锌层；"
            "拼焊件可含多条钢材价；废料为基地年月净价；"
            "加工费为工艺文件+工序依据。单元格可手工补填。",
        ),
        (
            "AI~AM",
            "关联ID（弱化）",
            "钢材价ID / 废料价ID / 加工费ID / 模具费ID / 检具费ID；" "拼焊件钢材价ID 按子片分行。",
        ),
        ("AN~AO", "利废", "是否从废料里出成品；利用哪类废料取板料式样扫描结果"),
        ("AP~AQ", "摆剪", "工序含「摆剪」则为是；加价为快照元/吨，已计入 Q"),
    ]
    for idx, (symbol, name, desc) in enumerate(note_rows, start=5):
        note.cell(idx, 1, symbol).font = note_font
        note.cell(idx, 2, name).font = note_font
        note.cell(idx, 3, desc).font = note_font
        for col in range(1, 4):
            note.cell(idx, col).border = thin
            note.cell(idx, col).alignment = left if col > 1 else center

    note.column_dimensions["A"].width = 14
    note.column_dimensions["B"].width = 18
    note.column_dimensions["C"].width = 52

    return sheet.title, len(rows)


def export_process_sheet(
    db: Session,
    rows: list[Any],
    *,
    bom_snapshot_repo: BomSnapshotRepository,
    bom_node_repo: BomNodeRepository,
    steel_price_repo: SteelPriceRepository,
    scrap_price_repo: StampingScrapPriceRepository,
    process_fee_repo: StampingProcessFeeRepository,
    year: int | None = None,
    month: int | None = None,
    location: str | None = None,
    stamping_supplier: str | None = None,
    vehicle_model: str | None = None,
    version: str | None = None,
    current_user: UserVO,
) -> tuple[bytes, str]:
    """
    工艺清单口径：对齐样例至 S 列 + 加工费 + 模具费 + 检具费 + 成品费率费用 + 最终价。
    普通件一行一件，N/O/P/J/S 保留 Excel 公式；
    直接成本 W=S+T+U+V；四项费用（仓储包装→管理→财务→税前利润）=W×费率；最终价=W+四项。
    拼焊件按子片分行（料厚/长/宽/单价各片一格），净重与成本等列合并；
    子片 N 引用本行 H/K/L/M，总成 P/O/S 引用各子片行。
    总成供货号：同基地 BOM 上溯 L0；多总成时逗号分隔。
    零件号后为短码（数字体末四位）。
    末尾 AC~AH 为人可读溯源（核对/手工补填），AI~AM 为关联 ID。
    """
    workbook = Workbook()
    write_process_sheet_to_workbook(
        workbook,
        db,
        rows,
        bom_snapshot_repo=bom_snapshot_repo,
        bom_node_repo=bom_node_repo,
        steel_price_repo=steel_price_repo,
        scrap_price_repo=scrap_price_repo,
        process_fee_repo=process_fee_repo,
        year=year,
        month=month,
        location=location,
        stamping_supplier=stamping_supplier,
        vehicle_model=vehicle_model,
        version=version,
    )

    content = workbook_to_xlsx_bytes(workbook, current_user)
    filename = build_process_sheet_filename(
        rows,
        year=year,
        month=month,
        location=location,
        stamping_supplier=stamping_supplier,
        vehicle_model=vehicle_model,
        version=version,
    )
    logger.info("导出冲压零件核价结果 %s 条 → %s", len(rows), filename)
    return content, filename
