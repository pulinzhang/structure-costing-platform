"""主表数据行：普通件一行；拼焊按子片分行再合并总成列。

拼焊合并列集合：`columns._TAILOR_MERGED_COLS`。
"""

from __future__ import annotations

from decimal import Decimal
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.common.money import Q4, excel_money_cell
from app.repositories.bom.bom.node import BomNodeRepository
from app.repositories.bom.bom.snapshot import BomSnapshotRepository
from app.repositories.costing.steel_price.steel_price_repository import SteelPriceRepository
from app.repositories.stamping.process_fee.repository import (
    StampingProcessFeeRepository,
)
from app.repositories.yishiyiyi.scrap_price.scrap_price_repository import (
    StampingScrapPriceRepository,
)
from app.services.costing.stamping_part_cost.export.bom_resolver import (
    resolve_assembly_supply_no,
)
from app.services.costing.stamping_part_cost.export.formatters import formula_number
from app.services.costing.stamping_part_cost.export.process_sheet.cells import (
    _fee_amount_cell,
    _finished_fee_formula_cells,
    _offal_and_swing_cells,
    _part_identity_cells,
    _part_size_formula,
    _steel_kg_for_material,
)
from app.services.costing.stamping_part_cost.export.process_sheet.columns import (
    _FINAL_COLS,
    _NUMBER_COLS,
    _PROCESS_COL_LETTERS,
    _RESULT_COLS,
    _SCOPE_WEAK_COLS,
    _TAILOR_MERGED_COLS,
    _TEXT_LEFT_COLS,
)
from app.services.costing.stamping_part_cost.export.tailor_weld_display import (
    blank_term_has_mass_inputs,
    build_blank_row_m_formula,
    build_steel_price_index,
    build_tailor_group_n_formula,
    build_tailor_group_o_formula,
    build_tailor_group_r_formula,
    load_blanks_by_part_id,
    resolve_part_cost_display,
)
from app.services.costing.stamping_part_cost.export.trace_lookups import load_trace_maps
from app.services.costing.stamping_part_cost.helpers import (
    DEFAULT_MATERIAL_LOSS_RATE,
    DEFAULT_SCRAP_RECOVERY_RATE,
    DEFAULT_STEEL_DENSITY,
    format_process_fee_trace,
    format_scrap_price_trace,
    format_steel_price_excel_trace,
    part_excel_row_col,
    to_decimal,
)
from app.services.costing.stamping_part_cost.helpers import (
    format_number as _num,
)


def write_process_data_rows(
    sheet: Worksheet,
    db: Session,
    rows: list[Any],
    *,
    bom_snapshot_repo: BomSnapshotRepository,
    bom_node_repo: BomNodeRepository,
    steel_price_repo: SteelPriceRepository,
    scrap_price_repo: StampingScrapPriceRepository,
    process_fee_repo: StampingProcessFeeRepository,
    warehouse_packaging_rate: Decimal,
    management_rate: Decimal,
    finance_rate: Decimal,
    pretax_profit_rate: Decimal,
    thin: Border,
    data_font: Font,
    weak_data_font: Font,
    fee_rate_font: Font,
    accent_font: Font,
    fill_final: PatternFill,
    fill_result: PatternFill,
    fill_fee_rate: PatternFill,
    fill_fee_rate_zebra: PatternFill,
    fill_zebra_trace: PatternFill,
    fill_zebra: PatternFill,
    center: Alignment,
    left: Alignment,
    right: Alignment,
    trace_cols: frozenset[str],
    id_cols: frozenset[str],
    fee_rate_cols: frozenset[str],
) -> int:
    """写入数据行，返回最后一行号（无数据时为 4）。"""
    # 数据行
    data_start = 5
    number_cols = _NUMBER_COLS
    result_cols = _RESULT_COLS
    final_cols = _FINAL_COLS
    text_left_cols = _TEXT_LEFT_COLS
    snapshot_l0_cache: dict[tuple[str, str, str], dict[str, str]] = {}
    fill_plain = PatternFill("solid", fgColor="FFFFFF")

    steel_by_id, scrap_by_id, fee_by_id = load_trace_maps(
        db,
        rows,
        steel_price_repo=steel_price_repo,
        scrap_price_repo=scrap_price_repo,
        process_fee_repo=process_fee_repo,
    )
    blanks_by_part = load_blanks_by_part_id(db, rows)
    steel_index = build_steel_price_index(db, rows, steel_repo=steel_price_repo)

    def paint_data_row(excel_row: int, values: dict[str, object], zebra: bool) -> None:
        for col_letter in _PROCESS_COL_LETTERS:
            cell = sheet[f"{col_letter}{excel_row}"]
            if col_letter in values:
                cell.value = values[col_letter]
            is_trace = col_letter in trace_cols
            is_fee_rate = col_letter in fee_rate_cols
            is_scope_weak = col_letter in _SCOPE_WEAK_COLS
            if is_trace or is_scope_weak:
                cell.font = weak_data_font
            elif is_fee_rate:
                cell.font = fee_rate_font
            elif col_letter in final_cols:
                cell.font = accent_font
            else:
                cell.font = data_font
            cell.border = thin
            if col_letter in final_cols:
                cell.fill = fill_final
            elif col_letter in result_cols:
                cell.fill = fill_result
            elif is_fee_rate:
                cell.fill = fill_fee_rate_zebra if zebra else fill_fee_rate
            elif is_trace or is_scope_weak:
                cell.fill = fill_zebra_trace if zebra else fill_plain
            elif zebra:
                cell.fill = fill_zebra
            else:
                cell.fill = fill_plain

            if col_letter in {"AH", "AI", "AJ"}:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif col_letter in id_cols:
                cell.alignment = center
            elif col_letter in text_left_cols or col_letter == "AC":
                cell.alignment = left
            elif col_letter in number_cols or col_letter == "A":
                cell.alignment = right if col_letter != "A" else center
            else:
                cell.alignment = center

            if col_letter == "J":
                cell.number_format = "0.0###"
            elif col_letter in {"K", "S", "T", "V", "W", "X"}:
                cell.number_format = "0.####"
            elif col_letter in {"M", "N"}:
                cell.number_format = "0.##"
            elif col_letter == "Q":
                cell.number_format = "0.00%"
            elif col_letter in {"P", "R", "U", "Y", "Z", "AA", "AB", "AC", "AD"}:
                cell.number_format = "0.0000"
            elif col_letter == "O" or col_letter in id_cols:
                cell.number_format = "0"

        multi_line = any(
            isinstance(values.get(col), str) and "\n" in str(values.get(col))
            for col in ("AH", "AI", "AJ")
        )
        sheet.row_dimensions[excel_row].height = 52 if multi_line else 18

    def merge_tailor_group(first_row: int, last_row: int) -> None:
        if last_row <= first_row:
            return
        for col_letter in _TAILOR_MERGED_COLS:
            sheet.merge_cells(f"{col_letter}{first_row}:{col_letter}{last_row}")

    excel_row = data_start
    for offset, row in enumerate(rows):
        density = to_decimal(getattr(row, "density", None)) or DEFAULT_STEEL_DENSITY
        loss_rate = (
            to_decimal(getattr(row, "material_loss_rate", None)) or DEFAULT_MATERIAL_LOSS_RATE
        )
        recovery = (
            to_decimal(getattr(row, "scrap_recovery_rate", None)) or DEFAULT_SCRAP_RECOVERY_RATE
        )
        display = resolve_part_cost_display(
            row,
            blanks_by_part.get(int(row.source_part_id or 0), []),
            steel_index,
        )
        scrap_ton = to_decimal(row.scrap_ton_price)
        scrap_kg = (scrap_ton / Decimal("1000")) if scrap_ton is not None else None
        assembly_supply_no = resolve_assembly_supply_no(
            db,
            bom_snapshot_repo=bom_snapshot_repo,
            bom_node_repo=bom_node_repo,
            location=row.location,
            vehicle_model=row.vehicle_model,
            version=row.version,
            part_number=row.part_number,
            snapshot_l0_cache=snapshot_l0_cache,
        )
        seq = row.seq_no if row.seq_no is not None else offset + 1
        zebra = offset % 2 == 1
        scrap_trace = format_scrap_price_trace(
            scrap_by_id.get(row.scrap_price_id) if row.scrap_price_id is not None else None
        )
        fee_trace = format_process_fee_trace(
            fee_by_id.get(row.process_fee_id) if row.process_fee_id is not None else None
        )
        process_fee_cell = _fee_amount_cell(row.process_fee_amount, 3)
        die_fee_cell = _fee_amount_cell(getattr(row, "die_fee_amount", None), 4)
        gauge_fee_cell = _fee_amount_cell(getattr(row, "gauge_fee_amount", None), 4)

        if display.is_tailor_welded and len(display.blank_terms) >= 2:
            terms = display.blank_terms
            first_row = excel_row
            last_row = excel_row + len(terms) - 1
            priced_rows = [
                first_row + index
                for index, term in enumerate(terms)
                if blank_term_has_mass_inputs(term) and term.steel_kg_price is not None
            ]
            can_formula = len(priced_rows) == len(terms)
            is_offal_nested = bool(getattr(row, "is_offal_nested", False))
            if is_offal_nested:
                group_n = ""
                group_o = ""
            else:
                group_n = build_tailor_group_n_formula(first_row, last_row)
                group_o = build_tailor_group_o_formula(first_row, last_row)
            if can_formula:
                group_r: Decimal | str = build_tailor_group_r_formula(
                    priced_rows,
                    first_row=first_row,
                    last_row=last_row,
                    loss_rate=loss_rate,
                    recovery=recovery,
                )
            else:
                net_per_part = to_decimal(row.material_net_per_part)
                group_r = net_per_part if net_per_part is not None else ""

            for index, term in enumerate(terms):
                blank_row = first_row + index
                blank_values: dict[str, object] = {
                    "J": (_num(term.thickness, 4) if term.thickness is not None else ""),
                    "L": _part_size_formula(first_row, blank_row),
                    "M": (_num(term.length, 2) if term.length is not None else ""),
                    "N": (_num(term.width, 2) if term.width is not None else ""),
                    "O": (term.parts_per_sheet if term.parts_per_sheet is not None else ""),
                    "P": (
                        build_blank_row_m_formula(blank_row, density)
                        if blank_term_has_mass_inputs(term)
                        else ""
                    ),
                    "S": excel_money_cell(term.steel_kg_price, Q4),
                    "AH": format_steel_price_excel_trace(term.steel),
                    "AK": term.steel.id if term.steel is not None else "",
                }
                if index == 0:
                    blank_values.update(
                        {
                            **_part_identity_cells(
                                row,
                                seq=seq,
                                assembly_supply_no=assembly_supply_no,
                            ),
                            "K": (_num(row.part_weight, 4) if row.part_weight is not None else ""),
                            "Q": group_n,
                            "R": group_o,
                            "T": excel_money_cell(scrap_kg, Q4),
                            "U": group_r,
                            "V": process_fee_cell,
                            "W": die_fee_cell,
                            "X": gauge_fee_cell,
                            **_finished_fee_formula_cells(
                                first_row,
                                warehouse_packaging_rate=warehouse_packaging_rate,
                                management_rate=management_rate,
                                finance_rate=finance_rate,
                                pretax_profit_rate=pretax_profit_rate,
                            ),
                            "AE": row.excel_filename or "",
                            "AF": row.excel_sheet_name or "",
                            "AG": part_excel_row_col(row),
                            "AI": scrap_trace,
                            "AJ": fee_trace,
                            "AL": (row.scrap_price_id if row.scrap_price_id is not None else ""),
                            "AM": (row.process_fee_id if row.process_fee_id is not None else ""),
                            "AN": (
                                row.die_fee_id
                                if getattr(row, "die_fee_id", None) is not None
                                else ""
                            ),
                            "AO": (
                                row.gauge_fee_id
                                if getattr(row, "gauge_fee_id", None) is not None
                                else ""
                            ),
                            **_offal_and_swing_cells(row),
                        }
                    )
                paint_data_row(blank_row, blank_values, zebra)
            merge_tailor_group(first_row, last_row)
            excel_row = last_row + 1
            continue

        steel_kg = _steel_kg_for_material(row)
        is_offal_nested = bool(getattr(row, "is_offal_nested", False))
        if is_offal_nested:
            # 利废件没有新板边角废料，也没有材料利用率。
            # Q 列是百分比格式：写 100 会显示成 10000%；写 1 才是 100%。此处留空，不展示。
            material_n = f"=K{excel_row}" if row.part_weight is not None else ""
            material_o = ""
            material_p = ""
            steel_q: Decimal | str = ""
            material_s = f"=K{excel_row}*T{excel_row}" if row.part_weight is not None else ""
            steel_trace = "利废：成品净重×废料单价（从废料里捡出的成品）"
        else:
            material_n = (
                f"=M{excel_row}*N{excel_row}*J{excel_row}"
                f"*{formula_number(density)}/1000000/O{excel_row}"
            )
            material_o = f"=K{excel_row}/P{excel_row}"
            material_p = f"=P{excel_row}-K{excel_row}"
            steel_q = excel_money_cell(steel_kg, Q4)
            material_s = (
                f"=P{excel_row}*S{excel_row}*(1+{formula_number(loss_rate)})"
                f"-P{excel_row}*{formula_number(loss_rate)}"
                f"*{formula_number(recovery)}*T{excel_row}"
                f"-R{excel_row}*T{excel_row}*{formula_number(recovery)}"
            )
            steel_trace = format_steel_price_excel_trace(
                steel_by_id.get(row.steel_price_id) if row.steel_price_id is not None else None
            )
        values: dict[str, object] = {
            **_part_identity_cells(
                row,
                seq=seq,
                assembly_supply_no=assembly_supply_no,
            ),
            "J": _num(row.thickness, 4) if row.thickness is not None else "",
            "K": _num(row.part_weight, 4) if row.part_weight is not None else "",
            "L": _part_size_formula(excel_row, excel_row),
            "M": _num(row.sheet_length, 2) if row.sheet_length is not None else "",
            "N": _num(row.sheet_width, 2) if row.sheet_width is not None else "",
            "O": row.parts_per_sheet if row.parts_per_sheet is not None else "",
            "P": material_n,
            "Q": material_o,
            "R": material_p,
            "S": steel_q,
            "T": excel_money_cell(scrap_kg, Q4),
            "U": material_s,
            "V": process_fee_cell,
            "W": die_fee_cell,
            "X": gauge_fee_cell,
            **_finished_fee_formula_cells(
                excel_row,
                warehouse_packaging_rate=warehouse_packaging_rate,
                management_rate=management_rate,
                finance_rate=finance_rate,
                pretax_profit_rate=pretax_profit_rate,
            ),
            "AE": row.excel_filename or "",
            "AF": row.excel_sheet_name or "",
            "AG": part_excel_row_col(row),
            "AH": steel_trace,
            "AI": scrap_trace,
            "AJ": fee_trace,
            "AK": row.steel_price_id if row.steel_price_id is not None else "",
            "AL": row.scrap_price_id if row.scrap_price_id is not None else "",
            "AM": row.process_fee_id if row.process_fee_id is not None else "",
            "AN": row.die_fee_id if getattr(row, "die_fee_id", None) is not None else "",
            "AO": row.gauge_fee_id if getattr(row, "gauge_fee_id", None) is not None else "",
            **_offal_and_swing_cells(row),
        }
        paint_data_row(excel_row, values, zebra)
        excel_row += 1

    last_data_row = excel_row - 1 if rows else 4
    return last_data_row
