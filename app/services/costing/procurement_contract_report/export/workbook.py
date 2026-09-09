"""协议 Excel 渲染引擎：只画表，不查库。"""

from __future__ import annotations

from typing import Protocol, cast

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.costing.procurement_contract_report.assemble import PriceRow, ReportBundle
from app.services.costing.procurement_contract_report.export.columns import (
    AGREEMENT_SHEET_TITLE,
    ALLOC_GROUP_HEADER,
    ALLOC_PRICE_HEADER,
    ALLOC_QTY_HEADER,
    ANNEX_LABEL,
    CURRENCY_LABEL,
    EMPTY_BODY,
    FIXED_HEADERS,
    MISSING_VALUE,
    NOTES_LABEL,
    NUMBER_LABEL,
    PAGE_NUMBER_FOOTER,
    PAY_SEPARATELY,
    PAY_SEPARATELY_HEADER,
    REMARK_COL_HEADER,
    REMARK_LABEL,
    TAX_EXCLUDED_NOTE,
    TITLE,
    TOOLING_FEE_HEADER,
    TOOLING_GROUP_HEADER,
)
from app.services.costing.procurement_contract_report.export.styles import (
    ANNEX_FONT,
    BOLD,
    CENTER,
    DATA_ROW_HEIGHT,
    HEADER_ROW_HEIGHT,
    INT_FORMAT,
    LEFT,
    META,
    PRICE_FORMAT,
    REMARK_LINE_HEIGHT,
    THIN,
    THOUSANDS_INT_FORMAT,
    TITLE_FONT,
    TITLE_ROW_HEIGHT,
    agreement_column_widths,
)
from app.services.costing.procurement_contract_report.export.values import (
    agreement_remarks,
    agreement_title,
    excel_int,
    excel_price,
    last_column,
    party_a_date_label,
    party_a_label,
    party_b_date_label,
    party_b_label,
    qty_unit,
    settlement_header,
)


class AgreementWorkbookEngine:
    def run(self, bundle: ReportBundle, company_name: str) -> Workbook:
        workbook = Workbook()
        agreement = workbook.active
        if not isinstance(agreement, Worksheet):
            raise RuntimeError("新建工作簿没有工作表")
        agreement.title = AGREEMENT_SHEET_TITLE
        _write_agreement_sheet(agreement, bundle, company_name)
        return workbook


def _write_agreement_sheet(sheet: Worksheet, bundle: ReportBundle, company_name: str) -> None:
    month_count = len(bundle.months)
    end_col = last_column(month_count)
    annex_cell = sheet.cell(1, 1, ANNEX_LABEL)
    annex_cell.font = ANNEX_FONT
    annex_cell.alignment = LEFT
    sheet.merge_cells(start_row=1, start_column=2, end_row=1, end_column=end_col)
    title_cell = sheet.cell(1, 2, agreement_title(bundle, TITLE))
    title_cell.font = TITLE_FONT
    title_cell.alignment = CENTER
    sheet.row_dimensions[1].height = TITLE_ROW_HEIGHT
    _write_meta_row(sheet, month_count, end_col)
    _write_price_headers(sheet, bundle.months)
    data_start = 6
    if bundle.price_rows:
        for offset, row in enumerate(bundle.price_rows):
            _write_price_data_row(sheet, data_start + offset, row)
        footer_row = data_start + len(bundle.price_rows)
    else:
        sheet.merge_cells(
            start_row=data_start, start_column=1, end_row=data_start, end_column=end_col
        )
        sheet.cell(data_start, 2, EMPTY_BODY)
        footer_row = data_start + 1
    _write_footer(sheet, footer_row, end_col, bundle, company_name)
    _apply_column_widths(sheet, agreement_column_widths(month_count))
    _tune_agreement_sheet(sheet)


def _write_meta_row(sheet: Worksheet, month_count: int, end_col: int) -> None:
    last_settlement = 6 + month_count
    tooling_start = last_settlement + 1
    number = sheet.cell(2, 1, NUMBER_LABEL)
    number.font = ANNEX_FONT
    number.alignment = META
    sheet.merge_cells(start_row=2, start_column=6, end_row=2, end_column=last_settlement)
    currency = sheet.cell(2, 6, CURRENCY_LABEL)
    currency.font = ANNEX_FONT
    currency.alignment = META
    sheet.merge_cells(start_row=2, start_column=tooling_start, end_row=2, end_column=end_col)
    tax = sheet.cell(2, tooling_start, TAX_EXCLUDED_NOTE)
    tax.font = ANNEX_FONT
    tax.alignment = META


def _write_price_headers(sheet: Worksheet, months: list[tuple[int, int]]) -> None:
    month_count = len(months)
    for column, label in FIXED_HEADERS:
        sheet.merge_cells(start_row=3, start_column=column, end_row=5, end_column=column)
        cell = sheet.cell(3, column, label)
        cell.alignment = CENTER
        cell.font = BOLD
    for index, (year, month) in enumerate(months):
        column = 7 + index
        sheet.merge_cells(start_row=3, start_column=column, end_row=5, end_column=column)
        cell = sheet.cell(3, column, settlement_header(year, month, month_count))
        cell.alignment = CENTER
        cell.font = BOLD
    tooling_start = 7 + month_count
    sheet.merge_cells(
        start_row=3,
        start_column=tooling_start,
        end_row=3,
        end_column=tooling_start + 3,
    )
    sheet.cell(3, tooling_start, TOOLING_GROUP_HEADER).alignment = CENTER
    sheet.cell(3, tooling_start).font = BOLD
    sheet.merge_cells(
        start_row=4,
        start_column=tooling_start,
        end_row=5,
        end_column=tooling_start,
    )
    sheet.cell(4, tooling_start, PAY_SEPARATELY_HEADER).alignment = CENTER
    sheet.merge_cells(
        start_row=4,
        start_column=tooling_start + 1,
        end_row=4,
        end_column=tooling_start + 3,
    )
    sheet.cell(4, tooling_start + 1, ALLOC_GROUP_HEADER).alignment = CENTER
    sheet.cell(5, tooling_start + 1, ALLOC_PRICE_HEADER).alignment = CENTER
    sheet.cell(5, tooling_start + 2, TOOLING_FEE_HEADER).alignment = CENTER
    sheet.cell(5, tooling_start + 3, ALLOC_QTY_HEADER).alignment = CENTER
    remark_col = tooling_start + 4
    sheet.merge_cells(start_row=3, start_column=remark_col, end_row=5, end_column=remark_col)
    sheet.cell(3, remark_col, REMARK_COL_HEADER).alignment = CENTER
    for column in range(1, remark_col + 1):
        for row in range(3, 6):
            sheet.cell(row, column).border = THIN
            sheet.cell(row, column).alignment = CENTER


def _write_price_data_row(sheet: Worksheet, excel_row: int, row: PriceRow) -> None:
    cells: list[tuple[str | int | float, str | None]] = [
        (row.seq_label, None),
        (row.internal_number or "", None),
        (row.part_number, None),
        (row.part_name or "", None),
        (excel_int(row.vehicle_qty), INT_FORMAT),
        (qty_unit(), None),
    ]
    for amount in row.month_amounts:
        cells.append((excel_price(amount.should_cost), PRICE_FORMAT))
    cells.extend(
        [
            (PAY_SEPARATELY, None),
            (excel_price(row.tooling_alloc_price), PRICE_FORMAT),
            (excel_int(row.tooling_fee), THOUSANDS_INT_FORMAT),
            (excel_int(row.tooling_alloc_qty), INT_FORMAT),
            ("", None),
        ]
    )
    for column, (value, number_format) in enumerate(cells, start=1):
        _write_data_cell(sheet, excel_row, column, value, number_format)
    sheet.row_dimensions[excel_row].height = DATA_ROW_HEIGHT
    name_cell = sheet.cell(excel_row, 4)
    name_cell.alignment = LEFT


def _write_data_cell(
    sheet: Worksheet,
    excel_row: int,
    column: int,
    value: str | int | float,
    number_format: str | None = None,
) -> None:
    written = sheet.cell(excel_row, column, value)
    written.border = THIN
    written.alignment = CENTER
    if number_format is not None and value != MISSING_VALUE:
        written.number_format = number_format


def _write_footer(
    sheet: Worksheet,
    start_row: int,
    last_col: int,
    bundle: ReportBundle,
    company_name: str,
) -> None:
    row = start_row
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    sheet.cell(row, 1, NOTES_LABEL).alignment = LEFT
    row += 1
    sheet.cell(row, 1, REMARK_LABEL)
    row += 1
    for remark in agreement_remarks(bundle, company_name):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        sheet.cell(row, 1, remark).alignment = LEFT
        if "\n" in remark:
            sheet.row_dimensions[row].height = REMARK_LINE_HEIGHT * (remark.count("\n") + 1)
        row += 1
    row += 1
    sheet.cell(row, 1, party_a_label(company_name))
    sheet.cell(row, 8, party_b_label(bundle))
    row += 1
    sheet.cell(row, 1, "代    表：")
    sheet.cell(row, 8, "代   表：")
    row += 2
    sheet.cell(row, 1, party_a_date_label())
    sheet.cell(row, 8, party_b_date_label())


def _apply_column_widths(sheet: Worksheet, widths: tuple[float, ...]) -> None:
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def _tune_agreement_sheet(sheet: Worksheet) -> None:
    sheet.row_dimensions[2].height = 20
    sheet.row_dimensions[3].height = HEADER_ROW_HEIGHT
    sheet.row_dimensions[4].height = HEADER_ROW_HEIGHT
    sheet.row_dimensions[5].height = 26
    sheet.freeze_panes = "A6"
    _set_print_layout(sheet, title_rows="1:5")


def _set_print_layout(sheet: Worksheet, *, title_rows: str) -> None:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    page_setup_pr = sheet.sheet_properties.pageSetUpPr
    if page_setup_pr is not None:
        page_setup_pr.fitToPage = True
    sheet.print_title_rows = title_rows
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = 0.4
    sheet.page_margins.right = 0.4
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.75
    sheet.page_margins.header = 0.3
    sheet.page_margins.footer = 0.28
    _set_page_number_footer(sheet, PAGE_NUMBER_FOOTER)


class _FooterPart(Protocol):
    text: str


class _FooterItem(Protocol):
    center: _FooterPart


def _set_page_number_footer(sheet: Worksheet, text: str) -> None:
    for item in (
        cast(_FooterItem | None, sheet.oddFooter),
        cast(_FooterItem | None, sheet.evenFooter),
    ):
        if item is not None:
            item.center.text = text
