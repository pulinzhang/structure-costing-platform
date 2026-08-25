"""钢材联动价格导出服务（含完整溯源，便于对账扯皮）。"""

from __future__ import annotations

import logging
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.common.audit import audit_status_label
from app.common.excel_workbook import resolve_worksheet
from app.common.export_filters import require_at_least_one_export_filter
from app.common.money import Q4, excel_money_cell
from app.common.timezone import now
from app.repositories.costing.steel_price.steel_price_repository import SteelPriceRepository
from app.services.costing.steel_price.spec_format import (
    build_spec,
    format_thickness,
    format_width,
)

logger = logging.getLogger("service.steel_price_export")

_FONT = "微软雅黑"

HEADERS = [
    "ID",
    "年度",
    "月份",
    "基地",
    "钢厂",
    "钢种",
    "牌号",
    "厚度(mm)",
    "宽度(mm)",
    "规格",
    "基准价",
    "厚度加减价",
    "宽度加减价",
    "锌层加减价",
    "加工费",
    "其他调整",
    "最终价/吨",
    # 基价溯源
    "基价ID",
    "基价文件名",
    "基价表页",
    "基价行号",
    "基价列号",
    # 厚度溯源
    "厚度政策ID",
    "厚度阶梯ID",
    "厚度文件名",
    "厚度表页",
    "厚度行号",
    "厚度列号",
    # 宽度溯源
    "宽度政策ID",
    "宽度阶梯ID",
    "宽度文件名",
    "宽度表页",
    "宽度行号",
    "宽度列号",
    # 锌层溯源
    "锌层政策ID",
    "锌层阶梯ID",
    "锌层文件名",
    "锌层表页",
    "锌层行号",
    "锌层列号",
    # 计算 / 审核
    "匹配状态",
    "计算状态",
    "计算说明",
    "计算批次",
    "公式版本",
    "计算时间",
    "审核状态",
    "审核人",
    "审核时间",
    "审核意见",
    "反审核人",
    "反审核时间",
    "反审核原因",
    "备注",
    # 审计
    "创建人",
    "创建时间",
    "最后修改人",
    "最后修改时间",
]

COL_COUNT = len(HEADERS)
WIDTHS = [
    8, 8, 8, 10, 10, 8, 16, 10, 10, 14,
    12, 12, 12, 12, 10, 10, 12,
    10, 28, 12, 10, 10,
    12, 12, 28, 12, 10, 10,
    12, 12, 28, 12, 10, 10,
    12, 12, 28, 12, 10, 10,
    14, 10, 24, 20, 10, 18,
    10, 10, 18, 16, 10, 18, 16, 20,
    12, 18, 12, 18,
]

# 列区间（1-based）
MONEY_COLS = {11, 12, 13, 14, 15, 16}
KPI_COL = 17
PROVENANCE_COLS = set(range(18, 41))
CALC_COLS = set(range(41, 55))
AUDIT_COLS = set(range(55, 59))
RIGHT_ALIGN_COLS = {
    1, 2, 3, 8, 9, *MONEY_COLS, KPI_COL,
    18, 21, 22, 23, 24, 27, 28, 29, 30, 33, 34, 35, 36, 39, 40,
}
CENTER_COLS = {1, 2, 3, 41, 42, 47}

GROUP_RANGES: list[tuple[int, int, str, str]] = [
    (1, 10, "规格主表", "37474F"),
    (11, 17, "价格构成（最终价 = 基价 ± 加减价 + 加工费 + 其他）", "EF6C00"),
    (18, 22, "基价溯源（扯皮备用）", "90A4AE"),
    (23, 28, "厚度溯源（扯皮备用）", "90A4AE"),
    (29, 34, "宽度溯源（扯皮备用）", "90A4AE"),
    (35, 40, "锌层溯源（扯皮备用）", "90A4AE"),
    (41, 54, "计算 / 审核", "546E7A"),
    (55, 58, "审计字段", "455A64"),
]


def _fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def _font(
    *,
    size: int = 10,
    bold: bool = False,
    color: str = "000000",
) -> Font:
    return Font(name=_FONT, size=size, bold=bold, color=color)


def _thin() -> Border:
    side = Side(style="thin", color="BFBFBF")
    return Border(left=side, right=side, top=side, bottom=side)


def _num(value: Any, quant: Any = Q4, *, allow_negative: bool = True) -> Any:
    return excel_money_cell(value, quant, allow_negative=allow_negative)


def _num_blank_zero(value: Any, quant: Any = Q4, *, allow_negative: bool = True) -> Any:
    """单价为 0 / 空时导出空白，避免表里一堆 0。"""
    if value is None:
        return ""
    cell = excel_money_cell(value, quant, allow_negative=allow_negative)
    if isinstance(cell, Decimal) and cell == 0:
        return ""
    if cell == "" or cell == 0:
        return ""
    return cell


def _dt(value: datetime | None) -> str:
    if value is None:
        return ""
    return f"{value:%Y-%m-%d %H:%M:%S}"


class SteelPriceExportService:
    """钢材联动价格 Excel 导出（核对风排版 + 完整溯源）。"""

    def __init__(self) -> None:
        self.repository = SteelPriceRepository()
        self._center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self._left = Alignment(horizontal="left", vertical="center")
        self._right = Alignment(horizontal="right", vertical="center")
        self._indent = Alignment(horizontal="left", vertical="center", indent=1)

    def export(
        self,
        db: Session,
        *,
        year: int | None = None,
        month: int | None = None,
        location: str | None = None,
        category_name: str | None = None,
        grade: str | None = None,
        steel_mill: str | None = None,
        audit_status: str | None = None,
    ) -> tuple[bytes, str]:
        """
        生成钢材联动价格 Excel。

        溯源列完整导出（基价 / 厚度 / 宽度 / 锌层：ID + 文件名 + 表页 + 行号 + 列号），
        供对账扯皮核对原始 Excel 单元格。

        Returns:
            (文件字节, 文件名)
        """
        require_at_least_one_export_filter(
            year, month, location, category_name, grade, steel_mill, audit_status
        )
        wb = Workbook()
        _title, record_count = self.write_to_workbook(
            wb,
            db,
            year=year,
            month=month,
            location=location,
            category_name=category_name,
            grade=grade,
            steel_mill=steel_mill,
            audit_status=audit_status,
        )
        scope_parts = [
            f"{year}年" if year else None,
            f"{month}月" if month else None,
            location,
            category_name,
            steel_mill,
            grade,
            audit_status_label(audit_status) if audit_status else None,
        ]
        scope = "_".join(p for p in scope_parts if p) or "筛选"
        stamp = now().strftime("%Y%m%d_%H%M%S")
        filename = f"钢材联动价格_{scope}_{stamp}.xlsx"

        output = BytesIO()
        wb.save(output)
        logger.info(
            "导出钢材联动价格 Excel 完成: filename=%s records=%d",
            filename,
            record_count,
        )
        return output.getvalue(), filename

    def write_to_workbook(
        self,
        wb: Workbook,
        db: Session,
        *,
        year: int | None = None,
        month: int | None = None,
        location: str | None = None,
        category_name: str | None = None,
        grade: str | None = None,
        steel_mill: str | None = None,
        audit_status: str | None = None,
        sheet_title: str = "钢材联动价",
    ) -> tuple[str, int]:
        records = self.repository.list_for_export(
            db,
            year=year,
            month=month,
            location=location,
            category_name=category_name,
            grade=grade,
            steel_mill=steel_mill,
            audit_status=audit_status,
        )
        filter_text = self._filter_text(
            year=year,
            month=month,
            location=location,
            category_name=category_name,
            grade=grade,
            steel_mill=steel_mill,
            audit_status=audit_status,
        )

        ws = resolve_worksheet(wb, sheet_title)

        self._write_banner(ws, filter_text=filter_text, record_count=len(records))
        header_row = self._write_headers(ws)

        for seq, record in enumerate(records, start=1):
            row_idx = header_row + seq
            row_data = [
                record.id,
                record.year,
                record.month,
                record.location or "",
                record.steel_mill or "",
                record.category_name or "",
                record.grade or "",
                format_thickness(record.thickness),
                format_width(record.width),
                build_spec(record.thickness, record.width),
                _num_blank_zero(record.base_price, Q4),
                _num_blank_zero(record.thickness_adjustment, Q4, allow_negative=True),
                _num_blank_zero(record.width_adjustment, Q4, allow_negative=True),
                _num_blank_zero(record.zinc_adjustment, Q4, allow_negative=True),
                _num_blank_zero(record.processing_fee, Q4),
                _num_blank_zero(record.other_adjustment, Q4, allow_negative=True),
                _num(record.final_price_per_ton, Q4),
                record.base_price_id or "",
                record.base_excel_filename or "",
                record.base_sheet_name or "",
                record.base_excel_row if record.base_excel_row is not None else "",
                record.base_excel_col if record.base_excel_col is not None else "",
                record.thickness_policy_id or "",
                record.thickness_ladder_item_id or "",
                record.thickness_excel_filename or "",
                record.thickness_sheet_name or "",
                record.thickness_excel_row if record.thickness_excel_row is not None else "",
                record.thickness_excel_col if record.thickness_excel_col is not None else "",
                record.width_policy_id or "",
                record.width_ladder_item_id or "",
                record.width_excel_filename or "",
                record.width_sheet_name or "",
                record.width_excel_row if record.width_excel_row is not None else "",
                record.width_excel_col if record.width_excel_col is not None else "",
                record.zinc_policy_id or "",
                record.zinc_ladder_item_id or "",
                record.zinc_excel_filename or "",
                record.zinc_sheet_name or "",
                record.zinc_excel_row if record.zinc_excel_row is not None else "",
                record.zinc_excel_col if record.zinc_excel_col is not None else "",
                record.match_status or "",
                record.calc_status or "",
                record.calc_message or "",
                record.calc_batch_id or "",
                record.formula_version or "",
                _dt(record.calculated_at),
                audit_status_label(record.audit_status),
                record.audited_by or "",
                _dt(record.audited_at),
                record.audit_remark or "",
                record.reverse_audited_by or "",
                _dt(record.reverse_audited_at),
                record.reverse_audit_remark or "",
                record.remark or "",
                record.created_by or "",
                _dt(record.created_at),
                record.updated_by or "",
                _dt(record.updated_at),
            ]
            self._write_data_row(ws, row_idx, row_data, seq=seq)

        for col, width in enumerate(WIDTHS, start=1):
            dim = ws.column_dimensions[get_column_letter(col)]
            dim.width = width
            if col in PROVENANCE_COLS:
                # 溯源默认折叠，需要时再点大纲展开
                dim.hidden = True
                dim.outline_level = 1

        ws.sheet_properties.outlinePr.summaryRight = True

        last_data = header_row + max(len(records), 1)
        # 冻结到价格区右侧，横向滚动时主表仍可见
        ws.freeze_panes = f"{get_column_letter(KPI_COL + 1)}{header_row + 1}"
        ws.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(COL_COUNT)}{last_data}"
        )
        ws.print_title_rows = f"1:{header_row}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.3
        ws.page_margins.right = 0.3
        ws.page_margins.top = 0.35
        ws.page_margins.bottom = 0.35
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 85

        return ws.title, len(records)

    @staticmethod
    def _filter_text(
        *,
        year: int | None,
        month: int | None,
        location: str | None,
        category_name: str | None,
        grade: str | None,
        steel_mill: str | None,
        audit_status: str | None,
    ) -> str:
        parts: list[str] = []
        if year:
            parts.append(f"{year}年")
        if month:
            parts.append(f"{month}月")
        if location:
            parts.append(location)
        if category_name:
            parts.append(category_name)
        if steel_mill:
            parts.append(steel_mill)
        if grade:
            parts.append(grade)
        if audit_status:
            parts.append(audit_status_label(audit_status))
        return " · ".join(parts) if parts else "全部筛选"

    def _write_banner(
        self,
        ws: Any,
        *,
        filter_text: str,
        record_count: int,
    ) -> None:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COL_COUNT)
        c1 = ws.cell(row=1, column=1, value="钢材联动价格核对表")
        c1.font = _font(size=15, bold=True, color="FFFFFF")
        c1.fill = _fill("1A237E")
        c1.alignment = self._indent
        ws.row_dimensions[1].height = 28

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=COL_COUNT)
        c2 = ws.cell(
            row=2,
            column=1,
            value=(
                f"{filter_text} · 共 {record_count} 条"
                f"　　导出 {now().strftime('%Y-%m-%d %H:%M')}"
            ),
        )
        c2.font = _font(size=9, color="455A64")
        c2.fill = _fill("E8EAF6")
        c2.alignment = self._indent
        ws.row_dimensions[2].height = 20

        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=COL_COUNT)
        c3 = ws.cell(
            row=3,
            column=1,
            value=(
                "说明：红色列为最终价/吨；右侧溯源列已默认折叠，仅对账扯皮时展开即可"
                "（文件名 · 表页 · 行列定位原始单元格）。"
            ),
        )
        c3.font = _font(size=9, color="B71C1C")
        c3.fill = _fill("FFF3E0")
        c3.alignment = self._indent
        ws.row_dimensions[3].height = 20

    def _write_headers(self, ws: Any) -> int:
        group_row = 4
        header_row = 5
        white = _font(size=9, bold=True, color="FFFFFF")
        header_font = _font(size=10, bold=True, color="FFFFFF")

        for start, end, label, color in GROUP_RANGES:
            if end > start:
                ws.merge_cells(
                    start_row=group_row,
                    start_column=start,
                    end_row=group_row,
                    end_column=end,
                )
            cell = ws.cell(row=group_row, column=start, value=label)
            cell.font = white
            cell.fill = _fill(color)
            cell.alignment = self._center
            for col in range(start, end + 1):
                g = ws.cell(row=group_row, column=col)
                g.fill = _fill(color)
                g.border = _thin()
        ws.row_dimensions[group_row].height = 18

        for col, title in enumerate(HEADERS, start=1):
            cell = ws.cell(row=header_row, column=col, value=title)
            cell.alignment = self._center
            cell.border = _thin()
            if col == KPI_COL:
                cell.font = header_font
                cell.fill = _fill("C62828")
            elif col in MONEY_COLS:
                cell.font = header_font
                cell.fill = _fill("EF6C00")
            elif col in PROVENANCE_COLS:
                # 溯源仅扯皮备用：灰底小字，弱化视觉
                cell.font = _font(size=8, bold=True, color="ECEFF1")
                cell.fill = _fill("B0BEC5")
            elif col in AUDIT_COLS:
                cell.font = header_font
                cell.fill = _fill("455A64")
            elif col in CALC_COLS:
                cell.font = header_font
                cell.fill = _fill("607D8B")
            else:
                cell.font = header_font
                cell.fill = _fill("546E7A")
        ws.row_dimensions[header_row].height = 28
        return header_row

    def _write_data_row(
        self,
        ws: Any,
        row_idx: int,
        row_data: list[Any],
        *,
        seq: int,
    ) -> None:
        stripe = "F7F7F7" if seq % 2 == 0 else "FFFFFF"
        money_stripe = "FFF8E1" if seq % 2 == 0 else "FFFDE7"
        kpi_fill = "FFEBEE" if seq % 2 == 0 else "FFCDD2"
        # 溯源：接近白底 + 浅灰字，不抢主表注意力
        prov_stripe = "FAFAFA" if seq % 2 == 0 else "F5F5F5"

        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = _thin()
            if col == KPI_COL:
                cell.fill = _fill(kpi_fill)
                cell.font = _font(size=10, bold=True, color="B71C1C")
                cell.alignment = self._right
            elif col in MONEY_COLS:
                cell.fill = _fill(money_stripe)
                cell.font = _font(size=9)
                cell.alignment = self._right
            elif col in PROVENANCE_COLS:
                cell.fill = _fill(prov_stripe)
                cell.font = _font(size=8, color="90A4AE")
                cell.alignment = (
                    self._right if col in RIGHT_ALIGN_COLS else self._left
                )
            else:
                cell.fill = _fill(stripe)
                cell.font = _font(size=9)
                if col in CENTER_COLS:
                    cell.alignment = self._center
                elif col in RIGHT_ALIGN_COLS:
                    cell.alignment = self._right
                else:
                    cell.alignment = self._left
        ws.row_dimensions[row_idx].height = 20
