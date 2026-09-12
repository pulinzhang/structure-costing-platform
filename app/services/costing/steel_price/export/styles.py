"""钢材联动价导出 — 色板、字体、边框。"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_FONT = "微软雅黑"


def fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def font(
    *,
    size: int = 10,
    bold: bool = False,
    color: str = "000000",
) -> Font:
    return Font(name=_FONT, size=size, bold=bold, color=color)


def thin() -> Border:
    side = Side(style="thin", color="BFBFBF")
    return Border(left=side, right=side, top=side, bottom=side)


@dataclass(frozen=True)
class SheetStyle:
    center: Alignment
    left: Alignment
    right: Alignment
    indent: Alignment


def build_sheet_style() -> SheetStyle:
    return SheetStyle(
        center=Alignment(horizontal="center", vertical="center", wrap_text=True),
        left=Alignment(horizontal="left", vertical="center"),
        right=Alignment(horizontal="right", vertical="center"),
        indent=Alignment(horizontal="left", vertical="center", indent=1),
    )
