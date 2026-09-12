"""钢材联动价导出 — openpyxl 渲染（不访问 DB）。"""

from __future__ import annotations

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.common.excel_workbook import excel_write_value
from app.common.timezone import now
from app.models.costing.steel_price import SteelPrice
from app.services.costing.steel_price.export.columns import (
    AUDIT_COLS,
    CALC_COLS,
    CENTER_COLS,
    COL_COUNT,
    GROUP_RANGES,
    HEADERS,
    KPI_COL,
    MONEY_COLS,
    PROVENANCE_COLS,
    RIGHT_ALIGN_COLS,
    WIDTHS,
)
from app.services.costing.steel_price.export.styles import SheetStyle, fill, font, thin
from app.services.costing.steel_price.export.values import row_values


def write_banner(
    ws: Worksheet,
    *,
    style: SheetStyle,
    filter_text: str,
    record_count: int,
) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COL_COUNT)
    c1 = ws.cell(row=1, column=1, value="钢材联动价格核对表")
    c1.font = font(size=15, bold=True, color="FFFFFF")
    c1.fill = fill("1A237E")
    c1.alignment = style.indent
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=COL_COUNT)
    c2 = ws.cell(
        row=2,
        column=1,
        value=(
            f"{filter_text} · 共 {record_count} 条" f"　　导出 {now().strftime('%Y-%m-%d %H:%M')}"
        ),
    )
    c2.font = font(size=9, color="455A64")
    c2.fill = fill("E8EAF6")
    c2.alignment = style.indent
    ws.row_dimensions[2].height = 20

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=COL_COUNT)
    c3 = ws.cell(
        row=3,
        column=1,
        value=(
            "说明：红色列为最终价/吨；右侧溯源列已默认折叠，仅对账扯皮时展开即可"
            "（文件名 · 表页 · 行列定位原始单元格）。"
        ),
    )
    c3.font = font(size=9, color="B71C1C")
    c3.fill = fill("FFF3E0")
    c3.alignment = style.indent
    ws.row_dimensions[3].height = 20


def write_headers(ws: Worksheet, *, style: SheetStyle) -> int:
    group_row = 4
    header_row = 5
    white = font(size=9, bold=True, color="FFFFFF")
    header_font = font(size=10, bold=True, color="FFFFFF")

    for start, end, label, color in GROUP_RANGES:
        if end > start:
            ws.merge_cells(
                start_row=group_row,
                start_column=start,
                end_row=group_row,
                end_column=end,
            )
        cell = ws.cell(row=group_row, column=start, value=label)
        cell.font = white
        cell.fill = fill(color)
        cell.alignment = style.center
        for col in range(start, end + 1):
            grouped = ws.cell(row=group_row, column=col)
            grouped.fill = fill(color)
            grouped.border = thin()
    ws.row_dimensions[group_row].height = 18

    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.alignment = style.center
        cell.border = thin()
        if col == KPI_COL:
            cell.font = header_font
            cell.fill = fill("C62828")
        elif col in MONEY_COLS:
            cell.font = header_font
            cell.fill = fill("EF6C00")
        elif col in PROVENANCE_COLS:
            cell.font = font(size=8, bold=True, color="ECEFF1")
            cell.fill = fill("B0BEC5")
        elif col in AUDIT_COLS:
            cell.font = header_font
            cell.fill = fill("455A64")
        elif col in CALC_COLS:
            cell.font = header_font
            cell.fill = fill("607D8B")
        else:
            cell.font = header_font
            cell.fill = fill("546E7A")
    ws.row_dimensions[header_row].height = 28
    return header_row


def write_data_row(
    ws: Worksheet,
    row_idx: int,
    data: list[object],
    *,
    style: SheetStyle,
    seq: int,
) -> None:
    stripe = "F7F7F7" if seq % 2 == 0 else "FFFFFF"
    money_stripe = "FFF8E1" if seq % 2 == 0 else "FFFDE7"
    kpi_fill = "FFEBEE" if seq % 2 == 0 else "FFCDD2"
    prov_stripe = "FAFAFA" if seq % 2 == 0 else "F5F5F5"

    for col, value in enumerate(data, start=1):
        cell = ws.cell(row=row_idx, column=col, value=excel_write_value(value))
        cell.border = thin()
        if col == KPI_COL:
            cell.fill = fill(kpi_fill)
            cell.font = font(size=10, bold=True, color="B71C1C")
            cell.alignment = style.right
        elif col in MONEY_COLS:
            cell.fill = fill(money_stripe)
            cell.font = font(size=9)
            cell.alignment = style.right
        elif col in PROVENANCE_COLS:
            cell.fill = fill(prov_stripe)
            cell.font = font(size=8, color="90A4AE")
            cell.alignment = style.right if col in RIGHT_ALIGN_COLS else style.left
        else:
            cell.fill = fill(stripe)
            cell.font = font(size=9)
            if col in CENTER_COLS:
                cell.alignment = style.center
            elif col in RIGHT_ALIGN_COLS:
                cell.alignment = style.right
            else:
                cell.alignment = style.left
    ws.row_dimensions[row_idx].height = 20


def fill_sheet(
    ws: Worksheet,
    records: list[SteelPrice],
    *,
    style: SheetStyle,
    filter_text: str,
) -> None:
    write_banner(ws, style=style, filter_text=filter_text, record_count=len(records))
    header_row = write_headers(ws, style=style)

    for seq, record in enumerate(records, start=1):
        row_idx = header_row + seq
        write_data_row(ws, row_idx, row_values(record), style=style, seq=seq)

    for col, width in enumerate(WIDTHS, start=1):
        dim = ws.column_dimensions[get_column_letter(col)]
        dim.width = width
        if col in PROVENANCE_COLS:
            dim.hidden = True
            dim.outline_level = 1

    outline_pr = ws.sheet_properties.outlinePr
    if outline_pr is not None:
        outline_pr.summaryRight = True

    last_data = header_row + max(len(records), 1)
    ws.freeze_panes = f"{get_column_letter(KPI_COL + 1)}{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(COL_COUNT)}{last_data}"
    ws.print_title_rows = f"1:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
