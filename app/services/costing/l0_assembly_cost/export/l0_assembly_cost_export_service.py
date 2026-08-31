"""成品月核价 Excel 导出（列表所见 + 材料明细 + 焊接费 + 间接费用，核对风排版）。

金额一律写入系统计算快照；明细 Sheet 仅作佐证，不用 Excel 公式回算。
"""

from __future__ import annotations

import logging
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.common.audit import audit_status_label
from app.common.excel_workbook import resolve_worksheet
from app.common.export_filters import require_at_least_one_export_filter
from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.money import Q4, excel_money_cell
from app.common.timezone import now
from app.models.costing.l0_assembly_cost import L0AssemblyCost
from app.repositories.costing.l0_assembly_cost.l0_assembly_cost_repository import (
    L0AssemblyCostRepository,
)
from app.schemas.vo.user_vo import UserVO

logger = logging.getLogger("service.l0_assembly_cost.export")

_FONT = "微软雅黑"
_CALC_LABEL = {
    "ok": "正常",
    "partial": "部分",
    "failed": "失败",
}
_MONEY_FORMAT = "0.0000"

# 列表所见列（与一览表对齐）
LIST_HEADERS = [
    "序号",
    "年度",
    "月份",
    "基地",
    "零件号",
    "零件名称",
    "项目",
    "版本",
    "材料",
    "焊接",
    "月总账",
    "间接",
    "总费用",
    "核算",
    "审核",
]
LIST_WIDTHS = [6, 8, 6, 10, 16, 18, 10, 10, 12, 12, 12, 12, 14, 8, 8]
LIST_MONEY_COLS = {9, 10, 11, 12, 13}  # 材料～总费用
LIST_KPI_COL = 13  # 总费用


def _fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def _font(
    *,
    size: int = 10,
    bold: bool = False,
    color: str = "000000",
) -> Font:
    return Font(name=_FONT, size=size, bold=bold, color=color)


def _num(value: Any) -> Any:
    return excel_money_cell(value, Q4, allow_negative=True)


def _dt(value: datetime | None) -> str:
    if value is None:
        return ""
    return f"{value:%Y-%m-%d %H:%M:%S}"


def _thin() -> Border:
    side = Side(style="thin", color="BFBFBF")
    return Border(left=side, right=side, top=side, bottom=side)


class L0AssemblyCostExportService:
    """按当前筛选导出成品核价（列表所见优先，排版对齐核对表）。"""

    def __init__(self) -> None:
        self.repository = L0AssemblyCostRepository()
        self._center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self._left = Alignment(horizontal="left", vertical="center")
        self._right = Alignment(horizontal="right", vertical="center")
        self._indent = Alignment(horizontal="left", vertical="center", indent=1)

    def export(
        self,
        db: Session,
        *,
        year: int | None = None,
        month: int | None = None,
        location: str | None = None,
        vehicle_model: str | None = None,
        bom_version: str | None = None,
        part_number: str | None = None,
        audit_status: str | None = None,
        calc_status: str | None = None,
        current_user: UserVO,
    ) -> tuple[bytes, str]:
        require_at_least_one_export_filter(
            year,
            month,
            location,
            vehicle_model,
            bom_version,
            part_number,
            audit_status,
            calc_status,
        )
        wb = Workbook()
        _title, record_count = self.write_to_workbook(
            wb,
            db,
            year=year,
            month=month,
            location=location,
            vehicle_model=vehicle_model,
            bom_version=bom_version,
            part_number=part_number,
            audit_status=audit_status,
            calc_status=calc_status,
        )

        scope_parts = [
            f"{year}年" if year else None,
            f"{month}月" if month else None,
            location,
            vehicle_model,
            bom_version,
        ]
        scope = "_".join(p for p in scope_parts if p) or "筛选"
        stamp = now().strftime("%Y%m%d_%H%M%S")
        filename = f"成品核价_{scope}_{stamp}.xlsx"
        content = workbook_to_xlsx_bytes(wb, current_user)
        logger.info(
            "导出成品核价 Excel: filename=%s records=%d",
            filename,
            record_count,
        )
        return content, filename

    def write_to_workbook(
        self,
        wb: Workbook,
        db: Session,
        *,
        year: int | None = None,
        month: int | None = None,
        location: str | None = None,
        vehicle_model: str | None = None,
        bom_version: str | None = None,
        part_number: str | None = None,
        audit_status: str | None = None,
        calc_status: str | None = None,
        list_sheet_title: str = "成品核价一览",
        items_sheet_title: str = "成品材料明细",
        indirects_sheet_title: str = "成品间接费用",
        welding_sheet_title: str = "成品焊接费",
    ) -> tuple[str, int]:
        records = self.repository.list_for_export(
            db,
            year=year,
            month=month,
            location=location,
            vehicle_model=vehicle_model,
            bom_version=bom_version,
            part_number=part_number,
            audit_status=audit_status,
            calc_status=calc_status,
        )
        filter_text = self._filter_text(
            year=year,
            month=month,
            location=location,
            vehicle_model=vehicle_model,
            bom_version=bom_version,
            part_number=part_number,
            audit_status=audit_status,
            calc_status=calc_status,
        )

        list_ws = resolve_worksheet(wb, list_sheet_title)
        self._write_list_sheet(list_ws, records, filter_text=filter_text)
        self._write_items_sheet(
            resolve_worksheet(wb, items_sheet_title),
            records,
            filter_text=filter_text,
        )
        from app.services.bom.welding_fee.export_service import (
            BomWeldingFeeExportService,
        )

        BomWeldingFeeExportService().write_to_workbook(
            wb,
            db,
            year=year,
            location=location,
            vehicle_model=vehicle_model,
            bom_version=bom_version,
            part_number=part_number,
            level=0,
            sheet_title=welding_sheet_title,
        )
        self._write_indirects_sheet(
            resolve_worksheet(wb, indirects_sheet_title),
            records,
            filter_text=filter_text,
        )
        return list_ws.title, len(records)

    @staticmethod
    def _filter_text(
        *,
        year: int | None,
        month: int | None,
        location: str | None,
        vehicle_model: str | None,
        bom_version: str | None,
        part_number: str | None,
        audit_status: str | None,
        calc_status: str | None,
    ) -> str:
        parts: list[str] = []
        if year:
            parts.append(f"{year}年")
        if month:
            parts.append(f"{month}月")
        if location:
            parts.append(location)
        if vehicle_model:
            parts.append(vehicle_model)
        if bom_version:
            parts.append(f"版本 {bom_version}")
        if part_number:
            parts.append(f"零件 {part_number}")
        if audit_status:
            parts.append(audit_status_label(audit_status))
        if calc_status:
            parts.append(_CALC_LABEL.get(calc_status, calc_status))
        return " · ".join(parts) if parts else "全部筛选"

    def _write_banner(
        self,
        ws: Any,
        *,
        title: str,
        filter_text: str,
        count_label: str,
        note: str,
        col_count: int,
    ) -> int:
        """返回表头起始行（数据从 header_row+1 起）。"""
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
        c1 = ws.cell(row=1, column=1, value=title)
        c1.font = _font(size=15, bold=True, color="FFFFFF")
        c1.fill = _fill("1A237E")
        c1.alignment = self._indent
        ws.row_dimensions[1].height = 28

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
        c2 = ws.cell(
            row=2,
            column=1,
            value=(f"{filter_text} · {count_label}" f"　　导出 {now().strftime('%Y-%m-%d %H:%M')}"),
        )
        c2.font = _font(size=9, color="455A64")
        c2.fill = _fill("E8EAF6")
        c2.alignment = self._indent
        ws.row_dimensions[2].height = 20

        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=col_count)
        c3 = ws.cell(row=3, column=1, value=note)
        c3.font = _font(size=9, color="B71C1C")
        c3.fill = _fill("FFF3E0")
        c3.alignment = self._indent
        ws.row_dimensions[3].height = 20
        return 4

    def _write_list_sheet(
        self,
        ws: Any,
        records: list[L0AssemblyCost],
        *,
        filter_text: str,
    ) -> None:
        ws.title = "成品核价一览"
        col_count = len(LIST_HEADERS)
        header_row = self._write_banner(
            ws,
            title="成品核价一览表",
            filter_text=filter_text,
            count_label=f"共 {len(records)} 条",
            note=(
                "说明：金额为系统核价快照（材料 / 焊接 / 月总账 / 间接 / 总费用）；"
                "明细 Sheet 仅作佐证，不以 Excel 公式回算。"
            ),
            col_count=col_count,
        )

        for col, title in enumerate(LIST_HEADERS, start=1):
            cell = ws.cell(row=header_row, column=col, value=title)
            cell.font = _font(size=10, bold=True, color="FFFFFF")
            cell.alignment = self._center
            cell.border = _thin()
            if col == LIST_KPI_COL:
                cell.fill = _fill("C62828")
            elif col in LIST_MONEY_COLS:
                cell.fill = _fill("EF6C00")
            else:
                cell.fill = _fill("546E7A")
        ws.row_dimensions[header_row].height = 22

        for seq, record in enumerate(records, start=1):
            row_idx = header_row + seq
            row_data = [
                seq,
                record.year,
                record.month,
                record.location or "",
                record.part_number or "",
                record.part_name or "",
                record.vehicle_model or "",
                record.bom_version or "",
                _num(record.material_amount),
                _num(record.welding_amount),
                _num(record.total_amount),
                _num(record.indirect_amount),
                _num(record.grand_total),
                _CALC_LABEL.get(record.calc_status, record.calc_status or ""),
                audit_status_label(record.audit_status),
            ]
            stripe = "F7F7F7" if seq % 2 == 0 else "FFFFFF"
            kpi_fill = "FFEBEE" if seq % 2 == 0 else "FFCDD2"
            for col, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.font = _font(size=9)
                cell.border = _thin()
                if col == LIST_KPI_COL:
                    cell.fill = _fill(kpi_fill)
                    cell.font = _font(size=10, bold=True, color="B71C1C")
                    cell.alignment = self._right
                    cell.number_format = _MONEY_FORMAT
                elif col in LIST_MONEY_COLS:
                    cell.fill = _fill(stripe)
                    cell.alignment = self._right
                    cell.number_format = _MONEY_FORMAT
                elif col in {1, 2, 3, 14, 15}:
                    cell.fill = _fill(stripe)
                    cell.alignment = self._center
                else:
                    cell.fill = _fill(stripe)
                    cell.alignment = self._left
            ws.row_dimensions[row_idx].height = 20

        for col, width in enumerate(LIST_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(col_count)}" f"{header_row + max(len(records), 1)}"
        )

    def _write_items_sheet(
        self,
        ws: Any,
        records: list[L0AssemblyCost],
        *,
        filter_text: str,
    ) -> None:
        headers = [
            "序号",
            "年度",
            "月份",
            "基地",
            "项目",
            "版本",
            "L0零件号",
            "层级",
            "子件零件号",
            "子件名称",
            "取价来源",
            "匹配说明",
            "单价",
            "有效用量",
            "行金额",
        ]
        header_row = self._write_banner(
            ws,
            title="成品核价 · 材料明细",
            filter_text=filter_text,
            count_label=f"共 {len(records)} 个成品",
            note="说明：含一口价剪枝行（skipped_by_bundle）；行金额为计算快照。",
            col_count=len(headers),
        )
        self._write_simple_header(ws, header_row, headers, money_cols={13, 14, 15})

        row_idx = header_row + 1
        seq = 0
        for record in records:
            for item in record.items or []:
                seq += 1
                row_data = [
                    seq,
                    record.year,
                    record.month,
                    record.location or "",
                    record.vehicle_model or "",
                    record.bom_version or "",
                    record.part_number or "",
                    item.level_label or "",
                    item.part_number or "",
                    item.part_name or "",
                    item.cost_source or "",
                    item.match_hint or "",
                    _num(item.unit_price),
                    _num(item.effective_qty),
                    _num(item.line_amount),
                ]
                self._append_styled_row(ws, row_idx, row_data, money_cols={13, 14, 15}, seq=seq)
                row_idx += 1

        widths = [6, 8, 6, 10, 10, 10, 16, 10, 16, 18, 14, 18, 12, 10, 12]
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = f"A{header_row + 1}"

    def _write_indirects_sheet(
        self,
        ws: Any,
        records: list[L0AssemblyCost],
        *,
        filter_text: str,
    ) -> None:
        headers = [
            "序号",
            "年度",
            "月份",
            "基地",
            "项目",
            "版本",
            "L0零件号",
            "代码",
            "名称",
            "金额",
            "费率",
            "基数",
        ]
        header_row = self._write_banner(
            ws,
            title="成品核价 · 间接费用",
            filter_text=filter_text,
            count_label=f"共 {len(records)} 个成品",
            note="说明：I1～I8 一一列出；费率类以月总账为基数（系统快照）。",
            col_count=len(headers),
        )
        self._write_simple_header(ws, header_row, headers, money_cols={10, 11, 12})

        row_idx = header_row + 1
        seq = 0
        for record in records:
            for ind in record.indirects or []:
                seq += 1
                row_data = [
                    seq,
                    record.year,
                    record.month,
                    record.location or "",
                    record.vehicle_model or "",
                    record.bom_version or "",
                    record.part_number or "",
                    ind.code or "",
                    ind.label or "",
                    _num(ind.amount),
                    _num(ind.rate) if ind.rate is not None else "",
                    _num(ind.base_amount) if ind.base_amount is not None else "",
                ]
                self._append_styled_row(ws, row_idx, row_data, money_cols={10, 11, 12}, seq=seq)
                row_idx += 1

        widths = [6, 8, 6, 10, 10, 10, 16, 14, 20, 12, 10, 12]
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = f"A{header_row + 1}"

    def _write_simple_header(
        self,
        ws: Any,
        header_row: int,
        headers: list[str],
        *,
        money_cols: set[int],
    ) -> None:
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col, value=title)
            cell.font = _font(size=10, bold=True, color="FFFFFF")
            cell.alignment = self._center
            cell.border = _thin()
            cell.fill = _fill("EF6C00") if col in money_cols else _fill("546E7A")
        ws.row_dimensions[header_row].height = 22

    def _append_styled_row(
        self,
        ws: Any,
        row_idx: int,
        row_data: list[Any],
        *,
        money_cols: set[int],
        seq: int,
    ) -> None:
        stripe = "F7F7F7" if seq % 2 == 0 else "FFFFFF"
        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.font = _font(size=9)
            cell.border = _thin()
            cell.fill = _fill(stripe)
            if col in money_cols:
                cell.alignment = self._right
                if value != "":
                    cell.number_format = _MONEY_FORMAT
            elif col == 1:
                cell.alignment = self._center
            else:
                cell.alignment = self._left
        ws.row_dimensions[row_idx].height = 20
