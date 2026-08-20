"""成品月核价 Excel 导出（主表 + 材料明细 + 间接费用）。"""

from __future__ import annotations

import logging
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from app.common.audit import audit_status_label
from app.common.money import Q4, excel_money_cell
from app.models.costing.l0_assembly_cost import L0AssemblyCost
from app.repositories.costing.l0_assembly_cost.l0_assembly_cost_repository import (
    L0AssemblyCostRepository,
)

logger = logging.getLogger("service.l0_assembly_cost.export")

_CALC_LABEL = {
    "ok": "正常",
    "partial": "部分",
    "failed": "失败",
}


def _num(value: Any) -> Any:
    return excel_money_cell(value, Q4, allow_negative=True)


def _dt(value: datetime | None) -> str:
    if value is None:
        return ""
    return value.strftime("%Y-%m-%d %H:%M:%S")


class L0AssemblyCostExportService:
    """按当前筛选导出成品核价结果。"""

    def __init__(self) -> None:
        self.repository = L0AssemblyCostRepository()

    def export(
        self,
        db: Session,
        *,
        year: int | None = None,
        month: int | None = None,
        location: str | None = None,
        vehicle_model: str | None = None,
        bom_version: str | None = None,
        part_number: str | None = None,
        audit_status: str | None = None,
        calc_status: str | None = None,
    ) -> tuple[bytes, str]:
        records = self.repository.list_for_export(
            db,
            year=year,
            month=month,
            location=location,
            vehicle_model=vehicle_model,
            bom_version=bom_version,
            part_number=part_number,
            audit_status=audit_status,
            calc_status=calc_status,
        )

        wb = Workbook()
        self._write_main_sheet(wb.active, records)
        self._write_items_sheet(wb.create_sheet("材料明细"), records)
        self._write_indirects_sheet(wb.create_sheet("间接费用"), records)

        parts = ["成品核价导出"]
        if year:
            parts.append(f"{year}年")
        if month:
            parts.append(f"{month}月")
        if location:
            parts.append(location)
        if vehicle_model:
            parts.append(vehicle_model)
        if bom_version:
            parts.append(bom_version)
        if audit_status:
            parts.append(audit_status_label(audit_status))
        if calc_status:
            parts.append(_CALC_LABEL.get(calc_status, calc_status))
        filename = "_".join(parts) + ".xlsx"

        output = BytesIO()
        wb.save(output)
        logger.info(
            "导出成品核价 Excel 完成: filename=%s records=%d",
            filename,
            len(records),
        )
        return output.getvalue(), filename

    def _write_main_sheet(
        self, ws: Any, records: list[L0AssemblyCost]
    ) -> None:
        ws.title = "成品主表"
        headers = [
            "ID",
            "年度",
            "月份",
            "基地",
            "项目",
            "版本",
            "零件号",
            "零件名称",
            "单车用量",
            "冲压小计",
            "外购小计",
            "一口价小计",
            "材料合计",
            "焊接费",
            "焊接其它",
            "月总账",
            "运输费",
            "仓储打包",
            "财务费用",
            "管理费用",
            "废品损失",
            "税前利润",
            "检具分摊",
            "夹具分摊",
            "分摊合计",
            "间接合计",
            "总费用",
            "财务费率",
            "管理费率",
            "废品费率",
            "税前利润率",
            "核算状态",
            "缺叶数",
            "计算说明",
            "公式版本",
            "计算时间",
            "审核状态",
            "审核人",
            "审核时间",
            "创建人",
            "创建时间",
            "修改人",
            "修改时间",
        ]
        self._style_header(ws, headers)
        right_cols = set(range(9, 32)) | {1, 2, 3, 33}

        for row_idx, record in enumerate(records, start=2):
            row_data = [
                record.id,
                record.year,
                record.month,
                record.location or "",
                record.vehicle_model or "",
                record.bom_version or "",
                record.part_number or "",
                record.part_name or "",
                _num(record.supply_qty),
                _num(record.stamping_amount),
                _num(record.purchased_amount),
                _num(record.bundle_amount),
                _num(record.material_amount),
                _num(record.welding_amount),
                _num(record.welding_other_fee),
                _num(record.total_amount),
                _num(record.unit_transport_fee),
                _num(record.warehouse_bundling_fee),
                _num(record.finished_finance_fee),
                _num(record.finished_management_fee),
                _num(record.finished_scrap_fee),
                _num(record.finished_pretax_profit_fee),
                _num(record.gauge_alloc_fee),
                _num(record.fixture_alloc_fee),
                _num(record.allocation_fee_total),
                _num(record.indirect_amount),
                _num(record.grand_total),
                _num(record.finance_rate),
                _num(record.management_rate),
                _num(record.scrap_rate),
                _num(record.pretax_profit_rate),
                _CALC_LABEL.get(record.calc_status, record.calc_status or ""),
                record.missing_leaf_count or 0,
                record.calc_message or "",
                record.formula_version or "",
                _dt(record.calculated_at),
                audit_status_label(record.audit_status),
                record.audited_by or "",
                _dt(record.audited_at),
                record.created_by or "",
                _dt(record.created_at),
                record.updated_by or "",
                _dt(record.updated_at),
            ]
            self._append_data_row(ws, row_idx, row_data, right_cols)

        widths = [
            8, 8, 6, 10, 10, 10, 16, 18, 10,
            12, 12, 12, 12, 12, 12, 12,
            10, 10, 10, 10, 10, 10, 10, 10, 10, 12, 12,
            10, 10, 10, 10,
            10, 8, 24, 10, 18,
            10, 10, 18, 10, 18, 10, 18,
        ]
        self._apply_widths(ws, widths)
        ws.freeze_panes = "A2"

    def _write_items_sheet(
        self, ws: Any, records: list[L0AssemblyCost]
    ) -> None:
        headers = [
            "主表ID",
            "年度",
            "月份",
            "基地",
            "项目",
            "版本",
            "L0零件号",
            "层级",
            "子件零件号",
            "子件名称",
            "节点角色",
            "供货类型",
            "取价来源",
            "匹配说明",
            "单价",
            "有效用量",
            "行金额",
            "排序",
        ]
        self._style_header(ws, headers)
        right_cols = {15, 16, 17, 18}
        row_idx = 2
        for record in records:
            for item in record.items or []:
                row_data = [
                    record.id,
                    record.year,
                    record.month,
                    record.location or "",
                    record.vehicle_model or "",
                    record.bom_version or "",
                    record.part_number or "",
                    item.level_label or "",
                    item.part_number or "",
                    item.part_name or "",
                    item.node_role or "",
                    item.supply_type or "",
                    item.cost_source or "",
                    item.match_hint or "",
                    _num(item.unit_price),
                    _num(item.effective_qty),
                    _num(item.line_amount),
                    item.sort_no,
                ]
                self._append_data_row(ws, row_idx, row_data, right_cols)
                row_idx += 1
        widths = [
            8, 8, 6, 10, 10, 10, 16, 8, 16, 18,
            10, 10, 14, 20, 12, 10, 12, 8,
        ]
        self._apply_widths(ws, widths)
        ws.freeze_panes = "A2"

    def _write_indirects_sheet(
        self, ws: Any, records: list[L0AssemblyCost]
    ) -> None:
        headers = [
            "主表ID",
            "年度",
            "月份",
            "基地",
            "项目",
            "版本",
            "L0零件号",
            "代码",
            "名称",
            "金额",
            "费率",
            "基数",
            "排序",
        ]
        self._style_header(ws, headers)
        right_cols = {10, 11, 12, 13}
        row_idx = 2
        for record in records:
            for ind in record.indirects or []:
                row_data = [
                    record.id,
                    record.year,
                    record.month,
                    record.location or "",
                    record.vehicle_model or "",
                    record.bom_version or "",
                    record.part_number or "",
                    ind.code or "",
                    ind.label or "",
                    _num(ind.amount),
                    _num(ind.rate) if ind.rate is not None else "",
                    _num(ind.base_amount) if ind.base_amount is not None else "",
                    ind.sort_no,
                ]
                self._append_data_row(ws, row_idx, row_data, right_cols)
                row_idx += 1
        widths = [8, 8, 6, 10, 10, 10, 16, 12, 18, 12, 10, 12, 8]
        self._apply_widths(ws, widths)
        ws.freeze_panes = "A2"

    @staticmethod
    def _style_header(ws: Any, headers: list[str]) -> None:
        header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        header_alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )
        ws.append(headers)
        ws.row_dimensions[1].height = 28
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

    @staticmethod
    def _append_data_row(
        ws: Any,
        row_idx: int,
        row_data: list[Any],
        right_cols: set[int],
    ) -> None:
        data_font = Font(name="微软雅黑", size=9, color="000000")
        left = Alignment(horizontal="left", vertical="center")
        right = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )
        fill = PatternFill(
            start_color="F2F2F2" if row_idx % 2 == 0 else "FFFFFF",
            end_color="F2F2F2" if row_idx % 2 == 0 else "FFFFFF",
            fill_type="solid",
        )
        ws.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.fill = fill
            cell.alignment = right if col_idx in right_cols else left

    @staticmethod
    def _apply_widths(ws: Any, widths: list[int]) -> None:
        for col_idx, width in enumerate(widths, start=1):
            letter = ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[letter].width = width
