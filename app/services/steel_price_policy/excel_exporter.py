"""
钢价政策 Excel 导出器
职责：纯 Excel 生成逻辑，不涉及数据库操作
"""

from io import BytesIO
from decimal import Decimal
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.repositories.steel_price_policy.dto import ExportRow


# =============================================================================
# 样式常量
# =============================================================================

HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")

DATA_FONT = Font(name="微软雅黑", size=10, color="000000")
DATA_ALIGNMENT = Alignment(horizontal="left", vertical="center")
DATA_ALIGNMENT_CENTER = Alignment(horizontal="center", vertical="center")
DATA_ALIGNMENT_RIGHT = Alignment(horizontal="right", vertical="center")

THIN_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)

EVEN_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
ODD_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

HEADERS = [
    "政策ID", "年度", "月份", "钢厂", "钢种大类", "政策类型",
    "牌号类型", "牌号/范围", "基准厚度",
    "明细ID", "子分组", "厚度条件", "宽度条件",
    "加减价(元/吨)", "基准", "排序",
]

COLUMN_WIDTHS = [10, 8, 8, 12, 12, 12, 10, 15, 10, 10, 15, 18, 18, 14, 8, 8]


# =============================================================================
# 格式化辅助函数
# =============================================================================

def format_decimal(value, digits=2) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        value = float(value)
    if isinstance(value, (int, float)):
        return f"{value:.{digits}f}"
    return str(value)


def format_thick_condition(row: ExportRow) -> str:
    """
    格式化厚度条件 - 直接显示数据库存储值。
    例如：thick_min=0.5, thick_min_operator='<' → 显示 "<0.50"
    """
    parts = []

    if row.thick_min is not None and row.thick_min_operator:
        parts.append(f"{row.thick_min_operator}{format_decimal(row.thick_min)}")

    if row.thick_max is not None and row.thick_max_operator:
        parts.append(f"{row.thick_max_operator}{format_decimal(row.thick_max)}")

    return ", ".join(parts) if parts else ""


def format_width_condition(row: ExportRow) -> str:
    """
    格式化宽度条件 - 直接显示数据库存储值。
    """
    parts = []

    if row.width_min is not None and row.width_min_operator:
        parts.append(f"{row.width_min_operator}{format_decimal(row.width_min)}")

    if row.width_max is not None and row.width_max_operator:
        parts.append(f"{row.width_max_operator}{format_decimal(row.width_max)}")

    return ", ".join(parts) if parts else ""


# =============================================================================
# 导出器类
# =============================================================================

class SteelPricePolicyExporter:
    """钢价政策 Excel 导出器"""

    def generate(
        self,
        records: list[ExportRow],
        year: Optional[int] = None,
        month: Optional[int] = None,
        steel_mill: Optional[str] = None,
    ) -> tuple[bytes, str]:
        """
        生成 Excel 文件字节数据。
        返回 (文件字节数据, 文件名)
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "钢价政策导出"

        self._write_headers(ws)
        self._write_rows(ws, records)
        self._set_column_widths(ws)
        ws.freeze_panes = "A2"

        filename = self._build_filename(year, month, steel_mill)
        output = BytesIO()
        wb.save(output)
        return output.getvalue(), filename

    def _write_headers(self, ws) -> None:
        """写入表头行"""
        ws.append(HEADERS)
        for col_idx, _ in enumerate(HEADERS, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
            cell.border = THIN_BORDER

    def _write_rows(self, ws, records: list[ExportRow]) -> None:
        """写入数据行"""
        for row_idx, row in enumerate(records, start=2):
            row_data = [
                row.policy_id or "",
                row.year or "",
                row.month or "",
                row.steel_mill or "",
                row.category_name or "",
                row.policy_type or "",
                row.grade_type or "",
                row.grade_value or "",
                format_decimal(row.base_thickness),
                row.item_id or "",
                row.target_sub_group or "",
                format_thick_condition(row),
                format_width_condition(row),
                row.price_adjustment if row.price_adjustment is not None else "",
                "是" if row.is_base == 1 else "",
                row.sort_order or 0,
            ]
            ws.append(row_data)

            fill = EVEN_FILL if row_idx % 2 == 0 else ODD_FILL
            for col_idx, _ in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                cell.fill = fill
                if col_idx in (1, 2, 3, 10, 15, 16):
                    cell.alignment = DATA_ALIGNMENT_RIGHT
                elif col_idx in (14,):
                    cell.alignment = DATA_ALIGNMENT_CENTER
                else:
                    cell.alignment = DATA_ALIGNMENT

    def _set_column_widths(self, ws) -> None:
        """设置列宽"""
        for col_idx, width in enumerate(COLUMN_WIDTHS, start=1):
            letter = ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[letter].width = width

    def _build_filename(
        self,
        year: Optional[int],
        month: Optional[int],
        steel_mill: Optional[str],
    ) -> str:
        """构建导出文件名"""
        parts = ["钢价政策导出"]
        if year:
            parts.append(f"{year}年")
        if month:
            parts.append(f"{month}月")
        if steel_mill:
            parts.append(steel_mill)
        return f"{'_'.join(parts)}.xlsx"
