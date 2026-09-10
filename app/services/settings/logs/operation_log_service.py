"""操作日志 Service。"""

from __future__ import annotations

import contextlib
import logging
from datetime import datetime
from typing import Any, TypedDict
from urllib.parse import unquote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.common.audit import display_name
from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.operation_log.context import OperationAuditScope
from app.common.operation_log.ip_location import lookup_ip_location
from app.common.operation_log.rules import (
    append_upload_filenames,
    describe_operation,
    format_audit_question,
    format_restore_scope,
)
from app.common.permissions.catalog import (
    CatalogFilterOption,
    catalog_subtree_for_menu,
    list_catalog_filter_options,
)
from app.common.timezone import now
from app.database import DatabaseConfig
from app.repositories.settings.logs.operation_log_repository import (
    OperationLogRepository,
)
from app.schemas.dto.operation_log_dto import OperationLogCreate, OperationLogQuery
from app.schemas.vo.operation_log_vo import (
    OperationLogCatalogOptionVO,
    OperationLogFilterOptionsVO,
    OperationLogVO,
)
from app.schemas.vo.user_vo import UserVO
from app.services.async_task.async_task_service import AsyncTaskService

logger = logging.getLogger("service.operation_log")

LOGIN_FILTER_CODE = "login:session"
_SENSITIVE_PARAM_PREFIXES = ("token=", "password=", "authorization=", "_t=")
_SCOPE_QUERY_KEYS = frozenset(
    {"location", "part_number", "year", "month", "view", "vehicle_model", "bom_version"}
)


class CatalogQueryFilter(TypedDict):
    login_only: bool
    permission_code: str | None
    permission_codes: list[str] | None
    menu_labels: list[str] | None


def _build_detail(
    query_string: str,
    path_params: dict[str, str] | None,
    audit_scope: OperationAuditScope | None = None,
    upload_filenames: list[str] | None = None,
) -> str | None:
    """文件名、提问、业务口径；不把缓存戳和重复查询串堆进对象列。"""
    parts: list[str] = []
    if upload_filenames:
        parts.extend(upload_filenames)
    question_text = format_audit_question(audit_scope or {})
    if question_text:
        parts.append(question_text)
    scope_text = format_restore_scope(audit_scope or {})
    if scope_text:
        parts.append(scope_text)
    if path_params:
        for key, value in path_params.items():
            lowered = key.lower()
            if lowered in {"token", "password", "authorization"}:
                continue
            parts.append(f"{key}={value}")
    if query_string and not scope_text and not upload_filenames and not question_text:
        raw = unquote(query_string)
        for item in raw.split("&"):
            if not item or item.lower().startswith(_SENSITIVE_PARAM_PREFIXES):
                continue
            key = item.split("=", 1)[0].strip().lower()
            if key in _SCOPE_QUERY_KEYS:
                continue
            parts.append(item)
    joined = " · ".join(parts)[:500]
    return joined or None


def _schedule_ip_lookup(log_id: int, client_ip: str | None) -> None:
    if not log_id or not client_ip:
        return
    AsyncTaskService.submit(enrich_operation_log_ip, log_id, client_ip)


def enrich_operation_log_ip(log_id: object, client_ip: object) -> None:
    """后台按 id + ip 回写国家/省/市；失败忽略。"""
    if not isinstance(log_id, int) or not isinstance(client_ip, str) or not client_ip:
        return
    session_factory = DatabaseConfig.get_session_factory()
    db = session_factory()
    repository = OperationLogRepository()
    try:
        location = repository.get_ip_location(db, client_ip)
        if location is None:
            location = lookup_ip_location(client_ip)
        if location is None:
            return
        repository.update_ip_location(
            db,
            log_id=log_id,
            client_ip=client_ip,
            country=location["country"],
            province=location["province"],
            city=location["city"],
        )
        db.commit()
    except Exception:
        logger.debug("操作日志 IP 归属地回写失败: id=%s ip=%s", log_id, client_ip, exc_info=True)
        with contextlib.suppress(Exception):
            db.rollback()
    finally:
        db.close()


class OperationLogService:
    """操作审计日志业务。"""

    def __init__(self) -> None:
        self.repository = OperationLogRepository()

    def record_from_request(
        self,
        *,
        user: UserVO | None,
        method: str,
        path: str,
        query_string: str,
        status_code: int,
        client_ip: str | None,
        permission_code: str | None = None,
        path_params: dict[str, str] | None = None,
        audit_scope: OperationAuditScope | None = None,
        upload_filenames: list[str] | None = None,
    ) -> None:
        """中间件调用：独立 Session 写入，失败不影响主请求。"""
        session_factory = DatabaseConfig.get_session_factory()
        db = session_factory()
        try:
            module, action, summary = describe_operation(
                method,
                path,
                permission_code,
                query_string=query_string,
                audit_scope=audit_scope,
            )
            summary = append_upload_filenames(summary, upload_filenames)
            detail = _build_detail(
                query_string,
                path_params,
                audit_scope,
                upload_filenames=upload_filenames,
            )

            row = self.repository.create(
                db,
                OperationLogCreate(
                    user_id=user.id if user else None,
                    username=user.username if user else None,
                    operator_name=display_name(user) if user else None,
                    method=method.upper(),
                    path=path.split("?", 1)[0][:255],
                    module=module,
                    action=action[:30],
                    permission_code=(permission_code or "")[:80] or None,
                    summary=summary[:255],
                    status_code=status_code,
                    success=status_code < 400,
                    client_ip=(client_ip or "")[:64] or None,
                    detail=detail,
                ),
            )
            db.commit()
            with contextlib.suppress(Exception):
                _schedule_ip_lookup(row.id, row.client_ip)
        except Exception:
            logger.exception("写入操作日志失败: %s %s", method, path)
            with contextlib.suppress(Exception):
                db.rollback()
        finally:
            db.close()

    def paginate(
        self,
        db: Session,
        query: OperationLogQuery,
    ) -> tuple[list[OperationLogVO], int, OperationLogFilterOptionsVO]:
        skip = (query.page - 1) * query.size
        catalog = _resolve_catalog_filter(query.menu_code, query.permission_code)
        rows, total = self.repository.paginate(
            db,
            username=query.username,
            permission_code=catalog["permission_code"],
            permission_codes=catalog["permission_codes"],
            menu_labels=catalog["menu_labels"],
            login_only=catalog["login_only"],
            module=query.module,
            action=query.action,
            method=query.method,
            success=query.success,
            created_from=query.created_from,
            created_to=query.created_to,
            skip=skip,
            limit=query.size,
        )
        options = _catalog_filter_options_vo()
        return [OperationLogVO.model_validate(r) for r in rows], total, options

    def export(
        self,
        db: Session,
        *,
        username: str | None = None,
        menu_code: str | None = None,
        permission_code: str | None = None,
        module: str | None = None,
        action: str | None = None,
        method: str | None = None,
        success: bool | None = None,
        created_from: datetime | None = None,
        created_to: datetime | None = None,
        current_user: UserVO,
    ) -> tuple[bytes, str]:
        catalog = _resolve_catalog_filter(menu_code, permission_code)
        rows = self.repository.list_for_export(
            db,
            username=username,
            permission_code=catalog["permission_code"],
            permission_codes=catalog["permission_codes"],
            menu_labels=catalog["menu_labels"],
            login_only=catalog["login_only"],
            module=module,
            action=action,
            method=method,
            success=success,
            created_from=created_from,
            created_to=created_to,
        )
        wb = Workbook()
        ws = wb.worksheets[0]
        ws.title = "操作日志"
        headers = [
            "时间",
            "操作人",
            "账号",
            "操作",
            "对象",
            "结果",
            "状态码",
            "IP",
            "国家",
            "省",
            "城市",
            "权限码",
            "方法",
            "路径",
        ]
        header_fill = PatternFill("solid", fgColor="1E1E1E")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(1, col, title)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row_idx, row in enumerate(rows, start=2):
            created = row.created_at.strftime("%Y-%m-%d %H:%M:%S") if row.created_at else ""
            values: list[Any] = [
                created,
                row.operator_name or "",
                row.username or "",
                row.summary or "",
                row.detail or "",
                "成功" if row.success else "失败",
                row.status_code,
                row.client_ip or "",
                row.ip_country or "",
                row.ip_province or "",
                row.ip_city or "",
                row.permission_code or "",
                row.method,
                row.path,
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row_idx, col, value)

        for col in ws.columns:
            max_len = 0
            letter = col[0].column_letter
            for cell in col[:200]:
                max_len = max(max_len, len(str(cell.value or "")))
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 48)

        content = workbook_to_xlsx_bytes(wb, current_user)
        stamp = now().strftime("%Y%m%d_%H%M%S")
        return content, f"操作日志_{stamp}.xlsx"


def _catalog_filter_options_vo() -> OperationLogFilterOptionsVO:
    menus, operations = list_catalog_filter_options()
    login_menu: CatalogFilterOption = {
        "code": LOGIN_FILTER_CODE,
        "label": "登录",
        "parent_code": None,
        "breadcrumb": "登录认证",
    }
    login_op: CatalogFilterOption = {
        "code": LOGIN_FILTER_CODE,
        "label": "登录",
        "parent_code": LOGIN_FILTER_CODE,
        "breadcrumb": "登录认证 / 登录",
    }
    return OperationLogFilterOptionsVO(
        menus=[OperationLogCatalogOptionVO.model_validate(login_menu)]
        + [OperationLogCatalogOptionVO.model_validate(item) for item in menus],
        operations=[OperationLogCatalogOptionVO.model_validate(login_op)]
        + [OperationLogCatalogOptionVO.model_validate(item) for item in operations],
    )


def _resolve_catalog_filter(
    menu_code: str | None,
    permission_code: str | None,
) -> CatalogQueryFilter:
    empty: CatalogQueryFilter = {
        "login_only": False,
        "permission_code": None,
        "permission_codes": None,
        "menu_labels": None,
    }
    perm = (permission_code or "").strip() or None
    menu = (menu_code or "").strip() or None
    if perm == LOGIN_FILTER_CODE or menu == LOGIN_FILTER_CODE:
        return {
            "login_only": True,
            "permission_code": None,
            "permission_codes": None,
            "menu_labels": None,
        }
    if perm:
        return {
            "login_only": False,
            "permission_code": perm,
            "permission_codes": None,
            "menu_labels": None,
        }
    if not menu:
        return empty
    op_codes, menu_labels = catalog_subtree_for_menu(menu)
    return {
        "login_only": False,
        "permission_code": None,
        "permission_codes": op_codes,
        "menu_labels": menu_labels,
    }
