"""
L0 核对导出 — openpyxl 工作簿写入（不访问数据库）
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter

from app.common.timezone import now
from app.models.bom_welding_process_breakdown import (
    BomWeldingProcessBreakdownLine,
    BomWeldingProcessBreakdownNode,
)
from app.services.bom.welding_process_breakdown.export.columns import (
    BILLING_COLS,
    CENTER_COLS,
    COL,
    COL_COUNT,
    DATA_H,
    FILTER_H,
    GROUP_DETAIL,
    GROUP_H,
    GROUP_MAIN,
    HEADER_H,
    HEADERS,
    MAIN_COLS,
    MAP_LABEL,
    MUTED_QTY_COLS,
    NOTE,
    NOTE_H,
    PARSE_LABEL,
    QTY_COLS,
    REPORT_TITLE,
    SHEET_TITLE,
    STD_COLS,
    TITLE_H,
    WIDTHS,
)
from app.services.bom.welding_process_breakdown.export.styles import (
    MAP_TONE,
    PARSE_TONE,
    SheetStyle,
    build_sheet_style,
    fill,
    font,
    tone,
)
from app.services.bom.welding_process_breakdown.export.values import (
    billing_qty_unit,
    qty_cell,
    standard_name,
    standard_unit,
)


class L0CheckWorkbookBuilder:
    """把 L0 节点列表渲染为核对用 xlsx bytes。"""

    def __init__(self) -> None:
        self.style = build_sheet_style()

    def build(
        self,
        nodes: list[BomWeldingProcessBreakdownNode],
        *,
        filter_text: str,
    ) -> tuple[bytes, int]:
        style = self.style
        wb = Workbook()
        ws = wb.worksheets[0]
        ws.title = SHEET_TITLE

        self._write_banner(ws, filter_text=filter_text, node_count=len(nodes))
        header_row = self._write_headers(ws)

        row_idx = header_row + 1
        export_row_count = 0
        for seq_no, node in enumerate(nodes, start=1):
            lines = sorted(node.lines or [], key=lambda ln: ln.line_no)
            main_fill = style.main_a if seq_no % 2 else style.main_b
            if not lines:
                self._write_row(
                    ws,
                    row_idx,
                    node,
                    seq_no=seq_no,
                    line=None,
                    write_main=True,
                    line_index=0,
                    main_fill=main_fill,
                    row_fill=style.zero,
                )
                ws.row_dimensions[row_idx].height = self._row_height(node, None)
                row_idx += 1
                export_row_count += 1
                continue

            group_start = row_idx
            for i, line in enumerate(lines):
                self._write_row(
                    ws,
                    row_idx,
                    node,
                    seq_no=seq_no,
                    line=line,
                    write_main=(i == 0),
                    line_index=i,
                    main_fill=main_fill,
                    row_fill=None,
                )
                ws.row_dimensions[row_idx].height = self._row_height(node, line)
                row_idx += 1
                export_row_count += 1
            if row_idx - 1 > group_start:
                self._merge_main(ws, group_start, row_idx - 1, main_fill)

        last_data = max(header_row, row_idx - 1)
        for idx, width in enumerate(WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        ws.freeze_panes = "F6"
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(COL_COUNT)}{last_data}"
        ws.print_title_rows = "1:5"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.3
        ws.page_margins.right = 0.3
        ws.page_margins.top = 0.35
        ws.page_margins.bottom = 0.35
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90

        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue(), export_row_count

    def _write_banner(
        self,
        ws: Any,
        *,
        filter_text: str,
        node_count: int,
    ) -> None:
        indent = Alignment(horizontal="left", vertical="center", indent=1)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COL_COUNT)
        c1 = ws.cell(row=1, column=1, value=REPORT_TITLE)
        c1.font = font(size=15, bold=True, color="FFFFFF")
        c1.fill = fill("1A237E")
        c1.alignment = indent
        ws.row_dimensions[1].height = TITLE_H

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=COL_COUNT)
        c2 = ws.cell(
            row=2,
            column=1,
            value=(
                f"{filter_text} · 共 {node_count} 个 L0"
                f"　　导出 {now().strftime('%Y-%m-%d %H:%M')}"
            ),
        )
        c2.font = font(size=9, color="455A64")
        c2.fill = fill("E8EAF6")
        c2.alignment = indent
        ws.row_dimensions[2].height = FILTER_H

        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=COL_COUNT)
        c3 = ws.cell(row=3, column=1, value=NOTE)
        c3.font = font(size=9, color="B71C1C")
        c3.fill = fill("FFF3E0")
        c3.alignment = indent
        ws.row_dimensions[3].height = NOTE_H

    def _write_headers(self, ws: Any) -> int:
        style = self.style
        group_row = 4
        header_row = 5
        white = font(size=9, bold=True, color="FFFFFF")
        header_font = font(bold=True, color="FFFFFF")
        center = style.center

        ws.merge_cells(
            start_row=group_row,
            start_column=1,
            end_row=group_row,
            end_column=MAIN_COLS,
        )
        ws.merge_cells(
            start_row=group_row,
            start_column=MAIN_COLS + 1,
            end_row=group_row,
            end_column=COL_COUNT,
        )
        main_group = ws.cell(row=group_row, column=1, value=GROUP_MAIN)
        main_group.font = white
        main_group.fill = fill("37474F")
        main_group.alignment = center
        for col in range(1, MAIN_COLS + 1):
            cell = ws.cell(row=group_row, column=col)
            cell.fill = fill("37474F")
            cell.border = style.thin if col < MAIN_COLS else style.split_main

        detail_group = ws.cell(
            row=group_row,
            column=MAIN_COLS + 1,
            value=GROUP_DETAIL,
        )
        detail_group.font = white
        detail_group.fill = fill("B71C1C")
        detail_group.alignment = center
        for col in range(MAIN_COLS + 1, COL_COUNT + 1):
            cell = ws.cell(row=group_row, column=col)
            cell.fill = fill("B71C1C")
            cell.border = (
                style.split_detail if col == MAIN_COLS + 1 else style.thin
            )
        ws.row_dimensions[group_row].height = GROUP_H

        for col, title in enumerate(HEADERS, start=1):
            cell = ws.cell(row=header_row, column=col, value=title)
            cell.font = header_font
            cell.alignment = center
            if col <= MAIN_COLS:
                cell.fill = fill("546E7A")
                cell.border = style.thin if col < MAIN_COLS else style.split_main
            elif col in BILLING_COLS:
                cell.fill = fill("EF6C00")
                cell.border = style.thin
            else:
                cell.fill = fill("C62828")
                cell.border = (
                    style.split_detail if col == MAIN_COLS + 1 else style.thin
                )
        ws.row_dimensions[header_row].height = HEADER_H
        return header_row

    @staticmethod
    def _row_height(
        node: BomWeldingProcessBreakdownNode,
        line: BomWeldingProcessBreakdownLine | None,
    ) -> float:
        weld = (node.weld_info_raw or "").strip()
        extra = 0.0
        if len(weld) > 40:
            extra = min(32.0, (len(weld) // 40) * 8.0)
        if line and line.matched_span and len(line.matched_span) > 24:
            extra = max(extra, 6.0)
        return DATA_H + extra

    def _merge_main(
        self,
        ws: Any,
        start_row: int,
        end_row: int,
        main_fill: PatternFill,
    ) -> None:
        style = self.style
        for col in range(1, MAIN_COLS + 1):
            ws.merge_cells(
                start_row=start_row,
                start_column=col,
                end_row=end_row,
                end_column=col,
            )
            cell = ws.cell(row=start_row, column=col)
            cell.fill = main_fill
            cell.border = style.thin if col < MAIN_COLS else style.split_main
            if col == COL["序号"]:
                cell.font = style.seq
                cell.alignment = style.center
            elif col == COL["零件号"]:
                cell.font = style.part
                cell.alignment = style.left
            elif col == COL["工艺数"]:
                cell.font = style.data
                cell.alignment = style.center
            else:
                cell.font = style.data
                cell.alignment = style.left

    def _write_row(
        self,
        ws: Any,
        row_idx: int,
        node: BomWeldingProcessBreakdownNode,
        *,
        seq_no: int,
        line: BomWeldingProcessBreakdownLine | None,
        write_main: bool,
        line_index: int,
        main_fill: PatternFill,
        row_fill: PatternFill | None,
    ) -> None:
        style = self.style
        parse_raw = node.parse_status or ""
        if write_main:
            main_values: list[Any] = [
                seq_no,
                node.location or "",
                node.vehicle_model or "",
                node.bom_version or "",
                node.part_number or "",
                node.part_name or "",
                node.weld_info_raw or "",
                int(node.process_line_count or 0),
                PARSE_LABEL.get(parse_raw, parse_raw),
            ]
        else:
            main_values = [""] * MAIN_COLS

        map_raw = ""
        if line is None:
            line_values: list[Any] = [""] * (COL_COUNT - MAIN_COLS)
        else:
            map_raw = line.map_status or ""
            billing_qty, billing_unit = billing_qty_unit(line)
            line_values = [
                line.line_no,
                line.process_token or "",
                line.matched_span or "",
                standard_name(line.process_token, line.standard_process_name),
                standard_unit(
                    line.process_token,
                    line.standard_process_name,
                    line.standard_process_unit,
                ),
                qty_cell(line.process_qty),
                line.process_qty_unit or "",
                qty_cell(line.process_qty_secondary),
                line.process_qty_secondary_unit or "",
                billing_qty,
                billing_unit,
                MAP_LABEL.get(map_raw, map_raw),
            ]

        values = [*main_values, *line_values]
        alt = bool(line_index % 2)
        detail = style.line_alt if alt else style.line
        focus = style.focus_alt if alt else style.focus
        billing = style.billing_alt if alt else style.billing_fill
        ok = (style.ok_font, style.ok_fill)
        warn = (style.warn_font, style.warn_fill)
        bad = (style.bad_font, style.bad_fill)

        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            if col == MAIN_COLS:
                cell.border = style.split_main
            elif col == MAIN_COLS + 1:
                cell.border = style.split_detail
            else:
                cell.border = style.thin

            if row_fill is not None:
                cell.fill = row_fill
                cell.font = (
                    style.accent
                    if col in (COL["序号"], COL["零件号"])
                    or col in STD_COLS | BILLING_COLS
                    else style.data
                )
            elif col <= MAIN_COLS:
                cell.fill = main_fill
                if col == COL["序号"]:
                    cell.font = style.seq
                elif col == COL["零件号"]:
                    cell.font = style.part
                elif col == COL["解析状态"]:
                    cell.font, cell.fill = tone(
                        parse_raw,
                        PARSE_TONE,
                        ok=ok,
                        warn=warn,
                        bad=bad,
                        default=(style.data, main_fill),
                    )
                else:
                    cell.font = style.data
            elif col in BILLING_COLS:
                cell.fill = billing
                cell.font = style.billing
            elif col in STD_COLS:
                cell.fill = focus
                cell.font = style.accent
            elif col == COL["映射状态"]:
                cell.font, cell.fill = tone(
                    map_raw,
                    MAP_TONE,
                    ok=ok,
                    warn=warn,
                    bad=bad,
                    default=(style.muted, detail),
                )
            elif col in MUTED_QTY_COLS:
                cell.fill = detail
                cell.font = style.muted
            else:
                cell.fill = detail
                cell.font = style.data

            if col in QTY_COLS:
                cell.alignment = style.right
                if value != "":
                    cell.number_format = "0.####"
            elif col in CENTER_COLS:
                cell.alignment = style.center
            else:
                cell.alignment = style.left
