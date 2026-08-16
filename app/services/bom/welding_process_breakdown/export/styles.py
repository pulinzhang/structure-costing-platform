"""
L0 核对导出 — openpyxl 色板与字体
"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_FONT_NAME = "微软雅黑"

PARSE_TONE = {
    "ok": "ok",
    "partial": "warn",
    "unparsed": "bad",
    "anomaly": "bad",
}
MAP_TONE = {
    "ready": "ok",
    "ambiguous": "warn",
    "unmapped": "bad",
}


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def font(
    *,
    size: int = 10,
    bold: bool = False,
    color: str = "263238",
) -> Font:
    return Font(name=_FONT_NAME, size=size, bold=bold, color=color)


def border(*, right: Side | None = None, left: Side | None = None) -> Border:
    thin = Side(style="thin", color="CFD8DC")
    return Border(
        left=left or thin,
        right=right or thin,
        top=thin,
        bottom=thin,
    )


def tone(
    raw: str,
    mapping: dict[str, str],
    *,
    ok: tuple[Font, PatternFill],
    warn: tuple[Font, PatternFill],
    bad: tuple[Font, PatternFill],
    default: tuple[Font, PatternFill],
) -> tuple[Font, PatternFill]:
    key = mapping.get(raw, "")
    if key == "ok":
        return ok
    if key == "warn":
        return warn
    if key == "bad":
        return bad
    return default


@dataclass(frozen=True)
class SheetStyle:
    data: Font
    muted: Font
    accent: Font
    billing: Font
    part: Font
    seq: Font
    ok_font: Font
    warn_font: Font
    bad_font: Font
    left: Alignment
    center: Alignment
    right: Alignment
    thin: Border
    split_main: Border
    split_detail: Border
    line: PatternFill
    line_alt: PatternFill
    focus: PatternFill
    focus_alt: PatternFill
    billing_fill: PatternFill
    billing_alt: PatternFill
    zero: PatternFill
    ok_fill: PatternFill
    warn_fill: PatternFill
    bad_fill: PatternFill
    main_a: PatternFill
    main_b: PatternFill


def build_sheet_style() -> SheetStyle:
    return SheetStyle(
        data=font(),
        muted=font(color="607D8B"),
        accent=font(bold=True, color="B71C1C"),
        billing=font(size=11, bold=True, color="E65100"),
        part=font(bold=True, color="1565C0"),
        seq=font(size=11, bold=True, color="37474F"),
        ok_font=font(bold=True, color="2E7D32"),
        warn_font=font(bold=True, color="EF6C00"),
        bad_font=font(bold=True, color="C62828"),
        left=Alignment(horizontal="left", vertical="center", wrap_text=True),
        center=Alignment(horizontal="center", vertical="center", wrap_text=True),
        right=Alignment(horizontal="right", vertical="center"),
        thin=border(),
        split_main=border(right=Side(style="medium", color="B71C1C")),
        split_detail=border(left=Side(style="medium", color="B71C1C")),
        line=fill("FFFFFF"),
        line_alt=fill("F7F9FC"),
        focus=fill("FFEBEE"),
        focus_alt=fill("FFCDD2"),
        billing_fill=fill("FFF3E0"),
        billing_alt=fill("FFE0B2"),
        zero=fill("FFCDD2"),
        ok_fill=fill("E8F5E9"),
        warn_fill=fill("FFF3E0"),
        bad_fill=fill("FFEBEE"),
        main_a=fill("FFFDE7"),
        main_b=fill("FFF8E1"),
    )
