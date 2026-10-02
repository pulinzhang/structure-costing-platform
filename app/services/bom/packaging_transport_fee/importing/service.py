"""BOM 包装运输费 Excel 导入 / 模板下载服务。"""

from __future__ import annotations

import logging
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from app.common.audit import audit_create, audit_update
from app.common.period import is_costing_year
from app.common.workbook_reject import EMPTY_FILE, NO_DATA_ROWS, format_workbook_reject
from app.config import settings
from app.exceptions.base import BusinessException
from app.parsers.bom.packaging_transport_fee_parser import (
    parse_packaging_transport_fee_excel,
)
from app.repositories.bom.bom.snapshot import BomSnapshotRepository
from app.repositories.bom.packaging_transport_fee_repository import (
    BomPackagingTransportFeeRepository,
)
from app.schemas.dto.bom_packaging_transport_fee_dto import (
    BomPackagingTransportFeeCreate,
    BomPackagingTransportFeeUpdate,
    quantize_fee,
)
from app.schemas.vo.bom_packaging_transport_fee_vo import (
    BomPackagingTransportFeeImportResultVO,
)
from app.schemas.vo.user_vo import UserVO
from app.services.file_trace.register import register_existing
from app.services.masterdata.vehicle_model.service import VehicleModelService
from app.utils.bom_packaging_transport_fee_archive import (
    save_packaging_transport_fee_upload,
)

logger = logging.getLogger("service.bom_packaging_transport_fee_import")

TEMPLATE_HEADERS = [
    "零件号",
    "产品",
    "包装类型",
    "单包装数",
    "器具长",
    "器具宽",
    "器具高",
    "上线库存数量",
    "在途周转数量",
    "公司库存数量",
    "总投入数量",
    "器具单价",
    "投入总金额",
    "单件包装费",
    "往返运输价格",
    "单件运输费",
    "器具年损维护费",
    "其它费用",
    "包装运输费合计",
    "备注",
]

TEMPLATE_SAMPLE: list[tuple[str | int | Decimal, ...]] = [
    (
        "501006870AA",
        "流水槽后部总成",
        "带轮料架",
        18,
        1600,
        1100,
        1150,
        16,
        16,
        48,
        80,
        Decimal("4200.0000"),
        Decimal("336000.0000"),
        Decimal("0.9800"),
        Decimal("500.0000"),
        Decimal("1.5500"),
        Decimal("0.4900"),
        Decimal("0.6994"),
        Decimal("3.7194"),
        "",
    ),
]


class BomPackagingTransportFeeImportService:
    """包装运输费同步导入。"""

    def __init__(self) -> None:
        self.repository = BomPackagingTransportFeeRepository()
        self.snapshot_repository = BomSnapshotRepository()

    def build_template(self) -> tuple[bytes, str]:
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = "包装运输费"

        header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        data_font = Font(name="微软雅黑", size=9)
        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )

        for col, title in enumerate(TEMPLATE_HEADERS, start=1):
            cell = sheet.cell(row=1, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        money_cols = {12, 13, 14, 15, 16, 17, 18, 19}
        for row_idx, sample in enumerate(TEMPLATE_SAMPLE, start=2):
            for col, value in enumerate(sample, start=1):
                cell = sheet.cell(row=row_idx, column=col, value=value)
                cell.font = data_font
                cell.border = thin_border
                if col in money_cols:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "0.0000"
                elif col in (4, 5, 6, 7, 8, 9, 10, 11):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        widths = [16, 18, 12, 10, 10, 10, 10, 12, 12, 12, 10, 12, 12, 12, 12, 12, 12, 10, 14, 12]
        from openpyxl.utils import get_column_letter

        for idx, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(idx)].width = width
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 30

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue(), "包装运输费导入模板.xlsx"

    def import_excel(
        self,
        db: Session,
        *,
        file_content: bytes,
        filename: str,
        location: str,
        year: int,
        vehicle_model: str,
        current_user: UserVO,
    ) -> BomPackagingTransportFeeImportResultVO:
        location = (location or "").strip()
        if not location:
            raise BusinessException("请选择采购基地")
        vehicle_model = (vehicle_model or "").strip()
        if not vehicle_model:
            raise BusinessException("请选择项目")
        if not is_costing_year(year):
            raise BusinessException("年度不正确")
        location, vehicle_model = VehicleModelService().require_overlaps_year(
            db, location, vehicle_model, year
        )

        bom_snapshots = self.snapshot_repository.list_by_location_model(db, location, vehicle_model)
        if not bom_snapshots:
            raise BusinessException(
                f"未找到材料清单：基地「{location}」项目「{vehicle_model}」，"
                "请先上传 BOM 后再导入包装运输费"
            )

        original_name = Path(filename or "").name
        if Path(original_name).suffix.lower() != ".xlsx":
            raise BusinessException("仅支持 .xlsx 文件")
        if not file_content:
            raise BusinessException(EMPTY_FILE)
        if len(file_content) > settings.MAX_UPLOAD_SIZE:
            max_mb = settings.MAX_UPLOAD_SIZE // (1024 * 1024)
            raise BusinessException(f"文件过大，最大支持 {max_mb}MB")

        # 导入前先归档源文件（对齐一事一议加减价/基价溯源规则）
        archive_path = save_packaging_transport_fee_upload(
            file_content=file_content,
            filename=original_name,
            year=year,
            location=location,
            vehicle_model=vehicle_model,
        )
        try:
            archived_rel = str(archive_path.relative_to(settings.UPLOAD_DIR)).replace("\\", "/")
        except ValueError:
            archived_rel = str(archive_path).replace("\\", "/")

        parse_result = parse_packaging_transport_fee_excel(file_content)
        if parse_result.errors:
            raise BusinessException(format_workbook_reject(parse_result.errors))
        if not parse_result.rows:
            raise BusinessException(NO_DATA_ROWS)

        sheet_name = parse_result.sheet_name
        created_count = 0
        updated_count = 0
        skipped_approved = 0
        create_audit = audit_create(current_user)
        update_audit = audit_update(current_user)

        try:
            for row in parse_result.rows:
                existing = self.repository.get_by_unique_key(
                    db, location, year, vehicle_model, row.part_number
                )
                if existing:
                    if existing.audit_status == "approved":
                        skipped_approved += 1
                        continue
                    self.repository.update(
                        db,
                        existing.id,
                        BomPackagingTransportFeeUpdate(
                            part_name=row.part_name,
                            packaging_type=row.packaging_type,
                            pieces_per_pack=row.pieces_per_pack,
                            fixture_length_mm=row.fixture_length_mm,
                            fixture_width_mm=row.fixture_width_mm,
                            fixture_height_mm=row.fixture_height_mm,
                            oem_stock_qty=row.oem_stock_qty,
                            transit_qty=row.transit_qty,
                            company_stock_qty=row.company_stock_qty,
                            total_fixture_qty=row.total_fixture_qty,
                            fixture_unit_price=row.fixture_unit_price,
                            fixture_invest_amount=row.fixture_invest_amount,
                            unit_packaging_fee=quantize_fee(row.unit_packaging_fee),
                            round_trip_transport_price=row.round_trip_transport_price,
                            unit_transport_fee=quantize_fee(row.unit_transport_fee),
                            fixture_maint_fee=quantize_fee(row.fixture_maint_fee),
                            other_fee=quantize_fee(row.other_fee),
                            packaging_transport_total=quantize_fee(row.packaging_transport_total),
                            remark=row.remark,
                            excel_filename=original_name,
                            sheet_name=sheet_name,
                            excel_row=row.row_number,
                        ),
                        updated_by=update_audit["updated_by"],
                    )
                    updated_count += 1
                else:
                    self.repository.create(
                        db,
                        BomPackagingTransportFeeCreate(
                            location=location,
                            year=year,
                            vehicle_model=vehicle_model,
                            part_number=row.part_number,
                            part_name=row.part_name,
                            packaging_type=row.packaging_type,
                            pieces_per_pack=row.pieces_per_pack,
                            fixture_length_mm=row.fixture_length_mm,
                            fixture_width_mm=row.fixture_width_mm,
                            fixture_height_mm=row.fixture_height_mm,
                            oem_stock_qty=row.oem_stock_qty,
                            transit_qty=row.transit_qty,
                            company_stock_qty=row.company_stock_qty,
                            total_fixture_qty=row.total_fixture_qty,
                            fixture_unit_price=row.fixture_unit_price,
                            fixture_invest_amount=row.fixture_invest_amount,
                            unit_packaging_fee=quantize_fee(row.unit_packaging_fee),
                            round_trip_transport_price=row.round_trip_transport_price,
                            unit_transport_fee=quantize_fee(row.unit_transport_fee),
                            fixture_maint_fee=quantize_fee(row.fixture_maint_fee),
                            other_fee=quantize_fee(row.other_fee),
                            packaging_transport_total=quantize_fee(row.packaging_transport_total),
                            remark=row.remark,
                            excel_filename=original_name,
                            sheet_name=sheet_name,
                            excel_row=row.row_number,
                        ),
                        created_by=create_audit["created_by"],
                        updated_by=create_audit["updated_by"],
                    )
                    created_count += 1
            register_existing(
                db,
                file_category="bom_packaging_transport_fee",
                stored_relpath=archived_rel,
                original_filename=original_name,
                current_user=current_user,
                location=location,
                year=year,
                vehicle_model=vehicle_model,
            )
            db.commit()
        except Exception:
            db.rollback()
            raise

        warnings: list[str] = []
        if skipped_approved:
            warnings.append(f"跳过已审核 {skipped_approved} 条（未覆盖）")
        message = (
            f"导入完成：基地「{location}」{year} 年项目「{vehicle_model}」，"
            f"新增 {created_count} 条，更新 {updated_count} 条"
        )
        if skipped_approved:
            message += f"，跳过已审核 {skipped_approved} 条"

        logger.info(
            "%s file=%s archived=%s created=%s updated=%s skipped_approved=%s",
            message,
            original_name,
            archived_rel,
            created_count,
            updated_count,
            skipped_approved,
        )
        return BomPackagingTransportFeeImportResultVO(
            location=location,
            year=year,
            vehicle_model=vehicle_model,
            created_count=created_count,
            updated_count=updated_count,
            total_count=created_count + updated_count,
            skipped_count=skipped_approved,
            skipped_approved_count=skipped_approved,
            warnings=warnings,
            message=message,
            source_filename=original_name,
        )
