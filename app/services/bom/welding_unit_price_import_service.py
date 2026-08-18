"""焊接单价 Excel 导入 / 模板下载服务。"""

from __future__ import annotations

import logging
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from app.common.audit import audit_create, audit_update
from app.parsers.bom.welding_unit_price_parser import parse_welding_unit_price_excel
from app.repositories.bom.welding_unit_price_repository import BomWeldingUnitPriceRepository
from app.schemas.dto.bom_welding_unit_price_dto import (
    BomWeldingUnitPriceCreate,
    BomWeldingUnitPriceUpdate,
    quantize_unit_price,
)
from app.schemas.vo.bom_welding_unit_price_vo import BomWeldingUnitPriceImportResultVO
from app.schemas.vo.user_vo import UserVO

logger = logging.getLogger("service.bom_welding_unit_price_import")

TEMPLATE_SAMPLE_ROWS: list[tuple[str, str, Decimal, str]] = [
    ("焊点", "个", Decimal("0.0999"), ""),
    ("植/螺柱/凸焊", "个", Decimal("0.1053"), ""),
    ("铆接", "个", Decimal("0.1053"), ""),
    ("CO2焊缝", "CM", Decimal("0.0700"), ""),
    ("折边/点焊胶", "MM", Decimal("0.0009"), ""),
    ("电镀", "平方", Decimal("12.0000"), ""),
    ("电泳（盐雾试验要求≤244H）", "平方", Decimal("12.0000"), ""),
    ("电泳（盐雾试验要求244H-720H）", "平方", Decimal("15.0000"), ""),
    ("电泳（盐雾试验要求＞720H）", "平方", Decimal("18.0000"), ""),
    ("压铆", "个", Decimal("0.1800"), ""),
    ("打铆钉", "个", Decimal("0.4500"), ""),
    ("二维码", "个", Decimal("0.1800"), ""),
    ("套袋", "个", Decimal("0.0900"), ""),
]


class BomWeldingUnitPriceImportService:
    """焊接单价同步导入（数据量小，不走异步任务）。"""

    def __init__(self) -> None:
        self.repository = BomWeldingUnitPriceRepository()

    def build_template(self) -> tuple[bytes, str]:
        """生成四列模板：名称 / 单位 / 费用 / 备注（费用=单位工费单价）。"""
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = "焊接单价"

        header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        data_font = Font(name="微软雅黑", size=9)
        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )

        headers = ["名称", "单位", "费用", "备注"]
        for col, title in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for row_idx, (name, unit, unit_price, remark) in enumerate(
            TEMPLATE_SAMPLE_ROWS, start=2
        ):
            values: list[str | Decimal] = [name, unit, unit_price, remark]
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_idx, column=col, value=value)
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal="center" if col in (2,) else ("right" if col == 3 else "left"),
                    vertical="center",
                )
                if col == 3:
                    cell.number_format = "0.0000"

        sheet.column_dimensions["A"].width = 36
        sheet.column_dimensions["B"].width = 10
        sheet.column_dimensions["C"].width = 12
        sheet.column_dimensions["D"].width = 20
        sheet.freeze_panes = "A2"

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue(), "焊接单价导入模板.xlsx"

    def import_excel(
        self,
        db: Session,
        *,
        file_content: bytes,
        filename: str,
        location: str,
        year: int,
        current_user: UserVO,
    ) -> BomWeldingUnitPriceImportResultVO:
        location = (location or "").strip()
        if not location:
            raise ValueError("请选择采购基地")
        if year < 2000 or year > 2100:
            raise ValueError("年度不正确")

        original_name = Path(filename or "").name
        if Path(original_name).suffix.lower() != ".xlsx":
            raise ValueError("仅支持 .xlsx 文件")
        if not file_content:
            raise ValueError("文件内容为空")

        parse_result = parse_welding_unit_price_excel(file_content)
        if parse_result.errors and not parse_result.rows:
            raise ValueError("；".join(parse_result.errors))
        if not parse_result.rows:
            raise ValueError("未解析到有效数据行")

        created_count = 0
        updated_count = 0
        skipped_approved = 0
        create_audit = audit_create(current_user)
        update_audit = audit_update(current_user)

        try:
            for row in parse_result.rows:
                unit_price = quantize_unit_price(row.unit_price)
                existing = self.repository.get_by_location_year_item(
                    db, location, year, row.item_name
                )
                if existing:
                    if existing.audit_status == "approved":
                        skipped_approved += 1
                        continue
                    self.repository.update(
                        db,
                        existing.id,
                        BomWeldingUnitPriceUpdate(
                            unit=row.unit,
                            unit_price=unit_price,
                            remark=row.remark,
                        ),
                        updated_by=update_audit["updated_by"],
                        commit=False,
                    )
                    updated_count += 1
                else:
                    self.repository.create(
                        db,
                        BomWeldingUnitPriceCreate(
                            location=location,
                            year=year,
                            item_name=row.item_name,
                            unit=row.unit,
                            unit_price=unit_price,
                            remark=row.remark,
                        ),
                        created_by=create_audit["created_by"],
                        updated_by=create_audit["updated_by"],
                        commit=False,
                    )
                    created_count += 1
            db.commit()
        except Exception:
            db.rollback()
            raise

        warnings = list(parse_result.errors)
        if skipped_approved:
            warnings.append(f"跳过已审核 {skipped_approved} 条（未覆盖）")
        message = (
            f"导入完成：基地「{location}」{year} 年，"
            f"新增 {created_count} 条，更新 {updated_count} 条"
        )
        if skipped_approved:
            message += f"，跳过已审核 {skipped_approved} 条"
        if parse_result.errors:
            message += f"；有 {len(parse_result.errors)} 行解析跳过"

        logger.info(
            "%s file=%s created=%s updated=%s skipped_approved=%s warnings=%s",
            message,
            original_name,
            created_count,
            updated_count,
            skipped_approved,
            len(warnings),
        )
        return BomWeldingUnitPriceImportResultVO(
            location=location,
            year=year,
            created_count=created_count,
            updated_count=updated_count,
            total_count=created_count + updated_count,
            skipped_count=len(parse_result.errors) + skipped_approved,
            warnings=warnings,
            message=message,
        )
