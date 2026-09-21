"""BOM 材料清单 Excel 导入模板生成。"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

CellValue = str | int | float | None


def _merge_vertical(sheet: Worksheet, col: int, start_row: int, end_row: int) -> None:
    if end_row > start_row:
        sheet.merge_cells(
            start_row=start_row,
            start_column=col,
            end_row=end_row,
            end_column=col,
        )


def build_bom_material_template() -> tuple[bytes, str]:
    """生成车身件材料清单样例（表头/总成/一级纵向合并，对齐实样）。"""
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "车身件材料清单"

    header_font = Font(name="微软雅黑", size=9, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    data_font = Font(name="微软雅黑", size=9)
    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 头信息横跨（对齐实样 A1:W1 / A2:W2 / 分区 meta）
    sheet.merge_cells("A1:W1")
    sheet["A1"] = "T28样例项目产品材料清单"
    sheet["A1"].font = Font(name="微软雅黑", size=14, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")

    sheet.merge_cells("A2:W2")
    sheet["A2"] = (
        "编制部门： 技术中心                    "
        "编制日期：2025-01-01  版本： D/00  编号：BOM-T28-DEMO"
    )

    sheet.merge_cells("A3:H3")
    sheet["A3"] = "项目代号：  T28"
    sheet.merge_cells("I3:M3")
    sheet["I3"] = "产品类别：  车身件"
    sheet.merge_cells("N3:O3")
    sheet["N3"] = "顾客名称：  奇瑞公司"
    sheet.merge_cells("P3:Q3")
    sheet["P3"] = "阶段标识： D"
    sheet.merge_cells("R3:U3")
    sheet["R3"] = "■批量生产清单"
    sheet.merge_cells("V3:W3")
    sheet["V3"] = "修订日期：2025-01-01"

    headers = [
        "序号",
        "总成（供货）零件内部号",
        "总成（供货）零件号",
        "总成（供货）零件名称",
        "总成图示",
        "总成单车用量",
        "焊接总成 信息",
        "一级子件内部号",
        "一级子件 零件号",
        "一级子件 零件名称",
        "一级子件单车用量",
        "焊接总成 信息",
        "二级子件内部号",
        "二级子件 零件号",
        "二级子件零件 名称",
        "单件图示",
        "二级子件单车用量",
        "材料",
        "料厚 （mm）",
        "材料厂家",
        "单件重量 （kg）",
        "检具是否 开发",
        "备注",
    ]
    for col, title in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin
    sheet.row_dimensions[4].height = 36

    # 行结构（与实样一致：总成跨多行；一级分总成跨其二级行）
    # R5: L0 + L1 自制叶子
    # R6-R7: 同一 L0 下的 L1 分总成 + 两个 L2
    # R8: 同一 L0 下另一 L1 自制叶子
    rows: list[list[CellValue]] = [
        [
            1,
            "ACA610",
            "501006873AA",
            "流水槽加强梁总成",
            None,
            1,
            "焊点：4 凸焊：2",
            "ACA610001",
            "501006875AA",
            "雨刮电机安装支架Ⅰ",
            1,
            "/",
            "/",
            "/",
            "/",
            None,
            1,
            "DC51D+Z",
            "T=1.2",
            "样例钢厂",
            0.096,
            "是",
            "自制件样例",
        ],
        [
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "ACA395N01",
            "501005123AA-1",
            "流水槽上加强梁本体总成",
            1,
            "凸焊：2",
            "/",
            "501005123AA",
            "流水槽上加强梁本体",
            None,
            1,
            "DC04",
            "T=0.8",
            "样例钢厂",
            1.25,
            "是",
            "一级分总成+二级自制",
        ],
        [
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "/",
            "N0111931",
            "焊接方螺母",
            None,
            2,
            None,
            "M6",
            None,
            0.008,
            "/",
            "标准件/外购样例",
        ],
        [
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "ACA610002",
            "501008900AA",
            "流水槽下加强梁",
            1,
            "/",
            "/",
            "/",
            "/",
            None,
            1,
            "HC340LA",
            "T=1.0",
            "样例钢厂",
            0.82,
            "是",
            "另一自制一级",
        ],
    ]

    for row_idx, values in enumerate(rows, start=5):
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_idx, column=col, value=value)
            cell.font = data_font
            cell.border = thin
            cell.alignment = center

    # 总成字段跨 R5-R8（同一供货总成）
    for col in (1, 2, 3, 4, 5, 6, 7):
        _merge_vertical(sheet, col, 5, 8)
    # 写回总成值到合并首格
    sheet.cell(row=5, column=1, value=1)
    sheet.cell(row=5, column=2, value="ACA610")
    sheet.cell(row=5, column=3, value="501006873AA")
    sheet.cell(row=5, column=4, value="流水槽加强梁总成")
    sheet.cell(row=5, column=6, value=1)
    sheet.cell(row=5, column=7, value="焊点：4 凸焊：2")

    # 一级分总成跨其二级行 R6-R7
    for col in (8, 9, 10, 11, 12):
        _merge_vertical(sheet, col, 6, 7)
    sheet.cell(row=6, column=8, value="ACA395N01")
    sheet.cell(row=6, column=9, value="501005123AA-1")
    sheet.cell(row=6, column=10, value="流水槽上加强梁本体总成")
    sheet.cell(row=6, column=11, value=1)
    sheet.cell(row=6, column=12, value="凸焊：2")

    for idx in range(1, 24):
        sheet.column_dimensions[get_column_letter(idx)].width = 12 if idx != 4 else 18
    sheet.column_dimensions["D"].width = 20
    sheet.column_dimensions["J"].width = 22
    sheet.column_dimensions["O"].width = 20
    sheet.freeze_panes = "A5"

    note = workbook.create_sheet("填写说明")
    _write_notes(
        note,
        [
            "本模板对齐「车身件产品材料清单」三级稀疏树：总成 / 一级 / 二级。",
            "标题与编制信息行横向合并；同一总成下多行纵向合并总成列。",
            "一级分总成若含多个二级子件，一级列须纵向合并到末行二级。",
            "板料写 T=1.2 判定为自制件；标准件写 M6 等判定为外购件。",
            "上传页须先选择采购基地；同基地+车型+版本再次导入将整树替换。简图可空。",
        ],
    )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), "材料清单导入模板.xlsx"


def _write_notes(sheet: Worksheet, lines: list[str]) -> None:
    sheet["A1"] = "填写说明"
    sheet["A1"].font = Font(name="微软雅黑", size=12, bold=True, color="C00000")
    for idx, line in enumerate(lines, start=3):
        sheet.cell(row=idx, column=1, value=f"{idx - 2}. {line}")
        sheet.cell(row=idx, column=1).font = Font(name="微软雅黑", size=10)
        sheet.cell(row=idx, column=1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[idx].height = 22
    sheet.column_dimensions["A"].width = 96
