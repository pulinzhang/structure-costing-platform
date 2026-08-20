"""BOM 包装运输费 Excel 导出服务。"""

from __future__ import annotations

import logging
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.common.money import Q4, excel_money_cell
from app.common.timezone import now
from app.repositories.bom.packaging_transport_fee_repository import (
    BomPackagingTransportFeeRepository,
)

logger = logging.getLogger("service.bom_packaging_transport_fee_export")

_AUDIT_LABELS: dict[str, str] = {
    "pending": "待审",
    "approved": "已通过",
    "rejected": "已驳回",
}


def _money_cell(value: Any) -> Decimal | str:
    if value is None:
        return ""
    return excel_money_cell(value, Q4)


class BomPackagingTransportFeeExportService:
    """包装运输费 Excel 导出。"""

    def __init__(self) -> None:
        self.repository = BomPackagingTransportFeeRepository()

    def export(
        self,
        db: Session,
        *,
        location: str | None = None,
        year: int | None = None,
        vehicle_model: str | None = None,
    ) -> tuple[bytes, str]:
        records = self.repository.list_for_export(
            db,
            location=location,
            year=year,
            vehicle_model=vehicle_model,
        )

        wb = Workbook()
        ws = wb.worksheets[0]
        ws.title = "包装运输费"

        header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        header_alignment = Alignment(horizontal="center", vertical="center")
        data_font = Font(name="微软雅黑", size=9)
        align_left = Alignment(horizontal="left", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )

        headers = [
            "基地",
            "年份",
            "项目",
            "零件号",
            "产品",
            "包装类型",
            "单包装数",
            "器具长",
            "器具宽",
            "器具高",
            "上线库存数量",
            "在途周转数量",
            "公司库存数量",
            "总投入数量",
            "器具单价",
            "投入总金额",
            "单件包装费",
            "往返运输价格",
            "单件运输费",
            "器具年损维护费",
            "其它费用",
            "包装运输费合计",
            "备注",
            "审核状态",
        ]
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        money_cols = {15, 16, 17, 18, 19, 20, 21, 22}
        for row_idx, record in enumerate(records, start=2):
            audit_label = _AUDIT_LABELS.get(record.audit_status, record.audit_status)
            values: list[Any] = [
                record.location,
                record.year,
                record.vehicle_model,
                record.part_number,
                record.part_name or "",
                record.packaging_type or "",
                record.pieces_per_pack if record.pieces_per_pack is not None else "",
                record.fixture_length_mm if record.fixture_length_mm is not None else "",
                record.fixture_width_mm if record.fixture_width_mm is not None else "",
                record.fixture_height_mm if record.fixture_height_mm is not None else "",
                record.oem_stock_qty if record.oem_stock_qty is not None else "",
                record.transit_qty if record.transit_qty is not None else "",
                record.company_stock_qty if record.company_stock_qty is not None else "",
                record.total_fixture_qty if record.total_fixture_qty is not None else "",
                _money_cell(record.fixture_unit_price),
                _money_cell(record.fixture_invest_amount),
                _money_cell(record.unit_packaging_fee),
                _money_cell(record.round_trip_transport_price),
                _money_cell(record.unit_transport_fee),
                _money_cell(record.fixture_maint_fee),
                _money_cell(record.other_fee),
                _money_cell(record.packaging_transport_total),
                record.remark or "",
                audit_label,
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.font = data_font
                cell.border = thin_border
                if col in money_cols:
                    cell.alignment = align_right
                    if value != "":
                        cell.number_format = "0.0000"
                elif col in (2, 6, 7, 8, 9, 10, 11, 12, 13, 14, 24):
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        for idx in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(idx)].width = 12
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 18
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:X{max(1, len(records) + 1)}"

        buffer = BytesIO()
        wb.save(buffer)
        content = buffer.getvalue()

        stamp = now().strftime("%Y%m%d_%H%M%S")
        parts = ["包装运输费"]
        if location:
            parts.append(location)
        if year:
            parts.append(f"{year}年")
        if vehicle_model:
            parts.append(vehicle_model)
        filename = "_".join(parts) + f"_{stamp}.xlsx"

        logger.info(
            "导出包装运输费 %s 条 location=%s year=%s vehicle=%s -> %s",
            len(records),
            location,
            year,
            vehicle_model,
            filename,
        )
        return content, filename
