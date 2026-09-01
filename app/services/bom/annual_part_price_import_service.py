"""BOM 年度零件价 Excel 导入 / 模板下载服务。"""

from __future__ import annotations

import logging
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from app.common.audit import audit_create, audit_update
from app.config import settings
from app.exceptions.base import BusinessException
from app.parsers.bom.annual_part_price_parser import (
    HEADER_PART_NAME,
    HEADER_PART_NUMBER,
    HEADER_REMARK,
    HEADER_SUPPLIER,
    HEADER_SUPPLY_TYPE,
    HEADER_UNIT,
    HEADER_UNIT_PRICE,
    parse_annual_part_price_excel,
)
from app.repositories.bom.annual_part_price_repository import (
    BomAnnualPartPriceRepository,
)
from app.schemas.dto.bom_annual_part_price_dto import (
    BomAnnualPartPriceCreate,
    BomAnnualPartPriceUpdate,
    quantize_unit_price,
)
from app.schemas.vo.bom_annual_part_price_vo import BomAnnualPartPriceImportResultVO
from app.schemas.vo.user_vo import UserVO
from app.services.bom.procurement_object.classify import SECTION_OTHER_LEAF
from app.services.bom.procurement_object.labels import SECTION_LABELS
from app.services.file_trace.path_utils import to_upload_relpath
from app.services.file_trace.register import register_existing
from app.utils.bom_annual_part_price_archive import save_annual_part_price_upload

logger = logging.getLogger("service.bom_annual_part_price_import")

_TEMPLATE_SUPPLY_TYPE = SECTION_LABELS[SECTION_OTHER_LEAF]
_TEMPLATE_ZERO_PRICE = Decimal("0.0000")


class BomAnnualPartPriceImportService:
    """年度零件价同步导入（不走异步任务）。"""

    def __init__(self) -> None:
        self.repository = BomAnnualPartPriceRepository()

    def build_template(self, db: Session, *, location: str, year: int) -> tuple[bytes, str]:
        """按基地采购对象「标准件/胶等」生成导入模板。"""
        from app.repositories.bom.procurement_object_repository import (
            BomProcurementObjectRepository,
        )

        location = (location or "").strip()
        if not location:
            raise BusinessException("请选择采购基地")
        if year < 2000 or year > 2100:
            raise BusinessException("年度不正确")

        procurement_object_repository = BomProcurementObjectRepository()
        if not procurement_object_repository.has_list_for_location(db, location):
            raise BusinessException("请先生成该基地的采购对象清单")

        source_rows = procurement_object_repository.list_objects_for_export(
            db,
            location=location,
            section=SECTION_OTHER_LEAF,
        )
        seen_part_numbers: set[str] = set()
        template_rows: list[tuple[str, str, str, Decimal, str, str, str]] = []
        for row in source_rows:
            part_number = (row.part_number or "").strip()
            if not part_number or part_number in seen_part_numbers:
                continue
            seen_part_numbers.add(part_number)
            unit = (row.qty_unit or "").strip() or "pcs"
            if unit.upper() == "PCS":
                unit = "pcs"
            template_rows.append(
                (
                    part_number,
                    (row.part_name or "").strip(),
                    unit,
                    _TEMPLATE_ZERO_PRICE,
                    _TEMPLATE_SUPPLY_TYPE,
                    "",
                    "",
                )
            )
        if not template_rows:
            raise BusinessException("该基地采购对象清单中没有标准件/胶等，无法生成模板")

        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = "年度零件价"

        header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        data_font = Font(name="微软雅黑", size=9)
        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )

        headers = [
            HEADER_PART_NUMBER,
            HEADER_PART_NAME,
            HEADER_UNIT,
            HEADER_UNIT_PRICE,
            HEADER_SUPPLY_TYPE,
            HEADER_SUPPLIER,
            HEADER_REMARK,
        ]
        for col, title in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for row_idx, sample in enumerate(template_rows, start=2):
            values: list[str | Decimal] = list(sample)
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_idx, column=col, value=value)
                cell.font = data_font
                cell.border = thin_border
                if col == 4:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "0.0000"
                elif col in (3, 5):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        widths = [18, 24, 8, 12, 12, 16, 20]
        for idx, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + idx)].width = width
        sheet.freeze_panes = "A2"

        buffer = BytesIO()
        workbook.save(buffer)
        filename = f"年度零件价导入模板_{location}_{year}.xlsx"
        return buffer.getvalue(), filename

    def import_excel(
        self,
        db: Session,
        *,
        file_content: bytes,
        filename: str,
        location: str,
        year: int,
        current_user: UserVO,
    ) -> BomAnnualPartPriceImportResultVO:
        location = (location or "").strip()
        if not location:
            raise BusinessException("请选择采购基地")
        if year < 2000 or year > 2100:
            raise BusinessException("年度不正确")

        original_name = Path(filename or "").name
        if Path(original_name).suffix.lower() != ".xlsx":
            raise BusinessException("仅支持 .xlsx 文件")
        if not file_content:
            raise BusinessException("文件内容为空")
        if len(file_content) > settings.MAX_UPLOAD_SIZE:
            max_mb = settings.MAX_UPLOAD_SIZE // (1024 * 1024)
            raise BusinessException(f"文件过大，最大支持 {max_mb}MB")

        archive_path = save_annual_part_price_upload(
            file_content=file_content,
            filename=original_name,
            year=year,
            location=location,
        )
        archived_rel = to_upload_relpath(archive_path)

        parse_result = parse_annual_part_price_excel(file_content)
        if parse_result.errors and not parse_result.rows:
            raise BusinessException("；".join(parse_result.errors))
        if not parse_result.rows:
            raise BusinessException("未解析到有效数据行")

        created_count = 0
        updated_count = 0
        skipped_approved = 0
        create_audit = audit_create(current_user)
        update_audit = audit_update(current_user)

        try:
            for row in parse_result.rows:
                unit_price = quantize_unit_price(row.unit_price)
                existing = self.repository.get_by_location_year_part_number(
                    db, location, year, row.part_number
                )
                if existing:
                    if existing.audit_status == "approved":
                        skipped_approved += 1
                        continue
                    self.repository.update(
                        db,
                        existing.id,
                        BomAnnualPartPriceUpdate(
                            part_name=row.part_name,
                            unit=row.unit,
                            unit_price=unit_price,
                            supply_type=row.supply_type,
                            supplier_name=row.supplier_name,
                            remark=row.remark,
                        ),
                        updated_by=update_audit["updated_by"],
                    )
                    updated_count += 1
                else:
                    self.repository.create(
                        db,
                        BomAnnualPartPriceCreate(
                            location=location,
                            year=year,
                            part_number=row.part_number,
                            part_name=row.part_name,
                            unit=row.unit,
                            unit_price=unit_price,
                            supply_type=row.supply_type,
                            supplier_name=row.supplier_name,
                            remark=row.remark,
                        ),
                        created_by=create_audit["created_by"],
                        updated_by=create_audit["updated_by"],
                    )
                    created_count += 1
            register_existing(
                db,
                file_category="bom_annual_part_price",
                stored_relpath=archived_rel,
                original_filename=original_name,
                current_user=current_user,
                location=location,
                year=year,
            )
            db.commit()
        except Exception:
            db.rollback()
            raise

        warnings = list(parse_result.errors)
        if skipped_approved:
            warnings.append(f"跳过已审核 {skipped_approved} 条（未覆盖）")
        message = (
            f"导入完成：基地「{location}」{year} 年，"
            f"新增 {created_count} 条，更新 {updated_count} 条"
        )
        if skipped_approved:
            message += f"，跳过已审核 {skipped_approved} 条"
        if parse_result.errors:
            message += f"；有 {len(parse_result.errors)} 行解析跳过"

        logger.info(
            "%s file=%s archived=%s created=%s updated=%s skipped_approved=%s warnings=%s",
            message,
            original_name,
            archived_rel,
            created_count,
            updated_count,
            skipped_approved,
            len(warnings),
        )
        return BomAnnualPartPriceImportResultVO(
            location=location,
            year=year,
            created_count=created_count,
            updated_count=updated_count,
            total_count=created_count + updated_count,
            skipped_count=len(parse_result.errors) + skipped_approved,
            warnings=warnings,
            message=message,
        )
