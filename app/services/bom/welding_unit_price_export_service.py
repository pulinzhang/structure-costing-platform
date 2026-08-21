"""焊接单价 Excel 导出服务。"""

from __future__ import annotations

import logging
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.common.excel_workbook import resolve_worksheet
from app.common.money import Q4, excel_money_cell
from app.common.timezone import now
from app.repositories.bom.welding_unit_price_repository import BomWeldingUnitPriceRepository

logger = logging.getLogger("service.bom_welding_unit_price_export")


def _unit_price_cell(value: Any) -> Decimal | str:
    """导出单价：保持 Decimal，不经 float。"""
    return excel_money_cell(value, Q4)


def _dt(value: datetime | None) -> str:
    if value is None:
        return ""
    return f"{value:%Y-%m-%d %H:%M:%S}"


class BomWeldingUnitPriceExportService:
    """焊接单价 Excel 导出。"""

    def __init__(self) -> None:
        self.repository = BomWeldingUnitPriceRepository()

    def export(
        self,
        db: Session,
        *,
        location: str | None = None,
        year: int | None = None,
    ) -> tuple[bytes, str]:
        wb = Workbook()
        _title, count = self.write_to_workbook(
            wb, db, location=location, year=year
        )
        buffer = BytesIO()
        wb.save(buffer)
        content = buffer.getvalue()

        stamp = now().strftime("%Y%m%d_%H%M%S")
        parts = ["焊接单价"]
        if location:
            parts.append(location)
        if year:
            parts.append(f"{year}年")
        filename = "_".join(parts) + f"_{stamp}.xlsx"

        logger.info(
            "导出焊接单价 %s 条 location=%s year=%s -> %s",
            count,
            location,
            year,
            filename,
        )
        return content, filename

    def write_to_workbook(
        self,
        wb: Workbook,
        db: Session,
        *,
        location: str | None = None,
        year: int | None = None,
        sheet_title: str = "BOM焊接单价",
    ) -> tuple[str, int]:
        records = self.repository.list_for_export(
            db,
            location=location,
            year=year,
        )
        ws = resolve_worksheet(wb, sheet_title)

        header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        data_font = Font(name="微软雅黑", size=9)
        align_left = Alignment(horizontal="left", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )

        headers = [
            "基地",
            "年份",
            "名称",
            "单位",
            "费用",
            "备注",
            "创建人",
            "创建时间",
            "更新人",
            "更新时间",
        ]
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for row_idx, record in enumerate(records, start=2):
            values: list[Any] = [
                record.location,
                record.year,
                record.item_name,
                record.unit,
                _unit_price_cell(record.unit_price),
                record.remark or "",
                record.created_by or "",
                _dt(record.created_at),
                record.updated_by or "",
                _dt(record.updated_at),
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.font = data_font
                cell.border = thin_border
                if col in (2, 4):
                    cell.alignment = align_center
                elif col == 5:
                    cell.alignment = align_right
                    cell.number_format = "0.0000"
                else:
                    cell.alignment = align_left

        widths = [12, 10, 36, 10, 12, 20, 12, 20, 12, 20]
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:J{max(1, len(records) + 1)}"
        return ws.title, len(records)
