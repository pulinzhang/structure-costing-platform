"""零件用在哪 Excel：与页面路径逐级一致，不含供货总成汇总页。"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.common.bom.manufacturing_category import short_part_number_digits
from app.common.excel_workbook import excel_write_value, resolve_worksheet
from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.money import Q4, optional_format_money
from app.common.timezone import now
from app.schemas.dto.bom_where_used_dto import BomWhereUsedQuery
from app.schemas.vo.bom_where_used_vo import BomWhereUsedNodeVO, BomWhereUsedPathVO, BomWhereUsedVO
from app.schemas.vo.user_vo import UserVO
from app.services.bom.where_used.service import BomWhereUsedService

_LEVEL_LABEL = {0: "供货总成", 1: "一级", 2: "二级"}
_HEADERS = ["路径", "基地", "项目", "版本", "层级", "短码", "零件号", "零件名称", "用量"]
_COL_WIDTHS = [8, 12, 12, 12, 14, 10, 20, 28, 16]
_HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
_TITLE_FONT = Font(name="微软雅黑", size=15, bold=True, color="FFFFFF")
_META_FONT = Font(name="微软雅黑", size=9, color="455A64")
_DATA_FONT = Font(name="微软雅黑", size=9)
_HIT_FONT = Font(name="微软雅黑", size=9, bold=True)
_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_META_FILL = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
_HIT_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
_ZEBRA_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
_ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


class BomWhereUsedExportService:
    """按当前查询导出路径逐级明细。"""

    def __init__(self) -> None:
        self._search = BomWhereUsedService()

    def export(self, db: Session, query: BomWhereUsedQuery, exporter: UserVO) -> tuple[bytes, str]:
        result = self._search.search(db, query)
        workbook = Workbook()
        sheet = resolve_worksheet(workbook, "零件用在哪")
        self._write_sheet(sheet, result, query)
        stamp = now().strftime("%Y%m%d_%H%M%S")
        token = (query.query or "").strip()[:40]
        filename = f"零件用在哪_{token}_{stamp}.xlsx"
        return workbook_to_xlsx_bytes(workbook, exporter), filename

    def _write_sheet(
        self, sheet: Worksheet, result: BomWhereUsedVO, query: BomWhereUsedQuery
    ) -> None:
        last_col = len(_HEADERS)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
        title_cell = sheet.cell(1, 1, "零件用在哪")
        title_cell.font = _TITLE_FONT
        title_cell.fill = _HEADER_FILL
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[1].height = 28
        for col in range(1, last_col + 1):
            sheet.cell(1, col).fill = _HEADER_FILL

        token = (query.query or "").strip()
        summary = result.summary
        meta = (
            f"查询「{token}」    "
            f"命中 {summary.hit_count} 条 · 清单 {summary.snapshot_count} 份 · 项目 {summary.vehicle_model_count} 个"
        )
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        meta_cell = sheet.cell(2, 1, meta)
        meta_cell.font = _META_FONT
        meta_cell.fill = _META_FILL
        meta_cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[2].height = 20
        for col in range(1, last_col + 1):
            sheet.cell(2, col).fill = _META_FILL

        header_row = 4
        for col, title in enumerate(_HEADERS, start=1):
            cell = sheet.cell(header_row, col, title)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _ALIGN_CENTER
            cell.border = _THIN
        sheet.row_dimensions[header_row].height = 22

        row_no = header_row + 1
        for path_index, path in enumerate(result.paths, start=1):
            path_start = row_no
            zebra = path_index % 2 == 0
            for node_index, node in enumerate(path.nodes):
                is_hit = node_index == 0
                values = [
                    path_index,
                    path.location,
                    path.vehicle_model,
                    path.bom_version,
                    _level_label(node, is_hit),
                    short_part_number_digits(node.part_number) or node.part_number,
                    node.part_number,
                    node.part_name or "",
                    _qty_label(path, node, is_hit),
                ]
                fill = _HIT_FILL if is_hit else (_ZEBRA_FILL if zebra else None)
                font = _HIT_FONT if is_hit else _DATA_FONT
                for col, value in enumerate(values, start=1):
                    cell = sheet.cell(row_no, col, excel_write_value(value))
                    cell.font = font
                    cell.border = _THIN
                    if fill is not None:
                        cell.fill = fill
                    if col in {1, 5, 6}:
                        cell.alignment = _ALIGN_CENTER
                    elif col == 9:
                        cell.alignment = _ALIGN_RIGHT
                    else:
                        cell.alignment = _ALIGN_LEFT
                sheet.row_dimensions[row_no].height = 20
                row_no += 1
            path_end = row_no - 1
            if path_end > path_start:
                for col in range(1, 5):
                    sheet.merge_cells(
                        start_row=path_start,
                        start_column=col,
                        end_row=path_end,
                        end_column=col,
                    )
                    sheet.cell(path_start, col).alignment = _ALIGN_CENTER

        for col, width in enumerate(_COL_WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = "A5"
        sheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(last_col)}{max(header_row, row_no - 1)}"
        )
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.3
        sheet.page_margins.right = 0.3
        sheet.page_margins.top = 0.35
        sheet.page_margins.bottom = 0.35
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 90
        sheet.print_title_rows = "1:4"


def _level_label(node: BomWhereUsedNodeVO, is_hit: bool) -> str:
    base = _LEVEL_LABEL.get(int(node.level), f"L{node.level}")
    return f"命中 · {base}" if is_hit else base


def _qty_label(path: BomWhereUsedPathVO, node: BomWhereUsedNodeVO, is_hit: bool) -> str:
    if is_hit:
        qty = optional_format_money(path.path_qty, Q4, 4, allow_negative=True) or "0"
        return f"单车 {qty}"
    if node.qty_text:
        return node.qty_text
    if node.qty is None:
        return ""
    return optional_format_money(node.qty, Q4, 4, allow_negative=True) or ""
