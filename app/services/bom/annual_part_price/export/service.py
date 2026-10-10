"""年度零件价 Excel 导出：核心列对齐导入 Parser；另含基地/年份/审核审计列。"""

from __future__ import annotations

import logging
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.common.audit import audit_status_label
from app.common.excel_workbook import ExcelCellValue, resolve_worksheet
from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.money import Q4, excel_money_cell
from app.common.timezone import now
from app.parsers.bom.annual_part_price_parser import (
    HEADER_CONVERSION_FACTOR,
    HEADER_PART_NAME,
    HEADER_PART_NUMBER,
    HEADER_PURCHASE_UNIT,
    HEADER_PURCHASE_UNIT_PRICE,
    HEADER_REMARK,
    HEADER_SPECIFICATION,
    HEADER_SUPPLIER,
    HEADER_SUPPLY_TYPE,
    HEADER_UNIT,
    HEADER_UNIT_PRICE,
)
from app.repositories.bom.annual_part_price_repository import (
    BomAnnualPartPriceRepository,
)
from app.schemas.vo.user_vo import UserVO

logger = logging.getLogger("service.bom_annual_part_price_export")


def _unit_price_cell(value: Decimal | None) -> Decimal | str:
    """导出单价：保持 Decimal，不经 float。"""
    return excel_money_cell(value, Q4)


def _dt(value: datetime | None) -> str:
    if value is None:
        return ""
    return f"{value:%Y-%m-%d %H:%M:%S}"


class BomAnnualPartPriceExportService:
    """年度零件价 Excel 导出（可再导入）。"""

    def __init__(self) -> None:
        self.repository = BomAnnualPartPriceRepository()

    def export(
        self,
        db: Session,
        *,
        location: str | None = None,
        year: int | None = None,
        current_user: UserVO,
    ) -> tuple[bytes, str]:
        wb = Workbook()
        _title, count = self.write_to_workbook(wb, db, location=location, year=year)
        content = workbook_to_xlsx_bytes(wb, current_user)

        stamp = now().strftime("%Y%m%d_%H%M%S")
        parts = ["年度零件价"]
        if location:
            parts.append(location)
        if year:
            parts.append(f"{year}年")
        filename = "_".join(parts) + f"_{stamp}.xlsx"

        logger.info(
            "导出年度零件价 %s 条 location=%s year=%s -> %s",
            count,
            location,
            year,
            filename,
        )
        return content, filename

    def write_to_workbook(
        self,
        wb: Workbook,
        db: Session,
        *,
        location: str | None = None,
        year: int | None = None,
        sheet_title: str = "年度零件价",
    ) -> tuple[str, int]:
        records = self.repository.list_for_export(
            db,
            location=location,
            year=year,
        )
        ws = resolve_worksheet(wb, sheet_title)

        header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        data_font = Font(name="微软雅黑", size=9)
        align_left = Alignment(horizontal="left", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )

        headers = [
            "基地",
            "年份",
            HEADER_PART_NUMBER,
            HEADER_PART_NAME,
            HEADER_SPECIFICATION,
            HEADER_PURCHASE_UNIT,
            HEADER_PURCHASE_UNIT_PRICE,
            HEADER_CONVERSION_FACTOR,
            HEADER_UNIT,
            HEADER_UNIT_PRICE,
            HEADER_SUPPLY_TYPE,
            HEADER_SUPPLIER,
            HEADER_REMARK,
            "审核状态",
            "创建人",
            "创建时间",
            "更新人",
            "更新时间",
        ]
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for row_idx, record in enumerate(records, start=2):
            values: list[ExcelCellValue] = [
                record.location,
                record.year,
                record.part_number,
                record.part_name or "",
                record.specification or "",
                record.purchase_unit or "",
                _unit_price_cell(record.purchase_unit_price)
                if record.purchase_unit_price is not None
                else "",
                record.conversion_factor,
                record.unit,
                _unit_price_cell(record.unit_price),
                record.supply_type,
                record.supplier_name or "",
                record.remark or "",
                audit_status_label(record.audit_status),
                record.created_by or "",
                _dt(record.created_at),
                record.updated_by or "",
                _dt(record.updated_at),
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.font = data_font
                cell.border = thin_border
                if col in (2, 6, 9, 11, 14):
                    cell.alignment = align_center
                elif col in (7, 10):
                    cell.alignment = align_right
                    cell.number_format = "0.0000"
                elif col == 8:
                    cell.alignment = align_center
                    cell.number_format = "0.######"
                else:
                    cell.alignment = align_left

        widths = [12, 10, 18, 24, 16, 12, 12, 12, 8, 12, 12, 16, 20, 10, 12, 20, 12, 20]
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        ws.freeze_panes = "A2"
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{max(1, len(records) + 1)}"
        return ws.title, len(records)
