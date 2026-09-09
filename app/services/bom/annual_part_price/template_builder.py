"""年度零件价导入模板 Excel 排版（不含业务查询）。"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.parsers.bom.annual_part_price_parser import (
    DEFAULT_CONVERSION_FACTOR,
    HEADER_CONVERSION_FACTOR,
    HEADER_PART_NAME,
    HEADER_PART_NUMBER,
    HEADER_PURCHASE_UNIT,
    HEADER_PURCHASE_UNIT_PRICE,
    HEADER_REMARK,
    HEADER_SPECIFICATION,
    HEADER_SUPPLIER,
    HEADER_SUPPLY_TYPE,
    HEADER_UNIT,
    HEADER_UNIT_PRICE,
)

_SHEET_TITLE = "年度零件价"
_NOTES_SHEET_TITLE = "填写说明"
_HEADERS = (
    HEADER_PART_NUMBER,
    HEADER_PART_NAME,
    HEADER_SPECIFICATION,
    HEADER_PURCHASE_UNIT,
    HEADER_PURCHASE_UNIT_PRICE,
    HEADER_CONVERSION_FACTOR,
    HEADER_UNIT,
    HEADER_UNIT_PRICE,
    HEADER_SUPPLY_TYPE,
    HEADER_SUPPLIER,
    HEADER_REMARK,
)
_COLUMN_WIDTHS = (18, 24, 16, 12, 12, 12, 10, 12, 12, 16, 20)
_PURCHASE_PRICE_COL = 5
_FACTOR_COL = 6
_UNIT_PRICE_COL = 8
_CENTER_COLS = frozenset({4, 7, 9})

_TEMPLATE_ZERO_PRICE = Decimal("0.0000")

_FILL_NOTES = (
    "「单位」「单价」是核价口径：成品核价用 单价 × BOM 用量。用量来自材料清单，单位须与 BOM 一致（pcs、mm 等）。不要把采购包装单位（打、瓶）填进「单位」。",
    "报价常按包装（打、瓶），BOM 按消耗（个、mm）。量纲不同时填写采购单位、采购单价、转换系数，系统写入核价单价，避免瓶价直接乘 mm。",
    "转换系数须大于 0，默认 1，最多 6 位小数。含义：1 个采购单位 = 转换系数 个核价单位。核价单价 = 采购单价 ÷ 转换系数。不要填 1/12 这类分数。",
    "例：120 元/打、BOM 按个用 → 系数填 12（1 打=12 个），核价单价 10 元/个。例：1 英寸=25.4 mm → 系数填 25.4。例：1 平方米=10000 平方厘米 → 系数填 10000。",
    "同单位只改「单价」即可，采购单价留空、系数保持 1。不同单位：改采购单位、填采购单价和转换系数；「单价」可仍为 0，有采购单价时系统按公式覆盖。",
    "「规格」从 BOM 料厚/规格原文预填，可改；不参与取价（仍按零件号匹配）。",
    "上传时选择基地+年度。本文件填完即可导入，不必删除「填写说明」页。",
)


@dataclass(frozen=True)
class AnnualPartPriceTemplateRow:
    """导入模板一行（与解析表头列顺序一致）。"""

    part_number: str
    part_name: str
    specification: str = ""
    purchase_unit: str = ""
    purchase_unit_price: Decimal | None = None
    conversion_factor: Decimal = DEFAULT_CONVERSION_FACTOR
    unit: str = "pcs"
    unit_price: Decimal = _TEMPLATE_ZERO_PRICE
    supply_type: str = ""
    supplier_name: str = ""
    remark: str = ""


def _write_notes(sheet: Worksheet, lines: tuple[str, ...]) -> None:
    sheet["A1"] = "填写说明"
    sheet["A1"].font = Font(name="微软雅黑", size=12, bold=True, color="C00000")
    for idx, line in enumerate(lines, start=3):
        sheet.cell(row=idx, column=1, value=f"{idx - 2}. {line}")
        sheet.cell(row=idx, column=1).font = Font(name="微软雅黑", size=10)
        sheet.cell(row=idx, column=1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[idx].height = 36
    sheet.column_dimensions["A"].width = 96


def render_import_template(rows: list[AnnualPartPriceTemplateRow]) -> bytes:
    """按解析器约定列写出导入模板工作簿。"""
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = _SHEET_TITLE

    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    data_font = Font(name="微软雅黑", size=9)
    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )

    for col, title in enumerate(_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, sample in enumerate(rows, start=2):
        values: list[str | Decimal | None] = [
            sample.part_number,
            sample.part_name,
            sample.specification,
            sample.purchase_unit,
            sample.purchase_unit_price,
            sample.conversion_factor,
            sample.unit,
            sample.unit_price,
            sample.supply_type,
            sample.supplier_name,
            sample.remark,
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_idx, column=col, value=value)
            cell.font = data_font
            cell.border = thin_border
            if col in (_UNIT_PRICE_COL, _PURCHASE_PRICE_COL):
                cell.alignment = align_right
                if value is not None:
                    cell.number_format = "0.0000"
            elif col == _FACTOR_COL:
                cell.alignment = align_center
                if value is not None:
                    cell.number_format = "0.######"
            elif col in _CENTER_COLS:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    for idx, width in enumerate(_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "A2"

    notes = workbook.create_sheet(_NOTES_SHEET_TITLE)
    _write_notes(notes, _FILL_NOTES)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
