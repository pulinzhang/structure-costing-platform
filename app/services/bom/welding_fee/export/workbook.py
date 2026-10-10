"""焊接费导出 — openpyxl 渲染（不访问 DB）。"""

from __future__ import annotations

from decimal import Decimal

from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.common.excel_workbook import excel_write_value
from app.common.timezone import now
from app.models.bom import BomWeldingFee, BomWeldingFeeItem
from app.services.bom.welding_fee.export.columns import (
    FILTER_H,
    GROUP_H,
    HEADER_H,
    MONEY_FORMAT,
    NOTE_H,
    TITLE_H,
    TOTAL_H,
    WeldingSheetLayout,
)
from app.services.bom.welding_fee.export.styles import (
    SheetStyle,
    audit_tone,
    fill,
    font,
    match_tone,
    status_tone,
)
from app.services.bom.welding_fee.export.values import (
    _count_label,
    detail_values,
    main_pairs,
    match_status_of,
    row_height,
)


def write_banner(
    ws: Worksheet,
    *,
    layout: WeldingSheetLayout,
    filter_text: str,
    records: list[BomWeldingFee],
) -> None:
    indent = Alignment(horizontal="left", vertical="center", indent=1)
    title = "焊接费" if layout.omit_scope else "焊接费核对表"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=layout.col_count)
    c1 = ws.cell(row=1, column=1, value=title)
    c1.font = font(size=15, bold=True, color="FFFFFF")
    c1.fill = fill("1A237E")
    c1.alignment = indent
    ws.row_dimensions[1].height = TITLE_H

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=layout.col_count)
    c2 = ws.cell(
        row=2,
        column=1,
        value=(
            f"{filter_text} · {_count_label(records)}"
            f"　　导出 {now().strftime('%Y-%m-%d %H:%M')}"
        ),
    )
    c2.font = font(size=9, color="455A64")
    c2.fill = fill("E8EAF6")
    c2.alignment = Alignment(
        horizontal="left",
        vertical="center",
        indent=1,
        wrap_text=True,
    )
    ws.row_dimensions[2].height = FILTER_H

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=layout.col_count)
    note = (
        "说明：范围已在表头锁定，表内不再重复年度/基地/项目/版本；"
        "左侧主表按节点合并，右侧为工艺明细。行费 = 计费量 × 分年单价。"
        "表头与末行为本表合计；末行公式不进筛选、不随筛选变化。"
        if layout.omit_scope
        else (
            "说明：左侧主表按节点合并，含各层级（见「层级」列）；"
            "右侧为工艺明细。行费 = 计费量 × 分年单价。"
            "表头与末行为本表合计；末行公式不进筛选、不随筛选变化。"
        )
    )
    c3 = ws.cell(row=3, column=1, value=note)
    c3.font = font(size=9, color="B71C1C")
    c3.fill = fill("FFF3E0")
    c3.alignment = indent
    ws.row_dimensions[3].height = NOTE_H


def write_headers(ws: Worksheet, *, layout: WeldingSheetLayout, style: SheetStyle) -> int:
    group_row = 4
    header_row = 5
    white = font(size=9, bold=True, color="FFFFFF")
    header_font = font(bold=True, color="FFFFFF")
    main_cols = layout.main_cols
    col_count = layout.col_count

    ws.merge_cells(
        start_row=group_row,
        start_column=1,
        end_row=group_row,
        end_column=main_cols,
    )
    ws.merge_cells(
        start_row=group_row,
        start_column=main_cols + 1,
        end_row=group_row,
        end_column=col_count,
    )
    main_group = ws.cell(row=group_row, column=1, value="节点主表")
    main_group.font = white
    main_group.fill = fill("37474F")
    main_group.alignment = style.center
    for col in range(1, main_cols + 1):
        cell = ws.cell(row=group_row, column=col)
        cell.fill = fill("37474F")
        cell.border = style.thin if col < main_cols else style.split_main

    detail_group = ws.cell(row=group_row, column=main_cols + 1, value="工艺明细 / 计费")
    detail_group.font = white
    detail_group.fill = fill("B71C1C")
    detail_group.alignment = style.center
    for col in range(main_cols + 1, col_count + 1):
        cell = ws.cell(row=group_row, column=col)
        cell.fill = fill("B71C1C")
        cell.border = style.split_detail if col == main_cols + 1 else style.thin
    ws.row_dimensions[group_row].height = GROUP_H

    fee_cols = layout.fee_cols
    for col, title in enumerate(layout.headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = header_font
        cell.alignment = style.center
        if col <= main_cols:
            cell.fill = fill("EF6C00") if col in fee_cols else fill("546E7A")
            cell.border = style.thin if col < main_cols else style.split_main
        elif col in fee_cols:
            cell.fill = fill("EF6C00")
            cell.border = style.thin
        else:
            cell.fill = fill("C62828")
            cell.border = style.split_detail if col == main_cols + 1 else style.thin
    ws.row_dimensions[header_row].height = HEADER_H
    return header_row


def write_data_row(
    ws: Worksheet,
    row_idx: int,
    record: BomWeldingFee,
    *,
    layout: WeldingSheetLayout,
    style: SheetStyle,
    seq_no: int,
    item: BomWeldingFeeItem | None,
    write_main: bool,
    line_index: int,
    main_fill: PatternFill,
    display_bom_version: str | None = None,
) -> None:
    col_index = layout.col
    main_cols = layout.main_cols
    detail_fill = fill("FFFFFF") if line_index % 2 == 0 else fill("FFEBEE")
    missing = int(record.missing_price_count or 0) > 0
    calc_status = record.calc_status or ""
    match_status = match_status_of(item)

    if write_main:
        values = [
            value
            for name, value in main_pairs(record, seq_no=seq_no, bom_version=display_bom_version)
            if name in col_index
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=excel_write_value(value))
            cell.fill = main_fill
            cell.border = style.thin if col < main_cols else style.split_main
            _style_main_cell(
                cell,
                col,
                layout=layout,
                style=style,
                calc_status=calc_status,
                audit_status=record.audit_status,
                missing=missing,
            )

    for offset, value in enumerate(detail_values(item)):
        col = main_cols + 1 + offset
        cell = ws.cell(row=row_idx, column=col, value=excel_write_value(value))
        cell.fill = detail_fill
        cell.border = style.split_detail if col == main_cols + 1 else style.thin
        _style_detail_cell(cell, col, layout=layout, style=style, match_status=match_status)


def _style_main_cell(
    cell: Cell,
    col: int,
    *,
    layout: WeldingSheetLayout,
    style: SheetStyle,
    calc_status: str,
    audit_status: str | None,
    missing: bool,
) -> None:
    col_index = layout.col
    if col == col_index["序号"]:
        cell.font = font(size=11, bold=True, color="37474F")
        cell.alignment = style.center
    elif col == col_index["零件号"]:
        cell.font = font(bold=True, color="1565C0")
        cell.alignment = style.left
    elif col == col_index["短代码"]:
        cell.font = font(size=11, bold=True, color="1565C0")
        cell.alignment = style.center
    elif col == col_index["焊接信息原文"]:
        cell.font = font(size=9, color="607D8B")
        cell.alignment = style.left
    elif col == col_index["系统计算"]:
        cell.font = font(size=11, bold=True, color="546E7A")
        cell.alignment = style.right
    elif col == col_index["其它费用"]:
        cell.font = font(size=11, color="37474F")
        cell.alignment = style.right
    elif col == col_index["确认单价"]:
        cell.font = font(size=11, bold=True, color="E65100")
        cell.alignment = style.right
        cell.fill = fill("FFF3E0")
    elif col == col_index["备注"]:
        cell.font = font(size=9, color="607D8B")
        cell.alignment = style.left
    elif col in layout.count_cols:
        cell.font = (
            font(bold=True, color="C62828") if (col == col_index["缺单价"] and missing) else font()
        )
        cell.alignment = style.center
    elif col == col_index["计算状态"]:
        cell.font, cell.fill = status_tone(calc_status)
        cell.alignment = style.center
    elif col == col_index["审核状态"]:
        cell.font, cell.fill = audit_tone(audit_status)
        cell.alignment = style.center
    else:
        cell.font = font()
        cell.alignment = style.left if col in layout.left_main_cols else style.center


def _style_detail_cell(
    cell: Cell,
    col: int,
    *,
    layout: WeldingSheetLayout,
    style: SheetStyle,
    match_status: str,
) -> None:
    col_index = layout.col
    if col == col_index["行费"]:
        cell.font = font(size=11, bold=True, color="E65100")
        cell.alignment = style.right
        cell.fill = fill("FFF3E0")
    elif col in {col_index["过程量"], col_index["单价"]}:
        cell.font = font()
        cell.alignment = style.right
    elif col == col_index["匹配状态"]:
        cell.font, cell.fill = match_tone(match_status)
        cell.alignment = style.center
    else:
        cell.font = font()
        cell.alignment = style.left if col == col_index["标准工艺"] else style.center


def merge_main(
    ws: Worksheet,
    start_row: int,
    end_row: int,
    main_fill: PatternFill,
    *,
    layout: WeldingSheetLayout,
    style: SheetStyle,
) -> None:
    for col in range(1, layout.main_cols + 1):
        ws.merge_cells(
            start_row=start_row,
            start_column=col,
            end_row=end_row,
            end_column=col,
        )
        cell = ws.cell(row=start_row, column=col)
        if col == layout.col["确认单价"]:
            cell.fill = fill("FFF3E0")
        else:
            cell.fill = main_fill
        cell.border = style.thin if col < layout.main_cols else style.split_main


def write_total_row(
    ws: Worksheet,
    *,
    layout: WeldingSheetLayout,
    style: SheetStyle,
    header_row: int,
    last_data: int,
) -> None:
    total_row = last_data + 1
    has_data = last_data > header_row
    data_start = header_row + 1
    kpi_cols = {layout.col["确认单价"], layout.col["行费"]}
    for col in range(1, layout.col_count + 1):
        cell = ws.cell(row=total_row, column=col, value=None)
        cell.font = font(size=10, bold=True)
        cell.fill = fill("ECEFF1")
        if col in layout.fee_cols:
            cell.alignment = style.right
            cell.number_format = MONEY_FORMAT
            cell.border = style.thin
        elif col == layout.main_cols:
            cell.alignment = style.center
            cell.border = style.split_main
        elif col == layout.main_cols + 1:
            cell.alignment = style.center
            cell.border = style.split_detail
        else:
            cell.alignment = style.center
            cell.border = style.thin
    label = ws.cell(row=total_row, column=1, value="合计")
    label.font = font(size=10, bold=True)
    label.fill = fill("ECEFF1")
    label.alignment = style.center
    for col in sorted(layout.fee_cols):
        letter = get_column_letter(col)
        cell = ws.cell(row=total_row, column=col)
        if has_data:
            cell.value = f"=SUM({letter}{data_start}:{letter}{last_data})"
        else:
            cell.value = Decimal("0.0000")
        cell.font = font(size=10, bold=True, color="B71C1C")
        cell.fill = fill("FFCDD2") if col in kpi_cols else fill("FFF3E0")
        cell.alignment = style.right
        cell.number_format = MONEY_FORMAT
        if col == layout.main_cols:
            cell.border = style.split_main
        elif col == layout.main_cols + 1:
            cell.border = style.split_detail
        else:
            cell.border = style.thin
    ws.row_dimensions[total_row].height = TOTAL_H


def fill_sheet(
    ws: Worksheet,
    records: list[BomWeldingFee],
    *,
    layout: WeldingSheetLayout,
    style: SheetStyle,
    filter_text: str,
    bom_version_by_id: dict[int, str] | None = None,
) -> int:
    write_banner(ws, layout=layout, filter_text=filter_text, records=records)
    header_row = write_headers(ws, layout=layout, style=style)

    row_idx = header_row + 1
    for seq_no, record in enumerate(records, start=1):
        items = sorted(
            list(record.items or []),
            key=lambda item: (item.line_no or 0, item.id or 0),
        )
        main_fill = fill("ECEFF1") if seq_no % 2 else fill("F5F7FA")
        display_bom_version = None
        record_id = getattr(record, "id", None)
        if record_id is not None:
            display_bom_version = (bom_version_by_id or {}).get(int(record_id))
        if not items:
            write_data_row(
                ws,
                row_idx,
                record,
                layout=layout,
                style=style,
                seq_no=seq_no,
                item=None,
                write_main=True,
                line_index=0,
                main_fill=main_fill,
                display_bom_version=display_bom_version,
            )
            ws.row_dimensions[row_idx].height = row_height(record)
            row_idx += 1
            continue

        group_start = row_idx
        for index, item in enumerate(items):
            write_data_row(
                ws,
                row_idx,
                record,
                layout=layout,
                style=style,
                seq_no=seq_no,
                item=item,
                write_main=(index == 0),
                line_index=index,
                main_fill=main_fill,
                display_bom_version=display_bom_version,
            )
            ws.row_dimensions[row_idx].height = row_height(record)
            row_idx += 1
        if row_idx - 1 > group_start:
            merge_main(ws, group_start, row_idx - 1, main_fill, layout=layout, style=style)

    last_data = max(header_row, row_idx - 1)
    write_total_row(
        ws,
        layout=layout,
        style=style,
        header_row=header_row,
        last_data=last_data,
    )
    for idx, width in enumerate(layout.widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(layout.col_count)}{last_data}"
    ws.print_title_rows = f"1:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    return last_data
