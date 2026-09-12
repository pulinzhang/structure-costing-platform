"""焊接费导出 — 色板、字体、状态着色。"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_FONT = "微软雅黑"


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def font(*, size: int = 10, bold: bool = False, color: str = "263238") -> Font:
    return Font(name=_FONT, size=size, bold=bold, color=color)


def border(*, right: Side | None = None, left: Side | None = None) -> Border:
    thin = Side(style="thin", color="CFD8DC")
    return Border(
        left=left or thin,
        right=right or thin,
        top=thin,
        bottom=thin,
    )


@dataclass(frozen=True)
class SheetStyle:
    left: Alignment
    center: Alignment
    right: Alignment
    thin: Border
    split_main: Border
    split_detail: Border


def build_sheet_style() -> SheetStyle:
    return SheetStyle(
        left=Alignment(horizontal="left", vertical="center", wrap_text=True),
        center=Alignment(horizontal="center", vertical="center", wrap_text=True),
        right=Alignment(horizontal="right", vertical="center"),
        thin=border(),
        split_main=border(right=Side(style="medium", color="B71C1C")),
        split_detail=border(left=Side(style="medium", color="B71C1C")),
    )


def status_tone(calc_status: str) -> tuple[Font, PatternFill]:
    if calc_status == "ok":
        return font(bold=True, color="2E7D32"), fill("E8F5E9")
    if calc_status == "partial":
        return font(bold=True, color="EF6C00"), fill("FFF3E0")
    if calc_status == "failed":
        return font(bold=True, color="C62828"), fill("FFEBEE")
    return font(), fill("FFFFFF")


def audit_tone(audit_status: str | None) -> tuple[Font, PatternFill]:
    if audit_status == "approved":
        return font(bold=True, color="2E7D32"), fill("E8F5E9")
    if audit_status == "pending":
        return font(bold=True, color="EF6C00"), fill("FFF8E1")
    if audit_status == "rejected":
        return font(bold=True, color="C62828"), fill("FFEBEE")
    return font(), fill("FFFFFF")


def match_tone(match_status: str) -> tuple[Font, PatternFill]:
    if match_status == "matched":
        return font(bold=True, color="2E7D32"), fill("E8F5E9")
    if match_status == "missing_price":
        return font(bold=True, color="EF6C00"), fill("FFF3E0")
    if match_status == "unmapped":
        return font(bold=True, color="607D8B"), fill("ECEFF1")
    return font(), fill("FFFFFF")
