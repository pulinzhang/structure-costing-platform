"""BOM 焊接费 Excel 导出（核对风排版）。"""

from __future__ import annotations

import logging
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.common.audit import audit_status_label
from app.common.bom.manufacturing_category import (
    short_part_number_digits,
    short_weld_part_number,
)
from app.common.excel_workbook import resolve_worksheet
from app.common.export_filters import require_at_least_one_export_filter
from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.money import Q4, excel_money_cell
from app.common.timezone import now
from app.models.bom import BomWeldingFee
from app.repositories.bom.welding_fee_repository import BomWeldingFeeRepository
from app.repositories.bom.welding_process_breakdown_repository import (
    BomWeldingProcessBreakdownRepository,
)
from app.schemas.dto.bom_welding_fee_dto import resolve_welding_fee_levels
from app.schemas.vo.user_vo import UserVO

logger = logging.getLogger("service.bom.welding_fee.export")

_FONT = "微软雅黑"

_MATCH_LABEL = {
    "matched": "已匹配",
    "missing_price": "缺单价",
    "unmapped": "未映射",
}
_CALC_LABEL = {
    "ok": "正常",
    "partial": "缺单价",
    "failed": "失败",
}

# 主表列（纵向合并）+ 明细分隔
MAIN_COLS = 19
COL_COUNT = 26
HEADERS = [
    "序号",
    "年度",
    "基地",
    "项目",
    "版本",
    "层级",
    "零件号",
    "短代码",
    "零件名称",
    "焊接信息原文",
    "工艺数",
    "命中",
    "缺单价",
    "系统计算",
    "其它费用",
    "确认单价",
    "备注",
    "计算状态",
    "审核状态",
    "标准工艺",
    "过程量",
    "过程量单位",
    "单价",
    "单价单位",
    "行费",
    "匹配状态",
]
WIDTHS = [
    6,
    8,
    10,
    10,
    10,
    8,
    16,
    8,
    18,
    36,
    8,
    8,
    8,
    14,
    12,
    14,
    24,
    10,
    10,
    18,
    10,
    10,
    10,
    10,
    10,
    10,
]

assert len(HEADERS) == COL_COUNT
assert len(WIDTHS) == COL_COUNT
COL = {name: i for i, name in enumerate(HEADERS, start=1)}
FEE_COLS = frozenset({COL["系统计算"], COL["其它费用"], COL["确认单价"], COL["行费"]})
COUNT_COLS = frozenset({COL["工艺数"], COL["命中"], COL["缺单价"]})
LEFT_MAIN_COLS = frozenset({COL["基地"], COL["项目"], COL["版本"], COL["零件名称"], COL["备注"]})
_SCOPE_HEADERS = frozenset({"年度", "基地", "项目", "版本"})
_MAIN_HEADERS = tuple(HEADERS[:MAIN_COLS])
_DETAIL_HEADERS = tuple(HEADERS[MAIN_COLS:])
_FEE_HEADER_NAMES = ("系统计算", "其它费用", "确认单价", "行费")
_COUNT_HEADER_NAMES = ("工艺数", "命中", "缺单价")
_LEFT_MAIN_HEADER_NAMES = ("基地", "项目", "版本", "零件名称", "备注")
_WIDTH_BY_HEADER = dict(zip(HEADERS, WIDTHS, strict=True))


class WeldingSheetLayout:
    """完整列或去掉范围列（采购册范围已在表头）。"""

    def __init__(self, *, omit_scope: bool) -> None:
        main = (
            [name for name in _MAIN_HEADERS if name not in _SCOPE_HEADERS]
            if omit_scope
            else list(_MAIN_HEADERS)
        )
        self.headers = main + list(_DETAIL_HEADERS)
        self.widths = [_WIDTH_BY_HEADER[name] for name in self.headers]
        self.main_cols = len(main)
        self.col_count = len(self.headers)
        self.col = {name: index for index, name in enumerate(self.headers, start=1)}
        self.fee_cols = frozenset(self.col[name] for name in _FEE_HEADER_NAMES)
        self.count_cols = frozenset(self.col[name] for name in _COUNT_HEADER_NAMES)
        self.left_main_cols = frozenset(
            self.col[name] for name in _LEFT_MAIN_HEADER_NAMES if name in self.col
        )
        self.omit_scope = omit_scope


TITLE_H = 28
FILTER_H = 20
NOTE_H = 20
GROUP_H = 18
HEADER_H = 22
DATA_H = 26


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(*, size: int = 10, bold: bool = False, color: str = "263238") -> Font:
    return Font(name=_FONT, size=size, bold=bold, color=color)


def _border(*, right: Side | None = None, left: Side | None = None) -> Border:
    thin = Side(style="thin", color="CFD8DC")
    return Border(
        left=left or thin,
        right=right or thin,
        top=thin,
        bottom=thin,
    )


def _num(value: Any) -> Any:
    return excel_money_cell(value, Q4)


def _level_label(level: int | None) -> str:
    if level is None:
        return ""
    return f"L{int(level)}"


def part_short_code(*, part_number: str | None, level: int | None) -> str:
    """L0 末四位；分总成带 -N 时保留字母尾（同拆解列表）。"""
    if int(level or 0) == 0:
        return short_part_number_digits(part_number) or ""
    return short_weld_part_number(part_number) or ""


class BomWeldingFeeExportService:
    """按筛选导出焊接费（主表合并 + 工艺明细分行）。"""

    def __init__(self) -> None:
        self.fee_repo = BomWeldingFeeRepository()
        self.breakdown_repo = BomWeldingProcessBreakdownRepository()
        self._left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        self._center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self._right = Alignment(horizontal="right", vertical="center")
        self._thin = _border()
        self._split_main = _border(right=Side(style="medium", color="B71C1C"))
        self._split_detail = _border(left=Side(style="medium", color="B71C1C"))
        self._layout = WeldingSheetLayout(omit_scope=False)

    def write_to_workbook(
        self,
        wb: Workbook,
        db: Session,
        *,
        year: int | None = None,
        location: str | None = None,
        vehicle_model: str | None = None,
        bom_version: str | None = None,
        part_number: str | None = None,
        audit_status: str | None = None,
        calc_status: str | None = None,
        level: int | None = None,
        levels: list[int] | None = None,
        ids: list[int] | None = None,
        sheet_title: str = "焊接费",
        omit_scope_columns: bool = False,
    ) -> tuple[str, int]:
        self._layout = WeldingSheetLayout(omit_scope=omit_scope_columns)
        resolved_levels = resolve_welding_fee_levels(level=level, levels=levels)
        records = self.fee_repo.list_for_export(
            db,
            year=year,
            location=location,
            vehicle_model=vehicle_model,
            bom_version=bom_version,
            part_number=part_number,
            audit_status=audit_status,
            calc_status=calc_status,
            levels=resolved_levels,
            ids=ids,
        )
        self._fill_missing_weld_info(db, records)

        ws = resolve_worksheet(wb, sheet_title)

        filter_text = self._filter_text(
            year=year,
            location=location,
            vehicle_model=vehicle_model,
            bom_version=bom_version,
            part_number=part_number,
            audit_status=audit_status,
            calc_status=calc_status,
            levels=resolved_levels,
            ids=ids,
        )
        self._write_banner(ws, filter_text=filter_text, fee_count=len(records))
        header_row = self._write_headers(ws)

        row_idx = header_row + 1
        for seq_no, record in enumerate(records, start=1):
            items = sorted(
                list(record.items or []),
                key=lambda item: (item.line_no or 0, item.id or 0),
            )
            main_fill = _fill("ECEFF1") if seq_no % 2 else _fill("F5F7FA")
            if not items:
                self._write_data_row(
                    ws,
                    row_idx,
                    record,
                    seq_no=seq_no,
                    item=None,
                    write_main=True,
                    line_index=0,
                    main_fill=main_fill,
                )
                ws.row_dimensions[row_idx].height = self._row_height(record)
                row_idx += 1
                continue

            group_start = row_idx
            for i, item in enumerate(items):
                self._write_data_row(
                    ws,
                    row_idx,
                    record,
                    seq_no=seq_no,
                    item=item,
                    write_main=(i == 0),
                    line_index=i,
                    main_fill=main_fill,
                )
                ws.row_dimensions[row_idx].height = self._row_height(record)
                row_idx += 1
            if row_idx - 1 > group_start:
                self._merge_main(ws, group_start, row_idx - 1, main_fill)

        last_data = max(header_row, row_idx - 1)
        layout = self._layout
        for idx, width in enumerate(layout.widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(layout.col_count)}{last_data}"
        ws.print_title_rows = f"1:{header_row}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.3
        ws.page_margins.right = 0.3
        ws.page_margins.top = 0.35
        ws.page_margins.bottom = 0.35
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90

        return ws.title, len(records)

    def export(
        self,
        db: Session,
        *,
        year: int | None = None,
        location: str | None = None,
        vehicle_model: str | None = None,
        bom_version: str | None = None,
        part_number: str | None = None,
        audit_status: str | None = None,
        calc_status: str | None = None,
        level: int | None = None,
        levels: list[int] | None = None,
        current_user: UserVO,
    ) -> tuple[bytes, str]:
        require_at_least_one_export_filter(
            year,
            location,
            vehicle_model,
            bom_version,
            part_number,
            audit_status,
            calc_status,
        )
        wb = Workbook()
        _title, fee_count = self.write_to_workbook(
            wb,
            db,
            year=year,
            location=location,
            vehicle_model=vehicle_model,
            bom_version=bom_version,
            part_number=part_number,
            audit_status=audit_status,
            calc_status=calc_status,
            level=level,
            levels=levels,
        )
        # header_row is always 5 (banner 1–3 + group 4 + headers 5)
        export_row_count = max(0, wb.worksheets[0].max_row - 5)

        content = workbook_to_xlsx_bytes(wb, current_user)

        scope_parts = [
            str(year) if year else None,
            location,
            vehicle_model,
            bom_version,
        ]
        scope = "_".join(p for p in scope_parts if p) or "全部"
        stamp = now().strftime("%Y%m%d_%H%M%S")
        filename = f"焊接费_{scope}_{stamp}.xlsx"
        logger.info(
            "导出焊接费 %s 条主表 / %s 行明细 / 文件 %s",
            fee_count,
            export_row_count,
            filename,
        )
        return content, filename

    def _write_banner(self, ws: Any, *, filter_text: str, fee_count: int) -> None:
        indent = Alignment(horizontal="left", vertical="center", indent=1)
        layout = self._layout
        title = "焊接费" if layout.omit_scope else "焊接费核对表"

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=layout.col_count)
        c1 = ws.cell(row=1, column=1, value=title)
        c1.font = _font(size=15, bold=True, color="FFFFFF")
        c1.fill = _fill("1A237E")
        c1.alignment = indent
        ws.row_dimensions[1].height = TITLE_H

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=layout.col_count)
        c2 = ws.cell(
            row=2,
            column=1,
            value=(
                f"{filter_text} · 共 {fee_count} 条" f"　　导出 {now().strftime('%Y-%m-%d %H:%M')}"
            ),
        )
        c2.font = _font(size=9, color="455A64")
        c2.fill = _fill("E8EAF6")
        c2.alignment = indent
        ws.row_dimensions[2].height = FILTER_H

        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=layout.col_count)
        note = (
            "说明：范围已在表头锁定，表内不再重复年度/基地/项目/版本；"
            "左侧主表按节点合并，右侧为工艺明细。行费 = 计费量 × 分年单价。"
            if layout.omit_scope
            else "说明：左侧主表按节点合并，含各层级（见「层级」列）；右侧为工艺明细。行费 = 计费量 × 分年单价。"
        )
        c3 = ws.cell(row=3, column=1, value=note)
        c3.font = _font(size=9, color="B71C1C")
        c3.fill = _fill("FFF3E0")
        c3.alignment = indent
        ws.row_dimensions[3].height = NOTE_H

    def _write_headers(self, ws: Any) -> int:
        group_row = 4
        header_row = 5
        white = _font(size=9, bold=True, color="FFFFFF")
        header_font = _font(bold=True, color="FFFFFF")
        layout = self._layout
        main_cols = layout.main_cols
        col_count = layout.col_count

        ws.merge_cells(
            start_row=group_row,
            start_column=1,
            end_row=group_row,
            end_column=main_cols,
        )
        ws.merge_cells(
            start_row=group_row,
            start_column=main_cols + 1,
            end_row=group_row,
            end_column=col_count,
        )
        main_group = ws.cell(row=group_row, column=1, value="节点主表")
        main_group.font = white
        main_group.fill = _fill("37474F")
        main_group.alignment = self._center
        for col in range(1, main_cols + 1):
            cell = ws.cell(row=group_row, column=col)
            cell.fill = _fill("37474F")
            cell.border = self._thin if col < main_cols else self._split_main

        detail_group = ws.cell(row=group_row, column=main_cols + 1, value="工艺明细 / 计费")
        detail_group.font = white
        detail_group.fill = _fill("B71C1C")
        detail_group.alignment = self._center
        for col in range(main_cols + 1, col_count + 1):
            cell = ws.cell(row=group_row, column=col)
            cell.fill = _fill("B71C1C")
            cell.border = self._split_detail if col == main_cols + 1 else self._thin
        ws.row_dimensions[group_row].height = GROUP_H

        fee_cols = layout.fee_cols
        for col, title in enumerate(layout.headers, start=1):
            cell = ws.cell(row=header_row, column=col, value=title)
            cell.font = header_font
            cell.alignment = self._center
            if col <= main_cols:
                cell.fill = _fill("EF6C00") if col in fee_cols else _fill("546E7A")
                cell.border = self._thin if col < main_cols else self._split_main
            elif col in fee_cols:
                cell.fill = _fill("EF6C00")
                cell.border = self._thin
            else:
                cell.fill = _fill("C62828")
                cell.border = self._split_detail if col == main_cols + 1 else self._thin
        ws.row_dimensions[header_row].height = HEADER_H
        return header_row

    def _write_data_row(
        self,
        ws: Any,
        row_idx: int,
        record: BomWeldingFee,
        *,
        seq_no: int,
        item: Any,
        write_main: bool,
        line_index: int,
        main_fill: PatternFill,
    ) -> None:
        layout = self._layout
        col_index = layout.col
        main_cols = layout.main_cols
        detail_fill = _fill("FFFFFF") if line_index % 2 == 0 else _fill("FFEBEE")
        missing = int(record.missing_price_count or 0) > 0
        calc_status = record.calc_status or ""
        match_status = (item.match_status if item else "") or ""

        if write_main:
            main_pairs: list[tuple[str, Any]] = [
                ("序号", seq_no),
                ("年度", record.year),
                ("基地", record.location),
                ("项目", record.vehicle_model),
                ("版本", record.bom_version),
                ("层级", _level_label(record.level)),
                ("零件号", record.part_number),
                (
                    "短代码",
                    part_short_code(part_number=record.part_number, level=record.level),
                ),
                ("零件名称", record.part_name or ""),
                ("焊接信息原文", (record.weld_info_raw or "").strip()),
                ("工艺数", record.process_line_count),
                ("命中", record.matched_line_count),
                ("缺单价", record.missing_price_count),
                ("系统计算", _num(record.total_welding_fee)),
                ("其它费用", _num(record.other_fee)),
                ("确认单价", _num(record.final_welding_fee)),
                ("备注", (record.remark or "").strip()),
                ("计算状态", _CALC_LABEL.get(calc_status, calc_status)),
                ("审核状态", audit_status_label(record.audit_status)),
            ]
            main_values = [value for name, value in main_pairs if name in col_index]
            for col, value in enumerate(main_values, start=1):
                cell = ws.cell(row=row_idx, column=col, value="" if value is None else value)
                cell.fill = main_fill
                cell.border = self._thin if col < main_cols else self._split_main
                if col == col_index["序号"]:
                    cell.font = _font(size=11, bold=True, color="37474F")
                    cell.alignment = self._center
                elif col == col_index["零件号"]:
                    cell.font = _font(bold=True, color="1565C0")
                    cell.alignment = self._left
                elif col == col_index["短代码"]:
                    cell.font = _font(size=11, bold=True, color="1565C0")
                    cell.alignment = self._center
                elif col == col_index["焊接信息原文"]:
                    cell.font = _font(size=9, color="607D8B")
                    cell.alignment = self._left
                elif col == col_index["系统计算"]:
                    cell.font = _font(size=11, bold=True, color="546E7A")
                    cell.alignment = self._right
                elif col == col_index["其它费用"]:
                    cell.font = _font(size=11, color="37474F")
                    cell.alignment = self._right
                elif col == col_index["确认单价"]:
                    cell.font = _font(size=11, bold=True, color="E65100")
                    cell.alignment = self._right
                    cell.fill = _fill("FFF3E0")
                elif col == col_index["备注"]:
                    cell.font = _font(size=9, color="607D8B")
                    cell.alignment = self._left
                elif col in layout.count_cols:
                    cell.font = (
                        _font(bold=True, color="C62828")
                        if (col == col_index["缺单价"] and missing)
                        else _font()
                    )
                    cell.alignment = self._center
                elif col == col_index["计算状态"]:
                    cell.font, cell.fill = self._status_tone(calc_status)
                    cell.alignment = self._center
                elif col == col_index["审核状态"]:
                    cell.font, cell.fill = self._audit_tone(record.audit_status)
                    cell.alignment = self._center
                else:
                    cell.font = _font()
                    cell.alignment = self._left if col in layout.left_main_cols else self._center

        detail_values: list[Any] = [
            (item.standard_process_name or item.process_token or "") if item else "",
            _num(item.process_qty) if item else "",
            (item.process_qty_unit or "") if item else "",
            _num(item.unit_price) if item and item.unit_price is not None else "",
            (item.unit or "") if item else "",
            _num(item.line_fee) if item and item.line_fee is not None else "",
            _MATCH_LABEL.get(match_status, match_status) if item else "",
        ]
        for offset, value in enumerate(detail_values):
            col = main_cols + 1 + offset
            cell = ws.cell(row=row_idx, column=col, value="" if value is None else value)
            cell.fill = detail_fill
            cell.border = self._split_detail if col == main_cols + 1 else self._thin
            if col == col_index["行费"]:
                cell.font = _font(size=11, bold=True, color="E65100")
                cell.alignment = self._right
                cell.fill = _fill("FFF3E0")
            elif col in {col_index["过程量"], col_index["单价"]}:
                cell.font = _font()
                cell.alignment = self._right
            elif col == col_index["匹配状态"]:
                cell.font, cell.fill = self._match_tone(match_status)
                cell.alignment = self._center
            else:
                cell.font = _font()
                cell.alignment = self._left if col == col_index["标准工艺"] else self._center

    def _merge_main(
        self,
        ws: Any,
        start_row: int,
        end_row: int,
        main_fill: PatternFill,
    ) -> None:
        layout = self._layout
        for col in range(1, layout.main_cols + 1):
            ws.merge_cells(
                start_row=start_row,
                start_column=col,
                end_row=end_row,
                end_column=col,
            )
            cell = ws.cell(row=start_row, column=col)
            if col == layout.col["确认单价"]:
                cell.fill = _fill("FFF3E0")
            else:
                cell.fill = main_fill
            cell.border = self._thin if col < layout.main_cols else self._split_main

    @staticmethod
    def _row_height(record: BomWeldingFee) -> float:
        weld = (record.weld_info_raw or "").strip()
        extra = 0.0
        if len(weld) > 40:
            extra = min(36.0, (len(weld) // 40) * 8.0)
        return DATA_H + extra

    @staticmethod
    def _status_tone(calc_status: str) -> tuple[Font, PatternFill]:
        if calc_status == "ok":
            return _font(bold=True, color="2E7D32"), _fill("E8F5E9")
        if calc_status == "partial":
            return _font(bold=True, color="EF6C00"), _fill("FFF3E0")
        if calc_status == "failed":
            return _font(bold=True, color="C62828"), _fill("FFEBEE")
        return _font(), _fill("FFFFFF")

    @staticmethod
    def _audit_tone(audit_status: str | None) -> tuple[Font, PatternFill]:
        if audit_status == "approved":
            return _font(bold=True, color="2E7D32"), _fill("E8F5E9")
        if audit_status == "pending":
            return _font(bold=True, color="EF6C00"), _fill("FFF8E1")
        if audit_status == "rejected":
            return _font(bold=True, color="C62828"), _fill("FFEBEE")
        return _font(), _fill("FFFFFF")

    @staticmethod
    def _match_tone(match_status: str) -> tuple[Font, PatternFill]:
        if match_status == "matched":
            return _font(bold=True, color="2E7D32"), _fill("E8F5E9")
        if match_status == "missing_price":
            return _font(bold=True, color="EF6C00"), _fill("FFF3E0")
        if match_status == "unmapped":
            return _font(bold=True, color="607D8B"), _fill("ECEFF1")
        return _font(), _fill("FFFFFF")

    @staticmethod
    def _filter_text(
        *,
        year: int | None,
        location: str | None,
        vehicle_model: str | None,
        bom_version: str | None,
        part_number: str | None,
        audit_status: str | None,
        calc_status: str | None,
        level: int | None = None,
        levels: list[int] | None = None,
        ids: list[int] | None = None,
    ) -> str:
        parts: list[str] = []
        if year is not None:
            parts.append(f"年度={year}")
        if location:
            parts.append(f"基地={location}")
        if vehicle_model:
            parts.append(f"项目={vehicle_model}")
        if bom_version:
            parts.append(f"版本={bom_version}")
        resolved = resolve_welding_fee_levels(level=level, levels=levels)
        if resolved:
            parts.append("层级=" + "+".join(_level_label(item) for item in resolved))
        if ids is not None:
            parts.append(f"核价卡焊费 {len(ids)} 条")
        if part_number:
            parts.append(f"零件号含={part_number}")
        if audit_status:
            parts.append(f"审核={audit_status_label(audit_status)}")
        if calc_status:
            parts.append(f"计算={_CALC_LABEL.get(calc_status, calc_status)}")
        return " · ".join(parts) if parts else "全部筛选"

    def _fill_missing_weld_info(self, db: Session, records: list[BomWeldingFee]) -> None:
        need_ids = [
            int(r.source_breakdown_node_id)
            for r in records
            if r.source_breakdown_node_id and not (r.weld_info_raw or "").strip()
        ]
        if not need_ids:
            return
        weld_map = self.breakdown_repo.map_weld_info_by_ids(db, need_ids)
        for record in records:
            if (record.weld_info_raw or "").strip():
                continue
            node_id = record.source_breakdown_node_id
            if not node_id:
                continue
            text = weld_map.get(int(node_id))
            if text:
                record.weld_info_raw = text
