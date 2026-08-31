"""采购对象清单 Excel 导出。"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Protocol, cast

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.common.timezone import now
from app.models.bom.procurement_object import BomProcurementObject
from app.repositories.bom.procurement_object_repository import BomProcurementObjectRepository
from app.schemas.dto.bom_procurement_object_dto import BomProcurementObjectQuery
from app.services.bom.procurement_object.classify import (
    SECTION_ANOMALY,
    SECTION_OBJECT,
    SECTION_OTHER_LEAF,
)
from app.services.bom.procurement_object.labels import (
    AUDIT_LABELS,
    CONFIRM_LABELS,
    KIND_LABELS,
    SECTION_LABELS,
    SECTION_SHEET_TITLES,
    SUBTYPE_LABELS,
)

_FONT = "微软雅黑"
_ExcelCell = str | int | None
_COL_COUNT = 20
_GROUP_ROW = 4
_HEADER_ROW = 5
_DATA_START = 6


class _HeaderFooterPart(Protocol):
    text: str


class _HeaderFooterItem(Protocol):
    left: _HeaderFooterPart
    center: _HeaderFooterPart
    right: _HeaderFooterPart


def _set_footer_text(
    item: _HeaderFooterItem | None,
    *,
    left: str,
    center: str,
    right: str,
) -> None:
    if item is None:
        return
    item.left.text = left
    item.center.text = center
    item.right.text = right


# 色板与冲压模具费导出一致（靛蓝标题 / 石板分组 / 品牌红判定）。
def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(*, size: int = 10, bold: bool = False, color: str = "263238") -> Font:
    return Font(name=_FONT, size=size, bold=bold, color=color)


def _thin_border(*, right: Side | None = None, left: Side | None = None) -> Border:
    thin = Side(style="thin", color="CFD8DC")
    return Border(
        left=left or thin,
        right=right or thin,
        top=thin,
        bottom=thin,
    )


_FILL_TITLE = _fill("1A237E")
_FILL_META = _fill("E8EAF6")
_FILL_NOTE = _fill("FFF3E0")
_FILL_GROUP_MAIN = _fill("37474F")
_FILL_GROUP_JUDGE = _fill("B71C1C")
_FILL_GROUP_EVIDENCE = _fill("607D8B")
_FILL_HEADER_MAIN = _fill("546E7A")
_FILL_HEADER_JUDGE = _fill("C62828")
_FILL_HEADER_KPI = _fill("EF6C00")
_FILL_HEADER_EVIDENCE = _fill("78909C")
_FILL_ZEBRA = _fill("ECEFF1")
_FILL_WHITE = _fill("F5F7FA")
_FILL_EMPTY = _fill("FAFAFA")
_FILL_L0 = _fill("ECEFF1")
_FILL_PENDING = _fill("FFF3E0")
_FILL_OUTSOURCED = _fill("FFEBEE")
_FILL_IN_HOUSE = _fill("F5F7FA")

_THIN = _thin_border()
_SPLIT_JUDGE = _thin_border(right=Side(style="medium", color="B71C1C"))
_SPLIT_EVIDENCE = _thin_border(left=Side(style="medium", color="B71C1C"))

# 对着原件 | 供货总成 | 判定 | 佐证 | 原件补充 | 范围
_GROUP_SPANS = (
    (1, 4, "对着原件"),
    (5, 6, "供货总成"),
    (7, 9, "判定"),
    (10, 14, "佐证"),
    (15, 17, "原件补充"),
    (18, 20, "范围"),
)
_GROUP_FILLS = (
    _FILL_GROUP_MAIN,
    _FILL_GROUP_MAIN,
    _FILL_GROUP_JUDGE,
    _FILL_GROUP_EVIDENCE,
    _FILL_GROUP_EVIDENCE,
    _FILL_GROUP_EVIDENCE,
)

_HEADERS = (
    "Excel行号",
    "工作表",
    "零件号",
    "零件名称",
    "供货总成零件号",
    "供货总成名称",
    "对象种类",
    "确认",
    "审核",
    "材料",
    "料厚",
    "厂家",
    "焊接信息",
    "备注",
    "退化亚类",
    "层级",
    "原始文件名",
    "基地",
    "车型",
    "BOM版本",
)
_WIDTHS = (
    12,
    12,
    18,
    22,
    18,
    22,
    12,
    10,
    10,
    12,
    10,
    14,
    28,
    22,
    12,
    8,
    28,
    10,
    10,
    12,
)
_L0_NO_COL = 5
_L0_NAME_COL = 6
_CONFIRM_COL = 8
_PRIMARY_COLS = frozenset({1, 2, 3, 4, 5, 6, 7, 8, 9})
_CENTER_COLS = frozenset({1, 2, 7, 8, 9, 16})
_KEY_COLS = frozenset({1, 3, 5})
_L0_COLS = (_L0_NO_COL, _L0_NAME_COL)

_JUDGE_END_COL = 9
_KPI_COLS = frozenset({8})
_EVIDENCE_START_COL = 10

_FONT_TITLE = _font(size=15, bold=True, color="FFFFFF")
_FONT_META = _font(size=9, color="455A64")
_FONT_NOTE = _font(size=9, color="B71C1C")
_FONT_GROUP = _font(size=9, bold=True, color="FFFFFF")
_FONT_HEADER = _font(bold=True, color="FFFFFF")
_FONT_HEADER_MUTED = _font(size=9, bold=True, color="FFFFFF")
_FONT_PRIMARY = _font(size=10, color="263238")
_FONT_PRIMARY_KEY = _font(size=11, bold=True, color="37474F")
_FONT_EVIDENCE = _font(size=9, color="455A64")

_SHEETS = (
    (SECTION_SHEET_TITLES[SECTION_OBJECT], SECTION_OBJECT, "1A237E"),
    (SECTION_SHEET_TITLES[SECTION_OTHER_LEAF], SECTION_OTHER_LEAF, "607D8B"),
    (SECTION_SHEET_TITLES[SECTION_ANOMALY], SECTION_ANOMALY, "EF6C00"),
)
_CONFIRM_FILLS = {
    "pending": _FILL_PENDING,
    "outsourced": _FILL_OUTSOURCED,
    "in_house": _FILL_IN_HOUSE,
}


class BomProcurementObjectExportService:
    def __init__(self) -> None:
        self.repository = BomProcurementObjectRepository()

    def export_excel(self, db: Session, query: BomProcurementObjectQuery) -> tuple[bytes, str]:
        rows = self.repository.list_objects_for_export(
            db,
            location=query.location,
            vehicle_model=query.vehicle_model,
            bom_version=query.bom_version,
            section=query.section,
            object_kind=query.object_kind,
            confirm_status=query.confirm_status,
            audit_status=query.audit_status,
            part_number=query.part_number,
            l0_part_number=query.l0_part_number,
            excel_row=query.excel_row,
        )
        stamp = now()
        workbook = Workbook()
        first_sheet = True
        for title, section, tab_color in _SHEETS:
            if query.section and query.section != section:
                continue
            sheet_rows = [row for row in rows if row.section == section]
            if first_sheet:
                sheet = workbook.worksheets[0]
                sheet.title = title
                first_sheet = False
            else:
                sheet = workbook.create_sheet(title)
            sheet.sheet_properties.tabColor = tab_color
            self._write_data_sheet(sheet, sheet_rows, query, stamp, title)

        buffer = BytesIO()
        workbook.save(buffer)
        filename = self._filename(query, stamp)
        return buffer.getvalue(), filename

    def _write_data_sheet(
        self,
        sheet: Worksheet,
        rows: list[BomProcurementObject],
        query: BomProcurementObjectQuery,
        stamp: datetime,
        sheet_title: str,
    ) -> None:
        self._write_banner(sheet, rows, query, stamp, sheet_title)
        self._write_group_headers(sheet)
        self._write_column_headers(sheet)
        if rows:
            self._write_rows(sheet, rows)
            self._merge_same_l0(sheet, rows)
            last_row = _DATA_START + len(rows) - 1
        else:
            last_row = _DATA_START
            sheet.merge_cells(
                start_row=_DATA_START,
                start_column=1,
                end_row=_DATA_START,
                end_column=_COL_COUNT,
            )
            empty = sheet.cell(row=_DATA_START, column=1, value="暂无数据")
            empty.font = _font(size=10, color="455A64")
            empty.fill = _FILL_EMPTY
            empty.alignment = Alignment(horizontal="center", vertical="center")
            for col in range(1, _COL_COUNT + 1):
                cell = sheet.cell(row=_DATA_START, column=col)
                cell.fill = _FILL_EMPTY
                cell.border = _THIN
            sheet.row_dimensions[_DATA_START].height = 24

        for idx, width in enumerate(_WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(idx)].width = width
        sheet.freeze_panes = "G6"
        sheet.auto_filter.ref = f"A{_HEADER_ROW}:{get_column_letter(_COL_COUNT)}{last_row}"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 90
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = 8
        sheet.page_setup.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.3
        sheet.page_margins.right = 0.3
        sheet.page_margins.top = 0.35
        sheet.page_margins.bottom = 0.35
        sheet.print_options.horizontalCentered = True
        sheet.print_title_rows = "1:5"
        sheet.print_title_cols = "A:F"
        sheet.print_area = f"A1:{get_column_letter(_COL_COUNT)}{last_row}"
        _set_footer_text(
            cast(_HeaderFooterItem | None, sheet.oddFooter),
            left=f"采购对象清单 · {sheet_title}",
            center="按 Excel 行号对照原件",
            right="第 &P 页 / 共 &N 页",
        )

    def _write_banner(
        self,
        sheet: Worksheet,
        rows: list[BomProcurementObject],
        query: BomProcurementObjectQuery,
        stamp: datetime,
        sheet_title: str,
    ) -> None:
        indent = Alignment(horizontal="left", vertical="center", indent=1)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=_COL_COUNT)
        title = sheet.cell(row=1, column=1, value=f"采购对象清单 · {sheet_title}")
        title.font = _FONT_TITLE
        title.fill = _FILL_TITLE
        title.alignment = indent
        for col in range(1, _COL_COUNT + 1):
            sheet.cell(row=1, column=col).fill = _FILL_TITLE
        sheet.row_dimensions[1].height = 28

        qa_text, qa_ok = self._qa_summary(rows)
        pending = sum(1 for row in rows if row.confirm_status == "pending")
        outsourced = sum(1 for row in rows if row.confirm_status == "outsourced")
        in_house = sum(1 for row in rows if row.confirm_status == "in_house")
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=_COL_COUNT)
        meta = sheet.cell(
            row=2,
            column=1,
            value=(
                f"{self._scope_text(query)} · {qa_text} · {len(rows)} 条"
                f"　　待定 {pending} · 外发 {outsourced} · 自制 {in_house}"
                f"　　导出 {stamp.strftime('%Y-%m-%d %H:%M')}"
            ),
        )
        meta.font = _FONT_META
        meta.fill = _FILL_META
        meta.alignment = indent
        for col in range(1, _COL_COUNT + 1):
            sheet.cell(row=2, column=col).fill = _FILL_META
        sheet.row_dimensions[2].height = 20

        sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=_COL_COUNT)
        note = sheet.cell(
            row=3,
            column=1,
            value=(
                "打开原件按「Excel行号」对照零件与供货总成；左侧冻结列为核对主键。"
                + ("" if qa_ok else " 检测未通过：仍可导出，禁止确认/审核。")
            ),
        )
        note.font = _FONT_NOTE
        note.fill = _FILL_NOTE
        note.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        for col in range(1, _COL_COUNT + 1):
            sheet.cell(row=3, column=col).fill = _FILL_NOTE
        sheet.row_dimensions[3].height = 36

    def _write_group_headers(self, sheet: Worksheet) -> None:
        center = Alignment(horizontal="center", vertical="center")
        for index, (start, end, label) in enumerate(_GROUP_SPANS):
            fill = _GROUP_FILLS[index]
            if start != end:
                sheet.merge_cells(
                    start_row=_GROUP_ROW,
                    start_column=start,
                    end_row=_GROUP_ROW,
                    end_column=end,
                )
            cell = sheet.cell(row=_GROUP_ROW, column=start, value=label)
            cell.font = _FONT_GROUP
            cell.fill = fill
            cell.alignment = center
            for col in range(start, end + 1):
                grouped = sheet.cell(row=_GROUP_ROW, column=col)
                grouped.fill = fill
                grouped.border = self._cell_border(col, is_group_start=False)
        sheet.row_dimensions[_GROUP_ROW].height = 18

    def _write_column_headers(self, sheet: Worksheet) -> None:
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col, title in enumerate(_HEADERS, start=1):
            cell = sheet.cell(row=_HEADER_ROW, column=col, value=title)
            cell.font = _FONT_HEADER if col < _EVIDENCE_START_COL else _FONT_HEADER_MUTED
            cell.fill = self._header_fill(col)
            cell.alignment = center
            cell.border = self._cell_border(col, is_group_start=False)
        sheet.row_dimensions[_HEADER_ROW].height = 22

    def _header_fill(self, col: int) -> PatternFill:
        if col in _KPI_COLS:
            return _FILL_HEADER_KPI
        if col <= 6:
            return _FILL_HEADER_MAIN
        if col <= _JUDGE_END_COL:
            return _FILL_HEADER_JUDGE
        return _FILL_HEADER_EVIDENCE

    def _cell_border(self, col: int, *, is_group_start: bool) -> Border:
        if col == _JUDGE_END_COL:
            border = _SPLIT_JUDGE
        elif col == _EVIDENCE_START_COL:
            border = _SPLIT_EVIDENCE
        else:
            border = _THIN
        if not is_group_start:
            return border
        return Border(
            left=border.left,
            right=border.right,
            top=Side(style="medium", color="B71C1C"),
            bottom=border.bottom,
        )

    def _write_rows(self, sheet: Worksheet, rows: list[BomProcurementObject]) -> None:
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        center = Alignment(horizontal="center", vertical="center")
        previous_l0 = object()
        for offset, row in enumerate(rows):
            row_idx = _DATA_START + offset
            zebra = offset % 2 == 1
            is_group_start = row.l0_part_number != previous_l0
            previous_l0 = row.l0_part_number
            header = row.object_list
            values: list[_ExcelCell] = [
                row.excel_row,
                row.excel_sheet_name or "",
                row.part_number,
                row.part_name or "",
                row.l0_part_number,
                row.l0_part_name or "",
                self._kind_label(row),
                CONFIRM_LABELS.get(row.confirm_status, row.confirm_status),
                AUDIT_LABELS.get(row.audit_status, row.audit_status),
                row.material or "",
                row.thickness_text or "",
                row.material_mill or "",
                row.weld_info or "",
                row.remark or "",
                SUBTYPE_LABELS.get(row.degenerate_subtype or "", row.degenerate_subtype or ""),
                row.level,
                row.excel_filename or (header.source_filename if header else ""),
                header.location if header else "",
                header.vehicle_model if header else "",
                header.bom_version if header else "",
            ]
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_idx, column=col, value=value)
                cell.border = self._cell_border(col, is_group_start=is_group_start and offset > 0)
                cell.alignment = center if col in _CENTER_COLS else left
                cell.fill = self._cell_fill(col, row, zebra)
                cell.font = self._cell_font(col)
            sheet.row_dimensions[row_idx].height = 24

    def _merge_same_l0(self, sheet: Worksheet, rows: list[BomProcurementObject]) -> None:
        """相邻且供货总成零件号、名称都相同的行，纵向合并这两列。"""
        if len(rows) < 2:
            return
        start = 0
        for index in range(1, len(rows) + 1):
            still_same = index < len(rows) and self._l0_merge_key(
                rows[index]
            ) == self._l0_merge_key(rows[start])
            if still_same:
                continue
            end = index - 1
            if end > start and self._l0_merge_key(rows[start])[0]:
                start_excel = _DATA_START + start
                end_excel = _DATA_START + end
                align = Alignment(horizontal="left", vertical="center", wrap_text=True)
                for col in _L0_COLS:
                    sheet.merge_cells(
                        start_row=start_excel,
                        start_column=col,
                        end_row=end_excel,
                        end_column=col,
                    )
                    cell = sheet.cell(row=start_excel, column=col)
                    cell.alignment = align
                    cell.fill = _FILL_L0
                    for excel_row in range(start_excel, end_excel + 1):
                        grouped = sheet.cell(row=excel_row, column=col)
                        grouped.fill = _FILL_L0
            start = index

    def _l0_merge_key(self, row: BomProcurementObject) -> tuple[str, str]:
        return (row.l0_part_number or "", row.l0_part_name or "")

    def _cell_fill(
        self,
        col: int,
        row: BomProcurementObject,
        zebra: bool,
    ) -> PatternFill:
        if col in _L0_COLS:
            return _FILL_L0
        if col == _CONFIRM_COL:
            return _CONFIRM_FILLS.get(row.confirm_status, _FILL_PENDING)
        if zebra:
            return _FILL_ZEBRA
        return _FILL_WHITE

    def _cell_font(self, col: int) -> Font:
        if col in _KEY_COLS:
            return _FONT_PRIMARY_KEY
        if col in _PRIMARY_COLS:
            return _FONT_PRIMARY
        return _FONT_EVIDENCE

    def _kind_label(self, row: BomProcurementObject) -> str:
        if row.object_kind:
            return KIND_LABELS.get(row.object_kind, row.object_kind)
        return SECTION_LABELS.get(row.section, "—")

    def _qa_summary(self, rows: list[BomProcurementObject]) -> tuple[str, bool]:
        headers = {id(row.object_list): row.object_list for row in rows if row.object_list}
        if not headers:
            return "尚未生成或当前筛选无行", True
        failed = [item for item in headers.values() if not item.qa_passed]
        if failed:
            return "检测未通过（仍可导出，禁止确认/审核）", False
        return "检测通过", True

    def _scope_text(self, query: BomProcurementObjectQuery) -> str:
        location = (query.location or "").strip() or "全部基地"
        model = (query.vehicle_model or "").strip() or "全部项目"
        version = (query.bom_version or "").strip() or "全部版本"
        return f"{location} · {model} · {version}"

    def _filename(self, query: BomProcurementObjectQuery, stamp: datetime) -> str:
        parts = ["采购对象清单"]
        if query.location:
            parts.append(query.location.strip())
        if query.vehicle_model:
            parts.append(query.vehicle_model.strip())
        if query.bom_version:
            parts.append(query.bom_version.strip())
        return "_".join(parts) + f"_{stamp.strftime('%Y%m%d_%H%M%S')}.xlsx"
