"""采购对象清单导出 — openpyxl 渲染（不访问 DB）。"""

from __future__ import annotations

from collections.abc import Callable
from datetime import datetime
from typing import Protocol, cast

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models.bom.procurement_object import BomProcurementObject
from app.schemas.dto.bom_procurement_object_dto import BomProcurementObjectQuery
from app.services.bom.procurement_object.columns import (
    CENTER_COLS,
    COL_COUNT,
    CONFIRM_COL,
    DATA_START,
    EVIDENCE_START_COL,
    FILENAME_COL,
    GROUP_ROW,
    GROUP_SPANS,
    HEADER_ROW,
    HEADERS,
    INTERNAL_NUMBER_COL,
    JUDGE_END_COL,
    KEY_COLS,
    KPI_COLS,
    L0_COLS,
    PRIMARY_COLS,
    SHEET_COL,
    SHORT_CODE_COL,
    WIDTHS,
)
from app.services.bom.procurement_object.mapping import L0Period, l0_period_for_row
from app.services.bom.procurement_object.styles import (
    CONFIRM_FILLS,
    FILL_EMPTY,
    FILL_HEADER_EVIDENCE,
    FILL_HEADER_JUDGE,
    FILL_HEADER_KPI,
    FILL_HEADER_MAIN,
    FILL_L0,
    FILL_META,
    FILL_NOTE,
    FILL_PENDING,
    FILL_TITLE,
    FILL_WHITE,
    FILL_ZEBRA,
    FONT_EVIDENCE,
    FONT_GROUP,
    FONT_HEADER,
    FONT_HEADER_MUTED,
    FONT_INTERNAL_NUMBER,
    FONT_META,
    FONT_NOTE,
    FONT_PRIMARY,
    FONT_PRIMARY_KEY,
    FONT_SHORT_CODE,
    FONT_TITLE,
    GROUP_FILLS,
    SPLIT_EVIDENCE,
    SPLIT_JUDGE,
    THIN,
    font,
)
from app.services.bom.procurement_object.values import (
    excel_filename_key,
    excel_sheet_key,
    l0_merge_key,
    qa_summary,
    row_values,
    scope_text,
)


class _HeaderFooterPart(Protocol):
    text: str


class _HeaderFooterItem(Protocol):
    left: _HeaderFooterPart
    center: _HeaderFooterPart
    right: _HeaderFooterPart


def _set_footer_text(
    item: _HeaderFooterItem | None,
    *,
    left: str,
    center: str,
    right: str,
) -> None:
    if item is None:
        return
    item.left.text = left
    item.center.text = center
    item.right.text = right


def cell_border(col: int, *, is_group_start: bool) -> Border:
    if col == JUDGE_END_COL:
        border = SPLIT_JUDGE
    elif col == EVIDENCE_START_COL:
        border = SPLIT_EVIDENCE
    else:
        border = THIN
    if not is_group_start:
        return border
    return Border(
        left=border.left,
        right=border.right,
        top=Side(style="medium", color="B71C1C"),
        bottom=border.bottom,
    )


def header_fill(col: int) -> PatternFill:
    if col in KPI_COLS:
        return FILL_HEADER_KPI
    if col <= 8:
        return FILL_HEADER_MAIN
    if col <= JUDGE_END_COL:
        return FILL_HEADER_JUDGE
    return FILL_HEADER_EVIDENCE


def cell_fill(col: int, row: BomProcurementObject, zebra: bool) -> PatternFill:
    if col in L0_COLS:
        return FILL_L0
    if col == CONFIRM_COL:
        return CONFIRM_FILLS.get(row.confirm_status, FILL_PENDING)
    if zebra:
        return FILL_ZEBRA
    return FILL_WHITE


def cell_font(col: int) -> Font:
    if col == SHORT_CODE_COL:
        return FONT_SHORT_CODE
    if col == INTERNAL_NUMBER_COL:
        return FONT_INTERNAL_NUMBER
    if col in KEY_COLS:
        return FONT_PRIMARY_KEY
    if col in PRIMARY_COLS:
        return FONT_PRIMARY
    return FONT_EVIDENCE


def write_banner(
    sheet: Worksheet,
    rows: list[BomProcurementObject],
    query: BomProcurementObjectQuery,
    stamp: datetime,
    sheet_title: str,
) -> None:
    indent = Alignment(horizontal="left", vertical="center", indent=1)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COL_COUNT)
    title = sheet.cell(row=1, column=1, value=f"采购对象清单 · {sheet_title}")
    title.font = FONT_TITLE
    title.fill = FILL_TITLE
    title.alignment = indent
    for col in range(1, COL_COUNT + 1):
        sheet.cell(row=1, column=col).fill = FILL_TITLE
    sheet.row_dimensions[1].height = 28

    qa_text, qa_ok = qa_summary(rows)
    pending = sum(1 for row in rows if row.confirm_status == "pending")
    outsourced = sum(1 for row in rows if row.confirm_status == "outsourced")
    in_house = sum(1 for row in rows if row.confirm_status == "in_house")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=COL_COUNT)
    meta = sheet.cell(
        row=2,
        column=1,
        value=(
            f"{scope_text(query)} · {qa_text} · {len(rows)} 条"
            f"　　待定 {pending} · 外发 {outsourced} · 自制 {in_house}"
            f"　　导出 {stamp.strftime('%Y-%m-%d %H:%M')}"
        ),
    )
    meta.font = FONT_META
    meta.fill = FILL_META
    meta.alignment = indent
    for col in range(1, COL_COUNT + 1):
        sheet.cell(row=2, column=col).fill = FILL_META
    sheet.row_dimensions[2].height = 20

    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=COL_COUNT)
    note = sheet.cell(
        row=3,
        column=1,
        value=(
            "打开原件按「Excel行号」对照；短码是零件号数字末四位，日常核对用。"
            + ("" if qa_ok else " 检测未通过：仍可导出与确认/审核；已审行跳过不改。")
        ),
    )
    note.font = FONT_NOTE
    note.fill = FILL_NOTE
    note.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    for col in range(1, COL_COUNT + 1):
        sheet.cell(row=3, column=col).fill = FILL_NOTE
    sheet.row_dimensions[3].height = 36


def write_group_headers(sheet: Worksheet) -> None:
    center = Alignment(horizontal="center", vertical="center")
    for index, (start, end, label) in enumerate(GROUP_SPANS):
        fill_color = GROUP_FILLS[index]
        if start != end:
            sheet.merge_cells(
                start_row=GROUP_ROW,
                start_column=start,
                end_row=GROUP_ROW,
                end_column=end,
            )
        cell = sheet.cell(row=GROUP_ROW, column=start, value=label)
        cell.font = FONT_GROUP
        cell.fill = fill_color
        cell.alignment = center
        for col in range(start, end + 1):
            grouped = sheet.cell(row=GROUP_ROW, column=col)
            grouped.fill = fill_color
            grouped.border = cell_border(col, is_group_start=False)
    sheet.row_dimensions[GROUP_ROW].height = 18


def write_column_headers(sheet: Worksheet) -> None:
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, title in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=col, value=title)
        cell.font = FONT_HEADER if col < EVIDENCE_START_COL else FONT_HEADER_MUTED
        cell.fill = header_fill(col)
        cell.alignment = center
        cell.border = cell_border(col, is_group_start=False)
    sheet.row_dimensions[HEADER_ROW].height = 22


def write_rows(
    sheet: Worksheet,
    rows: list[BomProcurementObject],
    *,
    period_map: dict[tuple[int, str], L0Period] | None = None,
) -> None:
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    center = Alignment(horizontal="center", vertical="center")
    previous_l0 = object()
    periods = period_map or {}
    for offset, row in enumerate(rows):
        row_idx = DATA_START + offset
        zebra = offset % 2 == 1
        is_group_start = row.l0_part_number != previous_l0
        previous_l0 = row.l0_part_number
        values = row_values(row, l0_period=l0_period_for_row(row, periods))
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_idx, column=col, value=value)
            cell.border = cell_border(col, is_group_start=is_group_start and offset > 0)
            cell.alignment = center if col in CENTER_COLS else left
            cell.fill = cell_fill(col, row, zebra)
            cell.font = cell_font(col)
        sheet.row_dimensions[row_idx].height = 24


def merge_consecutive(
    sheet: Worksheet,
    rows: list[BomProcurementObject],
    *,
    columns: tuple[int, ...],
    key_of: Callable[[BomProcurementObject], tuple[str, ...]],
    fill: PatternFill | None = None,
) -> None:
    if len(rows) < 2:
        return
    start = 0
    for index in range(1, len(rows) + 1):
        still_same = index < len(rows) and key_of(rows[index]) == key_of(rows[start])
        if still_same:
            continue
        end = index - 1
        key = key_of(rows[start])
        if end > start and key[0]:
            start_excel = DATA_START + start
            end_excel = DATA_START + end
            align = Alignment(horizontal="left", vertical="center", wrap_text=True)
            for col in columns:
                sheet.merge_cells(
                    start_row=start_excel,
                    start_column=col,
                    end_row=end_excel,
                    end_column=col,
                )
                cell = sheet.cell(row=start_excel, column=col)
                cell.alignment = align
                if fill is not None:
                    cell.fill = fill
                    for excel_row in range(start_excel, end_excel + 1):
                        sheet.cell(row=excel_row, column=col).fill = fill
        start = index


def write_data_sheet(
    sheet: Worksheet,
    rows: list[BomProcurementObject],
    query: BomProcurementObjectQuery,
    stamp: datetime,
    sheet_title: str,
    *,
    period_map: dict[tuple[int, str], L0Period] | None = None,
) -> None:
    write_banner(sheet, rows, query, stamp, sheet_title)
    write_group_headers(sheet)
    write_column_headers(sheet)
    if rows:
        write_rows(sheet, rows, period_map=period_map)
        merge_consecutive(sheet, rows, columns=L0_COLS, key_of=l0_merge_key, fill=FILL_L0)
        merge_consecutive(sheet, rows, columns=(FILENAME_COL,), key_of=excel_filename_key)
        merge_consecutive(sheet, rows, columns=(SHEET_COL,), key_of=excel_sheet_key)
        last_row = DATA_START + len(rows) - 1
    else:
        last_row = DATA_START
        sheet.merge_cells(
            start_row=DATA_START,
            start_column=1,
            end_row=DATA_START,
            end_column=COL_COUNT,
        )
        empty = sheet.cell(row=DATA_START, column=1, value="暂无数据")
        empty.font = font(size=10, color="455A64")
        empty.fill = FILL_EMPTY
        empty.alignment = Alignment(horizontal="center", vertical="center")
        for col in range(1, COL_COUNT + 1):
            cell = sheet.cell(row=DATA_START, column=col)
            cell.fill = FILL_EMPTY
            cell.border = THIN
        sheet.row_dimensions[DATA_START].height = 24

    for idx, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "I6"
    sheet.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(COL_COUNT)}{last_row}"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = 8
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = 0.3
    sheet.page_margins.right = 0.3
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.35
    sheet.print_options.horizontalCentered = True
    sheet.print_title_rows = "1:5"
    sheet.print_title_cols = "A:H"
    sheet.print_area = f"A1:{get_column_letter(COL_COUNT)}{last_row}"
    _set_footer_text(
        cast(_HeaderFooterItem | None, sheet.oddFooter),
        left=f"采购对象清单 · {sheet_title}",
        center="按 Excel 行号对照原件",
        right="第 &P 页 / 共 &N 页",
    )
