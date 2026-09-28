"""采购对象供应商导入模板：第一页清单现况 + 第三页填写说明。"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models.bom.procurement_object import BomProcurementObject
from app.parsers.bom.procurement_object_supplier_parser import (
    HEADER_AUDIT,
    HEADER_CONFIRM,
    HEADER_INTERNAL_NUMBER,
    HEADER_OBJECT_KIND,
    HEADER_PART_NAME,
    HEADER_PART_NUMBER,
    HEADER_SUPPLIER,
)
from app.services.bom.procurement_object.classify import SECTION_OBJECT
from app.services.bom.procurement_object.labels import (
    AUDIT_LABELS,
    CONFIRM_LABELS,
    KIND_LABELS,
    SECTION_SHEET_TITLES,
)

_SHEET_TITLE = SECTION_SHEET_TITLES[SECTION_OBJECT]
_NOTES_SHEET_TITLE = "填写说明"
_HEADERS = (
    HEADER_PART_NUMBER,
    HEADER_PART_NAME,
    HEADER_INTERNAL_NUMBER,
    HEADER_OBJECT_KIND,
    HEADER_CONFIRM,
    HEADER_AUDIT,
    HEADER_SUPPLIER,
)
_COLUMN_WIDTHS = (18, 22, 16, 14, 10, 10, 36)
_INTERNAL_NUMBER_COL = 3
_NOTES = (
    "第一页按零件号去重，是该范围采购对象的库现况。同一零件号只填一次供应商。",
    "「内部号」给人看，方便到 ERP 按编号找厂、对照。导入不读这一列。",
    "导入按零件号写回本范围全部采购对象行（含挂在不同供货总成下的重复件）。",
    "已指定的供应商写出工商名称；未指定留空。只改「供应商」列。可删掉不改的行；文件里没有的零件号库里不动。空单元格表示清空对方。",
    "第二页是供应商主数据目录（已启用），只对照，导入不读。可复制工商名称 / 缩写 / ERP 填回第一页。",
    "找不到的厂请到供应商列表补档（#/supplier-list），本文件不能建档。",
    "已审行若改供应商会整单拒绝，须先在清单页按范围反审核。自制行禁止填非空供应商。",
)


def unique_object_rows_by_part_number(
    rows: list[BomProcurementObject],
) -> list[BomProcurementObject]:
    seen: set[str] = set()
    unique: list[BomProcurementObject] = []
    for row in rows:
        key = row.part_number or ""
        if not key or key in seen:
            continue
        seen.add(key)
        unique.append(row)
    return unique


def _kind_label(row: BomProcurementObject) -> str:
    if row.object_kind:
        return KIND_LABELS.get(row.object_kind, row.object_kind)
    return ""


def write_object_sheet(workbook: Workbook, rows: list[BomProcurementObject]) -> None:
    sheet = workbook.worksheets[0]
    sheet.title = _SHEET_TITLE
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    data_font = Font(name="微软雅黑", size=9)
    muted_font = Font(name="微软雅黑", size=9, color="808080")
    align_left = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    for col, title in enumerate(_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    for row_idx, row in enumerate(rows, start=2):
        values: list[str | int | None] = [
            row.part_number,
            row.part_name,
            row.internal_number,
            _kind_label(row),
            CONFIRM_LABELS.get(row.confirm_status, row.confirm_status),
            AUDIT_LABELS.get(row.audit_status, row.audit_status),
            row.supplier_name,
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_idx, column=col, value=value)
            cell.font = muted_font if col == _INTERNAL_NUMBER_COL else data_font
            cell.border = thin_border
            cell.alignment = align_left
    for idx, width in enumerate(_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "A2"


def write_notes_sheet(workbook: Workbook) -> None:
    sheet: Worksheet = workbook.create_sheet(_NOTES_SHEET_TITLE)
    sheet["A1"] = "填写说明"
    sheet["A1"].font = Font(name="微软雅黑", size=12, bold=True, color="C00000")
    for idx, line in enumerate(_NOTES, start=3):
        cell = sheet.cell(row=idx, column=1, value=f"{idx - 2}. {line}")
        cell.font = Font(name="微软雅黑", size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[idx].height = 36
    sheet.column_dimensions["A"].width = 96
