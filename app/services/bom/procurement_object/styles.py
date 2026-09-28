"""采购对象列表导出 — 色板、字体、分隔线。"""

from __future__ import annotations

from openpyxl.styles import Border, Font, PatternFill, Side

from app.services.bom.procurement_object.classify import (
    SECTION_ANOMALY,
    SECTION_OBJECT,
    SECTION_OTHER_LEAF,
)
from app.services.bom.procurement_object.labels import SECTION_SHEET_TITLES

_FONT = "微软雅黑"


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def font(*, size: int = 10, bold: bool = False, color: str = "263238") -> Font:
    return Font(name=_FONT, size=size, bold=bold, color=color)


def thin_border(*, right: Side | None = None, left: Side | None = None) -> Border:
    thin = Side(style="thin", color="CFD8DC")
    return Border(
        left=left or thin,
        right=right or thin,
        top=thin,
        bottom=thin,
    )


FILL_TITLE = fill("1A237E")
FILL_META = fill("E8EAF6")
FILL_NOTE = fill("FFF3E0")
FILL_GROUP_MAIN = fill("37474F")
FILL_GROUP_JUDGE = fill("B71C1C")
FILL_GROUP_EVIDENCE = fill("607D8B")
FILL_HEADER_MAIN = fill("546E7A")
FILL_HEADER_JUDGE = fill("C62828")
FILL_HEADER_KPI = fill("EF6C00")
FILL_HEADER_EVIDENCE = fill("78909C")
FILL_ZEBRA = fill("ECEFF1")
FILL_WHITE = fill("F5F7FA")
FILL_EMPTY = fill("FAFAFA")
FILL_L0 = fill("ECEFF1")
FILL_PENDING = fill("FFF3E0")
FILL_OUTSOURCED = fill("FFEBEE")
FILL_IN_HOUSE = fill("F5F7FA")

THIN = thin_border()
SPLIT_JUDGE = thin_border(right=Side(style="medium", color="B71C1C"))
SPLIT_EVIDENCE = thin_border(left=Side(style="medium", color="B71C1C"))

GROUP_FILLS = (
    FILL_GROUP_MAIN,
    FILL_GROUP_MAIN,
    FILL_GROUP_JUDGE,
    FILL_GROUP_EVIDENCE,
    FILL_GROUP_EVIDENCE,
    FILL_GROUP_EVIDENCE,
)

FONT_TITLE = font(size=15, bold=True, color="FFFFFF")
FONT_META = font(size=9, color="455A64")
FONT_NOTE = font(size=9, color="B71C1C")
FONT_GROUP = font(size=9, bold=True, color="FFFFFF")
FONT_HEADER = font(bold=True, color="FFFFFF")
FONT_HEADER_MUTED = font(size=9, bold=True, color="FFFFFF")
FONT_PRIMARY = font(size=10, color="263238")
FONT_PRIMARY_KEY = font(size=11, bold=True, color="37474F")
FONT_SHORT_CODE = Font(name=_FONT, size=10, bold=True, color="1F4E79")
FONT_INTERNAL_NUMBER = Font(name=_FONT, size=9, color="808080")
FONT_EVIDENCE = font(size=9, color="455A64")

SHEETS = (
    (SECTION_SHEET_TITLES[SECTION_OBJECT], SECTION_OBJECT, "1A237E"),
    (SECTION_SHEET_TITLES[SECTION_OTHER_LEAF], SECTION_OTHER_LEAF, "607D8B"),
    (SECTION_SHEET_TITLES[SECTION_ANOMALY], SECTION_ANOMALY, "EF6C00"),
)
CONFIRM_FILLS = {
    "pending": FILL_PENDING,
    "outsourced": FILL_OUTSOURCED,
    "in_house": FILL_IN_HOUSE,
}
