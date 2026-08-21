"""钢价加减价政策 Excel 导入模板生成。"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet


def build_steel_price_policy_template() -> tuple[bytes, str]:
    """生成厚度加减价政策样例（标题/钢种/区间块合并，对齐首钢类排版）。"""
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "样例钢厂"

    header_font = Font(name="微软雅黑", size=9, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    data_font = Font(name="微软雅黑", size=9)
    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    center = Alignment(horizontal="center", vertical="center")

    sheet.merge_cells("A1:I1")
    sheet["A1"] = "样例钢厂厚度加减价（模板演示，非正式发布件）"
    sheet["A1"].font = Font(name="微软雅黑", size=12, bold=True)

    sheet.merge_cells("H2:I2")
    sheet["H2"] = "元/吨，不含税"
    sheet["H2"].alignment = center

    # 钢种跨 R3-R4；厚度标题横跨区间列（对齐首钢实样）
    sheet.merge_cells("A3:A4")
    sheet["A3"] = "钢种"
    sheet["A3"].font = Font(name="微软雅黑", size=10, bold=True)
    sheet["A3"].alignment = center

    sheet.merge_cells("B3:I3")
    sheet["B3"] = "厚度加减价-基准厚度1.0——元/吨"
    sheet["B3"].font = Font(name="微软雅黑", size=10, bold=True)
    sheet["B3"].alignment = center

    boundaries = ["≥0.5", "≥0.6", "≥0.7", "≥0.8", "≥0.9", "≥1.0", "≥1.2", "≥1.5"]
    for col, text in enumerate(boundaries, start=2):
        cell = sheet.cell(row=4, column=col, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin

    # A3:A4 已合并，只写主格 A3，避免对 MergedCell 赋值触发 mypy
    steel_header = sheet.cell(row=3, column=1)
    steel_header.border = thin
    steel_header.alignment = center
    steel_header.fill = header_fill
    steel_header.font = header_font
    for col in range(2, 10):
        title_cell = sheet.cell(row=3, column=col)
        title_cell.border = thin
        title_cell.fill = header_fill

    samples: list[tuple[str, list[int]]] = [
        ("DC01", [210, 180, 140, 120, 80, 0, -30, -50]),
        ("DC03", [210, 180, 140, 120, 80, 0, -30, -50]),
        ("DC04", [230, 190, 150, 120, 80, 0, -30, -50]),
        ("DC51D+Z", [260, 210, 160, 130, 90, 0, -20, -40]),
    ]
    for row_idx, (grade, prices) in enumerate(samples, start=5):
        grade_cell = sheet.cell(row=row_idx, column=1, value=grade)
        grade_cell.font = data_font
        grade_cell.border = thin
        grade_cell.alignment = center
        for col, price in enumerate(prices, start=2):
            cell = sheet.cell(row=row_idx, column=col, value=price)
            cell.font = data_font
            cell.border = thin
            cell.alignment = center

    sheet.column_dimensions["A"].width = 12
    for col in range(2, 10):
        sheet.column_dimensions[chr(64 + col)].width = 10

    note = workbook.create_sheet("填写说明")
    _write_notes(
        note,
        [
            "正式文件一般由奇瑞发布的「一事一议价格厚度加减价」Excel，按钢厂分 Sheet。",
            "标题行横向合并；「钢种」跨两行；厚度加减价标题横跨区间列。",
            "厚度块标题须含「厚度加减价」；下一行是 ≥0.5、≥0.6 等区间列头。",
            "本模板「样例钢厂」为可解析精简样例；生产导入请使用奇瑞原件。",
            "上传页须手工填写生效年度与月份；同键政策会 upsert 覆盖。",
        ],
    )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), "加减价政策导入模板.xlsx"


def _write_notes(sheet: Worksheet, lines: list[str]) -> None:
    sheet["A1"] = "填写说明"
    sheet["A1"].font = Font(name="微软雅黑", size=12, bold=True, color="C00000")
    for idx, line in enumerate(lines, start=3):
        sheet.cell(row=idx, column=1, value=f"{idx - 2}. {line}")
        sheet.cell(row=idx, column=1).font = Font(name="微软雅黑", size=10)
        sheet.cell(row=idx, column=1).alignment = Alignment(
            wrap_text=True, vertical="center"
        )
        sheet.row_dimensions[idx].height = 22
    sheet.column_dimensions["A"].width = 96
