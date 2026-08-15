"""
废料月度销售价格 Excel 导出

版式对齐业务原表：
- 每年一块：首列「{year}年」+ 1~12 月
- 有扣费的基地：3 行（大盘价 / 扣费 / 净价公式），首列合并，净价行浅蓝底
- 净价单元格使用公式：=大盘价单元格-扣费单元格
- 无扣费的基地：仅 1 行
- 跨年之间空一行
"""

from __future__ import annotations

import logging
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.common.scrap_price_locations import SCRAP_PRICE_LOCATIONS
from app.common.timezone import now
from app.repositories.yishiyiyi.scrap_price.scrap_price_repository import StampingScrapPriceRepository


logger = logging.getLogger("service.stamping_scrap_price_export")

_NET_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
_YEAR_FONT = Font(name="宋体", size=10, bold=True)
_MONTH_FONT = Font(name="宋体", size=10, bold=False)
_DATA_FONT = Font(name="宋体", size=10)
_CENTER = Alignment(horizontal="center", vertical="center")
_RIGHT = Alignment(horizontal="right", vertical="center")
_THIN_SIDE = Side(style="thin", color="000000")
_THICK_SIDE = Side(style="medium", color="000000")
_THIN = Border(
    left=_THIN_SIDE,
    right=_THIN_SIDE,
    top=_THIN_SIDE,
    bottom=_THIN_SIDE,
)
# 单价：千分位 + 两位小数
_PRICE_FORMAT = "#,##0.00"
_ROW_HEIGHT = 19
_BLOCK_LAST_COL = 13  # A + 1~12 月
_THREE_ROW_COMMENT = (
    "第1行：大盘价（含杂费账面价）\n"
    "第2行：扣费（运费/扣水/扣杂等）\n"
    "第3行：净价 = 大盘价 − 扣费"
)


def _to_number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return float(value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _set_price(cell, value: float | int | None) -> None:
    """写入单价并套用千分位数字格式。"""
    if value is None:
        return
    cell.value = float(value)
    cell.number_format = _PRICE_FORMAT


def _set_row_height(ws, row: int) -> None:
    ws.row_dimensions[row].height = _ROW_HEIGHT


def _apply_block_outer_border(
    ws,
    min_row: int,
    max_row: int,
    *,
    min_col: int = 1,
    max_col: int = _BLOCK_LAST_COL,
) -> None:
    """年块外框用粗线，内部仍为细线（合并格按整块外缘处理）。"""
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row=row, column=col)
            if isinstance(cell, MergedCell):
                continue
            cell.border = Border(
                left=_THICK_SIDE if col == min_col else _THIN_SIDE,
                right=_THICK_SIDE if col == max_col else _THIN_SIDE,
                top=_THICK_SIDE if row == min_row else _THIN_SIDE,
                bottom=_THICK_SIDE if row == max_row else _THIN_SIDE,
            )

    # 合并区域以外缘为准重写主格边框，否则底边/左边粗框会被合并吃掉
    for merged in ws.merged_cells.ranges:
        mr_min_row, mr_min_col = merged.min_row, merged.min_col
        mr_max_row, mr_max_col = merged.max_row, merged.max_col
        if mr_max_row < min_row or mr_min_row > max_row:
            continue
        if mr_max_col < min_col or mr_min_col > max_col:
            continue
        master = ws.cell(row=mr_min_row, column=mr_min_col)
        master.border = Border(
            left=_THICK_SIDE if mr_min_col == min_col else _THIN_SIDE,
            right=_THICK_SIDE if mr_max_col == max_col else _THIN_SIDE,
            top=_THICK_SIDE if mr_min_row == min_row else _THIN_SIDE,
            bottom=_THICK_SIDE if mr_max_row == max_row else _THIN_SIDE,
        )


class StampingScrapPriceExportService:
    """废料销售价宽表导出。"""

    def __init__(self) -> None:
        self.repository = StampingScrapPriceRepository()

    def export(
        self,
        db: Session,
        *,
        year: int | None = None,
    ) -> tuple[bytes, str]:
        years = [year] if year is not None else self.repository.list_years_asc(db)
        if not years:
            years = []

        wb = Workbook()
        ws = wb.active
        ws.title = "废料销售价"

        current_row = 1
        for index, year_value in enumerate(years):
            if index > 0:
                current_row += 1  # 跨年空行
            current_row = self._write_year_block(ws, db, year_value, current_row)

        if not years:
            ws.cell(row=1, column=1, value="暂无数据")

        ws.column_dimensions["A"].width = 12
        for col in range(2, 14):
            ws.column_dimensions[get_column_letter(col)].width = 11

        buffer = BytesIO()
        wb.save(buffer)
        content = buffer.getvalue()

        stamp = now().strftime("%Y%m%d_%H%M%S")
        if year is not None:
            filename = f"废料月度销售价_{year}年_{stamp}.xlsx"
        else:
            filename = f"废料月度销售价_汇总_{stamp}.xlsx"

        logger.info("导出废料销售价: years=%s bytes=%s", years, len(content))
        return content, filename

    def _write_year_block(self, ws, db: Session, year: int, start_row: int) -> int:
        records = self.repository.list_by_year(db, year)
        # location -> month -> record
        by_location: dict[str, dict[int, Any]] = {
            name: {} for name in SCRAP_PRICE_LOCATIONS
        }
        for record in records:
            if record.location not in by_location:
                by_location[record.location] = {}
            by_location[record.location][int(record.month)] = record

        row = start_row
        # 表头
        headers = [f"{year}年"] + [f"{m}月" for m in range(1, 13)]
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=text)
            # 首列年份加粗，月份不加粗
            cell.font = _YEAR_FONT if col == 1 else _MONTH_FONT
            cell.alignment = _CENTER
            cell.border = _THIN
        _set_row_height(ws, row)
        row += 1

        for location in SCRAP_PRICE_LOCATIONS:
            month_map = by_location.get(location, {})
            has_deduction = any(
                (_to_number(r.deduction_fee) or 0) > 0 for r in month_map.values()
            )
            if has_deduction:
                row = self._write_deduction_location(ws, row, location, month_map)
            else:
                row = self._write_plain_location(ws, row, location, month_map)

        # 整年块加粗外框（start_row ~ 最后一行）
        _apply_block_outer_border(ws, start_row, row - 1)
        return row

    def _write_plain_location(self, ws, row: int, location: str, month_map: dict) -> int:
        name_cell = ws.cell(row=row, column=1, value=location)
        name_cell.font = _DATA_FONT
        name_cell.alignment = _CENTER
        name_cell.border = _THIN
        _set_row_height(ws, row)

        for month in range(1, 13):
            record = month_map.get(month)
            cell = ws.cell(row=row, column=month + 1)
            cell.border = _THIN
            cell.alignment = _RIGHT
            cell.font = _DATA_FONT
            if record is None:
                continue
            market = _to_number(record.market_price)
            if market is not None:
                _set_price(cell, market)
        return row + 1

    def _write_deduction_location(
        self,
        ws,
        row: int,
        location: str,
        month_map: dict,
    ) -> int:
        market_row = row
        deduction_row = row + 1
        net_row = row + 2

        # 首列合并
        ws.merge_cells(
            start_row=market_row,
            start_column=1,
            end_row=net_row,
            end_column=1,
        )
        name_cell = ws.cell(row=market_row, column=1, value=location)
        name_cell.font = _DATA_FONT
        name_cell.alignment = _CENTER
        # 三行含义批注（仅有扣费的基地）
        comment = Comment(_THREE_ROW_COMMENT, "系统")
        comment.width = 220
        comment.height = 70
        name_cell.comment = comment

        for r in (market_row, deduction_row, net_row):
            ws.cell(row=r, column=1).border = _THIN
            _set_row_height(ws, r)

        for month in range(1, 13):
            col = month + 1
            col_letter = get_column_letter(col)
            record = month_map.get(month)
            market_cell = ws.cell(row=market_row, column=col)
            deduction_cell = ws.cell(row=deduction_row, column=col)
            net_cell = ws.cell(row=net_row, column=col)

            for cell in (market_cell, deduction_cell, net_cell):
                cell.border = _THIN
                cell.alignment = _RIGHT
                cell.font = _DATA_FONT
            net_cell.fill = _NET_FILL

            # 无该月数据：大盘/扣费/净价均保持空（勿写公式，否则 Excel 数字格式会显示 0）
            if record is None:
                continue

            market = _to_number(record.market_price)
            if market is None or market == 0:
                continue

            deduction = _to_number(record.deduction_fee) or 0.0
            _set_price(market_cell, market)
            if deduction != 0:
                _set_price(deduction_cell, deduction)

            net = market - deduction
            # 净价为 0：保持空，不写公式、不套数字格式（"" + #,##0.00 在 Excel 里会显示成 0）
            if net == 0:
                continue

            net_cell.value = f"={col_letter}{market_row}-{col_letter}{deduction_row}"
            net_cell.number_format = _PRICE_FORMAT

        return row + 3
