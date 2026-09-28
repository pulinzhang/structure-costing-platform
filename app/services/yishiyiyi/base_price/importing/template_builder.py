"""钢材基准价格 Excel 导入模板生成。"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.yishiyiyi.base_price.gate import reference_width_note

# 样例宽表表头，对齐发布价实样。不是系统基地花名册。
_SAMPLE_LOCATION_COLUMNS: tuple[tuple[str, str], ...] = (
    ("芜湖", "芜湖基地"),
    ("大连", "大连基地"),
    ("青岛", "青岛基地"),
    ("鄂尔多斯", "鄂尔多斯基地"),
)


def sample_location_lookup() -> dict[str, str]:
    """校验样例模板用：别名与规范名都映射到规范名。正式导入请用已启用主数据。"""
    lookup: dict[str, str] = {}
    for canonical, alias in _SAMPLE_LOCATION_COLUMNS:
        lookup[canonical] = canonical
        lookup[alias] = canonical
    return lookup


def build_steel_base_price_template() -> tuple[bytes, str]:
    """生成基准价格样例宽表（表头跨行/跨列合并，对齐发布价实样）。"""
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Sheet1"

    header_font = Font(name="微软雅黑", size=9, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    data_font = Font(name="微软雅黑", size=9)
    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    last_col = 13
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    sheet["A1"] = "密级：内部资料□  外部公开□     发放范围：奇瑞外协冲压厂（模板演示）"

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    sheet["A2"] = "外协钢板样例订单(样例月冲焊件原料)一事一议价格"
    sheet["A2"].font = Font(name="微软雅黑", size=12, bold=True)
    sheet["A2"].alignment = center

    sheet.merge_cells("L3:M3")
    sheet["L3"] = "元/吨，不含税"
    sheet["L3"].alignment = center

    # 固定列对齐发布价：序号 / 钢种 / 牌号 / 参考规格，各占一列，跨第 4-5 行
    sheet.merge_cells("A4:A5")
    sheet["A4"] = "序号"
    sheet.merge_cells("B4:B5")
    sheet["B4"] = "钢种"
    sheet.merge_cells("C4:C5")
    sheet["C4"] = "牌号"
    sheet.merge_cells("D4:D5")
    sheet["D4"] = "参考规格"

    # 基地名跨两列（第 4 行）；写发布价常用别名，导入时由 lookup 收成规范名
    start_col = 5
    for offset, (_canonical, alias) in enumerate(_SAMPLE_LOCATION_COLUMNS):
        col = start_col + offset * 2
        sheet.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        sheet.cell(row=4, column=col, value=alias)
        sheet.cell(row=5, column=col, value="中标钢厂")
        sheet.cell(row=5, column=col + 1, value="到厂原卷承兑价格")

    fee_col = start_col + len(_SAMPLE_LOCATION_COLUMNS) * 2
    sheet.merge_cells(start_row=4, start_column=fee_col, end_row=5, end_column=fee_col)
    sheet.cell(row=4, column=fee_col, value="加工费(加工中心成品自提)")

    for row in (4, 5):
        for col in range(1, last_col + 1):
            cell = sheet.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = thin

    # 末列加工费是样例单元格里的数字，导入只读这一列，不会另填默认值
    samples: list[tuple[int, str, str, str, str, int, str, int, str, int, str, int, int]] = [
        (
            1,
            "镀锌",
            "DC51D+Z",
            "1.0*1250",
            "VAMA",
            4123,
            "VAMA",
            4236,
            "山钢",
            4113,
            "VAMA",
            4291,
            120,
        ),
        (
            2,
            "镀锌",
            "DC53D+Z",
            "1.0*1250",
            "本钢",
            4521,
            "本钢",
            4465,
            "山钢",
            4480,
            "山钢",
            4665,
            120,
        ),
        (
            3,
            "冷轧",
            "DC04",
            "0.8*1250",
            "宝钢",
            3980,
            "宝钢",
            4050,
            "首钢",
            4010,
            "宝钢",
            4120,
            120,
        ),
        (
            4,
            "冷轧",
            "HC340LA",
            "1.2*1500",
            "马钢",
            4210,
            "马钢",
            4280,
            "唐钢",
            4195,
            "马钢",
            4350,
            120,
        ),
    ]
    for row_idx, sample in enumerate(samples, start=6):
        for col, cell_value in enumerate(sample, start=1):
            written: str | float = (
                cell_value if isinstance(cell_value, str) else float(str(cell_value))
            )
            cell = sheet.cell(row=row_idx, column=col, value=written)
            cell.font = data_font
            cell.alignment = center
            cell.border = thin

    for idx in range(1, last_col + 1):
        sheet.column_dimensions[get_column_letter(idx)].width = 14
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 12
    sheet.column_dimensions["M"].width = 22
    sheet.row_dimensions[4].height = 22
    sheet.row_dimensions[5].height = 32
    sheet.freeze_panes = "A6"

    note = workbook.create_sheet("填写说明")
    _write_notes(
        note,
        [
            "正式文件为奇瑞发布的「发布价格及中标钢厂」Excel（外协钢板订单一事一议价格）。",
            "序号、钢种、牌号、参考规格各占一列，跨第 4-5 行；基地名跨两列。",
            "第 5 行为「中标钢厂 / 到厂原卷承兑价格」；数据从第 6 行起。",
            "各基地价格列之后必须有表头含「加工费」的列（正式文件为「加工费(加工中心成品自提)」）。",
            "每个数据行的加工费必须填写数字（元/吨）。没有该列，或单元格为空、不是数字时，整批拒绝。",
            "系统不会改读旁边的到厂价格，也不会填 120 或其它默认加工费。样例行里的 120 只是单元格里的演示数字。",
            "加工费大于系统设置的基准价加工费上限时整批拒绝。",
            "与上月同一基地、牌号、规格相比，单价加加工费的涨跌超出系统设置的基准价月度价格变动范围（默认 ±300 元/吨）时整批拒绝。",
            "参考规格为厚度*宽度，厚度和宽度都不能为空。料厚大于系统设置的车身件料厚上限（默认 3.00 mm）时整批拒绝，等于上限可以写入。",
            reference_width_note(),
            "价格单位：元/吨，不含税；钢厂列不得填数字（防列错位）。",
            "上传页须手工填写生效年度与月份。",
            "本系统只导入基地页已启用的采购基地；宽表上其它基地列会跳过，样例四列只演示形状。",
        ],
    )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), "基准价格导入模板.xlsx"


def _write_notes(sheet: Worksheet, lines: list[str]) -> None:
    sheet["A1"] = "填写说明"
    sheet["A1"].font = Font(name="微软雅黑", size=12, bold=True, color="C00000")
    for idx, line in enumerate(lines, start=3):
        sheet.cell(row=idx, column=1, value=f"{idx - 2}. {line}")
        sheet.cell(row=idx, column=1).font = Font(name="微软雅黑", size=10)
        sheet.cell(row=idx, column=1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[idx].height = 22
    sheet.column_dimensions["A"].width = 96
