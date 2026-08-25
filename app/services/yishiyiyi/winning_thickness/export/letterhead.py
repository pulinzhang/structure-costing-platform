"""中标厚度价格一览 Excel — 外发抬头与打印页眉页脚。"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from typing import Protocol, cast

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.worksheet import Worksheet

from app.services.yishiyiyi.winning_thickness.export.styles import (
    CENTER,
    FORMULA_BANNER_FONT,
    FORMULA_FILL,
    META_BORDER,
    META_FILL,
    META_LABEL_FONT,
    META_VALUE_FONT,
    SUBTITLE_FILL,
    SUBTITLE_FONT,
    TITLE_FILL,
    TITLE_FONT,
    WRAP_CENTER,
)
from app.services.yishiyiyi.winning_thickness.export.values import (
    DOC_ORG,
    DOC_TITLE,
    PRICE_UNIT,
    TABLE_HEADER_ROW,
)


@dataclass(frozen=True, slots=True)
class WorkbookMeta:
    year: int
    month: int
    location: str
    row_count: int
    exported_at: datetime


def write_letterhead(
    ws: Worksheet,
    *,
    last_col: int,
    meta: WorkbookMeta,
    sheet_role: str,
    formula_note: str,
) -> int:
    fill_merged_row(
        ws,
        row=1,
        last_col=last_col,
        value=DOC_TITLE,
        font=TITLE_FONT,
        fill=TITLE_FILL,
        alignment=CENTER,
        height=34,
    )
    fill_merged_row(
        ws,
        row=2,
        last_col=last_col,
        value=f"{DOC_ORG}    {sheet_role}",
        font=SUBTITLE_FONT,
        fill=SUBTITLE_FILL,
        alignment=CENTER,
        height=20,
    )

    period = f"{meta.year}年{meta.month}月"
    date_text = meta.exported_at.strftime("%Y年%m月%d日")
    pairs = [
        ("基地", meta.location),
        ("期间", period),
        ("计价单位", PRICE_UNIT),
        ("记录数", str(meta.row_count)),
        ("导出日期", date_text),
    ]
    for index, (label, value) in enumerate(pairs):
        label_col = index * 2 + 1
        value_col = label_col + 1
        if value_col > last_col:
            break
        label_cell = ws.cell(row=3, column=label_col, value=label)
        label_cell.font = META_LABEL_FONT
        label_cell.alignment = CENTER
        value_cell = ws.cell(row=3, column=value_col, value=value)
        value_cell.font = META_VALUE_FONT
        value_cell.alignment = CENTER
    for col_idx in range(1, last_col + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.fill = META_FILL
        cell.border = META_BORDER
    ws.row_dimensions[3].height = 22

    fill_merged_row(
        ws,
        row=4,
        last_col=last_col,
        value=formula_note,
        font=FORMULA_BANNER_FONT,
        fill=FORMULA_FILL,
        alignment=WRAP_CENTER,
        height=28,
    )
    ws.row_dimensions[5].height = 10
    return TABLE_HEADER_ROW


def apply_sheet_chrome(
    ws: Worksheet,
    *,
    last_col: int,
    header_row: int,
    meta: WorkbookMeta,
    tab_color: str,
    freeze: bool = True,
) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab_color
    if freeze:
        ws.freeze_panes = f"A{header_row + 1}"
    ws.print_title_rows = f"1:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(
        left=0.5,
        right=0.5,
        top=0.6,
        bottom=0.6,
        header=0.25,
        footer=0.25,
    )
    ws.print_options.horizontalCentered = True
    period = f"{meta.location} {meta.year}年{meta.month}月"
    _set_header_footer_text(
        cast(_HeaderFooterItem | None, ws.oddHeader),
        left=DOC_TITLE,
        right=period,
    )
    _set_header_footer_text(
        cast(_HeaderFooterItem | None, ws.oddFooter),
        left="仅供核价核对使用",
        center="第 &P 页 / 共 &N 页",
        right=DOC_ORG,
    )
    last_letter = get_column_letter(last_col)
    last_row = max(ws.max_row, header_row)
    ws.print_area = f"A1:{last_letter}{last_row}"


class _HeaderFooterPart(Protocol):
    text: str


class _HeaderFooterItem(Protocol):
    left: _HeaderFooterPart
    center: _HeaderFooterPart
    right: _HeaderFooterPart


def _set_header_footer_text(
    item: _HeaderFooterItem | None,
    *,
    left: str = "",
    center: str = "",
    right: str = "",
) -> None:
    if item is None:
        return
    if left:
        item.left.text = left
    if center:
        item.center.text = center
    if right:
        item.right.text = right


def fill_merged_row(
    ws: Worksheet,
    *,
    row: int,
    last_col: int,
    value: str,
    font: Font,
    fill: PatternFill,
    alignment: Alignment,
    height: int,
) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    for col_idx in range(1, last_col + 1):
        cell = ws.cell(row=row, column=col_idx)
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment
        if col_idx == 1:
            cell.value = value
    ws.row_dimensions[row].height = height
