"""中标厚度价格一览 Excel — 计算用表（厚度上下限拆列）。"""

from __future__ import annotations

from decimal import Decimal

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.common.money import Q4, excel_money_cell
from app.schemas.vo.winning_thickness_vo import WinningThicknessRowVO
from app.services.yishiyiyi.winning_thickness.export.letterhead import (
    WorkbookMeta,
    apply_sheet_chrome,
    write_letterhead,
)
from app.services.yishiyiyi.winning_thickness.export.styles import (
    DATA_CENTER,
    DATA_FONT,
    DATA_LEFT,
    DATA_RIGHT,
    FINAL_COL_FILL,
    FINAL_FONT,
    FINAL_HEADER_FILL,
    HEADER_ALIGNMENT,
    HEADER_FILL,
    HEADER_FONT,
    PRICE_FORMAT,
    THICK_FORMAT,
    THIN_BORDER,
    zebra_fill,
)
from app.services.yishiyiyi.winning_thickness.export.values import (
    CALC_COLUMN_WIDTHS,
    CALC_FINAL_COL,
    CALC_HEADERS,
    CALC_HIDDEN_COLS,
    CALC_NUMBER_COLS,
    CALC_OPERATOR_COLS,
    CALC_PRICE_COLS,
    CALC_TAB_COLOR,
    CALC_THICK_COLS,
    excel_operator,
    excel_thickness,
)


def write_calc_sheet(
    ws: Worksheet,
    rows: list[WinningThicknessRowVO],
    meta: WorkbookMeta,
) -> None:
    last_col = len(CALC_HEADERS)
    header_row = write_letterhead(
        ws,
        last_col=last_col,
        meta=meta,
        sheet_role="计算用 · 可读厚度范围 + 公式匹配列（已隐藏）",
        formula_note=(
            "厚度范围与一览页相同，便于阅读。"
            "仅含当期有中标基价的牌号；"
            "最终价 = 基准价 + 加减价 + 加工费（Excel 公式，改左侧三项会自动重算）。"
        ),
    )
    write_calc_header(ws, header_row=header_row)
    data_start = header_row + 1
    for offset, row in enumerate(rows):
        write_calc_data_row(ws, data_start + offset, offset, row)
    apply_calc_column_widths(ws)
    apply_sheet_chrome(
        ws,
        last_col=last_col,
        header_row=header_row,
        meta=meta,
        tab_color=CALC_TAB_COLOR,
    )


def apply_calc_column_widths(ws: Worksheet) -> None:
    for col_idx, width in enumerate(CALC_COLUMN_WIDTHS, start=1):
        dim = ws.column_dimensions[get_column_letter(col_idx)]
        dim.width = width
        if col_idx in CALC_HIDDEN_COLS:
            dim.hidden = True


def write_calc_header(ws: Worksheet, *, header_row: int) -> None:
    for col_idx, title in enumerate(CALC_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = FINAL_HEADER_FILL if col_idx == CALC_FINAL_COL else HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER
    ws.row_dimensions[header_row].height = 28


def write_calc_data_row(
    ws: Worksheet,
    excel_row: int,
    data_index: int,
    row: WinningThicknessRowVO,
) -> None:
    values: list[Decimal | str | None] = [
        row.location or "",
        row.category_name or "",
        row.grade or "",
        row.steel_mill or "",
        excel_thickness(row.thick_min),
        excel_operator(row.thick_min_operator),
        excel_thickness(row.thick_max),
        excel_operator(row.thick_max_operator),
        row.thickness_range or "",
        excel_money_cell(row.base_price, Q4),
        excel_money_cell(row.thickness_adjustment, Q4, allow_negative=True),
        excel_money_cell(row.processing_fee, Q4),
        f"=J{excel_row}+K{excel_row}+L{excel_row}",
    ]
    stripe = zebra_fill(data_index)
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=excel_row, column=col_idx, value=value)
        cell.border = THIN_BORDER
        cell.fill = FINAL_COL_FILL if col_idx == CALC_FINAL_COL else stripe
        cell.font = FINAL_FONT if col_idx == CALC_FINAL_COL else DATA_FONT
        if col_idx in CALC_OPERATOR_COLS:
            cell.alignment = DATA_CENTER
        elif col_idx in CALC_NUMBER_COLS:
            cell.alignment = DATA_RIGHT
        else:
            cell.alignment = DATA_LEFT
        if col_idx in CALC_THICK_COLS and value not in ("", None):
            cell.number_format = THICK_FORMAT
        if col_idx in CALC_PRICE_COLS:
            cell.number_format = PRICE_FORMAT
