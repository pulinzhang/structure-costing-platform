"""中标厚度价格一览 Excel — 色板、字体、边框、数字格式。"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT_NAME = "微软雅黑"

THICK_FORMAT = "0.00"
PRICE_FORMAT = "#,##0.00"

INPUT_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
NOTE_FILL = PatternFill(start_color="F7F9FC", end_color="F7F9FC", fill_type="solid")
TITLE_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
SUBTITLE_FILL = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
META_FILL = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
FORMULA_FILL = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
FINAL_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
FINAL_COL_FILL = PatternFill(start_color="E8F0F8", end_color="E8F0F8", fill_type="solid")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
WEAK_HEADER_FILL = PatternFill(start_color="8FAADC", end_color="8FAADC", fill_type="solid")

HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
DATA_FONT = Font(name=FONT_NAME, size=10, color="000000")
WEAK_FONT = Font(name=FONT_NAME, size=8, color="808080")
TITLE_FONT = Font(name=FONT_NAME, size=18, bold=True, color="FFFFFF")
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, color="FFFFFF")
META_LABEL_FONT = Font(name=FONT_NAME, size=9, bold=True, color="5B7A9D")
META_VALUE_FONT = Font(name=FONT_NAME, size=10, bold=True, color="1F4E79")
SECTION_FONT = Font(name=FONT_NAME, size=11, bold=True, color="1F4E79")
# 示例页底部公式说明：辅助文案，字号刻意压低
NOTE_HEADING_FONT = Font(name=FONT_NAME, size=8, color="8A8A8A")
NOTE_FONT = Font(name=FONT_NAME, size=7, color="9A9A9A")
MUTED_FONT = Font(name=FONT_NAME, size=9, color="666666")
FORMULA_BANNER_FONT = Font(name=FONT_NAME, size=10, color="1F4E79")
FINAL_FONT = Font(name=FONT_NAME, size=10, bold=True, color="1F4E79")

HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")
DATA_LEFT = Alignment(horizontal="left", vertical="center")
DATA_RIGHT = Alignment(horizontal="right", vertical="center")
DATA_CENTER = Alignment(horizontal="center", vertical="center")
WRAP_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
META_BORDER = Border(
    left=Side(style="thin", color="C5D8EB"),
    right=Side(style="thin", color="C5D8EB"),
    top=Side(style="thin", color="C5D8EB"),
    bottom=Side(style="thin", color="C5D8EB"),
)


def zebra_fill(data_index: int) -> PatternFill:
    odd = data_index % 2 != 0
    color = "F7F9FC" if odd else "FFFFFF"
    return PatternFill(start_color=color, end_color=color, fill_type="solid")
