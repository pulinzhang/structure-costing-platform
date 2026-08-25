"""中标厚度价格一览 Excel — 使用示例页。"""

from __future__ import annotations

from decimal import Decimal

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.common.money import Q2, excel_money_cell
from app.schemas.vo.winning_thickness_vo import WinningThicknessRowVO
from app.services.yishiyiyi.winning_thickness.export.letterhead import (
    WorkbookMeta,
    apply_sheet_chrome,
    write_letterhead,
)
from app.services.yishiyiyi.winning_thickness.export.styles import (
    DATA_FONT,
    DATA_LEFT,
    DATA_RIGHT,
    FINAL_COL_FILL,
    FINAL_FONT,
    FINAL_HEADER_FILL,
    HEADER_ALIGNMENT,
    HEADER_FILL,
    HEADER_FONT,
    INPUT_FILL,
    MUTED_FONT,
    NOTE_FONT,
    NOTE_HEADING_FONT,
    PRICE_FORMAT,
    THICK_FORMAT,
    THIN_BORDER,
    WRAP_TOP,
    zebra_fill,
)
from app.services.yishiyiyi.winning_thickness.export.values import (
    CALC_SHEET_SUFFIX,
    EXAMPLE_TAB_COLOR,
    TABLE_HEADER_ROW,
    excel_sheet_ref,
    pick_example_rows,
    sample_thickness,
    scoped_sheet_name,
)

_HEADERS = [
    "基地",
    "钢种",
    "牌号",
    "钢厂",
    "厚度",
    "计算页行号",
    "基准价(元/吨)",
    "加减价(元/吨)",
    "加工费(元/吨)",
    "最终价(元/吨)",
]
_COLUMN_WIDTHS = [14, 14, 14, 12, 12, 12, 16, 16, 16, 16]
_THICK_COL = 5
_CALC_ROW_COL = 6
_PRICE_COLS = frozenset({7, 8, 9, 10})
_FINAL_COL = 10
_LETTERHEAD_COLS = 11


def write_example_sheet(
    ws: Worksheet,
    rows: list[WinningThicknessRowVO],
    meta: WorkbookMeta,
) -> None:
    examples = pick_example_rows(rows)
    header_row = write_letterhead(
        ws,
        last_col=_LETTERHEAD_COLS,
        meta=meta,
        sheet_role="使用示例 · 改黄色格料厚即可套价",
        formula_note=(
            "黄色格填写厚度；「计算页行号」为命中「计算用」页的行号；"
            "基准价、加减价、加工费按基地、牌号和厚度区间自动带出（不按钢厂筛选），"
            "最终价 = 三者之和。"
        ),
    )
    _write_header(ws, header_row)
    calc_start = TABLE_HEADER_ROW + 1
    calc_end = TABLE_HEADER_ROW + max(len(rows), 1)
    calc_sheet_name = scoped_sheet_name(
        year=meta.year,
        month=meta.month,
        location=meta.location,
        suffix=CALC_SHEET_SUFFIX,
    )
    data_start = header_row + 1
    if examples:
        for offset, row in enumerate(examples):
            _write_data_row(
                ws,
                data_start + offset,
                offset,
                row,
                calc_start=calc_start,
                calc_end=calc_end,
                calc_sheet_name=calc_sheet_name,
            )
        last_data_row = data_start + len(examples) - 1
    else:
        _write_data_row(
            ws,
            data_start,
            0,
            None,
            calc_start=calc_start,
            calc_end=calc_end,
            calc_sheet_name=calc_sheet_name,
        )
        last_data_row = data_start
        empty = ws.cell(row=data_start + 1, column=1, value="本次无数据可摘录。")
        empty.font = MUTED_FONT

    for col_idx, width in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 与主表示例表拉开间距；公式说明仅为辅助
    _write_formula_notes(ws, start_row=last_data_row + 5, data_start=data_start)
    apply_sheet_chrome(
        ws,
        last_col=_LETTERHEAD_COLS,
        header_row=header_row,
        meta=meta,
        tab_color=EXAMPLE_TAB_COLOR,
        freeze=False,
    )


def _write_header(ws: Worksheet, header_row: int) -> None:
    for col_idx, title in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = FINAL_HEADER_FILL if col_idx == _FINAL_COL else HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER
    ws.row_dimensions[header_row].height = 28


def _write_data_row(
    ws: Worksheet,
    excel_row: int,
    data_index: int,
    row: WinningThicknessRowVO | None,
    *,
    calc_start: int,
    calc_end: int,
    calc_sheet_name: str,
) -> None:
    thickness: Decimal | None = sample_thickness(row) if row is not None else None
    values: list[Decimal | str | int | None] = [
        (row.location or "") if row is not None else "",
        (row.category_name or "") if row is not None else "",
        (row.grade or "") if row is not None else "",
        (row.steel_mill or "") if row is not None else "",
        excel_money_cell(thickness, Q2) if thickness is not None else None,
        _lookup_calc_row_formula(excel_row, calc_start, calc_end, calc_sheet_name),
        _lookup_formula(excel_row, calc_start, calc_end, "J", calc_sheet_name),
        _lookup_formula(excel_row, calc_start, calc_end, "K", calc_sheet_name),
        _lookup_formula(excel_row, calc_start, calc_end, "L", calc_sheet_name),
        f'=IF(G{excel_row}="","",G{excel_row}+H{excel_row}+I{excel_row})',
    ]
    stripe = zebra_fill(data_index)
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=excel_row, column=col_idx, value=value)
        cell.border = THIN_BORDER
        cell.font = FINAL_FONT if col_idx == _FINAL_COL else DATA_FONT
        if col_idx == _THICK_COL:
            cell.fill = INPUT_FILL
            cell.alignment = DATA_RIGHT
            cell.number_format = THICK_FORMAT
        elif col_idx == _CALC_ROW_COL:
            cell.fill = stripe
            cell.alignment = DATA_RIGHT
        elif col_idx in _PRICE_COLS:
            cell.fill = FINAL_COL_FILL if col_idx == _FINAL_COL else stripe
            cell.alignment = DATA_RIGHT
            cell.number_format = PRICE_FORMAT
        else:
            cell.fill = stripe
            cell.alignment = DATA_LEFT


def _lookup_match_expr(
    excel_row: int,
    start: int,
    end: int,
    calc_sheet_name: str,
) -> str:
    sheet = excel_sheet_ref(calc_sheet_name)
    thickness = f"$E{excel_row}"
    loc = f"{sheet}!$A${start}:$A${end}"
    grade = f"{sheet}!$C${start}:$C${end}"
    thick_min = f"{sheet}!$E${start}:$E${end}"
    min_op = f"{sheet}!$F${start}:$F${end}"
    thick_max = f"{sheet}!$G${start}:$G${end}"
    max_op = f"{sheet}!$H${start}:$H${end}"
    keys = f"({loc}=$A{excel_row})*({grade}=$C{excel_row})"
    min_ok = (
        f'((({thick_min}="")+({min_op}=">=")*({thickness}>={thick_min})'
        f'+({min_op}=">")*({thickness}>{thick_min})'
        f'+({min_op}="")*({thick_min}<>"")*({thickness}>={thick_min}))>0)'
    )
    max_ok = (
        f'((({thick_max}="")+({max_op}="<=")*({thickness}<={thick_max})'
        f'+({max_op}="<")*({thickness}<{thick_max})'
        f'+({max_op}="")*({thick_max}<>"")*({thickness}<={thick_max}))>0)'
    )
    return f"{keys}*{min_ok}*{max_ok}"


def _lookup_formula(
    excel_row: int,
    start: int,
    end: int,
    value_col: str,
    calc_sheet_name: str,
) -> str:
    match = _lookup_match_expr(excel_row, start, end, calc_sheet_name)
    sheet = excel_sheet_ref(calc_sheet_name)
    values = f"{sheet}!${value_col}${start}:${value_col}${end}"
    return (
        f'=IF($E{excel_row}="","",'
        f'IF(SUMPRODUCT({match})=0,"",'
        f"SUMPRODUCT({match}*({values}))))"
    )


def _lookup_calc_row_formula(
    excel_row: int,
    start: int,
    end: int,
    calc_sheet_name: str,
) -> str:
    """命中「计算用」页的 Excel 行号，与基准价等同一套匹配条件。"""
    match = _lookup_match_expr(excel_row, start, end, calc_sheet_name)
    sheet = excel_sheet_ref(calc_sheet_name)
    row_numbers = f"ROW({sheet}!$A${start}:$A${end})"
    return (
        f'=IF($E{excel_row}="","",'
        f'IF(SUMPRODUCT({match})=0,"",'
        f"SUMPRODUCT({match}*({row_numbers}))))"
    )


def _write_formula_notes(ws: Worksheet, *, start_row: int, data_start: int) -> None:
    heading = ws.cell(row=start_row, column=1, value="公式说明")
    heading.font = NOTE_HEADING_FONT
    last_col = len(_HEADERS)
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=last_col)
    ws.row_dimensions[start_row].height = 12
    notes = [
        "黄色格为厚度输入；基地、牌号可改。匹配不按钢厂。",
        "最终价 = 基准价 + 加减价 + 加工费。"
        f' 例：=IF(G{data_start}="","",G{data_start}+H{data_start}+I{data_start})'
        " —— IF：基准价为空则结果为空，否则三列相加。",
        "计算页行号、基准价、加减价、加工费使用同一套匹配条件；"
        "取出字段不同：行号用 ROW，价格用「计算用」页 J/K/L 列。",
        '外层 IF($E="","", …)：厚度为空时不计算。',
        '内层 IF(SUMPRODUCT(条件)=0,"",SUMPRODUCT(条件*取值))：'
        "无命中行则空；有命中则对「条件×取值」求和，得到价格或行号。"
        "预期命中 1 行；若命中多行，数值会相加，需核对数据。",
        "SUMPRODUCT：对「计算用」页逐行求值。"
        "条件为 (基地相同)×(牌号相同)×(下限成立)×(上限成立)。"
        "Excel 中真=1、假=0；相乘表示「且」，相加表示「或」。"
        "四项均为 1 的行进入结果。",
        "「计算用」页 E–H 已隐藏（厚min、下限符号、厚max、上限符号）。"
        "阅读用 I 列「厚度范围」；公式用 E–H 判断厚度是否落入区间。",
        "下限匹配说明：展示文案为「0.60 ≤ t < 0.80」，"
        '存储拆为 厚min=0.60 与 下限符号=">="（文本，非 Excel 比较运算符）。'
        '公式用 (下限符号=">=") 或 (下限符号=">") 分支，再分别判断'
        " 厚度>=厚min 或 厚度>厚min。",
        "下限四段以加号连接（互斥，每行只成立一段），再整体与 0 比较："
        "① 厚min 空 → 无下限，成立；"
        '② 符号 ">=" → 厚度 ≥ 厚min；'
        '③ 符号 ">" → 厚度 > 厚min；'
        '④ 有厚min 且符号空 → 默认按 ">="。'
        "外层 (…)>0：加号累加的 0/1 之和大于 0 即下限成立。",
        '上限同理：厚max 空表示不限制；符号 "<=" / "<" 对应 ≤ / <；'
        '有厚max 且符号空时默认 "<="。常见区间为左闭右开（≥ 且 <）。',
        "核对方法：取消隐藏「计算用」E–H，对照厚min/符号/厚max，"
        "修改本页黄色厚度，查看「计算页行号」是否指向对应行。",
    ]
    for offset, text in enumerate(notes, start=1):
        excel_row = start_row + offset
        ws.merge_cells(
            start_row=excel_row,
            start_column=1,
            end_row=excel_row,
            end_column=last_col,
        )
        cell = ws.cell(row=excel_row, column=1, value=text)
        cell.font = NOTE_FONT
        cell.alignment = WRAP_TOP
        ws.row_dimensions[excel_row].height = 16 if len(text) < 80 else 22
