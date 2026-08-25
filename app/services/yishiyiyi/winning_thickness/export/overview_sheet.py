"""中标厚度价格一览 Excel — 外发阅读页。"""

from __future__ import annotations

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.common.money import Q4, excel_money_cell
from app.schemas.vo.winning_thickness_vo import WinningThicknessRowVO
from app.services.yishiyiyi.winning_thickness.export.letterhead import (
    WorkbookMeta,
    apply_sheet_chrome,
    write_letterhead,
)
from app.services.yishiyiyi.winning_thickness.export.styles import (
    DATA_FONT,
    DATA_LEFT,
    DATA_RIGHT,
    FINAL_COL_FILL,
    FINAL_HEADER_FILL,
    HEADER_ALIGNMENT,
    HEADER_FILL,
    HEADER_FONT,
    PRICE_FORMAT,
    THIN_BORDER,
    WEAK_FONT,
    WEAK_HEADER_FILL,
    zebra_fill,
)
from app.services.yishiyiyi.winning_thickness.export.values import (
    OVERVIEW_COLUMN_WIDTHS,
    OVERVIEW_FINAL_COL,
    OVERVIEW_HEADERS,
    OVERVIEW_PRICE_COLS,
    OVERVIEW_TAB_COLOR,
    OVERVIEW_WEAK_COLS,
    format_base_trace,
    format_location_category,
    format_thickness_trace,
)


def write_overview_sheet(
    ws: Worksheet,
    rows: list[WinningThicknessRowVO],
    meta: WorkbookMeta,
) -> None:
    last_col = len(OVERVIEW_HEADERS)
    header_row = write_letterhead(
        ws,
        last_col=last_col,
        meta=meta,
        sheet_role="外发阅读版 · 含基价与厚度政策溯源",
        formula_note=(
            "计价公式：最终价 = 基准价 + 厚度加减价 + 加工费。"
            "本表由系统按当期中标基价与厚度加减价政策自动生成，仅供核价核对。"
        ),
    )
    _write_header(ws, header_row)
    data_start = header_row + 1
    for offset, row in enumerate(rows):
        _write_data_row(ws, data_start + offset, offset, row)
    for col_idx, width in enumerate(OVERVIEW_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    apply_sheet_chrome(
        ws,
        last_col=last_col,
        header_row=header_row,
        meta=meta,
        tab_color=OVERVIEW_TAB_COLOR,
    )


def _write_header(ws: Worksheet, header_row: int) -> None:
    for col_idx, title in enumerate(OVERVIEW_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=title)
        cell.font = HEADER_FONT
        if col_idx == OVERVIEW_FINAL_COL:
            cell.fill = FINAL_HEADER_FILL
        elif col_idx in OVERVIEW_WEAK_COLS:
            cell.fill = WEAK_HEADER_FILL
        else:
            cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER
    ws.row_dimensions[header_row].height = 28


def _write_data_row(
    ws: Worksheet,
    excel_row: int,
    data_index: int,
    row: WinningThicknessRowVO,
) -> None:
    values = [
        format_location_category(row.location, row.category_name),
        row.grade,
        row.steel_mill,
        row.thickness_range,
        excel_money_cell(row.base_price, Q4),
        excel_money_cell(row.thickness_adjustment, Q4, allow_negative=True),
        excel_money_cell(row.processing_fee, Q4),
        excel_money_cell(row.final_price, Q4),
        format_base_trace(row),
        format_thickness_trace(row),
    ]
    stripe = zebra_fill(data_index)
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=excel_row, column=col_idx, value=value)
        cell.border = THIN_BORDER
        cell.fill = FINAL_COL_FILL if col_idx == OVERVIEW_FINAL_COL else stripe
        cell.font = WEAK_FONT if col_idx in OVERVIEW_WEAK_COLS else DATA_FONT
        cell.alignment = DATA_RIGHT if col_idx in OVERVIEW_PRICE_COLS else DATA_LEFT
        if col_idx in OVERVIEW_PRICE_COLS:
            cell.number_format = PRICE_FORMAT
