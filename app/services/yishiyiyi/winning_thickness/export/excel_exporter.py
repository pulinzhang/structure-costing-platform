"""中标钢厂厚度价格列表 Excel 生成（编排入口）。"""

from __future__ import annotations

import logging

from openpyxl import Workbook

from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.timezone import now
from app.schemas.vo.user_vo import UserVO
from app.schemas.vo.winning_thickness_vo import WinningThicknessRowVO
from app.services.yishiyiyi.winning_thickness.export.calc_sheet import write_calc_sheet
from app.services.yishiyiyi.winning_thickness.export.example_sheet import write_example_sheet
from app.services.yishiyiyi.winning_thickness.export.letterhead import WorkbookMeta
from app.services.yishiyiyi.winning_thickness.export.overview_sheet import write_overview_sheet
from app.services.yishiyiyi.winning_thickness.export.values import (
    CALC_SHEET_SUFFIX,
    EXAMPLE_SHEET_SUFFIX,
    OVERVIEW_SHEET_SUFFIX,
    build_filename,
    filter_rows_for_export,
    scoped_sheet_name,
)

logger = logging.getLogger("service.winning_thickness.export.excel")


class WinningThicknessExporter:
    """
    将一览行写成 xlsx（阅读页 + 计算页 + 示例页）。

    输入行须来自 assemble（基价表驱动）；本类只负责排版，不组装业务数据。
    """

    def generate(
        self,
        rows: list[WinningThicknessRowVO],
        *,
        year: int | None,
        month: int | None,
        location: str,
        current_user: UserVO,
    ) -> tuple[bytes, str]:
        export_rows = filter_rows_for_export(rows)
        if len(export_rows) != len(rows):
            logger.warning(
                "导出前丢弃无基价溯源行: dropped=%d kept=%d",
                len(rows) - len(export_rows),
                len(export_rows),
            )

        meta = WorkbookMeta(
            year=year,
            month=month,
            location=location or "",
            row_count=len(export_rows),
            exported_at=now(),
        )
        wb = Workbook()
        overview = wb.worksheets[0]
        overview.title = scoped_sheet_name(
            year=year,
            month=month,
            location=location,
            suffix=OVERVIEW_SHEET_SUFFIX,
        )
        write_overview_sheet(overview, export_rows, meta)

        calc_sheet = wb.create_sheet(
            scoped_sheet_name(
                year=year,
                month=month,
                location=location,
                suffix=CALC_SHEET_SUFFIX,
            )
        )
        write_calc_sheet(calc_sheet, export_rows, meta)

        example_sheet = wb.create_sheet(
            scoped_sheet_name(
                year=year,
                month=month,
                location=location,
                suffix=EXAMPLE_SHEET_SUFFIX,
            )
        )
        write_example_sheet(example_sheet, export_rows, meta)

        filename = build_filename(year=year, month=month, location=location or None)
        content = workbook_to_xlsx_bytes(wb, current_user)
        logger.info(
            "导出中标钢厂厚度价格列表完成: filename=%s rows=%d",
            filename,
            len(export_rows),
        )
        return content, filename
