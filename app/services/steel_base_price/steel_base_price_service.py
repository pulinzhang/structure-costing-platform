"""钢材基准价格业务服务。"""

import logging
from typing import Any

from sqlalchemy.orm import Session

from app.parsers.steel_base_price import parse_base_price_excel
from app.repositories.steel_base_price import SteelBasePriceRepository
from app.repositories.steel_base_price.dto import FilterOptions
from app.schemas.steel_base_price import (
    SteelBasePriceCreate,
    SteelBasePriceQuery,
    SteelBasePriceResponse,
    SteelBasePriceUpdate,
)


logger = logging.getLogger("service.steel_base_price")


class SteelBasePriceService:
    """钢材基准价格查询与导入业务。"""

    def __init__(self) -> None:
        self.repository = SteelBasePriceRepository()

    def paginate(
        self,
        db: Session,
        query: SteelBasePriceQuery,
    ) -> tuple[list[dict[str, Any]], int, FilterOptions]:
        """分页浏览基准价格。"""
        from app.repositories.steel_base_price.steel_base_price_repository import (
            BasePriceFilter,
        )

        skip = (query.page - 1) * query.size
        base_filter = BasePriceFilter(
            year=query.year,
            month=query.month,
            location=query.location.strip() if query.location else None,
            steel_mill=query.steel_mill.strip() if query.steel_mill else None,
            category_name=query.category_name.strip() if query.category_name else None,
            grade=query.grade.strip() if query.grade else None,
            skip=skip,
            limit=query.size,
        )
        records, total = self.repository.paginate(db, base_filter)
        rows = [
            SteelBasePriceResponse.model_validate(r).model_dump()
            for r in records
        ]
        return rows, total, self.repository.get_filter_options(db)

    def get_by_id(self, db: Session, record_id: int) -> dict[str, Any] | None:
        """获取单条记录。"""
        record = self.repository.get_by_id(db, record_id)
        if record is None:
            return None
        return SteelBasePriceResponse.model_validate(record).model_dump()

    def create(self, db: Session, body: SteelBasePriceCreate) -> tuple[dict, str | None]:
        """创建基准价格记录。"""
        try:
            data = body.model_dump()
            record = self.repository.create(db, data)
            db.commit()
            return SteelBasePriceResponse.model_validate(record).model_dump(), None
        except Exception as e:
            db.rollback()
            logger.exception("创建基准价格失败")
            return None, str(e)

    def update(self, db: Session, record_id: int, body: SteelBasePriceUpdate) -> tuple[dict | None, str | None]:
        """更新基准价格记录。"""
        try:
            data = body.model_dump(exclude_unset=True)
            if not data:
                return None, "没有提供任何更新字段"
            record = self.repository.update(db, record_id, data)
            if record is None:
                return None, "记录不存在"
            db.commit()
            return SteelBasePriceResponse.model_validate(record).model_dump(), None
        except Exception as e:
            db.rollback()
            logger.exception("更新基准价格失败: id=%s", record_id)
            return None, str(e)

    def delete(self, db: Session, record_id: int) -> tuple[bool, str | None]:
        """删除基准价格记录。"""
        try:
            success = self.repository.delete(db, record_id)
            if not success:
                return False, "记录不存在"
            db.commit()
            return True, None
        except Exception as e:
            db.rollback()
            logger.exception("删除基准价格失败: id=%s", record_id)
            return False, str(e)

    def batch_delete(
        self,
        db: Session,
        *,
        year: int,
        month: int,
        location: str | None = None,
        steel_mill: str | None = None,
    ) -> tuple[bool, str | None]:
        """按条件批量删除（年度+月份必须同时指定）。"""
        try:
            count = self.repository.delete_by_conditions(
                db,
                year=year,
                month=month,
                location=location,
                steel_mill=steel_mill,
            )
            db.commit()
            logger.info("批量删除基准价格: year=%s month=%s location=%s steel_mill=%s count=%s",
                year, month, location, steel_mill, count)
            return True, None
        except Exception as e:
            db.rollback()
            logger.exception("批量删除基准价格失败")
            return False, str(e)

    def import_excel(
        self,
        db: Session,
        *,
        file_content: bytes,
        filename: str,
        year: int,
        month: int,
    ) -> dict[str, Any]:
        """
        解析 Excel 并执行 upsert。

        Args:
            db: 数据库会话
            file_content: Excel 文件字节内容
            filename: 文件名（注入到 excel_filename 字段）
            year: 生效年度
            month: 生效月份

        Returns:
            导入结果（created/updated/total/warnings）
        """
        parse_result = parse_base_price_excel(file_content)

        # 注入文件名、年度、月份
        for record in parse_result.records:
            record["excel_filename"] = filename
            record["year"] = year
            record["month"] = month

        if not parse_result.records:
            return {
                "filename": filename,
                "created_count": 0,
                "updated_count": 0,
                "total": 0,
                "warnings": parse_result.warnings,
                "errors": parse_result.errors,
            }

        try:
            counts = self.repository.upsert(db, parse_result.records)
            db.commit()
        except Exception:
            db.rollback()
            logger.exception("导入基准价格失败: filename=%s", filename)
            raise

        logger.info(
            "导入基准价格完成: filename=%s, created=%d, updated=%d",
            filename,
            counts["created"],
            counts["updated"],
        )

        return {
            "filename": filename,
            "created_count": counts["created"],
            "updated_count": counts["updated"],
            "total": counts["created"] + counts["updated"],
            "warnings": parse_result.warnings,
            "errors": parse_result.errors,
        }

    def list_for_export(
        self,
        db: Session,
        *,
        year: int | None = None,
        month: int | None = None,
        location: str | None = None,
        steel_mill: str | None = None,
        skip: int = 0,
        limit: int | None = None,
    ) -> list[dict[str, Any]]:
        """导出预览数据（分批）。"""
        records = self.repository.list_for_export(
            db,
            year=year,
            month=month,
            location=location,
            steel_mill=steel_mill,
            skip=skip,
            limit=limit,
        )
        return [
            SteelBasePriceResponse.model_validate(r).model_dump()
            for r in records
        ]

    def count_for_export(
        self,
        db: Session,
        *,
        year: int | None = None,
        month: int | None = None,
        location: str | None = None,
        steel_mill: str | None = None,
    ) -> int:
        """导出数据统计。"""
        return self.repository.count_for_export(
            db,
            year=year,
            month=month,
            location=location,
            steel_mill=steel_mill,
        )

    def delete_by_conditions(
        self,
        db: Session,
        *,
        year: int,
        month: int,
        location: str | None = None,
        steel_mill: str | None = None,
    ) -> dict[str, Any]:
        """
        按条件批量删除（年度+月份必须同时指定）。

        Returns:
            {"deleted_count": N, "conditions": {...}}
        """
        count = self.repository.delete_by_conditions(
            db,
            year=year,
            month=month,
            location=location,
            steel_mill=steel_mill,
        )
        db.commit()
        return {
            "deleted_count": count,
            "conditions": {
                "year": year,
                "month": month,
                "location": location,
                "steel_mill": steel_mill,
            },
        }

    def generate_export_excel_bytes(
        self,
        db: Session,
        *,
        year: int | None = None,
        month: int | None = None,
        location: str | None = None,
        steel_mill: str | None = None,
    ) -> tuple[bytes, str]:
        """
        生成基准价格 Excel 导出文件。

        Returns:
            (文件字节数据, 文件名)
        """
        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        records = self.repository.list_for_export(
            db,
            year=year,
            month=month,
            location=location,
            steel_mill=steel_mill,
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "基准价格导出"

        # 表头样式
        header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        data_font = Font(name="微软雅黑", size=10, color="000000")
        data_alignment_left = Alignment(horizontal="left", vertical="center")
        data_alignment_right = Alignment(horizontal="right", vertical="center")

        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )

        headers = [
            "ID", "年度", "月份", "基地", "钢厂", "钢种大类",
            "牌号", "厚度(mm)", "宽度(mm)", "基准价格(元/吨)", "加工费(元/吨)",
            "文件名", "表页", "行号", "列号", "备注",
        ]

        # 写入表头
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # 写入数据
        for row_idx, record in enumerate(records, start=2):
            row_data = [
                record.id,
                record.year,
                record.month,
                record.location,
                record.steel_mill,
                record.category_name,
                record.grade,
                float(record.thick) if record.thick else "",
                float(record.width) if record.width else "",
                float(record.base_price) if record.base_price else "",
                float(record.processing_fee) if record.processing_fee else "",
                record.excel_filename,
                record.sheet_name,
                record.excel_row,
                record.excel_col if record.excel_col else "",
                record.remark or "",
            ]
            ws.append(row_data)

            fill = PatternFill(
                start_color="F2F2F2" if row_idx % 2 == 0 else "FFFFFF",
                end_color="F2F2F2" if row_idx % 2 == 0 else "FFFFFF",
                fill_type="solid",
            )
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                cell.fill = fill
                if col_idx in (1, 2, 3, 8, 9, 10, 11, 14, 15):
                    cell.alignment = data_alignment_right
                else:
                    cell.alignment = data_alignment_left

        # 设置列宽
        column_widths = [8, 8, 8, 12, 12, 10, 15, 12, 12, 14, 14, 30, 15, 8, 8, 20]
        for col_idx, width in enumerate(column_widths, start=1):
            letter = ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[letter].width = width

        ws.freeze_panes = "A2"

        # 构建文件名
        parts = ["基准价格导出"]
        if year:
            parts.append(f"{year}年")
        if month:
            parts.append(f"{month}月")
        if location:
            parts.append(location)
        if steel_mill:
            parts.append(steel_mill)
        filename = "_".join(parts) + ".xlsx"

        output = BytesIO()
        wb.save(output)
        return output.getvalue(), filename
