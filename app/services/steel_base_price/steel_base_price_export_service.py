"""钢材基准价格导出服务。"""

import logging
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from app.repositories.steel_base_price import SteelBasePriceRepository
from app.schemas.vo.steel_base_price_vo import SteelBasePriceResponse


logger = logging.getLogger("service.steel_base_price_export")


class SteelBasePriceExportService:
    """钢材基准价格导出业务。"""

    def __init__(self) -> None:
        self.repository = SteelBasePriceRepository()

    def export(
        self,
        db: Session,
        *,
        year: int | None = None,
        month: int | None = None,
        location: str | None = None,
        steel_mill: str | None = None,
        category_name: str | None = None,
        grade: str | None = None,
        spec: str | None = None,
    ) -> tuple[bytes, str]:
        """
        生成基准价格 Excel 导出文件。

        Returns:
            (文件字节数据, 文件名)
        """
        records = self.repository.list_for_export(
            db,
            year=year,
            month=month,
            location=location,
            steel_mill=steel_mill,
            category_name=category_name.strip() if category_name else None,
            grade=grade.strip() if grade else None,
            spec=spec,
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "基准价格导出"

        header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        data_font = Font(name="微软雅黑", size=10, color="000000")
        data_alignment_left = Alignment(horizontal="left", vertical="center")
        data_alignment_right = Alignment(horizontal="right", vertical="center")

        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )

        headers = [
            "ID", "年度", "月份", "基地", "钢厂", "钢种大类",
            "牌号", "规格", "厚度(mm)", "宽度(mm)", "基准价格(元/吨)", "加工费(元/吨)",
            "文件名", "表页", "行号", "列号", "备注",
        ]

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for row_idx, record in enumerate(records, start=2):
            row_data = [
                record.id,
                record.year,
                record.month,
                record.location,
                record.steel_mill,
                record.category_name,
                record.grade,
                record.spec or "",
                float(record.thick) if record.thick else "",
                int(float(record.width)) if record.width else "",
                float(record.base_price) if record.base_price else "",
                float(record.processing_fee) if record.processing_fee else "",
                record.excel_filename,
                record.sheet_name,
                record.excel_row,
                record.excel_col if record.excel_col else "",
                record.remark or "",
            ]
            ws.append(row_data)

            fill = PatternFill(
                start_color="F2F2F2" if row_idx % 2 == 0 else "FFFFFF",
                end_color="F2F2F2" if row_idx % 2 == 0 else "FFFFFF",
                fill_type="solid",
            )
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                cell.fill = fill
                if col_idx in (1, 2, 3, 9, 10, 11, 12, 15, 16):
                    cell.alignment = data_alignment_right
                else:
                    cell.alignment = data_alignment_left

        column_widths = [8, 8, 8, 12, 12, 10, 15, 15, 12, 12, 14, 14, 30, 15, 8, 8, 20]
        for col_idx, width in enumerate(column_widths, start=1):
            letter = ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[letter].width = width

        ws.freeze_panes = "A2"

        parts = ["基准价格导出"]
        if year:
            parts.append(f"{year}年")
        if month:
            parts.append(f"{month}月")
        if location:
            parts.append(location)
        if steel_mill:
            parts.append(steel_mill)
        filename = "_".join(parts) + ".xlsx"

        output = BytesIO()
        wb.save(output)
        logger.info("导出基准价格 Excel 完成: filename=%s, records=%d", filename, len(records))
        return output.getvalue(), filename
