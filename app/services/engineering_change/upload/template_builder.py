"""工程变更导入模板：青岛 T28 五种改法对照 + 后围同一总成合写。"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import cast

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from app.common.engineering_change_gates import (
    CHANGE_TYPE_ADD_PART,
    CHANGE_TYPE_NOT_INSTALL,
    CHANGE_TYPE_REVISE,
    CHANGE_TYPE_STOP,
    PART_LEVEL_ASSEMBLY,
    PART_LEVEL_L1,
    PART_LEVEL_L2,
)
from app.parsers.engineering_change.change_order_parser import CHANGE_ORDER_HEADERS
from app.services.bom.upload.template_builder import build_bom_material_template
from app.services.stamping.die.template_builder import COL_COUNT as DIE_COL_COUNT
from app.services.stamping.die.template_builder import build_die_layout_template
from app.services.stamping.upload.template_builder import build_stamping_process_template

TEMPLATE_FILENAME = "工程变更导入模板.xlsx"

LOCATION = "青岛"
VEHICLE_MODEL = "T28"
CUSTOMER_NAME = "奇瑞公司"
PERIOD_LABEL = "2025年07月"
CHANGE_VERSION = "EC-202507"
STAMPING_SUPPLIER = "烟台海众汽车零部件有限公司"
SLASH = "/"
ADD_PART_NUMBER = "501008903AA"
ADD_PART_NAME = "后围加强板"
REAR_PANEL_BEFORE = "501008897AA"
REAR_PANEL_AFTER = "501008897AB"
REAR_PANEL_NAME = "后围总成"
FIVE_TYPE_ADD_PART_NUMBER = "501007040AA"
_ASSEMBLY_CHANGE_ORDER_COLS = (3, 4, 5)

_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_DIE_THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_DATA_FONT = Font(name="微软雅黑", size=9)
_DIE_DATA_FONT = Font(name="宋体", size=12)
_DIE_NOTE_TITLE_FONT = Font(name="宋体", size=12, bold=True)
_DIE_NOTE_FONT = Font(name="宋体", size=10)
_DIE_NOTE_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
_NOT_ASSEMBLED_FILL = PatternFill("solid", fgColor="FFE4E4")
_ADD_PART_FILL = PatternFill("solid", fgColor="E2EFDA")
_DIE_ALLOC_NOTES = (
    "【模具费分摊说明】上传只读上方表格，不读本段。",
    "1. 先选采购基地并填写项目（新项目在本页填写）。模具方式或生产方式含「双槽 / 合并 / 一出二」→ 各序默认分摊 0.50；否则 1.00（单干）。左右件各记一半，避免同一副模记两次。",
    "2. 共模件若第一序单独开模（常见落料），上传后到模具清单把该序比例改成 1，后面共模序保持 0.50。",
    "3. 级进模写在生产方式。模具费不打 0.80（那是加工费折扣）。级进且共模仍是 0.50。",
    "4. 模具编号可空，系统按 DP-40-项目-流水自动补号；已填写的编号保持不变。",
    "5. 预估重量（T）= 长×宽×高×掏空×7.85÷10⁹。工序费 = 重量 × 单价 × 分摊比例 ÷ 分摊万台。",
)


@dataclass(frozen=True, slots=True)
class ChangeOrderSample:
    change_type: str
    before_assembly: str
    after_assembly: str | None
    assembly_name: str
    part_level: str
    before_part: str | None
    after_part: str | None
    part_name: str | None


@dataclass(frozen=True, slots=True)
class BomSample:
    seq_no: int
    internal_number: str
    part_number: str
    part_name: str
    weld_info: str
    l1_internal: str
    l1_part_number: str
    l1_part_name: str
    l1_qty: int | str
    l1_weld: str
    l2_internal: str
    l2_part_number: str
    l2_part_name: str
    l2_qty: int | str
    material: str
    thickness_text: str
    part_weight: float
    gauge_developed: str
    remark: str | None
    is_not_assembled: bool = False
    is_added_part: bool = False


@dataclass(frozen=True, slots=True)
class TemplateBook:
    change_orders: tuple[ChangeOrderSample, ...]
    bom: tuple[BomSample, ...]
    process: tuple[ProcessSample, ...]
    change_order_note: str


@dataclass(frozen=True, slots=True)
class ProcessSample:
    part_number: str
    part_name: str
    material: str
    thickness: float
    part_weight: float
    sheet_length: int
    sheet_width: int
    parts_per_sheet: int
    sheet_type: str
    process_seq: str
    process_code: str
    die_sets: int
    tonnage: int
    die_type: str


def _process_ops(
    part_number: str,
    part_name: str,
    material: str,
    thickness: float,
    part_weight: float,
    sheet_length: int,
    sheet_width: int,
    parts_per_sheet: int,
    sheet_type: str,
    die_type: str,
    operations: tuple[tuple[str, str, int], ...],
) -> tuple[ProcessSample, ...]:
    return tuple(
        ProcessSample(
            part_number,
            part_name,
            material,
            thickness,
            part_weight,
            sheet_length,
            sheet_width,
            parts_per_sheet,
            sheet_type,
            process_seq,
            process_code,
            1,
            tonnage,
            die_type,
        )
        for process_seq, process_code, tonnage in operations
    )


def _grouped_process_parts(
    process_samples: tuple[ProcessSample, ...],
) -> list[tuple[int, tuple[ProcessSample, ...]]]:
    groups: dict[str, list[ProcessSample]] = {}
    order: list[str] = []
    for sample in process_samples:
        if sample.part_number not in groups:
            groups[sample.part_number] = []
            order.append(sample.part_number)
        groups[sample.part_number].append(sample)
    return [
        (seq_no, tuple(groups[part_number])) for seq_no, part_number in enumerate(order, start=1)
    ]


def _assembly_change_order_key(sample: ChangeOrderSample) -> tuple[str, str, str]:
    return (sample.before_assembly, sample.after_assembly or "", sample.assembly_name)


def _merge_same_assembly_change_order_rows(
    sheet: Worksheet,
    change_orders: tuple[ChangeOrderSample, ...],
    first_row: int,
) -> None:
    start = 0
    while start < len(change_orders):
        end = start
        key = _assembly_change_order_key(change_orders[start])
        while (
            end + 1 < len(change_orders)
            and _assembly_change_order_key(change_orders[end + 1]) == key
        ):
            end += 1
        if end > start:
            for col in _ASSEMBLY_CHANGE_ORDER_COLS:
                _merge_vertical(sheet, col, first_row + start, first_row + end)
        start = end + 1


def _merge_vertical(sheet: Worksheet, col: int, start_row: int, end_row: int) -> None:
    if end_row > start_row:
        sheet.merge_cells(
            start_row=start_row,
            start_column=col,
            end_row=end_row,
            end_column=col,
        )


MIXED_CHANGE_ORDER_SAMPLES = (
    ChangeOrderSample(
        CHANGE_TYPE_REVISE,
        REAR_PANEL_BEFORE,
        REAR_PANEL_AFTER,
        REAR_PANEL_NAME,
        PART_LEVEL_L1,
        "501008899AA",
        "501008899AB",
        "后围内板",
    ),
    ChangeOrderSample(
        CHANGE_TYPE_REVISE,
        REAR_PANEL_BEFORE,
        REAR_PANEL_AFTER,
        REAR_PANEL_NAME,
        PART_LEVEL_L2,
        "501008901AA",
        "501008901AB",
        "后围左连接板",
    ),
    ChangeOrderSample(
        CHANGE_TYPE_NOT_INSTALL,
        REAR_PANEL_BEFORE,
        REAR_PANEL_AFTER,
        REAR_PANEL_NAME,
        PART_LEVEL_L1,
        "504000341AA",
        None,
        "左D柱下结构增强胶块",
    ),
    ChangeOrderSample(
        CHANGE_TYPE_ADD_PART,
        REAR_PANEL_BEFORE,
        REAR_PANEL_AFTER,
        REAR_PANEL_NAME,
        PART_LEVEL_L1,
        None,
        ADD_PART_NUMBER,
        ADD_PART_NAME,
    ),
)

MIXED_BOM_SAMPLES = (
    BomSample(
        seq_no=2,
        internal_number="ACA611",
        part_number=REAR_PANEL_AFTER,
        part_name=REAR_PANEL_NAME,
        weld_info="焊点：99 凸焊：7 涂胶：1550mm",
        l1_internal="ACA611001",
        l1_part_number="501008899AB",
        l1_part_name="后围内板",
        l1_qty="1",
        l1_weld=SLASH,
        l2_internal=SLASH,
        l2_part_number=SLASH,
        l2_part_name=SLASH,
        l2_qty=SLASH,
        material="DC01",
        thickness_text="T=0.7",
        part_weight=1.204,
        gauge_developed="是",
        remark=None,
    ),
    BomSample(
        seq_no=2,
        internal_number="ACA611",
        part_number=REAR_PANEL_AFTER,
        part_name=REAR_PANEL_NAME,
        weld_info="焊点：99 凸焊：7 涂胶：1550mm",
        l1_internal="ACA611N01",
        l1_part_number="501008901AB-1",
        l1_part_name="后围左连接板分总成",
        l1_qty="1",
        l1_weld="凸焊：2",
        l2_internal=SLASH,
        l2_part_number="501008901AB",
        l2_part_name="后围左连接板",
        l2_qty="1",
        material="HC180Y",
        thickness_text="T=1.5",
        part_weight=1.284,
        gauge_developed="是",
        remark="凸焊总成供货",
    ),
    BomSample(
        seq_no=2,
        internal_number="ACA611",
        part_number=REAR_PANEL_AFTER,
        part_name=REAR_PANEL_NAME,
        weld_info="焊点：99 凸焊：7 涂胶：1550mm",
        l1_internal="ACA611009",
        l1_part_number="504000341AA",
        l1_part_name="左D柱下结构增强胶块",
        l1_qty=0,
        l1_weld=SLASH,
        l2_internal=SLASH,
        l2_part_number=SLASH,
        l2_part_name=SLASH,
        l2_qty=SLASH,
        material="PA6+GF36",
        thickness_text="发泡料（环氧树脂等）",
        part_weight=0.136,
        gauge_developed="否",
        remark=CHANGE_TYPE_NOT_INSTALL,
        is_not_assembled=True,
    ),
    BomSample(
        seq_no=2,
        internal_number="ACA611",
        part_number=REAR_PANEL_AFTER,
        part_name=REAR_PANEL_NAME,
        weld_info="焊点：99 凸焊：7 涂胶：1550mm",
        l1_internal="ACA611011",
        l1_part_number=ADD_PART_NUMBER,
        l1_part_name=ADD_PART_NAME,
        l1_qty="1",
        l1_weld=SLASH,
        l2_internal=SLASH,
        l2_part_number=SLASH,
        l2_part_name=SLASH,
        l2_qty=SLASH,
        material="HC260Y",
        thickness_text="T=1.4",
        part_weight=0.343,
        gauge_developed=SLASH,
        remark=None,
        is_added_part=True,
    ),
)

MIXED_PROCESS_SAMPLES = (
    *_process_ops(
        "501008899AB",
        "后围内板",
        "DC01",
        0.7,
        1.204,
        1000,
        370,
        1,
        "条料",
        "铸造模",
        (
            ("OP10", "DR", 800),
            ("OP20", "TR+CTR+CPI", 400),
            ("OP30", "RST", 630),
            ("OP40", "TR+CTR+PI+CPI", 400),
        ),
    ),
    *_process_ops(
        "501008901AB",
        "后围左连接板",
        "HC180Y",
        1.5,
        1.284,
        860,
        520,
        2,
        "条料",
        "铸造模",
        (
            ("OP10", "DR", 800),
            ("OP20", "TR", 400),
            ("OP30", "TR+C/TR+SEP", 630),
            ("OP40", "FL+U/FL+RST", 630),
            ("OP50", "TR+PI+C/PI", 400),
        ),
    ),
    *_process_ops(
        ADD_PART_NUMBER,
        ADD_PART_NAME,
        "HC260Y",
        1.4,
        0.343,
        301,
        180,
        1,
        "卷料",
        "钢板模",
        (("OP10", "PRG", 315),),
    ),
)

FIVE_TYPE_CHANGE_ORDER_SAMPLES = (
    ChangeOrderSample(
        CHANGE_TYPE_STOP,
        "501009533AA",
        None,
        "发盖左铰链安装支架总成",
        PART_LEVEL_ASSEMBLY,
        None,
        None,
        None,
    ),
    ChangeOrderSample(
        CHANGE_TYPE_REVISE,
        "501006873AA",
        "501006873AB",
        "流水槽加强梁总成",
        PART_LEVEL_L1,
        "501006875AA",
        "501006875AB",
        "雨刮电机安装支架Ⅰ",
    ),
    ChangeOrderSample(
        CHANGE_TYPE_REVISE,
        "501009144AA",
        "501009144AB",
        "后防撞梁总成",
        PART_LEVEL_L2,
        "501009122AA",
        "501009122AB",
        "后保下安装支架",
    ),
    ChangeOrderSample(
        CHANGE_TYPE_NOT_INSTALL,
        "501008897AA",
        "501008897AB",
        "后围总成",
        PART_LEVEL_L1,
        "504000341AA",
        None,
        "左D柱下结构增强胶块",
    ),
    ChangeOrderSample(
        CHANGE_TYPE_ADD_PART,
        "501007029AA",
        "501007029AB",
        "中通道左后纵梁总成",
        PART_LEVEL_L1,
        None,
        FIVE_TYPE_ADD_PART_NUMBER,
        "中通道后纵梁加强板",
    ),
)

FIVE_TYPE_BOM_SAMPLES = (
    BomSample(
        seq_no=1,
        internal_number="ACA610",
        part_number="501006873AB",
        part_name="流水槽加强梁总成",
        weld_info="焊点：4 凸焊：2",
        l1_internal="ACA610001",
        l1_part_number="501006875AB",
        l1_part_name="雨刮电机安装支架Ⅰ",
        l1_qty="1",
        l1_weld=SLASH,
        l2_internal=SLASH,
        l2_part_number=SLASH,
        l2_part_name=SLASH,
        l2_qty=SLASH,
        material="DC51D+Z",
        thickness_text="T=1.2",
        part_weight=0.096,
        gauge_developed="是",
        remark=None,
    ),
    BomSample(
        seq_no=2,
        internal_number="ACA611",
        part_number="501008897AB",
        part_name="后围总成",
        weld_info="焊点：99 凸焊：7 涂胶：1550mm",
        l1_internal="ACA611009",
        l1_part_number="504000341AA",
        l1_part_name="左D柱下结构增强胶块",
        l1_qty=0,
        l1_weld=SLASH,
        l2_internal=SLASH,
        l2_part_number=SLASH,
        l2_part_name=SLASH,
        l2_qty=SLASH,
        material="PA6+GF36",
        thickness_text="发泡料（环氧树脂等）",
        part_weight=0.136,
        gauge_developed="否",
        remark=CHANGE_TYPE_NOT_INSTALL,
        is_not_assembled=True,
    ),
    BomSample(
        seq_no=3,
        internal_number="ACA612",
        part_number="501009144AB",
        part_name="后防撞梁总成",
        weld_info="焊缝：58条（1392mm) 凸焊：4",
        l1_internal="ACA612N01",
        l1_part_number="501009122AB-1",
        l1_part_name="后保下安装支架总成",
        l1_qty="2",
        l1_weld="凸焊:1",
        l2_internal=SLASH,
        l2_part_number="501009122AB",
        l2_part_name="后保下安装支架",
        l2_qty="1",
        material="HC180YD+Z",
        thickness_text="T=1.0",
        part_weight=0.041,
        gauge_developed="是",
        remark="凸焊总成供货",
    ),
    BomSample(
        seq_no=4,
        internal_number="ACA613",
        part_number="501007029AB",
        part_name="中通道左后纵梁总成",
        weld_info="焊点：7 焊缝：2条（52mm）",
        l1_internal="ACA613002",
        l1_part_number=FIVE_TYPE_ADD_PART_NUMBER,
        l1_part_name="中通道后纵梁加强板",
        l1_qty="1",
        l1_weld=SLASH,
        l2_internal=SLASH,
        l2_part_number=SLASH,
        l2_part_name=SLASH,
        l2_qty=SLASH,
        material="HC340/590DP",
        thickness_text="T=1.0",
        part_weight=0.235,
        gauge_developed=SLASH,
        remark=None,
        is_added_part=True,
    ),
)

FIVE_TYPE_PROCESS_SAMPLES = (
    ProcessSample(
        "501006875AB",
        "雨刮电机安装支架Ⅰ",
        "DC51D+Z",
        1.2,
        0.096,
        165,
        122,
        1,
        "卷料",
        "OP10",
        "PRG",
        1,
        315,
        "钢板模",
    ),
    ProcessSample(
        "501009122AB",
        "后保下安装支架",
        "HC180YD+Z",
        1,
        0.041,
        136,
        54,
        1,
        "卷料",
        "OP10",
        "PRG",
        1,
        250,
        "钢板模",
    ),
    ProcessSample(
        FIVE_TYPE_ADD_PART_NUMBER,
        "中通道后纵梁加强板",
        "HC340/590DP",
        1,
        0.235,
        149,
        230,
        1,
        "卷料",
        "OP10",
        "PRG",
        1,
        315,
        "钢板模",
    ),
)

CHANGE_ORDER_SAMPLES = FIVE_TYPE_CHANGE_ORDER_SAMPLES + MIXED_CHANGE_ORDER_SAMPLES
BOM_SAMPLES = (
    FIVE_TYPE_BOM_SAMPLES[0],
    *MIXED_BOM_SAMPLES,
    FIVE_TYPE_BOM_SAMPLES[2],
    FIVE_TYPE_BOM_SAMPLES[3],
)
PROCESS_SAMPLES = FIVE_TYPE_PROCESS_SAMPLES + MIXED_PROCESS_SAMPLES

COMBINED_CHANGE_ORDER_NOTE = (
    "一本对照样表，两种填法订在一起。后面三张表和日常导入模板相同。"
    f"基地 {LOCATION}，顾客 {CUSTOMER_NAME}，投产清单版本 D/00。"
    "上五行：五种改法，一行一个交货总成"
    f"（{CHANGE_TYPE_STOP} / 一级{CHANGE_TYPE_REVISE} / 二级{CHANGE_TYPE_REVISE} / "
    f"{CHANGE_TYPE_NOT_INSTALL} / {CHANGE_TYPE_ADD_PART}）。"
    f"下四行：同一个{REAR_PANEL_NAME} {REAR_PANEL_BEFORE} → {REAR_PANEL_AFTER} 同时四种改法；"
    "变更前总成、变更后总成、总成名称上下合并，只需填一次。"
    "先写原来的交货总成，再写变更类型和这一次动到的零件。"
    "材料清单只写变化的行；改图二级写出整条链。没变的零件不要抄。"
    f"{CHANGE_TYPE_STOP}：变更后总成空，材料清单不写这个总成。"
    f"{CHANGE_TYPE_NOT_INSTALL}：变更后零件空；用量 0、备注含「{CHANGE_TYPE_NOT_INSTALL}」。"
    f"{CHANGE_TYPE_ADD_PART}：加在原来的总成上；变更前零件空、变更后零件写新号。"
    "日常按交货总成拆开上传；整本对照不必一次进系统。"
)

COMBINED_BOOK = TemplateBook(
    change_orders=CHANGE_ORDER_SAMPLES,
    bom=BOM_SAMPLES,
    process=PROCESS_SAMPLES,
    change_order_note=COMBINED_CHANGE_ORDER_NOTE,
)

_ASSEMBLY_SKETCH_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf\xc0"
    b"\x00\x00\x00\x03\x00\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
)


def build_engineering_change_template(
    book: TemplateBook | None = None,
) -> tuple[bytes, str]:
    """运行时拼工程变更对照样表：五种改法 + 后围同一总成合写。"""
    samples = book or COMBINED_BOOK
    dest_wb = Workbook()
    order_sheet = dest_wb.active
    assert order_sheet is not None
    order_sheet.title = "工程变更单"
    _write_change_order(order_sheet, samples.change_orders, samples.change_order_note)

    bom_wb = load_workbook(BytesIO(build_bom_material_template()[0]))
    bom_sheet = _copy_sheet(bom_wb["车身件材料清单"], dest_wb, "车身件材料清单")
    _write_bom_header(bom_sheet)
    _clear_row_range(bom_sheet, 5, 200, 23)
    for offset, bom_sample in enumerate(samples.bom):
        _write_bom_row(bom_sheet, 5 + offset, bom_sample)
    _copy_sheet(bom_wb["填写说明"], dest_wb, "材料清单填写说明")

    process_wb = load_workbook(BytesIO(build_stamping_process_template()[0]))
    process_sheet = _copy_sheet(process_wb["冲压工艺确认表"], dest_wb, "冲压工艺确认表")
    _strip_process_sample(process_sheet)
    process_row = 7
    for seq_no, process_group in _grouped_process_parts(samples.process):
        process_row = _fill_process_part(
            process_sheet,
            process_group,
            start_row=process_row,
            seq_no=seq_no,
        )
    _copy_sheet(process_wb["填写说明"], dest_wb, "工艺确认填写说明")

    die_wb = load_workbook(BytesIO(build_die_layout_template()[0]))
    die_sheet = _copy_sheet(die_wb["零件汇总"], dest_wb, "零件汇总")
    _write_die_layout(die_sheet, samples.process)

    buffer = BytesIO()
    with TemporaryDirectory() as temp_dir:
        png_path = Path(temp_dir) / "assembly.png"
        png_path.write_bytes(_ASSEMBLY_SKETCH_PNG)
        _embed_assembly_sketches(bom_sheet, samples.bom, png_path)
        dest_wb.save(buffer)
    return buffer.getvalue(), TEMPLATE_FILENAME


def _write_change_order(
    sheet: Worksheet,
    change_orders: tuple[ChangeOrderSample, ...],
    change_order_note: str,
) -> None:
    title = Font(name="微软雅黑", size=14, bold=True)
    label = Font(name="微软雅黑", size=10)
    head = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4472C4")
    sheet["A1"] = "工程变更单"
    sheet["A1"].font = title
    sheet.merge_cells("A1:J1")
    sheet["A2"] = change_order_note
    sheet["A2"].font = Font(name="微软雅黑", size=9, color="666666")
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.merge_cells("A2:J2")
    sheet.row_dimensions[2].height = 148
    for col, header in enumerate(CHANGE_ORDER_HEADERS, start=1):
        cell = sheet.cell(4, col, header)
        cell.font = head
        cell.fill = fill
        cell.border = _THIN
        cell.alignment = _CENTER
    for offset, sample in enumerate(change_orders):
        excel_row = 5 + offset
        values = (
            VEHICLE_MODEL,
            PERIOD_LABEL,
            sample.before_assembly,
            sample.after_assembly or "",
            sample.assembly_name,
            sample.change_type,
            sample.part_level,
            sample.before_part or "",
            sample.after_part or "",
            sample.part_name or "",
        )
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(excel_row, col, value)
            cell.font = label
            cell.border = _THIN
            cell.alignment = _CENTER
            if col in {6, 7, 8, 9, 10}:
                if sample.change_type == CHANGE_TYPE_NOT_INSTALL:
                    cell.fill = _NOT_ASSEMBLED_FILL
                if sample.change_type == CHANGE_TYPE_ADD_PART:
                    cell.fill = _ADD_PART_FILL
        sheet.row_dimensions[excel_row].height = 28
    last_change_order_row = 4 + len(change_orders)
    sheet.merge_cells(f"A5:A{last_change_order_row}")
    sheet.merge_cells(f"B5:B{last_change_order_row}")
    _merge_same_assembly_change_order_rows(sheet, change_orders, 5)
    widths = (12, 14, 22, 22, 28, 12, 10, 22, 22, 28)
    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + col)].width = width
    sheet.freeze_panes = "A5"


def _write_bom_header(sheet: Worksheet) -> None:
    sheet["A1"] = f"{VEHICLE_MODEL}项目产品材料清单"
    sheet["A2"] = (
        "编制部门： 技术中心                    "
        f"编制日期：2025-07-01  版本： {CHANGE_VERSION}  编号：EC-T28-202507"
    )
    sheet["A3"] = f"项目代号：  {VEHICLE_MODEL}"
    sheet["I3"] = "产品类别：  车身件"
    sheet["N3"] = f"顾客名称：  {CUSTOMER_NAME}"
    sheet["P3"] = "阶段标识： D"
    sheet["R3"] = "■批量生产清单"
    sheet["V3"] = "修订日期：2025-07-01"


def _write_bom_row(sheet: Worksheet, row: int, sample: BomSample) -> None:
    fill = None
    if sample.is_not_assembled:
        fill = _NOT_ASSEMBLED_FILL
    elif sample.is_added_part:
        fill = _ADD_PART_FILL
    for col in range(1, 24):
        cell = sheet.cell(row, col)
        cell.font = _DATA_FONT
        cell.border = _THIN
        cell.alignment = _CENTER
        if fill is not None:
            cell.fill = fill
    sheet.row_dimensions[row].height = 22
    sheet.cell(row, 1, sample.seq_no)
    sheet.cell(row, 2, sample.internal_number)
    sheet.cell(row, 3, sample.part_number)
    sheet.cell(row, 4, sample.part_name)
    sheet.cell(row, 6, 1)
    sheet.cell(row, 7, sample.weld_info)
    sheet.cell(row, 8, sample.l1_internal)
    sheet.cell(row, 9, sample.l1_part_number)
    sheet.cell(row, 10, sample.l1_part_name)
    sheet.cell(row, 11, sample.l1_qty)
    sheet.cell(row, 12, sample.l1_weld)
    sheet.cell(row, 13, sample.l2_internal)
    sheet.cell(row, 14, sample.l2_part_number)
    sheet.cell(row, 15, sample.l2_part_name)
    sheet.cell(row, 17, sample.l2_qty)
    sheet.cell(row, 18, sample.material)
    sheet.cell(row, 19, sample.thickness_text)
    sheet.cell(row, 21, sample.part_weight)
    sheet.cell(row, 22, sample.gauge_developed)
    sheet.cell(row, 23, sample.remark)


def _embed_assembly_sketches(
    sheet: Worksheet, samples: tuple[BomSample, ...], png_path: Path
) -> None:
    seen: set[str] = set()
    for offset, sample in enumerate(samples):
        if sample.part_number in seen:
            continue
        seen.add(sample.part_number)
        copy_path = png_path.parent / f"assembly_{offset}.png"
        copy_path.write_bytes(png_path.read_bytes())
        image = ExcelImage(str(copy_path))
        image.width = 36
        image.height = 36
        sheet.add_image(image, f"E{5 + offset}")


def _strip_process_sample(sheet: Worksheet) -> None:
    sheet["D2"] = f"{VEHICLE_MODEL}项目冲压工艺确认表-工程变更"
    sheet["D3"] = STAMPING_SUPPLIER
    sheet["S3"] = CHANGE_VERSION
    max_row = sheet.max_row or 6
    _clear_row_range(sheet, 7, max(max_row, 7), 20)


def _fill_process_part(
    sheet: Worksheet,
    samples: tuple[ProcessSample, ...],
    *,
    start_row: int,
    seq_no: int,
) -> int:
    first = samples[0]
    end_row = start_row + len(samples) - 1
    for excel_row in range(start_row, end_row + 1):
        for col in range(1, 20):
            cell = sheet.cell(excel_row, col)
            cell.border = _THIN
            cell.font = _DATA_FONT
            cell.alignment = _CENTER
        sheet.row_dimensions[excel_row].height = 22
    for col in (1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 19):
        _merge_vertical(sheet, col, start_row, end_row)
    sheet.cell(start_row, 1, seq_no)
    sheet.cell(start_row, 2, first.part_number)
    sheet.cell(start_row, 3, first.part_name)
    sheet.cell(start_row, 5, first.material)
    sheet.cell(start_row, 6, first.thickness)
    sheet.cell(start_row, 7, first.part_weight)
    sheet.cell(start_row, 8, first.sheet_length)
    sheet.cell(start_row, 9, first.sheet_width)
    sheet.cell(start_row, 10, first.parts_per_sheet)
    sheet.cell(start_row, 13, first.sheet_type)
    sheet.cell(start_row, 19, first.die_type)
    for offset, sample in enumerate(samples):
        excel_row = start_row + offset
        sheet.cell(excel_row, 14, sample.process_seq)
        sheet.cell(excel_row, 15, sample.process_code)
        sheet.cell(excel_row, 16, sample.die_sets)
        sheet.cell(excel_row, 18, sample.tonnage)
    return end_row + 1


def _write_die_layout(
    sheet: Worksheet,
    process_samples: tuple[ProcessSample, ...],
) -> None:
    sheet.cell(1, 1, f"{VEHICLE_MODEL}项目冲压模具布点工艺清单")
    sheet.cell(1, 30, CHANGE_VERSION)
    _strip_die_below_header(sheet)
    excel_row = 6
    for seq_no, samples in _grouped_process_parts(process_samples):
        excel_row = _fill_die_part(sheet, samples, start_row=excel_row, seq_no=seq_no)
    _write_die_alloc_notes(sheet, excel_row - 1)


def _strip_die_below_header(sheet: Worksheet) -> None:
    max_row = max(sheet.max_row or 6, 40)
    _clear_row_range(sheet, 6, max_row, DIE_COL_COUNT)


def _fill_die_part(
    sheet: Worksheet,
    samples: tuple[ProcessSample, ...],
    *,
    start_row: int,
    seq_no: int,
) -> int:
    first = samples[0]
    end_row = start_row + len(samples) - 1
    for excel_row in range(start_row, end_row + 1):
        for col in range(1, DIE_COL_COUNT + 1):
            cell = sheet.cell(excel_row, col)
            cell.border = _DIE_THIN
            cell.font = _DIE_DATA_FONT
            cell.alignment = _CENTER
        sheet.row_dimensions[excel_row].height = 18.75
    for col in (1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 23, 25):
        _merge_vertical(sheet, col, start_row, end_row)
    sheet.cell(start_row, 1, seq_no)
    sheet.cell(start_row, 2, first.part_number)
    sheet.cell(start_row, 3, first.part_name)
    sheet.cell(start_row, 5, first.material)
    sheet.cell(start_row, 6, first.thickness)
    sheet.cell(start_row, 7, first.part_weight)
    sheet.cell(start_row, 8, first.sheet_length)
    sheet.cell(start_row, 9, first.sheet_width)
    sheet.cell(start_row, 10, first.sheet_type)
    sheet.cell(start_row, 11, first.parts_per_sheet)
    sheet.cell(start_row, 25, first.die_type)
    for offset, sample in enumerate(samples):
        excel_row = start_row + offset
        sheet.cell(excel_row, 14, sample.process_seq)
        sheet.cell(excel_row, 15, sample.process_code)
        sheet.cell(excel_row, 16, sample.die_sets)
        sheet.cell(excel_row, 24, sample.tonnage)
    return end_row + 1


def _write_die_alloc_notes(sheet: Worksheet, data_end_row: int) -> None:
    start_row = data_end_row + 2
    for offset, line in enumerate(_DIE_ALLOC_NOTES):
        excel_row = start_row + offset
        sheet.merge_cells(
            start_row=excel_row,
            start_column=1,
            end_row=excel_row,
            end_column=DIE_COL_COUNT,
        )
        cell = sheet.cell(excel_row, 1, line)
        cell.font = _DIE_NOTE_TITLE_FONT if offset == 0 else _DIE_NOTE_FONT
        cell.alignment = _DIE_NOTE_ALIGN
        sheet.row_dimensions[excel_row].height = 22


def _unmerge_in_rows(sheet: Worksheet, min_row: int, max_row: int) -> None:
    ranges = [
        str(item)
        for item in sheet.merged_cells.ranges
        if item.min_row >= min_row and item.max_row <= max_row
    ]
    for item in ranges:
        sheet.unmerge_cells(item)


def _clear_row_range(sheet: Worksheet, start_row: int, end_row: int, max_col: int) -> None:
    _unmerge_in_rows(sheet, start_row, end_row)
    for row in range(start_row, end_row + 1):
        for col in range(1, max_col + 1):
            sheet.cell(row, col).value = None


def _copy_sheet(source: Worksheet, dest_wb: Workbook, title: str) -> Worksheet:
    dest = cast(Worksheet, dest_wb.create_sheet(title))
    for row in source.iter_rows():
        for cell in row:
            cloned = dest.cell(cell.row, cell.column, cell.value)
            if cell.has_style:
                cloned.font = copy(cell.font)
                cloned.border = copy(cell.border)
                cloned.fill = copy(cell.fill)
                cloned.number_format = cell.number_format
                cloned.protection = copy(cell.protection)
                cloned.alignment = copy(cell.alignment)
    for merged in source.merged_cells.ranges:
        dest.merge_cells(str(merged))
    for letter, col_dim in source.column_dimensions.items():
        dest.column_dimensions[letter].width = col_dim.width
    for index, row_dim in source.row_dimensions.items():
        dest.row_dimensions[index].height = row_dim.height
    dest.freeze_panes = source.freeze_panes
    dest.sheet_view.zoomScale = source.sheet_view.zoomScale or 100
    dest.print_title_rows = source.print_title_rows
    dest.print_area = source.print_area
    return dest
