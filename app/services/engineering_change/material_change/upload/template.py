"""材料变更导入模板：按所选基地、项目和生效年月写两行举例。"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.parsers.engineering_change.material_change.codes import (
    HEADERS,
    INTERCHANGE_REQUIRED,
    SHEET_HELP,
    SHEET_ORDER,
)

PLACEHOLDER_SUPPLIER = "请填写冲压工厂"

_EXAMPLE_ROWS = (
    ("EX-MAT-AA", "样例零件甲", "DC01", 1.5, "DC03", 1.5),
    ("EX-MAT-AB", "样例零件乙", "DC01", 1.2, "DC03", 1.2),
)

_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
_DATA_FONT = Font(name="微软雅黑", size=9)
_TITLE_FONT = Font(name="微软雅黑", size=14, bold=True, color="1F4E79")
_HELP_FONT = Font(name="微软雅黑", size=10)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_WIDTHS = (10, 10, 14, 14, 36, 16, 22, 14, 12, 14, 12, 16)
_THICKNESS_COLUMNS = frozenset({9, 11})


@dataclass(frozen=True, slots=True)
class MaterialChangeTemplateMeta:
    location: str
    vehicle_model: str
    change_version: str
    period_label: str
    stamping_supplier: str = PLACEHOLDER_SUPPLIER


def build_material_change_template(meta: MaterialChangeTemplateMeta) -> tuple[bytes, str]:
    workbook = Workbook()
    help_sheet = workbook.active
    assert help_sheet is not None
    help_sheet.title = SHEET_HELP
    _write_help(help_sheet, meta)
    _write_order(workbook.create_sheet(SHEET_ORDER), meta)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    filename = f"材料变更导入模板_{meta.location}_{meta.vehicle_model}.xlsx"
    return buffer.getvalue(), filename


def _help_lines(meta: MaterialChangeTemplateMeta) -> tuple[str, ...]:
    return (
        "材料变更填写说明（格子即规格）",
        "",
        f"这一本按「{meta.location} / {meta.vehicle_model} / {meta.period_label}」来写。"
        f"变更版本是 {meta.change_version}，必须和生效年月对上。",
        "只准两张表：本表 +「工程变更单（材质）」。说明表只认表名。",
        "第 1 行必须是规定的十二列表头。除料厚外，格子必须是文字。禁止合并、公式、贴图。",
        "变更前料厚、变更后料厚必须是数字，大于 0，最多两位小数，例如 1.5。不要写成文字。",
        "变更版本 MAT- 加六位数字，且等于生效年月。生效年月可以写成 2026年6月 或 2026年06月，也可以直接选日期。",
        "互换性只准「不影响互换性」。变更前必须等于库里最新材料。",
        "下面两行是举例。零件号、零件名称、冲压工厂、材料和料厚都要换成这个项目正在用的，不要原样上传。",
        "样例零件号是 EX-MAT-AA、EX-MAT-AB。基地、项目、变更版本、生效年月已经按下载时选的填好。",
    )


def _write_help(sheet: Worksheet, meta: MaterialChangeTemplateMeta) -> None:
    lines = _help_lines(meta)
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = lines[0]
    sheet["A1"].font = _TITLE_FONT
    for index, line in enumerate(lines[1:], start=2):
        sheet.cell(index, 1, line).font = _HELP_FONT
        sheet.cell(index, 1).alignment = _LEFT
    sheet.column_dimensions["A"].width = 110


def _write_order(sheet: Worksheet, meta: MaterialChangeTemplateMeta) -> None:
    for col, title in enumerate(HEADERS, start=1):
        cell = sheet.cell(1, col, title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN
        sheet.column_dimensions[get_column_letter(col)].width = _WIDTHS[col - 1]
    for row_index, example in enumerate(_EXAMPLE_ROWS, start=2):
        (
            part_number,
            part_name,
            before_material,
            before_thickness,
            after_material,
            after_thickness,
        ) = example
        values: tuple[str | float, ...] = (
            meta.location,
            meta.vehicle_model,
            meta.change_version,
            meta.period_label,
            meta.stamping_supplier,
            part_number,
            part_name,
            before_material,
            before_thickness,
            after_material,
            after_thickness,
            INTERCHANGE_REQUIRED,
        )
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, col, value)
            cell.font = _DATA_FONT
            cell.alignment = _CENTER
            cell.border = _THIN
            if col in _THICKNESS_COLUMNS:
                cell.number_format = "0.00"
            else:
                cell.number_format = "@"
        sheet.row_dimensions[row_index].height = 22
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{1 + len(_EXAMPLE_ROWS)}"
