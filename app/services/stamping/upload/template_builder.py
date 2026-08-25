"""冲压工艺 Excel 导入模板生成。"""

from __future__ import annotations

from io import BytesIO
from typing import TypedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


class _ProcessSample(TypedDict):
    op: str
    code: str
    die_sets: int
    tonnage: int


class _BlankSegment(TypedDict):
    """拼焊子片：卷料宽/步距纵向切块（行数 ≠ 零件号合并行数）。"""

    start_row: int
    end_row: int
    length: float
    width: float


class _PartSample(TypedDict):
    start_row: int
    end_row: int
    seq: int
    part_number: str
    part_name: str
    material: str
    thickness: float | str
    weight: float
    length: float | None
    width: float | None
    count: int
    sheet_type: str
    die_type: str
    processes: list[_ProcessSample]
    blank_segments: list[_BlankSegment] | None


def _merge_vertical(sheet: Worksheet, col: int, start_row: int, end_row: int) -> None:
    if end_row > start_row:
        sheet.merge_cells(
            start_row=start_row,
            start_column=col,
            end_row=end_row,
            end_column=col,
        )


def build_stamping_process_template() -> tuple[bytes, str]:
    """生成冲压工艺确认表样例（含普通件 + 拼焊件，对齐海众拼焊切块规则）。"""
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "冲压工艺确认表"

    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    title_font = Font(name="微软雅黑", size=14, bold=True)
    data_font = Font(name="微软雅黑", size=9)
    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.merge_cells("D2:R2")
    sheet["D2"] = "T28样例项目冲压工艺确认表-模板"
    sheet["D2"].font = title_font
    sheet["D2"].alignment = center

    sheet["A3"] = "冲压承制公司："
    sheet.merge_cells("D3:F3")
    sheet["D3"] = "样例冲压有限公司"
    sheet.merge_cells("Q3:R3")
    sheet["Q3"] = "版本"
    sheet["Q3"].alignment = center
    sheet.merge_cells("S3:T3")
    sheet["S3"] = "V1"
    sheet["S3"].alignment = center

    vertical_headers: list[tuple[int, str]] = [
        (1, "序号"),
        (2, "零件号"),
        (3, "零件名称"),
        (4, "零件简图"),
        (5, "材质"),
        (6, "料厚"),
        (7, "净重"),
        (17, "模具编号"),
        (18, "冲压设备"),
        (19, "模具类型"),
    ]
    for col, title in vertical_headers:
        sheet.cell(row=5, column=col, value=title)
        _merge_vertical(sheet, col, 5, 6)

    sheet.cell(row=5, column=8, value="材料尺寸信息")
    sheet.merge_cells("H5:M5")
    sheet.cell(row=5, column=14, value="工艺")
    sheet.merge_cells("N5:P5")

    secondary: dict[int, str] = {
        8: "长（卷料宽）",
        9: "宽（卷料步距）",
        10: "可制件数",
        11: "单件消耗",
        12: "材料利用率",
        13: "料片式样",
        14: "工序名称",
        15: "工序",
        16: "模具套数",
    }
    for col, title in secondary.items():
        sheet.cell(row=6, column=col, value=title)

    for row in (5, 6):
        for col in range(1, 20):
            cell = sheet.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = thin

    parts: list[_PartSample] = [
        {
            "start_row": 7,
            "end_row": 9,
            "seq": 1,
            "part_number": "501006875AA",
            "part_name": "雨刮电机安装支架Ⅰ",
            "material": "DC51D+Z",
            "thickness": 1.2,
            "weight": 0.096,
            "length": 320.0,
            "width": 180.0,
            "count": 1,
            "sheet_type": "矩形料",
            "die_type": "钢板模",
            "processes": [
                {"op": "OP10", "code": "BL", "die_sets": 1, "tonnage": 400},
                {"op": "OP20", "code": "PI", "die_sets": 1, "tonnage": 400},
                {"op": "OP30", "code": "FL", "die_sets": 1, "tonnage": 630},
            ],
            "blank_segments": None,
        },
        {
            "start_row": 10,
            "end_row": 12,
            "seq": 2,
            "part_number": "501008899AA",
            "part_name": "后围内板",
            "material": "HC340/590DP",
            "thickness": 1.0,
            "weight": 4.52,
            "length": 1480.0,
            "width": 620.0,
            "count": 1,
            "sheet_type": "异形料",
            "die_type": "铸造模",
            "processes": [
                {"op": "OP10", "code": "DR", "die_sets": 1, "tonnage": 1000},
                {"op": "OP20", "code": "TR+PI", "die_sets": 1, "tonnage": 800},
                {"op": "OP30", "code": "RST", "die_sets": 1, "tonnage": 630},
            ],
            "blank_segments": None,
        },
        {
            "start_row": 13,
            "end_row": 14,
            "seq": 3,
            "part_number": "502011828AA",
            "part_name": "左B柱内板本体",
            "material": "HC600/980QP",
            "thickness": 1.0,
            "weight": 1.69,
            "length": 1290.0,
            "width": 442.5,
            "count": 1,
            "sheet_type": "直角梯形料",
            "die_type": "铸造模",
            "processes": [
                {"op": "OP10", "code": "DR", "die_sets": 1, "tonnage": 1300},
                {"op": "OP20", "code": "TR+PI", "die_sets": 1, "tonnage": 630},
            ],
            "blank_segments": None,
        },
        # 拼焊样例：零件号整段合并；卷料宽拆成两块（行数 ≠ 零件跨度）→ tailor_weld_parser 判定为拼焊
        {
            "start_row": 15,
            "end_row": 20,
            "seq": 4,
            "part_number": "551001690AA",
            "part_name": "左前门外窗框加强板",
            "material": "HC220YD+Z",
            "thickness": "0.70/1.00",
            "weight": 1.56,
            "length": None,
            "width": None,
            "count": 1,
            "sheet_type": "拼焊板：见2D图",
            "die_type": "铸造模",
            "processes": [
                {"op": "OP05", "code": "两块方料&拼焊", "die_sets": 1, "tonnage": 0},
                {"op": "OP10", "code": "DR", "die_sets": 1, "tonnage": 1250},
                {"op": "OP20", "code": "TR/PI", "die_sets": 1, "tonnage": 630},
                {"op": "OP30", "code": "UFL/RST/TR", "die_sets": 1, "tonnage": 800},
                {"op": "OP40", "code": "CPI/PI/RST", "die_sets": 1, "tonnage": 630},
            ],
            "blank_segments": [
                {"start_row": 15, "end_row": 17, "length": 550.0, "width": 465.0},
                {"start_row": 18, "end_row": 20, "length": 555.0, "width": 465.0},
            ],
        },
    ]

    # 普通件：零件字段整段纵向合并（含长/宽）；拼焊件：长宽按子片切块，其余仍整段合并
    common_merge_cols = (1, 2, 3, 4, 5, 6, 7, 10, 11, 12, 13, 19)

    for part in parts:
        start_row = part["start_row"]
        end_row = part["end_row"]
        blank_segments = part["blank_segments"]
        is_tailor = blank_segments is not None and len(blank_segments) >= 2

        for col in common_merge_cols:
            _merge_vertical(sheet, col, start_row, end_row)

        if is_tailor:
            assert blank_segments is not None
            for segment in blank_segments:
                _merge_vertical(sheet, 8, segment["start_row"], segment["end_row"])
                _merge_vertical(sheet, 9, segment["start_row"], segment["end_row"])
        else:
            _merge_vertical(sheet, 8, start_row, end_row)
            _merge_vertical(sheet, 9, start_row, end_row)

        for row in range(start_row, end_row + 1):
            for col in range(1, 20):
                cell = sheet.cell(row=row, column=col)
                cell.border = thin
                cell.font = data_font
                cell.alignment = center

        sheet.cell(row=start_row, column=1, value=part["seq"])
        sheet.cell(row=start_row, column=2, value=part["part_number"])
        sheet.cell(row=start_row, column=3, value=part["part_name"])
        sheet.cell(row=start_row, column=5, value=part["material"])
        sheet.cell(row=start_row, column=6, value=part["thickness"])
        sheet.cell(row=start_row, column=7, value=part["weight"])
        sheet.cell(row=start_row, column=10, value=part["count"])
        sheet.cell(row=start_row, column=13, value=part["sheet_type"])
        sheet.cell(row=start_row, column=19, value=part["die_type"])

        if is_tailor:
            assert blank_segments is not None
            for segment in blank_segments:
                sheet.cell(row=segment["start_row"], column=8, value=segment["length"])
                sheet.cell(row=segment["start_row"], column=9, value=segment["width"])
        else:
            sheet.cell(row=start_row, column=8, value=part["length"])
            sheet.cell(row=start_row, column=9, value=part["width"])

        for offset, process in enumerate(part["processes"]):
            row = start_row + offset
            if row > end_row:
                break
            sheet.cell(row=row, column=14, value=process["op"])
            sheet.cell(row=row, column=15, value=process["code"])
            sheet.cell(row=row, column=16, value=process["die_sets"])
            tonnage = process["tonnage"]
            if tonnage > 0:
                sheet.cell(row=row, column=18, value=tonnage)

    widths = [6, 16, 20, 10, 14, 8, 10, 12, 12, 8, 10, 10, 14, 10, 14, 8, 14, 10, 10]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.column_dimensions["S"].width = 8
    sheet.column_dimensions["T"].width = 8
    sheet.row_dimensions[5].height = 22
    sheet.row_dimensions[6].height = 28
    sheet.freeze_panes = "A7"

    note = workbook.create_sheet("填写说明")
    _write_notes(
        note,
        [
            "本模板对齐「车身件项目冲压工艺确认表」常见排版，含普通件与拼焊件样例。",
            "普通件：零件号与「长/卷料宽」合并行数相同。",
            "拼焊件：零件号整段合并，但「长/卷料宽」拆成多块（行数不等）；料厚写 0.70/1.00 按片对齐。",
            "拼焊样例见序号4（551001690AA）：H15:H17=550、H18:H20=555。",
            "上传页须先选择采购基地；同基地+车型+版本再次上传将覆盖。",
        ],
    )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), "冲压工艺导入模板.xlsx"


def _write_notes(sheet: Worksheet, lines: list[str]) -> None:
    sheet["A1"] = "填写说明"
    sheet["A1"].font = Font(name="微软雅黑", size=12, bold=True, color="C00000")
    for idx, line in enumerate(lines, start=3):
        sheet.cell(row=idx, column=1, value=f"{idx - 2}. {line}")
        sheet.cell(row=idx, column=1).font = Font(name="微软雅黑", size=10)
        sheet.cell(row=idx, column=1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[idx].height = 22
    sheet.column_dimensions["A"].width = 96
