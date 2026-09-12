"""冲压加工费导出 — 色板、字体、分隔线。"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_FONT = "微软雅黑"


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def font(*, size: int = 10, bold: bool = False, color: str = "263238") -> Font:
    return Font(name=_FONT, size=size, bold=bold, color=color)


def border(*, right: Side | None = None, left: Side | None = None) -> Border:
    thin = Side(style="thin", color="CFD8DC")
    return Border(
        left=left or thin,
        right=right or thin,
        top=thin,
        bottom=thin,
    )


@dataclass(frozen=True)
class SheetStyle:
    left: Alignment
    center: Alignment
    right: Alignment
    thin: Border
    split_main: Border
    split_detail: Border


def build_sheet_style() -> SheetStyle:
    return SheetStyle(
        left=Alignment(horizontal="left", vertical="center", wrap_text=True),
        center=Alignment(horizontal="center", vertical="center", wrap_text=True),
        right=Alignment(horizontal="right", vertical="center"),
        thin=border(),
        split_main=border(right=Side(style="medium", color="B71C1C")),
        split_detail=border(left=Side(style="medium", color="B71C1C")),
    )
