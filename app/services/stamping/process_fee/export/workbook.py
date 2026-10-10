"""冲压加工费导出 — openpyxl 渲染（不访问 DB）。"""

from __future__ import annotations

from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.common.excel_workbook import excel_write_value
from app.common.timezone import now
from app.models.stamping.process_fee import StampingProcessFee
from app.models.stamping.process_fee_item import StampingProcessFeeItem
from app.services.stamping.process_fee.export.columns import (
    BASIS_COL,
    COL_COUNT,
    DATA_H,
    FEE_HEADER_COLS,
    FILTER_H,
    FINAL_FEE_COL,
    GROUP_H,
    HEADER_H,
    HEADERS,
    LEFT_MAIN_COLS,
    LINE_FEE_COL,
    MAIN_COLS,
    NOTE_H,
    OTHER_FEE_COL,
    PART_NUMBER_COL,
    PROCESS_COUNT_COL,
    RATIO_COL,
    SEQ_COL,
    SHORT_CODE_COL,
    STROKE_PRICE_COL,
    SYSTEM_FEE_COL,
    THICKNESS_COL,
    TITLE_H,
    TONNAGE_COL,
    WIDTHS,
)
from app.services.stamping.process_fee.export.styles import SheetStyle, fill, font
from app.services.stamping.process_fee.export.values import detail_values, main_values, row_height


def write_banner(ws: Worksheet, *, filter_text: str, fee_count: int) -> None:
    indent = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COL_COUNT)
    c1 = ws.cell(row=1, column=1, value="冲压加工费核对表")
    c1.font = font(size=15, bold=True, color="FFFFFF")
    c1.fill = fill("1A237E")
    c1.alignment = indent
    ws.row_dimensions[1].height = TITLE_H

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=COL_COUNT)
    c2 = ws.cell(
        row=2,
        column=1,
        value=(
            f"{filter_text} · 共 {fee_count} 个零件" f"　　导出 {now().strftime('%Y-%m-%d %H:%M')}"
        ),
    )
    c2.font = font(size=9, color="455A64")
    c2.fill = fill("E8EAF6")
    c2.alignment = indent
    ws.row_dimensions[2].height = FILTER_H

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=COL_COUNT)
    c3 = ws.cell(
        row=3,
        column=1,
        value=(
            "说明：左侧主表按零件合并；右侧为工序明细。"
            "单序费用 = 冲次单价 × 费用比例（按工序核对吨位与单价）。"
        ),
    )
    c3.font = font(size=9, color="B71C1C")
    c3.fill = fill("FFF3E0")
    c3.alignment = indent
    ws.row_dimensions[3].height = NOTE_H


def write_headers(ws: Worksheet, *, style: SheetStyle) -> int:
    group_row = 4
    header_row = 5
    white = font(size=9, bold=True, color="FFFFFF")
    header_font = font(bold=True, color="FFFFFF")

    ws.merge_cells(
        start_row=group_row,
        start_column=1,
        end_row=group_row,
        end_column=MAIN_COLS,
    )
    ws.merge_cells(
        start_row=group_row,
        start_column=MAIN_COLS + 1,
        end_row=group_row,
        end_column=COL_COUNT,
    )
    main_group = ws.cell(row=group_row, column=1, value="零件主表")
    main_group.font = white
    main_group.fill = fill("37474F")
    main_group.alignment = style.center
    for col in range(1, MAIN_COLS + 1):
        cell = ws.cell(row=group_row, column=col)
        cell.fill = fill("37474F")
        cell.border = style.thin if col < MAIN_COLS else style.split_main

    detail_group = ws.cell(
        row=group_row,
        column=MAIN_COLS + 1,
        value="工序明细 / 核对（序 · 吨位 · 单价 · 单序费用）",
    )
    detail_group.font = white
    detail_group.fill = fill("B71C1C")
    detail_group.alignment = style.center
    for col in range(MAIN_COLS + 1, COL_COUNT + 1):
        cell = ws.cell(row=group_row, column=col)
        cell.fill = fill("B71C1C")
        cell.border = style.split_detail if col == MAIN_COLS + 1 else style.thin
    ws.row_dimensions[group_row].height = GROUP_H

    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = header_font
        cell.alignment = style.center
        if col <= MAIN_COLS:
            cell.fill = fill("EF6C00") if col in FEE_HEADER_COLS else fill("546E7A")
            cell.border = style.thin if col < MAIN_COLS else style.split_main
        elif col in FEE_HEADER_COLS:
            cell.fill = fill("EF6C00")
            cell.border = style.thin
        else:
            cell.fill = fill("C62828")
            cell.border = style.split_detail if col == MAIN_COLS + 1 else style.thin
    ws.row_dimensions[header_row].height = HEADER_H
    return header_row


def write_data_row(
    ws: Worksheet,
    row_idx: int,
    record: StampingProcessFee,
    *,
    style: SheetStyle,
    seq_no: int,
    item: StampingProcessFeeItem | None,
    write_main: bool,
    line_index: int,
    main_fill: PatternFill,
) -> None:
    detail_fill = fill("FFFFFF") if line_index % 2 == 0 else fill("FFEBEE")

    if write_main:
        for col, value in enumerate(main_values(record, seq_no=seq_no), start=1):
            cell = ws.cell(row=row_idx, column=col, value=excel_write_value(value))
            cell.fill = main_fill
            cell.border = style.thin if col < MAIN_COLS else style.split_main
            _style_main_cell(cell, col, style=style)

    for offset, value in enumerate(detail_values(item)):
        col = MAIN_COLS + 1 + offset
        cell = ws.cell(row=row_idx, column=col, value=excel_write_value(value))
        cell.fill = detail_fill
        cell.border = style.split_detail if col == MAIN_COLS + 1 else style.thin
        _style_detail_cell(cell, col, style=style)


def _style_main_cell(cell: Cell, col: int, *, style: SheetStyle) -> None:
    if col == SEQ_COL:
        cell.font = font(size=11, bold=True, color="37474F")
        cell.alignment = style.center
    elif col == PART_NUMBER_COL:
        cell.font = font(bold=True, color="1565C0")
        cell.alignment = style.left
    elif col == SHORT_CODE_COL:
        cell.font = font(size=11, bold=True, color="1565C0")
        cell.alignment = style.center
    elif col == FINAL_FEE_COL:
        cell.font = font(size=11, bold=True, color="E65100")
        cell.alignment = style.right
        cell.fill = fill("FFF3E0")
    elif col in {PROCESS_COUNT_COL, SYSTEM_FEE_COL, OTHER_FEE_COL}:
        cell.font = font(bold=True) if col == PROCESS_COUNT_COL else font()
        cell.alignment = style.right if col in {SYSTEM_FEE_COL, OTHER_FEE_COL} else style.center
    elif col == THICKNESS_COL:
        cell.font = font()
        cell.alignment = style.right
    else:
        cell.font = font()
        cell.alignment = style.left if col in LEFT_MAIN_COLS else style.center


def _style_detail_cell(cell: Cell, col: int, *, style: SheetStyle) -> None:
    if col == LINE_FEE_COL:
        cell.font = font(size=11, bold=True, color="E65100")
        cell.alignment = style.right
        cell.fill = fill("FFF3E0")
    elif col in {TONNAGE_COL, RATIO_COL, STROKE_PRICE_COL}:
        cell.font = font()
        cell.alignment = style.right
    elif col == BASIS_COL:
        cell.font = font(size=9, color="607D8B")
        cell.alignment = style.left
    else:
        cell.font = font()
        cell.alignment = style.center


def merge_main(
    ws: Worksheet,
    start_row: int,
    end_row: int,
    main_fill: PatternFill,
    *,
    style: SheetStyle,
) -> None:
    for col in range(1, MAIN_COLS + 1):
        ws.merge_cells(
            start_row=start_row,
            start_column=col,
            end_row=end_row,
            end_column=col,
        )
        cell = ws.cell(row=start_row, column=col)
        if col == FINAL_FEE_COL:
            cell.fill = fill("FFF3E0")
        else:
            cell.fill = main_fill
        cell.border = style.thin if col < MAIN_COLS else style.split_main


def fill_sheet(
    ws: Worksheet,
    records: list[StampingProcessFee],
    *,
    style: SheetStyle,
    filter_text: str,
) -> int:
    write_banner(ws, filter_text=filter_text, fee_count=len(records))
    header_row = write_headers(ws, style=style)

    row_idx = header_row + 1
    export_row_count = 0
    for seq_no, record in enumerate(records, start=1):
        items = sorted(
            list(record.process_items or []),
            key=lambda item: (
                item.process_order if item.process_order is not None else 10**9,
                item.id or 0,
            ),
        )
        main_fill = fill("ECEFF1") if seq_no % 2 else fill("F5F7FA")
        if not items:
            write_data_row(
                ws,
                row_idx,
                record,
                style=style,
                seq_no=seq_no,
                item=None,
                write_main=True,
                line_index=0,
                main_fill=main_fill,
            )
            ws.row_dimensions[row_idx].height = DATA_H
            row_idx += 1
            export_row_count += 1
            continue

        group_start = row_idx
        for index, item in enumerate(items):
            write_data_row(
                ws,
                row_idx,
                record,
                style=style,
                seq_no=seq_no,
                item=item,
                write_main=(index == 0),
                line_index=index,
                main_fill=main_fill,
            )
            ws.row_dimensions[row_idx].height = row_height(item)
            row_idx += 1
            export_row_count += 1
        if row_idx - 1 > group_start:
            merge_main(ws, group_start, row_idx - 1, main_fill, style=style)

    last_data = max(header_row, row_idx - 1)
    for idx, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(COL_COUNT)}{last_data}"
    ws.print_title_rows = f"1:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    return export_row_count
