"""冲压模具费导出 — openpyxl 渲染（不访问 DB）。"""

from __future__ import annotations

from decimal import Decimal

from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.common.excel_workbook import excel_write_value
from app.common.timezone import now
from app.models.stamping.die import StampingDieProcess
from app.models.stamping.die_fee import StampingDieFee
from app.models.stamping.die_fee_item import StampingDieFeeItem
from app.services.stamping.die_fee.export.columns import (
    BOLD_MONEY_COLS,
    CHECK_NUM_COLS,
    CHECK_WEIGHT_FORMAT,
    COL_CHECK_START,
    COL_CHECK_WEIGHT,
    COL_COUNT,
    COL_FEE_BASIS,
    COL_LINE_FEE,
    DATA_H,
    FEE_END_COL,
    FEE_HEADER_COLS,
    FEE_MONEY_COLS,
    FILTER_H,
    FINAL_FEE_COL,
    GROUP_H,
    HEADER_H,
    HEADERS,
    LEFT_MAIN_COLS,
    MAIN_COLS,
    MAIN_MONEY_COLS,
    NOTE_H,
    PART_NUMBER_COL,
    SEQ_COL,
    TITLE_H,
    WIDTHS,
)
from app.services.stamping.die_fee.export.styles import SheetStyle, fill, font
from app.services.stamping.die_fee.export.values import (
    _check_weight_formula,
    detail_values,
    main_values,
    row_height,
)


def write_banner(sheet: Worksheet, *, filter_text: str, fee_count: int) -> None:
    indent = Alignment(horizontal="left", vertical="center", indent=1)

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COL_COUNT)
    title = sheet.cell(row=1, column=1, value="冲压模具费核对表")
    title.font = font(size=15, bold=True, color="FFFFFF")
    title.fill = fill("1A237E")
    title.alignment = indent
    sheet.row_dimensions[1].height = TITLE_H

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=COL_COUNT)
    meta = sheet.cell(
        row=2,
        column=1,
        value=(
            f"{filter_text} · 共 {fee_count} 个零件" f"　　导出 {now().strftime('%Y-%m-%d %H:%M')}"
        ),
    )
    meta.font = font(size=9, color="455A64")
    meta.fill = fill("E8EAF6")
    meta.alignment = indent
    sheet.row_dimensions[2].height = FILTER_H

    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=COL_COUNT)
    note = sheet.cell(
        row=3,
        column=1,
        value=(
            "关注点是模具费：本序模具费 = 模具重量 × 单价 × 分摊比例 / 分摊万台；"
            "最终模具费 = 系统模具费 + 调整金额。"
            "最右侧「布点核对」仅辅助检查导入外形是否正确；"
            "核对重量列为公式（长×宽×高×掏空系数×7.85÷10^9），四位小数，不参与计费。"
        ),
    )
    note.font = font(size=9, color="B71C1C")
    note.fill = fill("FFF3E0")
    note.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    sheet.row_dimensions[3].height = NOTE_H


def write_headers(sheet: Worksheet, *, style: SheetStyle) -> int:
    group_row = 4
    header_row = 5
    white = font(size=9, bold=True, color="FFFFFF")
    header_font = font(bold=True, color="FFFFFF")

    sheet.merge_cells(
        start_row=group_row,
        start_column=1,
        end_row=group_row,
        end_column=MAIN_COLS,
    )
    sheet.merge_cells(
        start_row=group_row,
        start_column=MAIN_COLS + 1,
        end_row=group_row,
        end_column=FEE_END_COL,
    )
    sheet.merge_cells(
        start_row=group_row,
        start_column=COL_CHECK_START,
        end_row=group_row,
        end_column=COL_COUNT,
    )
    main_group = sheet.cell(row=group_row, column=1, value="零件主表")
    main_group.font = white
    main_group.fill = fill("37474F")
    main_group.alignment = style.center
    for col in range(1, MAIN_COLS + 1):
        cell = sheet.cell(row=group_row, column=col)
        cell.fill = fill("37474F")
        cell.border = style.thin if col < MAIN_COLS else style.split_main

    detail_group = sheet.cell(
        row=group_row,
        column=MAIN_COLS + 1,
        value="各序模具费（重量 · 比例 · 本序模具费）",
    )
    detail_group.font = white
    detail_group.fill = fill("B71C1C")
    detail_group.alignment = style.center
    for col in range(MAIN_COLS + 1, FEE_END_COL + 1):
        cell = sheet.cell(row=group_row, column=col)
        cell.fill = fill("B71C1C")
        cell.border = _detail_border(col, style)

    check_group = sheet.cell(
        row=group_row,
        column=COL_CHECK_START,
        value="布点核对（辅助）· 检查导入外形是否正确，不参与计费",
    )
    check_group.font = white
    check_group.fill = fill("607D8B")
    check_group.alignment = style.center
    for col in range(COL_CHECK_START, COL_COUNT + 1):
        cell = sheet.cell(row=group_row, column=col)
        cell.fill = fill("607D8B")
        cell.border = style.split_check if col == COL_CHECK_START else style.thin
    sheet.row_dimensions[group_row].height = GROUP_H

    for col, title in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=col, value=title)
        cell.font = header_font
        cell.alignment = style.center
        if col <= MAIN_COLS:
            cell.fill = fill("EF6C00") if col in FEE_HEADER_COLS else fill("546E7A")
            cell.border = style.thin if col < MAIN_COLS else style.split_main
        elif col <= FEE_END_COL:
            cell.fill = fill("EF6C00") if col in FEE_HEADER_COLS else fill("C62828")
            cell.border = _detail_border(col, style)
        else:
            cell.fill = fill("78909C")
            cell.font = font(size=9, bold=True, color="FFFFFF")
            cell.border = style.split_check if col == COL_CHECK_START else style.thin
    sheet.row_dimensions[header_row].height = HEADER_H
    return header_row


def _detail_border(col: int, style: SheetStyle) -> Border:
    if col == MAIN_COLS + 1:
        return style.split_detail
    if col == FEE_END_COL:
        return style.split_fee_end
    if col == COL_CHECK_START:
        return style.split_check
    return style.thin


def write_data_row(
    sheet: Worksheet,
    row_idx: int,
    record: StampingDieFee,
    *,
    style: SheetStyle,
    seq_no: int,
    item: StampingDieFeeItem | None,
    process: StampingDieProcess | None,
    write_main: bool,
    line_index: int,
    main_fill: PatternFill,
    density: Decimal,
) -> None:
    detail_fill = fill("FFFFFF") if line_index % 2 == 0 else fill("FFEBEE")
    check_fill = fill("FAFAFA") if line_index % 2 == 0 else fill("ECEFF1")

    if write_main:
        for col, value in enumerate(main_values(record, seq_no=seq_no), start=1):
            cell = sheet.cell(row=row_idx, column=col, value=excel_write_value(value))
            cell.fill = main_fill
            cell.border = style.thin if col < MAIN_COLS else style.split_main
            _style_main_cell(cell, col, style=style)

    for offset, value in enumerate(detail_values(item, process)):
        col = MAIN_COLS + 1 + offset
        cell = sheet.cell(row=row_idx, column=col, value=excel_write_value(value))
        is_check = col >= COL_CHECK_START
        cell.fill = check_fill if is_check else detail_fill
        cell.border = _detail_border(col, style)
        _style_detail_cell(cell, col, style=style, is_check=is_check)

    weight_cell = sheet.cell(row=row_idx, column=COL_CHECK_WEIGHT)
    weight_cell.fill = check_fill
    weight_cell.border = _detail_border(COL_CHECK_WEIGHT, style)
    weight_cell.font = font(size=9, color="78909C")
    weight_cell.alignment = style.right
    weight_cell.number_format = CHECK_WEIGHT_FORMAT
    weight_cell.value = _check_weight_formula(row_idx, density) if item is not None else ""


def _style_main_cell(cell: Cell, col: int, *, style: SheetStyle) -> None:
    if col == SEQ_COL:
        cell.font = font(size=11, bold=True, color="37474F")
        cell.alignment = style.center
    elif col == PART_NUMBER_COL:
        cell.font = font(bold=True, color="1565C0")
        cell.alignment = style.left
    elif col == FINAL_FEE_COL:
        cell.font = font(size=11, bold=True, color="E65100")
        cell.alignment = style.right
        cell.fill = fill("FFF3E0")
    elif col in MAIN_MONEY_COLS:
        cell.font = font(bold=True) if col in BOLD_MONEY_COLS else font()
        cell.alignment = style.right
    else:
        cell.font = font()
        cell.alignment = style.left if col in LEFT_MAIN_COLS else style.center


def _style_detail_cell(
    cell: Cell,
    col: int,
    *,
    style: SheetStyle,
    is_check: bool,
) -> None:
    if col == COL_LINE_FEE:
        cell.font = font(size=11, bold=True, color="E65100")
        cell.alignment = style.right
        cell.fill = fill("FFF3E0")
    elif col in FEE_MONEY_COLS:
        cell.font = font()
        cell.alignment = style.right
    elif col == COL_FEE_BASIS:
        cell.font = font(size=9, color="607D8B")
        cell.alignment = style.left
    elif col in CHECK_NUM_COLS:
        cell.font = font(size=9, color="78909C")
        cell.alignment = style.right
    else:
        cell.font = font(size=9, color="78909C") if is_check else font()
        cell.alignment = style.center


def merge_main(
    sheet: Worksheet,
    start_row: int,
    end_row: int,
    main_fill: PatternFill,
    *,
    style: SheetStyle,
) -> None:
    for col in range(1, MAIN_COLS + 1):
        sheet.merge_cells(
            start_row=start_row,
            start_column=col,
            end_row=end_row,
            end_column=col,
        )
        cell = sheet.cell(row=start_row, column=col)
        if col == FINAL_FEE_COL:
            cell.fill = fill("FFF3E0")
        else:
            cell.fill = main_fill
        cell.border = style.thin if col < MAIN_COLS else style.split_main


def fill_sheet(
    sheet: Worksheet,
    records: list[StampingDieFee],
    *,
    style: SheetStyle,
    process_by_id: dict[int, StampingDieProcess],
    filter_text: str,
    density: Decimal,
) -> int:
    write_banner(sheet, filter_text=filter_text, fee_count=len(records))
    header_row = write_headers(sheet, style=style)

    row_idx = header_row + 1
    export_row_count = 0
    for seq_no, record in enumerate(records, start=1):
        items = sorted(
            list(record.items or []),
            key=lambda item: (
                item.process_order if item.process_order is not None else 10**9,
                item.id or 0,
            ),
        )
        main_fill = fill("ECEFF1") if seq_no % 2 else fill("F5F7FA")
        if not items:
            write_data_row(
                sheet,
                row_idx,
                record,
                style=style,
                seq_no=seq_no,
                item=None,
                process=None,
                write_main=True,
                line_index=0,
                main_fill=main_fill,
                density=density,
            )
            sheet.row_dimensions[row_idx].height = DATA_H
            row_idx += 1
            export_row_count += 1
            continue

        group_start = row_idx
        for index, item in enumerate(items):
            process = (
                process_by_id.get(item.source_die_process_id)
                if item.source_die_process_id
                else None
            )
            write_data_row(
                sheet,
                row_idx,
                record,
                style=style,
                seq_no=seq_no,
                item=item,
                process=process,
                write_main=(index == 0),
                line_index=index,
                main_fill=main_fill,
                density=density,
            )
            sheet.row_dimensions[row_idx].height = row_height(item)
            row_idx += 1
            export_row_count += 1
        if row_idx - 1 > group_start:
            merge_main(sheet, group_start, row_idx - 1, main_fill, style=style)

    last_data = max(header_row, row_idx - 1)
    for idx, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width

    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(COL_COUNT)}{last_data}"
    sheet.print_title_rows = f"1:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = 0.3
    sheet.page_margins.right = 0.3
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.35
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f"A1:{get_column_letter(COL_COUNT)}{last_data}"
    return export_row_count
