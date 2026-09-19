"""冲压模具费导出 — 单元格取值、核对重量公式、筛选摘要。"""

from __future__ import annotations

from decimal import Decimal

from openpyxl.utils import get_column_letter

from app.common.audit import audit_status_label
from app.common.money import Q0, Q2, Q4, excel_money_cell
from app.common.timezone import now
from app.models.stamping.die import StampingDieProcess
from app.models.stamping.die_fee import StampingDieFee
from app.models.stamping.die_fee_item import StampingDieFeeItem
from app.services.stamping.die.weight_check import DIE_STEEL_DENSITY
from app.services.stamping.die_fee.export.columns import (
    COL_HEIGHT,
    COL_HOLLOW,
    COL_LENGTH,
    COL_WIDTH,
    DATA_H,
)


def money_cell(value: object) -> Decimal | str:
    return excel_money_cell(value, Q4, allow_negative=True)


def money_cell_q2(value: object) -> Decimal | str:
    return excel_money_cell(value, Q2, allow_negative=True)


def money_cell_whole(value: object) -> Decimal | str:
    return excel_money_cell(value, Q0, allow_negative=False)


def _check_weight_formula(row_idx: int) -> str:
    length_ref = f"{get_column_letter(COL_LENGTH)}{row_idx}"
    width_ref = f"{get_column_letter(COL_WIDTH)}{row_idx}"
    height_ref = f"{get_column_letter(COL_HEIGHT)}{row_idx}"
    hollow_ref = f"{get_column_letter(COL_HOLLOW)}{row_idx}"
    density = str(DIE_STEEL_DENSITY)
    return (
        f"=IF(COUNT({length_ref},{width_ref},{height_ref},{hollow_ref})=4,"
        f'{length_ref}*{width_ref}*{height_ref}*{hollow_ref}*{density}/1000000000,"")'
    )


def price_kind_label(kind: str | None) -> str:
    return "级进模" if kind == "progressive" else "钢板模"


def filter_text(
    *,
    year: int | None,
    location: str | None,
    vehicle_model: str | None,
    stamping_supplier: str | None,
    version: str | None,
    part_number: str | None,
    audit_status: str | None,
    only_zero: bool,
) -> str:
    parts: list[str] = []
    if year is not None:
        parts.append(f"年度={year}")
    if location:
        parts.append(f"基地={location}")
    if vehicle_model:
        parts.append(f"项目={vehicle_model}")
    if stamping_supplier:
        parts.append(f"供应商={stamping_supplier}")
    if version:
        parts.append(f"版本={version}")
    if part_number:
        parts.append(f"零件号含={part_number}")
    if audit_status:
        parts.append(f"审核={audit_status_label(audit_status)}")
    if only_zero:
        parts.append("仅最终为零")
    return " · ".join(parts) if parts else "全部筛选"


def build_filename(
    *,
    year: int | None,
    location: str | None,
    vehicle_model: str | None,
    stamping_supplier: str | None,
    version: str | None,
) -> str:
    stamp = now().strftime("%Y%m%d_%H%M%S")
    parts = ["冲压模具费"]
    if year is not None:
        parts.append(str(year))
    if location:
        parts.append(location)
    if vehicle_model:
        parts.append(vehicle_model)
    if stamping_supplier:
        parts.append(stamping_supplier)
    if version:
        parts.append(version)
    return "_".join(parts) + f"_{stamp}.xlsx"


def row_height(item: StampingDieFeeItem) -> float:
    basis = (item.basis_text or "").strip()
    extra = 0.0
    if len(basis) > 40:
        extra = min(36.0, (len(basis) // 40) * 8.0)
    return DATA_H + extra


def main_values(record: StampingDieFee, *, seq_no: int) -> list[object]:
    return [
        seq_no,
        record.year,
        record.location,
        record.vehicle_model,
        record.stamping_supplier or "",
        record.version or "",
        record.part_number,
        record.part_name or "",
        record.production_mode or "",
        price_kind_label(record.price_kind),
        record.material or "",
        money_cell(record.thickness),
        money_cell(record.unit_price),
        money_cell_whole(record.allocation_qty),
        money_cell(record.die_fee),
        money_cell(record.adjustment_amount),
        money_cell(record.final_die_fee),
        audit_status_label(record.audit_status),
    ]


def detail_values(
    item: StampingDieFeeItem | None,
    process: StampingDieProcess | None,
) -> list[object]:
    return [
        (item.process_seq or "") if item else "",
        (item.process_code or "") if item else "",
        money_cell(item.estimated_weight_t) if item else "",
        money_cell(item.alloc_ratio) if item else "",
        money_cell(item.line_die_fee) if item else "",
        (item.basis_text or "").strip() if item else "",
        (process.die_number or "").strip() if process else "",
        money_cell_q2(process.die_length) if process else "",
        money_cell_q2(process.die_width) if process else "",
        money_cell_q2(process.die_height) if process else "",
        money_cell(process.hollow_factor) if process else "",
    ]
