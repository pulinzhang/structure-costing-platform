"""冲压模具费 Excel 导出（主子表核对风排版，对齐冲压加工费）。

左侧零件主表纵向合并；中段各序模具费；右侧布点外形仅作导入核对。
"""

from __future__ import annotations

import logging
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.common.audit import audit_status_label
from app.common.excel_workbook import resolve_worksheet
from app.common.money import Q0, Q2, Q4, excel_money_cell
from app.common.timezone import now
from app.common.vehicle_model import normalize_vehicle_model
from app.exceptions.base import BusinessException
from app.models.stamping.die import StampingDieProcess
from app.models.stamping.die_fee import StampingDieFee
from app.models.stamping.die_fee_item import StampingDieFeeItem
from app.repositories.stamping.die_fee_repository import StampingDieFeeRepository
from app.repositories.stamping.stamping_die_repository import StampingDieProcessRepository
from app.services.stamping.die.weight_check import DIE_STEEL_DENSITY

logger = logging.getLogger("service.stamping_die_fee_export")

_FONT = "微软雅黑"

MAIN_COLS = 18
FEE_DETAIL_COLS = 6
FEE_END_COL = MAIN_COLS + FEE_DETAIL_COLS
COL_COUNT = 30
HEADERS = [
    "序号",
    "年度",
    "基地",
    "项目",
    "供应商",
    "版本",
    "零件号",
    "零件名称",
    "生产方式",
    "取价档",
    "材质",
    "料厚",
    "单价(万元/吨)",
    "分摊(万台)",
    "系统模具费",
    "调整金额",
    "最终模具费",
    "审核状态",
    "工序序号",
    "工序代码",
    "模具重量(T)",
    "分摊比例",
    "本序模具费",
    "费用计算依据",
    "模具编号",
    "长(mm)",
    "宽(mm)",
    "高(mm)",
    "掏空系数",
    "核对重量(T)",
]
WIDTHS = [
    6,
    8,
    10,
    12,
    18,
    10,
    16,
    18,
    12,
    10,
    10,
    8,
    14,
    12,
    14,
    12,
    14,
    10,
    10,
    10,
    12,
    10,
    14,
    40,
    16,
    10,
    10,
    10,
    10,
    14,
]

TITLE_H = 28
FILTER_H = 20
NOTE_H = 36
GROUP_H = 18
HEADER_H = 22
DATA_H = 24

_FEE_HEADER_COLS = {13, 14, 15, 16, 17, 23}
_MAIN_MONEY_COLS = {12, 13, 14, 15, 16, 17}
_COL_LINE_FEE = 23
_COL_FEE_BASIS = 24
_COL_CHECK_START = FEE_END_COL + 1
_COL_LENGTH = 26
_COL_WIDTH = 27
_COL_HEIGHT = 28
_COL_HOLLOW = 29
_COL_CHECK_WEIGHT = 30
_FEE_MONEY_COLS = {21, 22, 23}
_CHECK_NUM_COLS = {26, 27, 28, 29, 30}
_CHECK_WEIGHT_FORMAT = "0.0000"


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(*, size: int = 10, bold: bool = False, color: str = "263238") -> Font:
    return Font(name=_FONT, size=size, bold=bold, color=color)


def _border(*, right: Side | None = None, left: Side | None = None) -> Border:
    thin = Side(style="thin", color="CFD8DC")
    return Border(
        left=left or thin,
        right=right or thin,
        top=thin,
        bottom=thin,
    )


def _num(value: Any) -> Any:
    return excel_money_cell(value, Q4, allow_negative=True)


def _num_q2(value: Any) -> Any:
    return excel_money_cell(value, Q2, allow_negative=True)


def _num_whole(value: Any) -> Any:
    return excel_money_cell(value, Q0, allow_negative=False)


def _check_weight_formula(row_idx: int) -> str:
    length_ref = f"{get_column_letter(_COL_LENGTH)}{row_idx}"
    width_ref = f"{get_column_letter(_COL_WIDTH)}{row_idx}"
    height_ref = f"{get_column_letter(_COL_HEIGHT)}{row_idx}"
    hollow_ref = f"{get_column_letter(_COL_HOLLOW)}{row_idx}"
    density = str(DIE_STEEL_DENSITY)
    return (
        f"=IF(COUNT({length_ref},{width_ref},{height_ref},{hollow_ref})=4,"
        f'{length_ref}*{width_ref}*{height_ref}*{hollow_ref}*{density}/1000000000,"")'
    )


def _price_kind_label(kind: str | None) -> str:
    return "级进模" if kind == "progressive" else "钢板模"


class StampingDieFeeExportService:
    """冲压模具费 Excel 导出（主表合并 + 各序明细分行）。"""

    def __init__(self) -> None:
        self.fee_repo = StampingDieFeeRepository()
        self.process_repo = StampingDieProcessRepository()
        self._left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        self._center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self._right = Alignment(horizontal="right", vertical="center")
        self._thin = _border()
        self._split_main = _border(right=Side(style="medium", color="B71C1C"))
        self._split_detail = _border(left=Side(style="medium", color="B71C1C"))
        self._split_fee_end = _border(right=Side(style="medium", color="78909C"))
        self._split_check = _border(left=Side(style="medium", color="78909C"))

    def export(
        self,
        db: Session,
        *,
        year: int,
        location: str,
        vehicle_model: str,
        part_number: str | None,
        audit_status: str | None,
        only_zero: bool,
        stamping_supplier: str | None = None,
        version: str | None = None,
    ) -> tuple[bytes, str]:
        loc = location.strip()
        vehicle = normalize_vehicle_model(vehicle_model)
        if not loc or not vehicle:
            raise BusinessException("导出必须指定基地和项目")
        part = (part_number or "").strip() or None
        supplier = (stamping_supplier or "").strip() or None
        version_value = (version or "").strip() or None

        workbook = Workbook()
        _title, fee_count = self._write_to_workbook(
            workbook,
            db,
            year=year,
            location=loc,
            vehicle_model=vehicle,
            part_number=part,
            audit_status=audit_status,
            only_zero=only_zero,
            stamping_supplier=supplier,
            version=version_value,
        )

        buffer = BytesIO()
        workbook.save(buffer)
        stamp = now().strftime("%Y%m%d_%H%M%S")
        extra_bits = ""
        if supplier:
            extra_bits += f"_{supplier}"
        if version_value:
            extra_bits += f"_{version_value}"
        filename = f"冲压模具费_{year}_{loc}_{vehicle}{extra_bits}_{stamp}.xlsx"
        logger.info("导出冲压模具费 %s 条主表 -> %s", fee_count, filename)
        return buffer.getvalue(), filename

    def _write_to_workbook(
        self,
        workbook: Workbook,
        db: Session,
        *,
        year: int,
        location: str,
        vehicle_model: str,
        part_number: str | None,
        audit_status: str | None,
        only_zero: bool,
        stamping_supplier: str | None = None,
        version: str | None = None,
        sheet_title: str = "冲压模具费",
    ) -> tuple[str, int]:
        records = self.fee_repo.list_for_export(
            db,
            year=year,
            location=location,
            vehicle_model=vehicle_model,
            part_number=part_number,
            audit_status=audit_status,
            only_zero=only_zero,
            stamping_supplier=stamping_supplier,
            version=version,
        )
        process_by_id = self._load_source_processes(db, records)
        sheet = resolve_worksheet(workbook, sheet_title)
        filter_text = self._filter_text(
            year=year,
            location=location,
            vehicle_model=vehicle_model,
            stamping_supplier=stamping_supplier,
            version=version,
            part_number=part_number,
            audit_status=audit_status,
            only_zero=only_zero,
        )
        self._write_banner(sheet, filter_text=filter_text, fee_count=len(records))
        header_row = self._write_headers(sheet)

        row_idx = header_row + 1
        export_row_count = 0
        for seq_no, record in enumerate(records, start=1):
            items = sorted(
                list(record.items or []),
                key=lambda item: (
                    item.process_order if item.process_order is not None else 10**9,
                    item.id or 0,
                ),
            )
            main_fill = _fill("ECEFF1") if seq_no % 2 else _fill("F5F7FA")
            if not items:
                self._write_data_row(
                    sheet,
                    row_idx,
                    record,
                    seq_no=seq_no,
                    item=None,
                    process=None,
                    write_main=True,
                    line_index=0,
                    main_fill=main_fill,
                )
                sheet.row_dimensions[row_idx].height = DATA_H
                row_idx += 1
                export_row_count += 1
                continue

            group_start = row_idx
            for index, item in enumerate(items):
                process = (
                    process_by_id.get(item.source_die_process_id)
                    if item.source_die_process_id
                    else None
                )
                self._write_data_row(
                    sheet,
                    row_idx,
                    record,
                    seq_no=seq_no,
                    item=item,
                    process=process,
                    write_main=(index == 0),
                    line_index=index,
                    main_fill=main_fill,
                )
                sheet.row_dimensions[row_idx].height = self._row_height(item)
                row_idx += 1
                export_row_count += 1
            if row_idx - 1 > group_start:
                self._merge_main(sheet, group_start, row_idx - 1, main_fill)

        last_data = max(header_row, row_idx - 1)
        for idx, width in enumerate(WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(idx)].width = width

        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(COL_COUNT)}{last_data}"
        sheet.print_title_rows = f"1:{header_row}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.3
        sheet.page_margins.right = 0.3
        sheet.page_margins.top = 0.35
        sheet.page_margins.bottom = 0.35
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 90
        sheet.print_options.horizontalCentered = True
        sheet.print_area = f"A1:{get_column_letter(COL_COUNT)}{last_data}"

        logger.info(
            "写入冲压模具费 sheet=%s 主表=%s 明细行=%s",
            sheet.title,
            len(records),
            export_row_count,
        )
        return sheet.title, len(records)

    def _write_banner(self, sheet: Worksheet, *, filter_text: str, fee_count: int) -> None:
        indent = Alignment(horizontal="left", vertical="center", indent=1)

        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COL_COUNT)
        title = sheet.cell(row=1, column=1, value="冲压模具费核对表")
        title.font = _font(size=15, bold=True, color="FFFFFF")
        title.fill = _fill("1A237E")
        title.alignment = indent
        sheet.row_dimensions[1].height = TITLE_H

        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=COL_COUNT)
        meta = sheet.cell(
            row=2,
            column=1,
            value=(
                f"{filter_text} · 共 {fee_count} 个零件"
                f"　　导出 {now().strftime('%Y-%m-%d %H:%M')}"
            ),
        )
        meta.font = _font(size=9, color="455A64")
        meta.fill = _fill("E8EAF6")
        meta.alignment = indent
        sheet.row_dimensions[2].height = FILTER_H

        sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=COL_COUNT)
        note = sheet.cell(
            row=3,
            column=1,
            value=(
                "关注点是模具费：本序模具费 = 模具重量 × 单价 × 分摊比例 / 分摊万台；"
                "最终模具费 = 系统模具费 + 调整金额。"
                "最右侧「布点核对」仅辅助检查导入外形是否正确；"
                "核对重量列为公式（长×宽×高×掏空系数×7.85÷10^9），四位小数，不参与计费。"
            ),
        )
        note.font = _font(size=9, color="B71C1C")
        note.fill = _fill("FFF3E0")
        note.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        sheet.row_dimensions[3].height = NOTE_H

    def _write_headers(self, sheet: Worksheet) -> int:
        group_row = 4
        header_row = 5
        white = _font(size=9, bold=True, color="FFFFFF")
        header_font = _font(bold=True, color="FFFFFF")

        sheet.merge_cells(
            start_row=group_row,
            start_column=1,
            end_row=group_row,
            end_column=MAIN_COLS,
        )
        sheet.merge_cells(
            start_row=group_row,
            start_column=MAIN_COLS + 1,
            end_row=group_row,
            end_column=FEE_END_COL,
        )
        sheet.merge_cells(
            start_row=group_row,
            start_column=_COL_CHECK_START,
            end_row=group_row,
            end_column=COL_COUNT,
        )
        main_group = sheet.cell(row=group_row, column=1, value="零件主表")
        main_group.font = white
        main_group.fill = _fill("37474F")
        main_group.alignment = self._center
        for col in range(1, MAIN_COLS + 1):
            cell = sheet.cell(row=group_row, column=col)
            cell.fill = _fill("37474F")
            cell.border = self._thin if col < MAIN_COLS else self._split_main

        detail_group = sheet.cell(
            row=group_row,
            column=MAIN_COLS + 1,
            value="各序模具费（重量 · 比例 · 本序模具费）",
        )
        detail_group.font = white
        detail_group.fill = _fill("B71C1C")
        detail_group.alignment = self._center
        for col in range(MAIN_COLS + 1, FEE_END_COL + 1):
            cell = sheet.cell(row=group_row, column=col)
            cell.fill = _fill("B71C1C")
            if col == MAIN_COLS + 1:
                cell.border = self._split_detail
            elif col == FEE_END_COL:
                cell.border = self._split_fee_end
            else:
                cell.border = self._thin

        check_group = sheet.cell(
            row=group_row,
            column=_COL_CHECK_START,
            value="布点核对（辅助）· 检查导入外形是否正确，不参与计费",
        )
        check_group.font = white
        check_group.fill = _fill("607D8B")
        check_group.alignment = self._center
        for col in range(_COL_CHECK_START, COL_COUNT + 1):
            cell = sheet.cell(row=group_row, column=col)
            cell.fill = _fill("607D8B")
            cell.border = self._split_check if col == _COL_CHECK_START else self._thin
        sheet.row_dimensions[group_row].height = GROUP_H

        for col, title in enumerate(HEADERS, start=1):
            cell = sheet.cell(row=header_row, column=col, value=title)
            cell.font = header_font
            cell.alignment = self._center
            if col <= MAIN_COLS:
                cell.fill = _fill("EF6C00") if col in _FEE_HEADER_COLS else _fill("546E7A")
                cell.border = self._thin if col < MAIN_COLS else self._split_main
            elif col <= FEE_END_COL:
                cell.fill = _fill("EF6C00") if col in _FEE_HEADER_COLS else _fill("C62828")
                if col == MAIN_COLS + 1:
                    cell.border = self._split_detail
                elif col == FEE_END_COL:
                    cell.border = self._split_fee_end
                else:
                    cell.border = self._thin
            else:
                cell.fill = _fill("78909C")
                cell.font = _font(size=9, bold=True, color="FFFFFF")
                cell.border = self._split_check if col == _COL_CHECK_START else self._thin
        sheet.row_dimensions[header_row].height = HEADER_H
        return header_row

    def _load_source_processes(
        self,
        db: Session,
        records: list[StampingDieFee],
    ) -> dict[int, StampingDieProcess]:
        process_ids = [
            item.source_die_process_id
            for record in records
            for item in (record.items or [])
            if item.source_die_process_id
        ]
        processes = self.process_repo.list_by_ids(db, process_ids)
        return {process.id: process for process in processes}

    def _write_data_row(
        self,
        sheet: Worksheet,
        row_idx: int,
        record: StampingDieFee,
        *,
        seq_no: int,
        item: StampingDieFeeItem | None,
        process: StampingDieProcess | None,
        write_main: bool,
        line_index: int,
        main_fill: PatternFill,
    ) -> None:
        detail_fill = _fill("FFFFFF") if line_index % 2 == 0 else _fill("FFEBEE")
        check_fill = _fill("FAFAFA") if line_index % 2 == 0 else _fill("ECEFF1")

        if write_main:
            main_values: list[Any] = [
                seq_no,
                record.year,
                record.location,
                record.vehicle_model,
                record.stamping_supplier or "",
                record.version or "",
                record.part_number,
                record.part_name or "",
                record.production_mode or "",
                _price_kind_label(record.price_kind),
                record.material or "",
                _num(record.thickness),
                _num(record.unit_price),
                _num_whole(record.allocation_qty),
                _num(record.die_fee),
                _num(record.adjustment_amount),
                _num(record.final_die_fee),
                audit_status_label(record.audit_status),
            ]
            for col, value in enumerate(main_values, start=1):
                cell = sheet.cell(row=row_idx, column=col, value="" if value is None else value)
                cell.fill = main_fill
                cell.border = self._thin if col < MAIN_COLS else self._split_main
                if col == 1:
                    cell.font = _font(size=11, bold=True, color="37474F")
                    cell.alignment = self._center
                elif col == 7:
                    cell.font = _font(bold=True, color="1565C0")
                    cell.alignment = self._left
                elif col == 17:
                    cell.font = _font(size=11, bold=True, color="E65100")
                    cell.alignment = self._right
                    cell.fill = _fill("FFF3E0")
                elif col in _MAIN_MONEY_COLS:
                    cell.font = _font(bold=True) if col in {15, 16} else _font()
                    cell.alignment = self._right
                else:
                    cell.font = _font()
                    cell.alignment = (
                        self._left if col in {3, 4, 5, 6, 8, 9, 10, 11} else self._center
                    )

        detail_values: list[Any] = [
            (item.process_seq or "") if item else "",
            (item.process_code or "") if item else "",
            _num(item.estimated_weight_t) if item else "",
            _num(item.alloc_ratio) if item else "",
            _num(item.line_die_fee) if item else "",
            (item.basis_text or "").strip() if item else "",
            (process.die_number or "").strip() if process else "",
            _num_q2(process.die_length) if process else "",
            _num_q2(process.die_width) if process else "",
            _num_q2(process.die_height) if process else "",
            _num(process.hollow_factor) if process else "",
        ]
        for offset, value in enumerate(detail_values):
            col = MAIN_COLS + 1 + offset
            cell = sheet.cell(row=row_idx, column=col, value="" if value is None else value)
            is_check = col >= _COL_CHECK_START
            cell.fill = check_fill if is_check else detail_fill
            cell.border = self._detail_border(col)
            if col == _COL_LINE_FEE:
                cell.font = _font(size=11, bold=True, color="E65100")
                cell.alignment = self._right
                cell.fill = _fill("FFF3E0")
            elif col in _FEE_MONEY_COLS:
                cell.font = _font()
                cell.alignment = self._right
            elif col == _COL_FEE_BASIS:
                cell.font = _font(size=9, color="607D8B")
                cell.alignment = self._left
            elif col in _CHECK_NUM_COLS:
                cell.font = _font(size=9, color="78909C")
                cell.alignment = self._right
            else:
                cell.font = _font(size=9, color="78909C") if is_check else _font()
                cell.alignment = self._center

        weight_cell = sheet.cell(row=row_idx, column=_COL_CHECK_WEIGHT)
        weight_cell.fill = check_fill
        weight_cell.border = self._detail_border(_COL_CHECK_WEIGHT)
        weight_cell.font = _font(size=9, color="78909C")
        weight_cell.alignment = self._right
        weight_cell.number_format = _CHECK_WEIGHT_FORMAT
        if item is not None:
            weight_cell.value = _check_weight_formula(row_idx)
        else:
            weight_cell.value = ""

    def _detail_border(self, col: int) -> Border:
        if col == MAIN_COLS + 1:
            return self._split_detail
        if col == FEE_END_COL:
            return self._split_fee_end
        if col == _COL_CHECK_START:
            return self._split_check
        return self._thin

    def _merge_main(
        self,
        sheet: Worksheet,
        start_row: int,
        end_row: int,
        main_fill: PatternFill,
    ) -> None:
        for col in range(1, MAIN_COLS + 1):
            sheet.merge_cells(
                start_row=start_row,
                start_column=col,
                end_row=end_row,
                end_column=col,
            )
            cell = sheet.cell(row=start_row, column=col)
            if col == 17:
                cell.fill = _fill("FFF3E0")
            else:
                cell.fill = main_fill
            cell.border = self._thin if col < MAIN_COLS else self._split_main

    @staticmethod
    def _row_height(item: StampingDieFeeItem) -> float:
        basis = (item.basis_text or "").strip()
        extra = 0.0
        if len(basis) > 40:
            extra = min(36.0, (len(basis) // 40) * 8.0)
        return DATA_H + extra

    @staticmethod
    def _filter_text(
        *,
        year: int,
        location: str,
        vehicle_model: str,
        stamping_supplier: str | None,
        version: str | None,
        part_number: str | None,
        audit_status: str | None,
        only_zero: bool,
    ) -> str:
        parts = [
            f"年度={year}",
            f"基地={location}",
            f"项目={vehicle_model}",
        ]
        if stamping_supplier:
            parts.append(f"供应商={stamping_supplier}")
        if version:
            parts.append(f"版本={version}")
        if part_number:
            parts.append(f"零件号含={part_number}")
        if audit_status:
            parts.append(f"审核={audit_status_label(audit_status)}")
        if only_zero:
            parts.append("仅最终为零")
        return " · ".join(parts)
