"""冲压加工费 Excel 导出（主子表核对风排版）。

左侧主表按零件纵向合并；右侧工序明细一行一序：
工序序号 / 工序代码 / 吨位 / 冲次单价 / 单序费用。
"""

from __future__ import annotations

import logging
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.common.audit import audit_status_label
from app.common.excel_workbook import resolve_worksheet
from app.common.export_filters import require_at_least_one_export_filter
from app.common.money import Q3, Q4, excel_money_cell
from app.common.timezone import now
from app.models.stamping.process_fee import StampingProcessFee
from app.models.stamping.process_fee_item import StampingProcessFeeItem
from app.repositories.stamping.stamping_process_fee_repository import (
    StampingProcessFeeRepository,
)


logger = logging.getLogger("service.stamping_process_fee_export")

_FONT = "微软雅黑"

_TONNAGE_SOURCE_LABEL = {
    "self": "自有",
    "inherited": "继承",
    "swing_shear": "摆剪",
}

# 主表列（纵向合并）+ 工序明细分隔
MAIN_COLS = 14
COL_COUNT = 22
HEADERS = [
    "序号",
    "年度",
    "基地",
    "冲压工厂",
    "车型",
    "版本",
    "零件号",
    "零件名称",
    "材质",
    "料厚",
    "工序数",
    "系统计算加工费",
    "其它费用",
    "确认单价",
    "工序序号",
    "工序代码",
    "吨位(T)",
    "吨位来源",
    "费用比例",
    "冲次单价",
    "单序费用",
    "计算依据",
]
WIDTHS = [
    6,
    8,
    10,
    12,
    10,
    10,
    16,
    18,
    10,
    8,
    8,
    14,
    12,
    14,
    10,
    10,
    10,
    10,
    10,
    10,
    12,
    40,
]

TITLE_H = 28
FILTER_H = 20
NOTE_H = 20
GROUP_H = 18
HEADER_H = 22
DATA_H = 24


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(*, size: int = 10, bold: bool = False, color: str = "263238") -> Font:
    return Font(name=_FONT, size=size, bold=bold, color=color)


def _border(*, right: Side | None = None, left: Side | None = None) -> Border:
    thin = Side(style="thin", color="CFD8DC")
    return Border(
        left=left or thin,
        right=right or thin,
        top=thin,
        bottom=thin,
    )


def _num(value: Any, quant: Any = Q3) -> Any:
    return excel_money_cell(value, quant)


class StampingProcessFeeExportService:
    """冲压加工费 Excel 导出（主表合并 + 工序明细分行）。"""

    def __init__(self) -> None:
        self.repository = StampingProcessFeeRepository()
        self._left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        self._center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self._right = Alignment(horizontal="right", vertical="center")
        self._thin = _border()
        self._split_main = _border(right=Side(style="medium", color="B71C1C"))
        self._split_detail = _border(left=Side(style="medium", color="B71C1C"))

    def export(
        self,
        db: Session,
        *,
        year: int | None = None,
        location: str | None = None,
        stamping_supplier: str | None = None,
        vehicle_model: str | None = None,
        version: str | None = None,
        part_number: str | None = None,
        audit_status: str | None = None,
    ) -> tuple[bytes, str]:
        require_at_least_one_export_filter(
            year,
            location,
            stamping_supplier,
            vehicle_model,
            version,
            part_number,
            audit_status,
        )
        wb = Workbook()
        _title, fee_count = self.write_to_workbook(
            wb,
            db,
            year=year,
            location=location,
            stamping_supplier=stamping_supplier,
            vehicle_model=vehicle_model,
            version=version,
            part_number=part_number,
            audit_status=audit_status,
        )

        buffer = BytesIO()
        wb.save(buffer)
        content = buffer.getvalue()

        stamp = now().strftime("%Y%m%d_%H%M%S")
        parts = ["冲压加工费"]
        if year:
            parts.append(f"{year}年")
        if location:
            parts.append(location)
        if stamping_supplier:
            parts.append(stamping_supplier)
        filename = "_".join(parts) + f"_{stamp}.xlsx"
        logger.info(
            "导出冲压加工费 %s 条主表 -> %s",
            fee_count,
            filename,
        )
        return content, filename

    def write_to_workbook(
        self,
        wb: Workbook,
        db: Session,
        *,
        year: int | None = None,
        location: str | None = None,
        stamping_supplier: str | None = None,
        vehicle_model: str | None = None,
        version: str | None = None,
        part_number: str | None = None,
        audit_status: str | None = None,
        sheet_title: str = "冲压加工费",
    ) -> tuple[str, int]:
        records = self.repository.list_for_export(
            db,
            year=year,
            location=location,
            stamping_supplier=stamping_supplier,
            vehicle_model=vehicle_model,
            version=version,
            part_number=part_number,
            audit_status=audit_status,
        )

        ws = resolve_worksheet(wb, sheet_title)

        filter_text = self._filter_text(
            year=year,
            location=location,
            stamping_supplier=stamping_supplier,
            vehicle_model=vehicle_model,
            version=version,
            part_number=part_number,
            audit_status=audit_status,
        )
        self._write_banner(ws, filter_text=filter_text, fee_count=len(records))
        header_row = self._write_headers(ws)

        row_idx = header_row + 1
        export_row_count = 0
        for seq_no, record in enumerate(records, start=1):
            items = sorted(
                list(record.process_items or []),
                key=lambda item: (
                    item.process_order if item.process_order is not None else 10**9,
                    item.id or 0,
                ),
            )
            main_fill = _fill("ECEFF1") if seq_no % 2 else _fill("F5F7FA")
            if not items:
                self._write_data_row(
                    ws,
                    row_idx,
                    record,
                    seq_no=seq_no,
                    item=None,
                    write_main=True,
                    line_index=0,
                    main_fill=main_fill,
                )
                ws.row_dimensions[row_idx].height = DATA_H
                row_idx += 1
                export_row_count += 1
                continue

            group_start = row_idx
            for i, item in enumerate(items):
                self._write_data_row(
                    ws,
                    row_idx,
                    record,
                    seq_no=seq_no,
                    item=item,
                    write_main=(i == 0),
                    line_index=i,
                    main_fill=main_fill,
                )
                ws.row_dimensions[row_idx].height = self._row_height(item)
                row_idx += 1
                export_row_count += 1
            if row_idx - 1 > group_start:
                self._merge_main(ws, group_start, row_idx - 1, main_fill)

        last_data = max(header_row, row_idx - 1)
        for idx, width in enumerate(WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(COL_COUNT)}{last_data}"
        ws.print_title_rows = f"1:{header_row}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.3
        ws.page_margins.right = 0.3
        ws.page_margins.top = 0.35
        ws.page_margins.bottom = 0.35
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90

        logger.info(
            "写入冲压加工费 sheet=%s 主表=%s 明细行=%s",
            ws.title,
            len(records),
            export_row_count,
        )
        return ws.title, len(records)

    def _write_banner(self, ws: Any, *, filter_text: str, fee_count: int) -> None:
        indent = Alignment(horizontal="left", vertical="center", indent=1)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COL_COUNT)
        c1 = ws.cell(row=1, column=1, value="冲压加工费核对表")
        c1.font = _font(size=15, bold=True, color="FFFFFF")
        c1.fill = _fill("1A237E")
        c1.alignment = indent
        ws.row_dimensions[1].height = TITLE_H

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=COL_COUNT)
        c2 = ws.cell(
            row=2,
            column=1,
            value=(
                f"{filter_text} · 共 {fee_count} 个零件"
                f"　　导出 {now().strftime('%Y-%m-%d %H:%M')}"
            ),
        )
        c2.font = _font(size=9, color="455A64")
        c2.fill = _fill("E8EAF6")
        c2.alignment = indent
        ws.row_dimensions[2].height = FILTER_H

        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=COL_COUNT)
        c3 = ws.cell(
            row=3,
            column=1,
            value=(
                "说明：左侧主表按零件合并；右侧为工序明细。"
                "单序费用 = 冲次单价 × 费用比例（按工序核对吨位与单价）。"
            ),
        )
        c3.font = _font(size=9, color="B71C1C")
        c3.fill = _fill("FFF3E0")
        c3.alignment = indent
        ws.row_dimensions[3].height = NOTE_H

    def _write_headers(self, ws: Any) -> int:
        group_row = 4
        header_row = 5
        white = _font(size=9, bold=True, color="FFFFFF")
        header_font = _font(bold=True, color="FFFFFF")

        ws.merge_cells(
            start_row=group_row,
            start_column=1,
            end_row=group_row,
            end_column=MAIN_COLS,
        )
        ws.merge_cells(
            start_row=group_row,
            start_column=MAIN_COLS + 1,
            end_row=group_row,
            end_column=COL_COUNT,
        )
        main_group = ws.cell(row=group_row, column=1, value="零件主表")
        main_group.font = white
        main_group.fill = _fill("37474F")
        main_group.alignment = self._center
        for col in range(1, MAIN_COLS + 1):
            cell = ws.cell(row=group_row, column=col)
            cell.fill = _fill("37474F")
            cell.border = self._thin if col < MAIN_COLS else self._split_main

        detail_group = ws.cell(
            row=group_row,
            column=MAIN_COLS + 1,
            value="工序明细 / 核对（序 · 吨位 · 单价 · 单序费用）",
        )
        detail_group.font = white
        detail_group.fill = _fill("B71C1C")
        detail_group.alignment = self._center
        for col in range(MAIN_COLS + 1, COL_COUNT + 1):
            cell = ws.cell(row=group_row, column=col)
            cell.fill = _fill("B71C1C")
            cell.border = self._split_detail if col == MAIN_COLS + 1 else self._thin
        ws.row_dimensions[group_row].height = GROUP_H

        fee_cols = {12, 13, 14, 20, 21}  # 系统/其它/确认、冲次单价、单序费用
        for col, title in enumerate(HEADERS, start=1):
            cell = ws.cell(row=header_row, column=col, value=title)
            cell.font = header_font
            cell.alignment = self._center
            if col <= MAIN_COLS:
                cell.fill = _fill("EF6C00") if col in fee_cols else _fill("546E7A")
                cell.border = self._thin if col < MAIN_COLS else self._split_main
            elif col in fee_cols:
                cell.fill = _fill("EF6C00")
                cell.border = self._thin
            else:
                cell.fill = _fill("C62828")
                cell.border = self._split_detail if col == MAIN_COLS + 1 else self._thin
        ws.row_dimensions[header_row].height = HEADER_H
        return header_row

    def _write_data_row(
        self,
        ws: Any,
        row_idx: int,
        record: StampingProcessFee,
        *,
        seq_no: int,
        item: StampingProcessFeeItem | None,
        write_main: bool,
        line_index: int,
        main_fill: PatternFill,
    ) -> None:
        detail_fill = _fill("FFFFFF") if line_index % 2 == 0 else _fill("FFEBEE")

        if write_main:
            main_values: list[Any] = [
                seq_no,
                record.year,
                record.location,
                record.stamping_supplier,
                record.vehicle_model,
                record.version,
                record.part_number,
                record.part_name or "",
                record.material or "",
                _num(record.thickness, Q4),
                record.process_count,
                _num(record.processing_fee, Q3),
                _num(record.other_fee, Q3),
                _num(record.final_processing_fee, Q3),
            ]
            for col, value in enumerate(main_values, start=1):
                cell = ws.cell(row=row_idx, column=col, value="" if value is None else value)
                cell.fill = main_fill
                cell.border = self._thin if col < MAIN_COLS else self._split_main
                if col == 1:
                    cell.font = _font(size=11, bold=True, color="37474F")
                    cell.alignment = self._center
                elif col == 7:
                    cell.font = _font(bold=True, color="1565C0")
                    cell.alignment = self._left
                elif col == 14:
                    cell.font = _font(size=11, bold=True, color="E65100")
                    cell.alignment = self._right
                    cell.fill = _fill("FFF3E0")
                elif col in {11, 12, 13}:
                    cell.font = _font(bold=True) if col == 11 else _font()
                    cell.alignment = self._right if col in {12, 13} else self._center
                elif col == 10:
                    cell.font = _font()
                    cell.alignment = self._right
                else:
                    cell.font = _font()
                    cell.alignment = self._left if col in {3, 4, 5, 6, 8, 9} else self._center

        tonnage_source = (item.tonnage_source if item else "") or ""
        detail_values: list[Any] = [
            (item.process_seq or "") if item else "",
            (item.process_code or "") if item else "",
            item.resolved_tonnage if item and item.resolved_tonnage is not None else "",
            _TONNAGE_SOURCE_LABEL.get(tonnage_source, tonnage_source) if item else "",
            _num(item.cost_ratio, Q4) if item and item.cost_ratio is not None else "",
            _num(item.price_per_stroke, Q3) if item and item.price_per_stroke is not None else "",
            _num(item.line_fee, Q3) if item else "",
            (item.basis_text or "").strip() if item else "",
        ]
        for offset, value in enumerate(detail_values):
            col = MAIN_COLS + 1 + offset
            cell = ws.cell(row=row_idx, column=col, value="" if value is None else value)
            cell.fill = detail_fill
            cell.border = self._split_detail if col == MAIN_COLS + 1 else self._thin
            if col == 21:  # 单序费用
                cell.font = _font(size=11, bold=True, color="E65100")
                cell.alignment = self._right
                cell.fill = _fill("FFF3E0")
            elif col in {17, 19, 20}:  # 吨位、比例、冲次单价
                cell.font = _font()
                cell.alignment = self._right
            elif col == 22:  # 计算依据
                cell.font = _font(size=9, color="607D8B")
                cell.alignment = self._left
            else:
                cell.font = _font()
                cell.alignment = self._center

    def _merge_main(
        self,
        ws: Any,
        start_row: int,
        end_row: int,
        main_fill: PatternFill,
    ) -> None:
        for col in range(1, MAIN_COLS + 1):
            ws.merge_cells(
                start_row=start_row,
                start_column=col,
                end_row=end_row,
                end_column=col,
            )
            cell = ws.cell(row=start_row, column=col)
            if col == 14:
                cell.fill = _fill("FFF3E0")
            else:
                cell.fill = main_fill
            cell.border = self._thin if col < MAIN_COLS else self._split_main

    @staticmethod
    def _row_height(item: StampingProcessFeeItem) -> float:
        basis = (item.basis_text or "").strip()
        extra = 0.0
        if len(basis) > 40:
            extra = min(36.0, (len(basis) // 40) * 8.0)
        return DATA_H + extra

    @staticmethod
    def _filter_text(
        *,
        year: int | None,
        location: str | None,
        stamping_supplier: str | None,
        vehicle_model: str | None,
        version: str | None,
        part_number: str | None,
        audit_status: str | None,
    ) -> str:
        parts: list[str] = []
        if year is not None:
            parts.append(f"年度={year}")
        if location:
            parts.append(f"基地={location}")
        if stamping_supplier:
            parts.append(f"冲压工厂={stamping_supplier}")
        if vehicle_model:
            parts.append(f"车型={vehicle_model}")
        if version:
            parts.append(f"版本={version}")
        if part_number:
            parts.append(f"零件号含={part_number}")
        if audit_status:
            parts.append(f"审核={audit_status_label(audit_status)}")
        return " · ".join(parts) if parts else "全部筛选"
