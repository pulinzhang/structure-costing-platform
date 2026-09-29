"""钢材采购测算表：一行一件（拼焊子片分行），数量手填、吨为公式。"""

from __future__ import annotations

import logging
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from sqlalchemy.orm import Session

from app.common.bom.manufacturing_category import short_part_number_digits
from app.common.measures import KG_PER_TON, MM3_G_PER_CM3_TO_KG
from app.common.money import Q2, Q4
from app.common.timezone import now
from app.models.stamping import StampingPart, StampingProcess, StampingTailorWeldBlank
from app.schemas.vo.user_vo import UserVO
from app.services.settings.system.service import SettingsService
from app.services.stamping.part.export.constants import (
    _FONT_NAME,
    _HEADER_FILL,
    _LOCKED,
    _META_FILL,
    _MUTE_FILL,
    _MUTE_HEADER_FILL,
    _MUTE_SUB_HEADER_FILL,
    _QTY_FILL,
    _SHARED_HINT_FILL,
    _SUB_HEADER_FILL,
    _THIN,
    _TITLE_FILL,
    _TON_FILL,
    _UNLOCKED,
    _ZEBRA,
    DEFAULT_ROW_HEIGHT,
    STEEL_COL_COUNT,
    STEEL_DATA_START_ROW,
    STEEL_DENSITY_CELL,
    STEEL_MUTE_COLS,
    STEEL_QTY_COL,
    STEEL_SHARED_COLS,
    STEEL_SHEET_TITLE,
    STEEL_SHORT_CODE_COL,
    STEEL_TON_COL,
    STEEL_WEIGHT_COL,
    STEEL_WIDTHS,
)
from app.services.stamping.part.export.helpers import PartExportStyle, _num
from app.services.stamping.part.export.process_confirm import ProcessConfirmExport
from app.services.stamping.part.upload.cost_ratio_resolver import (
    has_shared_die_ratio,
    shared_die_labels,
)

logger = logging.getLogger("service.stamping.part.export")


class SteelPurchaseExport:
    _center: Alignment
    _left: Alignment
    _right: Alignment
    _header_font: Font
    _title_font: Font
    _meta_font: Font
    _data_font: Font
    _mute_header_font: Font
    _mute_data_font: Font
    _short_code_font: Font

    def __init__(self, confirm: ProcessConfirmExport, style: PartExportStyle) -> None:
        self._confirm = confirm
        self.repository = confirm.repository
        style.bind(self)

    def export_steel_purchase(
        self,
        db: Session,
        *,
        location: str,
        stamping_supplier: str | None = None,
        vehicle_model: str | None = None,
        version: str | None = None,
        part_number: str | None = None,
        current_user: UserVO,
    ) -> tuple[bytes, str]:
        location, stamping_supplier, vehicle_model, version, part_number, parts = (
            self._confirm._parts_for_export(
                db,
                location=location,
                stamping_supplier=stamping_supplier,
                vehicle_model=vehicle_model,
                version=version,
                part_number=part_number,
            )
        )
        density = SettingsService().get_steel_density(db)
        workbook = self._build_steel_purchase_workbook(
            parts,
            density=density,
            location=location,
            stamping_supplier=stamping_supplier,
            vehicle_model=vehicle_model,
            version=version,
        )
        filename = self._confirm._export_filename(
            "钢材采购测算",
            location,
            stamping_supplier,
            vehicle_model,
            version,
        )
        logger.info(
            "导出钢材采购测算 %s 件 / 文件 %s / 操作人 %s",
            len(parts),
            filename,
            current_user.username,
        )
        return self._confirm._workbook_to_plain_bytes(workbook), filename

    def _build_steel_purchase_workbook(
        self,
        parts: list[StampingPart],
        *,
        density: Any,
        location: str,
        stamping_supplier: str | None,
        vehicle_model: str | None,
        version: str | None,
    ) -> Workbook:
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = STEEL_SHEET_TITLE
        title_project = f"{vehicle_model}车身件" if vehicle_model else location
        title_mill = f"-{stamping_supplier}" if stamping_supplier else ""
        title = f"{title_project}钢材采购测算{title_mill}"
        row_count = self._steel_purchase_row_count(parts)
        self._write_steel_purchase_banner(
            sheet,
            title=title,
            location=location,
            stamping_supplier=stamping_supplier,
            vehicle_model=vehicle_model,
            version=version,
            part_count=row_count,
            density=density,
        )
        self._write_steel_purchase_headers(sheet)
        excel_row = STEEL_DATA_START_ROW
        for index, part in enumerate(parts):
            excel_row = self._write_steel_purchase_part(
                sheet,
                start_row=excel_row,
                part=part,
                zebra=index % 2 == 0,
            )
        last_data_row = excel_row - 1
        if last_data_row >= STEEL_DATA_START_ROW:
            self._write_steel_purchase_total(sheet, last_data_row)
            sheet.auto_filter.ref = f"A5:{get_column_letter(STEEL_COL_COUNT)}{last_data_row}"
        else:
            sheet.auto_filter.ref = f"A5:{get_column_letter(STEEL_COL_COUNT)}5"
        for idx, width in enumerate(STEEL_WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(idx)].width = width
        sheet.freeze_panes = f"{get_column_letter(STEEL_SHORT_CODE_COL + 1)}{STEEL_DATA_START_ROW}"
        self._confirm._apply_print_chrome(sheet)
        workbook.defined_names.add(
            DefinedName(
                name="SteelDensity",
                attr_text=f"'{STEEL_SHEET_TITLE}'!{STEEL_DENSITY_CELL}",
            )
        )
        self._protect_steel_purchase_sheet(sheet)
        self._write_steel_purchase_help(workbook)
        return workbook

    def _steel_purchase_row_count(self, parts: list[StampingPart]) -> int:
        total = 0
        for part in parts:
            blanks = self._confirm._sorted_blanks(part)
            if part.is_tailor_welded and blanks:
                total += len(blanks)
            else:
                total += 1
        return total

    def _write_steel_purchase_banner(
        self,
        ws: Any,
        *,
        title: str,
        location: str,
        stamping_supplier: str | None,
        vehicle_model: str | None,
        version: str | None,
        part_count: int,
        density: Any,
    ) -> None:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=STEEL_COL_COUNT)
        title_cell = ws.cell(row=1, column=1, value=title)
        title_cell.font = self._title_font
        title_cell.fill = _TITLE_FILL
        title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 28
        for col in range(1, STEEL_COL_COUNT + 1):
            ws.cell(row=1, column=col).fill = _TITLE_FILL
            ws.cell(row=1, column=col).border = _THIN

        scope_parts = [f"基地：{location}"]
        if stamping_supplier:
            scope_parts.append(f"冲压工厂：{stamping_supplier}")
        if vehicle_model:
            scope_parts.append(f"项目：{vehicle_model}")
        if version:
            scope_parts.append(f"版本：{version}")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=18)
        scope_cell = ws.cell(row=2, column=1, value="    ".join(scope_parts))
        scope_cell.font = self._meta_font
        scope_cell.fill = _META_FILL
        scope_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in range(1, STEEL_COL_COUNT + 1):
            cell = ws.cell(row=2, column=col)
            cell.fill = _META_FILL
            cell.border = _THIN
        count_cell = ws.cell(row=2, column=19, value=f"行数：{part_count}")
        count_cell.font = Font(name=_FONT_NAME, size=10, bold=True, color="1F4E79")
        count_cell.alignment = self._right
        ws.cell(row=2, column=20, value="密度(g/cm³)").font = self._meta_font
        ws.cell(row=2, column=20).alignment = self._right
        density_cell = ws.cell(row=2, column=STEEL_QTY_COL, value=_num(density, Q4))
        density_cell.font = Font(name=_FONT_NAME, size=10, bold=True, color="1F4E79")
        density_cell.alignment = self._right
        density_cell.protection = _UNLOCKED
        density_cell.fill = _QTY_FILL
        ws.cell(row=2, column=STEEL_TON_COL, value=now().strftime("%Y-%m-%d %H:%M"))
        ws.cell(row=2, column=STEEL_TON_COL).font = self._meta_font
        ws.cell(row=2, column=STEEL_TON_COL).alignment = self._right
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[3].height = 8

    def _write_steel_purchase_headers(self, ws: Any) -> None:
        groups = [
            (1, 5, "范围"),
            (6, 9, "零件"),
            (10, 16, "料片"),
            (17, 18, "标记"),
            (19, 20, "共模"),
            (21, 22, "测算"),
        ]
        fields = {
            1: "序号",
            2: "基地",
            3: "项目",
            4: "版本",
            5: "冲压工厂",
            6: "零件号",
            7: "短码",
            8: "零件名称",
            9: "子片",
            10: "牌号",
            11: "厚度",
            12: "长/卷料宽",
            13: "宽/卷料步距",
            14: "可制件数",
            15: "材料重",
            16: "零件重量",
            17: "是否利废",
            18: "废料副产品",
            19: "是否共模",
            20: "共模标记",
            21: "主机厂数量",
            22: "采购钢材(吨)",
        }
        for start, end, label in groups:
            is_mute = start in STEEL_MUTE_COLS
            header_fill = _MUTE_HEADER_FILL if is_mute else _HEADER_FILL
            header_font = self._mute_header_font if is_mute else self._header_font
            if start != end:
                ws.merge_cells(start_row=4, start_column=start, end_row=4, end_column=end)
            cell = ws.cell(row=4, column=start, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = self._center
            cell.border = _THIN
            for col in range(start, end + 1):
                ws.cell(row=4, column=col).fill = header_fill
                ws.cell(row=4, column=col).border = _THIN
                ws.cell(row=4, column=col).font = header_font
        for col, value in fields.items():
            is_mute = col in STEEL_MUTE_COLS
            cell = ws.cell(row=5, column=col, value=value)
            cell.font = self._mute_header_font if is_mute else self._header_font
            cell.fill = _MUTE_SUB_HEADER_FILL if is_mute else _SUB_HEADER_FILL
            cell.alignment = self._center
            cell.border = _THIN
        ws.row_dimensions[4].height = 22
        ws.row_dimensions[5].height = 22

    def _write_steel_purchase_part(
        self,
        ws: Any,
        *,
        start_row: int,
        part: StampingPart,
        zebra: bool,
    ) -> int:
        processes = self._confirm._sorted_processes(part)
        blanks = self._confirm._sorted_blanks(part)
        if part.is_tailor_welded and blanks:
            row = start_row
            for blank in blanks:
                self._write_steel_purchase_row(
                    ws,
                    excel_row=row,
                    part=part,
                    processes=processes,
                    blank=blank,
                    zebra=zebra,
                )
                row += 1
            return row
        self._write_steel_purchase_row(
            ws,
            excel_row=start_row,
            part=part,
            processes=processes,
            blank=None,
            zebra=zebra,
        )
        return start_row + 1

    def _write_steel_purchase_row(
        self,
        ws: Any,
        *,
        excel_row: int,
        part: StampingPart,
        processes: list[StampingProcess],
        blank: StampingTailorWeldBlank | None,
        zebra: bool,
    ) -> None:
        material = (blank.material if blank and blank.material else part.material) or ""
        thickness = blank.thickness if blank and blank.thickness is not None else part.thickness
        length = (
            blank.sheet_length if blank and blank.sheet_length is not None else part.sheet_length
        )
        width = blank.sheet_width if blank and blank.sheet_width is not None else part.sheet_width
        parts_per_sheet = (
            blank.parts_per_sheet
            if blank and blank.parts_per_sheet is not None
            else part.parts_per_sheet
        )
        part_weight = (
            blank.part_weight if blank and blank.part_weight is not None else part.part_weight
        )
        process_codes = [item.process_code for item in processes]
        cost_ratios = [item.cost_ratio for item in processes]
        labels = shared_die_labels(
            die_type=part.die_type,
            process_method=part.process_method,
            process_codes=process_codes,
            cost_ratios=cost_ratios,
        )
        is_shared = bool(labels) or has_shared_die_ratio(cost_ratios)
        values: dict[int, Any] = {
            1: part.seq_no,
            2: part.location,
            3: part.vehicle_model,
            4: part.version,
            5: part.stamping_supplier or "",
            6: part.part_number,
            7: short_part_number_digits(part.part_number) or "",
            8: part.part_name or "",
            9: blank.blank_part_number if blank else "",
            10: material,
            11: _num(thickness, Q4),
            12: _num(length, Q2),
            13: _num(width, Q2),
            14: parts_per_sheet,
            15: (
                f'=IF(OR(L{excel_row}="",M{excel_row}="",K{excel_row}=""),"",'
                f"L{excel_row}*M{excel_row}*K{excel_row}*SteelDensity/{MM3_G_PER_CM3_TO_KG})"
            ),
            16: _num(part_weight, Q4),
            17: "是" if part.is_offal_nested else "否",
            18: "是" if part.is_scrap_byproduct else "否",
            19: "是" if is_shared else "否",
            20: "、".join(labels),
            21: None,
            22: (
                f'=IF(OR(Q{excel_row}="是",R{excel_row}="是"),0,'
                f'IF(OR(O{excel_row}="",N{excel_row}="",N{excel_row}<=0,U{excel_row}=""),"",'
                f"O{excel_row}/N{excel_row}*U{excel_row}/{KG_PER_TON}))"
            ),
        }
        zebra_fill = _ZEBRA if zebra else None
        ws.row_dimensions[excel_row].height = DEFAULT_ROW_HEIGHT
        for col in range(1, STEEL_COL_COUNT + 1):
            cell = ws.cell(row=excel_row, column=col, value=values.get(col))
            cell.font = self._mute_data_font if col in STEEL_MUTE_COLS else self._data_font
            cell.border = _THIN
            cell.protection = _LOCKED
            if col in (6, 8, 9, 10, 20):
                cell.alignment = self._left
            elif col in (11, 12, 13, 14, 15, 16, 22):
                cell.alignment = self._right
            else:
                cell.alignment = self._center
            if col == STEEL_SHORT_CODE_COL:
                cell.font = self._short_code_font
            if col == STEEL_QTY_COL:
                cell.fill = _QTY_FILL
                cell.protection = _UNLOCKED
                cell.comment = Comment("在此填写", "系统")
            elif col == STEEL_TON_COL:
                cell.fill = _TON_FILL
                cell.number_format = "0.0000"
            elif col == STEEL_WEIGHT_COL:
                cell.number_format = "0.0000"
                if zebra_fill is not None:
                    cell.fill = zebra_fill
            elif is_shared and col in STEEL_SHARED_COLS:
                cell.fill = _SHARED_HINT_FILL
            elif col in STEEL_MUTE_COLS:
                cell.fill = _MUTE_FILL
            elif zebra_fill is not None:
                cell.fill = zebra_fill

    def _write_steel_purchase_total(self, ws: Any, last_data_row: int) -> None:
        total_row = last_data_row + 1
        ws.merge_cells(
            start_row=total_row,
            start_column=1,
            end_row=total_row,
            end_column=STEEL_TON_COL - 1,
        )
        for col in range(1, STEEL_COL_COUNT + 1):
            cell = ws.cell(row=total_row, column=col)
            cell.fill = _TITLE_FILL
            cell.font = self._header_font
            cell.border = _THIN
            cell.protection = _LOCKED
        label = ws.cell(row=total_row, column=1, value="合计")
        label.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        ton = ws.cell(
            row=total_row,
            column=STEEL_TON_COL,
            value=f"=SUM({get_column_letter(STEEL_TON_COL)}{STEEL_DATA_START_ROW}:"
            f"{get_column_letter(STEEL_TON_COL)}{last_data_row})",
        )
        ton.number_format = "0.0000"
        ton.alignment = self._right
        ws.row_dimensions[total_row].height = 20

    def _protect_steel_purchase_sheet(self, ws: Any) -> None:
        ws.protection.sheet = True

    def _write_steel_purchase_help(self, workbook: Workbook) -> None:
        sheet = workbook.create_sheet("填写说明")
        sheet["A1"] = "填写说明"
        sheet["A1"].font = self._title_font
        sheet["A1"].fill = _TITLE_FILL
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        for col in (1, 2):
            sheet.cell(row=1, column=col).fill = _TITLE_FILL
            sheet.cell(row=1, column=col).border = _THIN
        headers = ("项", "说明")
        for col, text in enumerate(headers, start=1):
            cell = sheet.cell(row=2, column=col, value=text)
            cell.font = self._header_font
            cell.fill = _SUB_HEADER_FILL
            cell.alignment = self._center
            cell.border = _THIN
        rows = [
            (
                "主机厂数量",
                "只填黄色列。材料重、采购吨由公式算出。取消保护：审阅 → 撤销工作表保护（无密码）。",
            ),
            ("共模", "是否共模为「是」时，成对件数量应相同。"),
            ("利废 / 废料副产品", "采购钢材(吨)为 0，不买新板。"),
            ("可制件数", "一张料能出几件成品。吨 = 材料重 ÷ 可制件数 × 数量 ÷ 1000。"),
            ("密度", "表头黄色格可改；改后材料重跟着变。"),
            ("拼焊", "子片分行，各行填同一总成件数。"),
        ]
        for index, (item, detail) in enumerate(rows, start=3):
            item_cell = sheet.cell(row=index, column=1, value=item)
            detail_cell = sheet.cell(row=index, column=2, value=detail)
            for cell in (item_cell, detail_cell):
                cell.font = self._data_font
                cell.border = _THIN
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            if index % 2 == 0:
                item_cell.fill = _ZEBRA
                detail_cell.fill = _ZEBRA
            sheet.row_dimensions[index].height = 32
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 72
        sheet.row_dimensions[1].height = 28
        self._confirm._apply_print_chrome(sheet)
