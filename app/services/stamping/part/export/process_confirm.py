"""工艺确认表：主表纵向合并 + 工序/拼焊子片拉开 + 简图。"""

from __future__ import annotations

import logging
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.period import format_period, require_export_as_of
from app.common.timezone import now
from app.exceptions import BusinessException
from app.models.stamping import StampingPart, StampingProcess, StampingTailorWeldBlank
from app.repositories.stamping.part.repository import StampingPartRepository
from app.schemas.vo.user_vo import UserVO
from app.services.stamping.part.export.constants import (
    _FONT_NAME,
    _HEADER_FILL,
    _META_FILL,
    _SUB_HEADER_FILL,
    _THIN,
    _TITLE_FILL,
    COL_COUNT,
    DATA_START_ROW,
    WIDTHS,
)
from app.services.stamping.part.export.helpers import PartExportStyle
from app.services.stamping.part.export.part_block import write_part_block
from app.services.stamping.part.export.sketch import embed_sketch

logger = logging.getLogger("service.stamping.part.export")


class ProcessConfirmExport:
    _center: Alignment
    _left: Alignment
    _right: Alignment
    _header_font: Font
    _title_font: Font
    _meta_font: Font
    _data_font: Font
    _mute_header_font: Font
    _mute_data_font: Font
    _short_code_font: Font

    def __init__(self, repository: StampingPartRepository, style: PartExportStyle) -> None:
        self.repository = repository
        style.bind(self)

    def export(
        self,
        db: Session,
        *,
        location: str,
        stamping_supplier: str | None = None,
        vehicle_model: str | None = None,
        version: str | None = None,
        part_number: str | None = None,
        year: int | None = None,
        month: int | None = None,
        include_expired: bool = False,
        current_user: UserVO,
    ) -> tuple[bytes, str]:
        if include_expired:
            as_of_year, as_of_month = None, None
            period_label = "全部"
        else:
            as_of_year, as_of_month = require_export_as_of(year, month)
            period_label = format_period(as_of_year, as_of_month)
        location, stamping_supplier, vehicle_model, version, part_number, parts = (
            self._parts_for_export(
                db,
                location=location,
                stamping_supplier=stamping_supplier,
                vehicle_model=vehicle_model,
                version=version,
                part_number=part_number,
                as_of_year=as_of_year,
                as_of_month=as_of_month,
                include_expired=include_expired,
            )
        )

        wb = Workbook()
        ws = wb.worksheets[0]
        ws.title = "冲压工艺确认表"

        title_project = f"{vehicle_model}车身件项目" if vehicle_model else location
        title_mill = f"-{stamping_supplier}" if stamping_supplier else ""
        title = f"{title_project}冲压工艺确认表{title_mill}"
        self._write_banner(
            ws,
            title=title,
            location=location,
            stamping_supplier=stamping_supplier,
            vehicle_model=vehicle_model,
            version=version,
            period_label=period_label,
            part_count=len(parts),
        )
        self._write_headers(ws)

        row_idx = DATA_START_ROW
        process_count = 0
        blank_count = 0
        sketch_count = 0
        sketch_buffers: list[BytesIO] = []
        for index, part in enumerate(parts, start=1):
            processes = self._sorted_processes(part)
            blanks = self._sorted_blanks(part)
            process_count += len(processes)
            blank_count += len(blanks)
            display_seq = part.seq_no if part.seq_no is not None else index
            used, has_sketch = self._write_part_block(
                ws,
                start_row=row_idx,
                seq_no=display_seq,
                part=part,
                processes=processes,
                blanks=blanks,
                zebra=index % 2 == 0,
                sketch_buffers=sketch_buffers,
            )
            if has_sketch:
                sketch_count += 1
            row_idx += used

        for idx, width in enumerate(WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = f"A{DATA_START_ROW}"
        self._apply_print_chrome(ws)

        try:
            content = workbook_to_xlsx_bytes(wb, current_user)
        finally:
            sketch_buffers.clear()

        filename = self._export_filename(
            "冲压工艺确认表",
            location,
            stamping_supplier,
            vehicle_model,
            version,
            period_label,
        )
        logger.info(
            "导出冲压工艺 %s 件 / %s 道工序 / %s 片子片 / %s 张简图 / 文件 %s",
            len(parts),
            process_count,
            blank_count,
            sketch_count,
            filename,
        )
        return content, filename

    def _write_banner(
        self,
        ws: Any,
        *,
        title: str,
        location: str,
        stamping_supplier: str | None,
        vehicle_model: str | None,
        version: str | None,
        period_label: str,
        part_count: int,
    ) -> None:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COL_COUNT)
        c1 = ws.cell(row=1, column=1, value=title)
        c1.font = self._title_font
        c1.fill = _TITLE_FILL
        c1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 28

        scope_parts = [f"基地：{location}"]
        if stamping_supplier:
            scope_parts.append(f"冲压工厂：{stamping_supplier}")
        if vehicle_model:
            scope_parts.append(f"项目：{vehicle_model}")
        if version:
            scope_parts.append(f"版本：{version}")
        scope_parts.append(f"命中月：{period_label}")
        scope_text = "    ".join(scope_parts)

        count_start = 22
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=count_start - 1)
        c2 = ws.cell(row=2, column=1, value=scope_text)
        c2.font = self._meta_font
        c2.fill = _META_FILL
        c2.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws.merge_cells(start_row=2, start_column=count_start, end_row=2, end_column=COL_COUNT)
        c3 = ws.cell(row=2, column=count_start, value=f"零件数量：{part_count}")
        c3.font = Font(name=_FONT_NAME, size=10, bold=True, color="1F4E79")
        c3.fill = _META_FILL
        c3.alignment = Alignment(horizontal="right", vertical="center", indent=1)

        for col in range(1, COL_COUNT + 1):
            cell = ws.cell(row=2, column=col)
            cell.fill = _META_FILL
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[3].height = 8

    def _write_headers(self, ws: Any) -> None:
        row4 = {
            1: "序号",
            2: "零件号",
            3: "短码",
            4: "零件名称",
            5: "冲压工厂",
            6: "项目",
            7: "版本",
            8: "零件简图",
            9: "材质",
            10: "料厚",
            11: "零件重量",
            12: "料片尺寸（mm）",
            14: "可制件数",
            15: "料片式样",
            16: "是否利废",
            17: "利用哪类废料",
            18: "是否摆剪",
            19: "工艺",
            23: "使用冲压设备吨位",
            24: "模具类型",
            25: "工艺方法",
            26: "自动化",
            27: "备注",
            28: "费用比例",
        }
        row5 = {
            12: "长/卷料宽",
            13: "宽/卷料步距",
            19: "工序名称",
            20: "工序",
            21: "套数",
            22: "模具编号",
        }
        merges = [
            (4, 1, 5, 1),
            (4, 2, 5, 2),
            (4, 3, 5, 3),
            (4, 4, 5, 4),
            (4, 5, 5, 5),
            (4, 6, 5, 6),
            (4, 7, 5, 7),
            (4, 8, 5, 8),
            (4, 9, 5, 9),
            (4, 10, 5, 10),
            (4, 11, 5, 11),
            (4, 12, 4, 13),
            (4, 14, 5, 14),
            (4, 15, 5, 15),
            (4, 16, 5, 16),
            (4, 17, 5, 17),
            (4, 18, 5, 18),
            (4, 19, 4, 22),
            (4, 23, 5, 23),
            (4, 24, 5, 24),
            (4, 25, 5, 25),
            (4, 26, 5, 26),
            (4, 27, 5, 27),
            (4, 28, 5, 28),
        ]
        for r1, c1, r2, c2 in merges:
            if (r1, c1) != (r2, c2):
                ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

        for col in range(1, COL_COUNT + 1):
            for row in (4, 5):
                cell = ws.cell(row=row, column=col)
                cell.font = self._header_font
                cell.fill = _HEADER_FILL if row == 4 else _SUB_HEADER_FILL
                cell.alignment = self._center
                cell.border = _THIN
        for col, value in row4.items():
            ws.cell(row=4, column=col, value=value)
        for col, value in row5.items():
            ws.cell(row=5, column=col, value=value)
        ws.row_dimensions[4].height = 22
        ws.row_dimensions[5].height = 22

    def _normalize_export_scope(
        self,
        *,
        location: str,
        stamping_supplier: str | None,
        vehicle_model: str | None,
        version: str | None,
        part_number: str | None,
    ) -> tuple[str, str | None, str | None, str | None, str | None]:
        location = location.strip()
        stamping_supplier = (stamping_supplier or "").strip() or None
        vehicle_model = (vehicle_model or "").strip() or None
        version = (version or "").strip() or None
        part_number = (part_number or "").strip() or None
        if not location:
            raise BusinessException("导出须选定基地")
        return location, stamping_supplier, vehicle_model, version, part_number

    def _parts_for_export(
        self,
        db: Session,
        *,
        location: str,
        stamping_supplier: str | None,
        vehicle_model: str | None,
        version: str | None,
        part_number: str | None,
        as_of_year: int | None = None,
        as_of_month: int | None = None,
        include_expired: bool = False,
    ) -> tuple[str, str | None, str | None, str | None, str | None, list[StampingPart]]:
        location, stamping_supplier, vehicle_model, version, part_number = (
            self._normalize_export_scope(
                location=location,
                stamping_supplier=stamping_supplier,
                vehicle_model=vehicle_model,
                version=version,
                part_number=part_number,
            )
        )
        parts = self.repository.list_for_export(
            db,
            location=location,
            stamping_supplier=stamping_supplier,
            vehicle_model=vehicle_model,
            version=version,
            part_number=part_number,
            as_of_year=as_of_year,
            as_of_month=as_of_month,
            include_expired=include_expired,
        )
        if not parts:
            raise BusinessException("当前筛选下没有零件，无法导出")
        return location, stamping_supplier, vehicle_model, version, part_number, parts

    def _sorted_processes(self, part: StampingPart) -> list[StampingProcess]:
        return sorted(
            list(part.processes or []),
            key=lambda item: (item.process_order or 0, item.id or 0),
        )

    def _sorted_blanks(self, part: StampingPart) -> list[StampingTailorWeldBlank]:
        return sorted(
            list(part.tailor_weld_blanks or []),
            key=lambda item: (item.blank_seq or 0, item.id or 0),
        )

    def _apply_print_chrome(self, ws: Any) -> None:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90

    def _export_filename(
        self,
        prefix: str,
        location: str,
        stamping_supplier: str | None,
        vehicle_model: str | None,
        version: str | None,
        period_label: str | None = None,
    ) -> str:
        scope = "_".join(
            part
            for part in (location, stamping_supplier, vehicle_model, version, period_label)
            if part
        )
        stamp = now().strftime("%Y%m%d_%H%M%S")
        return f"{prefix}_{scope}_{stamp}.xlsx"

    def _workbook_to_plain_bytes(self, workbook: Workbook) -> bytes:
        """测算簿不涉密，不加内控水印。工艺确认表仍走 workbook_to_xlsx_bytes。"""
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _write_part_block(
        self,
        ws: Any,
        *,
        start_row: int,
        seq_no: int,
        part: StampingPart,
        processes: list[StampingProcess],
        blanks: list[StampingTailorWeldBlank],
        zebra: bool,
        sketch_buffers: list[BytesIO],
    ) -> tuple[int, bool]:
        return write_part_block(
            ws,
            start_row=start_row,
            seq_no=seq_no,
            part=part,
            processes=processes,
            blanks=blanks,
            zebra=zebra,
            sketch_buffers=sketch_buffers,
            data_font=self._data_font,
            short_code_font=self._short_code_font,
            center=self._center,
            left=self._left,
            right=self._right,
        )

    def _embed_sketch(
        self,
        ws: Any,
        *,
        part: StampingPart,
        start_row: int,
        n_rows: int,
        sketch_buffers: list[BytesIO],
    ) -> bool:
        return embed_sketch(
            ws,
            part=part,
            start_row=start_row,
            n_rows=n_rows,
            sketch_buffers=sketch_buffers,
        )
