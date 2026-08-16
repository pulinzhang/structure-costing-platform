"""冲次单价 Excel 导出服务。"""

from __future__ import annotations

import logging
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from app.common.timezone import now
from app.repositories.stamping.tonnage_price_repository import StampingTonnagePriceRepository


logger = logging.getLogger("service.stamping_tonnage_price_export")


def _num(value: Any, digits: int = 3) -> Any:
    if value is None:
        return ""
    return round(float(value), digits)


def _dt(value: Any) -> Any:
    if value is None:
        return ""
    return value.strftime("%Y-%m-%d %H:%M:%S") if hasattr(value, "strftime") else str(value)


class StampingTonnagePriceExportService:
    """冲次单价 Excel 导出。"""

    def __init__(self) -> None:
        self.repository = StampingTonnagePriceRepository()

    def export(
        self,
        db: Session,
        *,
        location: str | None = None,
        year: int | None = None,
    ) -> tuple[bytes, str]:
        records = self.repository.list_for_export(
            db,
            location=location,
            year=year,
        )

        wb = Workbook()
        ws = wb.worksheets[0]
        ws.title = "冲次单价"

        header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        data_font = Font(name="微软雅黑", size=9)
        align_left = Alignment(horizontal="left", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )

        headers = [
            "基地",
            "年份",
            "压机吨位(T)",
            "单冲次价格(元)",
            "创建人",
            "创建时间",
            "更新人",
            "更新时间",
        ]
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for row_idx, record in enumerate(records, start=2):
            values = [
                record.location,
                record.year,
                record.tonnage,
                _num(record.price_per_stroke, 3),
                record.created_by or "",
                _dt(record.created_at),
                record.updated_by or "",
                _dt(record.updated_at),
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.font = data_font
                cell.border = thin_border
                if col in (2, 3):
                    cell.alignment = align_center
                elif col == 4:
                    cell.alignment = align_right
                    cell.number_format = "0.000"
                else:
                    cell.alignment = align_left

        widths = [12, 10, 14, 16, 12, 20, 12, 20]
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + idx)].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:H{max(1, len(records) + 1)}"

        buffer = BytesIO()
        wb.save(buffer)
        content = buffer.getvalue()

        stamp = now().strftime("%Y%m%d_%H%M%S")
        parts = ["冲次单价"]
        if location:
            parts.append(location)
        if year:
            parts.append(f"{year}年")
        filename = "_".join(parts) + f"_{stamp}.xlsx"

        logger.info(
            "导出冲次单价 %s 条 location=%s year=%s -> %s",
            len(records),
            location,
            year,
            filename,
        )
        return content, filename
