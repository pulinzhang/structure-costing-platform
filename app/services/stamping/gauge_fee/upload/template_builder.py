"""冲压检具费导入模板：按模具布点带出零件并填默认值。"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.common.system_defaults import DEFAULT_ALLOCATION_QTY
from app.models.stamping.gauge_fee import DEFAULT_GAUGE_FEE
from app.parsers.stamping.gauge_fee_parser import (
    HEADER_ALLOCATION_QTY,
    HEADER_GAUGE_FEE,
    HEADER_PART_NAME,
    HEADER_PART_NUMBER,
    HEADER_STATUS,
)
from app.repositories.stamping.die.repository import StampingDiePartScopeIdentity


def build_gauge_fee_template(
    parts: list[StampingDiePartScopeIdentity],
    *,
    location: str,
    vehicle_model: str,
) -> tuple[bytes, str]:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "检具费"

    fill = PatternFill("solid", fgColor="1E1E1E")
    font = Font(color="FFFFFF", bold=True, size=10)
    headers = [
        HEADER_PART_NUMBER,
        HEADER_PART_NAME,
        HEADER_GAUGE_FEE,
        HEADER_ALLOCATION_QTY,
        HEADER_STATUS,
    ]
    for col, title in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, part in enumerate(parts, start=2):
        sheet.cell(row=row_idx, column=1, value=part["part_number"])
        sheet.cell(row=row_idx, column=2, value=part["part_name"] or "")
        sheet.cell(row=row_idx, column=3, value=DEFAULT_GAUGE_FEE)
        sheet.cell(row=row_idx, column=4, value=DEFAULT_ALLOCATION_QTY)
        sheet.cell(row=row_idx, column=5, value="生效")
        sheet.cell(row=row_idx, column=3).number_format = "0.0000"
        sheet.cell(row=row_idx, column=4).number_format = "0.0000"

    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 22
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 22
    sheet.column_dimensions["E"].width = 10
    sheet.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    filename = f"冲压检具费导入模板_{location}_{vehicle_model}.xlsx"
    return buffer.getvalue(), filename
