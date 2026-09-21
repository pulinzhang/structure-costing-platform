"""冲压检具费 Excel 导出。"""

from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.money import Q0, Q4, excel_money_cell
from app.common.timezone import now
from app.common.vehicle_model import normalize_vehicle_model
from app.models.stamping.gauge_fee import STATUS_EFFECTIVE, STATUS_INVALID, StampingGaugeFee
from app.repositories.stamping.gauge_fee.repository import StampingGaugeFeeRepository
from app.schemas.vo.user_vo import UserVO

_COL_COUNT = 7
_DATA_START = 5
_FONT = "微软雅黑"
_TITLE_FILL = PatternFill("solid", fgColor="1A1A22")
_META_FILL = PatternFill("solid", fgColor="F4E8E8")
_HEADER_FILL = PatternFill("solid", fgColor="C62828")
_MONEY_HEADER_FILL = PatternFill("solid", fgColor="8B1515")
_MONEY_FILL = PatternFill("solid", fgColor="FFF5F5")
_ZEBRA = PatternFill("solid", fgColor="FAFAFA")
_THIN = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
_HEADERS = (
    "基地",
    "项目",
    "零件号",
    "零件名称",
    "检具费（万元）",
    "项目分摊数量（万台）",
    "状态",
)
_WIDTHS = (12, 14, 18, 22, 16, 22, 10)
_STATUS_LABELS = {
    STATUS_EFFECTIVE: "生效",
    STATUS_INVALID: "失效",
}


class StampingGaugeFeeExportService:
    def __init__(self) -> None:
        self.repository = StampingGaugeFeeRepository()

    def export(
        self,
        db: Session,
        *,
        location: str | None,
        vehicle_model: str | None,
        part_number: str | None,
        status: str | None,
        current_user: UserVO,
    ) -> tuple[bytes, str]:
        loc = (location or "").strip() or None
        vehicle = normalize_vehicle_model(vehicle_model) if vehicle_model else None
        part = (part_number or "").strip() or None
        records = self.repository.list_for_export(
            db,
            location=loc,
            vehicle_model=vehicle,
            part_number=part,
            status=status,
        )

        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = "检具费"

        stamp = now()
        self._write_banner(sheet, loc=loc, vehicle=vehicle, count=len(records), stamp=stamp)
        self._write_headers(sheet)
        self._write_rows(sheet, records)

        last_row = max(_DATA_START - 1 + len(records), _DATA_START)
        for idx, width in enumerate(_WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(idx)].width = width
        sheet.freeze_panes = f"A{_DATA_START}"
        sheet.auto_filter.ref = f"A{_DATA_START - 1}:{get_column_letter(_COL_COUNT)}{last_row}"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 110
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_options.horizontalCentered = True
        sheet.print_title_rows = "1:4"
        sheet.print_area = f"A1:{get_column_letter(_COL_COUNT)}{last_row}"

        content = workbook_to_xlsx_bytes(workbook, current_user)
        name_parts = ["冲压检具费"]
        if loc:
            name_parts.append(loc)
        if vehicle:
            name_parts.append(vehicle)
        filename = "_".join(name_parts) + f"_{stamp.strftime('%Y%m%d_%H%M%S')}.xlsx"
        return content, filename

    def _write_banner(
        self,
        sheet: Worksheet,
        *,
        loc: str | None,
        vehicle: str | None,
        count: int,
        stamp: datetime,
    ) -> None:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=_COL_COUNT)
        title = sheet.cell(row=1, column=1, value="冲压检具费一览")
        title.font = Font(name=_FONT, size=16, bold=True, color="FFFFFF")
        title.fill = _TITLE_FILL
        title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in range(1, _COL_COUNT + 1):
            sheet.cell(row=1, column=col).fill = _TITLE_FILL
        sheet.row_dimensions[1].height = 32

        scope_bits = [
            f"基地：{loc or '全部'}",
            f"项目：{vehicle or '全部'}",
            f"共 {count} 条",
            f"导出时间：{stamp.strftime('%Y-%m-%d %H:%M')}",
        ]
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=_COL_COUNT)
        meta = sheet.cell(row=2, column=1, value="    ".join(scope_bits))
        meta.font = Font(name=_FONT, size=10, color="6B1A1A")
        meta.fill = _META_FILL
        meta.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in range(1, _COL_COUNT + 1):
            sheet.cell(row=2, column=col).fill = _META_FILL
        sheet.row_dimensions[2].height = 22
        sheet.row_dimensions[3].height = 8

    def _write_headers(self, sheet: Worksheet) -> None:
        header_row = _DATA_START - 1
        header_font = Font(name=_FONT, size=10, bold=True, color="FFFFFF")
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col, title in enumerate(_HEADERS, start=1):
            cell = sheet.cell(row=header_row, column=col, value=title)
            cell.font = header_font
            cell.fill = _MONEY_HEADER_FILL if col in (5, 6) else _HEADER_FILL
            cell.alignment = center
            cell.border = _THIN
        sheet.row_dimensions[header_row].height = 28

    def _write_rows(self, sheet: Worksheet, records: list[StampingGaugeFee]) -> None:
        data_font = Font(name=_FONT, size=10, color="222222")
        money_font = Font(name=_FONT, size=10, bold=True, color="8B1515")
        left = Alignment(horizontal="left", vertical="center")
        center = Alignment(horizontal="center", vertical="center")
        right = Alignment(horizontal="right", vertical="center")
        for offset, record in enumerate(records):
            row_idx = _DATA_START + offset
            zebra = offset % 2 == 1
            values = [
                record.location,
                record.vehicle_model,
                record.part_number,
                record.part_name or "",
                excel_money_cell(record.gauge_fee, Q4),
                excel_money_cell(record.allocation_qty, Q0),
                _STATUS_LABELS.get(record.status, record.status),
            ]
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_idx, column=col, value=value)
                cell.font = money_font if col in (5, 6) else data_font
                cell.border = _THIN
                if col == 5:
                    cell.alignment = right
                    cell.number_format = "0.0000"
                    cell.fill = _MONEY_FILL
                elif col == 6:
                    cell.alignment = right
                    cell.number_format = "0"
                    cell.fill = _MONEY_FILL
                elif col == 7:
                    cell.alignment = center
                    if zebra:
                        cell.fill = _ZEBRA
                else:
                    cell.alignment = left
                    if zebra:
                        cell.fill = _ZEBRA
            sheet.row_dimensions[row_idx].height = 20
