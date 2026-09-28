"""模具列表导出（对齐冲压工艺确认表版式）。"""

from __future__ import annotations

from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.money import Q2, Q4, excel_money_cell
from app.common.period import format_period, require_export_as_of
from app.common.timezone import now
from app.common.vehicle_model import normalize_vehicle_model
from app.exceptions.base import BusinessException
from app.models.stamping.die import StampingDiePart, StampingDieProcess
from app.repositories.stamping.die.repository import StampingDiePartRepository
from app.schemas.vo.user_vo import UserVO

COL_COUNT = 32
DATA_START_ROW = 6
DEFAULT_ROW_HEIGHT = 18

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_SUB_HEADER_FILL = PatternFill("solid", fgColor="2E75B6")
_TITLE_FILL = PatternFill("solid", fgColor="1F4E79")
_META_FILL = PatternFill("solid", fgColor="D6EAF8")
_ZEBRA = PatternFill("solid", fgColor="F2F2F2")
_RATIO_FILL = PatternFill("solid", fgColor="FFF2CC")
_RATIO_HEADER_FILL = PatternFill("solid", fgColor="C65911")
_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_FONT_NAME = "微软雅黑"

WIDTHS = [
    6,
    16,
    20,
    18,
    10,
    12,
    8,
    8,
    10,
    10,
    10,
    8,
    11,
    11,
    10,
    8,
    6,
    9,
    9,
    9,
    9,
    10,
    10,
    8,
    10,
    10,
    10,
    12,
    8,
    18,
    10,
    14,
]

# 零件级列（多道工序时纵向合并）
_PART_COLS = (1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 23, 25, 26, 27, 28, 29, 30)
_RATIO_COL = 31
_DIE_NUMBER_COL = 32


class StampingDieExportService:
    def __init__(self) -> None:
        self.repository = StampingDiePartRepository()
        self._center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self._left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        self._right = Alignment(horizontal="right", vertical="center")
        self._header_font = Font(name=_FONT_NAME, size=10, bold=True, color="FFFFFF")
        self._title_font = Font(name=_FONT_NAME, size=14, bold=True, color="FFFFFF")
        self._meta_font = Font(name=_FONT_NAME, size=10, color="1F4E79")
        self._data_font = Font(name=_FONT_NAME, size=9, color="000000")

    def export_workbook(
        self,
        db: Session,
        *,
        location: str,
        vehicle_model: str,
        part_number: str | None = None,
        year: int | None = None,
        month: int | None = None,
        include_expired: bool = False,
        current_user: UserVO,
    ) -> tuple[bytes, str]:
        loc = (location or "").strip()
        vehicle = normalize_vehicle_model(vehicle_model) or ""
        if not loc or not vehicle:
            raise BusinessException("导出须选择基地和项目")
        if include_expired:
            as_of_year, as_of_month = None, None
            period_label = "全部"
        else:
            as_of_year, as_of_month = require_export_as_of(year, month)
            period_label = format_period(as_of_year, as_of_month)
        part_filter = (part_number or "").strip() or None
        parts = self.repository.list_for_export(
            db,
            location=loc,
            vehicle_model=vehicle,
            part_number=part_filter,
            as_of_year=as_of_year,
            as_of_month=as_of_month,
            include_expired=include_expired,
        )

        workbook = Workbook()
        sheet = workbook.active
        if not isinstance(sheet, Worksheet):
            sheet = workbook.create_sheet("模具列表")
        sheet.title = "模具列表"

        process_count = sum(len(part.processes or []) for part in parts)
        self._write_banner(
            sheet,
            location=loc,
            vehicle_model=vehicle,
            period_label=period_label,
            part_count=len(parts),
            process_count=process_count,
        )
        self._write_headers(sheet)

        row_idx = DATA_START_ROW
        for seq_no, part in enumerate(parts, start=1):
            processes = list(part.processes or [])
            used = self._write_part_block(
                sheet,
                start_row=row_idx,
                seq_no=part.seq_no if part.seq_no is not None else seq_no,
                part=part,
                processes=processes,
                zebra=seq_no % 2 == 0,
            )
            row_idx += used

        last_row = max(row_idx - 1, DATA_START_ROW - 1)
        for idx, width in enumerate(WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(idx)].width = width
        sheet.freeze_panes = f"A{DATA_START_ROW}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_options.horizontalCentered = True
        sheet.print_title_rows = "1:5"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 90
        if last_row >= DATA_START_ROW:
            sheet.print_area = f"A1:{get_column_letter(COL_COUNT)}{last_row}"

        content = workbook_to_xlsx_bytes(workbook, current_user)
        stamp = now().strftime("%Y%m%d_%H%M%S")
        filename = f"冲压模具布点工艺清单_{loc}_{vehicle}_{period_label}_{stamp}.xlsx"
        return content, filename

    def _write_banner(
        self,
        sheet: Worksheet,
        *,
        location: str,
        vehicle_model: str,
        period_label: str,
        part_count: int,
        process_count: int,
    ) -> None:
        title = f"{vehicle_model}项目冲压模具布点工艺清单"
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COL_COUNT)
        title_cell = sheet.cell(row=1, column=1, value=title)
        title_cell.font = self._title_font
        title_cell.fill = _TITLE_FILL
        title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        sheet.row_dimensions[1].height = 28

        scope_bits = [f"基地：{location}", f"项目：{vehicle_model}", f"命中月：{period_label}"]
        count_start = 26
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=count_start - 1)
        scope_cell = sheet.cell(
            row=2,
            column=1,
            value="    ".join(scope_bits),
        )
        scope_cell.font = self._meta_font
        scope_cell.fill = _META_FILL
        scope_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        sheet.merge_cells(start_row=2, start_column=count_start, end_row=2, end_column=COL_COUNT)
        count_cell = sheet.cell(
            row=2,
            column=count_start,
            value=f"零件：{part_count}    工序：{process_count}",
        )
        count_cell.font = Font(name=_FONT_NAME, size=10, bold=True, color="1F4E79")
        count_cell.fill = _META_FILL
        count_cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)

        for col in range(1, COL_COUNT + 1):
            sheet.cell(row=2, column=col).fill = _META_FILL
        sheet.row_dimensions[2].height = 20
        sheet.row_dimensions[3].height = 8

    def _write_headers(self, sheet: Worksheet) -> None:
        row4 = {
            1: "序号",
            2: "零件号",
            3: "零件名称",
            4: "",
            5: "",
            6: "材质",
            7: "料厚\n(mm)",
            8: "净重\n(kg)",
            9: "预估板料尺寸 (mm)",
            11: "板料形状",
            12: "可制件数",
            13: "材料消耗\n(kg)",
            14: "材料利用率",
            15: "预估工艺",
            18: "预估模具尺寸 (mm)",
            21: "掏空系数",
            22: "预估重量\n(T)",
            23: "重量小计\n(T)",
            24: "设备吨位\n(T)",
            25: "模具类型",
            26: "模具方式",
            27: "生产方式",
            28: "端拾器",
            29: "零件等级",
            30: "备注",
            31: "分摊比例",
            32: "模具编号",
        }
        row5 = {
            9: "长",
            10: "宽",
            15: "工序名称",
            16: "工序",
            17: "套数",
            18: "长度",
            19: "宽度",
            20: "高度",
        }
        merges = [
            (4, 1, 5, 1),
            (4, 2, 5, 2),
            (4, 3, 5, 3),
            (4, 4, 5, 4),
            (4, 5, 5, 5),
            (4, 6, 5, 6),
            (4, 7, 5, 7),
            (4, 8, 5, 8),
            (4, 9, 4, 10),
            (4, 11, 5, 11),
            (4, 12, 5, 12),
            (4, 13, 5, 13),
            (4, 14, 5, 14),
            (4, 15, 4, 17),
            (4, 18, 4, 20),
            (4, 21, 5, 21),
            (4, 22, 5, 22),
            (4, 23, 5, 23),
            (4, 24, 5, 24),
            (4, 25, 5, 25),
            (4, 26, 5, 26),
            (4, 27, 5, 27),
            (4, 28, 5, 28),
            (4, 29, 5, 29),
            (4, 30, 5, 30),
            (4, 31, 5, 31),
            (4, 32, 5, 32),
        ]
        for start_row, start_col, end_row, end_col in merges:
            if (start_row, start_col) != (end_row, end_col):
                sheet.merge_cells(
                    start_row=start_row,
                    start_column=start_col,
                    end_row=end_row,
                    end_column=end_col,
                )

        for col in range(1, COL_COUNT + 1):
            for row in (4, 5):
                cell = sheet.cell(row=row, column=col)
                cell.font = self._header_font
                if col == _RATIO_COL:
                    cell.fill = _RATIO_HEADER_FILL
                elif row == 4:
                    cell.fill = _HEADER_FILL
                else:
                    cell.fill = _SUB_HEADER_FILL
                cell.alignment = self._center
                cell.border = _THIN
        for col, value in row4.items():
            sheet.cell(row=4, column=col, value=value)
        for col, value in row5.items():
            sheet.cell(row=5, column=col, value=value)
        sheet.row_dimensions[4].height = 24
        sheet.row_dimensions[5].height = 22

    def _write_part_block(
        self,
        sheet: Worksheet,
        *,
        start_row: int,
        seq_no: int,
        part: StampingDiePart,
        processes: list[StampingDieProcess],
        zebra: bool,
    ) -> int:
        n_rows = max(len(processes), 1)
        end_row = start_row + n_rows - 1
        fill = _ZEBRA if zebra else None

        for offset in range(n_rows):
            excel_row = start_row + offset
            sheet.row_dimensions[excel_row].height = DEFAULT_ROW_HEIGHT
            for col in range(1, COL_COUNT + 1):
                cell = sheet.cell(row=excel_row, column=col, value=None)
                cell.font = self._data_font
                cell.border = _THIN
                cell.alignment = self._center
                if col == _RATIO_COL:
                    cell.fill = _RATIO_FILL
                elif fill is not None:
                    cell.fill = fill

        for col in _PART_COLS:
            _merge_if_needed(sheet, start_row, end_row, col)

        sheet.cell(row=start_row, column=1, value=seq_no)
        sheet.cell(row=start_row, column=2, value=part.part_number or "").alignment = self._left
        sheet.cell(row=start_row, column=3, value=part.part_name or "").alignment = self._left
        sheet.cell(row=start_row, column=4, value="")
        sheet.cell(row=start_row, column=5, value="")
        sheet.cell(row=start_row, column=6, value=part.material or "")
        _set_num(sheet, start_row, 7, part.thickness, Q4, "0.0000")
        _set_num(sheet, start_row, 8, part.part_weight, Q4, "0.0000")
        _set_num(sheet, start_row, 9, part.sheet_length, Q2, "0.00")
        _set_num(sheet, start_row, 10, part.sheet_width, Q2, "0.00")
        sheet.cell(row=start_row, column=11, value=part.sheet_shape or "")
        sheet.cell(row=start_row, column=12, value=part.parts_per_sheet)
        _set_num(sheet, start_row, 13, part.material_consumption, Q4, "0.0000")
        _set_num(sheet, start_row, 14, part.material_utilization, Q4, "0.00%")
        _set_num(sheet, start_row, 23, part.estimated_weight_subtotal_t, Q4, "0.0000")
        sheet.cell(row=start_row, column=25, value=part.die_type or "")
        sheet.cell(row=start_row, column=26, value=part.die_method or "")
        sheet.cell(row=start_row, column=27, value=part.production_mode or "")
        sheet.cell(row=start_row, column=28, value=part.needs_end_effector or "")
        sheet.cell(row=start_row, column=29, value=part.part_grade or "")
        sheet.cell(row=start_row, column=30, value=part.remark or "").alignment = self._left

        if not processes:
            return n_rows

        for offset, process in enumerate(processes):
            excel_row = start_row + offset
            sheet.cell(row=excel_row, column=15, value=process.process_seq)
            sheet.cell(row=excel_row, column=16, value=process.process_code or "")
            sheet.cell(row=excel_row, column=17, value=process.die_sets)
            _set_num(sheet, excel_row, 18, process.die_length, Q2, "0.00")
            _set_num(sheet, excel_row, 19, process.die_width, Q2, "0.00")
            _set_num(sheet, excel_row, 20, process.die_height, Q2, "0.00")
            _set_num(sheet, excel_row, 21, process.hollow_factor, Q4, "0.0000")
            _set_num(sheet, excel_row, 22, process.estimated_weight_t, Q4, "0.0000")
            sheet.cell(row=excel_row, column=24, value=process.tonnage)
            _set_num(sheet, excel_row, _RATIO_COL, process.alloc_ratio, Q2, "0.00")
            sheet.cell(row=excel_row, column=_RATIO_COL).fill = _RATIO_FILL
            sheet.cell(row=excel_row, column=_DIE_NUMBER_COL, value=process.die_number or "")
        return n_rows


def _set_num(
    sheet: Worksheet,
    row: int,
    col: int,
    value: object,
    quant: Decimal,
    number_format: str,
) -> None:
    cell = sheet.cell(row=row, column=col, value=excel_money_cell(value, quant))
    cell.alignment = Alignment(horizontal="right", vertical="center")
    if cell.value != "":
        cell.number_format = number_format


def _merge_if_needed(sheet: Worksheet, start_row: int, end_row: int, col: int) -> None:
    if end_row > start_row:
        sheet.merge_cells(
            start_row=start_row,
            start_column=col,
            end_row=end_row,
            end_column=col,
        )
