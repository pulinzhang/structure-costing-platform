"""冲压零件 Excel 导出。

工艺确认表：主表纵向合并 + 工序/拼焊子片拉开 + 简图。
钢材采购测算：同一取数排序，一行一件（拼焊按子片分行），数量手填、吨为公式。
"""

from __future__ import annotations

import logging
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from sqlalchemy.orm import Session

from app.common.bom.manufacturing_category import short_part_number_digits
from app.common.export_watermark import workbook_to_xlsx_bytes
from app.common.money import Q2, Q4, excel_money_cell
from app.common.timezone import now
from app.config import settings
from app.exceptions import BusinessException
from app.models.stamping import StampingPart, StampingProcess, StampingTailorWeldBlank
from app.repositories.stamping.stamping_part_repository import StampingPartRepository
from app.schemas.vo.user_vo import UserVO
from app.services.settings.system.settings_service import SettingsService
from app.services.stamping.upload.cost_ratio_resolver import has_shared_die_ratio, shared_die_labels

logger = logging.getLogger("service.stamping.part_export")

COL_COUNT = 28
DATA_START_ROW = 6
SKETCH_COL = 8
SHORT_CODE_COL = 3
DEFAULT_ROW_HEIGHT = 18
# 简图列加宽，约 130px，给 PNG 留出可视区域
SKETCH_COL_WIDTH = 18
SKETCH_MAX_WIDTH_PX = 118
SKETCH_MIN_HEIGHT_PX = 88
SKETCH_PAD_PT = 8
_PX_TO_PT = 72 / 96

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_SUB_HEADER_FILL = PatternFill("solid", fgColor="2E75B6")
_TITLE_FILL = PatternFill("solid", fgColor="1F4E79")
_META_FILL = PatternFill("solid", fgColor="D6EAF8")
_ZEBRA = PatternFill("solid", fgColor="F2F2F2")
_QTY_FILL = PatternFill("solid", fgColor="FFF2CC")
_TON_FILL = PatternFill("solid", fgColor="E2EFDA")
_SHARED_HINT_FILL = PatternFill("solid", fgColor="FCE4D6")
_MUTE_FILL = PatternFill("solid", fgColor="F4F4F4")
_MUTE_HEADER_FILL = PatternFill("solid", fgColor="C8C8C8")
_MUTE_SUB_HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
_UNLOCKED = Protection(locked=False)
_LOCKED = Protection(locked=True)
_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_FONT_NAME = "微软雅黑"

WIDTHS = [
    6,
    16,
    8,
    18,
    14,
    10,
    12,
    SKETCH_COL_WIDTH,
    14,
    10,
    10,
    12,
    12,
    10,
    12,
    12,
    14,
    10,
    8,
    14,
    10,
    10,
    10,
    10,
    16,
    10,
    10,
    12,
]


STEEL_COL_COUNT = 22
STEEL_DATA_START_ROW = 6
STEEL_SHORT_CODE_COL = 7
STEEL_QTY_COL = 21
STEEL_TON_COL = 22
STEEL_WEIGHT_COL = 15
STEEL_SHARED_COLS = (19, 20)
STEEL_MUTE_COLS = frozenset({1, 2, 3, 4, 5})
STEEL_DENSITY_CELL = "$U$2"
STEEL_SHEET_TITLE = "钢材测算"
STEEL_WIDTHS = [
    5,
    6,
    8,
    5,
    10,
    16,
    8,
    18,
    14,
    14,
    8,
    10,
    10,
    10,
    12,
    10,
    8,
    10,
    8,
    14,
    12,
    14,
]


class StampingPartExportService:
    """冲压零件导出：工艺确认表，以及钢材采购测算。"""

    def __init__(self) -> None:
        self.repository = StampingPartRepository()
        self._center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self._left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        self._right = Alignment(horizontal="right", vertical="center")
        self._header_font = Font(name=_FONT_NAME, size=10, bold=True, color="FFFFFF")
        self._title_font = Font(name=_FONT_NAME, size=14, bold=True, color="FFFFFF")
        self._meta_font = Font(name=_FONT_NAME, size=10, color="1F4E79")
        self._data_font = Font(name=_FONT_NAME, size=9, color="000000")
        self._mute_header_font = Font(name=_FONT_NAME, size=9, color="6B6B6B")
        self._mute_data_font = Font(name=_FONT_NAME, size=8, color="8A8A8A")
        self._short_code_font = Font(name=_FONT_NAME, size=10, bold=True, color="1F4E79")

    def export(
        self,
        db: Session,
        *,
        location: str,
        stamping_supplier: str | None = None,
        vehicle_model: str | None = None,
        version: str | None = None,
        part_number: str | None = None,
        current_user: UserVO,
    ) -> tuple[bytes, str]:
        location, stamping_supplier, vehicle_model, version, part_number, parts = (
            self._parts_for_export(
                db,
                location=location,
                stamping_supplier=stamping_supplier,
                vehicle_model=vehicle_model,
                version=version,
                part_number=part_number,
            )
        )

        wb = Workbook()
        ws = wb.worksheets[0]
        ws.title = "冲压工艺确认表"

        title_project = f"{vehicle_model}车身件项目" if vehicle_model else location
        title_mill = f"-{stamping_supplier}" if stamping_supplier else ""
        title = f"{title_project}冲压工艺确认表{title_mill}"
        self._write_banner(
            ws,
            title=title,
            location=location,
            stamping_supplier=stamping_supplier,
            vehicle_model=vehicle_model,
            version=version,
            part_count=len(parts),
        )
        self._write_headers(ws)

        row_idx = DATA_START_ROW
        process_count = 0
        blank_count = 0
        sketch_count = 0
        sketch_buffers: list[BytesIO] = []
        for seq_no, part in enumerate(parts, start=1):
            processes = self._sorted_processes(part)
            blanks = self._sorted_blanks(part)
            process_count += len(processes)
            blank_count += len(blanks)
            used, has_sketch = self._write_part_block(
                ws,
                start_row=row_idx,
                seq_no=seq_no,
                part=part,
                processes=processes,
                blanks=blanks,
                zebra=seq_no % 2 == 0,
                sketch_buffers=sketch_buffers,
            )
            if has_sketch:
                sketch_count += 1
            row_idx += used

        for idx, width in enumerate(WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = f"A{DATA_START_ROW}"
        self._apply_print_chrome(ws)

        try:
            content = workbook_to_xlsx_bytes(wb, current_user)
        finally:
            sketch_buffers.clear()

        filename = self._export_filename(
            "冲压工艺确认表",
            location,
            stamping_supplier,
            vehicle_model,
            version,
        )
        logger.info(
            "导出冲压工艺 %s 件 / %s 道工序 / %s 片子片 / %s 张简图 / 文件 %s",
            len(parts),
            process_count,
            blank_count,
            sketch_count,
            filename,
        )
        return content, filename

    def export_steel_purchase(
        self,
        db: Session,
        *,
        location: str,
        stamping_supplier: str | None = None,
        vehicle_model: str | None = None,
        version: str | None = None,
        part_number: str | None = None,
        current_user: UserVO,
    ) -> tuple[bytes, str]:
        location, stamping_supplier, vehicle_model, version, part_number, parts = (
            self._parts_for_export(
                db,
                location=location,
                stamping_supplier=stamping_supplier,
                vehicle_model=vehicle_model,
                version=version,
                part_number=part_number,
            )
        )
        density = SettingsService().get_steel_density(db)
        workbook = self._build_steel_purchase_workbook(
            parts,
            density=density,
            location=location,
            stamping_supplier=stamping_supplier,
            vehicle_model=vehicle_model,
            version=version,
        )
        filename = self._export_filename(
            "钢材采购测算",
            location,
            stamping_supplier,
            vehicle_model,
            version,
        )
        logger.info(
            "导出钢材采购测算 %s 件 / 文件 %s / 操作人 %s",
            len(parts),
            filename,
            current_user.username,
        )
        return self._workbook_to_plain_bytes(workbook), filename

    def _build_steel_purchase_workbook(
        self,
        parts: list[StampingPart],
        *,
        density: Any,
        location: str,
        stamping_supplier: str | None,
        vehicle_model: str | None,
        version: str | None,
    ) -> Workbook:
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = STEEL_SHEET_TITLE
        title_project = f"{vehicle_model}车身件" if vehicle_model else location
        title_mill = f"-{stamping_supplier}" if stamping_supplier else ""
        title = f"{title_project}钢材采购测算{title_mill}"
        row_count = self._steel_purchase_row_count(parts)
        self._write_steel_purchase_banner(
            sheet,
            title=title,
            location=location,
            stamping_supplier=stamping_supplier,
            vehicle_model=vehicle_model,
            version=version,
            part_count=row_count,
            density=density,
        )
        self._write_steel_purchase_headers(sheet)
        excel_row = STEEL_DATA_START_ROW
        for index, part in enumerate(parts):
            excel_row = self._write_steel_purchase_part(
                sheet,
                start_row=excel_row,
                part=part,
                zebra=index % 2 == 0,
            )
        last_data_row = excel_row - 1
        if last_data_row >= STEEL_DATA_START_ROW:
            self._write_steel_purchase_total(sheet, last_data_row)
            sheet.auto_filter.ref = f"A5:{get_column_letter(STEEL_COL_COUNT)}{last_data_row}"
        else:
            sheet.auto_filter.ref = f"A5:{get_column_letter(STEEL_COL_COUNT)}5"
        for idx, width in enumerate(STEEL_WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(idx)].width = width
        sheet.freeze_panes = f"{get_column_letter(STEEL_SHORT_CODE_COL + 1)}{STEEL_DATA_START_ROW}"
        self._apply_print_chrome(sheet)
        workbook.defined_names.add(
            DefinedName(
                name="SteelDensity",
                attr_text=f"'{STEEL_SHEET_TITLE}'!{STEEL_DENSITY_CELL}",
            )
        )
        self._protect_steel_purchase_sheet(sheet)
        self._write_steel_purchase_help(workbook)
        return workbook

    def _write_banner(
        self,
        ws: Any,
        *,
        title: str,
        location: str,
        stamping_supplier: str | None,
        vehicle_model: str | None,
        version: str | None,
        part_count: int,
    ) -> None:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COL_COUNT)
        c1 = ws.cell(row=1, column=1, value=title)
        c1.font = self._title_font
        c1.fill = _TITLE_FILL
        c1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 28

        scope_parts = [f"基地：{location}"]
        if stamping_supplier:
            scope_parts.append(f"冲压工厂：{stamping_supplier}")
        if vehicle_model:
            scope_parts.append(f"项目：{vehicle_model}")
        if version:
            scope_parts.append(f"版本：{version}")
        scope_text = "    ".join(scope_parts)

        count_start = 22
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=count_start - 1)
        c2 = ws.cell(row=2, column=1, value=scope_text)
        c2.font = self._meta_font
        c2.fill = _META_FILL
        c2.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws.merge_cells(start_row=2, start_column=count_start, end_row=2, end_column=COL_COUNT)
        c3 = ws.cell(row=2, column=count_start, value=f"零件数量：{part_count}")
        c3.font = Font(name=_FONT_NAME, size=10, bold=True, color="1F4E79")
        c3.fill = _META_FILL
        c3.alignment = Alignment(horizontal="right", vertical="center", indent=1)

        for col in range(1, COL_COUNT + 1):
            cell = ws.cell(row=2, column=col)
            cell.fill = _META_FILL
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[3].height = 8

    def _write_headers(self, ws: Any) -> None:
        row4 = {
            1: "序号",
            2: "零件号",
            3: "短码",
            4: "零件名称",
            5: "冲压工厂",
            6: "项目",
            7: "版本",
            8: "零件简图",
            9: "材质",
            10: "料厚",
            11: "零件重量",
            12: "料片尺寸（mm）",
            14: "可制件数",
            15: "料片式样",
            16: "是否利废",
            17: "利用哪类废料",
            18: "是否摆剪",
            19: "工艺",
            23: "使用冲压设备吨位",
            24: "模具类型",
            25: "工艺方法",
            26: "自动化",
            27: "备注",
            28: "费用比例",
        }
        row5 = {
            12: "长/卷料宽",
            13: "宽/卷料步距",
            19: "工序名称",
            20: "工序",
            21: "套数",
            22: "模具编号",
        }
        merges = [
            (4, 1, 5, 1),
            (4, 2, 5, 2),
            (4, 3, 5, 3),
            (4, 4, 5, 4),
            (4, 5, 5, 5),
            (4, 6, 5, 6),
            (4, 7, 5, 7),
            (4, 8, 5, 8),
            (4, 9, 5, 9),
            (4, 10, 5, 10),
            (4, 11, 5, 11),
            (4, 12, 4, 13),
            (4, 14, 5, 14),
            (4, 15, 5, 15),
            (4, 16, 5, 16),
            (4, 17, 5, 17),
            (4, 18, 5, 18),
            (4, 19, 4, 22),
            (4, 23, 5, 23),
            (4, 24, 5, 24),
            (4, 25, 5, 25),
            (4, 26, 5, 26),
            (4, 27, 5, 27),
            (4, 28, 5, 28),
        ]
        for r1, c1, r2, c2 in merges:
            if (r1, c1) != (r2, c2):
                ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

        for col in range(1, COL_COUNT + 1):
            for row in (4, 5):
                cell = ws.cell(row=row, column=col)
                cell.font = self._header_font
                cell.fill = _HEADER_FILL if row == 4 else _SUB_HEADER_FILL
                cell.alignment = self._center
                cell.border = _THIN
        for col, value in row4.items():
            ws.cell(row=4, column=col, value=value)
        for col, value in row5.items():
            ws.cell(row=5, column=col, value=value)
        ws.row_dimensions[4].height = 22
        ws.row_dimensions[5].height = 22

    def _normalize_export_scope(
        self,
        *,
        location: str,
        stamping_supplier: str | None,
        vehicle_model: str | None,
        version: str | None,
        part_number: str | None,
    ) -> tuple[str, str | None, str | None, str | None, str | None]:
        location = location.strip()
        stamping_supplier = (stamping_supplier or "").strip() or None
        vehicle_model = (vehicle_model or "").strip() or None
        version = (version or "").strip() or None
        part_number = (part_number or "").strip() or None
        if not location:
            raise BusinessException("导出须选定基地")
        return location, stamping_supplier, vehicle_model, version, part_number

    def _parts_for_export(
        self,
        db: Session,
        *,
        location: str,
        stamping_supplier: str | None,
        vehicle_model: str | None,
        version: str | None,
        part_number: str | None,
    ) -> tuple[str, str | None, str | None, str | None, str | None, list[StampingPart]]:
        location, stamping_supplier, vehicle_model, version, part_number = (
            self._normalize_export_scope(
                location=location,
                stamping_supplier=stamping_supplier,
                vehicle_model=vehicle_model,
                version=version,
                part_number=part_number,
            )
        )
        parts = self.repository.list_for_export(
            db,
            location=location,
            stamping_supplier=stamping_supplier,
            vehicle_model=vehicle_model,
            version=version,
            part_number=part_number,
        )
        return location, stamping_supplier, vehicle_model, version, part_number, parts

    def _sorted_processes(self, part: StampingPart) -> list[StampingProcess]:
        return sorted(
            list(part.processes or []),
            key=lambda item: (item.process_order or 0, item.id or 0),
        )

    def _sorted_blanks(self, part: StampingPart) -> list[StampingTailorWeldBlank]:
        return sorted(
            list(part.tailor_weld_blanks or []),
            key=lambda item: (item.blank_seq or 0, item.id or 0),
        )

    def _apply_print_chrome(self, ws: Any) -> None:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90

    def _export_filename(
        self,
        prefix: str,
        location: str,
        stamping_supplier: str | None,
        vehicle_model: str | None,
        version: str | None,
    ) -> str:
        scope = "_".join(
            part for part in (location, stamping_supplier, vehicle_model, version) if part
        )
        stamp = now().strftime("%Y%m%d_%H%M%S")
        return f"{prefix}_{scope}_{stamp}.xlsx"

    def _workbook_to_plain_bytes(self, workbook: Workbook) -> bytes:
        """测算簿不涉密，不加内控水印。工艺确认表仍走 workbook_to_xlsx_bytes。"""
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _steel_purchase_row_count(self, parts: list[StampingPart]) -> int:
        total = 0
        for part in parts:
            blanks = self._sorted_blanks(part)
            if part.is_tailor_welded and blanks:
                total += len(blanks)
            else:
                total += 1
        return total

    def _write_steel_purchase_banner(
        self,
        ws: Any,
        *,
        title: str,
        location: str,
        stamping_supplier: str | None,
        vehicle_model: str | None,
        version: str | None,
        part_count: int,
        density: Any,
    ) -> None:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=STEEL_COL_COUNT)
        title_cell = ws.cell(row=1, column=1, value=title)
        title_cell.font = self._title_font
        title_cell.fill = _TITLE_FILL
        title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 28
        for col in range(1, STEEL_COL_COUNT + 1):
            ws.cell(row=1, column=col).fill = _TITLE_FILL
            ws.cell(row=1, column=col).border = _THIN

        scope_parts = [f"基地：{location}"]
        if stamping_supplier:
            scope_parts.append(f"冲压工厂：{stamping_supplier}")
        if vehicle_model:
            scope_parts.append(f"项目：{vehicle_model}")
        if version:
            scope_parts.append(f"版本：{version}")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=18)
        scope_cell = ws.cell(row=2, column=1, value="    ".join(scope_parts))
        scope_cell.font = self._meta_font
        scope_cell.fill = _META_FILL
        scope_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in range(1, STEEL_COL_COUNT + 1):
            cell = ws.cell(row=2, column=col)
            cell.fill = _META_FILL
            cell.border = _THIN
        count_cell = ws.cell(row=2, column=19, value=f"行数：{part_count}")
        count_cell.font = Font(name=_FONT_NAME, size=10, bold=True, color="1F4E79")
        count_cell.alignment = self._right
        ws.cell(row=2, column=20, value="密度(g/cm³)").font = self._meta_font
        ws.cell(row=2, column=20).alignment = self._right
        density_cell = ws.cell(row=2, column=STEEL_QTY_COL, value=_num(density, Q4))
        density_cell.font = Font(name=_FONT_NAME, size=10, bold=True, color="1F4E79")
        density_cell.alignment = self._right
        density_cell.protection = _UNLOCKED
        density_cell.fill = _QTY_FILL
        ws.cell(row=2, column=STEEL_TON_COL, value=now().strftime("%Y-%m-%d %H:%M"))
        ws.cell(row=2, column=STEEL_TON_COL).font = self._meta_font
        ws.cell(row=2, column=STEEL_TON_COL).alignment = self._right
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[3].height = 8

    def _write_steel_purchase_headers(self, ws: Any) -> None:
        groups = [
            (1, 5, "范围"),
            (6, 9, "零件"),
            (10, 16, "料片"),
            (17, 18, "标记"),
            (19, 20, "共模"),
            (21, 22, "测算"),
        ]
        fields = {
            1: "序号",
            2: "基地",
            3: "项目",
            4: "版本",
            5: "冲压工厂",
            6: "零件号",
            7: "短码",
            8: "零件名称",
            9: "子片",
            10: "牌号",
            11: "厚度",
            12: "长/卷料宽",
            13: "宽/卷料步距",
            14: "可制件数",
            15: "材料重",
            16: "零件重量",
            17: "是否利废",
            18: "废料副产品",
            19: "是否共模",
            20: "共模标记",
            21: "主机厂数量",
            22: "采购钢材(吨)",
        }
        for start, end, label in groups:
            is_mute = start in STEEL_MUTE_COLS
            header_fill = _MUTE_HEADER_FILL if is_mute else _HEADER_FILL
            header_font = self._mute_header_font if is_mute else self._header_font
            if start != end:
                ws.merge_cells(start_row=4, start_column=start, end_row=4, end_column=end)
            cell = ws.cell(row=4, column=start, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = self._center
            cell.border = _THIN
            for col in range(start, end + 1):
                ws.cell(row=4, column=col).fill = header_fill
                ws.cell(row=4, column=col).border = _THIN
                ws.cell(row=4, column=col).font = header_font
        for col, value in fields.items():
            is_mute = col in STEEL_MUTE_COLS
            cell = ws.cell(row=5, column=col, value=value)
            cell.font = self._mute_header_font if is_mute else self._header_font
            cell.fill = _MUTE_SUB_HEADER_FILL if is_mute else _SUB_HEADER_FILL
            cell.alignment = self._center
            cell.border = _THIN
        ws.row_dimensions[4].height = 22
        ws.row_dimensions[5].height = 22

    def _write_steel_purchase_part(
        self,
        ws: Any,
        *,
        start_row: int,
        part: StampingPart,
        zebra: bool,
    ) -> int:
        processes = self._sorted_processes(part)
        blanks = self._sorted_blanks(part)
        if part.is_tailor_welded and blanks:
            row = start_row
            for blank in blanks:
                self._write_steel_purchase_row(
                    ws,
                    excel_row=row,
                    part=part,
                    processes=processes,
                    blank=blank,
                    zebra=zebra,
                )
                row += 1
            return row
        self._write_steel_purchase_row(
            ws,
            excel_row=start_row,
            part=part,
            processes=processes,
            blank=None,
            zebra=zebra,
        )
        return start_row + 1

    def _write_steel_purchase_row(
        self,
        ws: Any,
        *,
        excel_row: int,
        part: StampingPart,
        processes: list[StampingProcess],
        blank: StampingTailorWeldBlank | None,
        zebra: bool,
    ) -> None:
        material = (blank.material if blank and blank.material else part.material) or ""
        thickness = blank.thickness if blank and blank.thickness is not None else part.thickness
        length = (
            blank.sheet_length if blank and blank.sheet_length is not None else part.sheet_length
        )
        width = blank.sheet_width if blank and blank.sheet_width is not None else part.sheet_width
        parts_per_sheet = (
            blank.parts_per_sheet
            if blank and blank.parts_per_sheet is not None
            else part.parts_per_sheet
        )
        part_weight = (
            blank.part_weight if blank and blank.part_weight is not None else part.part_weight
        )
        process_codes = [item.process_code for item in processes]
        cost_ratios = [item.cost_ratio for item in processes]
        labels = shared_die_labels(
            die_type=part.die_type,
            process_method=part.process_method,
            process_codes=process_codes,
            cost_ratios=cost_ratios,
        )
        is_shared = bool(labels) or has_shared_die_ratio(cost_ratios)
        values: dict[int, Any] = {
            1: part.seq_no,
            2: part.location,
            3: part.vehicle_model,
            4: part.version,
            5: part.stamping_supplier or "",
            6: part.part_number,
            7: short_part_number_digits(part.part_number) or "",
            8: part.part_name or "",
            9: blank.blank_part_number if blank else "",
            10: material,
            11: _num(thickness, Q4),
            12: _num(length, Q2),
            13: _num(width, Q2),
            14: parts_per_sheet,
            15: (
                f'=IF(OR(L{excel_row}="",M{excel_row}="",K{excel_row}=""),"",'
                f"L{excel_row}*M{excel_row}*K{excel_row}*SteelDensity/1000000)"
            ),
            16: _num(part_weight, Q4),
            17: "是" if part.is_offal_nested else "否",
            18: "是" if part.is_scrap_byproduct else "否",
            19: "是" if is_shared else "否",
            20: "、".join(labels),
            21: None,
            22: (
                f'=IF(OR(Q{excel_row}="是",R{excel_row}="是"),0,'
                f'IF(OR(O{excel_row}="",N{excel_row}="",N{excel_row}<=0,U{excel_row}=""),"",'
                f"O{excel_row}/N{excel_row}*U{excel_row}/1000))"
            ),
        }
        zebra_fill = _ZEBRA if zebra else None
        ws.row_dimensions[excel_row].height = DEFAULT_ROW_HEIGHT
        for col in range(1, STEEL_COL_COUNT + 1):
            cell = ws.cell(row=excel_row, column=col, value=values.get(col))
            cell.font = self._mute_data_font if col in STEEL_MUTE_COLS else self._data_font
            cell.border = _THIN
            cell.protection = _LOCKED
            if col in (6, 8, 9, 10, 20):
                cell.alignment = self._left
            elif col in (11, 12, 13, 14, 15, 16, 22):
                cell.alignment = self._right
            else:
                cell.alignment = self._center
            if col == STEEL_SHORT_CODE_COL:
                cell.font = self._short_code_font
            if col == STEEL_QTY_COL:
                cell.fill = _QTY_FILL
                cell.protection = _UNLOCKED
                cell.comment = Comment("在此填写", "系统")
            elif col == STEEL_TON_COL:
                cell.fill = _TON_FILL
                cell.number_format = "0.0000"
            elif col == STEEL_WEIGHT_COL:
                cell.number_format = "0.0000"
                if zebra_fill is not None:
                    cell.fill = zebra_fill
            elif is_shared and col in STEEL_SHARED_COLS:
                cell.fill = _SHARED_HINT_FILL
            elif col in STEEL_MUTE_COLS:
                cell.fill = _MUTE_FILL
            elif zebra_fill is not None:
                cell.fill = zebra_fill

    def _write_steel_purchase_total(self, ws: Any, last_data_row: int) -> None:
        total_row = last_data_row + 1
        ws.merge_cells(
            start_row=total_row,
            start_column=1,
            end_row=total_row,
            end_column=STEEL_TON_COL - 1,
        )
        for col in range(1, STEEL_COL_COUNT + 1):
            cell = ws.cell(row=total_row, column=col)
            cell.fill = _TITLE_FILL
            cell.font = self._header_font
            cell.border = _THIN
            cell.protection = _LOCKED
        label = ws.cell(row=total_row, column=1, value="合计")
        label.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        ton = ws.cell(
            row=total_row,
            column=STEEL_TON_COL,
            value=f"=SUM({get_column_letter(STEEL_TON_COL)}{STEEL_DATA_START_ROW}:"
            f"{get_column_letter(STEEL_TON_COL)}{last_data_row})",
        )
        ton.number_format = "0.0000"
        ton.alignment = self._right
        ws.row_dimensions[total_row].height = 20

    def _protect_steel_purchase_sheet(self, ws: Any) -> None:
        ws.protection.sheet = True

    def _write_steel_purchase_help(self, workbook: Workbook) -> None:
        sheet = workbook.create_sheet("填写说明")
        sheet["A1"] = "填写说明"
        sheet["A1"].font = self._title_font
        sheet["A1"].fill = _TITLE_FILL
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        for col in (1, 2):
            sheet.cell(row=1, column=col).fill = _TITLE_FILL
            sheet.cell(row=1, column=col).border = _THIN
        headers = ("项", "说明")
        for col, text in enumerate(headers, start=1):
            cell = sheet.cell(row=2, column=col, value=text)
            cell.font = self._header_font
            cell.fill = _SUB_HEADER_FILL
            cell.alignment = self._center
            cell.border = _THIN
        rows = [
            (
                "主机厂数量",
                "只填黄色列。材料重、采购吨由公式算出。取消保护：审阅 → 撤销工作表保护（无密码）。",
            ),
            ("共模", "是否共模为「是」时，成对件数量应相同。"),
            ("利废 / 废料副产品", "采购钢材(吨)为 0，不买新板。"),
            ("可制件数", "一张料能出几件成品。吨 = 材料重 ÷ 可制件数 × 数量 ÷ 1000。"),
            ("密度", "表头黄色格可改；改后材料重跟着变。"),
            ("拼焊", "子片分行，各行填同一总成件数。"),
        ]
        for index, (item, detail) in enumerate(rows, start=3):
            item_cell = sheet.cell(row=index, column=1, value=item)
            detail_cell = sheet.cell(row=index, column=2, value=detail)
            for cell in (item_cell, detail_cell):
                cell.font = self._data_font
                cell.border = _THIN
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            if index % 2 == 0:
                item_cell.fill = _ZEBRA
                detail_cell.fill = _ZEBRA
            sheet.row_dimensions[index].height = 32
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 72
        sheet.row_dimensions[1].height = 28
        self._apply_print_chrome(sheet)

    def _write_part_block(
        self,
        ws: Any,
        *,
        start_row: int,
        seq_no: int,
        part: StampingPart,
        processes: list[StampingProcess],
        blanks: list[StampingTailorWeldBlank],
        zebra: bool,
        sketch_buffers: list[BytesIO],
    ) -> tuple[int, bool]:
        n_rows = max(len(processes), len(blanks), 1)
        end_row = start_row + n_rows - 1
        fill = _ZEBRA if zebra else None

        for offset in range(n_rows):
            excel_row = start_row + offset
            ws.row_dimensions[excel_row].height = DEFAULT_ROW_HEIGHT
            for col in range(1, COL_COUNT + 1):
                cell = ws.cell(row=excel_row, column=col, value=None)
                cell.font = self._data_font
                cell.border = _THIN
                cell.alignment = self._center
                if fill is not None:
                    cell.fill = fill

        part_cols: tuple[int, ...] = (
            1,
            2,
            3,
            4,
            5,
            6,
            7,
            8,
            9,
            11,
            14,
            15,
            16,
            17,
            18,
            24,
            25,
            26,
            27,
        )
        if not blanks:
            part_cols = (*part_cols, 10, 12, 13)
        for col in part_cols:
            _merge_if_needed(ws, start_row, end_row, col)

        ws.cell(row=start_row, column=1, value=seq_no).alignment = self._center
        ws.cell(row=start_row, column=2, value=part.part_number).alignment = self._left
        short_cell = ws.cell(
            row=start_row,
            column=SHORT_CODE_COL,
            value=short_part_number_digits(part.part_number) or "",
        )
        short_cell.alignment = self._center
        short_cell.font = self._short_code_font
        ws.cell(row=start_row, column=4, value=part.part_name or "").alignment = self._left
        ws.cell(row=start_row, column=5, value=part.stamping_supplier or "").alignment = self._left
        ws.cell(row=start_row, column=6, value=part.vehicle_model or "").alignment = self._center
        ws.cell(row=start_row, column=7, value=part.version or "").alignment = self._center
        ws.cell(row=start_row, column=9, value=part.material or "")
        ws.cell(row=start_row, column=11, value=_num(part.part_weight, Q4)).alignment = self._right
        ws.cell(row=start_row, column=14, value=part.parts_per_sheet)
        ws.cell(row=start_row, column=15, value=part.sheet_type or "")
        ws.cell(row=start_row, column=16, value="是" if part.is_offal_nested else "否")
        ws.cell(row=start_row, column=17, value=part.offal_source or "")
        ws.cell(row=start_row, column=18, value="是" if part.has_swing_shear else "否")
        ws.cell(row=start_row, column=24, value=part.die_type or "")
        ws.cell(row=start_row, column=25, value=part.process_method or "")
        ws.cell(row=start_row, column=26, value=part.automation or "")
        ws.cell(row=start_row, column=27, value=part.remark or "").alignment = self._left

        if blanks:
            spans = _split_spans(n_rows, len(blanks))
            for blank, (rel_start, rel_end) in zip(blanks, spans, strict=False):
                abs_start = start_row + rel_start
                abs_end = start_row + rel_end
                for col in (10, 12, 13):
                    _merge_if_needed(ws, abs_start, abs_end, col)
                thickness = blank.thickness if blank.thickness is not None else part.thickness
                length = blank.sheet_length if blank.sheet_length is not None else part.sheet_length
                width = blank.sheet_width if blank.sheet_width is not None else part.sheet_width
                ws.cell(row=abs_start, column=10, value=_num(thickness, Q4)).alignment = self._right
                ws.cell(row=abs_start, column=12, value=_num(length, Q2)).alignment = self._right
                ws.cell(row=abs_start, column=13, value=_num(width, Q2)).alignment = self._right
        else:
            ws.cell(
                row=start_row, column=10, value=_num(part.thickness, Q4)
            ).alignment = self._right
            ws.cell(
                row=start_row, column=12, value=_num(part.sheet_length, Q2)
            ).alignment = self._right
            ws.cell(
                row=start_row, column=13, value=_num(part.sheet_width, Q2)
            ).alignment = self._right

        for offset, process in enumerate(processes):
            excel_row = start_row + offset
            ws.cell(row=excel_row, column=19, value=process.process_seq)
            ws.cell(row=excel_row, column=20, value=process.process_code or "")
            ws.cell(row=excel_row, column=21, value=process.die_sets)
            ws.cell(row=excel_row, column=22, value=process.die_number or "")
            ws.cell(row=excel_row, column=23, value=process.tonnage)
            ratio_cell = ws.cell(
                row=excel_row,
                column=28,
                value=_num(process.cost_ratio, Q2),
            )
            ratio_cell.alignment = self._right

        has_sketch = self._embed_sketch(
            ws,
            part=part,
            start_row=start_row,
            n_rows=n_rows,
            sketch_buffers=sketch_buffers,
        )
        return n_rows, has_sketch

    def _embed_sketch(
        self,
        ws: Any,
        *,
        part: StampingPart,
        start_row: int,
        n_rows: int,
        sketch_buffers: list[BytesIO],
    ) -> bool:
        """把零件简图锚定到短码后的简图列合并格左上角，按格缩放；缺文件则跳过。"""
        sketch_path = _resolve_sketch_path(part)
        if sketch_path is None:
            return False
        try:
            payload = BytesIO(sketch_path.read_bytes())
            xl_img = XlImage(payload)
            payload.seek(0)
            sketch_buffers.append(payload)
        except Exception as exc:
            logger.warning("读取零件简图失败 %s: %s", part.part_number, exc)
            return False

        orig_w = float(xl_img.width or 0)
        orig_h = float(xl_img.height or 0)
        if orig_w <= 0 or orig_h <= 0:
            return False

        available_h = max(
            SKETCH_MIN_HEIGHT_PX,
            n_rows * DEFAULT_ROW_HEIGHT / _PX_TO_PT - SKETCH_PAD_PT,
        )
        scale = min(SKETCH_MAX_WIDTH_PX / orig_w, available_h / orig_h, 1.0)
        xl_img.width = max(1, int(orig_w * scale))
        xl_img.height = max(1, int(orig_h * scale))

        needed_pt = xl_img.height * _PX_TO_PT + SKETCH_PAD_PT
        current_pt = n_rows * DEFAULT_ROW_HEIGHT
        if needed_pt > current_pt:
            each_pt = needed_pt / n_rows
            for offset in range(n_rows):
                ws.row_dimensions[start_row + offset].height = each_pt

        xl_img.anchor = f"{get_column_letter(SKETCH_COL)}{start_row}"
        ws.add_image(xl_img)
        return True


def _resolve_sketch_path(part: StampingPart) -> Path | None:
    """按上传约定取简图：UPLOAD_DIR/{基地}/{车型}/{版本}/{零件号}.png。"""
    location = (part.location or "").strip()
    vehicle_model = (part.vehicle_model or "").strip()
    version = (part.version or "").strip()
    part_number = (part.part_number or "").strip()
    if not location or not vehicle_model or not version or not part_number:
        return None
    safe_location = location.replace("/", "_").replace("\\", "_")
    safe_vehicle = vehicle_model.replace("/", "_").replace("\\", "_")
    safe_version = version.replace("/", "_").replace("\\", "_").replace(" ", "_")
    safe_name = "".join(ch if ch.isalnum() or ch in ("_", "-") else "_" for ch in part_number)
    if not safe_name:
        return None
    candidate = (
        settings.UPLOAD_DIR / safe_location / safe_vehicle / safe_version / f"{safe_name}.png"
    )
    if candidate.is_file():
        return candidate
    return None


def _num(value: Any, quant: Any) -> Any:
    return excel_money_cell(value, quant)


def _merge_if_needed(ws: Any, start_row: int, end_row: int, col: int) -> None:
    if end_row > start_row:
        ws.merge_cells(
            start_row=start_row,
            start_column=col,
            end_row=end_row,
            end_column=col,
        )


def _split_spans(n_rows: int, n_blanks: int) -> list[tuple[int, int]]:
    """把零件占用行按子片数切成连续块（0-based 闭区间）。"""
    if n_blanks <= 0:
        return []
    base, extra = divmod(n_rows, n_blanks)
    spans: list[tuple[int, int]] = []
    start = 0
    for index in range(n_blanks):
        size = max(base + (1 if index < extra else 0), 1)
        end = min(start + size - 1, n_rows - 1)
        spans.append((start, end))
        start = end + 1
        if start >= n_rows:
            while len(spans) < n_blanks:
                spans.append((n_rows - 1, n_rows - 1))
            break
    return spans
