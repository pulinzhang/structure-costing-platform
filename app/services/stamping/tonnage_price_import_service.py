"""冲次单价 Excel 导入 / 模板下载服务。"""

from __future__ import annotations

import logging
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from app.common.audit import audit_create, audit_update
from app.config import settings
from app.exceptions.base import BusinessException
from app.parsers.stamping.tonnage_price_parser import parse_tonnage_price_excel
from app.repositories.stamping.tonnage_price_repository import StampingTonnagePriceRepository
from app.schemas.dto.stamping_tonnage_price_dto import (
    StampingTonnagePriceCreate,
    StampingTonnagePriceUpdate,
    quantize_price_per_stroke,
)
from app.schemas.vo.stamping_tonnage_price_vo import StampingTonnagePriceImportResultVO
from app.schemas.vo.user_vo import UserVO
from app.services.file_trace.path_utils import to_upload_relpath
from app.services.file_trace.register import register_existing
from app.utils.stamping_tonnage_price_archive import save_tonnage_price_upload

logger = logging.getLogger("service.stamping_tonnage_price_import")

TEMPLATE_SAMPLE_ROWS: list[tuple[str, str, Decimal]] = [
    ("400T", "次", Decimal("0.564")),
    ("500T", "次", Decimal("0.622")),
    ("630T", "次", Decimal("0.713")),
    ("800T", "次", Decimal("1.026")),
    ("1000T", "次", Decimal("1.283")),
    ("1200T", "次", Decimal("1.539")),
]


class StampingTonnagePriceImportService:
    """冲次单价同步导入（数据量小，不走异步任务）。"""

    def __init__(self) -> None:
        self.repository = StampingTonnagePriceRepository()

    def build_template(self) -> tuple[bytes, str]:
        """生成三列模板：名称 / 单位 / 单价。"""
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = "冲次单价"

        header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        data_font = Font(name="微软雅黑", size=9)
        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )

        headers = ["名称", "单位", "单价"]
        for col, title in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for row_idx, (name, unit, price) in enumerate(TEMPLATE_SAMPLE_ROWS, start=2):
            values: list[str | Decimal] = [name, unit, price]
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_idx, column=col, value=value)
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal="center" if col < 3 else "right",
                    vertical="center",
                )
                if col == 3:
                    cell.number_format = "0.000"

        sheet.column_dimensions["A"].width = 12
        sheet.column_dimensions["B"].width = 10
        sheet.column_dimensions["C"].width = 12
        sheet.freeze_panes = "A2"

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue(), "冲次单价导入模板.xlsx"

    def import_excel(
        self,
        db: Session,
        *,
        file_content: bytes,
        filename: str,
        location: str,
        year: int,
        current_user: UserVO,
    ) -> StampingTonnagePriceImportResultVO:
        location = (location or "").strip()
        if not location:
            raise BusinessException("请选择采购基地")
        if year < 2000 or year > 2100:
            raise BusinessException("年度不正确")

        original_name = Path(filename or "").name
        if Path(original_name).suffix.lower() != ".xlsx":
            raise BusinessException("仅支持 .xlsx 文件")
        if not file_content:
            raise BusinessException("文件内容为空")
        if len(file_content) > settings.MAX_UPLOAD_SIZE:
            max_mb = settings.MAX_UPLOAD_SIZE // (1024 * 1024)
            raise BusinessException(f"文件过大，最大支持 {max_mb}MB")

        archive_path = save_tonnage_price_upload(
            file_content=file_content,
            filename=original_name,
            year=year,
            location=location,
        )
        archived_rel = to_upload_relpath(archive_path)

        parse_result = parse_tonnage_price_excel(file_content)
        if parse_result.errors and not parse_result.rows:
            raise BusinessException("；".join(parse_result.errors))
        if not parse_result.rows:
            raise BusinessException("未解析到有效数据行")

        created_count = 0
        updated_count = 0
        skipped_approved = 0
        create_audit = audit_create(current_user)
        update_audit = audit_update(current_user)

        try:
            for row in parse_result.rows:
                price = quantize_price_per_stroke(row.price_per_stroke)
                existing = self.repository.get_by_location_year_tonnage(
                    db, location, year, row.tonnage
                )
                if existing:
                    if existing.audit_status == "approved":
                        skipped_approved += 1
                        continue
                    self.repository.update(
                        db,
                        existing.id,
                        StampingTonnagePriceUpdate(price_per_stroke=price),
                        updated_by=update_audit["updated_by"],
                    )
                    updated_count += 1
                else:
                    self.repository.create(
                        db,
                        StampingTonnagePriceCreate(
                            location=location,
                            year=year,
                            tonnage=row.tonnage,
                            price_per_stroke=price,
                        ),
                        created_by=create_audit["created_by"],
                        updated_by=create_audit["updated_by"],
                    )
                    created_count += 1
            register_existing(
                db,
                file_category="stamping_tonnage_price",
                stored_relpath=archived_rel,
                original_filename=original_name,
                current_user=current_user,
                location=location,
                year=year,
            )
            db.commit()
        except Exception:
            db.rollback()
            raise

        warnings = list(parse_result.errors)
        if skipped_approved:
            warnings.append(f"跳过已审核 {skipped_approved} 条（未覆盖）")
        message = (
            f"导入完成：基地「{location}」{year} 年，"
            f"新增 {created_count} 条，更新 {updated_count} 条"
        )
        if skipped_approved:
            message += f"，跳过已审核 {skipped_approved} 条"
        if parse_result.errors:
            message += f"；有 {len(parse_result.errors)} 行解析跳过"

        logger.info(
            "%s file=%s archived=%s created=%s updated=%s skipped_approved=%s warnings=%s",
            message,
            original_name,
            archived_rel,
            created_count,
            updated_count,
            skipped_approved,
            len(warnings),
        )
        return StampingTonnagePriceImportResultVO(
            location=location,
            year=year,
            created_count=created_count,
            updated_count=updated_count,
            total_count=created_count + updated_count,
            skipped_count=len(parse_result.errors) + skipped_approved,
            warnings=warnings,
            message=message,
        )
