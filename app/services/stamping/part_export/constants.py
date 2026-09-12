"""工艺确认表 / 钢材测算列宽与配色。改列数先改 COL_COUNT / STEEL_COL_COUNT 与 WIDTHS。"""

from __future__ import annotations

from openpyxl.styles import Border, PatternFill, Protection, Side

COL_COUNT = 28
DATA_START_ROW = 6
SKETCH_COL = 8
SHORT_CODE_COL = 3
DEFAULT_ROW_HEIGHT = 18
# 简图列加宽，约 130px，给 PNG 留出可视区域
SKETCH_COL_WIDTH = 18
SKETCH_MAX_WIDTH_PX = 118
SKETCH_MIN_HEIGHT_PX = 88
SKETCH_PAD_PT = 8
_PX_TO_PT = 72 / 96

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_SUB_HEADER_FILL = PatternFill("solid", fgColor="2E75B6")
_TITLE_FILL = PatternFill("solid", fgColor="1F4E79")
_META_FILL = PatternFill("solid", fgColor="D6EAF8")
_ZEBRA = PatternFill("solid", fgColor="F2F2F2")
_QTY_FILL = PatternFill("solid", fgColor="FFF2CC")
_TON_FILL = PatternFill("solid", fgColor="E2EFDA")
_SHARED_HINT_FILL = PatternFill("solid", fgColor="FCE4D6")
_MUTE_FILL = PatternFill("solid", fgColor="F4F4F4")
_MUTE_HEADER_FILL = PatternFill("solid", fgColor="C8C8C8")
_MUTE_SUB_HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
_UNLOCKED = Protection(locked=False)
_LOCKED = Protection(locked=True)
_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_FONT_NAME = "微软雅黑"

WIDTHS = [
    6,
    16,
    8,
    18,
    14,
    10,
    12,
    SKETCH_COL_WIDTH,
    14,
    10,
    10,
    12,
    12,
    10,
    12,
    12,
    14,
    10,
    8,
    14,
    10,
    10,
    10,
    10,
    16,
    10,
    10,
    12,
]


STEEL_COL_COUNT = 22
STEEL_DATA_START_ROW = 6
STEEL_SHORT_CODE_COL = 7
STEEL_QTY_COL = 21
STEEL_TON_COL = 22
STEEL_WEIGHT_COL = 15
STEEL_SHARED_COLS = (19, 20)
STEEL_MUTE_COLS = frozenset({1, 2, 3, 4, 5})
STEEL_DENSITY_CELL = "$U$2"
STEEL_SHEET_TITLE = "钢材测算"
STEEL_WIDTHS = [
    5,
    6,
    8,
    5,
    10,
    16,
    8,
    18,
    14,
    14,
    8,
    10,
    10,
    10,
    12,
    10,
    8,
    10,
    8,
    14,
    12,
    14,
]
