"""工艺确认表 — 零件块（主表合并 + 工序/拼焊子片）。"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill

from app.common.bom.manufacturing_category import short_part_number_digits
from app.common.money import Q2, Q4
from app.models.stamping import StampingPart, StampingProcess, StampingTailorWeldBlank
from app.services.stamping.part_export.constants import (
    _THIN,
    _ZEBRA,
    COL_COUNT,
    DEFAULT_ROW_HEIGHT,
    SHORT_CODE_COL,
)
from app.services.stamping.part_export.helpers import _merge_if_needed, _num, _split_spans
from app.services.stamping.part_export.sketch import embed_sketch


def write_part_block(
    ws: Any,
    *,
    start_row: int,
    seq_no: int,
    part: StampingPart,
    processes: list[StampingProcess],
    blanks: list[StampingTailorWeldBlank],
    zebra: bool,
    sketch_buffers: list[BytesIO],
    data_font: Font,
    short_code_font: Font,
    center: Alignment,
    left: Alignment,
    right: Alignment,
) -> tuple[int, bool]:
    n_rows = max(len(processes), len(blanks), 1)
    end_row = start_row + n_rows - 1
    fill: PatternFill | None = _ZEBRA if zebra else None

    for offset in range(n_rows):
        excel_row = start_row + offset
        ws.row_dimensions[excel_row].height = DEFAULT_ROW_HEIGHT
        for col in range(1, COL_COUNT + 1):
            cell = ws.cell(row=excel_row, column=col, value=None)
            cell.font = data_font
            cell.border = _THIN
            cell.alignment = center
            if fill is not None:
                cell.fill = fill

    part_cols: tuple[int, ...] = (
        1,
        2,
        3,
        4,
        5,
        6,
        7,
        8,
        9,
        11,
        14,
        15,
        16,
        17,
        18,
        24,
        25,
        26,
        27,
    )
    if not blanks:
        part_cols = (*part_cols, 10, 12, 13)
    for col in part_cols:
        _merge_if_needed(ws, start_row, end_row, col)

    ws.cell(row=start_row, column=1, value=seq_no).alignment = center
    ws.cell(row=start_row, column=2, value=part.part_number).alignment = left
    short_cell = ws.cell(
        row=start_row,
        column=SHORT_CODE_COL,
        value=short_part_number_digits(part.part_number) or "",
    )
    short_cell.alignment = center
    short_cell.font = short_code_font
    ws.cell(row=start_row, column=4, value=part.part_name or "").alignment = left
    ws.cell(row=start_row, column=5, value=part.stamping_supplier or "").alignment = left
    ws.cell(row=start_row, column=6, value=part.vehicle_model or "").alignment = center
    ws.cell(row=start_row, column=7, value=part.version or "").alignment = center
    ws.cell(row=start_row, column=9, value=part.material or "")
    ws.cell(row=start_row, column=11, value=_num(part.part_weight, Q4)).alignment = right
    ws.cell(row=start_row, column=14, value=part.parts_per_sheet)
    ws.cell(row=start_row, column=15, value=part.sheet_type or "")
    ws.cell(row=start_row, column=16, value="是" if part.is_offal_nested else "否")
    ws.cell(row=start_row, column=17, value=part.offal_source or "")
    ws.cell(row=start_row, column=18, value="是" if part.has_swing_shear else "否")
    ws.cell(row=start_row, column=24, value=part.die_type or "")
    ws.cell(row=start_row, column=25, value=part.process_method or "")
    ws.cell(row=start_row, column=26, value=part.automation or "")
    ws.cell(row=start_row, column=27, value=part.remark or "").alignment = left

    if blanks:
        spans = _split_spans(n_rows, len(blanks))
        for blank, (rel_start, rel_end) in zip(blanks, spans, strict=False):
            abs_start = start_row + rel_start
            abs_end = start_row + rel_end
            for col in (10, 12, 13):
                _merge_if_needed(ws, abs_start, abs_end, col)
            thickness = blank.thickness if blank.thickness is not None else part.thickness
            length = blank.sheet_length if blank.sheet_length is not None else part.sheet_length
            width = blank.sheet_width if blank.sheet_width is not None else part.sheet_width
            ws.cell(row=abs_start, column=10, value=_num(thickness, Q4)).alignment = right
            ws.cell(row=abs_start, column=12, value=_num(length, Q2)).alignment = right
            ws.cell(row=abs_start, column=13, value=_num(width, Q2)).alignment = right
    else:
        ws.cell(row=start_row, column=10, value=_num(part.thickness, Q4)).alignment = right
        ws.cell(row=start_row, column=12, value=_num(part.sheet_length, Q2)).alignment = right
        ws.cell(row=start_row, column=13, value=_num(part.sheet_width, Q2)).alignment = right

    for offset, process in enumerate(processes):
        excel_row = start_row + offset
        ws.cell(row=excel_row, column=19, value=process.process_seq)
        ws.cell(row=excel_row, column=20, value=process.process_code or "")
        ws.cell(row=excel_row, column=21, value=process.die_sets)
        ws.cell(row=excel_row, column=22, value=process.die_number or "")
        ws.cell(row=excel_row, column=23, value=process.tonnage)
        ratio_cell = ws.cell(
            row=excel_row,
            column=28,
            value=_num(process.cost_ratio, Q2),
        )
        ratio_cell.alignment = right

    has_sketch = embed_sketch(
        ws,
        part=part,
        start_row=start_row,
        n_rows=n_rows,
        sketch_buffers=sketch_buffers,
    )
    return n_rows, has_sketch
