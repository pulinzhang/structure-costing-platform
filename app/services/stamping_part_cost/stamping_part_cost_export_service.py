"""冲压零件价格核算 Excel 导出。"""

from __future__ import annotations

import logging
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from app.common.audit import audit_status_label
from app.repositories.bom.bom_repository import BomNodeRepository, BomSnapshotRepository
from app.repositories.stamping.scrap_price_repository import StampingScrapPriceRepository
from app.repositories.stamping.stamping_process_fee_repository import (
    StampingProcessFeeRepository,
)
from app.repositories.stamping_part_cost.stamping_part_cost_repository import (
    StampingPartCostRepository,
)
from app.repositories.steel_price.steel_price_repository import SteelPriceRepository
from app.services.stamping_part_cost.helpers import (
    DEFAULT_MATERIAL_LOSS_RATE,
    DEFAULT_SCRAP_RECOVERY_RATE,
    DEFAULT_STEEL_DENSITY,
    format_number as _num,
    format_process_fee_trace,
    format_scrap_price_trace,
    format_steel_price_excel_trace,
    location_candidates,
    part_excel_row_col,
    to_decimal,
)


logger = logging.getLogger("service.stamping_part_cost_export")

ExportFormat = Literal["detail", "process_sheet"]


_THIN = Border(
    left=Side(style="thin", color="333333"),
    right=Side(style="thin", color="333333"),
    top=Side(style="thin", color="333333"),
    bottom=Side(style="thin", color="333333"),
)

_WEAK_HEADER_FILL = PatternFill("solid", fgColor="455A64")
_WEAK_HEADER_FONT = Font(name="微软雅黑", size=8, bold=True, color="B0BEC5")
_WEAK_DATA_FONT = Font(name="微软雅黑", size=8, color="9E9E9E")


def _dt(value: Any) -> str:
    if value is None:
        return ""
    return value.strftime("%Y-%m-%d %H:%M:%S") if hasattr(value, "strftime") else str(value)


def _formula_number(value: Decimal) -> str:
    """写入公式的小数，去掉多余尾零。"""
    text = format(value, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _build_part_to_l0_map(nodes: list) -> dict[str, str]:
    """
    零件号 → L0 供货总成零件号。
    同一冲压件挂在多个总成下时，收集全部 L0，英文逗号分隔（去重、保序）。
    """
    by_id = {node.id: node for node in nodes}
    roots_by_part: dict[str, list[str]] = {}

    for node in nodes:
        part_number = (node.part_number or "").strip()
        if not part_number:
            continue
        current = node
        visited: set[int] = set()
        while current.parent_id is not None:
            if current.id in visited:
                break
            visited.add(current.id)
            parent = by_id.get(current.parent_id)
            if parent is None:
                break
            current = parent
        root_part = (current.part_number or "").strip()
        if not root_part:
            continue
        bucket = roots_by_part.setdefault(part_number, [])
        if root_part not in bucket:
            bucket.append(root_part)

    return {
        part_number: ",".join(roots)
        for part_number, roots in roots_by_part.items()
    }


class StampingPartCostExportService:
    """导出零件价格核算结果。"""

    def __init__(self) -> None:
        self.repository = StampingPartCostRepository()
        self.bom_snapshot_repo = BomSnapshotRepository()
        self.bom_node_repo = BomNodeRepository()
        self.steel_price_repo = SteelPriceRepository()
        self.scrap_price_repo = StampingScrapPriceRepository()
        self.process_fee_repo = StampingProcessFeeRepository()

    def export(
        self,
        db: Session,
        *,
        year: int | None = None,
        month: int | None = None,
        location: str | None = None,
        stamping_supplier: str | None = None,
        vehicle_model: str | None = None,
        version: str | None = None,
        part_number: str | None = None,
        audit_status: str | None = None,
        export_format: ExportFormat = "process_sheet",
    ) -> tuple[bytes, str]:
        rows = self.repository.list_for_export(
            db,
            year=year,
            month=month,
            location=location,
            stamping_supplier=stamping_supplier,
            vehicle_model=vehicle_model,
            version=version,
            part_number=part_number,
            audit_status=audit_status,
        )
        if export_format == "detail":
            return self._export_detail(db, rows)
        return self._export_process_sheet(db, rows)

    def _resolve_assembly_supply_no(
        self,
        db: Session,
        *,
        location: str | None,
        vehicle_model: str | None,
        version: str | None,
        part_number: str | None,
        snapshot_l0_cache: dict[tuple[str, str, str], dict[str, str]],
    ) -> str:
        """
        沿 BOM parent_id 上溯至 L0，取其零件号作为总成供货号。
        同一零件挂多个 L0 时逗号分隔。BOM 查询须基地相同（兼容「青岛/青岛基地」写法）。

        匹配顺序：
        1) 同基地 + vehicle_model + version(=bom_version) 精确命中
        2) 同基地 + vehicle_model 下回退（工艺版本与 BOM 版本号可能不一致）
        """
        loc = (location or "").strip()
        model = (vehicle_model or "").strip()
        bom_version = (version or "").strip()
        pn = (part_number or "").strip()
        if not loc or not model or not pn:
            return ""

        part_to_l0: dict[str, str] | None = None
        resolved_key: tuple[str, str, str] | None = None

        # 1) 精确版本
        if bom_version:
            for candidate_location in location_candidates(loc):
                key = (candidate_location, model, bom_version)
                if key in snapshot_l0_cache:
                    part_to_l0 = snapshot_l0_cache[key]
                    resolved_key = key
                    break
                snapshot = self.bom_snapshot_repo.get_by_unique_key(
                    db,
                    location=candidate_location,
                    vehicle_model=model,
                    bom_version=bom_version,
                )
                if snapshot is None:
                    continue
                part_to_l0 = _build_part_to_l0_map(
                    self.bom_node_repo.list_by_snapshot(db, snapshot.id)
                )
                snapshot_l0_cache[key] = part_to_l0
                resolved_key = key
                break

        # 2) 同车型最新 BOM 回退
        if part_to_l0 is None or (part_to_l0 is not None and pn not in part_to_l0):
            for candidate_location in location_candidates(loc):
                fallback_key = (candidate_location, model, "__latest__")
                if fallback_key in snapshot_l0_cache:
                    fallback_map = snapshot_l0_cache[fallback_key]
                else:
                    snapshots = self.bom_snapshot_repo.list_by_location_model(
                        db,
                        location=candidate_location,
                        vehicle_model=model,
                    )
                    if not snapshots:
                        snapshot_l0_cache[fallback_key] = {}
                        continue
                    # 优先选包含该零件号的 snapshot；否则取最新
                    chosen = None
                    for snap in snapshots:
                        nodes = self.bom_node_repo.list_by_snapshot(db, snap.id)
                        mapping = _build_part_to_l0_map(nodes)
                        snap_key = (
                            candidate_location,
                            model,
                            (snap.bom_version or "").strip() or f"id:{snap.id}",
                        )
                        snapshot_l0_cache[snap_key] = mapping
                        if pn in mapping:
                            chosen = mapping
                            break
                        if chosen is None:
                            chosen = mapping
                    snapshot_l0_cache[fallback_key] = chosen or {}
                    fallback_map = snapshot_l0_cache[fallback_key]

                if fallback_map:
                    part_to_l0 = fallback_map
                    resolved_key = fallback_key
                    if pn in fallback_map:
                        break

        if not part_to_l0:
            empty_key = (loc, model, bom_version or "__latest__")
            snapshot_l0_cache.setdefault(empty_key, {})
            return ""

        return part_to_l0.get(pn, "")

    def _export_process_sheet(self, db: Session, rows: list) -> tuple[bytes, str]:
        """
        工艺清单口径：对齐样例至 R 列 + 加工费 + 最终价。
        一行一个零件，M/N/O/I/R/最终价保留 Excel 公式。
        总成供货号：同基地 BOM 上溯 L0；多总成时逗号分隔。
        末尾 U~Z 为人可读溯源（核对/手工补填），AA~AC 为关联 ID（钢材/废料/加工费）。
        """
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "单件"

        last_col = 29  # AC
        last_col_letter = "AC"
        trace_cols = frozenset(
            {"U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC"}
        )
        id_cols = frozenset({"AA", "AB", "AC"})

        # —— 样式 ——
        thin = Border(
            left=Side(style="thin", color="B0B0B0"),
            right=Side(style="thin", color="B0B0B0"),
            top=Side(style="thin", color="B0B0B0"),
            bottom=Side(style="thin", color="B0B0B0"),
        )
        title_font = Font(name="微软雅黑", size=14, bold=True, color="1F1F1F")
        meta_font = Font(name="微软雅黑", size=9, color="666666")
        group_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        header_font = Font(name="微软雅黑", size=9, bold=True, color="1F1F1F")
        weak_header_font = Font(name="微软雅黑", size=8, bold=True, color="757575")
        data_font = Font(name="微软雅黑", size=9, color="000000")
        weak_data_font = Font(name="微软雅黑", size=8, color="9E9E9E")
        accent_font = Font(name="微软雅黑", size=9, bold=True, color="C62828")

        fill_title = PatternFill("solid", fgColor="F5F5F5")
        fill_group_id = PatternFill("solid", fgColor="424242")
        fill_group_product = PatternFill("solid", fgColor="455A64")
        fill_group_process = PatternFill("solid", fgColor="1565C0")
        fill_group_fee = PatternFill("solid", fgColor="2E7D32")
        fill_group_final = PatternFill("solid", fgColor="C62828")
        fill_group_trace = PatternFill("solid", fgColor="78909C")
        fill_header = PatternFill("solid", fgColor="ECEFF1")
        fill_header_trace = PatternFill("solid", fgColor="ECEFF1")
        fill_zebra = PatternFill("solid", fgColor="FAFAFA")
        fill_zebra_trace = PatternFill("solid", fgColor="F5F5F5")
        fill_result = PatternFill("solid", fgColor="FFF3E0")
        fill_final = PatternFill("solid", fgColor="FFEBEE")

        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center")
        right = Alignment(horizontal="right", vertical="center")

        # 范围摘要 + 核价口径（本批快照参数）
        meta_parts: list[str] = []
        density_text = "7.85"
        loss_rate_text = "0.003"
        loss_permille_text = "3"
        recovery_text = "0.97"
        recovery_percent_text = "97"
        if rows:
            sample = rows[0]
            if sample.year and sample.month:
                meta_parts.append(f"{sample.year}年{sample.month}月")
            if sample.location:
                meta_parts.append(str(sample.location))
            if sample.stamping_supplier:
                meta_parts.append(str(sample.stamping_supplier))
            density = to_decimal(getattr(sample, "density", None))
            loss_rate = to_decimal(getattr(sample, "material_loss_rate", None))
            recovery = to_decimal(getattr(sample, "scrap_recovery_rate", None))
            if density is not None:
                density_text = _formula_number(density)
            if loss_rate is not None:
                loss_rate_text = _formula_number(loss_rate)
                loss_permille_text = _formula_number(loss_rate * Decimal("1000"))
            if recovery is not None:
                recovery_text = _formula_number(recovery)
                recovery_percent_text = _formula_number(recovery * Decimal("100"))
        meta_text = " · ".join(meta_parts) if meta_parts else "按筛选条件导出"
        formula_hint = (
            f"核价口径："
            f"M=L×W×T×ρ_steel÷1,000,000÷n；"
            f"O=M−H；"
            f"R=M×P×(1+λ)−M×λ×η×Q−O×η×Q；"
            f"最终价=R+S。"
            f"其中 ρ_steel={density_text} g/cm³，"
            f"λ={loss_rate_text}（{loss_permille_text}‰），"
            f"η={recovery_text}（{recovery_percent_text}%），"
            f"P/Q 为元/kg（吨价÷1000），"
            f"÷1,000,000 为 mm·g/cm³ → kg 量纲换算。"
        )

        # Row 1 标题
        sheet.merge_cells(f"A1:{last_col_letter}1")
        title_cell = sheet["A1"]
        title_cell.value = "冲压零件工艺清单核价"
        title_cell.font = title_font
        title_cell.fill = fill_title
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        for col in range(1, last_col + 1):
            sheet.cell(1, col).fill = fill_title
            sheet.cell(1, col).border = thin
        sheet.row_dimensions[1].height = 28

        # Row 2 说明
        sheet.merge_cells(f"A2:{last_col_letter}2")
        meta_cell = sheet["A2"]
        meta_cell.value = (
            f"{meta_text}　|　共 {len(rows)} 件　|　{formula_hint}"
            f"　|　U~Z 人可读溯源可核对/补填；AA~AC 为关联 ID"
        )
        meta_cell.font = meta_font
        meta_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        for col in range(1, last_col + 1):
            sheet.cell(2, col).fill = fill_title
            sheet.cell(2, col).border = thin
        sheet.row_dimensions[2].height = 48

        # Row 3 分组表头
        group_fills = {
            range(1, 3): fill_group_id,
            range(3, 10): fill_group_product,
            range(10, 19): fill_group_process,
            range(19, 20): fill_group_fee,
            range(20, 21): fill_group_final,
            range(21, 27): fill_group_trace,
            range(27, 30): fill_group_trace,
        }
        for merge_range, label in (
            ("A3:B3", "标识"),
            ("C3:I3", "产品信息"),
            ("J3:R3", "工艺清单（材料）"),
            ("S3:S3", "加工费"),
            ("T3:T3", "最终价"),
            ("U3:Z3", "溯源"),
            ("AA3:AC3", "关联ID"),
        ):
            sheet.merge_cells(merge_range)
            cell = sheet[merge_range.split(":")[0]]
            cell.value = label
            cell.font = group_font
            cell.alignment = center
        for col in range(1, last_col + 1):
            cell = sheet.cell(3, col)
            cell.border = thin
            cell.font = group_font
            cell.alignment = center
            for col_range, fill in group_fills.items():
                if col in col_range:
                    cell.fill = fill
                    break
        sheet.row_dimensions[3].height = 22

        # Row 4 列名
        headers = [
            ("A", "序号"),
            ("B", "总成供货号"),
            ("C", "零件号"),
            ("D", "零件名称"),
            ("E", "零件简图"),
            ("F", "材质"),
            ("G", "料厚"),
            ("H", "净重"),
            ("I", "零件尺寸"),
            ("J", "长"),
            ("K", "宽(步距)"),
            ("L", "可制件数"),
            ("M", "预估单件材料消耗(KG)"),
            ("N", "材料利用率"),
            ("O", "废料重量"),
            ("P", "钢板单价(元/KG)"),
            ("Q", "废料单价(元/KG)"),
            ("R", "原材料成本"),
            ("S", "加工费"),
            ("T", "最终价"),
            ("U", "来源Excel"),
            ("V", "工作表"),
            ("W", "行/列"),
            ("X", "钢材价溯源"),
            ("Y", "废料价溯源"),
            ("Z", "加工费溯源"),
            ("AA", "钢材价ID"),
            ("AB", "废料价ID"),
            ("AC", "加工费ID"),
        ]
        for col_letter, title in headers:
            cell = sheet[f"{col_letter}4"]
            cell.value = title
            is_trace = col_letter in trace_cols
            cell.font = weak_header_font if is_trace else header_font
            cell.fill = fill_header_trace if is_trace else fill_header
            cell.alignment = center
            cell.border = thin
        sheet.row_dimensions[4].height = 40

        # 数据行
        data_start = 5
        number_cols = {"G", "H", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T"}
        result_cols = {"R"}
        final_cols = {"T"}
        text_left_cols = {"B", "C", "D", "E", "F", "I"}
        snapshot_l0_cache: dict[tuple[str, str, str], dict[str, str]] = {}

        steel_ids = sorted(
            {
                int(row.steel_price_id)
                for row in rows
                if row.steel_price_id is not None
            }
        )
        steel_by_id = {
            steel.id: steel
            for steel in self.steel_price_repo.list_by_ids(db, steel_ids)
        }
        scrap_ids = sorted(
            {
                int(row.scrap_price_id)
                for row in rows
                if row.scrap_price_id is not None
            }
        )
        scrap_by_id = {
            scrap.id: scrap
            for scrap in self.scrap_price_repo.list_by_ids(db, scrap_ids)
        }
        fee_ids = sorted(
            {
                int(row.process_fee_id)
                for row in rows
                if row.process_fee_id is not None
            }
        )
        fee_by_id = {
            fee.id: fee for fee in self.process_fee_repo.list_by_ids(db, fee_ids)
        }

        for offset, row in enumerate(rows):
            excel_row = data_start + offset
            density = to_decimal(getattr(row, "density", None)) or DEFAULT_STEEL_DENSITY
            loss_rate = (
                to_decimal(getattr(row, "material_loss_rate", None))
                or DEFAULT_MATERIAL_LOSS_RATE
            )
            recovery = (
                to_decimal(getattr(row, "scrap_recovery_rate", None))
                or DEFAULT_SCRAP_RECOVERY_RATE
            )
            steel_ton = to_decimal(row.steel_ton_price)
            scrap_ton = to_decimal(row.scrap_ton_price)
            steel_kg = (steel_ton / Decimal("1000")) if steel_ton is not None else None
            scrap_kg = (scrap_ton / Decimal("1000")) if scrap_ton is not None else None
            assembly_supply_no = self._resolve_assembly_supply_no(
                db,
                location=row.location,
                vehicle_model=row.vehicle_model,
                version=row.version,
                part_number=row.part_number,
                snapshot_l0_cache=snapshot_l0_cache,
            )

            seq = row.seq_no if row.seq_no is not None else offset + 1
            values = {
                "A": seq,
                "B": assembly_supply_no,
                "C": row.part_number or "",
                "D": row.part_name or "",
                "E": "",
                "F": row.material or "",
                "G": _num(row.thickness, 4) if row.thickness is not None else "",
                "H": _num(row.part_weight, 4) if row.part_weight is not None else "",
                "I": f'=F{excel_row}&"-"&G{excel_row}',
                "J": _num(row.sheet_length, 2) if row.sheet_length is not None else "",
                "K": _num(row.sheet_width, 2) if row.sheet_width is not None else "",
                "L": row.parts_per_sheet if row.parts_per_sheet is not None else "",
                "M": (
                    f"=J{excel_row}*K{excel_row}*G{excel_row}"
                    f"*{_formula_number(density)}/1000000/L{excel_row}"
                ),
                "N": f"=H{excel_row}/M{excel_row}",
                "O": f"=M{excel_row}-H{excel_row}",
                "P": float(steel_kg) if steel_kg is not None else "",
                "Q": float(scrap_kg) if scrap_kg is not None else "",
                "R": (
                    f"=M{excel_row}*P{excel_row}*(1+{_formula_number(loss_rate)})"
                    f"-M{excel_row}*{_formula_number(loss_rate)}"
                    f"*{_formula_number(recovery)}*Q{excel_row}"
                    f"-O{excel_row}*Q{excel_row}*{_formula_number(recovery)}"
                ),
                "S": (
                    _num(row.process_fee_amount, 3)
                    if row.process_fee_amount is not None
                    else ""
                ),
                "T": f"=R{excel_row}+S{excel_row}",
                "U": row.excel_filename or "",
                "V": row.excel_sheet_name or "",
                "W": part_excel_row_col(row),
                "X": format_steel_price_excel_trace(
                    steel_by_id.get(row.steel_price_id)
                    if row.steel_price_id is not None
                    else None
                ),
                "Y": format_scrap_price_trace(
                    scrap_by_id.get(row.scrap_price_id)
                    if row.scrap_price_id is not None
                    else None
                ),
                "Z": format_process_fee_trace(
                    fee_by_id.get(row.process_fee_id)
                    if row.process_fee_id is not None
                    else None
                ),
                "AA": (
                    row.steel_price_id
                    if row.steel_price_id is not None
                    else ""
                ),
                "AB": (
                    row.scrap_price_id
                    if row.scrap_price_id is not None
                    else ""
                ),
                "AC": (
                    row.process_fee_id
                    if row.process_fee_id is not None
                    else ""
                ),
            }

            zebra = offset % 2 == 1
            for col_letter, value in values.items():
                cell = sheet[f"{col_letter}{excel_row}"]
                cell.value = value
                is_trace = col_letter in trace_cols
                if is_trace:
                    cell.font = weak_data_font
                elif col_letter in final_cols:
                    cell.font = accent_font
                else:
                    cell.font = data_font
                cell.border = thin
                if col_letter in final_cols:
                    cell.fill = fill_final
                elif col_letter in result_cols:
                    cell.fill = fill_result
                elif is_trace:
                    cell.fill = (
                        fill_zebra_trace
                        if zebra
                        else PatternFill("solid", fgColor="FFFFFF")
                    )
                elif zebra:
                    cell.fill = fill_zebra

                if col_letter in {"X", "Y", "Z"}:
                    cell.alignment = Alignment(
                        horizontal="left", vertical="center", wrap_text=True
                    )
                elif col_letter in id_cols:
                    cell.alignment = center
                elif col_letter in text_left_cols or col_letter in {"U", "V"}:
                    cell.alignment = left
                elif col_letter in number_cols or col_letter == "A":
                    cell.alignment = right if col_letter != "A" else center
                else:
                    cell.alignment = center

                # 数字格式（含公式列）
                if col_letter in {"G", "H", "P", "Q", "S"}:
                    cell.number_format = "0.####"
                elif col_letter in {"J", "K"}:
                    cell.number_format = "0.##"
                elif col_letter == "N":
                    cell.number_format = "0.00%"
                elif col_letter in {"M", "O", "R", "T"}:
                    cell.number_format = "0.0000"
                elif col_letter == "L":
                    cell.number_format = "0"
                elif col_letter in id_cols:
                    cell.number_format = "0"

            multi_line = any(
                isinstance(values.get(col), str) and "\n" in values[col]
                for col in ("X", "Y", "Z")
            )
            sheet.row_dimensions[excel_row].height = 52 if multi_line else 18

        widths = {
            "A": 6,
            "B": 22,
            "C": 14,
            "D": 20,
            "E": 10,
            "F": 14,
            "G": 8,
            "H": 9,
            "I": 14,
            "J": 8,
            "K": 10,
            "L": 10,
            "M": 14,
            "N": 11,
            "O": 10,
            "P": 13,
            "Q": 13,
            "R": 12,
            "S": 10,
            "T": 11,
            "U": 28,
            "V": 14,
            "W": 8,
            "X": 42,
            "Y": 28,
            "Z": 36,
            "AA": 10,
            "AB": 10,
            "AC": 10,
        }
        for col, width in widths.items():
            sheet.column_dimensions[col].width = width

        last_data_row = data_start + len(rows) - 1 if rows else 4
        sheet.freeze_panes = "D5"  # 冻结至零件号列（A~C：序号/总成供货号/零件号）
        if rows:
            sheet.auto_filter.ref = f"A4:{last_col_letter}{last_data_row}"
        sheet.print_title_rows = "1:4"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0

        # 附录：核价口径定义（次于「单件」工作表，打开文件默认见单件）
        note = workbook.create_sheet("核价口径")
        note_font_title = Font(name="微软雅黑", size=12, bold=True, color="1F1F1F")
        note_font = Font(name="微软雅黑", size=10, color="333333")
        note_header = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        note_header_fill = PatternFill("solid", fgColor="455A64")
        note["A1"] = "冲压零件核价口径定义"
        note["A1"].font = note_font_title
        note.merge_cells("A1:C1")
        note.row_dimensions[1].height = 24
        note["A2"] = (
            f"本批次参数：钢材密度 ρ_steel={density_text} g/cm³；"
            f"材料损耗率 λ={loss_rate_text}（{loss_permille_text}‰）；"
            f"废料回收率 η={recovery_text}（{recovery_percent_text}%）"
        )
        note["A2"].font = meta_font
        note.merge_cells("A2:C2")

        note_headers = ("符号", "定义", "计算关系 / 单位")
        for col, text in enumerate(note_headers, start=1):
            cell = note.cell(4, col, text)
            cell.font = note_header
            cell.fill = note_header_fill
            cell.alignment = center
            cell.border = thin

        note_rows = [
            ("B", "总成供货号", "同基地 BOM 沿 parent 上溯至 L0；多总成逗号分隔"),
            ("C", "零件号", "本行核算零件号"),
            ("L, W, T", "板料长、宽(步距)、料厚", "mm；工艺清单预估板料尺寸"),
            ("n", "可制件数", "件/板"),
            ("ρ_steel", "钢材密度", f"{density_text} g/cm³"),
            ("1,000,000", "量纲换算系数", "mm³·(g/cm³) → kg"),
            ("M", "预估单件材料消耗", "M = L×W×T×ρ_steel / 1,000,000 / n　(kg)"),
            ("H", "零件净重", "kg"),
            ("O", "废料重量", "O = M − H　(kg)"),
            ("P", "钢板单价", "元/kg（吨价÷1000）"),
            ("Q", "废料单价", "元/kg（吨价÷1000）"),
            ("λ", "材料损耗率", f"{loss_rate_text}（{loss_permille_text}‰）"),
            ("η", "废料回收率", f"{recovery_text}（{recovery_percent_text}%）"),
            ("①", "钢材采购额", "M × P × (1+λ)"),
            ("②", "损耗废料扣除", "M × λ × η × Q"),
            ("③", "边角废料扣除", "O × η × Q"),
            ("R", "原材料成本", "R = ① − ② − ③"),
            ("S", "加工费", "元/件；取自加工费核算结果"),
            ("最终价", "单件制件价格", "R + S"),
            (
                "U~Z",
                "溯源（弱化）",
                "按文件/表/行定位核对；钢材含基价·厚度·宽度·锌层；废料为基地年月净价；"
                "加工费为工艺文件+工序依据。单元格可手工补填。",
            ),
            (
                "AA~AC",
                "关联ID（弱化）",
                "钢材价ID / 废料价ID / 加工费ID，供系统回链关联展示。",
            ),
        ]
        for idx, (symbol, name, desc) in enumerate(note_rows, start=5):
            note.cell(idx, 1, symbol).font = note_font
            note.cell(idx, 2, name).font = note_font
            note.cell(idx, 3, desc).font = note_font
            for col in range(1, 4):
                note.cell(idx, col).border = thin
                note.cell(idx, col).alignment = left if col > 1 else center

        note.column_dimensions["A"].width = 14
        note.column_dimensions["B"].width = 18
        note.column_dimensions["C"].width = 52

        # 确保打开文件时激活「单件」
        workbook.active = sheet

        buffer = BytesIO()
        workbook.save(buffer)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"冲压零件工艺清单核价_{stamp}.xlsx"
        logger.info("导出工艺清单格式 %s 条 → %s", len(rows), filename)
        return buffer.getvalue(), filename

    def _export_detail(self, db: Session, rows: list) -> tuple[bytes, str]:
        """明细导出：业务字段 + 末尾弱化溯源/关联 ID + 审核元数据。"""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "零件价格核算"

        headers = [
            "年度",
            "月份",
            "基地",
            "冲压工厂",
            "车型",
            "版本",
            "零件号",
            "零件名称",
            "材质",
            "料厚",
            "单重kg",
            "料片长",
            "料片宽",
            "可制件数",
            "密度ρ_steel",
            "材料损耗率λ",
            "废料回收率η",
            "材料消耗M kg",
            "净重H kg",
            "废料重量O kg",
            "材料利用率%",
            "钢厂",
            "钢材吨价",
            "废料吨价",
            "①钢材采购 M×P×(1+λ)",
            "②损耗废料扣除 M×λ×η×Q",
            "③边角废料扣除 O×η×Q",
            "废料扣除合计 ②+③",
            "原材料成本R ①−②−③",
            "加工费",
            "单件成本",
            "计算状态",
            "失败原因",
            "审核状态",
            "计算批次",
            "公式版本",
            "计算时间",
            # 溯源（弱化）
            "来源Excel",
            "工作表",
            "行/列",
            "钢材价溯源",
            "废料价溯源",
            "加工费溯源",
            "钢材价ID",
            "废料价ID",
            "加工费ID",
        ]
        weak_from_col = 38  # 来源Excel 起
        header_fill = PatternFill("solid", fgColor="1E1E1E")
        header_font = Font(color="FFFFFF", bold=True, size=10)
        for col, title in enumerate(headers, start=1):
            cell = sheet.cell(1, col, title)
            if col >= weak_from_col:
                cell.fill = _WEAK_HEADER_FILL
                cell.font = _WEAK_HEADER_FONT
            else:
                cell.fill = header_fill
                cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = _THIN

        steel_ids = sorted(
            {
                int(row.steel_price_id)
                for row in rows
                if row.steel_price_id is not None
            }
        )
        steel_by_id = {
            steel.id: steel
            for steel in self.steel_price_repo.list_by_ids(db, steel_ids)
        }
        scrap_ids = sorted(
            {
                int(row.scrap_price_id)
                for row in rows
                if row.scrap_price_id is not None
            }
        )
        scrap_by_id = {
            scrap.id: scrap
            for scrap in self.scrap_price_repo.list_by_ids(db, scrap_ids)
        }
        fee_ids = sorted(
            {
                int(row.process_fee_id)
                for row in rows
                if row.process_fee_id is not None
            }
        )
        fee_by_id = {
            fee.id: fee for fee in self.process_fee_repo.list_by_ids(db, fee_ids)
        }

        for row_index, row in enumerate(rows, start=2):
            values = [
                row.year,
                row.month,
                row.location,
                row.stamping_supplier,
                row.vehicle_model,
                row.version,
                row.part_number,
                row.part_name or "",
                row.material or "",
                _num(row.thickness, 4),
                _num(row.part_weight, 4),
                _num(row.sheet_length, 2),
                _num(row.sheet_width, 2),
                row.parts_per_sheet if row.parts_per_sheet is not None else "",
                _num(row.density, 4),
                _num(getattr(row, "material_loss_rate", None), 6),
                _num(getattr(row, "scrap_recovery_rate", None), 6),
                _num(row.blank_weight_kg, 10),
                _num(row.finished_total_kg, 10),
                _num(row.scrap_weight_kg, 10),
                _num(row.material_utilization_pct, 2),
                row.steel_mill or "",
                _num(row.steel_ton_price, 2),
                _num(row.scrap_ton_price, 2),
                _num(row.material_amount_sheet, 4),
                _num(getattr(row, "loss_scrap_deduction", None), 4),
                _num(getattr(row, "process_scrap_deduction", None), 4),
                _num(row.scrap_amount_sheet, 4),
                _num(row.material_net_per_part, 4),
                _num(row.process_fee_amount, 3),
                _num(row.unit_cost, 4),
                row.calc_status or "",
                row.calc_fail_reason or "",
                audit_status_label(row.audit_status),
                row.calc_batch_id or "",
                row.formula_version or "",
                _dt(row.calculated_at),
                row.excel_filename or "",
                row.excel_sheet_name or "",
                part_excel_row_col(row),
                format_steel_price_excel_trace(
                    steel_by_id.get(row.steel_price_id)
                    if row.steel_price_id is not None
                    else None
                ),
                format_scrap_price_trace(
                    scrap_by_id.get(row.scrap_price_id)
                    if row.scrap_price_id is not None
                    else None
                ),
                format_process_fee_trace(
                    fee_by_id.get(row.process_fee_id)
                    if row.process_fee_id is not None
                    else None
                ),
                row.steel_price_id if row.steel_price_id is not None else "",
                row.scrap_price_id if row.scrap_price_id is not None else "",
                row.process_fee_id if row.process_fee_id is not None else "",
            ]
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row_index, col, value)
                cell.border = _THIN
                if col >= weak_from_col:
                    cell.font = _WEAK_DATA_FONT
                    cell.alignment = Alignment(
                        vertical="center",
                        wrap_text=col in {41, 42, 43},
                    )
                else:
                    cell.alignment = Alignment(vertical="center")

        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 32
        buffer = BytesIO()
        workbook.save(buffer)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"冲压零件价格核算明细_{stamp}.xlsx"
        logger.info("导出明细格式 %s 条 → %s", len(rows), filename)
        return buffer.getvalue(), filename
