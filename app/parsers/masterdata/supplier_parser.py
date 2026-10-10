"""供应商主数据 Excel 解析器（纯函数，无副作用）。

表头：缩写 | 全称/工商名称 | 对外代码 | 别名
禁止调用 Service / Repository。
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

HEADER_ABBREVIATION = "缩写"
HEADER_FULL_NAME = "全称"
HEADER_ERP_CODE = "对外代码"
HEADER_ALIASES = "别名"
HEADER_STATUS = "状态"

_ALIAS_SEPARATORS = ("、", "，", ",", "\n", "\r", ";")


@dataclass(frozen=True, slots=True)
class SupplierParsedRow:
    """解析后的一行供应商。"""

    row_number: int
    abbreviation: str
    full_name: str
    erp_code: str
    aliases: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class SupplierParseResult:
    """解析结果。"""

    rows: list[SupplierParsedRow]
    errors: list[str]


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _split_aliases(raw: str) -> list[str]:
    text = raw.strip()
    if not text:
        return []
    pieces = [text]
    for separator in _ALIAS_SEPARATORS:
        next_pieces: list[str] = []
        for piece in pieces:
            next_pieces.extend(piece.split(separator))
        pieces = next_pieces
    result: list[str] = []
    for piece in pieces:
        item = piece.strip()
        if not item or item in result:
            continue
        result.append(item)
    return result


def _find_header_row(sheet: Worksheet) -> tuple[int, dict[str, int]] | None:
    aliases = {
        HEADER_ABBREVIATION: {HEADER_ABBREVIATION},
        HEADER_FULL_NAME: {HEADER_FULL_NAME, "工商名称", "全称"},
        HEADER_ERP_CODE: {HEADER_ERP_CODE, "ERP 代码", "ERP代码", "ERP"},
        HEADER_ALIASES: {HEADER_ALIASES},
    }
    required = {HEADER_ABBREVIATION, HEADER_FULL_NAME, HEADER_ERP_CODE}
    max_scan = min(sheet.max_row or 1, 20)
    max_col = min(sheet.max_column or 1, 20)
    for row_idx in range(1, max_scan + 1):
        mapping: dict[str, int] = {}
        for col_idx in range(1, max_col + 1):
            text = _cell_text(sheet.cell(row_idx, col_idx).value)
            for key, names in aliases.items():
                if text in names and key not in mapping:
                    mapping[key] = col_idx
        if required.issubset(mapping.keys()):
            return row_idx, mapping
    return None


def parse_supplier_excel(file_content: bytes) -> SupplierParseResult:
    """解析供应商导入表。结构错误写入 errors；业务门禁留给 Service。"""
    errors: list[str] = []
    rows: list[SupplierParsedRow] = []
    try:
        workbook = load_workbook(BytesIO(file_content), data_only=True, read_only=True)
    except Exception:
        return SupplierParseResult(rows=[], errors=["无法读取 Excel 文件"])

    sheet = workbook.worksheets[0]
    found = _find_header_row(sheet)
    if found is None:
        workbook.close()
        return SupplierParseResult(
            rows=[],
            errors=["未找到表头，需包含：缩写、全称、对外代码"],
        )
    header_row, mapping = found
    alias_col = mapping.get(HEADER_ALIASES)

    for row_idx in range(header_row + 1, (sheet.max_row or header_row) + 1):
        abbreviation = _cell_text(sheet.cell(row_idx, mapping[HEADER_ABBREVIATION]).value)
        full_name = _cell_text(sheet.cell(row_idx, mapping[HEADER_FULL_NAME]).value)
        erp_code = _cell_text(sheet.cell(row_idx, mapping[HEADER_ERP_CODE]).value)
        alias_raw = (
            _cell_text(sheet.cell(row_idx, alias_col).value) if alias_col is not None else ""
        )
        if not abbreviation and not full_name and not erp_code and not alias_raw:
            continue

        row_errors: list[str] = []
        if not abbreviation:
            row_errors.append("缩写不能为空")
        elif len(abbreviation) > 50:
            row_errors.append("缩写最长 50 字")
        if not full_name:
            row_errors.append("全称不能为空")
        elif len(full_name) > 100:
            row_errors.append("全称最长 100 字")
        if not erp_code:
            row_errors.append("对外代码不能为空")
        elif len(erp_code) > 50:
            row_errors.append("对外代码最长 50 字")

        alias_items = _split_aliases(alias_raw)
        for alias in alias_items:
            if len(alias) > 100:
                row_errors.append("别名最长 100 字")
                break

        if row_errors:
            errors.append(f"第 {row_idx} 行：" + "；".join(row_errors))
            continue
        rows.append(
            SupplierParsedRow(
                row_number=row_idx,
                abbreviation=abbreviation,
                full_name=full_name,
                erp_code=erp_code,
                aliases=tuple(alias_items),
            )
        )

    workbook.close()
    return SupplierParseResult(rows=rows, errors=errors)
