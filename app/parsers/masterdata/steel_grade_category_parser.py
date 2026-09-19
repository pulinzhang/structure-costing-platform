"""牌号分类 Excel 解析器（纯函数，无副作用）。

表头：大类 | 牌号 | 备注
多余列（大类序号、审计列等）忽略。
禁止调用 Service / Repository。
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

HEADER_CATEGORY_NAME = "大类"
HEADER_GRADE = "牌号"
HEADER_REMARK = "备注"


@dataclass(frozen=True, slots=True)
class SteelGradeCategoryParsedRow:
    """解析后的一行牌号分类。"""

    row_number: int
    category_name: str
    grade: str
    remark: str | None


@dataclass(frozen=True, slots=True)
class SteelGradeCategoryParseResult:
    """解析结果。"""

    rows: list[SteelGradeCategoryParsedRow]
    errors: list[str]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_header_row(sheet: Any) -> tuple[int, dict[str, int]] | None:
    aliases = {
        HEADER_CATEGORY_NAME: {HEADER_CATEGORY_NAME, "钢种大类", "分类"},
        HEADER_GRADE: {HEADER_GRADE, "钢材牌号"},
        HEADER_REMARK: {HEADER_REMARK},
    }
    required = {HEADER_CATEGORY_NAME, HEADER_GRADE}
    max_scan = min(sheet.max_row or 1, 20)
    max_col = min(sheet.max_column or 1, 20)
    for row_idx in range(1, max_scan + 1):
        mapping: dict[str, int] = {}
        for col_idx in range(1, max_col + 1):
            text = _cell_text(sheet.cell(row_idx, col_idx).value)
            for key, names in aliases.items():
                if text in names and key not in mapping:
                    mapping[key] = col_idx
        if required.issubset(mapping.keys()):
            return row_idx, mapping
    return None


def parse_steel_grade_category_excel(file_content: bytes) -> SteelGradeCategoryParseResult:
    """解析牌号分类导入表。结构错误写入 errors；业务门禁留给 Service。"""
    errors: list[str] = []
    rows: list[SteelGradeCategoryParsedRow] = []
    try:
        workbook = load_workbook(BytesIO(file_content), data_only=True, read_only=True)
    except Exception:
        return SteelGradeCategoryParseResult(rows=[], errors=["无法读取 Excel 文件"])

    sheet = workbook.worksheets[0]
    found = _find_header_row(sheet)
    if found is None:
        workbook.close()
        return SteelGradeCategoryParseResult(
            rows=[],
            errors=["未找到表头，需包含：大类、牌号"],
        )
    header_row, mapping = found
    remark_col = mapping.get(HEADER_REMARK)

    for row_idx in range(header_row + 1, (sheet.max_row or header_row) + 1):
        category_name = _cell_text(sheet.cell(row_idx, mapping[HEADER_CATEGORY_NAME]).value)
        grade = _cell_text(sheet.cell(row_idx, mapping[HEADER_GRADE]).value)
        remark_raw = (
            _cell_text(sheet.cell(row_idx, remark_col).value) if remark_col is not None else ""
        )
        if not category_name and not grade and not remark_raw:
            continue

        row_errors: list[str] = []
        if not category_name:
            row_errors.append("大类不能为空")
        if not grade:
            row_errors.append("牌号不能为空")
        elif len(grade) > 50:
            row_errors.append("牌号最长 50 字")

        if row_errors:
            errors.append(f"第 {row_idx} 行：" + "；".join(row_errors))
            continue
        rows.append(
            SteelGradeCategoryParsedRow(
                row_number=row_idx,
                category_name=category_name,
                grade=grade,
                remark=remark_raw or None,
            )
        )

    workbook.close()
    return SteelGradeCategoryParseResult(rows=rows, errors=errors)
