"""项目主数据 Excel 解析器（纯函数，无副作用）。

表头：基地 | 项目代号 | 生效年 | 生效月 | 失效年 | 失效月 | 当前供货台数
禁止调用 Service / Repository。
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

HEADER_LOCATION = "基地"
HEADER_VEHICLE_MODEL = "项目代号"
HEADER_EFFECTIVE_YEAR = "生效年"
HEADER_EFFECTIVE_MONTH = "生效月"
HEADER_EXPIRY_YEAR = "失效年"
HEADER_EXPIRY_MONTH = "失效月"
HEADER_CURRENT_SUPPLIED_UNITS = "当前供货台数"


@dataclass(frozen=True, slots=True)
class VehicleModelParsedRow:
    """解析后的一行项目。"""

    row_number: int
    location: str
    vehicle_model: str
    effective_year: int
    effective_month: int
    expiry_year: int | None
    expiry_month: int | None
    current_supplied_units: int


@dataclass(frozen=True, slots=True)
class VehicleModelParseResult:
    """解析结果。"""

    rows: list[VehicleModelParsedRow]
    errors: list[str]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip()
    if not text:
        return None
    if text.isdigit() or (text.startswith("-") and text[1:].isdigit()):
        return int(text)
    return None


def _find_header_row(sheet: Any) -> tuple[int, dict[str, int]] | None:
    aliases = {
        HEADER_LOCATION: {HEADER_LOCATION},
        HEADER_VEHICLE_MODEL: {HEADER_VEHICLE_MODEL, "项目"},
        HEADER_EFFECTIVE_YEAR: {HEADER_EFFECTIVE_YEAR},
        HEADER_EFFECTIVE_MONTH: {HEADER_EFFECTIVE_MONTH},
        HEADER_EXPIRY_YEAR: {HEADER_EXPIRY_YEAR},
        HEADER_EXPIRY_MONTH: {HEADER_EXPIRY_MONTH},
        HEADER_CURRENT_SUPPLIED_UNITS: {HEADER_CURRENT_SUPPLIED_UNITS, "供货台数"},
    }
    required = {
        HEADER_LOCATION,
        HEADER_VEHICLE_MODEL,
        HEADER_EFFECTIVE_YEAR,
        HEADER_EFFECTIVE_MONTH,
        HEADER_EXPIRY_YEAR,
        HEADER_EXPIRY_MONTH,
        HEADER_CURRENT_SUPPLIED_UNITS,
    }
    max_scan = min(sheet.max_row or 1, 20)
    max_col = min(sheet.max_column or 1, 20)
    for row_idx in range(1, max_scan + 1):
        mapping: dict[str, int] = {}
        for col_idx in range(1, max_col + 1):
            text = _cell_text(sheet.cell(row_idx, col_idx).value)
            for key, names in aliases.items():
                if text in names and key not in mapping:
                    mapping[key] = col_idx
        if required.issubset(mapping.keys()):
            return row_idx, mapping
    return None


def parse_vehicle_model_excel(file_content: bytes) -> VehicleModelParseResult:
    """解析项目导入表。结构错误写入 errors；业务门禁留给 Service。"""
    errors: list[str] = []
    rows: list[VehicleModelParsedRow] = []
    try:
        workbook = load_workbook(BytesIO(file_content), data_only=True, read_only=True)
    except Exception:
        return VehicleModelParseResult(rows=[], errors=["无法读取 Excel 文件"])

    sheet = workbook.worksheets[0]
    found = _find_header_row(sheet)
    if found is None:
        workbook.close()
        return VehicleModelParseResult(
            rows=[],
            errors=[
                "未找到表头，需包含：基地、项目代号、生效年、生效月、失效年、失效月、当前供货台数"
            ],
        )
    header_row, mapping = found

    for row_idx in range(header_row + 1, (sheet.max_row or header_row) + 1):
        location = _cell_text(sheet.cell(row_idx, mapping[HEADER_LOCATION]).value)
        vehicle_model = _cell_text(sheet.cell(row_idx, mapping[HEADER_VEHICLE_MODEL]).value)
        effective_year_raw = sheet.cell(row_idx, mapping[HEADER_EFFECTIVE_YEAR]).value
        effective_month_raw = sheet.cell(row_idx, mapping[HEADER_EFFECTIVE_MONTH]).value
        expiry_year_raw = sheet.cell(row_idx, mapping[HEADER_EXPIRY_YEAR]).value
        expiry_month_raw = sheet.cell(row_idx, mapping[HEADER_EXPIRY_MONTH]).value
        supplied_raw = sheet.cell(row_idx, mapping[HEADER_CURRENT_SUPPLIED_UNITS]).value

        expiry_year_text = _cell_text(expiry_year_raw)
        expiry_month_text = _cell_text(expiry_month_raw)
        supplied_text = _cell_text(supplied_raw)
        if (
            not location
            and not vehicle_model
            and effective_year_raw in (None, "")
            and effective_month_raw in (None, "")
            and not expiry_year_text
            and not expiry_month_text
            and not supplied_text
        ):
            continue

        row_errors: list[str] = []
        if not location:
            row_errors.append("基地不能为空")
        if not vehicle_model:
            row_errors.append("项目代号不能为空")

        effective_year = _parse_int(effective_year_raw)
        effective_month = _parse_int(effective_month_raw)
        if effective_year is None:
            row_errors.append("生效年必须为整数")
        if effective_month is None or not (1 <= effective_month <= 12):
            row_errors.append("生效月必须为 1–12")

        has_expiry_year = bool(expiry_year_text)
        has_expiry_month = bool(expiry_month_text)
        expiry_year: int | None = None
        expiry_month: int | None = None
        if has_expiry_year or has_expiry_month:
            if has_expiry_year != has_expiry_month:
                row_errors.append("失效年与失效月必须同空或同有")
            else:
                expiry_year = _parse_int(expiry_year_raw)
                expiry_month = _parse_int(expiry_month_raw)
                if expiry_year is None:
                    row_errors.append("失效年必须为整数")
                if expiry_month is None or not (1 <= expiry_month <= 12):
                    row_errors.append("失效月必须为 1–12")

        current_supplied_units = 0
        if supplied_text:
            parsed_units = _parse_int(supplied_raw)
            if parsed_units is None or parsed_units < 0:
                row_errors.append("当前供货台数必须为 ≥0 整数")
            else:
                current_supplied_units = parsed_units

        if row_errors:
            errors.append(f"第 {row_idx} 行：" + "；".join(row_errors))
            continue
        if effective_year is None or effective_month is None:
            continue
        rows.append(
            VehicleModelParsedRow(
                row_number=row_idx,
                location=location,
                vehicle_model=vehicle_model,
                effective_year=effective_year,
                effective_month=effective_month,
                expiry_year=expiry_year,
                expiry_month=expiry_month,
                current_supplied_units=current_supplied_units,
            )
        )

    workbook.close()
    return VehicleModelParseResult(rows=rows, errors=errors)
