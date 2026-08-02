"""PolicyAssembler：封装所有"业务语义解析"逻辑。

职责：
- 块定位（locate_thickness_block / locate_width_block）
- 边界推断（collect_boundaries / infer_upper_bound）
- grade_type 推断（infer_grade_type）
- 阶梯项构建（build_items_for_row / build_item）
- 政策构建（build_policies_from_block）
- category 推断（category_from_grade_name / infer_block_category）

所有方法均依赖 SheetScanner 完成底层读取，自己只负责"语义"层面。
"""

from decimal import Decimal
from typing import Any, Dict, List, Optional, Tuple

from openpyxl.worksheet.worksheet import Worksheet

from app.parsers.steel_price_policy.boundary_parser import _Boundary, _parse_boundary
from app.parsers.steel_price_policy.models import ParsedLadderItem, ParsedPolicy
from app.parsers.steel_price_policy.models.policy import (
    GRADE_TYPE_CATEGORY,
    GRADE_TYPE_GRADE,
    GRADE_TYPE_THICKNESS_RANGE,
    GRADE_TYPE_WIDTH_RANGE,
    POLICY_TYPE_THICKNESS,
    POLICY_TYPE_WIDTH,
)
from app.parsers.steel_price_policy.value_parser import _to_int_price
from app.parsers.steel_price_policy.handlers.constants import GRADE_LABELS
from .sheet_scanner import SheetScanner

_WIDTH_RANGE_KEYWORDS = ("≤", "≥", "＜", "＞", "W", "宽")
_CATEGORY_NAMES = frozenset({"冷轧钢", "镀锌钢", "酸洗钢", "热轧钢"})


class PolicyAssembler:
    """政策装配器：封装业务语义解析逻辑。"""

    def __init__(self, scanner: Optional[SheetScanner] = None):
        self.scanner = scanner or SheetScanner()

    # ------------------------------------------------------------------ #
    # 块定位（5步 / 3步定位协议）
    # ------------------------------------------------------------------ #

    def locate_thickness_block(
        self,
        sheet: Worksheet,
        label_row: int,
        block_end: int,
        max_col: int,
    ) -> Optional[Dict[str, Any]]:
        """执行厚度块的 5 步定位，返回块描述或 None。"""
        title_row = self.scanner.find_thickness_title_row(
            sheet, label_row, block_end, scan_cols=list(range(2, max_col + 1))
        )
        if title_row is None:
            return None

        header_row = self.scanner.find_header_row(
            sheet, title_row + 1, block_end, scan_cols=list(range(2, max_col + 1))
        )
        if header_row is None:
            return None

        price_cols = self.scanner.collect_price_cols(
            sheet, header_row, start_col=2, end_col=max_col
        )
        if not price_cols:
            return None

        data_end = self.scanner.trim_block_end(sheet, header_row, block_end, max_col)
        base_thickness = self.scanner.infer_base_thickness(
            sheet, header_row, label_row, max_col
        )
        category = self.scanner.infer_category(sheet, label_row, max_col)

        return {
            "category_name": category,
            "policy_type": "THICKNESS",
            "grade_col": 1,
            "header_row": header_row,
            "price_cols": price_cols,
            "data_row_start": header_row + 1,
            "data_row_end": data_end,
            "base_thickness": base_thickness,
            "grade_blacklist": GRADE_LABELS,
        }

    def locate_width_block(
        self,
        sheet: Worksheet,
        title_row: int,
        block_end: int,
        grade_col: int,
        data_col_start: int,
        max_col: int,
        *,
        category: str = "冷轧钢",
        header_row_offset: int = 1,
        scan_cols: Optional[List[int]] = None,
    ) -> Optional[Dict[str, Any]]:
        """执行宽度块的 3 步定位，返回块描述。"""
        header_row = title_row + header_row_offset
        if header_row > block_end:
            return None

        price_cols = self.scanner.collect_price_cols(
            sheet, header_row, start_col=data_col_start, end_col=max_col
        )
        if not price_cols:
            return None

        data_start = header_row + 1
        effective_scan_cols = list(scan_cols or range(1, max_col + 1))
        if grade_col not in effective_scan_cols:
            effective_scan_cols.append(grade_col)
        data_end = self.scanner.trim_block_end(
            sheet, header_row, block_end, max_col,
            scan_cols=effective_scan_cols,
            skip_grade_keywords=True,
            grade_col=grade_col,
        )

        return {
            "category_name": category,
            "policy_type": "WIDTH",
            "grade_col": grade_col,
            "header_row": header_row,
            "price_cols": price_cols,
            "data_row_start": data_start,
            "data_row_end": data_end,
            "base_thickness": None,
            "grade_blacklist": GRADE_LABELS,
        }

    # ------------------------------------------------------------------ #
    # 边界收集与推断
    # ------------------------------------------------------------------ #

    def collect_boundaries(
        self,
        sheet: Worksheet,
        header_row: int,
        price_cols: List[int],
    ) -> Dict[int, _Boundary]:
        """从列头行收集所有边界定义。"""
        result = {}
        for c in price_cols:
            v = sheet.cell(row=header_row, column=c).value
            b = _parse_boundary(v)
            if b is not None:
                result[c] = b
        return result

    def infer_upper_bound(
        self,
        boundary: _Boundary,
        sorted_cols: List[int],
        col_idx: int,
        boundaries: Dict[int, _Boundary],
    ) -> _Boundary:
        """推断完整边界：
        1. 若 min_val is None 且是第一列，设为 min=0
        2. 若 min_val is None 且非第一列，从前一列的 max_val 推断
        3. 若 max_val is None 且有下一列，从下一列的 min_val 推断
        4. 若 max_val is None 且是最后一列，设为 max=∞（None 表示无上限）
        """
        min_val = boundary.min_val
        min_op = boundary.min_operator
        max_val = boundary.max_val
        max_op = boundary.max_operator

        # 1. 推断下界
        if min_val is None:
            if col_idx == 0:
                # 第一列，宽度从 0 开始
                min_val = Decimal("0")
                min_op = ">="
            else:
                # 从前一列的 max 推断
                prev_c = sorted_cols[col_idx - 1]
                if prev_c in boundaries:
                    prev_boundary = boundaries[prev_c]
                    if prev_boundary.max_val is not None:
                        min_val = prev_boundary.max_val
                        # 操作符：若前一列是 <=，则 min 用 >；若是 <，则 min 用 >=
                        min_op = ">" if prev_boundary.max_operator == "<=" else ">="

        # 2. 推断上界
        if max_val is None:
            if col_idx < len(sorted_cols) - 1:
                # 有下一列，从下一列的 min_val 推断
                next_c = sorted_cols[col_idx + 1]
                if next_c in boundaries:
                    next_boundary = boundaries[next_c]
                    if next_boundary.min_val is not None:
                        max_val = next_boundary.min_val
                        max_op = "<"
            # else: 最后一列且无上界，max_val 保持 None（表示 ∞）

        return _Boundary(
            min_val=min_val,
            min_operator=min_op,
            max_val=max_val,
            max_operator=max_op,
        )

    # ------------------------------------------------------------------ #
    # grade_type 推断
    # ------------------------------------------------------------------ #

    def infer_grade_type(self, grade: str, policy_type: str) -> str:
        """推断 grade 的语义类型。"""
        if policy_type == "THICKNESS":
            return GRADE_TYPE_GRADE
        grade_stripped = grade.strip()
        if any(kw in grade_stripped for kw in _WIDTH_RANGE_KEYWORDS):
            return GRADE_TYPE_WIDTH_RANGE
        if grade_stripped in _CATEGORY_NAMES:
            return GRADE_TYPE_CATEGORY
        return GRADE_TYPE_GRADE

    # ------------------------------------------------------------------ #
    # 阶梯项构建
    # ------------------------------------------------------------------ #

    def build_items_for_row(
        self,
        sheet: Worksheet,
        r: int,
        price_cols: List[int],
        boundaries: Dict[int, _Boundary],
        policy_type: str = "THICKNESS",
        sort_order: int = 0,
    ) -> Tuple[List[ParsedLadderItem], Any]:
        """构建行内阶梯明细项列表。返回 (items, extra)。"""
        items: List[ParsedLadderItem] = []
        sorted_cols = sorted(price_cols)
        for i, c in enumerate(sorted_cols):
            if c not in boundaries:
                continue

            price_cell = sheet.cell(row=r, column=c)
            if price_cell.value is None:
                start_row = self.scanner.find_start_row(sheet, r, c)
                if start_row is not None:
                    start_val = sheet.cell(row=start_row, column=c).value
                    if start_val is not None:
                        price = _to_int_price(start_val)
                    else:
                        price = self.scanner.find_value_above(sheet, start_row, c)
                else:
                    price = 0
            else:
                price = _to_int_price(price_cell.value)

            boundary = self.infer_upper_bound(
                boundaries[c], sorted_cols, i, boundaries
            )

            if policy_type == "THICKNESS":
                items.append(self._make_thickness_item(boundary, price, sort_order))
            else:
                items.append(self._make_width_item(boundary, price, sort_order))

        return items, None

    @staticmethod
    def _make_thickness_item(
        boundary: _Boundary, price: int, sort_order: int
    ) -> ParsedLadderItem:
        return ParsedLadderItem(
            target_sub_group=None,
            thick_min=boundary.min_val,
            thick_min_operator=boundary.min_operator,
            thick_max=boundary.max_val,
            thick_max_operator=boundary.max_operator,
            width_min=None,
            width_min_operator=">=",
            width_max=None,
            width_max_operator="<",
            price_adjustment=price,
            is_base=False,
            sort_order=sort_order,
        )

    @staticmethod
    def _make_width_item(
        boundary: _Boundary, price: int, sort_order: int
    ) -> ParsedLadderItem:
        return ParsedLadderItem(
            target_sub_group=None,
            thick_min=None,
            thick_min_operator=">=",
            thick_max=None,
            thick_max_operator="<",
            width_min=boundary.min_val,
            width_min_operator=boundary.min_operator,
            width_max=boundary.max_val,
            width_max_operator=boundary.max_operator,
            price_adjustment=price,
            is_base=False,
            sort_order=sort_order,
        )

    def build_item(
        self,
        boundary: _Boundary,
        price: int,
        width_boundary: Optional[_Boundary] = None,
    ) -> ParsedLadderItem:
        """构建阶梯明细项（单维度或双维度）。"""
        if width_boundary is not None:
            return ParsedLadderItem(
                target_sub_group=None,
                thick_min=boundary.min_val,
                thick_min_operator=boundary.min_operator,
                thick_max=boundary.max_val,
                thick_max_operator=boundary.max_operator,
                width_min=width_boundary.min_val,
                width_min_operator=width_boundary.min_operator,
                width_max=width_boundary.max_val,
                width_max_operator=width_boundary.max_operator,
                price_adjustment=price,
                is_base=False,
                sort_order=0,
            )
        else:
            return ParsedLadderItem(
                target_sub_group=None,
                thick_min=boundary.min_val,
                thick_min_operator=boundary.min_operator,
                thick_max=boundary.max_val,
                thick_max_operator=boundary.max_operator,
                width_min=None,
                width_min_operator=">=",
                width_max=None,
                width_max_operator="<",
                price_adjustment=price,
                is_base=False,
                sort_order=0,
            )

    # ------------------------------------------------------------------ #
    # 政策构建（单块 → ParsedPolicy 列表）
    # ------------------------------------------------------------------ #

    def build_policies_from_block(
        self,
        sheet: Worksheet,
        steel_mill: str,
        block: Dict[str, Any],
    ) -> List[ParsedPolicy]:
        """给定块描述，构建 ParsedPolicy 列表。"""
        grade_col = block.get("grade_col")
        header_row = block["header_row"]
        price_cols = list(block["price_cols"])
        data_start = block["data_row_start"]
        data_end = block["data_row_end"]
        category_name = block.get("category_name", "冷轧钢")
        base_thickness = block.get("base_thickness")
        grade_blacklist = block.get("grade_blacklist", GRADE_LABELS)
        policy_type = block.get("policy_type", POLICY_TYPE_THICKNESS)

        boundaries = self.collect_boundaries(sheet, header_row, price_cols)
        if not boundaries:
            return []

        sorted_price_cols = sorted(price_cols)
        policies: List[ParsedPolicy] = []
        last_grade: Optional[str] = None

        for r in range(data_start, data_end + 1):
            if grade_col is None:
                continue
            grade = self.scanner.text_or_none(sheet, r, grade_col)
            if not grade:
                grade = last_grade
            else:
                last_grade = grade
            if not grade:
                continue
            if grade in grade_blacklist:
                continue

            grade_type = self.infer_grade_type(grade, policy_type)
            items, _ = self.build_items_for_row(
                sheet, r, sorted_price_cols, boundaries, policy_type, sort_order=r
            )
            if not items:
                continue

            policies.append(
                ParsedPolicy(
                    steel_mill=steel_mill,
                    category_name=category_name,
                    grade_type=grade_type,
                    grade_value=grade,
                    policy_type=policy_type,
                    base_thickness=base_thickness,
                    items=items,
                )
            )
        return policies

    # ------------------------------------------------------------------ #
    # category 推断
    # ------------------------------------------------------------------ #

    _GALVANIZED_SUFFIXES: tuple[str, ...] = ("+Z", "+ZF", "+ZM", "+ZA", "+AZ")
    _PICKLED_SUFFIXES: tuple[str, ...] = ("+A",)
    _PICKLED_GRADES: frozenset[str] = frozenset({
        "SPHC", "SPHD", "SPHE", "SPHF",
        "SAPH370", "SAPH400", "SAPH440",
        "QStE340TM", "QStE380TM", "QStE420TM", "QStE460TM", "QStE500TM",
        "QSTE420TM", "QSTE460TM", "QSTE500TM",
        "B510L", "B590L",
    })

    def category_from_grade_name(self, grade: str) -> str:
        """按钢种名后缀/名称推断 category。"""
        if not grade:
            return "冷轧钢"
        for suf in self._GALVANIZED_SUFFIXES:
            if grade.endswith(suf):
                return "镀锌钢"
        if grade in self._PICKLED_GRADES:
            return "酸洗钢"
        for suf in self._PICKLED_SUFFIXES:
            if grade.endswith(suf):
                return "酸洗钢"
        return "冷轧钢"

    def infer_block_category(
        self,
        sheet: Worksheet,
        block: Dict[str, Any],
        *,
        max_col: int = 14,
    ) -> str:
        """综合'字面 hint'与'grade 后缀'推断块的 category。"""
        label_row = block.get("header_row", 0) - 1
        literal_cat = block.get("category_name") or "冷轧钢"
        if literal_cat != "冷轧钢":
            return literal_cat

        data_start = block["data_row_start"]
        data_end = block["data_row_end"]
        grade_col = block.get("grade_col", 1)
        galvanized = pickled = cold = 0
        for r in range(data_start, data_end + 1):
            v = sheet.cell(row=r, column=grade_col).value
            if v is None:
                continue
            t = str(v).strip()
            if not t or t in GRADE_LABELS:
                continue
            cat = self.category_from_grade_name(t)
            if cat == "镀锌钢":
                galvanized += 1
            elif cat == "酸洗钢":
                pickled += 1
            else:
                cold += 1

        if galvanized and not cold:
            return "镀锌钢"
        if pickled and not cold:
            return "酸洗钢"
        if galvanized and cold:
            if galvanized > cold:
                return "镀锌钢"
            if pickled > cold:
                return "酸洗钢"
            return "冷轧钢"
        if pickled and cold:
            return "酸洗钢" if pickled > cold else "冷轧钢"
        return "冷轧钢"