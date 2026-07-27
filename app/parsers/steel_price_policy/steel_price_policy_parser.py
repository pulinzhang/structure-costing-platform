"""钢价政策 Excel 文件解析器。

架构（编排器 + Sheet 处理器）：

    SteelPricePolicyExcelParser  ←  编排器：负责加载 Excel、聚合结果
        ↓
        parse_steel_price_policy_sheet(sheet, sheet_name)
        ↓
    find_handler(sheet_name).parse(sheet, steel_mill, warnings)
        ↓
    BaseSheetHandler.parse() → List[ParsedPolicy]

子模块:
    value_parser.py     单元格值解析（数字、字符串、str）
    boundary_parser.py  列头区间解析（≥0.5, ≤950mm, 1000-1600 等）
    handlers/           每个 Sheet 排版一个处理器
    models/             ParsedPolicy, ParsedLadderItem, SheetParseResult

约定：
- 一次只处理一张 Sheet：调用方（service）控制循环，便于持久化单张 Sheet
- 新增钢厂只需新增 handler，找不到处理器时回退到 DefaultSheetHandler
"""

import io
from pathlib import Path
from typing import List, Union

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.utils.logging_config import get_logger

from .handlers import find_handler
from .models import SteelPricePolicyParseResult


logger = get_logger("parser.steel_price_policy")


# ============================================================================
# 单 Sheet 解析（核心 API，一次处理一张 Sheet）
# ============================================================================


def parse_steel_price_policy_sheet(
    sheet: Worksheet,
    sheet_name: str,
    warnings: List[str],
) -> List:
    """解析一张 Sheet，返回 ParsedPolicy 列表。

    设计要点：
    - 一次只处理一张 Sheet（一个钢厂）
    - 调用方（service 层）负责 db.commit() 持久化
    - 调用方负责多 Sheet 的迭代
    - 新增钢厂时只需新增 handler，本函数无需修改

    Args:
        sheet: openpyxl Worksheet 对象
        sheet_name: Sheet 名字（== 钢厂名）
        warnings: 调用方的警告列表（用作告警输入）

    Returns:
        该 Sheet 包含的所有 ParsedPolicy
    """
    handler = find_handler(sheet_name)
    try:
        policies = handler.parse(sheet, sheet_name, warnings)
        logger.info(
            f"钢厂 [{sheet_name}] 解析完成，共 {len(policies)} 条政策"
        )
        return policies
    except Exception as e:
        msg = f"钢厂 [{sheet_name}] 解析失败: {e}"
        logger.error(msg)
        warnings.append(msg)
        return []


# ============================================================================
# 多 Sheet 编排（便捷 API，一次解析整张 Excel 并聚合）
# ============================================================================


def list_sheets(file: Union[str, Path, bytes]) -> List[str]:
    """列出 Excel 文件中所有 Sheet 名称（用于前端下拉选择）。"""
    wb, _ = _load_workbook_or_bytes(file)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def parse_steel_price_policy_excel(
    file: Union[str, Path, bytes],
    *,
    year: int,
    month: str,
) -> SteelPricePolicyParseResult:
    """整张 Excel 解析（便捷 API）。

    内部按 Sheet 依次调用 parse_steel_price_policy_sheet，
    并聚合到 SteelPricePolicyParseResult。

    等价于：
        result = SteelPricePolicyParseResult(year=..., month=...)
        for sheet in wb:
            policies = parse_steel_price_policy_sheet(sheet, sheet.name, result.warnings)
            result.policies.extend(policies)
    """
    parser = SteelPricePolicyExcelParser(year=year, month=month)
    return parser.parse(file)


# ============================================================================
# 编排器
# ============================================================================


class SteelPricePolicyExcelParser:
    """钢价政策 Excel 编排器（按 Sheet 依次解析，聚合结果）。"""

    def __init__(self, *, year: int, month: str):
        self.year = year
        self.month = month
        self.result = SteelPricePolicyParseResult()

    def parse(self, file: Union[str, Path, bytes]) -> SteelPricePolicyParseResult:
        logger.info(f"开始解析钢价政策 Excel（{self.year}年{self.month}）")
        self.result = SteelPricePolicyParseResult(
            year=self.year, month=self.month
        )

        wb, _ = _load_workbook_or_bytes(file)
        try:
            self.result.sheet_names = list(wb.sheetnames)
            # === 关键：按 Sheet 逐张处理 ===
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                policies = parse_steel_price_policy_sheet(
                    sheet, sheet_name, self.result.warnings
                )
                self.result.policies.extend(policies)
            self._log_summary()
            return self.result
        finally:
            wb.close()

    def _log_summary(self) -> None:
        n_t = sum(1 for p in self.result.policies if p.policy_type == "THICKNESS")
        n_w = sum(1 for p in self.result.policies if p.policy_type == "WIDTH")
        logger.info("=== 钢价政策解析完成 ===")
        logger.info(
            f"钢厂 sheet 数: {len(self.result.sheet_names)}, "
            f"厚度政策: {n_t}, 宽度政策: {n_w}, "
            f"警告: {len(self.result.warnings)} 条"
        )


# ============================================================================
# Excel 加载辅助（兼容 path / bytes 两种入参）
# ============================================================================


def _load_workbook_or_bytes(file):
    """加载 workbook。"""
    if isinstance(file, (str, Path)):
        path = Path(file)
        wb = load_workbook(str(path), data_only=True, read_only=True)
        return wb, path
    if isinstance(file, bytes):
        wb = load_workbook(io.BytesIO(file), data_only=True, read_only=True)
        return wb, None
    raise TypeError(f"file must be path or bytes, got {type(file).__name__}")
