"""
车身供货 BOM Excel 解析（编排入口）

流程：
1. 加载工作簿
2. 解析头信息 → meta_parser
3. 解析树节点 → node_parser
4. 归属图示 → sketch_assigner
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.parsers.bom.meta_parser import parse_bom_meta
from app.parsers.bom.models import BomParsedMeta, BomParsedNode, BomParseResult
from app.parsers.bom.node_parser import parse_bom_nodes
from app.parsers.bom.sketch_assigner import assign_sketch_images

__all__ = [
    "BomParseResult",
    "BomParsedMeta",
    "BomParsedNode",
    "parse_bom_excel",
]


def parse_bom_excel(source: str | Path | bytes, *, sheet_name: str | None = None) -> BomParseResult:
    """解析车身件材料清单 Excel，返回头信息 + 扁平节点列表（含 parent_index）。"""
    if isinstance(source, bytes | bytearray):
        file_bytes = bytes(source)
        wb = load_workbook(BytesIO(file_bytes), data_only=True)
    else:
        path = Path(source)
        file_bytes = path.read_bytes()
        wb = load_workbook(str(path), data_only=True)

    if not wb.worksheets:
        raise ValueError("Excel 中没有工作表")
    if sheet_name:
        matching = [sheet for sheet in wb.worksheets if sheet.title == sheet_name]
        if not matching:
            raise ValueError(f"Excel 中没有工作表 {sheet_name}")
        ws = matching[0]
    else:
        ws = wb.worksheets[0]
    meta = parse_bom_meta(ws)
    nodes, row_l0, row_part_sketch = parse_bom_nodes(ws)
    image_count = assign_sketch_images(
        file_bytes, nodes, row_l0, row_part_sketch, sheet_name=ws.title
    )
    return BomParseResult(
        meta=meta,
        nodes=nodes,
        sheet_name=ws.title,
        image_count=image_count,
    )
