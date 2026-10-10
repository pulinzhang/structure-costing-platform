"""生产样例版式：双行表头 + 固定列。"""

from __future__ import annotations

from decimal import Decimal

from openpyxl.worksheet.worksheet import Worksheet

from app.common.workbook_reject import NO_DATA_ROWS, invalid_amount

from .cells import (
    PROD_COL,
    build_row,
    cell_text,
    is_skip_row,
    normalize_header,
    parse_int,
    parse_money,
)
from .types import PackagingTransportFeeParsedRow, PackagingTransportFeeParseResult


def looks_like_production_header(sheet: Worksheet, row_idx: int) -> bool:
    texts = [
        normalize_header(cell_text(sheet.cell(row_idx, c).value))
        for c in range(1, min(28, (sheet.max_column or 1) + 1))
    ]
    joined = "|".join(texts)
    has_part = any(t in ("料号", "零件号") for t in texts)
    has_pack = any("包装类型" in t for t in texts)
    has_total = any("包装运输费合计" in t for t in texts) or any("包装运输费" in t for t in texts)
    # 生产样例：料号 + 包装类型，合计可能在同行或下一行合并区
    if has_part and has_pack:
        return True
    return has_part and has_total and "单价" not in joined


def parse_production_sheet(sheet: Worksheet, header_row: int) -> PackagingTransportFeeParseResult:
    rows: list[PackagingTransportFeeParsedRow] = []
    errors: list[str] = []
    seen: set[str] = set()
    data_start = header_row + 1
    # 若下一行是子表头（长/宽/高），再跳一行
    next_texts = [
        normalize_header(cell_text(sheet.cell(header_row + 1, c).value)) for c in range(6, 15)
    ]
    if any(t.startswith("长") or t == "器具单价" or "总投入" in t for t in next_texts):
        data_start = header_row + 2

    for row_idx in range(data_start, (sheet.max_row or data_start) + 1):
        part_number = cell_text(sheet.cell(row_idx, PROD_COL["part_number"]).value)
        part_name = cell_text(sheet.cell(row_idx, PROD_COL["part_name"]).value)
        if is_skip_row(part_number, part_name):
            continue
        if not part_number:
            errors.append(f"第 {row_idx} 行「料号」是空的。请填写料号。")
            continue
        if not (part_number[0].isdigit() or part_number[0].isalpha()):
            continue
        if part_number in seen:
            errors.append(f"第 {row_idx} 行料号「{part_number}」和上面重复。每个料号只保留一行。")
            continue

        pkg_fee = parse_money(sheet.cell(row_idx, PROD_COL["unit_packaging_fee"]).value)
        trans_fee = parse_money(sheet.cell(row_idx, PROD_COL["unit_transport_fee"]).value)
        maint_fee = parse_money(
            sheet.cell(row_idx, PROD_COL["fixture_maint_fee"]).value,
            allow_empty=True,
        )
        other_fee = parse_money(
            sheet.cell(row_idx, PROD_COL["other_fee"]).value,
            allow_empty=True,
        )
        total = parse_money(sheet.cell(row_idx, PROD_COL["packaging_transport_total"]).value)
        if pkg_fee is None:
            errors.append(
                invalid_amount(
                    row_idx,
                    "单件包装费",
                    sheet.cell(row_idx, PROD_COL["unit_packaging_fee"]).value,
                    "0.9800",
                )
            )
            continue
        if trans_fee is None:
            errors.append(
                invalid_amount(
                    row_idx,
                    "单件运输费",
                    sheet.cell(row_idx, PROD_COL["unit_transport_fee"]).value,
                    "1.5500",
                )
            )
            continue
        if total is None:
            errors.append(
                invalid_amount(
                    row_idx,
                    "包装运输费合计",
                    sheet.cell(row_idx, PROD_COL["packaging_transport_total"]).value,
                    "3.7194",
                )
            )
            continue

        seen.add(part_number)
        rows.append(
            build_row(
                row_number=row_idx,
                part_number=part_number,
                part_name=part_name or None,
                packaging_type=cell_text(sheet.cell(row_idx, PROD_COL["packaging_type"]).value)
                or None,
                pieces_per_pack=parse_int(sheet.cell(row_idx, PROD_COL["pieces_per_pack"]).value),
                fixture_length_mm=parse_int(
                    sheet.cell(row_idx, PROD_COL["fixture_length_mm"]).value
                ),
                fixture_width_mm=parse_int(sheet.cell(row_idx, PROD_COL["fixture_width_mm"]).value),
                fixture_height_mm=parse_int(
                    sheet.cell(row_idx, PROD_COL["fixture_height_mm"]).value
                ),
                oem_stock_qty=parse_int(sheet.cell(row_idx, PROD_COL["oem_stock_qty"]).value),
                transit_qty=parse_int(sheet.cell(row_idx, PROD_COL["transit_qty"]).value),
                company_stock_qty=parse_int(
                    sheet.cell(row_idx, PROD_COL["company_stock_qty"]).value
                ),
                total_fixture_qty=parse_int(
                    sheet.cell(row_idx, PROD_COL["total_fixture_qty"]).value
                ),
                fixture_unit_price=parse_money(
                    sheet.cell(row_idx, PROD_COL["fixture_unit_price"]).value,
                    allow_empty=True,
                ),
                fixture_invest_amount=parse_money(
                    sheet.cell(row_idx, PROD_COL["fixture_invest_amount"]).value,
                    allow_empty=True,
                ),
                unit_packaging_fee=pkg_fee,
                round_trip_transport_price=parse_money(
                    sheet.cell(row_idx, PROD_COL["round_trip_transport_price"]).value,
                    allow_empty=True,
                ),
                unit_transport_fee=trans_fee,
                fixture_maint_fee=maint_fee or Decimal("0"),
                other_fee=other_fee or Decimal("0"),
                packaging_transport_total=total,
                remark=None,
            )
        )

    if not rows and not errors:
        errors.append(NO_DATA_ROWS)
    return PackagingTransportFeeParseResult(
        rows=rows,
        errors=errors,
        sheet_name=getattr(sheet, "title", None),
    )
