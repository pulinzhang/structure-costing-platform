"""
BOM 包装运输费 Excel 解析器（纯函数，无副作用）

支持两种版式：
1) 生产样例：双行表头（料号 / 包装类型 / … / 包装运输费合计），基地/年/项目由表单传入
2) 平台扁平模板：单行表头（零件号、单件包装费、…）

版式实现见 packaging_transport_fee/。对外仍从此文件导入 parse_packaging_transport_fee_excel。
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.common.workbook_reject import EMPTY_FILE, NO_WORKSHEET, UNREADABLE_EXCEL
from app.parsers.bom.packaging_transport_fee.cells import (
    H_PKG_FEE,
    PROD_COL,
    cell_text,
    normalize_header,
)
from app.parsers.bom.packaging_transport_fee.flat import find_flat_header, parse_flat_sheet
from app.parsers.bom.packaging_transport_fee.production import (
    looks_like_production_header,
    parse_production_sheet,
)
from app.parsers.bom.packaging_transport_fee.types import (
    PackagingTransportFeeParsedRow,
    PackagingTransportFeeParseResult,
)


def parse_packaging_transport_fee_excel(
    file_content: bytes,
) -> PackagingTransportFeeParseResult:
    """解析包装运输费 Excel（.xlsx）。"""
    if not file_content:
        return PackagingTransportFeeParseResult(rows=[], errors=[EMPTY_FILE])

    try:
        workbook = load_workbook(BytesIO(file_content), data_only=True, read_only=True)
    except Exception:
        return PackagingTransportFeeParseResult(rows=[], errors=[UNREADABLE_EXCEL])

    try:
        if not workbook.worksheets:
            return PackagingTransportFeeParseResult(rows=[], errors=[NO_WORKSHEET])
        sheet = workbook.worksheets[0]

        # 优先识别生产双行表头
        max_scan = min(sheet.max_row or 1, 10)
        for row_idx in range(1, max_scan + 1):
            if looks_like_production_header(sheet, row_idx):
                # 若同行已有「单件包装费」文字且列映射完整，走扁平
                flat = find_flat_header(sheet)
                if (
                    flat
                    and flat[0] == row_idx
                    and H_PKG_FEE in flat[1]
                    and flat[1][H_PKG_FEE] != PROD_COL["unit_packaging_fee"]
                ):
                    return parse_flat_sheet(sheet, flat[0], flat[1])
                # 生产样例：料号在 B，包装费在 O（15）
                cell_o = normalize_header(
                    cell_text(sheet.cell(row_idx, PROD_COL["unit_packaging_fee"]).value)
                )
                if "单件包装费" in cell_o or cell_o == "":
                    return parse_production_sheet(sheet, row_idx)

        flat = find_flat_header(sheet)
        if flat:
            return parse_flat_sheet(sheet, flat[0], flat[1])

        return PackagingTransportFeeParseResult(
            rows=[],
            errors=[
                "未找到表头。请下载本页模板，第一行保留「零件号」「单件包装费」"
                "「单件运输费」「包装运输费合计」；或使用生产样例（含「料号」「包装类型」）。列名不要改。"
            ],
        )
    finally:
        workbook.close()


__all__ = [
    "PackagingTransportFeeParsedRow",
    "PackagingTransportFeeParseResult",
    "parse_packaging_transport_fee_excel",
]
