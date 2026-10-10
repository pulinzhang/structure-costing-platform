"""
BOM 年度零件价 Excel 解析器（纯函数，无副作用）

约定列：零件号 | 零件名称 | 规格 | 采购单位 | 采购单价 | 转换系数 | 单位 | 单价 | 原供货类型 | 供应商 | 备注
基地 / 年度由上传表单传入，不从 Excel 读取。
只读第一张工作表；填写说明页忽略。
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.common.money import Q4, Q6, quantize_money
from app.common.workbook_reject import (
    EMPTY_FILE,
    NO_DATA_ROWS,
    NO_WORKSHEET,
    UNREADABLE_EXCEL,
    describe_cell,
    invalid_amount,
)

HEADER_PART_NUMBER = "零件号"
HEADER_PART_NAME = "零件名称"
HEADER_SPECIFICATION = "规格"
HEADER_PURCHASE_UNIT = "采购单位"
HEADER_PURCHASE_UNIT_PRICE = "采购单价"
HEADER_CONVERSION_FACTOR = "转换系数"
HEADER_UNIT = "单位"
HEADER_UNIT_PRICE = "单价"
HEADER_SUPPLY_TYPE = "原供货类型"
HEADER_SUPPLIER = "供应商"
HEADER_REMARK = "备注"

DEFAULT_CONVERSION_FACTOR = Decimal("1")
_ZERO = Decimal("0")


@dataclass(frozen=True)
class AnnualPartPriceParsedRow:
    """解析后的一行年度零件价"""

    row_number: int
    part_number: str
    part_name: str | None
    specification: str | None
    purchase_unit: str | None
    purchase_unit_price: Decimal | None
    conversion_factor: Decimal
    unit: str
    unit_price: Decimal
    supply_type: str
    supplier_name: str | None
    remark: str | None


@dataclass(frozen=True)
class AnnualPartPriceParseResult:
    """解析结果"""

    rows: list[AnnualPartPriceParsedRow]
    errors: list[str]


def _round4(value: Decimal) -> Decimal:
    return quantize_money(value, Q4)


def parse_conversion_factor(value: object) -> Decimal:
    """转换系数：空则 1；须 > 0，最多 6 位小数（1 采购单位 = N 核价单位）。"""
    if value is None:
        return DEFAULT_CONVERSION_FACTOR
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text:
            return DEFAULT_CONVERSION_FACTOR
        if "/" in text:
            raise ValueError("转换系数不能填分数，请填小数，例如 25.4")
        value = text
    number = quantize_money(value, Q6)
    if number <= 0:
        raise ValueError("转换系数须大于 0")
    return number


def format_conversion_factor(value: Decimal) -> str:
    """展示用：量化到 6 位后去掉末尾 0，不强制补齐小数位。"""
    number = parse_conversion_factor(value)
    return f"{number:.6f}".rstrip("0").rstrip(".")


def parse_optional_purchase_unit_price(value: object) -> Decimal | None:
    """采购单价：空为未填；0 是有效报价。"""
    if value is None or value == "":
        return None
    if isinstance(value, str) and not value.strip():
        return None
    number = _parse_unit_price(value)
    if number is None:
        raise ValueError("采购单价无效")
    return number


def resolve_annual_unit_price(
    *,
    purchase_unit_price: Decimal | None,
    conversion_factor: Decimal,
    unit_price: Decimal | None,
) -> Decimal:
    """有采购单价时：核价单价 = 采购单价 ÷ 转换系数；单价 0/空视为占位。"""
    factor = parse_conversion_factor(conversion_factor)
    if purchase_unit_price is None:
        if unit_price is None:
            raise ValueError("单价无效")
        return _round4(unit_price)
    derived = _round4(purchase_unit_price / factor)
    if unit_price is None or unit_price == _ZERO:
        return derived
    filled = _round4(unit_price)
    if filled != derived:
        raise ValueError(
            f"单价与采购单价÷转换系数不一致（表里单价 {filled}，按采购单价算出 {derived}）。"
            "请改成同一个数，或把单价留空、填 0，让系统按采购单价计算。"
        )
    return derived


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_unit_price(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        if isinstance(value, Decimal):
            number = value
        elif isinstance(value, int | float):
            number = Decimal(str(value))
        else:
            text = str(value).strip().replace(",", "")
            if not text:
                return None
            number = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    if number < 0:
        return None
    return _round4(number)


def _find_header_row(sheet: Worksheet) -> tuple[int, dict[str, int]] | None:
    """定位表头行，返回 (行号, 列名→列号)。"""
    aliases = {
        HEADER_PART_NUMBER: {HEADER_PART_NUMBER, "零件号", "料号", "part_number"},
        HEADER_PART_NAME: {HEADER_PART_NAME, "零件名称", "名称", "part_name"},
        HEADER_SPECIFICATION: {HEADER_SPECIFICATION, "specification"},
        HEADER_PURCHASE_UNIT: {HEADER_PURCHASE_UNIT, "purchase_unit"},
        HEADER_PURCHASE_UNIT_PRICE: {HEADER_PURCHASE_UNIT_PRICE, "purchase_unit_price"},
        HEADER_CONVERSION_FACTOR: {HEADER_CONVERSION_FACTOR, "conversion_factor"},
        HEADER_UNIT: {HEADER_UNIT, "单位", "unit"},
        HEADER_UNIT_PRICE: {
            HEADER_UNIT_PRICE,
            "单价",
            "价格",
            "单位价格",
            "unit_price",
        },
        HEADER_SUPPLY_TYPE: {
            HEADER_SUPPLY_TYPE,
            "供货类型",
            "原供货类型",
            "supply_type",
        },
        HEADER_SUPPLIER: {HEADER_SUPPLIER, "供应商", "供方", "supplier_name"},
        HEADER_REMARK: {HEADER_REMARK, "备注", "remark"},
    }
    max_scan = min(sheet.max_row or 1, 20)
    max_col = min(sheet.max_column or 1, 20)

    for row_idx in range(1, max_scan + 1):
        mapping: dict[str, int] = {}
        for col_idx in range(1, max_col + 1):
            text = _cell_text(sheet.cell(row_idx, col_idx).value)
            if not text:
                continue
            for field, names in aliases.items():
                if text in names and field not in mapping:
                    mapping[field] = col_idx
        if (
            HEADER_PART_NUMBER in mapping
            and HEADER_UNIT_PRICE in mapping
            and HEADER_SUPPLY_TYPE in mapping
        ):
            return row_idx, mapping
    return None


def parse_annual_part_price_excel(file_content: bytes) -> AnnualPartPriceParseResult:
    """
    解析年度零件价 Excel。
    仅支持 .xlsx（由调用方校验扩展名）。只读 worksheets[0]。
    """
    if not file_content:
        return AnnualPartPriceParseResult(rows=[], errors=[EMPTY_FILE])

    try:
        workbook = load_workbook(BytesIO(file_content), data_only=True, read_only=True)
    except Exception:
        return AnnualPartPriceParseResult(rows=[], errors=[UNREADABLE_EXCEL])

    try:
        if not workbook.worksheets:
            return AnnualPartPriceParseResult(rows=[], errors=[NO_WORKSHEET])
        sheet = workbook.worksheets[0]

        header = _find_header_row(sheet)
        if not header:
            return AnnualPartPriceParseResult(
                rows=[],
                errors=[
                    "未找到表头。请下载模板，第一行列名保持为："
                    f"{HEADER_PART_NUMBER} / {HEADER_PART_NAME} / "
                    f"{HEADER_SPECIFICATION} / {HEADER_PURCHASE_UNIT} / "
                    f"{HEADER_PURCHASE_UNIT_PRICE} / {HEADER_CONVERSION_FACTOR} / "
                    f"{HEADER_UNIT} / {HEADER_UNIT_PRICE} / {HEADER_SUPPLY_TYPE} / "
                    f"{HEADER_SUPPLIER} / {HEADER_REMARK}"
                ],
            )

        header_row, columns = header
        part_number_col = columns[HEADER_PART_NUMBER]
        part_name_col = columns.get(HEADER_PART_NAME)
        specification_col = columns.get(HEADER_SPECIFICATION)
        purchase_unit_col = columns.get(HEADER_PURCHASE_UNIT)
        purchase_price_col = columns.get(HEADER_PURCHASE_UNIT_PRICE)
        factor_col = columns.get(HEADER_CONVERSION_FACTOR)
        unit_col = columns.get(HEADER_UNIT)
        price_col = columns[HEADER_UNIT_PRICE]
        supply_type_col = columns[HEADER_SUPPLY_TYPE]
        supplier_col = columns.get(HEADER_SUPPLIER)
        remark_col = columns.get(HEADER_REMARK)

        rows: list[AnnualPartPriceParsedRow] = []
        errors: list[str] = []
        seen_parts: set[str] = set()

        for row_idx in range(header_row + 1, (sheet.max_row or header_row) + 1):
            part_number = _cell_text(sheet.cell(row_idx, part_number_col).value)
            part_name = (
                _cell_text(sheet.cell(row_idx, part_name_col).value) if part_name_col else ""
            )
            specification = (
                _cell_text(sheet.cell(row_idx, specification_col).value)
                if specification_col
                else ""
            )
            purchase_unit_raw = (
                _cell_text(sheet.cell(row_idx, purchase_unit_col).value)
                if purchase_unit_col
                else ""
            )
            purchase_price_raw = (
                sheet.cell(row_idx, purchase_price_col).value if purchase_price_col else None
            )
            factor_raw = sheet.cell(row_idx, factor_col).value if factor_col else None
            unit_raw = _cell_text(sheet.cell(row_idx, unit_col).value) if unit_col else ""
            price_raw = sheet.cell(row_idx, price_col).value
            supply_type = _cell_text(sheet.cell(row_idx, supply_type_col).value)
            supplier_raw = (
                _cell_text(sheet.cell(row_idx, supplier_col).value) if supplier_col else ""
            )
            remark_raw = _cell_text(sheet.cell(row_idx, remark_col).value) if remark_col else ""

            if not part_number and (price_raw is None or price_raw == ""):
                continue

            if not part_number:
                errors.append(f"第 {row_idx} 行「零件号」是空的。请填写零件号。")
                continue

            if not supply_type:
                errors.append(f"第 {row_idx} 行「原供货类型」是空的。请填写供货类型。")
                continue

            if len(supply_type) > 50:
                errors.append(f"第 {row_idx} 行「原供货类型」超过 50 个字。请缩短到 50 字以内。")
                continue

            try:
                purchase_unit_price = parse_optional_purchase_unit_price(purchase_price_raw)
            except ValueError:
                errors.append(invalid_amount(row_idx, "采购单价", purchase_price_raw, "12.5000"))
                continue

            try:
                conversion_factor = parse_conversion_factor(factor_raw)
            except (InvalidOperation, ValueError) as exc:
                errors.append(
                    f"第 {row_idx} 行「转换系数」填了「{describe_cell(factor_raw)}」。{exc}"
                )
                continue

            unit_price_raw = _parse_unit_price(price_raw)
            if unit_price_raw is None and purchase_unit_price is None:
                shown = describe_cell(price_raw)
                if shown == "空":
                    errors.append(
                        f"第 {row_idx} 行「单价」是空的，也没有采购单价。"
                        "请填单价，或改填采购单价（转换系数空着则按 1 计算）。"
                    )
                else:
                    errors.append(
                        f"第 {row_idx} 行「单价」填了「{shown}」，不是有效金额，也没有采购单价。"
                        "请改成数字，例如 1.2500；或改填采购单价。"
                    )
                continue

            try:
                unit_price = resolve_annual_unit_price(
                    purchase_unit_price=purchase_unit_price,
                    conversion_factor=conversion_factor,
                    unit_price=unit_price_raw,
                )
            except ValueError as exc:
                errors.append(f"第 {row_idx} 行：{exc}")
                continue

            if part_number in seen_parts:
                errors.append(
                    f"第 {row_idx} 行零件号「{part_number}」和上面重复。每个零件号只保留一行。"
                )
                continue
            seen_parts.add(part_number)

            unit = unit_raw or "pcs"
            if unit.upper() == "PCS":
                unit = "pcs"
            purchase_unit = purchase_unit_raw[:20] if purchase_unit_raw else None
            if purchase_unit and purchase_unit.upper() == "PCS":
                purchase_unit = "pcs"

            rows.append(
                AnnualPartPriceParsedRow(
                    row_number=row_idx,
                    part_number=part_number[:50],
                    part_name=(part_name[:200] if part_name else None),
                    specification=(specification[:100] if specification else None),
                    purchase_unit=purchase_unit,
                    purchase_unit_price=purchase_unit_price,
                    conversion_factor=conversion_factor,
                    unit=unit[:20],
                    unit_price=unit_price,
                    supply_type=supply_type[:50],
                    supplier_name=(supplier_raw[:100] if supplier_raw else None),
                    remark=remark_raw or None,
                )
            )
    finally:
        workbook.close()

    if not rows and not errors:
        errors.append(NO_DATA_ROWS)

    return AnnualPartPriceParseResult(rows=rows, errors=errors)
