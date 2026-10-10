"""采购对象供应商导入 Parser。只读第一张工作表。"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

HEADER_PART_NUMBER = "零件号"
HEADER_PART_NAME = "零件名称"
HEADER_INTERNAL_NUMBER = "内部号"
HEADER_OBJECT_KIND = "对象种类"
HEADER_CONFIRM = "确认"
HEADER_AUDIT = "审核"
HEADER_SUPPLIER = "供应商"

_REQUIRED_HEADERS = (
    HEADER_PART_NUMBER,
    HEADER_SUPPLIER,
)


@dataclass(frozen=True)
class ProcurementObjectSupplierParsedRow:
    row_number: int
    part_number: str
    supplier_name: str | None


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_header_row(sheet: Worksheet) -> tuple[int, dict[str, int]] | None:
    wanted = {HEADER_PART_NUMBER, HEADER_SUPPLIER}
    max_scan = min(sheet.max_row or 1, 10)
    max_col = min(sheet.max_column or 1, 20)
    for row_idx in range(1, max_scan + 1):
        mapping: dict[str, int] = {}
        for col_idx in range(1, max_col + 1):
            text = _cell_text(sheet.cell(row_idx, col_idx).value)
            if text in wanted and text not in mapping:
                mapping[text] = col_idx
        if all(name in mapping for name in _REQUIRED_HEADERS):
            return row_idx, mapping
    return None


def parse_procurement_object_supplier_excel(
    file_content: bytes,
) -> list[ProcurementObjectSupplierParsedRow]:
    workbook = load_workbook(BytesIO(file_content), data_only=True, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        found = _find_header_row(sheet)
        if found is None:
            raise ValueError("未找到表头，请使用下载模板第一页（零件号 / 供应商）")
        header_row, columns = found
        rows: list[ProcurementObjectSupplierParsedRow] = []
        for row_idx in range(header_row + 1, (sheet.max_row or header_row) + 1):
            part_number = _cell_text(sheet.cell(row_idx, columns[HEADER_PART_NUMBER]).value)
            supplier_name = _cell_text(sheet.cell(row_idx, columns[HEADER_SUPPLIER]).value) or None
            if not part_number and supplier_name is None:
                continue
            if not part_number:
                raise ValueError(f"Excel 第{row_idx}行：零件号不能为空")
            rows.append(
                ProcurementObjectSupplierParsedRow(
                    row_number=row_idx,
                    part_number=part_number,
                    supplier_name=supplier_name,
                )
            )
        return rows
    finally:
        workbook.close()
