"""
焊接工费单价 Excel 解析器（纯函数，无副作用）

约定四列：名称 | 单位 | 费用 | 备注
（「费用」列表示单位工费单价）
示例：
  焊点 | 个 | 0.0999 |
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.common.bom.welding_standard_process import validate_welding_unit_price_catalog
from app.common.money import Q4, quantize_money
from app.common.workbook_reject import (
    EMPTY_FILE,
    NO_DATA_ROWS,
    NO_WORKSHEET,
    UNREADABLE_EXCEL,
    invalid_amount,
)

HEADER_NAME = "名称"
HEADER_UNIT = "单位"
HEADER_UNIT_PRICE = "费用"
HEADER_REMARK = "备注"


@dataclass(frozen=True)
class WeldingUnitPriceParsedRow:
    """解析后的一行焊接工费单价"""

    row_number: int
    item_name: str
    unit: str
    unit_price: Decimal
    remark: str | None


@dataclass(frozen=True)
class WeldingUnitPriceParseResult:
    """解析结果"""

    rows: list[WeldingUnitPriceParsedRow]
    errors: list[str]


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _round4(value: Decimal) -> Decimal:
    return quantize_money(value, Q4)


def _parse_unit_price(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        if isinstance(value, Decimal):
            number = value
        elif isinstance(value, int | float):
            number = Decimal(str(value))
        else:
            text = str(value).strip().replace(",", "")
            if not text:
                return None
            number = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    if number < 0:
        return None
    return _round4(number)


def _find_header_row(sheet: Worksheet) -> tuple[int, dict[str, int]] | None:
    """定位表头行，返回 (行号, 列名→列号)。"""
    aliases = {
        HEADER_NAME: {HEADER_NAME, "名称", "项目", "费用项目", "操作项目"},
        HEADER_UNIT: {HEADER_UNIT, "单位"},
        HEADER_UNIT_PRICE: {
            HEADER_UNIT_PRICE,
            "费用",
            "单价",
            "工费",
            "工费单价",
            "单位价格",
            "价格",
        },
        HEADER_REMARK: {HEADER_REMARK, "备注", "说明"},
    }
    max_scan = min(sheet.max_row or 1, 20)
    max_col = min(sheet.max_column or 1, 20)

    for row_idx in range(1, max_scan + 1):
        mapping: dict[str, int] = {}
        for col_idx in range(1, max_col + 1):
            text = _cell_text(sheet.cell(row_idx, col_idx).value)
            if not text:
                continue
            for field, names in aliases.items():
                if text in names and field not in mapping:
                    mapping[field] = col_idx
        if HEADER_NAME in mapping and HEADER_UNIT_PRICE in mapping:
            return row_idx, mapping
    return None


def parse_welding_unit_price_excel(file_content: bytes) -> WeldingUnitPriceParseResult:
    """
    解析焊接工费单价 Excel。
    仅支持 .xlsx（由调用方校验扩展名）。
    """
    if not file_content:
        return WeldingUnitPriceParseResult(rows=[], errors=[EMPTY_FILE])

    try:
        workbook = load_workbook(BytesIO(file_content), data_only=True, read_only=True)
    except Exception:
        return WeldingUnitPriceParseResult(rows=[], errors=[UNREADABLE_EXCEL])

    try:
        if not workbook.worksheets:
            return WeldingUnitPriceParseResult(rows=[], errors=[NO_WORKSHEET])
        sheet = workbook.worksheets[0]

        header = _find_header_row(sheet)
        if not header:
            return WeldingUnitPriceParseResult(
                rows=[],
                errors=[
                    "未找到表头。请下载模板，第一行列名保持为："
                    f"{HEADER_NAME} / {HEADER_UNIT} / {HEADER_UNIT_PRICE} / {HEADER_REMARK}"
                ],
            )

        header_row, columns = header
        name_col = columns[HEADER_NAME]
        unit_col = columns.get(HEADER_UNIT)
        price_col = columns[HEADER_UNIT_PRICE]
        remark_col = columns.get(HEADER_REMARK)

        rows: list[WeldingUnitPriceParsedRow] = []
        errors: list[str] = []
        seen_names: set[str] = set()

        for row_idx in range(header_row + 1, (sheet.max_row or header_row) + 1):
            name_raw = _cell_text(sheet.cell(row_idx, name_col).value)
            unit_raw = _cell_text(sheet.cell(row_idx, unit_col).value) if unit_col else ""
            price_raw = sheet.cell(row_idx, price_col).value
            remark_raw = _cell_text(sheet.cell(row_idx, remark_col).value) if remark_col else ""

            if not name_raw and (price_raw is None or price_raw == ""):
                continue

            if not name_raw:
                errors.append(
                    f"第 {row_idx} 行「名称」是空的。请按标准工艺目录填写，不要自造名称。"
                )
                continue

            unit_price = _parse_unit_price(price_raw)
            if unit_price is None:
                errors.append(invalid_amount(row_idx, "费用", price_raw, "0.0999"))
                continue

            unit = unit_raw or "个"
            if name_raw in seen_names:
                errors.append(f"第 {row_idx} 行名称「{name_raw}」和上面重复。每个名称只保留一行。")
                continue
            try:
                validate_welding_unit_price_catalog(name_raw, unit)
            except ValueError as exc:
                errors.append(f"第 {row_idx} 行：{exc}")
                continue
            seen_names.add(name_raw)

            rows.append(
                WeldingUnitPriceParsedRow(
                    row_number=row_idx,
                    item_name=name_raw,
                    unit=unit,
                    unit_price=unit_price,
                    remark=remark_raw or None,
                )
            )
    finally:
        workbook.close()

    if not rows and not errors:
        errors.append(NO_DATA_ROWS)

    return WeldingUnitPriceParseResult(rows=rows, errors=errors)
