"""解析「工程变更单（材质）」：只读形，不查库。"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.datetime import from_excel
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.common.engineering_change_gates import quoted
from app.common.period import MONTH_MAX, MONTH_MIN
from app.exceptions.base import BusinessException
from app.parsers.engineering_change.material_change.codes import (
    ECM_CELL_NOT_EXACT,
    ECM_CELL_TYPE,
    ECM_EFFECTIVE_FORMAT,
    ECM_EMPTY_ROW,
    ECM_EXTRA_CELL,
    ECM_FORMULA,
    ECM_HEADER,
    ECM_MERGED,
    ECM_SHEET_IMAGE,
    ECM_SHEET_REQUIRED,
    ECM_THICKNESS_FORMAT,
    HEADERS,
    SHEET_HELP,
    SHEET_ORDER,
    ecm_row,
    ecm_sheet,
)

_STRING_TYPES = frozenset({"s", "inlineStr", "str"})
_PERIOD_RE = re.compile(r"^(\d{4})年(\d{1,2})月$")
_THICKNESS_INDEXES = frozenset({8, 10})


@dataclass(slots=True)
class MaterialChangeParsedRow:
    excel_row: int
    location: str
    vehicle_model: str
    change_version: str
    effective_period_text: str
    effective_year: int
    effective_month: int
    stamping_supplier: str
    part_number: str
    part_name: str
    before_material: str
    before_thickness: str
    after_material: str
    after_thickness: str
    interchangeability: str


def parse_material_change_workbook(file_content: bytes) -> list[MaterialChangeParsedRow]:
    workbook = load_workbook(BytesIO(file_content), data_only=False)
    try:
        return _parse_open_workbook(workbook)
    finally:
        workbook.close()


def _parse_open_workbook(workbook: Workbook) -> list[MaterialChangeParsedRow]:
    names = list(workbook.sheetnames)
    if names != [SHEET_HELP, SHEET_ORDER]:
        found = "、".join(f"「{name}」" for name in names) or "（没有表）"
        raise BusinessException(
            ecm_sheet(
                ECM_SHEET_REQUIRED,
                f"现在的表是{found}。请改成先「{SHEET_HELP}」，再「{SHEET_ORDER}」。",
            )
        )
    sheet = workbook[SHEET_ORDER]
    if sheet.sheet_state != "visible":
        raise BusinessException(
            ecm_sheet(ECM_SHEET_REQUIRED, "「工程变更单（材质）」被隐藏了，请取消隐藏后再传。")
        )
    if sheet.merged_cells.ranges:
        raise BusinessException(ecm_sheet(ECM_MERGED, "选中合并过的格子，取消合并后再传。"))
    if getattr(sheet, "_images", None):
        raise BusinessException(
            ecm_sheet(ECM_SHEET_IMAGE, "图片在「工程变更单（材质）」上，删掉后再传。")
        )
    _assert_no_formula(sheet)
    header = [_cell_header(sheet.cell(1, col), col) for col in range(1, 13)]
    if tuple(header) != HEADERS:
        for index, (got, want) in enumerate(zip(header, HEADERS, strict=True)):
            if got != want:
                raise BusinessException(
                    ecm_sheet(
                        ECM_HEADER,
                        f"第 1 行第 {index + 1} 列写成{quoted(got)}，请改成{quoted(want)}。",
                    )
                )
        raise BusinessException(ecm_sheet(ECM_HEADER))
    rows = _read_data_rows(sheet)
    if not rows:
        raise BusinessException(ecm_sheet(ECM_EMPTY_ROW))
    _assert_used_range(sheet, len(rows))
    return rows


def _cell_header(cell: Cell | MergedCell, col: int) -> str:
    expected = HEADERS[col - 1]
    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
        raise BusinessException(
            ecm_sheet(ECM_FORMULA, f"第 1 行「{expected}」是公式。请改成文字「{expected}」。")
        )
    value = cell.value
    if not isinstance(value, str):
        raise BusinessException(
            ecm_sheet(ECM_HEADER, f"第 1 行「{expected}」不是文字。请写成「{expected}」。")
        )
    return value


def _assert_no_formula(sheet: Worksheet) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f" or (
                isinstance(cell.value, str) and cell.value.startswith("=")
            ):
                place = cell.coordinate or "某一格"
                raise BusinessException(
                    ecm_sheet(ECM_FORMULA, f"{place} 是公式。请改成算好的文字后再传。")
                )


def _read_data_rows(sheet: Worksheet) -> list[MaterialChangeParsedRow]:
    parsed: list[MaterialChangeParsedRow] = []
    ended = False
    max_scan = max(sheet.max_row or 1, 2)
    for excel_row in range(2, max_scan + 1):
        values = [sheet.cell(excel_row, col).value for col in range(1, 13)]
        if all(value is None or value == "" for value in values):
            ended = True
            continue
        if ended:
            raise BusinessException(
                ecm_sheet(
                    ECM_EMPTY_ROW,
                    f"第 {excel_row} 行是空的，后面又有字。请删掉空行，让零件行连在一起。",
                )
            )
        cells = [sheet.cell(excel_row, col) for col in range(1, 13)]
        texts = [_read_data_cell(cells[index], excel_row, index) for index in range(12)]
        period_text = texts[3]
        match = _PERIOD_RE.match(period_text)
        if match is None:
            raise BusinessException(ecm_sheet(ECM_EFFECTIVE_FORMAT))
        parsed.append(
            MaterialChangeParsedRow(
                excel_row=excel_row,
                location=texts[0],
                vehicle_model=texts[1],
                change_version=texts[2],
                effective_period_text=period_text,
                effective_year=int(match.group(1)),
                effective_month=int(match.group(2)),
                stamping_supplier=texts[4],
                part_number=texts[5],
                part_name=texts[6],
                before_material=texts[7],
                before_thickness=texts[8],
                after_material=texts[9],
                after_thickness=texts[10],
                interchangeability=texts[11],
            )
        )
    return parsed


def _read_data_cell(cell: Cell | MergedCell, excel_row: int, index: int) -> str:
    column = HEADERS[index]
    if index == 3:
        return _read_period_cell(cell, excel_row)
    if index in _THICKNESS_INDEXES:
        return _read_thickness_cell(cell, excel_row, column)
    return _read_string_cell(cell, excel_row, column)


def _read_period_cell(cell: Cell | MergedCell, excel_row: int) -> str:
    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
        raise BusinessException(
            ecm_row(ECM_FORMULA, excel_row, "生效年月", "是公式。请改成文字或日期后再传。")
        )
    value = cell.value
    if value is None or value == "":
        raise BusinessException(
            ecm_row(
                ECM_CELL_NOT_EXACT,
                excel_row,
                "生效年月",
                "是空的。请填上「2026年06月」，或选一个日期。",
            )
        )
    if isinstance(value, str):
        if value != value.strip():
            raise BusinessException(
                ecm_row(
                    ECM_CELL_NOT_EXACT,
                    excel_row,
                    "生效年月",
                    "有空格或换行。请写成「2026年06月」，或直接选日期。",
                )
            )
        year_month = _year_month_from_text(value)
        if year_month is None:
            raise BusinessException(
                ecm_row(
                    ECM_EFFECTIVE_FORMAT,
                    excel_row,
                    "生效年月",
                    f"写成{quoted(value)}。请写成「2026年6月」或「2026年06月」，也可以直接选日期。",
                )
            )
        year, month = year_month
        return f"{year:04d}年{month:02d}月"
    year_month = _year_month_from_excel_date(cell)
    if year_month is None:
        raise BusinessException(
            ecm_row(
                ECM_EFFECTIVE_FORMAT,
                excel_row,
                "生效年月",
                f"写成{quoted(value)}。请写成「2026年6月」或「2026年06月」，也可以直接选日期。",
            )
        )
    year, month = year_month
    return f"{year:04d}年{month:02d}月"


def _year_month_from_text(value: str) -> tuple[int, int] | None:
    match = _PERIOD_RE.match(value)
    if match is None:
        return None
    month = int(match.group(2))
    if month < MONTH_MIN or month > MONTH_MAX:
        return None
    return int(match.group(1)), month


def _year_month_from_excel_date(cell: Cell | MergedCell) -> tuple[int, int] | None:
    value = cell.value
    if isinstance(value, datetime):
        return value.year, value.month
    if isinstance(value, date):
        return value.year, value.month
    if isinstance(value, bool) or not isinstance(value, int | float):
        return None
    number_format = cell.number_format or ""
    if cell.data_type != "d" and not is_date_format(number_format):
        return None
    try:
        converted = from_excel(value)
    except (ValueError, OverflowError, TypeError):
        return None
    if isinstance(converted, datetime):
        return converted.year, converted.month
    if isinstance(converted, date):
        return converted.year, converted.month
    return None


def _read_thickness_cell(cell: Cell | MergedCell, excel_row: int, column: str) -> str:
    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
        raise BusinessException(
            ecm_row(ECM_FORMULA, excel_row, column, "是公式。请改成算好的数字后再传。")
        )
    value = cell.value
    if value is None or value == "":
        raise BusinessException(
            ecm_row(
                ECM_THICKNESS_FORMAT,
                excel_row,
                column,
                "是空的。请填大于 0 的数字，最多两位小数，例如 1.5。",
            )
        )
    if isinstance(value, str):
        raise BusinessException(
            ecm_row(
                ECM_THICKNESS_FORMAT,
                excel_row,
                column,
                f"现在是文字{quoted(value)}。请改成数字，最多两位小数，例如 1.5。",
            )
        )
    if isinstance(value, bool):
        raise BusinessException(
            ecm_row(
                ECM_THICKNESS_FORMAT,
                excel_row,
                column,
                "现在是勾选框。请改成数字，最多两位小数，例如 1.5。",
            )
        )
    if isinstance(value, datetime | date) or cell.data_type == "d":
        raise BusinessException(
            ecm_row(
                ECM_THICKNESS_FORMAT,
                excel_row,
                column,
                "现在是日期。请改成数字，最多两位小数，例如 1.5。",
            )
        )
    if not isinstance(value, int | float):
        raise BusinessException(
            ecm_row(
                ECM_THICKNESS_FORMAT,
                excel_row,
                column,
                "不是数字。请改成数字，最多两位小数，例如 1.5。",
            )
        )
    shown = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
    try:
        parsed = Decimal(str(value))
    except InvalidOperation:
        parsed = None
    exponent = parsed.as_tuple().exponent if parsed is not None else None
    if parsed is None or not parsed.is_finite() or not isinstance(exponent, int):
        raise BusinessException(
            ecm_row(
                ECM_THICKNESS_FORMAT,
                excel_row,
                column,
                f"现在是 {shown}。请改成大于 0 的数字，最多两位小数，例如 1.5。",
            )
        )
    if exponent < -2:
        raise BusinessException(
            ecm_row(
                ECM_THICKNESS_FORMAT,
                excel_row,
                column,
                f"现在是 {shown}，小数超过两位。请改成最多两位小数，例如 1.5。",
            )
        )
    if parsed <= 0:
        raise BusinessException(
            ecm_row(
                ECM_THICKNESS_FORMAT,
                excel_row,
                column,
                f"现在是 {shown}。请改成大于 0 的数字，最多两位小数，例如 1.5。",
            )
        )
    text = format(parsed, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _read_string_cell(cell: Cell | MergedCell, excel_row: int, column: str) -> str:
    _assert_plain_string(cell, excel_row, column)
    value = cell.value
    if not isinstance(value, str):
        raise BusinessException(ecm_row(ECM_CELL_TYPE, excel_row, column, _cell_kind(cell)))
    if not value or value != value.strip():
        raise BusinessException(
            ecm_row(
                ECM_CELL_NOT_EXACT,
                excel_row,
                column,
                "有空格、换行，或是空的。请填上文字，并删掉多余空白后再传。",
            )
        )
    return value


def _cell_kind(cell: Cell | MergedCell) -> str:
    if cell.value is None or cell.value == "":
        return "是空的。请填上文字后再传。"
    if isinstance(cell.value, bool):
        return "是勾选框。请改成文字后再传。"
    if cell.data_type == "d":
        return "是日期。请改成文字，例如「2026年06月」。"
    if isinstance(cell.value, int | float) or cell.data_type == "n":
        return f"现在是数字 {cell.value}。请改成文字后再传。"
    return "不是文字。请改成文字后再传。"


def _assert_plain_string(cell: Cell | MergedCell, excel_row: int, column: str) -> None:
    if cell.data_type == "f":
        raise BusinessException(
            ecm_row(ECM_FORMULA, excel_row, column, "是公式。请改成算好的文字后再传。")
        )
    if cell.value is None:
        raise BusinessException(
            ecm_row(ECM_CELL_NOT_EXACT, excel_row, column, "是空的。请填上后再传。")
        )
    if isinstance(cell.value, bool) or not isinstance(cell.value, str):
        raise BusinessException(ecm_row(ECM_CELL_TYPE, excel_row, column, _cell_kind(cell)))
    if cell.data_type not in _STRING_TYPES:
        raise BusinessException(ecm_row(ECM_CELL_TYPE, excel_row, column, _cell_kind(cell)))


def _assert_used_range(sheet: Worksheet, data_count: int) -> None:
    last_row = 1 + data_count
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None or cell.value == "":
                continue
            if cell.column is None or cell.row is None:
                continue
            if cell.column > 12 or cell.row > last_row:
                place = cell.coordinate or "某一格"
                raise BusinessException(
                    ecm_sheet(ECM_EXTRA_CELL, f"{place} 还有字。请删掉后再传。")
                )
