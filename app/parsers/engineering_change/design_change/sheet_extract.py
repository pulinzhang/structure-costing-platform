"""按表名从工程变更抽出业务工作表字节。"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.worksheet.worksheet import Worksheet

from app.common.engineering_change_gates import EC_SHEET_REQUIRED, listed
from app.exceptions.base import BusinessException

SHEET_CHANGE_ORDER = "设计变更单"
SHEET_CHANGE_ORDER_HINT = "设计变更填写说明"
SHEET_BOM = "车身件材料清单"
SHEET_BOM_HINT = "材料清单填写说明"
SHEET_PROCESS = "冲压工艺确认表"
SHEET_PROCESS_HINT = "工艺确认填写说明"

REQUIRED_SHEETS = (SHEET_CHANGE_ORDER, SHEET_BOM, SHEET_PROCESS)


def extract_named_sheets(file_content: bytes) -> dict[str, bytes]:
    """
    按表名抽出工作表为独立 xlsx 字节。
    缺设计变更单 / 车身件材料清单 / 冲压工艺确认表 → EC_SHEET_REQUIRED。
    不抽「零件汇总」：模具布点逻辑建立，不读 Excel。
    """
    workbook = load_workbook(BytesIO(file_content), data_only=False)
    names = set(workbook.sheetnames)
    missing = [name for name in REQUIRED_SHEETS if name not in names]
    if missing:
        raise BusinessException(
            f"{EC_SHEET_REQUIRED}：缺 {listed(missing)}。"
            f"表名必须是：{SHEET_CHANGE_ORDER}、{SHEET_BOM}、{SHEET_PROCESS}。"
        )

    result: dict[str, bytes] = {}
    for sheet_name in workbook.sheetnames:
        if sheet_name not in {
            SHEET_CHANGE_ORDER,
            SHEET_BOM,
            SHEET_BOM_HINT,
            SHEET_PROCESS,
            SHEET_PROCESS_HINT,
        }:
            continue
        result[sheet_name] = _sheet_to_workbook_bytes(workbook[sheet_name])
    return result


def _sheet_to_workbook_bytes(source: Worksheet) -> bytes:
    dest_wb = Workbook()
    dest = dest_wb.active
    assert dest is not None
    dest.title = source.title
    for row in source.iter_rows():
        for cell in row:
            dest.cell(cell.row, cell.column, cell.value)
    for merged in source.merged_cells.ranges:
        dest.merge_cells(str(merged))
    buffer = BytesIO()
    with TemporaryDirectory() as temp_dir:
        _copy_sheet_images(source, dest, Path(temp_dir))
        dest_wb.save(buffer)
    return buffer.getvalue()


def _copy_sheet_images(source: Worksheet, dest: Worksheet, temp_dir: Path) -> None:
    for index, image in enumerate(list(getattr(source, "_images", []))):
        data_fn = getattr(image, "_data", None)
        payload = data_fn() if callable(data_fn) else None
        if not payload:
            continue
        png_path = temp_dir / f"sheet_img_{index}.bin"
        png_path.write_bytes(payload)
        copied = ExcelImage(str(png_path))
        if image.width:
            copied.width = image.width
        if image.height:
            copied.height = image.height
        dest.add_image(copied)
        anchor = getattr(image, "anchor", None)
        if anchor is not None:
            copied.anchor = anchor
