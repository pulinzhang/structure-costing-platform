"""变更单解析：十列；项目/生效年月、同一总成的变更前/后总成与总成名称可合并。"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from app.common.engineering_change_gates import EC_CHANGE_ORDER_ROW_COUNT
from app.exceptions.base import BusinessException
from app.parsers.engineering_change.design_change.sheet_extract import SHEET_CHANGE_ORDER

CHANGE_ORDER_HEADERS = (
    "项目",
    "生效年月",
    "变更前总成",
    "变更后总成",
    "总成名称",
    "变更类型",
    "零件层级",
    "变更前零件",
    "变更后零件",
    "零件名称",
)


@dataclass(slots=True)
class ChangeOrderRowParsed:
    """变更单一行。"""

    project_code: str
    effective_period_text: str
    change_type: str
    before_assembly_part_number: str
    after_assembly_part_number: str | None
    assembly_name: str | None
    part_level: str
    before_part_number: str | None
    after_part_number: str | None
    part_name: str | None
    excel_row: int
    before_assembly_span: tuple[int, int] | None = None
    after_assembly_span: tuple[int, int] | None = None


def parse_change_order_sheet(file_content: bytes) -> list[ChangeOrderRowParsed]:
    """解析变更单工作表；至少一行。一本可多行。"""
    workbook = load_workbook(BytesIO(file_content), data_only=True)
    sheet = (
        workbook[SHEET_CHANGE_ORDER]
        if SHEET_CHANGE_ORDER in workbook.sheetnames
        else workbook.active
    )
    assert sheet is not None
    header_row = _find_header_row(sheet)
    if header_row is None:
        raise BusinessException(
            f"{EC_CHANGE_ORDER_ROW_COUNT}：没找到工程变更单十列表头"
            "（项目、生效年月、变更前总成…）。"
        )

    data_rows: list[ChangeOrderRowParsed] = []
    for row_index in range(header_row + 1, sheet.max_row + 1):
        values = [
            _merged_cell_value(sheet, row_index, col)
            for col in range(1, len(CHANGE_ORDER_HEADERS) + 1)
        ]
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        data_rows.append(
            ChangeOrderRowParsed(
                project_code=_as_text(values[0]),
                effective_period_text=_as_text(values[1]),
                before_assembly_part_number=_as_text(values[2]),
                after_assembly_part_number=_as_optional_text(values[3]),
                assembly_name=_as_optional_text(values[4]),
                change_type=_as_text(values[5]),
                part_level=_as_text(values[6]),
                before_part_number=_as_optional_text(values[7]),
                after_part_number=_as_optional_text(values[8]),
                part_name=_as_optional_text(values[9]),
                excel_row=row_index,
                before_assembly_span=_vertical_merge_span(sheet, row_index, 3),
                after_assembly_span=_vertical_merge_span(sheet, row_index, 4),
            )
        )

    if not data_rows:
        raise BusinessException(f"{EC_CHANGE_ORDER_ROW_COUNT}。工程变更单表头下面没有数据行。")
    return data_rows


def _find_header_row(sheet: Worksheet) -> int | None:
    for row_index in range(1, min(sheet.max_row, 20) + 1):
        cells = [
            _cell_text(sheet.cell(row_index, col))
            for col in range(1, len(CHANGE_ORDER_HEADERS) + 1)
        ]
        if cells == list(CHANGE_ORDER_HEADERS):
            return row_index
    return None


def _vertical_merge_span(sheet: Worksheet, row: int, col: int) -> tuple[int, int] | None:
    for merged in sheet.merged_cells.ranges:
        if merged.min_col != col or merged.max_col != col:
            continue
        if merged.min_row <= row <= merged.max_row and merged.max_row > merged.min_row:
            return (merged.min_row, merged.max_row)
    return None


def _merged_cell_value(sheet: Worksheet, row: int, col: int) -> object | None:
    cell = sheet.cell(row, col)
    value = cell.value
    if value is not None and str(value).strip() != "":
        return value
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
            top_left = sheet.cell(merged.min_row, merged.min_col)
            return top_left.value
    return value


def _cell_text(cell: Cell | MergedCell) -> str:
    if cell.value is None:
        return ""
    return str(cell.value).strip()


def _as_text(value: object | None) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _as_optional_text(value: object | None) -> str | None:
    text = _as_text(value)
    return text or None
