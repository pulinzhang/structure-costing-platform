"""按表名从工程变更合册抽出业务工作表字节。"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.common.engineering_change_gates import EC_SHEET_REQUIRED
from app.exceptions.base import BusinessException

SHEET_CHANGE_ORDER = "工程变更单"
SHEET_BOM = "车身件材料清单"
SHEET_BOM_HINT = "材料清单填写说明"
SHEET_PROCESS = "冲压工艺确认表"
SHEET_PROCESS_HINT = "工艺确认填写说明"
SHEET_DIE = "零件汇总"

REQUIRED_SHEETS = (SHEET_CHANGE_ORDER, SHEET_BOM, SHEET_PROCESS)


def extract_named_sheets(file_content: bytes) -> dict[str, bytes]:
    """
    按表名抽出工作表为独立 xlsx 字节。
    缺工程变更单 / 车身件材料清单 / 冲压工艺确认表 → EC_SHEET_REQUIRED。
    「零件汇总」可缺。
    """
    workbook = load_workbook(BytesIO(file_content), data_only=False)
    names = set(workbook.sheetnames)
    missing = [name for name in REQUIRED_SHEETS if name not in names]
    if missing:
        raise BusinessException(EC_SHEET_REQUIRED)

    result: dict[str, bytes] = {}
    for sheet_name in workbook.sheetnames:
        if sheet_name not in {
            SHEET_CHANGE_ORDER,
            SHEET_BOM,
            SHEET_BOM_HINT,
            SHEET_PROCESS,
            SHEET_PROCESS_HINT,
            SHEET_DIE,
        }:
            continue
        result[sheet_name] = _sheet_to_workbook_bytes(workbook[sheet_name])
    return result


def _sheet_to_workbook_bytes(source: Worksheet) -> bytes:
    dest_wb = Workbook()
    dest = dest_wb.active
    assert dest is not None
    dest.title = source.title
    for row in source.iter_rows():
        for cell in row:
            dest.cell(cell.row, cell.column, cell.value)
    for merged in source.merged_cells.ranges:
        dest.merge_cells(str(merged))
    for image in list(getattr(source, "_images", [])):
        dest.add_image(image)
    buffer = BytesIO()
    dest_wb.save(buffer)
    return buffer.getvalue()
