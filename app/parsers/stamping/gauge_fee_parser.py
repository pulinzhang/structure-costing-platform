"""冲压检具费 Excel 解析（纯函数）。"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.common.money import Q4, quantize_money
from app.common.part_number import normalize_part_number_token
from app.models.stamping.gauge_fee import (
    STATUS_EFFECTIVE,
    STATUS_INVALID,
)

HEADER_PART_NUMBER = "零件号"
HEADER_PART_NAME = "零件名称"
HEADER_GAUGE_FEE = "检具费（万元）"
HEADER_ALLOCATION_QTY = "项目分摊数量（万台）"
HEADER_STATUS = "状态"

_HEADER_ALIASES: dict[str, str] = {
    "零件号": HEADER_PART_NUMBER,
    "零件编号": HEADER_PART_NUMBER,
    "零件名称": HEADER_PART_NAME,
    "名称": HEADER_PART_NAME,
    "检具费（万元）": HEADER_GAUGE_FEE,
    "检具费": HEADER_GAUGE_FEE,
    "项目分摊数量（万台）": HEADER_ALLOCATION_QTY,
    "项目分摊数量": HEADER_ALLOCATION_QTY,
    "分摊数量": HEADER_ALLOCATION_QTY,
    "状态": HEADER_STATUS,
}


@dataclass(frozen=True)
class GaugeFeeParsedRow:
    row_number: int
    part_number: str
    gauge_fee: Decimal
    allocation_qty: Decimal
    status: str


@dataclass(frozen=True)
class GaugeFeeParseResult:
    rows: list[GaugeFeeParsedRow]
    errors: list[str]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_header(value: Any) -> str:
    text = _cell_text(value).replace(" ", "").replace("\n", "")
    return _HEADER_ALIASES.get(text, text)


def _parse_required_money(value: Any, field_label: str) -> Decimal:
    if value is None or _cell_text(value) == "":
        raise ValueError(f"「{field_label}」不能为空")
    try:
        number = quantize_money(value, Q4, allow_negative=True)
    except (ValueError, InvalidOperation) as exc:
        raise ValueError(f"「{field_label}」不是有效数字") from exc
    if number < 0:
        raise ValueError(f"「{field_label}」不能为负")
    return number


def _parse_status(value: Any) -> str:
    text = _cell_text(value)
    if not text:
        return STATUS_EFFECTIVE
    mapping = {
        "生效": STATUS_EFFECTIVE,
        "失效": STATUS_INVALID,
        STATUS_EFFECTIVE: STATUS_EFFECTIVE,
        STATUS_INVALID: STATUS_INVALID,
    }
    status = mapping.get(text)
    if status is None:
        raise ValueError("状态须为生效或失效")
    return status


def parse_gauge_fee_excel(file_content: bytes) -> GaugeFeeParseResult:
    workbook = load_workbook(BytesIO(file_content), data_only=True, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        header_map: dict[str, int] = {}
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not first_row:
            return GaugeFeeParseResult(rows=[], errors=["模板没有表头"])
        for index, cell in enumerate(first_row):
            key = _normalize_header(cell)
            if key in {
                HEADER_PART_NUMBER,
                HEADER_PART_NAME,
                HEADER_GAUGE_FEE,
                HEADER_ALLOCATION_QTY,
                HEADER_STATUS,
            }:
                header_map[key] = index
        if HEADER_PART_NUMBER not in header_map:
            return GaugeFeeParseResult(rows=[], errors=["缺少列「零件号」"])
        missing_headers = [
            title for title in (HEADER_GAUGE_FEE, HEADER_ALLOCATION_QTY) if title not in header_map
        ]
        if missing_headers:
            labels = "、".join(f"「{title}」" for title in missing_headers)
            return GaugeFeeParseResult(rows=[], errors=[f"缺少列{labels}"])

        rows: list[GaugeFeeParsedRow] = []
        errors: list[str] = []
        seen: set[str] = set()
        for excel_row, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if values is None or all(item is None or str(item).strip() == "" for item in values):
                continue
            raw_number = _cell_text(
                values[header_map[HEADER_PART_NUMBER]]
                if header_map[HEADER_PART_NUMBER] < len(values)
                else None
            )
            part_number = normalize_part_number_token(raw_number)
            if not part_number:
                errors.append(f"第 {excel_row} 行：零件号不能为空")
                continue
            if part_number in seen:
                errors.append(f"第 {excel_row} 行：零件号「{part_number}」重复")
                continue
            seen.add(part_number)
            try:
                gauge_fee = _parse_required_money(
                    values[header_map[HEADER_GAUGE_FEE]]
                    if header_map[HEADER_GAUGE_FEE] < len(values)
                    else None,
                    HEADER_GAUGE_FEE,
                )
                allocation_qty = _parse_required_money(
                    values[header_map[HEADER_ALLOCATION_QTY]]
                    if header_map[HEADER_ALLOCATION_QTY] < len(values)
                    else None,
                    HEADER_ALLOCATION_QTY,
                )
                status = _parse_status(
                    values[header_map[HEADER_STATUS]]
                    if HEADER_STATUS in header_map and header_map[HEADER_STATUS] < len(values)
                    else None
                )
            except ValueError as exc:
                errors.append(f"第 {excel_row} 行：{exc}")
                continue
            rows.append(
                GaugeFeeParsedRow(
                    row_number=excel_row,
                    part_number=part_number,
                    gauge_fee=gauge_fee,
                    allocation_qty=allocation_qty,
                    status=status,
                )
            )
        return GaugeFeeParseResult(rows=rows, errors=errors)
    finally:
        workbook.close()
