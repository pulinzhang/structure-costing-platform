"""
冲次单价 Excel 解析器（纯函数，无副作用）

约定三列：名称 | 单位 | 单价
示例：
  400T | 次 | 0.564
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.common.money import Q3, quantize_money

HEADER_NAME = "名称"
HEADER_UNIT = "单位"
HEADER_PRICE = "单价"

_TONNAGE_RE = re.compile(r"^\s*(\d+)\s*[TtｔＴ]?\s*$")


@dataclass(frozen=True)
class TonnagePriceParsedRow:
    """解析后的一行冲次单价"""

    row_number: int
    tonnage: int
    unit: str
    price_per_stroke: Decimal
    name_raw: str


@dataclass(frozen=True)
class TonnagePriceParseResult:
    """解析结果"""

    rows: list[TonnagePriceParsedRow]
    errors: list[str]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _round3(value: Decimal) -> Decimal:
    return quantize_money(value, Q3)


def _parse_tonnage(name: str) -> int | None:
    match = _TONNAGE_RE.match(name or "")
    if not match:
        return None
    tonnage = int(match.group(1))
    return tonnage if tonnage > 0 else None


def _parse_price(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        if isinstance(value, Decimal):
            number = value
        elif isinstance(value, int | float):
            number = Decimal(str(value))
        else:
            text = str(value).strip().replace(",", "")
            if not text:
                return None
            number = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    if number < 0:
        return None
    return _round3(number)


def _find_header_row(sheet: Any) -> tuple[int, dict[str, int]] | None:
    """定位表头行，返回 (行号, 列名→列号)。"""
    aliases = {
        HEADER_NAME: {HEADER_NAME, "名称", "吨位", "压机吨位", "吨位名称"},
        HEADER_UNIT: {HEADER_UNIT, "单位"},
        HEADER_PRICE: {HEADER_PRICE, "单价", "单冲次价格", "价格"},
    }
    max_scan = min(sheet.max_row or 1, 20)
    max_col = min(sheet.max_column or 1, 20)

    for row_idx in range(1, max_scan + 1):
        mapping: dict[str, int] = {}
        for col_idx in range(1, max_col + 1):
            text = _cell_text(sheet.cell(row_idx, col_idx).value)
            if not text:
                continue
            for field, names in aliases.items():
                if text in names and field not in mapping:
                    mapping[field] = col_idx
        if HEADER_NAME in mapping and HEADER_PRICE in mapping:
            return row_idx, mapping
    return None


def parse_tonnage_price_excel(file_content: bytes) -> TonnagePriceParseResult:
    """
    解析冲次单价 Excel。
    仅支持 .xlsx（由调用方校验扩展名）。
    """
    if not file_content:
        return TonnagePriceParseResult(rows=[], errors=["文件内容为空"])

    try:
        workbook = load_workbook(BytesIO(file_content), data_only=True, read_only=True)
    except Exception as exc:
        return TonnagePriceParseResult(rows=[], errors=[f"无法读取 Excel：{exc}"])

    try:
        if not workbook.worksheets:
            return TonnagePriceParseResult(rows=[], errors=["Excel 中没有工作表"])
        sheet = workbook.worksheets[0]

        header = _find_header_row(sheet)
        if not header:
            return TonnagePriceParseResult(
                rows=[],
                errors=[f"未找到表头，请使用三列：{HEADER_NAME} / {HEADER_UNIT} / {HEADER_PRICE}"],
            )

        header_row, columns = header
        name_col = columns[HEADER_NAME]
        unit_col = columns.get(HEADER_UNIT)
        price_col = columns[HEADER_PRICE]

        rows: list[TonnagePriceParsedRow] = []
        errors: list[str] = []
        seen_tonnage: set[int] = set()

        for row_idx in range(header_row + 1, (sheet.max_row or header_row) + 1):
            name_raw = _cell_text(sheet.cell(row_idx, name_col).value)
            unit_raw = _cell_text(sheet.cell(row_idx, unit_col).value) if unit_col else ""
            price_raw = sheet.cell(row_idx, price_col).value

            if not name_raw and (price_raw is None or price_raw == ""):
                continue

            tonnage = _parse_tonnage(name_raw)
            if tonnage is None:
                errors.append(f"第 {row_idx} 行：名称「{name_raw}」无法识别为吨位（如 400T）")
                continue

            price = _parse_price(price_raw)
            if price is None:
                errors.append(f"第 {row_idx} 行：单价无效（{price_raw!r}）")
                continue

            unit = unit_raw or "次"
            if tonnage in seen_tonnage:
                errors.append(f"第 {row_idx} 行：吨位 {tonnage}T 重复")
                continue
            seen_tonnage.add(tonnage)

            rows.append(
                TonnagePriceParsedRow(
                    row_number=row_idx,
                    tonnage=tonnage,
                    unit=unit,
                    price_per_stroke=price,
                    name_raw=name_raw,
                )
            )
    finally:
        workbook.close()

    if not rows and not errors:
        errors.append("未解析到有效数据行")

    return TonnagePriceParseResult(rows=rows, errors=errors)
