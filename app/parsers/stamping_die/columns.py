"""模具布点表头：找 Sheet、展开合并格、解析列映射。"""

from __future__ import annotations

import re

from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.exceptions.base import BusinessException

from .types import DieColumnMap

MAX_COL = 32  # 含模板插入的「模具编号」列；旧 A～AC 样本仍可解析
_HEADER_SCAN_ROWS = 12
_CHILD_HEADER_TOKENS = (
    "工序名称",
    "工序内容",
    "长度",
    "宽度",
    "高度",
    "模具长",
    "模具宽",
    "模具高",
)


def find_sheet(workbook: Workbook) -> Worksheet:
    for candidate in workbook.worksheets:
        for row_idx in range(1, min(_HEADER_SCAN_ROWS, (candidate.max_row or 1) + 1)):
            for col in range(1, MAX_COL + 1):
                text = str(candidate.cell(row_idx, col).value or "")
                if "零件号" in text:
                    return candidate
    if not workbook.worksheets:
        raise BusinessException("Excel 没有工作表")
    return workbook.worksheets[0]


def build_merge_lookup(
    sheet: Worksheet,
    *,
    vertical_only: bool,
    max_row: int | None,
) -> dict[tuple[int, int], object]:
    """
    vertical_only：只展开「同一列」的纵向合并，避免横向合并把模具方式灌进工序列。
    max_row：表头扫描时只收录落在该行及以内的合并。
    """
    lookup: dict[tuple[int, int], object] = {}
    for rng in sheet.merged_cells.ranges:
        if rng.min_col > MAX_COL:
            continue
        if max_row is not None and rng.min_row > max_row:
            continue
        if vertical_only and rng.min_col != rng.max_col:
            continue
        top = sheet.cell(rng.min_row, rng.min_col).value
        last_col = min(rng.max_col, MAX_COL)
        last_row = rng.max_row if max_row is None else min(rng.max_row, max_row)
        for row in range(rng.min_row, last_row + 1):
            for col in range(rng.min_col, last_col + 1):
                lookup[(row, col)] = top
    return lookup


def cell(
    sheet: Worksheet,
    lookup: dict[tuple[int, int], object],
    row: int,
    col: int,
) -> object:
    if (row, col) in lookup:
        return lookup[(row, col)]
    return sheet.cell(row, col).value


def find_header_rows(
    sheet: Worksheet,
    lookup: dict[tuple[int, int], object],
) -> tuple[int, int | None, int]:
    max_row = min(sheet.max_row or 1, _HEADER_SCAN_ROWS)
    for row in range(1, max_row + 1):
        for col in range(1, MAX_COL + 1):
            text = str(cell(sheet, lookup, row, col) or "").replace("\n", "")
            if "零件号" not in text:
                continue
            child_row: int | None = None
            next_row = row + 1
            if next_row <= (sheet.max_row or 1):
                child_blob = "".join(
                    _compact(cell(sheet, lookup, next_row, item)) for item in range(1, MAX_COL + 1)
                )
                if any(token in child_blob for token in _CHILD_HEADER_TOKENS):
                    child_row = next_row
            data_start = (child_row or row) + 1
            return row, child_row, data_start
    raise BusinessException("未识别到模具布点表头（需要「零件号」列）")


def _compact(value: object) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def _column_path(
    sheet: Worksheet,
    lookup: dict[tuple[int, int], object],
    header_row: int,
    child_row: int | None,
    col: int,
) -> str:
    parent = _compact(cell(sheet, lookup, header_row, col))
    child = _compact(cell(sheet, lookup, child_row, col)) if child_row else ""
    if parent and child and child not in parent:
        return parent + child
    return parent or child


def resolve_columns(
    sheet: Worksheet,
    lookup: dict[tuple[int, int], object],
    header_row: int,
    child_row: int | None,
) -> DieColumnMap:
    columns: DieColumnMap = {}
    taken: set[int] = set()

    def assign(key: str, col: int) -> None:
        if key not in columns and col not in taken:
            columns[key] = col  # type: ignore[literal-required]
            taken.add(col)

    paths = {
        col: _column_path(sheet, lookup, header_row, child_row, col)
        for col in range(1, MAX_COL + 1)
    }
    # 先匹配更具体的词，避免「长度」被「长」吃掉、工序列被「工序名称」吃掉。
    for col, path in paths.items():
        if not path or "简图" in path or "图片" in path:
            continue
        if "零件号" in path:
            assign("part_number", col)
        elif "零件名称" in path:
            assign("part_name", col)
        elif "工序名称" in path:
            assign("process_seq", col)
        elif "工序内容" in path:
            assign("process_code", col)
        elif path == "套数" or path.endswith("套数"):
            assign("die_sets", col)
        elif "模具编号" in path:
            assign("die_number", col)
        elif "模具长" in path:
            assign("die_length", col)
        elif "模具宽" in path:
            assign("die_width", col)
        elif "模具高" in path:
            assign("die_height", col)
        elif path == "工序" or path.endswith("工序"):
            assign("process_code", col)
        elif "掏空" in path:
            assign("hollow_factor", col)
        elif "重量小计" in path:
            assign("estimated_weight_subtotal_t", col)
        elif "预估重量" in path or "模具重量" in path:
            assign("estimated_weight_t", col)
        elif "吨位" in path:
            assign("tonnage", col)
        elif "模具类型" in path:
            assign("die_type", col)
        elif "模具方式" in path or "工艺方式" in path:
            assign("die_method", col)
        elif "生产方式" in path or "自动化" in path and "凸焊" not in path:
            assign("production_mode", col)
        elif "端拾器" in path:
            assign("needs_end_effector", col)
        elif "零件等级" in path or "难易度" in path:
            assign("part_grade", col)
        elif path == "备注" or path.endswith("备注"):
            assign("remark", col)
        elif "材质" in path or "牌号" in path:
            assign("material", col)
        elif "料厚" in path or "厚度" in path:
            assign("thickness", col)
        elif "净重" in path or "零件重量" in path:
            assign("part_weight", col)
        elif "板料形状" in path or "料片式样" in path:
            assign("sheet_shape", col)
        elif "可制" in path and "件" in path:
            assign("parts_per_sheet", col)
        elif "材料消耗" in path:
            assign("material_consumption", col)
        elif "利用率" in path:
            assign("material_utilization", col)
        elif path == "序号" or path.endswith("序号"):
            assign("seq_no", col)
        elif "长度" in path:
            assign("die_length", col)
        elif "宽度" in path:
            assign("die_width", col)
        elif "高度" in path:
            assign("die_height", col)
        elif "长" in path and ("板料" in path or "料片" in path or path.endswith("长")):
            assign("sheet_length", col)
        elif "宽" in path and ("板料" in path or "料片" in path or path.endswith("宽")):
            assign("sheet_width", col)
    return columns
