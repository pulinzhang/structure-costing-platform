"""把 .xls / .xlsx 收成同一套 sheetnames + [] + iter_rows。"""

from __future__ import annotations

import io
from typing import Any

import openpyxl
import xlrd


class XlsSheetWrapper:
    def __init__(self, sheet: xlrd.sheet.Sheet):
        self._sheet = sheet

    @property
    def title(self) -> Any:
        return self._sheet.name

    def iter_rows(
        self, min_row: int = 1, max_row: int | None = None, values_only: bool = True
    ) -> list[list[Any]]:
        if not values_only:
            raise ValueError("xls 兼容读取只返回单元格值")
        results: list[list[Any]] = []
        row_count = self._sheet.nrows
        actual_min = max(0, min_row - 1)
        actual_max = row_count if max_row is None else min(max_row, row_count)
        for row_idx in range(actual_min, actual_max):
            row_values: list[Any] = []
            for col_idx in range(self._sheet.ncols):
                row_values.append(_xls_cell_value(self._sheet, row_idx, col_idx))
            results.append(row_values)
        return results


class XlsWorkbookWrapper:
    def __init__(self, book: xlrd.Book):
        self._sheets: dict[str, XlsSheetWrapper] = {
            sheet.name: XlsSheetWrapper(sheet) for sheet in book.sheets()
        }

    @property
    def sheetnames(self) -> list[str]:
        return list(self._sheets.keys())

    def __getitem__(self, name: str) -> XlsSheetWrapper:
        if name not in self._sheets:
            raise KeyError(f"Sheet '{name}' not found in workbook")
        return self._sheets[name]


def _xls_cell_value(sheet: xlrd.sheet.Sheet, row_idx: int, col_idx: int) -> Any:
    cell = sheet.cell(row_idx, col_idx)
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            date_tuple = xlrd.xldate_as_tuple(cell.value, sheet.book.datemode)
            if date_tuple[3:] == (0, 0, 0):
                return date_tuple[:3]
            return date_tuple
        except Exception:
            return cell.value
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_ERROR, xlrd.XL_CELL_BLANK):
        return None
    if cell.ctype == xlrd.XL_CELL_TEXT:
        return cell.value
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        value = cell.value
        if value == int(value):
            return int(value)
        return value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    return cell.value


def _is_xls_file(file_content: bytes) -> bool:
    return len(file_content) >= 4 and file_content[:4] == b"\xd0\xcf\x11\xe0"


def open_workbook(file_content: bytes) -> Any:
    if _is_xls_file(file_content):
        return XlsWorkbookWrapper(xlrd.open_workbook(file_contents=file_content))
    return openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
