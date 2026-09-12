"""阶梯项构建：按行读加减价、补全边界后生成 ParsedLadderItem。"""

from __future__ import annotations

from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from app.parsers.yishiyiyi.policy.boundary_parser import _Boundary
from app.parsers.yishiyiyi.policy.models import ParsedLadderItem
from app.parsers.yishiyiyi.policy.models.policy import (
    GRADE_TYPE_CATEGORY,
    GRADE_TYPE_GRADE,
    GRADE_TYPE_WIDTH_RANGE,
    POLICY_TYPE_THICKNESS,
)
from app.parsers.yishiyiyi.policy.value_parser import (
    _to_int_price,
    is_carry_forward_price,
)

from .locate_blocks import BlockLocator
from .sheet_scanner import SheetScanner

_WIDTH_RANGE_KEYWORDS = ("≤", "≥", "＜", "＞", "W", "宽")
_CATEGORY_NAMES = frozenset({"冷轧钢", "镀锌钢", "酸洗钢", "热轧钢"})


class ItemBuilder:
    """行内阶梯明细与单条阶梯项装配。"""

    def __init__(self, scanner: SheetScanner, locator: BlockLocator) -> None:
        self.scanner = scanner
        self._locator = locator

    def infer_grade_type(self, grade: str, policy_type: str) -> str:
        """推断 grade 的语义类型。"""
        if policy_type == POLICY_TYPE_THICKNESS:
            return GRADE_TYPE_CATEGORY
        grade_stripped = grade.strip()
        if any(kw in grade_stripped for kw in _WIDTH_RANGE_KEYWORDS):
            return GRADE_TYPE_WIDTH_RANGE
        if grade_stripped in _CATEGORY_NAMES:
            return GRADE_TYPE_CATEGORY
        return GRADE_TYPE_GRADE

    def build_items_for_row(
        self,
        sheet: Worksheet,
        r: int,
        price_cols: list[int],
        boundaries: dict[int, _Boundary],
        policy_type: str = "THICKNESS",
        sort_order: int = 0,
    ) -> tuple[list[ParsedLadderItem], Any]:
        """构建行内阶梯明细项列表。返回 (items, extra)。"""
        items: list[ParsedLadderItem] = []
        sorted_cols = sorted(price_cols)
        last_ladder_col = next((col for col in reversed(sorted_cols) if col in boundaries), None)
        for i, c in enumerate(sorted_cols):
            if c not in boundaries:
                continue

            price, price_row, raw_value = self._read_row_price(sheet, r, c)
            # 厚度末档填「/」或「无」：页内约定沿用上一阶梯，不得落成 0
            if (
                policy_type == POLICY_TYPE_THICKNESS
                and c == last_ladder_col
                and items
                and is_carry_forward_price(raw_value)
            ):
                price = items[-1].price_adjustment

            boundary = self._locator.infer_upper_bound(
                boundaries[c], sorted_cols, i, boundaries, policy_type
            )

            if policy_type == "THICKNESS":
                items.append(self._make_thickness_item(boundary, price, sort_order, price_row, c))
            else:
                items.append(self._make_width_item(boundary, price, sort_order, price_row, c))

        return items, None

    def _read_row_price(
        self,
        sheet: Worksheet,
        row_index: int,
        col: int,
    ) -> tuple[int, int, object]:
        """读取阶梯加减价，返回 (price, excel_row, raw_value)。"""
        raw_value = sheet.cell(row=row_index, column=col).value
        if raw_value is not None:
            return _to_int_price(raw_value), row_index, raw_value

        start_row = self.scanner.find_start_row(sheet, row_index, col)
        if start_row is None:
            return 0, row_index, None

        start_val = sheet.cell(row=start_row, column=col).value
        if start_val is not None:
            return _to_int_price(start_val), start_row, start_val

        price, price_row = self.scanner.find_value_above_with_row(sheet, start_row, col)
        above_raw = sheet.cell(row=price_row, column=col).value
        return price, price_row, above_raw

    @staticmethod
    def _make_thickness_item(
        boundary: _Boundary, price: int, sort_order: int, excel_row: int, excel_col: int
    ) -> ParsedLadderItem:
        return ParsedLadderItem(
            target_sub_group=None,
            thick_min=boundary.min_val,
            thick_min_operator=boundary.min_operator,
            thick_max=boundary.max_val,
            thick_max_operator=boundary.max_operator,
            width_min=None,
            width_min_operator=">=",
            width_max=None,
            width_max_operator="<",
            price_adjustment=price,
            is_base=False,
            sort_order=sort_order,
            excel_row=excel_row,
            excel_col=excel_col,
        )

    @staticmethod
    def _make_width_item(
        boundary: _Boundary, price: int, sort_order: int, excel_row: int, excel_col: int
    ) -> ParsedLadderItem:
        return ParsedLadderItem(
            target_sub_group=None,
            thick_min=None,
            thick_min_operator=">=",
            thick_max=None,
            thick_max_operator="<",
            width_min=boundary.min_val,
            width_min_operator=boundary.min_operator,
            width_max=boundary.max_val,
            width_max_operator=boundary.max_operator,
            price_adjustment=price,
            is_base=False,
            sort_order=sort_order,
            excel_row=excel_row,
            excel_col=excel_col,
        )

    def build_item(
        self,
        boundary: _Boundary,
        price: int,
        width_boundary: _Boundary | None = None,
        excel_row: int = 0,
        excel_col: int = 0,
    ) -> ParsedLadderItem:
        """构建阶梯明细项（单维度或双维度）。"""
        if width_boundary is not None:
            return ParsedLadderItem(
                target_sub_group=None,
                thick_min=boundary.min_val,
                thick_min_operator=boundary.min_operator,
                thick_max=boundary.max_val,
                thick_max_operator=boundary.max_operator,
                width_min=width_boundary.min_val,
                width_min_operator=width_boundary.min_operator,
                width_max=width_boundary.max_val,
                width_max_operator=width_boundary.max_operator,
                price_adjustment=price,
                is_base=False,
                sort_order=0,
                excel_row=excel_row,
                excel_col=excel_col,
            )
        return ParsedLadderItem(
            target_sub_group=None,
            thick_min=boundary.min_val,
            thick_min_operator=boundary.min_operator,
            thick_max=boundary.max_val,
            thick_max_operator=boundary.max_operator,
            width_min=None,
            width_min_operator=">=",
            width_max=None,
            width_max_operator="<",
            price_adjustment=price,
            is_base=False,
            sort_order=0,
            excel_row=excel_row,
            excel_col=excel_col,
        )
