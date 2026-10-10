"""12-page product briefing (full intro). Conversion deck is build_gm_intro_pptx.py."""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import build_gm_intro_pptx as deck
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches

deck.SLIDE_COUNT = 12

W = deck.W
H = deck.H
ORANGE = deck.ORANGE
INK = deck.INK
WHITE = deck.WHITE
SOFT = deck.SOFT
MUTED = deck.MUTED
DIM = deck.DIM
SAND = deck.SAND
LINE = deck.LINE
FONT_EN = deck.FONT_EN
COMPANY = deck.COMPANY
APP_NAME = deck.APP_NAME
APP_NAME_SHORT = deck.APP_NAME_SHORT


def make() -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    t = deck.textbox
    card = deck.card
    kicker = deck.kicker
    title = deck.title_block
    footer = deck.footer
    dark_footer = deck.dark_footer
    hline = deck.hline
    rect = deck.rect

    s = prs.slides.add_slide(blank)
    deck.paint_dark(s)
    t(
        s,
        Inches(0.7),
        Inches(0.42),
        Inches(6),
        Inches(0.36),
        COMPANY,
        size=16,
        color=WHITE,
        bold=True,
        spacing=40,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    hline(s, Inches(0.7), Inches(1.28), Inches(1.6), ORANGE, 2.25)
    t(
        s,
        Inches(0.7),
        Inches(1.48),
        Inches(12),
        Inches(0.32),
        "从靠人做表，到公司能力",
        size=14,
        color=ORANGE,
    )
    t(
        s,
        Inches(0.7),
        Inches(2.0),
        Inches(12),
        Inches(1.0),
        APP_NAME,
        size=40,
        color=WHITE,
        bold=True,
    )
    t(
        s,
        Inches(0.7),
        Inches(3.2),
        Inches(11.4),
        Inches(1.5),
        "管理人员是公司的建设者。\n核价，是数字化建设、AI 落地最好的抓手。",
        size=22,
        color=SOFT,
        line_pt=36,
    )
    deck.cover_trial_access(s)
    dark_footer(s, 1)

    s = prs.slides.add_slide(blank)
    deck.paint_light(s)
    kicker(s, "01   THE LEVER")
    title(s, "数字化、AI 要落地，核价是最好的抓手")
    t(
        s,
        Inches(0.7),
        Inches(1.5),
        Inches(12),
        Inches(0.55),
        "建设不是再上一个系统。是把公司最痛、最有规则、立刻能见效的事做成能力。",
        size=16,
        color=MUTED,
    )
    levers = [
        ("有规则", "钢价政策、工艺、BOM、设变，本来就有企业口径。\n适合写成引擎，而不是口头 AI。"),
        ("有真账", "开给奇瑞的总成价、付给二级商的采购价，每月都要对账。\n建设立刻碰到钱。"),
        ("立刻见效", "一事一议以前要熟练人员做好几天。\n上传后引擎约 20 秒。AI 问的是这本真账。"),
    ]
    for i, (head, body) in enumerate(levers):
        x = Inches(0.7 + i * 4.15)
        card(s, x, Inches(2.25), Inches(3.95), Inches(4.3))
        t(
            s,
            x + Inches(0.28),
            Inches(2.55),
            Inches(3.4),
            Inches(0.5),
            head,
            size=22,
            color=INK,
            bold=True,
        )
        t(
            s,
            x + Inches(0.28),
            Inches(3.25),
            Inches(3.4),
            Inches(2.8),
            body,
            size=15,
            color=MUTED,
            line_pt=26,
        )
    footer(s, 2)

    s = prs.slides.add_slide(blank)
    deck.paint_light(s)
    kicker(s, "02   WHAT IT IS")
    title(s, APP_NAME)
    t(
        s,
        Inches(0.7),
        Inches(1.52),
        Inches(12),
        Inches(0.7),
        "政策、BOM、工艺、变更纳入同一套规则。算出总成与采购对象价格，直接出协议；AI 解读动因。",
        size=16,
        color=MUTED,
        line_pt=26,
    )
    items = [
        ("Policy", "钢价一事一议"),
        ("BOM", "结构与外购"),
        ("Process", "工艺与模具"),
        ("Change", "工程变更"),
        ("Costing", "件价与总成"),
        ("Contract", "采购协议"),
    ]
    for i, (en, zh) in enumerate(items):
        col, row = i % 3, i // 3
        x = Inches(0.7 + col * 4.15)
        y = Inches(2.4 + row * 2.15)
        card(s, x, y, Inches(3.95), Inches(1.95))
        t(
            s,
            x + Inches(0.28),
            y + Inches(0.32),
            Inches(3.4),
            Inches(0.32),
            en,
            size=12,
            color=ORANGE,
            font=FONT_EN,
            spacing=80,
        )
        t(
            s,
            x + Inches(0.28),
            y + Inches(0.8),
            Inches(3.4),
            Inches(0.7),
            zh,
            size=20,
            color=INK,
            bold=True,
        )
    footer(s, 3)

    s = prs.slides.add_slide(blank)
    deck.paint_light(s)
    kicker(s, "03   SHIFT")
    title(s, "不是把 Excel 做得更复杂")
    card(s, Inches(0.7), Inches(1.7), Inches(5.7), Inches(4.85))
    t(s, Inches(1.05), Inches(2.05), Inches(5), Inches(0.3), "不是", size=13, color=DIM)
    t(
        s,
        Inches(1.05),
        Inches(2.5),
        Inches(5.1),
        Inches(1.0),
        "再做一张更大的表",
        size=24,
        color=INK,
        bold=True,
    )
    t(
        s,
        Inches(1.05),
        Inches(3.8),
        Inches(4.95),
        Inches(2.2),
        "公式在私人文件夹。口径靠骨干记忆。对奇瑞、对二级商，各执一份表，只能扯皮。",
        size=15,
        color=MUTED,
        line_pt=26,
    )
    card(s, Inches(6.9), Inches(1.7), Inches(5.7), Inches(4.85))
    t(s, Inches(7.25), Inches(2.05), Inches(5), Inches(0.3), "而是", size=13, color=ORANGE)
    t(
        s,
        Inches(7.25),
        Inches(2.5),
        Inches(5.1),
        Inches(1.0),
        "把核价变成公司能力",
        size=24,
        color=INK,
        bold=True,
    )
    t(
        s,
        Inches(7.25),
        Inches(3.8),
        Inches(5.0),
        Inches(2.2),
        "规则、依据、结果留在平台。人可以交接。这是数字化建设真正落地的那一步。",
        size=15,
        color=MUTED,
        line_pt=26,
    )
    footer(s, 4)

    s = prs.slides.add_slide(blank)
    deck.paint_light(s)
    kicker(s, "04   TWO BOOKS")
    title(s, "两本钱要对齐，设变不能再是糊涂账")
    cards3 = [
        ("商务", "总成价", "开票给奇瑞的售价。谈价、开票、对账看这一本。"),
        ("采购", "采购对象", "向二级商采购的价。该付多少、协议怎么签看这一本。"),
        ("工程", "当期用当期", "Excel 里设变是糊涂账。变更进系统，开票和应付一起跟上。"),
    ]
    for i, (who, head, body) in enumerate(cards3):
        x = Inches(0.7 + i * 4.15)
        card(s, x, Inches(1.7), Inches(3.95), Inches(4.85))
        t(s, x + Inches(0.28), Inches(2.05), Inches(3.4), Inches(0.32), who, size=13, color=ORANGE)
        t(
            s,
            x + Inches(0.28),
            Inches(2.5),
            Inches(3.4),
            Inches(0.85),
            head,
            size=22,
            color=INK,
            bold=True,
        )
        t(
            s,
            x + Inches(0.28),
            Inches(3.55),
            Inches(3.4),
            Inches(2.5),
            body,
            size=15,
            color=MUTED,
            line_pt=26,
        )
    footer(s, 5)

    s = prs.slides.add_slide(blank)
    deck.paint_light(s)
    kicker(s, "05   FIRST VALUE")
    title(s, "第一价值：降低对人的依赖")
    card(s, Inches(0.7), Inches(1.7), Inches(5.7), Inches(4.85))
    t(s, Inches(1.05), Inches(2.05), Inches(5), Inches(0.3), "以前", size=13, color=DIM)
    t(
        s,
        Inches(1.05),
        Inches(2.5),
        Inches(5.1),
        Inches(1.15),
        "熟练人员\n好几天",
        size=32,
        color=INK,
        bold=True,
        line_pt=42,
    )
    t(
        s,
        Inches(1.05),
        Inches(4.4),
        Inches(4.95),
        Inches(1.7),
        "一事一议、总成、采购对象靠骨干手感。\n人一走，账就停。",
        size=15,
        color=MUTED,
        line_pt=26,
    )
    card(s, Inches(6.9), Inches(1.7), Inches(5.7), Inches(4.85))
    t(s, Inches(7.25), Inches(2.05), Inches(5), Inches(0.3), "现在", size=13, color=ORANGE)
    t(
        s,
        Inches(7.25),
        Inches(2.5),
        Inches(5.1),
        Inches(1.15),
        "上传依据\n约 20 秒",
        size=32,
        color=INK,
        bold=True,
        line_pt=42,
    )
    t(
        s,
        Inches(7.25),
        Inches(4.4),
        Inches(5.0),
        Inches(1.7),
        "规则在引擎里。人请假，总成账和采购账照样能出。\n能力属于公司。",
        size=15,
        color=MUTED,
        line_pt=26,
    )
    footer(s, 6)

    s = prs.slides.add_slide(blank)
    deck.paint_light(s)
    kicker(s, "06   SECOND VALUE")
    title(s, "第二价值：溯源清楚，不必再扯皮")
    traces = [
        ("哪份文件", "当时用的 Excel 能找回，不是私人拷贝。"),
        ("谁操作的", "上传、计算、变更有日志。对账问记录。"),
        ("当期哪版", "钢价、BOM、工艺、设变按月对齐。"),
        ("同一套规则", "开给奇瑞和付给二级商，同一引擎。"),
    ]
    for i, (head, body) in enumerate(traces):
        col, row = i % 2, i // 2
        x = Inches(0.7 + col * 6.3)
        y = Inches(1.65 + row * 2.5)
        card(s, x, y, Inches(6.05), Inches(2.28))
        t(
            s,
            x + Inches(0.32),
            y + Inches(0.35),
            Inches(5.4),
            Inches(0.5),
            head,
            size=22,
            color=INK,
            bold=True,
        )
        t(
            s,
            x + Inches(0.32),
            y + Inches(1.05),
            Inches(5.4),
            Inches(0.85),
            body,
            size=16,
            color=MUTED,
            line_pt=26,
        )
    footer(s, 7)

    s = prs.slides.add_slide(blank)
    deck.paint_light(s)
    kicker(s, "07   VALUE")
    title(s, "四项价值，进的是公司而不是个人")
    values = [
        ("01", "少绑人", "核价能力属于公司。熟练骨干好几天，变成约 20 秒。"),
        ("02", "少扯皮", "文件、操作、当期依据可对。对奇瑞、对二级商、对财务同一本账。"),
        ("03", "算得一致", "同一套规则、同一套口径。\n总成售价和采购应付不再各算各的。"),
        ("04", "每月跟得上", "钢价、BOM、工艺、工程变更进入系统。\n当月依据，当月计算。"),
    ]
    for i, (num, head, body) in enumerate(values):
        col, row = i % 2, i // 2
        x = Inches(0.7 + col * 6.3)
        y = Inches(1.65 + row * 2.5)
        card(s, x, y, Inches(6.05), Inches(2.28))
        t(
            s,
            x + Inches(0.32),
            y + Inches(0.28),
            Inches(1.2),
            Inches(0.32),
            num,
            size=13,
            color=ORANGE,
            font=FONT_EN,
        )
        t(
            s,
            x + Inches(0.32),
            y + Inches(0.65),
            Inches(5.4),
            Inches(0.45),
            head,
            size=22,
            color=INK,
            bold=True,
        )
        t(
            s,
            x + Inches(0.32),
            y + Inches(1.2),
            Inches(5.4),
            Inches(0.8),
            body,
            size=14,
            color=MUTED,
            line_pt=22,
        )
    footer(s, 8)

    s = prs.slides.add_slide(blank)
    deck.paint_light(s)
    kicker(s, "08   AI")
    title(s, "引擎把账算准，AI 把账讲明白")
    t(
        s,
        Inches(0.7),
        Inches(1.48),
        Inches(12),
        Inches(0.4),
        "这才是 AI 落地：不是让模型编价格，是让管理人员直接问公司的账。",
        size=16,
        color=MUTED,
    )
    questions = [
        "这个月总成开给奇瑞的价，为什么涨了？",
        "向二级商采购的对象，哪些涨幅最大？",
        "工程变更之后，开票价和采购价各动了多少？",
        "这个总成和上个月差在哪里？",
        "把本月成本变化整理一下。",
    ]
    for i, q in enumerate(questions):
        y = Inches(2.0 + i * 0.9)
        rect(s, Inches(0.7), y, Inches(11.93), Inches(0.8), deck.WHITE, LINE)
        rect(s, Inches(0.7), y, Inches(0.055), Inches(0.8), ORANGE)
        t(
            s,
            Inches(1.05),
            y + Inches(0.14),
            Inches(0.7),
            Inches(0.5),
            f"Q{i + 1}",
            size=12,
            color=ORANGE,
            font=FONT_EN,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        t(
            s,
            Inches(1.9),
            y + Inches(0.14),
            Inches(10.3),
            Inches(0.5),
            q,
            size=17,
            color=INK,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    footer(s, 9)

    s = prs.slides.add_slide(blank)
    deck.paint_light(s)
    kicker(s, "09   FLOW")
    title(s, "从找表，到直接问")
    before = ["人找 Excel", "找数据", "找公式", "做分析", "得出结论"]
    after = ["直接问", "系统取数", "核价引擎计算", "AI 解释"]
    t(s, Inches(0.7), Inches(1.55), Inches(5.5), Inches(0.32), "过去", size=13, color=DIM)
    t(s, Inches(7.1), Inches(1.55), Inches(5.5), Inches(0.32), "现在", size=13, color=ORANGE)
    for i, label in enumerate(before):
        y = Inches(1.95 + i * 0.9)
        rect(s, Inches(0.7), y, Inches(5.5), Inches(0.78), SAND, LINE)
        t(
            s,
            Inches(1.0),
            y + Inches(0.14),
            Inches(4.9),
            Inches(0.5),
            f"{i + 1:02d}    {label}",
            size=16,
            color=MUTED,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    for i, label in enumerate(after):
        y = Inches(1.95 + i * 1.12)
        rect(s, Inches(7.1), y, Inches(5.5), Inches(0.98), deck.WHITE, LINE)
        rect(s, Inches(7.1), y, Inches(5.5), Inches(0.045), ORANGE)
        t(
            s,
            Inches(7.45),
            y + Inches(0.22),
            Inches(4.9),
            Inches(0.54),
            f"{i + 1:02d}    {label}",
            size=18,
            color=INK,
            bold=True,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    footer(s, 10)

    s = prs.slides.add_slide(blank)
    deck.paint_light(s)
    kicker(s, "10   ONE LEDGER")
    title(s, "采购、财务、管理人员，看见同一本账")
    pillars = [
        ("Data", "同一套数据", "文件可溯，不是私人拷贝。"),
        ("Rule", "同一套规则", "开给奇瑞、付给二级商，同一引擎。"),
        ("Result", "同一套结果", "当月依据，当月结果。对账不必扯皮。"),
    ]
    for i, (en, zh, extra) in enumerate(pillars):
        x = Inches(0.7 + i * 4.15)
        card(s, x, Inches(1.7), Inches(3.95), Inches(4.85))
        t(
            s,
            x + Inches(0.32),
            Inches(2.1),
            Inches(3.3),
            Inches(0.36),
            en,
            size=13,
            color=ORANGE,
            font=FONT_EN,
            spacing=80,
        )
        t(
            s,
            x + Inches(0.32),
            Inches(2.7),
            Inches(3.3),
            Inches(1.4),
            zh,
            size=26,
            color=INK,
            bold=True,
        )
        t(
            s,
            x + Inches(0.32),
            Inches(4.4),
            Inches(3.3),
            Inches(1.6),
            extra,
            size=15,
            color=MUTED,
            line_pt=26,
        )
    footer(s, 11)

    s = prs.slides.add_slide(blank)
    deck.paint_dark(s)
    hline(s, Inches(0.7), Inches(1.4), Inches(1.6), ORANGE, 2.25)
    t(
        s,
        Inches(0.7),
        Inches(1.65),
        Inches(12),
        Inches(0.32),
        "ONE LINE",
        size=13,
        color=ORANGE,
        font=FONT_EN,
        spacing=160,
    )
    t(
        s,
        Inches(0.7),
        Inches(2.15),
        Inches(12),
        Inches(1.85),
        "先把核价做成公司能力。\n数字化建设和 AI 落地，就有了抓手。",
        size=28,
        color=WHITE,
        bold=True,
        line_pt=44,
    )
    t(
        s,
        Inches(0.7),
        Inches(4.25),
        Inches(12),
        Inches(1.2),
        "少绑人、少扯皮。开给奇瑞多少、付给二级商多少，约 20 秒能算清。",
        size=18,
        color=SOFT,
    )
    dark_footer(s, 12)

    out = Path(__file__).resolve().parents[1] / "docs" / f"{APP_NAME}-产品介绍.pptx"
    prs.save(str(out))
    return out


if __name__ == "__main__":
    path = make()
    print(path)
