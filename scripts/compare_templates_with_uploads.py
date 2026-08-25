"""对比下载模板与 uploads 实样：表头关键词、合并结构、Parser 可用性。"""

from __future__ import annotations

import tempfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
UPLOADS = ROOT / "uploads"


def _cell_text(sheet: Worksheet, row: int, col: int) -> str:
    value = sheet.cell(row, col).value
    if value is None:
        return ""
    return str(value).replace("\n", "").replace("\r", "").strip()


def _merge_keys(sheet: Worksheet, max_row: int = 20) -> list[str]:
    keys: list[str] = []
    for merged in sorted(sheet.merged_cells.ranges, key=lambda item: (item.min_row, item.min_col)):
        if merged.min_row <= max_row:
            keys.append(str(merged))
    return keys


def _has_all(haystack: str, needles: list[str]) -> list[str]:
    return [item for item in needles if item not in haystack]


def check_stamping() -> list[str]:
    from app.parsers.stamping.stamping_parser import parse_stamping_excel
    from app.services.stamping.upload.template_builder import (
        build_stamping_process_template,
    )

    issues: list[str] = []
    content, name = build_stamping_process_template()
    tpl = load_workbook(BytesIO(content), data_only=False)["冲压工艺确认表"]

    real_path = sorted((UPLOADS / "芜湖" / "T13J" / "V3").glob("*.xlsx"))[0]
    real = load_workbook(real_path, data_only=False).active
    assert real is not None

    # 表头关键词
    tpl_header = " ".join(_cell_text(tpl, 5, col) + _cell_text(tpl, 6, col) for col in range(1, 20))
    real_header = " ".join(
        _cell_text(real, 5, col) + _cell_text(real, 6, col) for col in range(1, 20)
    )
    required = ["零件号", "零件名称", "材质", "料厚", "工序名称", "工序", "冲压设备"]
    missing_tpl = _has_all(tpl_header, required)
    missing_real = _has_all(real_header, required)
    if missing_tpl:
        issues.append(f"冲压模板缺表头: {missing_tpl}")
    if missing_real:
        issues.append(f"冲压实样缺表头(对照基准异常): {missing_real}")

    # 关键合并模式
    tpl_merges = set(_merge_keys(tpl, 14))
    must_have = {"H5:M5", "N5:P5", "A5:A6", "B5:B6", "B7:B9"}
    missing_merge = sorted(must_have - tpl_merges)
    if missing_merge:
        issues.append(f"冲压模板缺关键合并: {missing_merge}")

    # 实样也有同类模式
    real_merges = " ".join(_merge_keys(real, 20))
    for token in ("H5:M5", "N5:P5", "A5:A6", "B5:B6"):
        if token not in real_merges:
            issues.append(f"冲压实样缺少预期合并 {token}")

    path = Path(tempfile.gettempdir()) / name
    path.write_bytes(content)
    parsed = parse_stamping_excel(str(path))
    if len(parsed.parts) < 3 or len(parsed.processes) < 6:
        issues.append(
            f"冲压模板解析不足: parts={len(parsed.parts)} processes={len(parsed.processes)}"
        )

    print(
        f"[冲压] 模板合并={len(tpl.merged_cells.ranges)} "
        f"实样合并={len(real.merged_cells.ranges)} "
        f"解析 parts={len(parsed.parts)} processes={len(parsed.processes)} "
        f"-> {'OK' if not issues else 'FAIL'}"
    )
    return issues


def check_bom() -> list[str]:
    from app.parsers.bom.bom_parser import parse_bom_excel
    from app.services.bom.bom_template_builder import build_bom_material_template

    issues: list[str] = []
    content, _name = build_bom_material_template()
    tpl = load_workbook(BytesIO(content), data_only=False)["车身件材料清单"]

    real_path = next((UPLOADS / "bom" / "青岛" / "T28" / "D_00").glob("*.xlsx"))
    real = load_workbook(real_path, data_only=False).active
    assert real is not None

    tpl_h = " ".join(_cell_text(tpl, 4, col) for col in range(1, 24))
    real_h = " ".join(_cell_text(real, 4, col) for col in range(1, 24))
    required = ["总成", "一级", "二级", "材料", "料厚"]
    for label, text in (("模板", tpl_h), ("实样", real_h)):
        missing = _has_all(text, required)
        if missing:
            issues.append(f"BOM{label}缺表头关键词: {missing}")

    tpl_merges = set(_merge_keys(tpl, 8))
    must_have = {"A1:W1", "A2:W2", "A3:H3", "A5:A8", "H6:H7"}
    missing_merge = sorted(must_have - tpl_merges)
    if missing_merge:
        issues.append(f"BOM模板缺关键合并: {missing_merge}")

    real_merges = set(_merge_keys(real, 8))
    for token in ("A1:W1", "A2:W2", "A3:H3"):
        if token not in real_merges:
            issues.append(f"BOM实样缺少预期合并 {token}")

    parsed = parse_bom_excel(content)
    if len(parsed.nodes) < 4 or parsed.meta.vehicle_model != "T28":
        issues.append(
            f"BOM模板解析异常: nodes={len(parsed.nodes)} vehicle={parsed.meta.vehicle_model}"
        )

    print(
        f"[BOM] 模板合并={len(tpl.merged_cells.ranges)} "
        f"实样合并={len(real.merged_cells.ranges)} "
        f"解析 nodes={len(parsed.nodes)} vehicle={parsed.meta.vehicle_model} "
        f"ver={parsed.meta.bom_version} -> {'OK' if not issues else 'FAIL'}"
    )
    return issues


def check_base_price() -> list[str]:
    from app.parsers.yishiyiyi.base_price import parse_base_price_excel
    from app.services.yishiyiyi.base_price.importing.template_builder import (
        build_steel_base_price_template,
    )

    issues: list[str] = []
    content, _name = build_steel_base_price_template()
    tpl = load_workbook(BytesIO(content), data_only=False)["Sheet1"]

    real_path = sorted((UPLOADS / "steel-price" / "2025" / "01").glob("*base-price*.xlsx"))[0]
    real = load_workbook(real_path, data_only=False).active
    assert real is not None

    for label, sheet in (("模板", tpl), ("实样", real)):
        joined = " ".join(_cell_text(sheet, row, col) for row in (4, 5) for col in range(1, 16))
        required = ["序号", "钢种", "牌号", "参考规格", "芜湖", "青岛", "钢厂", "到厂原卷"]
        missing = _has_all(joined, required)
        if missing:
            issues.append(f"基价{label}缺关键词: {missing}")

    tpl_merges = set(_merge_keys(tpl, 9))
    must_have = {"A4:A5", "B4:B5", "C4:D5", "E4:E5", "F4:G4", "H4:I4", "C6:D6"}
    missing_merge = sorted(must_have - tpl_merges)
    if missing_merge:
        issues.append(f"基价模板缺关键合并: {missing_merge}")

    real_merges = set(_merge_keys(real, 8))
    for token in ("A4:A5", "B4:B5", "F4:G4"):
        if token not in real_merges:
            issues.append(f"基价实样缺少预期合并 {token}")

    parsed = parse_base_price_excel(content)
    locations = {str(row.get("location")) for row in parsed.records}
    if len(parsed.records) < 8 or parsed.errors:
        issues.append(f"基价模板解析异常: records={len(parsed.records)} errors={parsed.errors}")
    if not {"芜湖", "青岛", "大连", "鄂尔多斯"}.issubset(locations):
        issues.append(f"基价模板基地不全: {sorted(locations)}")

    print(
        f"[基价] 模板合并={len(tpl.merged_cells.ranges)} "
        f"实样合并={len(real.merged_cells.ranges)} "
        f"解析 records={len(parsed.records)} locations={sorted(locations)} "
        f"-> {'OK' if not issues else 'FAIL'}"
    )
    return issues


def check_policy() -> list[str]:
    from app.parsers.yishiyiyi.policy import parse_steel_price_policy_excel
    from app.services.yishiyiyi.policy.importing.template_builder import (
        build_steel_price_policy_template,
    )

    issues: list[str] = []
    content, _name = build_steel_price_policy_template()
    tpl = load_workbook(BytesIO(content), data_only=False)["样例钢厂"]

    real_path = sorted((UPLOADS / "steel-price" / "2025" / "01").glob("*policy*.xlsx"))[0]
    real_wb = load_workbook(real_path, data_only=False)
    # 首钢排版与模板最接近
    real = real_wb["首钢"] if "首钢" in real_wb.sheetnames else real_wb.worksheets[0]

    tpl_text = " ".join(_cell_text(tpl, row, col) for row in range(1, 6) for col in range(1, 10))
    real_text = " ".join(_cell_text(real, row, col) for row in range(1, 6) for col in range(1, 16))
    if "厚度加减价" not in tpl_text:
        issues.append("政策模板缺「厚度加减价」标题")
    if "≥0.5" not in tpl_text:
        issues.append("政策模板缺厚度区间列头")
    if "厚度加减价" not in real_text and "厚度" not in real_text:
        issues.append("政策实样缺厚度相关标题(对照基准异常)")

    tpl_merges = set(_merge_keys(tpl, 5))
    must_have = {"A1:I1", "A3:A4", "B3:I3"}
    missing_merge = sorted(must_have - tpl_merges)
    if missing_merge:
        issues.append(f"政策模板缺关键合并: {missing_merge}")

    real_merges = set(_merge_keys(real, 5))
    # 首钢: A3:A4 钢种, B3:L3 厚度标题
    if "A3:A4" not in real_merges:
        issues.append("政策实样(首钢)缺少 A3:A4 钢种合并")

    parsed = parse_steel_price_policy_excel(content, year=2025, month=1)
    thickness_count = sum(1 for item in parsed.policies if item.policy_type == "THICKNESS")
    if thickness_count < 4:
        issues.append(f"政策模板解析不足: thickness={thickness_count}")

    print(
        f"[政策] 模板合并={len(tpl.merged_cells.ranges)} "
        f"实样({real.title})合并={len(real.merged_cells.ranges)} "
        f"解析 thickness={thickness_count} -> {'OK' if not issues else 'FAIL'}"
    )
    return issues


def check_flat_templates() -> list[str]:
    issues: list[str] = []

    from app.parsers.stamping.tonnage_price_parser import parse_tonnage_price_excel
    from app.services.stamping.tonnage_price_import_service import (
        StampingTonnagePriceImportService,
    )

    tonnage = parse_tonnage_price_excel(StampingTonnagePriceImportService().build_template()[0])
    print(f"[冲次] rows={len(tonnage.rows)} -> {'OK' if len(tonnage.rows) >= 3 else 'FAIL'}")
    if len(tonnage.rows) < 3:
        issues.append("冲次模板解析不足")

    from app.parsers.bom.welding_unit_price_parser import parse_welding_unit_price_excel
    from app.services.bom.welding_unit_price_import_service import (
        BomWeldingUnitPriceImportService,
    )

    welding = parse_welding_unit_price_excel(BomWeldingUnitPriceImportService().build_template()[0])
    print(f"[焊接] rows={len(welding.rows)} -> {'OK' if len(welding.rows) >= 3 else 'FAIL'}")
    if len(welding.rows) < 3:
        issues.append("焊接模板解析不足")

    from app.parsers.bom.annual_part_price_parser import parse_annual_part_price_excel
    from app.services.bom.annual_part_price_import_service import (
        BomAnnualPartPriceImportService,
    )

    annual = parse_annual_part_price_excel(BomAnnualPartPriceImportService().build_template()[0])
    print(f"[年度零件价] rows={len(annual.rows)} -> {'OK' if len(annual.rows) >= 3 else 'FAIL'}")
    if len(annual.rows) < 3:
        issues.append("年度零件价模板解析不足")

    from app.parsers.bom.packaging_transport_fee_parser import (
        parse_packaging_transport_fee_excel,
    )
    from app.services.bom.packaging_transport_fee_import_service import (
        BomPackagingTransportFeeImportService,
    )

    packaging = parse_packaging_transport_fee_excel(
        BomPackagingTransportFeeImportService().build_template()[0]
    )
    print(
        f"[包装运输费] rows={len(packaging.rows)} -> "
        f"{'OK' if len(packaging.rows) >= 1 else 'FAIL'}"
    )
    if len(packaging.rows) < 1:
        issues.append("包装运输费模板解析不足")

    return issues


def main() -> None:
    all_issues: list[str] = []
    all_issues.extend(check_stamping())
    all_issues.extend(check_bom())
    all_issues.extend(check_base_price())
    all_issues.extend(check_policy())
    all_issues.extend(check_flat_templates())

    print()
    if all_issues:
        print("MISMATCH / FAIL:")
        for item in all_issues:
            print(" -", item)
        raise SystemExit(1)
    print("MATCH OK — 下载模板与 uploads 实样结构/关键词/合并模式一致，且均可被 Parser 解析")


if __name__ == "__main__":
    main()
