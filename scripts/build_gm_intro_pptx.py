"""Generate PwC-style GM intro deck (16:9)."""

from __future__ import annotations

from contextlib import suppress
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

# PwC
ORANGE = RGBColor(0xD0, 0x4A, 0x02)
INK = RGBColor(0x2D, 0x2D, 0x2D)
CHARCOAL = RGBColor(0x2B, 0x2B, 0x2B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xFA, 0xF7, 0xF2)
SAND = RGBColor(0xF3, 0xEE, 0xE6)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
DIM = RGBColor(0x8A, 0x8A, 0x8A)
LINE = RGBColor(0xE6, 0xE0, 0xD6)
SOFT = RGBColor(0xC8, 0xC8, 0xC8)

W = Inches(13.333)
H = Inches(7.5)
FONT_CN = "微软雅黑"
FONT_EN = "Arial"
COMPANY = "普炳软件"
APP_NAME = "结构件自动核价平台"
APP_NAME_SHORT = "核价平台"
APP_NAME_EN = "Structural Parts Automated Costing Platform"
TRIAL_URL = "http://117.143.214.90:8000/"
TRIAL_USERNAME = "admin"
TRIAL_PASSWORD = "admin123"

SLIDE_COUNT = 9


def rect(slide, left, top, width, height, fill: RGBColor, line: RGBColor | None = None) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    if line is not None:
        shape.line.fill.solid()
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill


def hline(slide, left, top, width, color: RGBColor, weight_pt: float = 0.75) -> None:
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, left, top, left + width, top)
    connector.line.color.rgb = color
    connector.line.width = Pt(weight_pt)


def textbox(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size: float,
    color: RGBColor,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    font: str = FONT_CN,
    anchor=MSO_ANCHOR.TOP,
    spacing: int | None = None,
    line_pt: float | None = None,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    with suppress(Exception):
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    p = tf.paragraphs[0]
    chunks = text.split("\n")
    first = True
    for chunk in chunks:
        if first:
            first = False
            p.alignment = align
        else:
            p = tf.add_paragraph()
            p.alignment = align
        run = p.add_run()
        run.text = chunk
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        rpr = run._r.get_or_add_rPr()
        ea = OxmlElement("a:ea")
        ea.set("typeface", FONT_CN)
        rpr.append(ea)
        latin = OxmlElement("a:latin")
        latin.set("typeface", font)
        rpr.append(latin)
        if spacing is not None:
            rpr.set("spc", str(spacing))
        if line_pt is not None:
            p_pr = p._p.get_or_add_pPr()
            ln_spc = OxmlElement("a:lnSpc")
            spc_pts = OxmlElement("a:spcPts")
            spc_pts.set("val", str(int(line_pt * 100)))
            ln_spc.append(spc_pts)
            p_pr.append(ln_spc)


def paint_light(slide) -> None:
    rect(slide, Inches(0), Inches(0), W, H, CREAM)
    rect(slide, Inches(0), Inches(0), Inches(0.12), H, ORANGE)


def paint_dark(slide) -> None:
    rect(slide, Inches(0), Inches(0), W, H, CHARCOAL)
    rect(slide, Inches(0), Inches(0), Inches(0.12), H, ORANGE)


def footer(slide, index: int) -> None:
    hline(slide, Inches(0.7), Inches(7.12), Inches(11.93), LINE, 0.6)
    textbox(
        slide,
        Inches(0.7),
        Inches(7.18),
        Inches(6),
        Inches(0.28),
        f"{COMPANY}  ·  {APP_NAME_SHORT}",
        size=10,
        color=DIM,
        font=FONT_CN,
        spacing=40,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    textbox(
        slide,
        Inches(10.4),
        Inches(7.18),
        Inches(2.2),
        Inches(0.28),
        f"{index} / {SLIDE_COUNT}",
        size=10,
        color=DIM,
        font=FONT_EN,
        align=PP_ALIGN.RIGHT,
        spacing=40,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def kicker(slide, text: str, y: float = 0.38) -> None:
    rect(slide, Inches(0.7), Inches(y + 0.08), Inches(0.28), Inches(0.06), ORANGE)
    textbox(
        slide,
        Inches(1.1),
        Inches(y),
        Inches(10),
        Inches(0.28),
        text,
        size=11,
        color=ORANGE,
        font=FONT_EN,
        spacing=200,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def title_block(slide, title: str, y: float = 0.72) -> None:
    textbox(
        slide,
        Inches(0.7),
        Inches(y),
        Inches(12),
        Inches(0.7),
        title,
        size=28,
        color=INK,
        bold=True,
        spacing=20,
    )


def card(slide, left, top, width, height) -> None:
    rect(slide, left, top, width, height, WHITE, LINE)
    rect(slide, left, top, width, Inches(0.045), ORANGE)


def cover_trial_access(slide) -> None:
    """封面试用入口。产品介绍与简介共用。"""
    textbox(
        slide,
        Inches(0.7),
        Inches(5.2),
        Inches(11.4),
        Inches(0.4),
        f"试用地址  {TRIAL_URL}",
        size=16,
        color=ORANGE,
    )
    textbox(
        slide,
        Inches(0.7),
        Inches(5.65),
        Inches(11.4),
        Inches(0.4),
        f"默认账号  {TRIAL_USERNAME} / {TRIAL_PASSWORD}",
        size=16,
        color=SOFT,
    )


def dark_footer(slide, index: int) -> None:
    textbox(
        slide,
        Inches(0.7),
        Inches(6.85),
        Inches(8),
        Inches(0.3),
        COMPANY,
        size=12,
        color=SOFT,
        font=FONT_CN,
    )
    textbox(
        slide,
        Inches(10.4),
        Inches(6.85),
        Inches(2.2),
        Inches(0.3),
        f"{index} / {SLIDE_COUNT}",
        size=12,
        color=SOFT,
        font=FONT_EN,
        align=PP_ALIGN.RIGHT,
    )


def make() -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # 01 cover
    s = prs.slides.add_slide(blank)
    paint_dark(s)
    textbox(
        s,
        Inches(0.7),
        Inches(0.42),
        Inches(6),
        Inches(0.36),
        COMPANY,
        size=16,
        color=WHITE,
        bold=True,
        spacing=40,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    hline(s, Inches(0.7), Inches(1.28), Inches(1.6), ORANGE, 2.25)
    textbox(
        s,
        Inches(0.7),
        Inches(1.48),
        Inches(12),
        Inches(0.32),
        "从靠人做表，到公司能力",
        size=14,
        color=ORANGE,
    )
    textbox(
        s,
        Inches(0.7),
        Inches(2.0),
        Inches(12),
        Inches(1.0),
        APP_NAME,
        size=40,
        color=WHITE,
        bold=True,
    )
    textbox(
        s,
        Inches(0.7),
        Inches(3.2),
        Inches(11.4),
        Inches(1.6),
        "管理人员是公司的建设者。\n进出两头的钱，不能再交给几个人的表。",
        size=22,
        color=SOFT,
        line_pt=36,
    )
    cover_trial_access(s)
    dark_footer(s, 1)

    # 02 crime / evidence
    s = prs.slides.add_slide(blank)
    paint_light(s)
    kicker(s, "01   THE COST")
    title_block(s, "现在核价绑在人身上，对账只能扯皮")
    card(s, Inches(0.7), Inches(1.7), Inches(5.7), Inches(4.85))
    textbox(
        s, Inches(1.05), Inches(2.05), Inches(5), Inches(0.3), "对人依赖", size=13, color=ORANGE
    )
    textbox(
        s,
        Inches(1.05),
        Inches(2.5),
        Inches(5.1),
        Inches(1.15),
        "会做表的人一走，\n账就停",
        size=24,
        color=INK,
        bold=True,
        line_pt=36,
    )
    textbox(
        s,
        Inches(1.05),
        Inches(4.2),
        Inches(4.95),
        Inches(1.9),
        "一事一议、总成、采购对象在骨干电脑里。\n人请假、离职，公司没有第二本账。",
        size=15,
        color=MUTED,
        line_pt=26,
    )
    card(s, Inches(6.9), Inches(1.7), Inches(5.7), Inches(4.85))
    textbox(s, Inches(7.25), Inches(2.05), Inches(5), Inches(0.3), "扯皮", size=13, color=ORANGE)
    textbox(
        s,
        Inches(7.25),
        Inches(2.5),
        Inches(5.1),
        Inches(1.15),
        "各执一份 Excel，\n说不清哪一版",
        size=24,
        color=INK,
        bold=True,
        line_pt=36,
    )
    textbox(
        s,
        Inches(7.25),
        Inches(4.2),
        Inches(5.0),
        Inches(1.9),
        "对奇瑞、对二级商、对内部财务，各拿各的表。用过哪份文件、谁改的、当期哪版，对不上就只能吵。",
        size=15,
        color=MUTED,
        line_pt=26,
    )
    footer(s, 2)

    # 03 two books + change
    s = prs.slides.add_slide(blank)
    paint_light(s)
    kicker(s, "02   THE MONEY")
    title_block(s, "漏的是两本钱，吵的是设变这本糊涂账")
    cards3 = [
        (
            "商务",
            "总成价",
            "开票给奇瑞的售价。谈价、开票、对账看这一本。\n算不准就是少收或超承诺。",
        ),
        (
            "采购",
            "采购对象",
            "向二级商采购的价。该付多少、协议怎么签看这一本。\n算不准就是多付或付旧件。",
        ),
        (
            "工程",
            "设变",
            "Excel 里改一层漏一层。\n商务可能还按旧总成开票，采购可能还按旧对象付款。",
        ),
    ]
    for i, (who, title, body) in enumerate(cards3):
        x = Inches(0.7 + i * 4.15)
        card(s, x, Inches(1.7), Inches(3.95), Inches(4.85))
        textbox(
            s, x + Inches(0.28), Inches(2.05), Inches(3.4), Inches(0.32), who, size=13, color=ORANGE
        )
        textbox(
            s,
            x + Inches(0.28),
            Inches(2.5),
            Inches(3.4),
            Inches(0.7),
            title,
            size=22,
            color=INK,
            bold=True,
        )
        textbox(
            s,
            x + Inches(0.28),
            Inches(3.5),
            Inches(3.4),
            Inches(2.5),
            body,
            size=15,
            color=MUTED,
            line_pt=26,
        )
    footer(s, 3)

    # 04 so not buying is indefensible
    s = prs.slides.add_slide(blank)
    paint_light(s)
    kicker(s, "03   THE CHOICE")
    title_block(s, "不做成公司能力，就是把公司的钱交给个人记忆")
    card(s, Inches(0.7), Inches(1.7), Inches(5.7), Inches(4.85))
    textbox(s, Inches(1.05), Inches(2.05), Inches(5), Inches(0.3), "继续 Excel", size=13, color=DIM)
    textbox(
        s,
        Inches(1.05),
        Inches(2.5),
        Inches(5.1),
        Inches(1.2),
        "进出两头的账\n在私人文件夹里",
        size=22,
        color=INK,
        bold=True,
        line_pt=34,
    )
    textbox(
        s,
        Inches(1.05),
        Inches(4.15),
        Inches(4.95),
        Inches(1.9),
        "人走、表散、设变对不上。开给奇瑞的价和付给二级商的价，对账时只能各执一词。\n管理人员说不过去。",
        size=15,
        color=MUTED,
        line_pt=26,
    )
    card(s, Inches(6.9), Inches(1.7), Inches(5.7), Inches(4.85))
    textbox(
        s, Inches(7.25), Inches(2.05), Inches(5), Inches(0.3), "做成能力", size=13, color=ORANGE
    )
    textbox(
        s,
        Inches(7.25),
        Inches(2.5),
        Inches(5.1),
        Inches(1.2),
        "规则、依据、结果\n留在公司",
        size=22,
        color=INK,
        bold=True,
        line_pt=34,
    )
    textbox(
        s,
        Inches(7.25),
        Inches(4.15),
        Inches(5.0),
        Inches(1.9),
        "不是再做一张更大的表。是人可以交接，依据可以拿出，两本钱可以对上。\n这才是建设。",
        size=15,
        color=MUTED,
        line_pt=26,
    )
    footer(s, 4)

    # 05 first value
    s = prs.slides.add_slide(blank)
    paint_light(s)
    kicker(s, "04   FIRST VALUE")
    title_block(s, "第一价值：核价不再绑在几个人身上")
    card(s, Inches(0.7), Inches(1.7), Inches(5.7), Inches(4.85))
    textbox(s, Inches(1.05), Inches(2.05), Inches(5), Inches(0.3), "以前", size=13, color=DIM)
    textbox(
        s,
        Inches(1.05),
        Inches(2.5),
        Inches(5.1),
        Inches(1.15),
        "熟练人员\n好几天",
        size=32,
        color=INK,
        bold=True,
        line_pt=42,
    )
    textbox(
        s,
        Inches(1.05),
        Inches(4.4),
        Inches(4.95),
        Inches(1.7),
        "一事一议靠骨干手感。人一走，这件事就停。\n能力属于个人。",
        size=15,
        color=MUTED,
        line_pt=26,
    )
    card(s, Inches(6.9), Inches(1.7), Inches(5.7), Inches(4.85))
    textbox(s, Inches(7.25), Inches(2.05), Inches(5), Inches(0.3), "现在", size=13, color=ORANGE)
    textbox(
        s,
        Inches(7.25),
        Inches(2.5),
        Inches(5.1),
        Inches(1.15),
        "上传依据\n约 20 秒",
        size=32,
        color=INK,
        bold=True,
        line_pt=42,
    )
    textbox(
        s,
        Inches(7.25),
        Inches(4.4),
        Inches(5.0),
        Inches(1.7),
        "规则在引擎里。人请假，总成账和采购账照样能出。能力属于公司。",
        size=15,
        color=MUTED,
        line_pt=26,
    )
    footer(s, 5)

    # 06 second value
    s = prs.slides.add_slide(blank)
    paint_light(s)
    kicker(s, "05   SECOND VALUE")
    title_block(s, "第二价值：依据能拿出，就不必再扯皮")
    traces = [
        ("哪份文件", "当时用的 Excel 能找回，不是私人拷贝。"),
        ("谁操作的", "上传、计算、变更有日志。对账问记录。"),
        ("当期哪版", "钢价、BOM、工艺、设变按月对齐。"),
        ("同一套规则", "开给奇瑞和付给二级商，同一引擎。"),
    ]
    for i, (title, body) in enumerate(traces):
        col = i % 2
        row = i // 2
        x = Inches(0.7 + col * 6.3)
        y = Inches(1.65 + row * 2.5)
        card(s, x, y, Inches(6.05), Inches(2.28))
        textbox(
            s,
            x + Inches(0.32),
            y + Inches(0.35),
            Inches(5.4),
            Inches(0.5),
            title,
            size=22,
            color=INK,
            bold=True,
        )
        textbox(
            s,
            x + Inches(0.32),
            y + Inches(1.05),
            Inches(5.4),
            Inches(0.85),
            body,
            size=16,
            color=MUTED,
            line_pt=26,
        )
    footer(s, 6)

    # 07 lever AFTER the pain
    s = prs.slides.add_slide(blank)
    paint_light(s)
    kicker(s, "06   THE LEVER")
    title_block(s, "所以数字化、AI 落地，核价是最好的抓手")
    textbox(
        s,
        Inches(0.7),
        Inches(1.5),
        Inches(12),
        Inches(0.55),
        "建设不是再上一个聊天工具。是把最痛、最有规则、立刻碰到钱的事做成公司能力。",
        size=16,
        color=MUTED,
    )
    levers = [
        ("有规则", "钢价政策、工艺、BOM、设变，本来就有企业口径。\n适合写成引擎，而不是口头 AI。"),
        ("有真账", "开给奇瑞的总成价、付给二级商的采购价，每月都要对。\n建设立刻碰到钱。"),
        ("能见效", "账能复算、人能交接、依据能拿出。\nAI 问的是这本真账，不是编一个价。"),
    ]
    for i, (title, body) in enumerate(levers):
        x = Inches(0.7 + i * 4.15)
        card(s, x, Inches(2.25), Inches(3.95), Inches(4.3))
        textbox(
            s,
            x + Inches(0.28),
            Inches(2.55),
            Inches(3.4),
            Inches(0.5),
            title,
            size=22,
            color=INK,
            bold=True,
        )
        textbox(
            s,
            x + Inches(0.28),
            Inches(3.25),
            Inches(3.4),
            Inches(2.8),
            body,
            size=15,
            color=MUTED,
            line_pt=26,
        )
    footer(s, 7)

    # 08 outline only — hunger for demo
    s = prs.slides.add_slide(blank)
    paint_light(s)
    kicker(s, "07   THE PLATFORM")
    title_block(s, "平台管这些依据，算出那两本账")
    textbox(
        s,
        Inches(0.7),
        Inches(1.5),
        Inches(12),
        Inches(0.5),
        "细则见面看。这一页只回答：它究竟管什么。",
        size=16,
        color=MUTED,
    )
    items = [
        ("Policy", "钢价一事一议"),
        ("BOM", "结构与外购"),
        ("Process", "工艺与模具"),
        ("Change", "工程变更"),
        ("Costing", "件价与总成"),
        ("Contract", "采购协议"),
    ]
    for i, (en, zh) in enumerate(items):
        col = i % 3
        row = i // 3
        x = Inches(0.7 + col * 4.15)
        y = Inches(2.2 + row * 2.25)
        card(s, x, y, Inches(3.95), Inches(2.05))
        textbox(
            s,
            x + Inches(0.28),
            y + Inches(0.35),
            Inches(3.4),
            Inches(0.32),
            en,
            size=12,
            color=ORANGE,
            font=FONT_EN,
            spacing=80,
        )
        textbox(
            s,
            x + Inches(0.28),
            y + Inches(0.85),
            Inches(3.4),
            Inches(0.7),
            zh,
            size=20,
            color=INK,
            bold=True,
        )
    footer(s, 8)

    # 09 close — come see
    s = prs.slides.add_slide(blank)
    paint_dark(s)
    hline(s, Inches(0.7), Inches(1.25), Inches(1.6), ORANGE, 2.25)
    textbox(
        s,
        Inches(0.7),
        Inches(1.5),
        Inches(12),
        Inches(0.32),
        "COME SEE",
        size=13,
        color=ORANGE,
        font=FONT_EN,
        spacing=160,
    )
    textbox(
        s,
        Inches(0.7),
        Inches(1.95),
        Inches(12),
        Inches(1.35),
        "请来看三件事。",
        size=32,
        color=WHITE,
        bold=True,
    )
    sees = [
        "人走了，账还在不在。",
        "依据能不能当场拿出。",
        "约 20 秒，两本账能不能出清。",
    ]
    for i, line in enumerate(sees):
        y = Inches(3.4 + i * 0.7)
        textbox(
            s,
            Inches(0.7),
            y,
            Inches(0.7),
            Inches(0.5),
            f"{i + 1:02d}",
            size=14,
            color=ORANGE,
            font=FONT_EN,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        textbox(
            s,
            Inches(1.5),
            y,
            Inches(11),
            Inches(0.5),
            line,
            size=20,
            color=WHITE,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    textbox(
        s,
        Inches(0.7),
        Inches(5.7),
        Inches(12),
        Inches(0.7),
        "先把核价做成公司能力。数字化建设和 AI 落地，就有了抓手。",
        size=16,
        color=SOFT,
    )
    dark_footer(s, 9)

    out = Path(__file__).resolve().parents[1] / "docs" / f"{APP_NAME}-简介.pptx"
    prs.save(str(out))
    return out


if __name__ == "__main__":
    path = make()
    print(path)
