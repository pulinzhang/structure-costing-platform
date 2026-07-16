import openpyxl

wb = openpyxl.load_workbook('excelSample/副本奇瑞T28车身件项目冲压工艺确认表-烟台海众.xlsm', data_only=True)
ws = wb.active

# 检查几个典型零件的数据完整性
print('=== 检查工序数据完整性 ===\n')

# 零件1 (row 6-9)
print('零件 501008898AA (Row 6-9):')
for row in range(6, 10):
    print(f'  Row {row}: ', end='')
    for col in [12, 13, 14, 15, 16, 17, 18]:
        val = ws.cell(row=row, column=col).value
        print(f'{chr(64+col)}={val}', end=' ')
    print()

# 零件4 (row 27-31) - 这个工序信息不完整
print('\n零件 501008902AA (Row 27-31):')
for row in range(27, 32):
    print(f'  Row {row}: ', end='')
    for col in [12, 13, 14, 15, 16, 17, 18]:
        val = ws.cell(row=row, column=col).value
        print(f'{chr(64+col)}={val}', end=' ')
    print()

# 检查数据是否跨行
print('\n=== 检查跨行数据 ===')
# 零件1的工艺方法在第几列
part1_process_method = ws.cell(row=6, column=17).value
print(f'零件1 成形方法 (Q6): {part1_process_method}')
for row in range(7, 10):
    val = ws.cell(row=row, column=17).value
    print(f'Row {row} Q列: {val}')

# 看看有没有工序行有部分数据的情况
print('\n=== 工序行完整数据检查 ===')
incomplete_processes = []
for row_idx in range(6, ws.max_row + 1):
    process_seq = ws.cell(row=row_idx, column=12).value
    if process_seq and str(process_seq).strip().startswith('OP'):
        process_code = ws.cell(row=row_idx, column=13).value
        die_sets = ws.cell(row=row_idx, column=14).value
        tonnage = ws.cell(row=row_idx, column=15).value

        # 检查是否不完整
        if not process_code and not die_sets and not tonnage:
            part_info = ws.cell(row=row_idx, column=2).value
            if part_info:
                incomplete_processes.append({
                    'row': row_idx,
                    'part': part_info,
                    'process': process_seq
                })

print(f'不完整工序行数: {len(incomplete_processes)}')
if incomplete_processes:
    print('前5个不完整工序:')
    for p in incomplete_processes[:5]:
        print(f"  Row {p['row']}: {p['part']} - {p['process']}")
