"""冲次单价导出与导入：核心三列名称/单位/单价；列表导出另含基地/年份/审核列。"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook

from app.common.audit import SYSTEM_USER
from app.parsers.stamping.tonnage_price_parser import parse_tonnage_price_excel
from app.services.stamping.tonnage_price_export_service import StampingTonnagePriceExportService
from app.services.stamping.tonnage_price_import_service import StampingTonnagePriceImportService


def test_import_template_is_parseable() -> None:
    content, filename = StampingTonnagePriceImportService().build_template()
    assert filename.endswith(".xlsx")

    result = parse_tonnage_price_excel(content)
    assert not result.errors
    assert len(result.rows) >= 1
    assert result.rows[0].tonnage == 400
    assert result.rows[0].unit == "次"
    assert result.rows[0].price_per_stroke == Decimal("0.564")


def test_export_includes_scope_audit_and_roundtrips() -> None:
    records = [
        SimpleNamespace(
            location="芜湖",
            year=2026,
            tonnage=400,
            price_per_stroke=Decimal("0.564"),
            audit_status="approved",
            created_by="张三",
            created_at=None,
            updated_by="李四",
            updated_at=None,
        ),
        SimpleNamespace(
            location="芜湖",
            year=2026,
            tonnage=800,
            price_per_stroke=Decimal("1.026"),
            audit_status="pending",
            created_by="张三",
            created_at=None,
            updated_by=None,
            updated_at=None,
        ),
    ]
    service = StampingTonnagePriceExportService()
    service.repository = MagicMock()
    service.repository.list_for_export.return_value = records

    content, filename = service.export(
        MagicMock(),
        location="芜湖",
        year=2026,
        current_user=SYSTEM_USER,
    )
    assert "冲次单价" in filename
    assert "芜湖" in filename

    workbook = load_workbook(BytesIO(content), data_only=True)
    sheet = workbook.worksheets[0]
    headers = [sheet.cell(1, col).value for col in range(1, 11)]
    assert headers == [
        "基地",
        "年份",
        "名称",
        "单位",
        "单价",
        "审核状态",
        "创建人",
        "创建时间",
        "更新人",
        "更新时间",
    ]
    assert sheet.cell(2, 1).value == "芜湖"
    assert sheet.cell(2, 2).value == 2026
    assert sheet.cell(2, 3).value == "400T"
    assert sheet.cell(2, 4).value == "次"
    assert sheet.cell(2, 6).value == "已通过"
    workbook.close()

    result = parse_tonnage_price_excel(content)
    assert not result.errors
    assert [(row.tonnage, row.unit, row.price_per_stroke) for row in result.rows] == [
        (400, "次", Decimal("0.564")),
        (800, "次", Decimal("1.026")),
    ]


def test_legacy_export_headers_still_parseable() -> None:
    """旧导出列名仍可解析，避免历史文件完全不可用。"""
    wb = Workbook()
    ws = wb.active
    ws.append(["基地", "年份", "压机吨位(T)", "单冲次价格(元)"])
    ws.append(["芜湖", 2026, 400, 0.564])
    ws.append(["芜湖", 2026, 500, 0.622])
    buffer = BytesIO()
    wb.save(buffer)

    result = parse_tonnage_price_excel(buffer.getvalue())
    assert not result.errors
    assert [row.tonnage for row in result.rows] == [400, 500]
