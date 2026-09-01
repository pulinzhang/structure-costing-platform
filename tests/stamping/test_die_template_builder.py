"""模具布点导入模板对齐样本 A～AC 表头，且能被 Parser 读取。"""

from __future__ import annotations

import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.parsers.stamping_die.parser import parse_die_layout_excel
from app.services.stamping.die.template_builder import build_die_layout_template


def test_die_layout_template_matches_sample_header_and_parses() -> None:
    content, filename = build_die_layout_template()
    assert filename.endswith(".xlsx")

    path = Path(tempfile.gettempdir()) / filename
    path.write_bytes(content)
    workbook = load_workbook(path)
    try:
        sheet = workbook.active
        assert sheet.title == "零件汇总"
        assert sheet["A1"].value == "T13J项目冲压模具布点工艺清单"
        assert sheet["AA1"].value == "版本"
        assert sheet["AC1"].value == "V0-01"
        assert sheet["A3"].value == "序号"
        assert sheet["B3"].value == "产品信息"
        assert sheet["H3"].value == "工艺清单"
        assert sheet["B4"].value == "零件号"
        assert sheet["D4"].value == "简图"
        assert sheet["N5"].value == "工序名称"
        assert sheet["Q5"].value == "长度"
        assert sheet["AC4"].value == "备注"
        assert sheet["AD1"].value is None
        assert sheet["AD4"].value is None

        merged = {str(rng) for rng in sheet.merged_cells.ranges}
        assert "A1:Z2" in merged
        assert "B3:G3" in merged
        assert "H3:AC3" in merged
        assert "A3:A5" in merged
        assert "H4:I4" in merged
        assert "N4:P4" in merged
        assert "Q4:S4" in merged
        assert "B6:B10" in merged
        assert "A12:AC12" in merged
        note_blob = " ".join(
            str(sheet.cell(row=row, column=1).value or "") for row in range(12, 22)
        )
        assert "模具费分摊" in note_blob
        assert "双槽" in note_blob
        assert "一出二" in note_blob
        assert "0.50" in note_blob
        assert "0.80" in note_blob
        assert "第一序" in note_blob
        assert "工艺确认表" in note_blob
        assert sheet["B12"].value is None

        last_col = max(
            cell.column
            for row in sheet.iter_rows(min_row=1, max_row=10)
            for cell in row
            if cell.value
        )
        assert last_col <= 29
        assert get_column_letter(last_col) <= "AC"
    finally:
        workbook.close()

    parsed = parse_die_layout_excel(str(path), filename=filename)
    assert parsed.vehicle_model == "T13J"
    assert len(parsed.parts) == 1
    part = parsed.parts[0]
    assert part.part_number == "502007709AA"
    assert [item.process_seq for item in part.processes] == [
        "OP05",
        "OP10",
        "OP20",
        "OP30",
        "OP40",
    ]
    assert part.processes[0].process_code == "BL"
    assert part.die_type == "钢板模"
    assert part.die_method == "双槽"
