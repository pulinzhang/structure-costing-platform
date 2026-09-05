"""工艺确认表可以没有「模具编号」列表头。"""

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.parsers.stamping.stamping_parser import parse_stamping_excel
from app.services.stamping.upload.template_builder import build_stamping_process_template


def test_parse_without_die_number_header(tmp_path: Path) -> None:
    content, _filename = build_stamping_process_template()
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["冲压工艺确认表"]
    for row in sheet.iter_rows(min_row=1, max_row=6, min_col=17, max_col=17):
        for cell in row:
            if cell.value and "模具编号" in str(cell.value):
                cell.value = None
    path = tmp_path / "process-no-die-number.xlsx"
    workbook.save(path)
    workbook.close()

    result = parse_stamping_excel(str(path))
    assert "die_number" not in result.col_map
    assert result.parts
    assert all(item.die_number is None for item in result.processes)
