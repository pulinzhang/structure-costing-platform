"""冲压工艺确认表导出含短码与是否摆剪。"""

from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from app.services.stamping.part.export.constants import (
    COL_COUNT,
    DATA_START_ROW,
    SHORT_CODE_COL,
    STEEL_COL_COUNT,
    STEEL_DATA_START_ROW,
    STEEL_MUTE_COLS,
    STEEL_QTY_COL,
    STEEL_SHORT_CODE_COL,
    STEEL_TON_COL,
    STEEL_WEIGHT_COL,
    STEEL_WIDTHS,
    WIDTHS,
)
from app.services.stamping.part.export.service import StampingPartExportService


def test_process_confirm_export_headers_include_short_code_and_swing_shear() -> None:
    assert COL_COUNT == 28
    assert len(WIDTHS) == COL_COUNT
    workbook = Workbook()
    sheet = workbook.active
    StampingPartExportService()._write_headers(sheet)
    assert sheet.cell(4, 2).value == "零件号"
    assert sheet.cell(4, 3).value == "短码"
    assert sheet.cell(4, 4).value == "零件名称"
    assert sheet.cell(4, 16).value == "是否利废"
    assert sheet.cell(4, 17).value == "利用哪类废料"
    assert sheet.cell(4, 18).value == "是否摆剪"
    assert sheet.cell(4, 19).value == "工艺"
    assert sheet.cell(4, 28).value == "费用比例"


def test_process_confirm_export_writes_last_four_digit_short_code() -> None:
    workbook = Workbook()
    sheet = workbook.active
    part = SimpleNamespace(
        part_number="552000656AA",
        part_name="后背门左铰链加强板",
        stamping_supplier="烟台海众汽车零部件有限公司",
        vehicle_model="T28",
        version="V3",
        location="青岛",
        material="HC340/590DPD+Z",
        part_weight=None,
        parts_per_sheet=4,
        sheet_type="条料",
        is_offal_nested=False,
        offal_source=None,
        has_swing_shear=False,
        die_type="铸造模",
        process_method="双槽",
        automation="手工线",
        remark=None,
        thickness=None,
        sheet_length=None,
        sheet_width=None,
    )
    process = SimpleNamespace(
        process_seq="OP05",
        process_code="BL",
        die_sets=1,
        die_number=None,
        tonnage=400,
        cost_ratio=None,
    )
    StampingPartExportService()._write_part_block(
        sheet,
        start_row=DATA_START_ROW,
        seq_no=1,
        part=part,
        processes=[process],
        blanks=[],
        zebra=False,
        sketch_buffers=[],
    )
    assert sheet.cell(DATA_START_ROW, SHORT_CODE_COL).value == "0656"


def _steel_part(**overrides: object) -> SimpleNamespace:
    fields = {
        "seq_no": 12,
        "location": "青岛",
        "vehicle_model": "T28",
        "version": "V3",
        "stamping_supplier": "烟台海众汽车零部件有限公司",
        "part_number": "552000656AA",
        "part_name": "后背门左铰链加强板",
        "material": "HC340/590DPD+Z",
        "thickness": Decimal("0.70"),
        "sheet_length": Decimal("1000"),
        "sheet_width": Decimal("400"),
        "sheet_type": "条料",
        "parts_per_sheet": 2,
        "part_weight": Decimal("1.2000"),
        "is_offal_nested": False,
        "offal_source": None,
        "is_scrap_byproduct": False,
        "has_swing_shear": False,
        "die_type": "铸造模",
        "process_method": "双槽",
        "automation": "手工线",
        "remark": None,
        "is_tailor_welded": False,
        "processes": [],
        "tailor_weld_blanks": [],
    }
    fields.update(overrides)
    return SimpleNamespace(**fields)


def test_steel_purchase_headers_have_short_code_and_qty() -> None:
    assert len(STEEL_WIDTHS) == STEEL_COL_COUNT
    workbook = Workbook()
    sheet = workbook.active
    StampingPartExportService()._write_steel_purchase_headers(sheet)
    assert sheet.cell(4, 1).value == "范围"
    assert sheet.cell(4, 6).value == "零件"
    assert sheet.cell(4, 21).value == "测算"
    assert sheet.cell(5, 6).value == "零件号"
    assert sheet.cell(5, STEEL_SHORT_CODE_COL).value == "短码"
    assert sheet.cell(5, STEEL_QTY_COL).value == "主机厂数量"
    assert sheet.cell(5, STEEL_TON_COL).value == "采购钢材(吨)"
    for col in STEEL_MUTE_COLS:
        assert sheet.cell(5, col).font.color.rgb[-6:] == "6B6B6B"


def test_steel_purchase_row_writes_short_code_and_formulas() -> None:
    workbook = Workbook()
    sheet = workbook.active
    process = SimpleNamespace(
        process_order=1,
        id=1,
        process_code="双槽",
        die_number="M-1",
        tonnage=400,
        cost_ratio=Decimal("0.50"),
    )
    part = _steel_part(processes=[process])
    StampingPartExportService()._write_steel_purchase_row(
        sheet,
        excel_row=STEEL_DATA_START_ROW,
        part=part,
        processes=[process],
        blank=None,
        zebra=False,
    )
    assert sheet.cell(STEEL_DATA_START_ROW, STEEL_SHORT_CODE_COL).value == "0656"
    assert sheet.cell(STEEL_DATA_START_ROW, 1).font.color.rgb[-6:] == "8A8A8A"
    assert sheet.cell(STEEL_DATA_START_ROW, 6).font.color.rgb[-6:] == "000000"
    assert sheet.cell(STEEL_DATA_START_ROW, 19).value == "是"
    assert "双槽" in (sheet.cell(STEEL_DATA_START_ROW, 20).value or "")
    weight_formula = sheet.cell(STEEL_DATA_START_ROW, STEEL_WEIGHT_COL).value
    ton_formula = sheet.cell(STEEL_DATA_START_ROW, STEEL_TON_COL).value
    assert isinstance(weight_formula, str) and "SteelDensity" in weight_formula
    assert isinstance(ton_formula, str) and 'Q6="是"' in ton_formula
    assert sheet.cell(STEEL_DATA_START_ROW, STEEL_QTY_COL).value is None


def test_steel_purchase_offal_ton_formula_is_zero_branch() -> None:
    workbook = Workbook()
    sheet = workbook.active
    part = _steel_part(is_offal_nested=True, process_method="单干")
    StampingPartExportService()._write_steel_purchase_row(
        sheet,
        excel_row=STEEL_DATA_START_ROW,
        part=part,
        processes=[],
        blank=None,
        zebra=False,
    )
    assert sheet.cell(STEEL_DATA_START_ROW, 17).value == "是"
    assert sheet.cell(STEEL_DATA_START_ROW, STEEL_TON_COL).value.startswith(
        '=IF(OR(Q6="是",R6="是"),0,'
    )


def test_steel_purchase_workbook_freezes_after_short_code() -> None:
    service = StampingPartExportService()
    part = _steel_part()
    workbook = service._build_steel_purchase_workbook(
        [part],
        density=Decimal("7.8500"),
        location="青岛",
        stamping_supplier=None,
        vehicle_model="T28",
        version="V3",
    )
    sheet = workbook[workbook.sheetnames[0]]
    assert workbook.sheetnames[1] == "填写说明"
    assert sheet.freeze_panes == "H6"
    assert sheet["U2"].value == Decimal("7.8500")
    assert "SteelDensity" in workbook.defined_names


def test_steel_purchase_plain_bytes_have_no_watermark_images() -> None:
    service = StampingPartExportService()
    workbook = service._build_steel_purchase_workbook(
        [_steel_part()],
        density=Decimal("7.8500"),
        location="青岛",
        stamping_supplier=None,
        vehicle_model="T28",
        version="V3",
    )
    payload = service._workbook_to_plain_bytes(workbook)
    loaded = load_workbook(BytesIO(payload))
    for sheet in loaded.worksheets:
        assert not sheet._images
