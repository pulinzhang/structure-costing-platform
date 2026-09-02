"""工序吨位列纵向合并时，后续工序应读到左上角吨位。"""

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.parsers.stamping.stamping_parser import parse_stamping_excel
from app.services.stamping.upload.template_builder import build_stamping_process_template

_TONNAGE_COL = 18


def _write_process_workbook(path: Path) -> None:
    content, _filename = build_stamping_process_template()
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["冲压工艺确认表"]

    sheet.merge_cells(
        start_row=7,
        start_column=_TONNAGE_COL,
        end_row=8,
        end_column=_TONNAGE_COL,
    )
    sheet.cell(row=12, column=_TONNAGE_COL).value = None

    workbook.save(path)
    workbook.close()


def test_merged_tonnage_fills_later_processes(tmp_path: Path) -> None:
    excel_path = tmp_path / "process.xlsx"
    _write_process_workbook(excel_path)

    result = parse_stamping_excel(str(excel_path))
    by_part: dict[str, list[int | None]] = {}
    for process in result.processes:
        by_part.setdefault(process.part_number, []).append(process.tonnage)

    assert by_part["501006875AA"] == [400, 400, 630]
    assert by_part["501008899AA"] == [1000, 800, None]
