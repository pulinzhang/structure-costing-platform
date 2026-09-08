"""项目主数据：门禁、跨基地同代号、上传整单、基地引用保护、Parser 隔离。"""

from __future__ import annotations

import inspect
from collections.abc import Generator
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, event
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.audit import SYSTEM_USER
from app.database import Base, DatabaseConfig
from app.exceptions.base import BusinessException
from app.models.masterdata.location import Location, LocationAlias
from app.models.masterdata.vehicle_model import VehicleModel
from app.models.stamping.gauge_fee import StampingGaugeFee
from app.parsers.masterdata.vehicle_model_parser import parse_vehicle_model_excel
from app.schemas.dto.location_dto import LocationUpdateDTO
from app.schemas.dto.vehicle_model_dto import VehicleModelCreateDTO, VehicleModelUpdateDTO
from app.services.masterdata.location import LocationService
from app.services.masterdata.vehicle_model import VehicleModelService


@pytest.fixture
def memory_db() -> Generator[Session, None, None]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _enable_fk(dbapi_connection: object, _connection_record: object) -> None:
        cursor = dbapi_connection.cursor()  # type: ignore[attr-defined]
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    Base.metadata.create_all(
        bind=engine,
        tables=[
            Location.__table__,
            LocationAlias.__table__,
            VehicleModel.__table__,
            StampingGaugeFee.__table__,
        ],
    )
    session_factory = sessionmaker(bind=engine)
    session = session_factory()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def _seed(db: Session) -> None:
    LocationService().ensure_seed_if_empty(db)


def _create(
    db: Session,
    *,
    location: str = "青岛",
    vehicle_model: str = "T13J",
    effective_year: int = 2020,
    effective_month: int = 1,
    expiry_year: int | None = None,
    expiry_month: int | None = None,
) -> None:
    VehicleModelService().create(
        db,
        VehicleModelCreateDTO(
            location=location,
            vehicle_model=vehicle_model,
            effective_year=effective_year,
            effective_month=effective_month,
            expiry_year=expiry_year,
            expiry_month=expiry_month,
        ),
        current_user=SYSTEM_USER,
    )


def _workbook_bytes(rows: list[tuple[object, ...]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    headers = ["基地", "项目代号", "生效年", "生效月", "失效年", "失效月", "当前供货台数"]
    for col, title in enumerate(headers, start=1):
        sheet.cell(row=1, column=col, value=title)
    for row_idx, values in enumerate(rows, start=2):
        for col, value in enumerate(values, start=1):
            sheet.cell(row=row_idx, column=col, value=value)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def test_create_uppercases_and_allows_same_code_on_other_location(memory_db: Session) -> None:
    _seed(memory_db)
    service = VehicleModelService()
    vo = service.create(
        memory_db,
        VehicleModelCreateDTO(
            location="青岛",
            vehicle_model="t13j",
            effective_year=2026,
            effective_month=1,
        ),
        current_user=SYSTEM_USER,
    )
    assert vo.vehicle_model == "T13J"
    assert vo.location == "青岛"
    assert vo.current_supplied_units == 0
    other = service.create(
        memory_db,
        VehicleModelCreateDTO(
            location="大连",
            vehicle_model="T13J",
            effective_year=2026,
            effective_month=1,
        ),
        current_user=SYSTEM_USER,
    )
    assert other.location == "大连"
    with pytest.raises(BusinessException, match="该基地已有此项目"):
        service.create(
            memory_db,
            VehicleModelCreateDTO(
                location="青岛",
                vehicle_model="T13J",
                effective_year=2027,
                effective_month=1,
            ),
            current_user=SYSTEM_USER,
        )


def test_create_rejects_placeholder_and_disabled_location(memory_db: Session) -> None:
    _seed(memory_db)
    service = VehicleModelService()
    with pytest.raises(BusinessException, match="项目代号无效"):
        service.create(
            memory_db,
            VehicleModelCreateDTO(
                location="青岛",
                vehicle_model="UNKNOWN",
                effective_year=2026,
                effective_month=1,
            ),
            current_user=SYSTEM_USER,
        )
    LocationService().update(
        memory_db,
        "青岛",
        LocationUpdateDTO(is_enabled=False),
        current_user=SYSTEM_USER,
    )
    with pytest.raises(BusinessException, match="基地已禁用"):
        service.create(
            memory_db,
            VehicleModelCreateDTO(
                location="青岛",
                vehicle_model="T28",
                effective_year=2026,
                effective_month=1,
            ),
            current_user=SYSTEM_USER,
        )


def test_require_gates_and_list_page_zero_project_cards(memory_db: Session) -> None:
    _seed(memory_db)
    _create(memory_db, vehicle_model="T13J", effective_year=2026, effective_month=6)
    service = VehicleModelService()
    assert service.require_registered(memory_db, "青岛", "t13j") == "T13J"
    assert service.require_introducible(memory_db, "青岛", "T13J", 2026, 6) == ("青岛", "T13J")
    with pytest.raises(BusinessException, match="尚未生效"):
        service.require_introducible(memory_db, "青岛", "T13J", 2026, 5)
    assert service.require_overlaps_year(memory_db, "青岛", "T13J", 2026) == ("青岛", "T13J")
    with pytest.raises(BusinessException, match="无交集"):
        service.require_overlaps_year(memory_db, "青岛", "T13J", 2025)
    assert service.list_introducible(memory_db, "青岛", 2026, 6) == ["T13J"]
    assert service.list_introducible(memory_db, "青岛", 2026, 5) == []
    assert service.list_overlaps_year(memory_db, "青岛", 2026) == ["T13J"]
    assert service.list_overlaps_year(memory_db, "青岛", 2025) == []
    assert service.list_registered(memory_db, "青岛") == ["T13J"]
    assert service.list_filter_codes(memory_db, "青岛") == ["T13J"]
    assert service.list_filter_codes(memory_db) == ["T13J"]

    page = service.list_page(memory_db)
    names = [card.location for card in page.locations]
    assert names == ["大连", "芜湖", "鄂尔多斯", "青岛"]
    qingdao = next(card for card in page.locations if card.location == "青岛")
    dalian = next(card for card in page.locations if card.location == "大连")
    assert qingdao.project_count == 1
    assert dalian.project_count == 0
    assert dalian.introducible_count == 0


def test_list_filter_codes_unions_across_locations(memory_db: Session) -> None:
    _seed(memory_db)
    _create(memory_db, location="青岛", vehicle_model="T13J")
    _create(memory_db, location="大连", vehicle_model="T13J")
    _create(memory_db, location="青岛", vehicle_model="T28")
    service = VehicleModelService()
    assert service.list_filter_codes(memory_db, "青岛") == ["T13J", "T28"]
    assert service.list_filter_codes(memory_db, "大连") == ["T13J"]
    assert service.list_filter_codes(memory_db) == ["T13J", "T28"]


def test_require_introducible_by_code_unique(memory_db: Session) -> None:
    _seed(memory_db)
    _create(memory_db, location="青岛", vehicle_model="T13J", effective_year=2020)
    service = VehicleModelService()
    assert service.require_introducible_by_code(memory_db, "t13j", 2026, 6) == ("青岛", "T13J")
    with pytest.raises(BusinessException, match="没有该项目"):
        service.require_introducible_by_code(memory_db, "T99", 2026, 6)
    with pytest.raises(BusinessException, match="尚未生效"):
        service.require_introducible_by_code(memory_db, "T13J", 2019, 1)


def test_require_introducible_by_code_ambiguous(memory_db: Session) -> None:
    _seed(memory_db)
    _create(memory_db, location="青岛", vehicle_model="T13J")
    _create(memory_db, location="大连", vehicle_model="T13J")
    with pytest.raises(BusinessException, match="多个可引入基地"):
        VehicleModelService().require_introducible_by_code(memory_db, "T13J", 2026, 6)


def test_require_introducible_by_code_skips_disabled_location(memory_db: Session) -> None:
    _seed(memory_db)
    _create(memory_db, location="青岛", vehicle_model="T13J")
    _create(memory_db, location="大连", vehicle_model="T13J")
    LocationService().update(
        memory_db,
        "大连",
        LocationUpdateDTO(is_enabled=False),
        current_user=SYSTEM_USER,
    )
    assert VehicleModelService().require_introducible_by_code(memory_db, "T13J", 2026, 6) == (
        "青岛",
        "T13J",
    )


def test_disabled_location_card_introducible_is_zero(memory_db: Session) -> None:
    _seed(memory_db)
    _create(memory_db, effective_year=2000, effective_month=1)
    LocationService().update(
        memory_db,
        "青岛",
        LocationUpdateDTO(is_enabled=False),
        current_user=SYSTEM_USER,
    )
    page = VehicleModelService().list_page(memory_db)
    qingdao = next(card for card in page.locations if card.location == "青岛")
    assert qingdao.is_enabled is False
    assert qingdao.project_count == 1
    assert qingdao.introducible_count == 0
    item = next(row for row in page.items if row.location == "青岛")
    assert item.period_status == "可引入"


def test_upload_skip_existing_and_in_file_conflict(memory_db: Session) -> None:
    _seed(memory_db)
    service = VehicleModelService()
    first = service.process_upload(
        memory_db,
        _workbook_bytes([("青岛", "T13J", 2026, 1, None, None)]),
        current_user=SYSTEM_USER,
    )
    assert first.created_count == 1
    assert first.skipped_count == 0
    second = service.process_upload(
        memory_db,
        _workbook_bytes(
            [
                ("青岛", "T13J", 2026, 1, None, None),
                ("青岛", "t13j", 2026, 1, None, None),
            ]
        ),
        current_user=SYSTEM_USER,
    )
    assert second.created_count == 0
    assert second.skipped_count == 2
    with pytest.raises(BusinessException, match="不一致"):
        service.process_upload(
            memory_db,
            _workbook_bytes(
                [
                    ("青岛", "T28", 2026, 1, None, None),
                    ("青岛", "T28", 2027, 1, None, None),
                ]
            ),
            current_user=SYSTEM_USER,
        )
    assert service.list_registered(memory_db, "青岛") == ["T13J"]


def test_upload_and_create_current_supplied_units(memory_db: Session) -> None:
    _seed(memory_db)
    service = VehicleModelService()
    created = service.process_upload(
        memory_db,
        _workbook_bytes([("青岛", "T13J", 2026, 1, None, None, 8500)]),
        current_user=SYSTEM_USER,
    )
    assert created.created_count == 1
    page = service.list_page(memory_db)
    row = next(item for item in page.items if item.vehicle_model == "T13J")
    assert row.current_supplied_units == 8500
    with pytest.raises(BusinessException, match="当前供货台数"):
        service.process_upload(
            memory_db,
            _workbook_bytes([("大连", "T28", 2026, 1, None, None, -1)]),
            current_user=SYSTEM_USER,
        )
    with pytest.raises(BusinessException, match="不一致"):
        service.process_upload(
            memory_db,
            _workbook_bytes(
                [
                    ("大连", "T28", 2026, 1, None, None, 100),
                    ("大连", "T28", 2026, 1, None, None, 200),
                ]
            ),
            current_user=SYSTEM_USER,
        )
    vo = service.create(
        memory_db,
        VehicleModelCreateDTO(
            location="大连",
            vehicle_model="T99",
            effective_year=2026,
            effective_month=1,
            current_supplied_units=3,
        ),
        current_user=SYSTEM_USER,
    )
    assert vo.current_supplied_units == 3


def test_upload_alias_location_and_disabled_fails_whole_file(memory_db: Session) -> None:
    _seed(memory_db)
    LocationService().update(
        memory_db,
        "青岛",
        LocationUpdateDTO(is_enabled=False),
        current_user=SYSTEM_USER,
    )
    with pytest.raises(BusinessException, match="基地已禁用"):
        VehicleModelService().process_upload(
            memory_db,
            _workbook_bytes(
                [
                    ("大连", "T28", 2026, 1, None, None),
                    ("青岛基地", "T13J", 2026, 1, None, None),
                ]
            ),
            current_user=SYSTEM_USER,
        )
    assert VehicleModelService().list_registered(memory_db, "大连") == []


def test_update_period_and_delete_integrity_message(memory_db: Session) -> None:
    _seed(memory_db)
    _create(memory_db)
    service = VehicleModelService()
    updated = service.update(
        memory_db,
        "青岛",
        "T13J",
        VehicleModelUpdateDTO(
            effective_year=2026,
            effective_month=3,
            expiry_year=2026,
            expiry_month=12,
            current_supplied_units=12000,
        ),
        current_user=SYSTEM_USER,
    )
    assert updated.effective_month == 3
    assert updated.expiry_month == 12
    assert updated.current_supplied_units == 12000

    def _raise_integrity(*_args: object, **_kwargs: object) -> None:
        raise IntegrityError("DELETE", {}, Exception("fk"))

    service.repository.delete = MagicMock(side_effect=_raise_integrity)  # type: ignore[method-assign]
    with pytest.raises(BusinessException, match="仍被引用，不能删除"):
        service.delete(memory_db, "青岛", "T13J")


def test_delete_blocked_when_child_row_exists(memory_db: Session) -> None:
    _seed(memory_db)
    _create(memory_db)
    memory_db.add(
        StampingGaugeFee(
            location="青岛",
            vehicle_model="T13J",
            part_number="P1",
            gauge_fee=Decimal("0.8"),
        )
    )
    memory_db.commit()
    updated = VehicleModelService().update(
        memory_db,
        "青岛",
        "T13J",
        VehicleModelUpdateDTO(
            effective_year=2024,
            effective_month=1,
            expiry_year=None,
            expiry_month=None,
            current_supplied_units=0,
        ),
        current_user=SYSTEM_USER,
    )
    assert updated.effective_year == 2024
    with pytest.raises(BusinessException, match="仍被引用，不能删除"):
        VehicleModelService().delete(memory_db, "青岛", "T13J")


def test_sqlite_rebuild_table_preserves_child_fk(memory_db: Session) -> None:
    from sqlalchemy import text

    _seed(memory_db)
    _create(memory_db)
    memory_db.add(
        StampingGaugeFee(
            location="青岛",
            vehicle_model="T13J",
            part_number="P1",
            gauge_fee=Decimal("0.8"),
        )
    )
    memory_db.commit()
    engine = memory_db.get_bind()
    with engine.connect() as conn:
        autocommit = conn.execution_options(isolation_level="AUTOCOMMIT")
        autocommit.exec_driver_sql("PRAGMA foreign_keys=OFF")
        try:
            DatabaseConfig._sqlite_rebuild_table(autocommit, "stamping_gauge_fees")
        finally:
            autocommit.exec_driver_sql("PRAGMA foreign_keys=ON")
        assert DatabaseConfig._sqlite_has_vehicle_model_fk(autocommit, "stamping_gauge_fees")
        count = autocommit.execute(text("SELECT COUNT(*) FROM stamping_gauge_fees")).scalar()
        assert count == 1
    memory_db.expire_all()
    with pytest.raises(BusinessException, match="仍被引用，不能删除"):
        VehicleModelService().delete(memory_db, "青岛", "T13J")


def test_ensure_vehicle_model_child_fks_is_idempotent(memory_db: Session) -> None:
    engine = memory_db.get_bind()
    DatabaseConfig._ensure_vehicle_model_child_fks(engine)
    DatabaseConfig._ensure_vehicle_model_child_fks(engine)
    with engine.connect() as conn:
        assert DatabaseConfig._sqlite_has_vehicle_model_fk(conn, "stamping_gauge_fees")


def test_location_delete_blocked_when_project_exists(memory_db: Session) -> None:
    _seed(memory_db)
    _create(memory_db)
    with pytest.raises(BusinessException, match="仍被引用，不能删除"):
        LocationService().delete(memory_db, "青岛")


def test_assert_matches_selected(memory_db: Session) -> None:
    service = VehicleModelService()
    service.assert_matches_selected("青岛t13j", "T13J")
    with pytest.raises(BusinessException, match="不一致"):
        service.assert_matches_selected("T13J", "T28")


def test_export_matches_import_template_and_parses(memory_db: Session) -> None:
    _seed(memory_db)
    _create(
        memory_db,
        location="芜湖",
        vehicle_model="T13J",
        effective_year=2025,
        effective_month=1,
        expiry_year=2026,
        expiry_month=12,
    )
    _create(memory_db, location="青岛", vehicle_model="T28", effective_year=2025, effective_month=1)
    service = VehicleModelService()
    content, filename = service.export_excel(memory_db, current_user=SYSTEM_USER)
    assert filename.startswith("项目一览_")
    assert filename.endswith(".xlsx")
    workbook = load_workbook(BytesIO(content))
    headers = [cell.value for cell in workbook.worksheets[0][1]]
    workbook.close()
    assert headers[:7] == [
        "基地",
        "项目代号",
        "生效年",
        "生效月",
        "失效年",
        "失效月",
        "当前供货台数",
    ]
    parsed = parse_vehicle_model_excel(content)
    assert parsed.errors == []
    assert [
        (
            row.location,
            row.vehicle_model,
            row.effective_year,
            row.effective_month,
            row.expiry_year,
            row.expiry_month,
            row.current_supplied_units,
        )
        for row in parsed.rows
    ] == [
        ("芜湖", "T13J", 2025, 1, 2026, 12, 0),
        ("青岛", "T28", 2025, 1, None, None, 0),
    ]


def test_export_filters_by_location(memory_db: Session) -> None:
    _seed(memory_db)
    _create(memory_db, location="芜湖", vehicle_model="T13J")
    _create(memory_db, location="青岛", vehicle_model="T28")
    service = VehicleModelService()
    content, _filename = service.export_excel(
        memory_db,
        current_user=SYSTEM_USER,
        location="青岛",
    )
    parsed = parse_vehicle_model_excel(content)
    assert [(row.location, row.vehicle_model) for row in parsed.rows] == [("青岛", "T28")]
    with pytest.raises(BusinessException, match="没有这座基地"):
        service.export_excel(memory_db, current_user=SYSTEM_USER, location="不存在")


def test_parser_does_not_import_vehicle_model_service() -> None:
    from app.parsers.masterdata import vehicle_model_parser

    source = inspect.getsource(vehicle_model_parser)
    assert "VehicleModelService" not in source
    assert "LocationService" not in source
    assert "app.services" not in source
    assert "app.repositories" not in source
