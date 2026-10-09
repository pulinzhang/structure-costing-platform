"""材料变更模板按所选基地、项目和生效年月写举例，并能被解析。"""

from io import BytesIO

from openpyxl import load_workbook

from app.parsers.engineering_change.material_change.codes import HEADERS, SHEET_HELP, SHEET_ORDER
from app.parsers.engineering_change.material_change.order_parser import (
    parse_material_change_workbook,
)
from app.services.engineering_change.material_change.upload.template import (
    PLACEHOLDER_SUPPLIER,
    MaterialChangeTemplateMeta,
    build_material_change_template,
)


def test_template_examples_follow_selected_project() -> None:
    meta = MaterialChangeTemplateMeta(
        location="大连",
        vehicle_model="DEMO",
        change_version="MAT-202610",
        period_label="2026年10月",
    )
    content, filename = build_material_change_template(meta)
    assert filename == "材料变更导入模板_大连_DEMO.xlsx"

    workbook = load_workbook(BytesIO(content), data_only=False)
    assert workbook.sheetnames == [SHEET_HELP, SHEET_ORDER]
    order = workbook[SHEET_ORDER]
    assert [order.cell(1, col).value for col in range(1, 13)] == list(HEADERS)
    assert order.cell(2, 1).value == "大连"
    assert order.cell(2, 2).value == "DEMO"
    assert order.cell(2, 3).value == "MAT-202610"
    assert order.cell(2, 4).value == "2026年10月"
    assert order.cell(2, 5).value == PLACEHOLDER_SUPPLIER
    assert order.cell(2, 6).value == "EX-MAT-AA"
    assert order.cell(3, 6).value == "EX-MAT-AB"
    assert order.cell(2, 9).value == 1.5
    assert order.cell(2, 11).value == 1.5
    assert order.cell(2, 12).value == "不影响互换性"
    assert order.cell(4, 1).value is None
    visible = " ".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value
    )
    workbook.close()
    assert "芜湖" not in visible
    assert "T13J" not in visible
    assert "振华" not in visible

    rows = parse_material_change_workbook(content)
    assert len(rows) == 2
    assert rows[0].location == "大连"
    assert rows[0].vehicle_model == "DEMO"
    assert rows[0].change_version == "MAT-202610"
    assert rows[0].effective_year == 2026
    assert rows[0].effective_month == 10
    assert rows[0].part_number == "EX-MAT-AA"
    assert rows[1].part_number == "EX-MAT-AB"
    assert rows[0].before_thickness == "1.5"
    assert rows[1].before_thickness == "1.2"
