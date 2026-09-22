"""工程变更三条门禁与变更单基础校验。"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.common.engineering_change_gates import (
    CHANGE_TYPE_ADD_PART,
    CHANGE_TYPE_NOT_INSTALL,
    CHANGE_TYPE_REVISE,
    CHANGE_TYPE_STOP,
    EC_ADD_ALREADY_ON_TREE,
    EC_AFTER_ASSEMBLY_EXISTS,
    EC_BEFORE_ASSEMBLY_INVALID,
    EC_CHANGE_ORDER_ASSEMBLY_REPEAT,
    EC_CHANGE_ORDER_ROW_COUNT,
    EC_EFFECTIVE_MISMATCH,
    EC_GATE_ENGINE_BREAKDOWN,
    EC_GATE_ENGINE_PROCUREMENT,
    EC_GATE_FULL_GENERATE,
    EC_PART_NOT_ON_TREE,
    EC_PROJECT_MISMATCH,
    PeriodRoot,
    is_sop_alignable_root,
    snapshot_blocks_full_generate,
)
from app.exceptions.base import BusinessException
from app.parsers.engineering_change.design_change.change_order_parser import (
    parse_change_order_sheet,
)
from app.services.engineering_change.design_change.upload.template_builder import (
    CHANGE_ORDER_SAMPLES,
    PROCESS_SAMPLES,
    SCENARIO_SAMPLES,
    TEMPLATE_FILENAME,
    build_engineering_change_template,
)


def _root(
    *,
    effective_year: int = 2025,
    effective_month: int = 1,
    expiry_year: int | None = None,
    expiry_month: int | None = None,
) -> PeriodRoot:
    return PeriodRoot(
        effective_year=effective_year,
        effective_month=effective_month,
        expiry_year=expiry_year,
        expiry_month=expiry_month,
    )


def test_sop_alignable_only_sop_month_active() -> None:
    sop = _root(effective_year=2025, effective_month=1)
    ec = _root(effective_year=2025, effective_month=7)
    expired = _root(
        effective_year=2025,
        effective_month=1,
        expiry_year=2025,
        expiry_month=6,
    )
    assert is_sop_alignable_root(sop, sop_year=2025, sop_month=1, as_of_year=2025, as_of_month=8)
    assert not is_sop_alignable_root(ec, sop_year=2025, sop_month=1, as_of_year=2025, as_of_month=8)
    assert not is_sop_alignable_root(
        expired, sop_year=2025, sop_month=1, as_of_year=2025, as_of_month=8
    )


def test_snapshot_blocks_full_generate_on_ec_or_expired() -> None:
    sop = _root(effective_year=2025, effective_month=1)
    assert not snapshot_blocks_full_generate(
        [sop], sop_year=2025, sop_month=1, as_of_year=2025, as_of_month=3
    )
    ec = _root(effective_year=2025, effective_month=7)
    assert snapshot_blocks_full_generate(
        [sop, ec], sop_year=2025, sop_month=1, as_of_year=2025, as_of_month=8
    )
    expired = _root(
        effective_year=2025,
        effective_month=1,
        expiry_year=2025,
        expiry_month=6,
    )
    assert snapshot_blocks_full_generate(
        [expired], sop_year=2025, sop_month=1, as_of_year=2025, as_of_month=8
    )


def test_gate_message_constants() -> None:
    assert "工程变更" in EC_GATE_FULL_GENERATE
    assert "拆解" in EC_GATE_ENGINE_BREAKDOWN
    assert "采购对象" in EC_GATE_ENGINE_PROCUREMENT


def test_template_change_order_examples_and_parses_as_book() -> None:
    content, filename = build_engineering_change_template()
    assert filename == TEMPLATE_FILENAME
    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook.sheetnames == [
        "工程变更单",
        "工程变更填写说明",
        "车身件材料清单",
        "材料清单填写说明",
        "冲压工艺确认表",
        "工艺确认填写说明",
    ]
    instructions = workbook["工程变更填写说明"]
    instruction_text = " ".join(str(instructions.cell(row, 1).value or "") for row in range(1, 10))
    assert "项目代号" in instruction_text
    assert "同一个字" in instruction_text
    assert "采购对象" in instruction_text
    change_order = workbook["工程变更单"]
    types = [change_order.cell(row, 6).value for row in range(5, 5 + len(CHANGE_ORDER_SAMPLES))]
    assert types == [sample.change_type for sample in CHANGE_ORDER_SAMPLES]
    assert change_order.cell(4 + len(CHANGE_ORDER_SAMPLES), 6).value == CHANGE_TYPE_ADD_PART
    process = workbook["冲压工艺确认表"]
    process_numbers: list[str] = []
    excel_row = 7
    for sample in PROCESS_SAMPLES:
        if sample.part_number in process_numbers:
            continue
        assert process.cell(excel_row, 2).value == sample.part_number
        process_numbers.append(sample.part_number)
        excel_row += sum(1 for item in PROCESS_SAMPLES if item.part_number == sample.part_number)
    assert process_numbers == list(dict.fromkeys(sample.part_number for sample in PROCESS_SAMPLES))
    merged = {str(item) for item in change_order.merged_cells.ranges}
    workbook.close()
    rows = parse_change_order_sheet(content)
    assert len(rows) == len(CHANGE_ORDER_SAMPLES)
    assert [row.change_type for row in rows] == [
        sample.change_type for sample in CHANGE_ORDER_SAMPLES
    ]
    last_row = 4 + len(CHANGE_ORDER_SAMPLES)
    assert merged >= {
        f"A5:A{last_row}",
        f"B5:B{last_row}",
    }
    assert "C5:C9" not in merged
    assert [row.before_assembly_part_number for row in rows] == [
        sample.before_assembly for sample in CHANGE_ORDER_SAMPLES
    ]
    assert len({row.before_assembly_part_number for row in rows}) == len(rows)
    after_numbers = [
        row.after_assembly_part_number for row in rows if row.after_assembly_part_number
    ]
    assert len(after_numbers) == len(set(after_numbers))


def test_template_passes_import_change_order_and_bom_l0_gates() -> None:
    from app.parsers.bom.bom_parser import parse_bom_excel
    from app.parsers.engineering_change.design_change.sheet_extract import SHEET_BOM
    from app.services.engineering_change.design_change.upload.gates.validate import (
        validate_bom_l0_matches_change_orders,
        validate_change_order_rows,
    )

    content, _ = build_engineering_change_template()
    rows = parse_change_order_sheet(content)
    validate_change_order_rows(
        rows,
        page_vehicle_model="T28",
        page_effective_year=2025,
        page_effective_month=7,
    )
    parsed = parse_bom_excel(content, sheet_name=SHEET_BOM)
    validate_bom_l0_matches_change_orders(parsed.nodes, rows)


def test_change_order_accepts_written_ac_not_inferred_ab() -> None:
    from app.parsers.engineering_change.design_change.change_order_parser import (
        ChangeOrderRowParsed,
    )
    from app.services.engineering_change.design_change.upload.gates.validate import (
        validate_change_order_row,
    )

    change_order = ChangeOrderRowParsed(
        project_code="T28",
        effective_period_text="2025年07月",
        change_type=CHANGE_TYPE_REVISE,
        before_assembly_part_number="501006873AA",
        after_assembly_part_number="501006873AC",
        assembly_name="流水槽加强梁总成",
        part_level="一级",
        before_part_number="501006875AA",
        after_part_number="501006875AC",
        part_name="雨刮电机安装支架Ⅰ",
        excel_row=5,
    )
    year, month = validate_change_order_row(
        change_order,
        page_vehicle_model="T28",
        page_effective_year=2025,
        page_effective_month=7,
    )
    assert (year, month) == (2025, 7)


def test_change_order_project_mismatch_is_first_gate() -> None:
    from app.parsers.engineering_change.design_change.change_order_parser import (
        ChangeOrderRowParsed,
    )
    from app.services.engineering_change.design_change.upload.gates.validate import (
        validate_change_order_project,
    )

    change_order = ChangeOrderRowParsed(
        project_code="T13J",
        effective_period_text="2025年07月",
        change_type=CHANGE_TYPE_REVISE,
        before_assembly_part_number="501006873AA",
        after_assembly_part_number="501006873AB",
        assembly_name="流水槽加强梁总成",
        part_level="总成",
        before_part_number=None,
        after_part_number=None,
        part_name=None,
        excel_row=5,
    )
    with pytest.raises(BusinessException) as exc:
        validate_change_order_project(
            [change_order],
            page_vehicle_model="T28",
            page_effective_year=2025,
            page_effective_month=7,
        )
    assert str(exc.value).startswith(EC_PROJECT_MISMATCH)


def test_change_order_one_upload_one_period() -> None:
    from app.parsers.engineering_change.design_change.change_order_parser import (
        ChangeOrderRowParsed,
    )
    from app.services.engineering_change.design_change.upload.gates.validate import (
        validate_change_order_project,
    )

    first = ChangeOrderRowParsed(
        project_code="T28",
        effective_period_text="2025年07月",
        change_type=CHANGE_TYPE_REVISE,
        before_assembly_part_number="501006873AA",
        after_assembly_part_number="501006873AB",
        assembly_name="流水槽加强梁总成",
        part_level="总成",
        before_part_number=None,
        after_part_number=None,
        part_name=None,
        excel_row=5,
    )
    second = ChangeOrderRowParsed(
        project_code="T28",
        effective_period_text="2025年08月",
        change_type=CHANGE_TYPE_REVISE,
        before_assembly_part_number="501008897AA",
        after_assembly_part_number="501008897AB",
        assembly_name="后围总成",
        part_level="总成",
        before_part_number=None,
        after_part_number=None,
        part_name=None,
        excel_row=6,
    )
    with pytest.raises(BusinessException) as exc:
        validate_change_order_project(
            [first, second],
            page_vehicle_model="T28",
            page_effective_year=2025,
            page_effective_month=7,
        )
    assert str(exc.value).startswith(EC_EFFECTIVE_MISMATCH)
    assert "一次上传只能有一个变更年月" in str(exc.value)
    assert "2025年07月" in str(exc.value)
    assert "2025年08月" in str(exc.value)


def _change_order_xlsx(
    rows: list[tuple[object, ...]],
    merges: list[str] | None = None,
) -> bytes:
    from openpyxl import Workbook

    from app.parsers.engineering_change.design_change.change_order_parser import (
        CHANGE_ORDER_HEADERS,
    )

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "工程变更单"
    for col, header in enumerate(CHANGE_ORDER_HEADERS, start=1):
        sheet.cell(1, col, header)
    for excel_row, values in enumerate(rows, start=2):
        for col, value in enumerate(values, start=1):
            sheet.cell(excel_row, col, value)
    for merge in merges or []:
        sheet.merge_cells(merge)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_same_assembly_two_scenarios_must_merge_cells() -> None:
    from app.services.engineering_change.design_change.upload.gates.validate import (
        validate_change_order_rows,
    )

    content = _change_order_xlsx(
        [
            (
                "T28",
                "2025年07月",
                "501006873AA",
                "501006873AB",
                "流水槽加强梁总成",
                CHANGE_TYPE_REVISE,
                "一级",
                "501006875AA",
                "501006875AB",
                "雨刮电机安装支架Ⅰ",
            ),
            (
                "T28",
                "2025年07月",
                "501006873AA",
                "501006873AB",
                "流水槽加强梁总成",
                CHANGE_TYPE_NOT_INSTALL,
                "一级",
                "504000341AA",
                None,
                "左胶块",
            ),
        ]
    )
    rows = parse_change_order_sheet(content)
    with pytest.raises(BusinessException) as exc:
        validate_change_order_rows(
            rows,
            page_vehicle_model="T28",
            page_effective_year=2025,
            page_effective_month=7,
        )
    assert str(exc.value).startswith(EC_CHANGE_ORDER_ASSEMBLY_REPEAT)
    assert "上下合并" in str(exc.value)
    assert "场景" not in str(exc.value)


def test_same_assembly_two_scenarios_merged_is_one_identity() -> None:
    from app.services.engineering_change.design_change.upload.gates.validate import (
        validate_change_order_rows,
    )

    content = _change_order_xlsx(
        [
            (
                "T28",
                "2025年07月",
                "501006873AA",
                "501006873AB",
                "流水槽加强梁总成",
                CHANGE_TYPE_REVISE,
                "一级",
                "501006875AA",
                "501006875AB",
                "雨刮电机安装支架Ⅰ",
            ),
            (
                "T28",
                "2025年07月",
                None,
                None,
                None,
                CHANGE_TYPE_NOT_INSTALL,
                "一级",
                "504000341AA",
                None,
                "左胶块",
            ),
        ],
        merges=["C2:C3", "D2:D3", "E2:E3"],
    )
    rows = parse_change_order_sheet(content)
    assert rows[0].before_assembly_span == (2, 3)
    assert rows[1].before_assembly_span == (2, 3)
    assert rows[0].after_assembly_span == (2, 3)
    assert rows[1].after_assembly_part_number == "501006873AB"
    validate_change_order_rows(
        rows,
        page_vehicle_model="T28",
        page_effective_year=2025,
        page_effective_month=7,
    )


def test_same_assembly_nonconsecutive_is_not_unique() -> None:
    from app.services.engineering_change.design_change.upload.gates.validate import (
        validate_change_order_rows,
    )

    content = _change_order_xlsx(
        [
            (
                "T28",
                "2025年07月",
                "501006873AA",
                "501006873AB",
                "流水槽加强梁总成",
                CHANGE_TYPE_REVISE,
                "一级",
                "501006875AA",
                "501006875AB",
                "雨刮电机安装支架Ⅰ",
            ),
            (
                "T28",
                "2025年07月",
                "501008897AA",
                "501008897AB",
                "后围总成",
                CHANGE_TYPE_NOT_INSTALL,
                "一级",
                "504000341AA",
                None,
                "左胶块",
            ),
            (
                "T28",
                "2025年07月",
                "501006873AA",
                "501006873AB",
                "流水槽加强梁总成",
                CHANGE_TYPE_ADD_PART,
                "一级",
                None,
                "501007040AA",
                "中通道后纵梁加强板",
            ),
        ]
    )
    rows = parse_change_order_sheet(content)
    with pytest.raises(BusinessException) as exc:
        validate_change_order_rows(
            rows,
            page_vehicle_model="T28",
            page_effective_year=2025,
            page_effective_month=7,
        )
    assert str(exc.value).startswith(EC_CHANGE_ORDER_ASSEMBLY_REPEAT)
    assert "不要隔开写" in str(exc.value)
    assert "必须唯一" not in str(exc.value)


def test_stop_after_assembly_excluded_from_uniqueness() -> None:
    from app.services.engineering_change.design_change.upload.gates.validate import (
        validate_change_order_rows,
    )

    content = _change_order_xlsx(
        [
            (
                "T28",
                "2025年07月",
                "501009533AA",
                None,
                "发盖左铰链安装支架总成",
                CHANGE_TYPE_STOP,
                "总成",
                None,
                None,
                None,
            ),
            (
                "T28",
                "2025年07月",
                "501009532AA",
                None,
                "发盖右铰链安装支架总成",
                CHANGE_TYPE_STOP,
                "总成",
                None,
                None,
                None,
            ),
        ]
    )
    rows = parse_change_order_sheet(content)
    validate_change_order_rows(
        rows,
        page_vehicle_model="T28",
        page_effective_year=2025,
        page_effective_month=7,
    )


def test_before_assembly_missing_rejects() -> None:
    from app.services.engineering_change.design_change.upload.gates.validate import (
        validate_before_assembly_active,
    )

    with pytest.raises(BusinessException) as exc:
        validate_before_assembly_active(
            None,
            effective_year=2025,
            effective_month=7,
            part_number="501006873AA",
        )
    assert str(exc.value).startswith(EC_BEFORE_ASSEMBLY_INVALID)


def test_after_assembly_already_exists_rejects() -> None:
    from types import SimpleNamespace

    from app.services.engineering_change.design_change.upload.gates.validate import (
        validate_after_assembly_absent,
    )

    with pytest.raises(BusinessException) as exc:
        validate_after_assembly_absent(
            SimpleNamespace(part_number="501006873AB"),
            part_number="501006873AB",
        )
    assert str(exc.value).startswith(EC_AFTER_ASSEMBLY_EXISTS)


def test_before_part_missing_and_after_part_already_on_tree() -> None:
    from types import SimpleNamespace

    from app.parsers.engineering_change.design_change.change_order_parser import (
        ChangeOrderRowParsed,
    )
    from app.services.engineering_change.design_change.upload.gates.validate import (
        validate_change_order_before_after,
    )

    change_order = ChangeOrderRowParsed(
        project_code="T28",
        effective_period_text="2025年07月",
        change_type=CHANGE_TYPE_REVISE,
        before_assembly_part_number="501006873AA",
        after_assembly_part_number="501006873AB",
        assembly_name="流水槽加强梁总成",
        part_level="一级",
        before_part_number="501006875AA",
        after_part_number="501006875AB",
        part_name=None,
        excel_row=5,
    )
    old_nodes = [
        SimpleNamespace(part_number="501006873AA", level=0),
    ]
    with pytest.raises(BusinessException) as exc:
        validate_change_order_before_after(
            change_order,
            old_nodes=old_nodes,
            after_assembly=None,
        )
    assert str(exc.value).startswith(EC_PART_NOT_ON_TREE)

    old_nodes.append(SimpleNamespace(part_number="501006875AA", level=1))
    old_nodes.append(SimpleNamespace(part_number="501006875AB", level=1))
    with pytest.raises(BusinessException) as exc:
        validate_change_order_before_after(
            change_order,
            old_nodes=old_nodes,
            after_assembly=None,
        )
    assert str(exc.value).startswith(EC_ADD_ALREADY_ON_TREE)


def test_template_bom_l0_has_assembly_sketch() -> None:
    from app.parsers.bom.bom_parser import parse_bom_excel
    from app.parsers.engineering_change.design_change.sheet_extract import SHEET_BOM

    content, _ = build_engineering_change_template()
    parsed = parse_bom_excel(content, sheet_name=SHEET_BOM)
    l0_groups: dict[str, list] = {}
    for node in parsed.nodes:
        if int(node.level) != 0:
            continue
        l0_groups.setdefault(node.part_number, []).append(node)
    assert l0_groups
    assert all(any(node.sketch_bytes for node in group) for group in l0_groups.values())
    assert all(
        any((node.internal_number or "").strip() for node in group) for group in l0_groups.values()
    )


def test_template_bom_l1_leaves_have_part_sketches() -> None:
    from app.common.engineering_change_gates import NOT_INSTALL_REMARK, is_not_install_row
    from app.parsers.bom.bom_parser import parse_bom_excel
    from app.parsers.engineering_change.design_change.sheet_extract import SHEET_BOM

    content, _ = build_engineering_change_template()
    parsed = parse_bom_excel(content, sheet_name=SHEET_BOM)
    parents = {node.parent_index for node in parsed.nodes if node.parent_index is not None}
    leaves = [
        node
        for index, node in enumerate(parsed.nodes)
        if int(node.level) == 1 and index not in parents
    ]
    assert leaves
    required = [
        node
        for node in leaves
        if not is_not_install_row(node.qty, node.remark, part_number=node.part_number)
    ]
    assert required
    assert all(node.sketch_bytes for node in required)
    skipped = [node for node in leaves if node.part_number == "504000341AA"]
    assert skipped
    assert all(not node.sketch_bytes for node in skipped)
    assert all(NOT_INSTALL_REMARK in (node.remark or "") for node in skipped)


def test_template_process_sheet_parses_new_part_numbers() -> None:
    from app.services.engineering_change.design_change.upload.gates.process_gates import (
        parse_process_confirmation,
    )

    content, _ = build_engineering_change_template()
    numbers = {
        (part.part_number or "").strip()
        for part in parse_process_confirmation(content).parts
        if (part.part_number or "").strip()
    }
    assert numbers == {sample.part_number for sample in PROCESS_SAMPLES}


def test_template_process_sheet_has_part_sketches() -> None:
    import tempfile
    from pathlib import Path

    from app.parsers.engineering_change.design_change.sheet_extract import (
        SHEET_PROCESS,
        extract_named_sheets,
    )
    from app.parsers.stamping.stamping_parser import parse_stamping_excel

    content, _ = build_engineering_change_template()
    sheets = extract_named_sheets(content)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(sheets[SHEET_PROCESS])
        tmp_path = tmp.name
    try:
        parsed = parse_stamping_excel(tmp_path)
    finally:
        Path(tmp_path).unlink(missing_ok=True)
    expected = {sample.part_number for sample in PROCESS_SAMPLES}
    assert set(parsed.images_by_part) == expected
    assert all(parsed.images_by_part[number] for number in expected)


def test_template_has_no_die_sheet() -> None:
    from app.parsers.engineering_change.design_change.sheet_extract import extract_named_sheets

    content, _ = build_engineering_change_template()
    sheets = extract_named_sheets(content)
    assert "零件汇总" not in sheets
    workbook = load_workbook(BytesIO(content), data_only=True)
    assert "零件汇总" not in workbook.sheetnames
    workbook.close()


@pytest.mark.parametrize(
    "scenario",
    SCENARIO_SAMPLES,
    ids=[sample.filename for sample in SCENARIO_SAMPLES],
)
def test_scenario_template_parses_and_passes_gates(scenario) -> None:
    from app.parsers.bom.bom_parser import parse_bom_excel
    from app.parsers.engineering_change.design_change.sheet_extract import SHEET_BOM
    from app.services.engineering_change.design_change.upload.gates.process_gates import (
        parse_process_confirmation,
    )
    from app.services.engineering_change.design_change.upload.gates.validate import (
        validate_bom_l0_matches_change_orders,
        validate_change_order_rows,
    )

    content, _ = build_engineering_change_template(scenario.book)
    workbook = load_workbook(BytesIO(content), data_only=True)
    change_order = workbook["工程变更单"]
    merged = {str(item) for item in change_order.merged_cells.ranges}
    workbook.close()
    assert "A5:A5" not in merged
    assert "B5:B5" not in merged
    rows = parse_change_order_sheet(content)
    assert len(rows) == 1
    sample = scenario.book.change_orders[0]
    assert rows[0].change_type == sample.change_type
    assert rows[0].before_assembly_part_number == sample.before_assembly
    assert rows[0].after_assembly_part_number == sample.after_assembly
    validate_change_order_rows(
        rows,
        page_vehicle_model="T28",
        page_effective_year=2025,
        page_effective_month=7,
    )
    parsed = parse_bom_excel(content, sheet_name=SHEET_BOM)
    validate_bom_l0_matches_change_orders(parsed.nodes, rows)
    numbers = {
        (part.part_number or "").strip()
        for part in parse_process_confirmation(content).parts
        if (part.part_number or "").strip()
    }
    assert numbers == {item.part_number for item in scenario.book.process}


def test_change_type_revise_constant() -> None:
    assert CHANGE_TYPE_REVISE == "改图"


def test_empty_change_order_rejected() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "工程变更单"
    headers = (
        "项目",
        "生效年月",
        "变更前总成",
        "变更后总成",
        "总成名称",
        "变更类型",
        "零件层级",
        "变更前零件",
        "变更后零件",
        "零件名称",
    )
    for col, header in enumerate(headers, start=1):
        sheet.cell(1, col, header)
    buffer = BytesIO()
    workbook.save(buffer)
    with pytest.raises(BusinessException) as exc:
        parse_change_order_sheet(buffer.getvalue())
    assert str(exc.value).startswith(EC_CHANGE_ORDER_ROW_COUNT)


def test_slice_bom_nodes_for_l0_remaps_parent_index() -> None:
    from app.parsers.bom.models import BomParsedNode
    from app.services.engineering_change.design_change.upload.gates.bom_slice import (
        slice_bom_nodes_for_l0,
    )

    nodes = [
        BomParsedNode(level=0, parent_index=None, part_number="A-AB", sort_no=0),
        BomParsedNode(level=1, parent_index=0, part_number="A1", sort_no=1),
        BomParsedNode(level=0, parent_index=None, part_number="B-AB", sort_no=2),
        BomParsedNode(level=1, parent_index=2, part_number="B1", sort_no=3),
    ]
    sliced = slice_bom_nodes_for_l0(nodes, "B-AB")
    assert [node.part_number for node in sliced] == ["B-AB", "B1"]
    assert sliced[0].parent_index is None
    assert sliced[1].parent_index == 0


def test_validate_bom_l0_matches_change_orders() -> None:
    from app.common.engineering_change_gates import CHANGE_TYPE_STOP
    from app.parsers.bom.models import BomParsedNode
    from app.parsers.engineering_change.design_change.change_order_parser import (
        ChangeOrderRowParsed,
    )
    from app.services.engineering_change.design_change.upload.gates.validate import (
        validate_bom_l0_matches_change_orders,
    )

    change_orders = [
        ChangeOrderRowParsed(
            project_code="T28",
            effective_period_text="2025年07月",
            change_type=CHANGE_TYPE_STOP,
            before_assembly_part_number="533AA",
            after_assembly_part_number=None,
            assembly_name="x",
            part_level="总成",
            before_part_number=None,
            after_part_number=None,
            part_name=None,
            excel_row=5,
        ),
        ChangeOrderRowParsed(
            project_code="T28",
            effective_period_text="2025年07月",
            change_type=CHANGE_TYPE_REVISE,
            before_assembly_part_number="873AA",
            after_assembly_part_number="873AB",
            assembly_name="x",
            part_level="一级",
            before_part_number="875AA",
            after_part_number="875AB",
            part_name="y",
            excel_row=6,
        ),
    ]
    nodes = [
        BomParsedNode(level=0, parent_index=None, part_number="873AB", sort_no=0),
        BomParsedNode(level=1, parent_index=0, part_number="875AB", sort_no=1),
    ]
    validate_bom_l0_matches_change_orders(nodes, change_orders)
    with pytest.raises(BusinessException):
        validate_bom_l0_matches_change_orders(nodes, change_orders[:1])
