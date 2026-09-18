"""工程变更三条门禁与变更单基础校验。"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.common.engineering_change_gates import (
    CHANGE_TYPE_ADD_PART,
    CHANGE_TYPE_REVISE,
    EC_CHANGE_ORDER_ROW_COUNT,
    EC_GATE_ENGINE_BREAKDOWN,
    EC_GATE_FULL_GENERATE,
    PeriodRoot,
    is_sop_alignable_root,
    snapshot_blocks_full_generate,
)
from app.exceptions.base import BusinessException
from app.parsers.engineering_change.change_order_parser import parse_change_order_sheet
from app.services.engineering_change.upload.template_builder import (
    CHANGE_ORDER_SAMPLES,
    PROCESS_SAMPLES,
    TEMPLATE_FILENAME,
    build_engineering_change_template,
)


def _root(
    *,
    effective_year: int = 2025,
    effective_month: int = 1,
    expiry_year: int | None = None,
    expiry_month: int | None = None,
) -> PeriodRoot:
    return PeriodRoot(
        effective_year=effective_year,
        effective_month=effective_month,
        expiry_year=expiry_year,
        expiry_month=expiry_month,
    )


def test_sop_alignable_only_sop_month_active() -> None:
    sop = _root(effective_year=2025, effective_month=1)
    ec = _root(effective_year=2025, effective_month=7)
    expired = _root(
        effective_year=2025,
        effective_month=1,
        expiry_year=2025,
        expiry_month=6,
    )
    assert is_sop_alignable_root(sop, sop_year=2025, sop_month=1, as_of_year=2025, as_of_month=8)
    assert not is_sop_alignable_root(ec, sop_year=2025, sop_month=1, as_of_year=2025, as_of_month=8)
    assert not is_sop_alignable_root(
        expired, sop_year=2025, sop_month=1, as_of_year=2025, as_of_month=8
    )


def test_snapshot_blocks_full_generate_on_ec_or_expired() -> None:
    sop = _root(effective_year=2025, effective_month=1)
    assert not snapshot_blocks_full_generate(
        [sop], sop_year=2025, sop_month=1, as_of_year=2025, as_of_month=3
    )
    ec = _root(effective_year=2025, effective_month=7)
    assert snapshot_blocks_full_generate(
        [sop, ec], sop_year=2025, sop_month=1, as_of_year=2025, as_of_month=8
    )
    expired = _root(
        effective_year=2025,
        effective_month=1,
        expiry_year=2025,
        expiry_month=6,
    )
    assert snapshot_blocks_full_generate(
        [expired], sop_year=2025, sop_month=1, as_of_year=2025, as_of_month=8
    )


def test_gate_message_constants() -> None:
    assert "工程变更" in EC_GATE_FULL_GENERATE
    assert "拆解" in EC_GATE_ENGINE_BREAKDOWN


def test_template_change_order_examples_and_parses_as_book() -> None:
    content, filename = build_engineering_change_template()
    assert filename == TEMPLATE_FILENAME
    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook.sheetnames == [
        "工程变更单",
        "车身件材料清单",
        "材料清单填写说明",
        "冲压工艺确认表",
        "工艺确认填写说明",
        "零件汇总",
    ]
    change_order = workbook["工程变更单"]
    types = [change_order.cell(row, 6).value for row in range(5, 5 + len(CHANGE_ORDER_SAMPLES))]
    assert types == [sample.change_type for sample in CHANGE_ORDER_SAMPLES]
    assert change_order.cell(4 + len(CHANGE_ORDER_SAMPLES), 6).value == CHANGE_TYPE_ADD_PART
    process = workbook["冲压工艺确认表"]
    process_numbers: list[str] = []
    excel_row = 7
    for sample in PROCESS_SAMPLES:
        if sample.part_number in process_numbers:
            continue
        assert process.cell(excel_row, 2).value == sample.part_number
        process_numbers.append(sample.part_number)
        excel_row += sum(1 for item in PROCESS_SAMPLES if item.part_number == sample.part_number)
    assert process_numbers == list(dict.fromkeys(sample.part_number for sample in PROCESS_SAMPLES))
    merged = {str(item) for item in change_order.merged_cells.ranges}
    workbook.close()
    rows = parse_change_order_sheet(content)
    assert len(rows) == len(CHANGE_ORDER_SAMPLES)
    assert [row.change_type for row in rows] == [
        sample.change_type for sample in CHANGE_ORDER_SAMPLES
    ]
    last_row = 4 + len(CHANGE_ORDER_SAMPLES)
    assert merged >= {
        f"A5:A{last_row}",
        f"B5:B{last_row}",
        "C10:C13",
        "D10:D13",
        "E10:E13",
    }
    assert "C5:C9" not in merged
    assert [row.before_assembly_part_number for row in rows[:5]] == [
        sample.before_assembly for sample in CHANGE_ORDER_SAMPLES[:5]
    ]
    assert all(row.before_assembly_part_number == "501008897AA" for row in rows[5:])
    assert all(row.after_assembly_part_number == "501008897AB" for row in rows[5:])


def test_template_bom_l0_has_assembly_sketch() -> None:
    from app.parsers.bom.bom_parser import parse_bom_excel
    from app.parsers.engineering_change.sheet_extract import SHEET_BOM

    content, _ = build_engineering_change_template()
    parsed = parse_bom_excel(content, sheet_name=SHEET_BOM)
    l0_groups: dict[str, list] = {}
    for node in parsed.nodes:
        if int(node.level) != 0:
            continue
        l0_groups.setdefault(node.part_number, []).append(node)
    assert l0_groups
    assert all(any(node.sketch_bytes for node in group) for group in l0_groups.values())
    assert all(
        any((node.internal_number or "").strip() for node in group) for group in l0_groups.values()
    )


def test_template_process_sheet_parses_new_part_numbers() -> None:
    from app.parsers.engineering_change.sheet_extract import SHEET_PROCESS, extract_named_sheets
    from app.services.engineering_change.upload.gates.process_gates import (
        _parse_process_part_numbers,
    )

    content, _ = build_engineering_change_template()
    sheets = extract_named_sheets(content)
    numbers = _parse_process_part_numbers(sheets[SHEET_PROCESS])
    assert numbers == {sample.part_number for sample in PROCESS_SAMPLES}


def test_change_type_revise_constant() -> None:
    assert CHANGE_TYPE_REVISE == "改图"


def test_empty_change_order_rejected() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "工程变更单"
    headers = (
        "项目",
        "生效年月",
        "变更前总成",
        "变更后总成",
        "总成名称",
        "变更类型",
        "零件层级",
        "变更前零件",
        "变更后零件",
        "零件名称",
    )
    for col, header in enumerate(headers, start=1):
        sheet.cell(1, col, header)
    buffer = BytesIO()
    workbook.save(buffer)
    with pytest.raises(BusinessException) as exc:
        parse_change_order_sheet(buffer.getvalue())
    assert str(exc.value) == EC_CHANGE_ORDER_ROW_COUNT


def test_slice_bom_nodes_for_l0_remaps_parent_index() -> None:
    from app.parsers.bom.models import BomParsedNode
    from app.services.engineering_change.upload.gates.bom_slice import slice_bom_nodes_for_l0

    nodes = [
        BomParsedNode(level=0, parent_index=None, part_number="A-AB", sort_no=0),
        BomParsedNode(level=1, parent_index=0, part_number="A1", sort_no=1),
        BomParsedNode(level=0, parent_index=None, part_number="B-AB", sort_no=2),
        BomParsedNode(level=1, parent_index=2, part_number="B1", sort_no=3),
    ]
    sliced = slice_bom_nodes_for_l0(nodes, "B-AB")
    assert [node.part_number for node in sliced] == ["B-AB", "B1"]
    assert sliced[0].parent_index is None
    assert sliced[1].parent_index == 0


def test_validate_bom_l0_matches_change_orders() -> None:
    from app.common.engineering_change_gates import CHANGE_TYPE_STOP
    from app.parsers.bom.models import BomParsedNode
    from app.parsers.engineering_change.change_order_parser import ChangeOrderRowParsed
    from app.services.engineering_change.upload.gates.validate import (
        validate_bom_l0_matches_change_orders,
    )

    change_orders = [
        ChangeOrderRowParsed(
            project_code="T28",
            effective_period_text="2025年07月",
            change_type=CHANGE_TYPE_STOP,
            before_assembly_part_number="533AA",
            after_assembly_part_number=None,
            assembly_name="x",
            part_level="总成",
            before_part_number=None,
            after_part_number=None,
            part_name=None,
            excel_row=5,
        ),
        ChangeOrderRowParsed(
            project_code="T28",
            effective_period_text="2025年07月",
            change_type=CHANGE_TYPE_REVISE,
            before_assembly_part_number="873AA",
            after_assembly_part_number="873AB",
            assembly_name="x",
            part_level="一级",
            before_part_number="875AA",
            after_part_number="875AB",
            part_name="y",
            excel_row=6,
        ),
    ]
    nodes = [
        BomParsedNode(level=0, parent_index=None, part_number="873AB", sort_no=0),
        BomParsedNode(level=1, parent_index=0, part_number="875AB", sort_no=1),
    ]
    validate_bom_l0_matches_change_orders(nodes, change_orders)
    with pytest.raises(BusinessException):
        validate_bom_l0_matches_change_orders(nodes, change_orders[:1])
