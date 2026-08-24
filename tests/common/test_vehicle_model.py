"""车型代号归一化测试。"""

from __future__ import annotations

from openpyxl import Workbook

from app.common.vehicle_model import normalize_vehicle_model
from app.parsers.bom.meta_parser import parse_bom_meta
from app.parsers.stamping.header_resolver import HeaderResolver


def test_normalize_vehicle_model_strips_chinese_prefix_and_suffix() -> None:
    assert normalize_vehicle_model("青岛T28") == "T28"
    assert normalize_vehicle_model("T28样例") == "T28"
    assert normalize_vehicle_model("T28") == "T28"
    assert normalize_vehicle_model("  青岛 T28  ") == "T28"
    assert normalize_vehicle_model("青岛") is None
    assert normalize_vehicle_model(None) is None


def test_stamping_title_prefix_normalized_to_t28() -> None:
    sheet = Workbook().active
    sheet["D2"] = "青岛T28项目冲压工艺确认表-V1"

    resolver = HeaderResolver(lambda row, col: sheet.cell(row=row, column=col).value)
    assert resolver.find_stamping_vehicle_model(sheet) == "T28"


def test_bom_meta_project_code_with_chinese_prefix() -> None:
    sheet = Workbook().active
    sheet["A3"] = "项目代号：青岛T28"

    meta = parse_bom_meta(sheet)
    assert meta.vehicle_model == "T28"
