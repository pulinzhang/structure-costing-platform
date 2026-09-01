"""导出 Excel 内控水印。"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.common.export_watermark import (
    apply_workbook_watermark,
    format_export_watermark_text,
    workbook_to_xlsx_bytes,
)
from app.schemas.vo.user_vo import UserVO

_STAMP = datetime(2026, 1, 31, 17, 0)


def _exporter() -> UserVO:
    return UserVO(
        id=1,
        username="zhangsan",
        real_name="张三",
        department="采购部",
    )


def test_format_includes_department_name_and_time() -> None:
    text = format_export_watermark_text(_exporter(), exported_at=_STAMP)
    assert text == "导出人：采购部-张三，时间：2026-01-31 17:00"


def test_format_without_department_uses_display_name() -> None:
    user = UserVO(id=2, username="lisi")
    text = format_export_watermark_text(user, exported_at=_STAMP)
    assert text == "导出人：lisi，时间：2026-01-31 17:00"


def test_workbook_bytes_contain_print_watermark_and_images() -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet["A1"] = "核价"
    content = workbook_to_xlsx_bytes(workbook, _exporter(), exported_at=_STAMP)
    loaded = load_workbook(BytesIO(content))
    out = loaded.worksheets[0]
    text = format_export_watermark_text(_exporter(), exported_at=_STAMP)
    assert text in (out.oddHeader.center.text or "")
    assert text in (out.oddFooter.center.text or "")
    assert len(out._images) >= 1
    loaded.close()


def test_print_watermark_appends_existing_footer_left() -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.oddFooter.left.text = "仅供核价核对使用"
    sheet.oddFooter.center.text = "第 &P 页 / 共 &N 页"
    text = format_export_watermark_text(_exporter(), exported_at=_STAMP)
    apply_workbook_watermark(workbook, text)
    assert "仅供核价核对使用" in (sheet.oddFooter.left.text or "")
    assert text in (sheet.oddFooter.left.text or "")
    assert text in (sheet.oddHeader.center.text or "")
    apply_workbook_watermark(workbook, text)
    assert (sheet.oddFooter.left.text or "").count(text) == 1
