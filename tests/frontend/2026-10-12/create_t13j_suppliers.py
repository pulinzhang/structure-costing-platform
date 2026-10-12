"""从前端新增四家冲压承制公司。全称和缩写读自工艺确认表表头，不写库。

先开开发服务（http://127.0.0.1:8000），在仓库根目录执行：

    python -X utf8 tests/frontend/2026-10-12/create_t13j_suppliers.py

全称是「冲压承制公司」右侧格子。缩写是标题「工艺确认表-」后面的短名。
工艺表没有对外代码，新增表单必填，测试里用 CQ001 / QX001 / ZH001 / BQ001。
管理员走登录页 admin / admin123。
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import Page, sync_playwright

BASE = "http://127.0.0.1:8000"
SAMPLE_DIR = Path(__file__).resolve().parents[3] / "excel_sample" / "T13J"
REPO_ROOT = Path(__file__).resolve().parents[3]
ERP_BY_ABBREVIATION = {
    "常庆": "CQ001",
    "全鑫": "QX001",
    "振华": "ZH001",
    "贝奇": "BQ001",
}


def read_suppliers() -> list[tuple[str, str, str]]:
    sys.path.insert(0, str(REPO_ROOT))
    from app.parsers.stamping.header_resolver import HeaderResolver

    rows: list[tuple[str, str, str]] = []
    for path in sorted(SAMPLE_DIR.glob("*工艺确认表*.xlsx")):
        workbook = load_workbook(path, data_only=True)
        sheet = workbook.worksheets[0]

        def cell_value(row_idx: int, col_idx: int, current=sheet) -> object:
            return current.cell(row_idx, col_idx).value

        resolver = HeaderResolver(cell_value)
        full_name = (resolver.find_stamping_supplier(sheet) or "").strip()
        title = str(sheet["D2"].value or "")
        workbook.close()
        marker = "工艺确认表-"
        if marker not in title or not full_name:
            raise SystemExit(f"{path.name} 没有读到冲压承制公司或标题短名")
        abbreviation = title.split(marker, 1)[1].strip()
        erp_code = ERP_BY_ABBREVIATION.get(abbreviation)
        if not erp_code:
            raise SystemExit(f"{path.name} 的缩写「{abbreviation}」没有测试用对外代码")
        rows.append((abbreviation, full_name, erp_code))
    if len(rows) != 4:
        raise SystemExit(f"期望 4 家供应商，读到 {len(rows)} 家")
    return rows


def login(page: Page) -> None:
    page.goto(f"{BASE}/login", wait_until="domcontentloaded")
    page.fill("#loginUsername", "admin")
    page.fill("#loginPassword", "admin123")
    page.click("#btnLogin")
    page.wait_for_url("**/app**", timeout=20000)


def open_supplier_list(page: Page) -> None:
    with page.expect_response(
        lambda response: response.request.method == "GET"
        and "/api/masterdata/suppliers" in response.url
        and "page=" in response.url,
        timeout=20000,
    ) as listed:
        page.goto(f"{BASE}/app#/supplier-list", wait_until="domcontentloaded")
    response = listed.value
    if response.status != 200:
        raise RuntimeError(f"供应商列表打开失败: {response.status} {response.text()[:300]}")
    page.wait_for_selector("#supplierBtnCreate", timeout=20000)


def create_supplier(page: Page, abbreviation: str, full_name: str, erp_code: str) -> None:
    page.click("#supplierBtnCreate")
    page.wait_for_selector("#supplierCreateAbbreviation", state="visible", timeout=10000)
    page.fill("#supplierCreateAbbreviation", abbreviation)
    page.fill("#supplierCreateFullName", full_name)
    page.fill("#supplierCreateErpCode", erp_code)
    with page.expect_response(
        lambda response: response.request.method == "POST"
        and response.url.rstrip("/").endswith("/api/masterdata/suppliers"),
        timeout=20000,
    ) as created:
        page.click("#submitSupplierCreateBtn")
    response = created.value
    if response.status != 201:
        raise RuntimeError(f"新增 {abbreviation} 失败: {response.status} {response.text()[:400]}")
    page.wait_for_function(
        "() => document.getElementById('supplierCreateDrawer')?.style.display === 'none'",
        timeout=10000,
    )
    row = page.locator("#supplierTableBody tr", has_text=full_name)
    row.wait_for(timeout=10000)
    text = row.inner_text()
    if abbreviation not in text or erp_code not in text:
        raise RuntimeError(f"列表行没有缩写或对外代码: {text}")


def main() -> None:
    suppliers = read_suppliers()
    for abbreviation, full_name, erp_code in suppliers:
        print(f"read: {abbreviation} / {full_name} / {erp_code}")
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        page = browser.new_page()
        login(page)
        open_supplier_list(page)
        for abbreviation, full_name, erp_code in suppliers:
            create_supplier(page, abbreviation, full_name, erp_code)
            print(f"created: {abbreviation}")
        info = page.locator("#supplierPageInfo").inner_text().strip()
        if info != "共 4 条":
            raise SystemExit(f"列表条数不是 4: {info}")
        browser.close()
    print("ok: 4 家供应商已从新增抽屉建档")


if __name__ == "__main__":
    main()
