"""采购价格只读报表：窗口、汇总、工装、导出。"""

from __future__ import annotations

from collections.abc import Generator
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base
from app.exceptions.base import BusinessException, NotFoundException
from app.models.bom.node import BomNode, BomSnapshot
from app.models.bom.procurement_object import BomProcurementObject, BomProcurementObjectList
from app.models.costing.assembly_cost import AssemblyCost
from app.models.masterdata.location import Location, LocationAlias
from app.models.masterdata.supplier import Supplier, SupplierAlias
from app.models.masterdata.vehicle_model import VehicleModel
from app.models.settings.system_config import SystemConfig
from app.models.stamping.die_fee import StampingDieFee
from app.models.stamping.gauge_fee import StampingGaugeFee
from app.repositories.bom.bom_repository import BomNodeRepository, BomSnapshotRepository
from app.repositories.bom.procurement_object_repository import BomProcurementObjectRepository
from app.schemas.dto.bom_dto import BomNodeWrite, BomSnapshotWrite
from app.schemas.dto.bom_procurement_object_dto import (
    BomProcurementObjectListWrite,
    BomProcurementObjectWrite,
)
from app.schemas.dto.procurement_contract_report_dto import ProcurementContractReportQuery
from app.services.costing.procurement_contract_report.assemble import (
    MonthAmount,
    PriceRow,
    seq_sort_key,
    sum_vehicle_total,
)
from app.services.costing.procurement_contract_report.export import (
    ProcurementContractReportExportOrchestrator,
)
from app.services.costing.procurement_contract_report.export.styles import (
    agreement_column_widths,
)
from app.services.costing.procurement_contract_report.export.values import agreement_remarks
from app.services.costing.procurement_contract_report.service import (
    ProcurementContractReportService,
)
from app.services.costing.procurement_contract_report.window import expand_months


@pytest.fixture()
def db() -> Generator[Session, None, None]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _enable_fk(dbapi_connection: object, _connection_record: object) -> None:
        cursor = dbapi_connection.cursor()  # type: ignore[attr-defined]
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    Base.metadata.create_all(
        bind=engine,
        tables=[
            Location.__table__,
            LocationAlias.__table__,
            VehicleModel.__table__,
            Supplier.__table__,
            SupplierAlias.__table__,
            BomSnapshot.__table__,
            BomNode.__table__,
            BomProcurementObjectList.__table__,
            BomProcurementObject.__table__,
            AssemblyCost.__table__,
            StampingDieFee.__table__,
            StampingGaugeFee.__table__,
            SystemConfig.__table__,
        ],
    )
    factory = sessionmaker(bind=engine)
    session = factory()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def _seed_masters(db: Session) -> None:
    db.add(Location(location="芜湖", is_enabled=True, sort_order=1))
    db.flush()
    db.add(
        VehicleModel(
            location="芜湖",
            vehicle_model="T13J",
            effective_year=2020,
            effective_month=1,
            current_supplied_units=0,
        )
    )
    db.add(
        Supplier(
            full_name="烟台海众汽车部件有限公司",
            abbreviation="烟台海众",
            erp_code="HZ001",
            is_enabled=True,
            sort_order=1,
        )
    )
    db.add(
        Supplier(
            full_name="安徽振华汽车部件有限公司",
            abbreviation="安徽振华",
            erp_code="ZH001",
            is_enabled=True,
            sort_order=2,
        )
    )
    db.flush()


def _seed_tree(db: Session) -> tuple[BomSnapshot, BomNode, BomNode, BomNode]:
    snapshot = BomSnapshotRepository().create(
        db,
        BomSnapshotWrite(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="D/09",
            source_filename="t13j.xlsx",
            created_by="tester",
            updated_by="tester",
        ),
    )
    nodes = BomNodeRepository()
    l0 = nodes.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot.id,
            level=0,
            sort_no=1,
            seq_no=5,
            part_number="551003086AA",
            part_name="右前门上铰链加强板总成",
            internal_number="ACA674",
            qty=Decimal("1"),
            node_role="supply",
            is_supply_unit=True,
            created_by="tester",
            updated_by="tester",
        ),
    )
    stamping = nodes.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot.id,
            parent_id=l0.id,
            level=1,
            sort_no=2,
            part_number="551003096AA",
            part_name="右前门上铰链加强板",
            internal_number="ACA674001",
            qty=Decimal("1"),
            node_role="make",
            created_by="tester",
            updated_by="tester",
        ),
    )
    subassembly = nodes.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot.id,
            parent_id=l0.id,
            level=1,
            sort_no=3,
            part_number="551001137AA",
            part_name="左后门上铰链安装板总成",
            internal_number="ACA333N01",
            qty=Decimal("1"),
            node_role="structure",
            created_by="tester",
            updated_by="tester",
        ),
    )
    return snapshot, l0, stamping, subassembly


def _object_write(
    *,
    list_id: int,
    kind: str,
    node: BomNode,
    l0: BomNode,
    supplier_name: str | None,
    excel_row: int,
    confirm_status: str = "outsourced",
) -> BomProcurementObjectWrite:
    return BomProcurementObjectWrite(
        list_id=list_id,
        section="object",
        object_kind=kind,
        degenerate_subtype=None,
        source_node_id=node.id,
        part_number=node.part_number,
        part_name=node.part_name,
        internal_number=node.internal_number,
        level=node.level,
        node_role=node.node_role,
        qty=node.qty,
        qty_unit="pcs",
        l0_part_number=l0.part_number,
        l0_part_name=l0.part_name,
        parent_part_number=l0.part_number if node.level > 0 else None,
        parent_part_name=l0.part_name if node.level > 0 else None,
        excel_filename="t13j.xlsx",
        excel_sheet_name="材料清单",
        excel_row=excel_row,
        child_count=0,
        weld_info=None,
        remark=None,
        has_outsource_remark=False,
        material=None,
        thickness_text=None,
        material_mill=None,
        data_gap=None,
        alert_code=None,
        confirm_status=confirm_status,
        confirmed_by="tester",
        confirmed_at=datetime(2026, 1, 1),
        confirm_remark=None,
        supplier_name=supplier_name,
        audit_status="pending",
        created_by="tester",
        updated_by="tester",
    )


def _seed_pack(db: Session) -> tuple[BomNode, BomNode]:
    _seed_masters(db)
    snapshot, l0, stamping, subassembly = _seed_tree(db)
    repo = BomProcurementObjectRepository()
    header = repo.create_list(
        db,
        BomProcurementObjectListWrite(
            snapshot_id=snapshot.id,
            location="芜湖",
            vehicle_model="T13J",
            bom_version="D/09",
            source_filename="t13j.xlsx",
            qa_passed=True,
            qa_error_text=None,
            l0_count=1,
            l0_with_children_count=1,
            l2_count=0,
            subassembly_child_sum=0,
            object_l0_degenerate_count=0,
            object_l1_subassembly_count=1,
            object_l1_stamping_count=1,
            other_leaf_count=0,
            anomaly_count=0,
            generated_at=datetime(2026, 1, 1),
            generated_by="tester",
            created_by="tester",
            updated_by="tester",
        ),
    )
    repo.bulk_create_objects(
        db,
        [
            _object_write(
                list_id=header.id,
                kind="l1_stamping",
                node=stamping,
                l0=l0,
                supplier_name="烟台海众汽车部件有限公司",
                excel_row=10,
            ),
            _object_write(
                list_id=header.id,
                kind="l1_subassembly",
                node=subassembly,
                l0=l0,
                supplier_name="安徽振华汽车部件有限公司",
                excel_row=11,
            ),
        ],
    )
    for node, amount in (
        (stamping, Decimal("5.5250")),
        (subassembly, Decimal("29.5236")),
    ):
        db.add(
            AssemblyCost(
                year=2026,
                month=1,
                location="芜湖",
                vehicle_model="T13J",
                bom_version="D/09",
                source_node_id=node.id,
                part_number=node.part_number,
                part_name=node.part_name,
                should_cost=amount,
                total_amount=amount,
            )
        )
    db.add(
        StampingDieFee(
            year=2026,
            location="芜湖",
            vehicle_model="T13J",
            part_number="551003096AA",
            price_kind="steel_plate",
            unit_price=Decimal("1.2000"),
            allocation_qty=Decimal("10"),
            die_fee=Decimal("0.4000"),
            final_die_fee=Decimal("0.4000"),
        )
    )
    db.add(
        StampingGaugeFee(
            location="芜湖",
            vehicle_model="T13J",
            part_number="551003096AA",
            gauge_fee=Decimal("0.8000"),
            allocation_qty=Decimal("10"),
            status="effective",
        )
    )
    db.flush()
    return stamping, subassembly


def _query(**overrides: object) -> ProcurementContractReportQuery:
    payload: dict[str, object] = {
        "location": "芜湖",
        "vehicle_model": "T13J",
        "bom_version": "D/09",
        "start_year": 2026,
        "start_month": 1,
        "end_year": 2026,
        "end_month": 1,
        "page": 1,
        "size": 50,
    }
    payload.update(overrides)
    return ProcurementContractReportQuery.model_validate(payload)


def test_agreement_column_widths_follow_month_count() -> None:
    one = agreement_column_widths(1)
    many = agreement_column_widths(3)
    assert len(one) == 12
    assert len(many) == 14
    assert one[6] == 11.0
    assert many[6] == many[7] == many[8] == 14.0
    assert one[3] == 36.0
    assert one[7] < one[9]


def test_expand_months_cross_year() -> None:
    months = expand_months(start_year=2025, start_month=11, end_year=2026, end_month=2)
    assert months == [(2025, 11), (2025, 12), (2026, 1), (2026, 2)]


def test_expand_months_rejects_over_12() -> None:
    with pytest.raises(BusinessException, match="最多 12 个月"):
        expand_months(start_year=2026, start_month=1, end_year=2027, end_month=2)


def _seed_agreement_settings(
    db: Session,
    *,
    company_name: str = "",
    remarks: str = "",
) -> None:
    db.add(
        SystemConfig(
            id=1,
            app_name="test",
            company_name=company_name,
            procurement_agreement_remarks=remarks,
        )
    )
    db.flush()


def test_agreement_remarks_empty_without_template() -> None:
    bundle = SimpleNamespace(months=[(2025, 1), (2025, 8)])
    assert agreement_remarks(bundle) == ()  # type: ignore[arg-type]
    assert agreement_remarks(bundle, "某厂", "  ") == ()  # type: ignore[arg-type]


def test_agreement_remarks_uses_settings_template() -> None:
    bundle = SimpleNamespace(months=[(2026, 2), (2026, 3)])
    remarks = agreement_remarks(
        bundle,  # type: ignore[arg-type]
        company_name="某冲压厂",
        remarks_template=(
            "甲方{company_name}有效期{start_year}年{start_month}月1日至"
            "{end_year}年{end_month}月{end_day}日。\n\n"
            "留空{unknown_token}。"
        ),
    )
    assert remarks == (
        "甲方某冲压厂有效期2026年2月1日至2026年3月31日。",
        "留空{unknown_token}。",
    )


def test_seq_sort_key_orders_by_l0_then_child() -> None:
    labels = ["10-1", "5-2", "5-1", "5", ""]
    ordered = sorted(labels, key=seq_sort_key)
    assert ordered == ["5", "5-1", "5-2", "10-1", ""]


def test_list_price_rows_and_tooling(db: Session) -> None:
    _seed_pack(db)
    rows, total, filters = ProcurementContractReportService().paginate(db, _query())
    assert total == 2
    assert filters.qa_passed is True
    assert filters.window_month_count == 1
    assert [row.seq_label for row in rows] == ["5-1", "5-2"]
    by_part = {row.part_number: row for row in rows}
    stamping = by_part["551003096AA"]
    subassembly = by_part["551001137AA"]
    assert stamping.seq_label == "5-1"
    assert stamping.internal_number == "ACA674001"
    assert stamping.supplier_name == "烟台海众汽车部件有限公司"
    assert stamping.month_amounts[0].should_cost == Decimal("5.5250")
    assert stamping.tooling_alloc_price == Decimal("0.4800")
    assert stamping.tooling_fee == Decimal("48000.0000")
    assert stamping.tooling_alloc_qty == Decimal("100000")
    assert subassembly.seq_label == "5-2"
    assert subassembly.month_amounts[0].should_cost == Decimal("29.5236")
    assert subassembly.tooling_alloc_price is None
    assert subassembly.tooling_fee is None
    assert subassembly.tooling_alloc_qty is None
    dumped = subassembly.model_dump()
    assert dumped["tooling_alloc_price"] == "/"
    assert dumped["tooling_fee"] == "/"
    assert dumped["tooling_alloc_qty"] == "/"
    assert filters.vehicle_total == Decimal("35.0486")
    assert filters.vehicle_total_skipped == 0


def test_list_fuzzy_matches_part_number_and_name(db: Session) -> None:
    _seed_pack(db)
    by_number, total_number, _filters = ProcurementContractReportService().paginate(
        db,
        _query(part_number="096"),
    )
    assert total_number == 1
    assert by_number[0].part_number == "551003096AA"
    by_name, total_name, _name_filters = ProcurementContractReportService().paginate(
        db,
        _query(part_name="加强板"),
    )
    assert total_name == 1
    assert by_name[0].part_number == "551003096AA"


def test_open_defaults_use_latest_list_and_cost_month(db: Session) -> None:
    _seed_pack(db)
    defaults = ProcurementContractReportService().get_open_defaults(db)
    assert defaults.location == "芜湖"
    assert defaults.vehicle_model == "T13J"
    assert defaults.bom_version == "D/09"
    assert defaults.start_year == 2026
    assert defaults.start_month == 1
    assert defaults.end_year == 2026
    assert defaults.end_month == 1


def test_sum_vehicle_total_sums_all_window_months() -> None:
    priced = PriceRow(
        seq_label="5-1",
        internal_number=None,
        part_number="A",
        part_name=None,
        vehicle_qty=Decimal("2"),
        object_kind=None,
        first_l0_part_number="L0",
        month_amounts=[
            MonthAmount(year=2026, month=1, should_cost=Decimal("1.5000")),
            MonthAmount(year=2026, month=2, should_cost=Decimal("2.5000")),
        ],
        tooling_alloc_price=None,
        tooling_fee=None,
        tooling_alloc_qty=None,
    )
    missing = PriceRow(
        seq_label="5-2",
        internal_number=None,
        part_number="B",
        part_name=None,
        vehicle_qty=Decimal("1"),
        object_kind=None,
        first_l0_part_number="L0",
        month_amounts=[
            MonthAmount(year=2026, month=1, should_cost=None),
            MonthAmount(year=2026, month=2, should_cost=None),
        ],
        tooling_alloc_price=None,
        tooling_fee=None,
        tooling_alloc_qty=None,
    )
    total, skipped = sum_vehicle_total([priced, missing], [(2026, 1), (2026, 2)])
    assert total == Decimal("8.0000")
    assert skipped == 1


def test_missing_list_is_not_found(db: Session) -> None:
    _seed_masters(db)
    with pytest.raises(NotFoundException, match="请先生成采购对象清单"):
        ProcurementContractReportService().paginate(db, _query())


def test_conflict_same_part_two_prices(db: Session) -> None:
    stamping, _sub = _seed_pack(db)
    db.add(
        AssemblyCost(
            year=2026,
            month=1,
            location="芜湖",
            vehicle_model="T13J",
            bom_version="D/09",
            source_node_id=stamping.id + 99999,
            part_number="ghost",
            should_cost=Decimal("9.0000"),
            total_amount=Decimal("9.0000"),
        )
    )
    nodes = BomNodeRepository()
    snapshot = db.query(BomSnapshot).one()
    l0 = db.query(BomNode).filter(BomNode.level == 0).one()
    twin = nodes.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot.id,
            parent_id=l0.id,
            level=1,
            sort_no=9,
            part_number="551003096AA",
            part_name="右前门上铰链加强板",
            internal_number="ACA674001",
            qty=Decimal("1"),
            node_role="make",
            created_by="tester",
            updated_by="tester",
        ),
    )
    header = db.query(BomProcurementObjectList).one()
    BomProcurementObjectRepository().bulk_create_objects(
        db,
        [
            _object_write(
                list_id=header.id,
                kind="l1_stamping",
                node=twin,
                l0=l0,
                supplier_name="烟台海众汽车部件有限公司",
                excel_row=20,
            )
        ],
    )
    db.add(
        AssemblyCost(
            year=2026,
            month=1,
            location="芜湖",
            vehicle_model="T13J",
            bom_version="D/09",
            source_node_id=twin.id,
            part_number="551003096AA",
            should_cost=Decimal("6.0000"),
            total_amount=Decimal("6.0000"),
        )
    )
    db.flush()
    with pytest.raises(BusinessException, match="多个结算价格"):
        ProcurementContractReportService().paginate(db, _query())


def test_export_matches_template_shape(db: Session) -> None:
    _seed_pack(db)
    _seed_agreement_settings(
        db,
        company_name="测试甲方",
        remarks=(
            "5、上述价格执行有效期为{start_year}年{start_month}月1日起至"
            "{end_year}年{end_month}月{end_day}日。"
        ),
    )
    content, filename = ProcurementContractReportExportOrchestrator().export(
        db,
        _query(supplier_name="烟台海众汽车部件有限公司"),
    )
    assert filename.endswith(".xlsx")
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["批量零部件价格协议"]
    assert sheet["A1"].value == "附件二"
    assert sheet["B1"].value == "批量零部件价格协议"
    assert sheet["A2"].value == "编号："
    assert sheet["F2"].value == "币种：人民币"
    assert sheet["H2"].value == "价格均不含增值税"
    assert sheet.column_dimensions["D"].width == 36
    assert sheet.row_dimensions[6].height == 22
    assert sheet["A3"].value == "序号"
    assert sheet["G3"].value == "结算价格\n（元）"
    assert sheet["H3"].value == "工装费"
    assert sheet["A6"].value == "5-1"
    assert sheet["C6"].value == "551003096AA"
    assert sheet["E6"].value == 1
    assert sheet["E6"].number_format == "0"
    assert isinstance(sheet["G6"].value, int | float)
    assert sheet["G6"].value == 5.53
    assert sheet["G6"].number_format == "0.00"
    assert isinstance(sheet["I6"].value, int | float)
    assert sheet["I6"].value == 0.48
    assert sheet["I6"].number_format == "0.00"
    assert sheet["J6"].value == 48000
    assert sheet["J6"].number_format == "#,##0"
    assert sheet["K6"].value == 100000
    assert sheet["K6"].number_format == "0"
    assert not sheet._images
    assert workbook.sheetnames == ["批量零部件价格协议"]
    clause_5 = next(
        str(sheet.cell(row, 1).value)
        for row in range(1, 40)
        if str(sheet.cell(row, 1).value or "").startswith("5、上述价格执行有效期")
    )
    assert clause_5 == "5、上述价格执行有效期为2026年1月1日起至2026年1月31日。"
    assert sheet.freeze_panes == "A6"
    assert sheet.column_dimensions["A"].width == 8
    assert sheet.column_dimensions["G"].width == 11
    assert sheet.column_dimensions["H"].width == 8
    assert sheet.column_dimensions["J"].width == 12
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.oddFooter.center.text == "第 &P/&N 页"
    assert sheet.evenFooter.center.text == "第 &P/&N 页"
    date_labels = {
        str(sheet.cell(row, column).value)
        for row in range(1, 50)
        for column in (1, 8)
        if sheet.cell(row, column).value
    }
    assert "日    期：" in date_labels
    assert "日   期：" in date_labels


def test_export_missing_tooling_is_slash(db: Session) -> None:
    _seed_pack(db)
    content, _filename = ProcurementContractReportExportOrchestrator().export(
        db,
        _query(),
    )
    sheet = load_workbook(BytesIO(content))["批量零部件价格协议"]
    by_part = {str(sheet.cell(row, 3).value): row for row in range(6, 8)}
    sub_row = by_part["551001137AA"]
    assert sheet.cell(sub_row, 9).value == "/"
    assert sheet.cell(sub_row, 10).value == "/"
    assert sheet.cell(sub_row, 11).value == "/"


def test_degenerate_seq_has_no_dash(db: Session) -> None:
    _seed_masters(db)
    snapshot = BomSnapshotRepository().create(
        db,
        BomSnapshotWrite(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="D/09",
            source_filename="t13j.xlsx",
            created_by="tester",
            updated_by="tester",
        ),
    )
    l0 = BomNodeRepository().create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot.id,
            level=0,
            sort_no=1,
            seq_no=5,
            part_number="L0-ONLY",
            part_name="退化供货件",
            internal_number="ACA000",
            qty=Decimal("2"),
            node_role="supply",
            is_supply_unit=True,
            created_by="tester",
            updated_by="tester",
        ),
    )
    repo = BomProcurementObjectRepository()
    header = repo.create_list(
        db,
        BomProcurementObjectListWrite(
            snapshot_id=snapshot.id,
            location="芜湖",
            vehicle_model="T13J",
            bom_version="D/09",
            source_filename="t13j.xlsx",
            qa_passed=True,
            qa_error_text=None,
            l0_count=1,
            l0_with_children_count=0,
            l2_count=0,
            subassembly_child_sum=0,
            object_l0_degenerate_count=1,
            object_l1_subassembly_count=0,
            object_l1_stamping_count=0,
            other_leaf_count=0,
            anomaly_count=0,
            generated_at=datetime(2026, 1, 1),
            generated_by="tester",
            created_by="tester",
            updated_by="tester",
        ),
    )
    repo.bulk_create_objects(
        db,
        [
            _object_write(
                list_id=header.id,
                kind="l0_degenerate",
                node=l0,
                l0=l0,
                supplier_name="烟台海众汽车部件有限公司",
                excel_row=8,
            )
        ],
    )
    db.flush()
    rows, total, _filters = ProcurementContractReportService().paginate(db, _query())
    assert total == 1
    assert rows[0].seq_label == "5"
    assert rows[0].vehicle_qty == Decimal("2")
    assert rows[0].month_amounts[0].should_cost is None
    assert rows[0].month_amounts[0].model_dump()["should_cost"] == "/"
    content, _filename = ProcurementContractReportExportOrchestrator().export(db, _query())
    sheet = load_workbook(BytesIO(content))["批量零部件价格协议"]
    assert sheet["G6"].value == "/"
    assert not sheet._images


def test_pending_objects_appear_in_default_list(db: Session) -> None:
    _seed_masters(db)
    snapshot, l0, stamping, _sub = _seed_tree(db)
    repo = BomProcurementObjectRepository()
    header = repo.create_list(
        db,
        BomProcurementObjectListWrite(
            snapshot_id=snapshot.id,
            location="芜湖",
            vehicle_model="T13J",
            bom_version="D/09",
            source_filename="t13j.xlsx",
            qa_passed=True,
            qa_error_text=None,
            l0_count=1,
            l0_with_children_count=1,
            l2_count=0,
            subassembly_child_sum=0,
            object_l0_degenerate_count=0,
            object_l1_subassembly_count=0,
            object_l1_stamping_count=1,
            other_leaf_count=0,
            anomaly_count=0,
            generated_at=datetime(2026, 1, 1),
            generated_by="tester",
            created_by="tester",
            updated_by="tester",
        ),
    )
    repo.bulk_create_objects(
        db,
        [
            _object_write(
                list_id=header.id,
                kind="l1_stamping",
                node=stamping,
                l0=l0,
                supplier_name=None,
                excel_row=10,
                confirm_status="pending",
            )
        ],
    )
    db.flush()
    service = ProcurementContractReportService()
    rows, total, filters = service.paginate(db, _query())
    assert total == 1
    assert filters.empty_hint is None
    assert rows[0].part_number == stamping.part_number


def _seed_l0_cost(db: Session, node: BomNode, amount: Decimal) -> None:
    db.add(
        AssemblyCost(
            year=2026,
            month=1,
            location="芜湖",
            vehicle_model="T13J",
            bom_version="D/09",
            source_node_id=node.id,
            part_number=node.part_number,
            part_name=node.part_name,
            should_cost=amount,
            total_amount=amount,
        )
    )
    db.flush()


def test_assembly_view_one_tree_root_not_objects(db: Session) -> None:
    _seed_pack(db)
    l0 = db.query(BomNode).filter(BomNode.level == 0).one()
    _seed_l0_cost(db, l0, Decimal("40.0000"))
    rows, total, filters = ProcurementContractReportService().paginate(
        db, _query(view="assembly", supplier_name="烟台海众汽车部件有限公司")
    )
    assert total == 1
    assert filters.view == "assembly"
    assert filters.specified_supplier_count == 0
    assert filters.supplier_name is None
    assert rows[0].part_number == "551003086AA"
    assert rows[0].part_name == "右前门上铰链加强板总成"
    assert rows[0].internal_number == "ACA674"
    assert rows[0].seq_label == "5"
    assert "-" not in rows[0].seq_label
    assert rows[0].vehicle_qty == Decimal("1")
    assert rows[0].month_amounts[0].should_cost == Decimal("40.0000")
    assert rows[0].tooling_alloc_price is None


def test_assembly_view_does_not_need_object_list(db: Session) -> None:
    _seed_masters(db)
    _snapshot, l0, _stamping, _sub = _seed_tree(db)
    _seed_l0_cost(db, l0, Decimal("12.3400"))
    rows, total, filters = ProcurementContractReportService().paginate(db, _query(view="assembly"))
    assert total == 1
    assert filters.empty_hint is None
    assert rows[0].part_number == l0.part_number


def test_assembly_missing_snapshot_is_not_found(db: Session) -> None:
    _seed_masters(db)
    with pytest.raises(NotFoundException, match="未找到对应材料清单"):
        ProcurementContractReportService().paginate(db, _query(view="assembly"))


def test_assembly_export_blank_party_b_and_no_supplier_in_filename(db: Session) -> None:
    _seed_pack(db)
    l0 = db.query(BomNode).filter(BomNode.level == 0).one()
    _seed_l0_cost(db, l0, Decimal("40.0000"))
    content, filename = ProcurementContractReportExportOrchestrator().export(
        db, _query(view="assembly")
    )
    assert "总成" in filename
    assert "烟台海众" not in filename
    sheet = load_workbook(BytesIO(content))["批量零部件价格协议"]
    assert sheet["A6"].value == "5"
    assert sheet["B6"].value == "ACA674"
    assert sheet["C6"].value == "551003086AA"
    assert sheet["G6"].value == 40.0
    assert sheet["I6"].value == "/"
    party_b = {str(sheet.cell(row, 8).value) for row in range(1, 50) if sheet.cell(row, 8).value}
    assert "乙   方：" in party_b
    assert not any("烟台海众" in item or "安徽振华" in item for item in party_b)
