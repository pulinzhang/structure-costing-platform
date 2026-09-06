"""总成核价计算报告 Excel 口径与筛选。"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.schemas.vo.assembly_cost_vo import (
    AssemblyCostCalcCardTraceVO,
    AssemblyCostCalcLogVO,
)
from app.services.costing.assembly_cost.calc_report.builder import (
    CalcReportWorkbookBuilder,
)
from app.services.costing.assembly_cost.calc_report.match import (
    match_calc_report_card,
)


def _card(**overrides: object) -> AssemblyCostCalcCardTraceVO:
    payload = {
        "part_number": "5010001AA",
        "part_name": "总成",
        "level": 0,
        "is_tree_root": True,
        "object_kind": None,
        "status": "ok",
        "should_cost": Decimal("12.3400"),
        "grand_total": Decimal("13.0000"),
        "material_amount": Decimal("10.0000"),
        "welding_amount": Decimal("2.3400"),
        "missing_leaf_count": 0,
        "message": "核算成功",
        "missing_parts": [],
    }
    payload.update(overrides)
    return AssemblyCostCalcCardTraceVO.model_validate(payload)


def _log(cards: list[AssemblyCostCalcCardTraceVO]) -> AssemblyCostCalcLogVO:
    return AssemblyCostCalcLogVO.model_validate(
        {
            "id": 7,
            "year": 2026,
            "month": 9,
            "location": "芜湖",
            "vehicle_model": "T13J",
            "bom_version": "D/09",
            "calc_batch_id": "abc123",
            "status": "partial",
            "message": "部分成功：1 张卡缺价",
            "subject_total": 2,
            "ok_count": 1,
            "partial_count": 1,
            "failed_count": 0,
            "skipped_approved_count": 0,
            "detail_text": "2026-09  芜湖/T13J/D/09\n完成。",
            "cards": cards,
            "started_at": datetime(2026, 9, 2, 9, 0, 0),
            "finished_at": datetime(2026, 9, 2, 9, 12, 0),
            "operator_name": "张三",
        }
    )


def test_match_calc_report_card_filters() -> None:
    root = _card()
    child = _card(
        part_number="5010002AA",
        part_name="分总成",
        level=1,
        is_tree_root=False,
        status="partial",
        missing_parts=["P-MISS（缺价）"],
        message="缺价 1",
    )
    assert match_calc_report_card(root, status="ok")
    assert not match_calc_report_card(root, status="failed")
    assert match_calc_report_card(child, level=1)
    assert not match_calc_report_card(root, level=1)
    assert match_calc_report_card(root, tree_root="yes")
    assert not match_calc_report_card(child, tree_root="yes")
    assert match_calc_report_card(child, keyword="p-miss")
    assert not match_calc_report_card(root, keyword="没有")


def test_write_calc_report_workbook_sheets() -> None:
    cards = [
        _card(),
        _card(
            part_number="5010002AA",
            part_name="分总成",
            level=1,
            is_tree_root=False,
            status="partial",
            missing_leaf_count=1,
            missing_parts=["P-MISS（缺价）"],
        ),
    ]
    log = _log(cards)
    wb = Workbook()
    CalcReportWorkbookBuilder().write(
        wb,
        log,
        [cards[1]],
        status="partial",
        tree_root="no",
    )
    buffer = BytesIO()
    wb.save(buffer)
    loaded = load_workbook(BytesIO(buffer.getvalue()))
    assert [sheet.title for sheet in loaded.worksheets] == ["封面摘要", "核价卡", "过程说明"]
    cover = loaded["封面摘要"]
    assert cover["A1"].value == "总成核价计算报告"
    assert "芜湖" in str(cover["B4"].value)
    cards_sheet = loaded["核价卡"]
    assert cards_sheet["D4"].value == "零件号"
    assert cards_sheet["E4"].value == "短码"
    assert cards_sheet["D5"].value == "5010002AA"
    assert cards_sheet["E5"].value == "0002"
    assert cards_sheet["H5"].value == "部分成功"
    assert cards_sheet["L5"].value == Decimal("13.0000")
    assert "P-MISS" in str(cards_sheet["O5"].value)
    assert "芜湖/T13J" in str(loaded["过程说明"]["A3"].value)
