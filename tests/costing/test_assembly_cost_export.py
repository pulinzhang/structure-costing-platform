"""成品核价导出：零件短码列。"""

from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from app.services.costing.assembly_cost.codes import (
    item_short_code,
    part_short_code,
)
from app.services.costing.assembly_cost.export.assembly_cost_export_service import (
    AssemblyCostExportService,
)
from app.services.costing.assembly_cost.export.workbook.builder import (
    collect_welding_fee_ids,
)
from app.services.costing.assembly_cost.export.workbook.columns import (
    LIST_HEADERS,
    LIST_KPI_COL,
    LIST_MONEY_COLS,
    PROCUREMENT_FILE_PREFIX,
    PROCUREMENT_LIST_HEADERS,
    PROCUREMENT_LIST_KPI_COL,
    PROCUREMENT_LIST_MONEY_COLS,
    PROCUREMENT_LIST_SHEET,
)
from app.services.costing.assembly_cost.export.workbook.labels import (
    confirm_status_label,
    cost_source_label,
    match_hint_label,
    object_kind_label,
    tree_level_label,
)


def test_short_code_column_after_part_number() -> None:
    assert LIST_HEADERS[4] == "零件号"
    assert LIST_HEADERS[5] == "短码"
    assert LIST_KPI_COL == 14
    assert {10, 11, 12, 13, 14} == LIST_MONEY_COLS


def test_procurement_list_headers() -> None:
    assert PROCUREMENT_LIST_HEADERS[1] == "短码"
    assert PROCUREMENT_LIST_HEADERS[6] == "材料成本"
    assert PROCUREMENT_LIST_HEADERS[7] == "焊接费/电泳黑漆等"
    assert PROCUREMENT_LIST_HEADERS[8] == "总成本"
    assert PROCUREMENT_LIST_KPI_COL == 9
    assert {7, 8, 9} == PROCUREMENT_LIST_MONEY_COLS
    assert "应成本" not in PROCUREMENT_LIST_HEADERS
    assert "报价" not in PROCUREMENT_LIST_HEADERS
    assert "价差" not in PROCUREMENT_LIST_HEADERS
    assert "年度" not in PROCUREMENT_LIST_HEADERS
    assert "基地" not in PROCUREMENT_LIST_HEADERS


def test_part_short_code_by_level() -> None:
    assert part_short_code(part_number="501029961AA", level=0) == "9961"
    assert part_short_code(part_number="501005123AA-1", level=1) == "5123AA-1"
    assert part_short_code(part_number="501005123AA", level=1) == "5123"
    assert part_short_code(part_number=None, level=0) == ""


def test_item_short_code_only_for_stamping_source() -> None:
    assert (
        item_short_code(
            part_number="501000011AA",
            level=1,
            cost_source="stamping",
        )
        == "0011"
    )
    assert (
        item_short_code(
            part_number="FQ371B06F36",
            level=2,
            cost_source="annual",
        )
        == ""
    )
    assert part_short_code(part_number="FQ371B06F36", level=2) == "0636"


def test_procurement_export_uses_chinese_labels() -> None:
    assert object_kind_label("l1_stamping") == "冲压件"
    assert object_kind_label("l1_subassembly") == "分总成"
    assert object_kind_label("l0_degenerate") == "退化供货件"
    assert tree_level_label(1) == "第1层"
    assert cost_source_label("stamping") == "冲压零件价"
    assert cost_source_label("skipped_by_bundle") == "一口价未计价"
    assert match_hint_label("bundle_parent:501000001AA") == "一口价父件 501000001AA"
    assert match_hint_label("multi_supplier_unit_cost") == "多家工厂单价不一致"
    assert confirm_status_label("outsourced") == "外发"


def test_collect_welding_fee_ids_dedupes_and_skips_empty() -> None:
    records = [
        SimpleNamespace(welding_fee_id=10),
        SimpleNamespace(welding_fee_id=None),
        SimpleNamespace(welding_fee_id=10),
        SimpleNamespace(welding_fee_id=7),
    ]
    assert collect_welding_fee_ids(records) == [10, 7]


def test_write_to_workbook_passes_welding_ids_not_level() -> None:
    service = AssemblyCostExportService()
    record = SimpleNamespace(
        welding_fee_id=88,
        procurement_object_id=None,
        items=[],
        indirects=[],
        year=2025,
        month=12,
        location="青岛",
        vehicle_model="T2AA",
        bom_version="D/09",
        part_number="501000001AA",
        part_name="总成",
        level=0,
        material_amount=0,
        welding_amount=0,
        total_amount=0,
        indirect_amount=0,
        grand_total=0,
        should_cost=0,
        vendor_quote=None,
        gap_amount=None,
        object_kind=None,
        calc_status="ok",
        audit_status="pending",
    )
    service.builder.repository.list_for_export = MagicMock(return_value=[record])
    captured: dict[str, object] = {}

    class _FakeWelding:
        def write_to_workbook(self, *args: object, **kwargs: object) -> tuple[str, int]:
            captured.update(kwargs)
            return "成品焊接费", 1

    with patch(
        "app.services.bom.welding_fee.export_service.BomWeldingFeeExportService",
        _FakeWelding,
    ):
        from openpyxl import Workbook

        service.write_to_workbook(
            Workbook(),
            MagicMock(),
            year=2025,
            month=12,
            location="青岛",
            vehicle_model="T2AA",
            bom_version="D/09",
            view="finance",
        )

    assert captured.get("ids") == [88]
    assert captured.get("omit_scope_columns") is False
    assert "level" not in captured


def test_procurement_workbook_skips_indirect_sheet() -> None:
    service = AssemblyCostExportService()
    record = SimpleNamespace(
        welding_fee_id=3,
        procurement_object_id=9,
        items=[],
        indirects=[],
        year=2025,
        month=12,
        location="青岛",
        vehicle_model="T2AA",
        bom_version="D/09",
        part_number="501000002AA",
        part_name="分总成",
        level=1,
        material_amount=0,
        welding_amount=0,
        total_amount=0,
        indirect_amount=0,
        grand_total=0,
        should_cost=0,
        vendor_quote=None,
        gap_amount=None,
        object_kind="l1_subassembly",
        calc_status="ok",
        audit_status="pending",
    )
    service.builder.repository.list_for_export = MagicMock(return_value=[record])

    class _FakeWelding:
        def write_to_workbook(self, *args: object, **kwargs: object) -> tuple[str, int]:
            return "采购焊接费", 1

    with (
        patch(
            "app.services.bom.welding_fee.export_service.BomWeldingFeeExportService",
            _FakeWelding,
        ),
        patch(
            "app.services.costing.assembly_cost.export.workbook.builder.BomProcurementObjectRepository"
        ) as repo_cls,
    ):
        repo_cls.return_value.map_confirm_status_by_ids.return_value = {9: "pending"}
        from openpyxl import Workbook

        wb = Workbook()
        service.write_to_workbook(
            wb,
            MagicMock(),
            year=2025,
            month=12,
            location="青岛",
            vehicle_model="T2AA",
            bom_version="D/09",
            view="procurement",
        )

    assert "成品间接费用" not in wb.sheetnames
    assert PROCUREMENT_LIST_SHEET in wb.sheetnames
    list_sheet = wb[PROCUREMENT_LIST_SHEET]
    assert list_sheet.cell(row=4, column=8).value == "焊接费/电泳黑漆等"
    assert list_sheet.cell(row=5, column=8).value == ""


def test_export_procurement_filename_prefix() -> None:
    service = AssemblyCostExportService()
    service.write_to_workbook = MagicMock(return_value=(PROCUREMENT_LIST_SHEET, 1))
    with patch(
        "app.services.costing.assembly_cost.export.assembly_cost_export_service.workbook_to_xlsx_bytes",
        return_value=b"xlsx",
    ):
        _content, filename = service.export(
            MagicMock(),
            year=2025,
            month=12,
            location="青岛",
            vehicle_model="T2AA",
            bom_version="D/09",
            view="procurement",
            current_user=MagicMock(),
        )
    assert filename.startswith(f"{PROCUREMENT_FILE_PREFIX}_")
    assert "2025年12月" in filename
    assert "应成本" not in filename


def _procurement_record_with_items() -> SimpleNamespace:
    return SimpleNamespace(
        welding_fee_id=3,
        procurement_object_id=9,
        items=[
            SimpleNamespace(
                node_level=1,
                part_number="501000011AA",
                part_name="内板",
                cost_source="stamping",
                match_hint=None,
                unit_price=1,
                effective_qty=2,
                line_amount=2,
            ),
            SimpleNamespace(
                node_level=1,
                part_number="FQ371B06F36",
                part_name="标准件",
                cost_source="annual",
                match_hint=None,
                unit_price=3,
                effective_qty=1,
                line_amount=3,
            ),
        ],
        indirects=[],
        year=2025,
        month=12,
        location="青岛",
        vehicle_model="T2AA",
        bom_version="D/09",
        part_number="501000002AA",
        part_name="分总成",
        level=1,
        material_amount=0,
        welding_amount=0,
        total_amount=0,
        indirect_amount=0,
        grand_total=0,
        should_cost=0,
        vendor_quote=None,
        gap_amount=None,
        object_kind="l1_subassembly",
        calc_status="ok",
        audit_status="pending",
    )


def test_procurement_items_sheet_merges_object_identity() -> None:
    service = AssemblyCostExportService()
    service.builder.repository.list_for_export = MagicMock(
        return_value=[_procurement_record_with_items()]
    )

    class _FakeWelding:
        def write_to_workbook(self, *args: object, **kwargs: object) -> tuple[str, int]:
            return "采购焊接费", 1

    with (
        patch(
            "app.services.bom.welding_fee.export_service.BomWeldingFeeExportService",
            _FakeWelding,
        ),
        patch(
            "app.services.costing.assembly_cost.export.workbook.builder.BomProcurementObjectRepository"
        ) as repo_cls,
    ):
        repo_cls.return_value.map_confirm_status_by_ids.return_value = {9: "pending"}
        from openpyxl import Workbook

        wb = Workbook()
        service.write_to_workbook(
            wb,
            MagicMock(),
            year=2025,
            month=12,
            location="青岛",
            vehicle_model="T2AA",
            bom_version="D/09",
            view="procurement",
        )

    sheet = wb["材料明细"]
    assert sheet.cell(row=4, column=1).value == "采购对象"
    assert sheet.cell(row=4, column=5).value == "子件构成（佐证）"
    assert sheet.cell(row=5, column=1).value == "序号"
    assert "年度" not in [sheet.cell(row=5, column=col).value for col in range(1, 14)]
    merged = {str(item) for item in sheet.merged_cells.ranges}
    assert "A6:A7" in merged
    assert "C6:C7" in merged
    assert "D6:D7" in merged
    assert "G6:G7" not in merged
    assert sheet.cell(row=6, column=1).value == 1
    assert sheet.cell(row=6, column=3).value == "501000002AA"
    assert sheet.cell(row=6, column=7).value == "501000011AA"
    assert sheet.cell(row=6, column=6).value == "0011"
    assert sheet.cell(row=7, column=7).value == "FQ371B06F36"
    assert sheet.cell(row=7, column=6).value == ""


def test_procurement_export_requires_period_scope() -> None:
    from app.exceptions.base import BusinessException

    service = AssemblyCostExportService()
    try:
        service.export(
            MagicMock(),
            year=2025,
            month=12,
            location="青岛",
            vehicle_model="T2AA",
            bom_version=None,
            view="procurement",
            current_user=MagicMock(),
        )
    except BusinessException as exc:
        assert "年度" in str(exc) or "版本" in str(exc)
    else:
        raise AssertionError("expected BusinessException")
