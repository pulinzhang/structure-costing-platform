"""仅价格列 + 附页钢厂映射。"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from app.parsers.yishiyiyi.base_price import parse_base_price_excel
from app.parsers.yishiyiyi.base_price.mill_sheet import extract_grade_mill_map
from app.parsers.yishiyiyi.base_price.xls_compat import open_workbook


def _price_only_workbook() -> bytes:
    workbook = Workbook()
    sheet1 = workbook.active
    sheet1.title = "Sheet1"
    sheet1["A4"] = "序号"
    sheet1["B4"] = "钢种"
    sheet1["C4"] = "牌号"
    sheet1["D4"] = "参考规格"
    sheet1["E4"] = "芜湖"
    sheet1["F4"] = "大连基地"
    sheet1["E5"] = "到厂原卷承兑价格"
    sheet1["F5"] = "到厂原卷承兑价格"
    sheet1["A6"] = 1
    sheet1["B6"] = "镀锌"
    sheet1["C6"] = "DC51D+Z"
    sheet1["D6"] = "1.0*1250"
    sheet1["E6"] = 4096
    sheet1["F6"] = 4090

    sheet2 = workbook.create_sheet("Sheet2")
    sheet2["A1"] = "2月订单芜湖基地钢板集采牌号对应钢厂如下："
    sheet2["A2"] = "马钢:"
    sheet2["A3"] = "DC51D+Z、DC53D+Z"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_extract_grade_mill_map_from_sheet2() -> None:
    content = _price_only_workbook()
    workbook = open_workbook(content)
    location_hint, mapping = extract_grade_mill_map(workbook)
    assert location_hint == "芜湖"
    assert mapping["DC51D+Z"] == "马钢"


def test_price_only_sheet2_mill_applies_to_all_sheet1_locations() -> None:
    content = _price_only_workbook()
    lookup = {"芜湖": "芜湖", "大连基地": "大连", "大连": "大连"}
    result = parse_base_price_excel(
        content,
        location_lookup=lookup,
        carry_mills={("大连", "DC51D+Z"): "山钢"},
    )
    assert not result.errors
    assert any("基地以主表列头为准" in warning for warning in result.warnings)
    by_location = {row["location"]: row for row in result.records}
    assert by_location["芜湖"]["steel_mill"] == "马钢"
    assert int(by_location["芜湖"]["base_price"]) == 4096
    assert by_location["大连"]["steel_mill"] == "马钢"
    assert int(by_location["大连"]["base_price"]) == 4090


def test_carry_mill_only_when_grade_missing_from_sheet2() -> None:
    workbook = Workbook()
    sheet1 = workbook.active
    sheet1.title = "Sheet1"
    sheet1["A4"] = "序号"
    sheet1["B4"] = "钢种"
    sheet1["C4"] = "牌号"
    sheet1["D4"] = "参考规格"
    sheet1["E4"] = "青岛基地"
    sheet1["E5"] = "到厂原卷承兑价格"
    sheet1["A6"] = 1
    sheet1["B6"] = "镀锌"
    sheet1["C6"] = "DC51D+Z"
    sheet1["D6"] = "1.0*1250"
    sheet1["E6"] = 3880
    sheet1["A7"] = 2
    sheet1["B7"] = "冷轧"
    sheet1["C7"] = "DC01"
    sheet1["D7"] = "1.0*1250"
    sheet1["E7"] = 3500
    sheet2 = workbook.create_sheet("Sheet2")
    sheet2["A1"] = "2月订单芜湖基地钢板集采牌号对应钢厂如下："
    sheet2["A2"] = "马钢:"
    sheet2["A3"] = "DC51D+Z"
    buffer = BytesIO()
    workbook.save(buffer)
    result = parse_base_price_excel(
        buffer.getvalue(),
        location_lookup={"青岛基地": "青岛"},
        carry_mills={("青岛", "DC01"): "本钢"},
    )
    by_grade = {row["grade"]: row for row in result.records}
    assert by_grade["DC51D+Z"]["location"] == "青岛"
    assert by_grade["DC51D+Z"]["steel_mill"] == "马钢"
    assert by_grade["DC01"]["steel_mill"] == "本钢"


def test_swapped_mill_price_cells() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A4"] = "序号"
    sheet["B4"] = "钢种"
    sheet["C4"] = "牌号"
    sheet["D4"] = "参考规格"
    sheet["E4"] = "芜湖"
    sheet["F4"] = "芜湖"
    sheet["E5"] = "中标钢厂"
    sheet["F5"] = "到厂原卷承兑价格"
    sheet["A6"] = 1
    sheet["B6"] = "镀锌"
    sheet["C6"] = "DC51D+Z"
    sheet["D6"] = "1.0*1250"
    sheet["E6"] = 4049
    sheet["F6"] = "VAMA"
    buffer = BytesIO()
    workbook.save(buffer)
    result = parse_base_price_excel(
        buffer.getvalue(),
        location_lookup={"芜湖": "芜湖"},
    )
    assert len(result.records) == 1
    assert result.records[0]["steel_mill"] == "VAMA"
    assert int(result.records[0]["base_price"]) == 4049


def test_mixed_price_only_and_mill_columns() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A4"] = "序号"
    sheet["B4"] = "钢种"
    sheet["C4"] = "牌号"
    sheet["D4"] = "参考规格"
    sheet["E4"] = "芜湖"
    sheet["F4"] = "青岛基地"
    sheet["E5"] = "到厂原卷承兑价格"
    sheet["F5"] = "钢厂"
    sheet["G5"] = "到厂原卷承兑价格"
    sheet["A6"] = 1
    sheet["B6"] = "镀锌"
    sheet["C6"] = "DC51D+Z"
    sheet["D6"] = "1.0*1250"
    sheet["E6"] = 4096
    sheet["F6"] = "山钢"
    sheet["G6"] = 4080
    sheet2 = workbook.create_sheet("Sheet2")
    sheet2["A1"] = "芜湖基地牌号对应钢厂如下："
    sheet2["A2"] = "马钢:"
    sheet2["A3"] = "DC51D+Z"
    buffer = BytesIO()
    workbook.save(buffer)
    result = parse_base_price_excel(
        buffer.getvalue(),
        location_lookup={"芜湖": "芜湖", "青岛基地": "青岛"},
    )
    by_location = {row["location"]: row for row in result.records}
    assert by_location["芜湖"]["steel_mill"] == "马钢"
    assert int(by_location["芜湖"]["base_price"]) == 4096
    assert by_location["青岛"]["steel_mill"] == "山钢"
    assert int(by_location["青岛"]["base_price"]) == 4080


def test_sample_price_only_file_resolves_wuhu_mill() -> None:
    root = (
        Path(__file__).resolve().parents[2]
        / "excel_sample"
        / "25年-26年一事一议价格"
        / "2025年1月-12月一事一议价格"
        / "25年2月订单3月冲焊件及废料价格"
    )
    path = root / "25年2月订单（25年3月原料）发布价格（奇瑞）.xlsx"
    if not path.is_file():
        return
    content = path.read_bytes()
    result = parse_base_price_excel(
        content,
        location_lookup={"芜湖": "芜湖", "大连基地": "大连", "青岛基地": "青岛"},
        carry_mills={
            ("大连", "DC51D+Z"): "山钢",
            ("青岛", "DC51D+Z"): "山钢",
        },
    )
    assert not result.errors
    assert result.records
    sample = next(
        row for row in result.records if row["grade"] == "DC51D+Z" and row["location"] == "芜湖"
    )
    assert sample["steel_mill"] == "马钢"


def test_june_sample_qingdao_uses_sheet2_mill_despite_wuhu_title() -> None:
    root = (
        Path(__file__).resolve().parents[2]
        / "excel_sample"
        / "25年-26年一事一议价格"
        / "2025年1月-12月一事一议价格"
        / "25年6月订单7月冲焊件及废料价格"
    )
    path = root / "25年6月订单（25年7月原料）发布价格.xlsx"
    if not path.is_file():
        return
    result = parse_base_price_excel(
        path.read_bytes(),
        location_lookup={"芜湖": "芜湖", "青岛基地": "青岛"},
    )
    assert not result.errors
    locations = {str(row["location"]) for row in result.records}
    assert locations == {"芜湖", "青岛"}
    wuhu = next(
        row for row in result.records if row["grade"] == "DC51D+Z" and row["location"] == "芜湖"
    )
    qingdao = next(
        row for row in result.records if row["grade"] == "DC51D+Z" and row["location"] == "青岛"
    )
    assert wuhu["steel_mill"] == "马钢"
    assert int(wuhu["base_price"]) == 4096
    assert qingdao["steel_mill"] == "马钢"
    assert int(qingdao["base_price"]) == 3880
