"""中标钢厂厚度价格一览 Excel：计算页 + 使用示例页。"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.schemas.vo.winning_thickness_vo import WinningThicknessRowVO
from app.services.yishiyiyi.winning_thickness.export.excel_exporter import (
    WinningThicknessExporter,
)
from app.services.yishiyiyi.winning_thickness.export.values import (
    CALC_SHEET_SUFFIX,
    EXAMPLE_SHEET_SUFFIX,
    OVERVIEW_SHEET_SUFFIX,
    scoped_sheet_name,
)

_HEADER_ROW = 6
_DATA_START = 7


def _row(
    *,
    grade: str,
    steel_mill: str,
    category_name: str = "冷轧",
    location: str = "芜湖",
    thickness_range: str = "0.60 ≤ t < 0.80",
    thick_min: Decimal | None = Decimal("0.60"),
    thick_min_operator: str | None = ">=",
    thick_max: Decimal | None = Decimal("0.80"),
    thick_max_operator: str | None = "<",
    base_price: Decimal = Decimal("3700"),
    adjustment: Decimal = Decimal("160"),
    processing_fee: Decimal = Decimal("120"),
) -> WinningThicknessRowVO:
    return WinningThicknessRowVO(
        base_price=base_price,
        thickness_adjustment=adjustment,
        processing_fee=processing_fee,
        final_price=base_price + adjustment + processing_fee,
        category_name=category_name,
        steel_mill=steel_mill,
        thickness_range=thickness_range,
        thick_min=thick_min,
        thick_min_operator=thick_min_operator,
        thick_max=thick_max,
        thick_max_operator=thick_max_operator,
        grade=grade,
        year=2026,
        month=1,
        location=location,
        base_price_id=11,
        thickness_policy_id=22,
    )


def test_export_has_overview_calc_and_example_sheets() -> None:
    content, filename = WinningThicknessExporter().generate(
        [
            _row(grade="DC01", steel_mill="马钢"),
            _row(
                grade="DC04",
                steel_mill="宝钢",
                thick_min=None,
                thick_min_operator=None,
                thick_max=None,
                thick_max_operator=None,
                thickness_range="-",
                adjustment=Decimal("0"),
            ),
        ],
        year=2026,
        month=1,
        location="芜湖",
    )

    assert filename == "中标钢厂厚度价格一览表_2026年_1月_芜湖.xlsx"
    workbook = load_workbook(BytesIO(content))
    overview_name = scoped_sheet_name(
        year=2026, month=1, location="芜湖", suffix=OVERVIEW_SHEET_SUFFIX
    )
    calc_name = scoped_sheet_name(year=2026, month=1, location="芜湖", suffix=CALC_SHEET_SUFFIX)
    example_name = scoped_sheet_name(
        year=2026, month=1, location="芜湖", suffix=EXAMPLE_SHEET_SUFFIX
    )
    assert workbook.sheetnames == [overview_name, calc_name, example_name]

    overview = workbook[overview_name]
    assert overview.cell(1, 1).value == "中标钢厂厚度价格一览表"
    assert overview.cell(3, 2).value == "芜湖"
    assert overview.cell(3, 4).value == "2026年1月"
    assert overview.cell(3, 8).value == "2"
    assert "最终价 = 基准价 + 厚度加减价 + 加工费" in str(overview.cell(4, 1).value)
    assert overview.cell(_HEADER_ROW, 9).value == "基价溯源"
    assert overview.cell(_DATA_START, 2).value == "DC01"
    assert Decimal(str(overview.cell(_DATA_START, 8).value)) == Decimal("3980")

    calc = workbook[calc_name]
    assert calc.cell(1, 1).value == "中标钢厂厚度价格一览表"
    assert calc.cell(_HEADER_ROW, 1).value == "基地"
    assert calc.cell(_HEADER_ROW, 2).value == "钢种"
    assert calc.cell(_HEADER_ROW, 9).value == "厚度范围"
    assert calc.cell(_HEADER_ROW, 13).value == "最终价(元/吨)"
    header_text = "".join(str(calc.cell(_HEADER_ROW, col).value or "") for col in range(1, 14))
    assert "溯源" not in header_text
    assert calc.column_dimensions["E"].hidden is True
    assert calc.column_dimensions["F"].hidden is True
    assert calc.column_dimensions["G"].hidden is True
    assert calc.column_dimensions["H"].hidden is True
    assert calc.column_dimensions["I"].hidden is not True
    assert calc.cell(_DATA_START, 1).value == "芜湖"
    assert calc.cell(_DATA_START, 2).value == "冷轧"
    assert calc.cell(_DATA_START, 3).value == "DC01"
    assert calc.cell(_DATA_START, 4).value == "马钢"
    assert Decimal(str(calc.cell(_DATA_START, 5).value)) == Decimal("0.60")
    assert calc.cell(_DATA_START, 6).value == ">="
    assert Decimal(str(calc.cell(_DATA_START, 7).value)) == Decimal("0.80")
    assert calc.cell(_DATA_START, 8).value == "<"
    assert calc.cell(_DATA_START, 9).value == "0.60 ≤ t < 0.80"
    assert calc.cell(_DATA_START, 13).value == f"=J{_DATA_START}+K{_DATA_START}+L{_DATA_START}"
    second = _DATA_START + 1
    assert calc.cell(second, 5).value is None
    assert calc.cell(second, 6).value is None
    assert calc.cell(second, 7).value is None
    assert calc.cell(second, 8).value is None
    assert calc.cell(second, 9).value == "-"
    assert calc.cell(second, 13).value == f"=J{second}+K{second}+L{second}"

    example = workbook[example_name]
    assert example.cell(1, 1).value == "中标钢厂厚度价格一览表"
    assert example.cell(_HEADER_ROW, 1).value == "基地"
    assert example.cell(_HEADER_ROW, 5).value == "厚度"
    assert example.cell(_HEADER_ROW, 6).value == "计算页行号"
    assert example.cell(_HEADER_ROW, 10).value == "最终价(元/吨)"
    assert example.cell(_DATA_START, 3).value == "DC01"
    calc_row_lookup = str(example.cell(_DATA_START, 6).value)
    assert calc_row_lookup.startswith("=IF($E7=")
    assert "ROW(" in calc_row_lookup
    assert f"'{calc_name}'!$A$" in calc_row_lookup
    lookup = str(example.cell(_DATA_START, 7).value)
    assert lookup.startswith("=IF($E7=")
    assert f"'{calc_name}'!$A$" in lookup
    assert f"'{calc_name}'!$C$" in lookup
    assert f"'{calc_name}'!$J$" in lookup
    assert f"'{calc_name}'!$B$" not in lookup
    assert f"'{calc_name}'!$D$" not in lookup
    assert f"'{calc_name}'!$I$" not in lookup
    assert str(example.cell(_DATA_START, 10).value).startswith("=IF(G7=")
    formula_notes = " ".join(
        str(example.cell(row, 1).value or "") for row in range(1, example.max_row + 1)
    )
    assert "公式说明" in formula_notes
    assert "计算页行号" in formula_notes
    assert "SUMPRODUCT" in formula_notes
    assert "下限" in formula_notes
    assert "最终价 = 基准价 + 加减价 + 加工费" in formula_notes
    workbook.close()


def test_example_sheet_without_rows_still_explains_formula() -> None:
    content, _filename = WinningThicknessExporter().generate(
        [],
        year=2026,
        month=8,
        location="青岛",
    )
    workbook = load_workbook(BytesIO(content))
    example_name = scoped_sheet_name(
        year=2026, month=8, location="青岛", suffix=EXAMPLE_SHEET_SUFFIX
    )
    calc_name = scoped_sheet_name(year=2026, month=8, location="青岛", suffix=CALC_SHEET_SUFFIX)
    example = workbook[example_name]
    joined = " ".join(
        str(example.cell(row, 1).value or "") for row in range(1, example.max_row + 1)
    )
    assert example.cell(1, 1).value == "中标钢厂厚度价格一览表"
    assert example.cell(3, 2).value == "青岛"
    assert example.cell(3, 8).value == "0"
    assert "本次无数据可摘录" in joined
    assert "公式说明" in joined
    assert workbook[calc_name].cell(_HEADER_ROW, 13).value == "最终价(元/吨)"
    workbook.close()
