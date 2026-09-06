"""厚度末档填「/」或「无」时沿用上一阶梯加减价。"""

from __future__ import annotations

from openpyxl import Workbook

from app.parsers.yishiyiyi.policy.handlers.base_handler.policy_assembler import (
    PolicyAssembler,
)
from app.parsers.yishiyiyi.policy.models.policy import POLICY_TYPE_THICKNESS
from app.parsers.yishiyiyi.policy.value_parser import is_carry_forward_price


def _build_row(last_value: object) -> list[int]:
    sheet = Workbook().active
    assert sheet is not None
    sheet.cell(row=1, column=1, value="钢种")
    sheet.cell(row=1, column=2, value=0.5)
    sheet.cell(row=1, column=3, value="≥0.6")
    sheet.cell(row=1, column=4, value="≥1.0")
    sheet.cell(row=2, column=1, value="DC01")
    sheet.cell(row=2, column=2, value=100)
    sheet.cell(row=2, column=3, value=50)
    sheet.cell(row=2, column=4, value=last_value)

    assembler = PolicyAssembler()
    price_cols = [2, 3, 4]
    boundaries = assembler.collect_boundaries_thickness(sheet, 1, price_cols)
    items, _ = assembler.build_items_for_row(
        sheet,
        2,
        price_cols,
        boundaries,
        POLICY_TYPE_THICKNESS,
        sort_order=2,
    )
    return [item.price_adjustment for item in items]


def test_is_carry_forward_price_markers() -> None:
    assert is_carry_forward_price("/")
    assert is_carry_forward_price("／")
    assert is_carry_forward_price("无")
    assert is_carry_forward_price(" / ")
    assert not is_carry_forward_price(0)
    assert not is_carry_forward_price("0")
    assert not is_carry_forward_price(-20)
    assert not is_carry_forward_price(None)


def test_last_ladder_slash_inherits_previous_price() -> None:
    assert _build_row("/") == [100, 50, 50]


def test_last_ladder_none_text_inherits_previous_price() -> None:
    assert _build_row("无") == [100, 50, 50]


def test_last_ladder_fullwidth_slash_inherits_previous_price() -> None:
    assert _build_row("／") == [100, 50, 50]


def test_last_ladder_explicit_zero_stays_zero() -> None:
    assert _build_row(0) == [100, 50, 0]


def test_last_ladder_numeric_keeps_value() -> None:
    assert _build_row(-20) == [100, 50, -20]
