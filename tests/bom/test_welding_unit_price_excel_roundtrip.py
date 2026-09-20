"""焊接单价导出与导入：核心列名称/单位/费用/备注；列表导出另含基地/年份/审核列。"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

from openpyxl import load_workbook

from app.common.audit import SYSTEM_USER
from app.parsers.bom.welding_unit_price_parser import parse_welding_unit_price_excel
from app.services.bom.welding_unit_price.export.service import BomWeldingUnitPriceExportService
from app.services.bom.welding_unit_price.import_service import BomWeldingUnitPriceImportService


def test_import_template_is_parseable() -> None:
    content, filename = BomWeldingUnitPriceImportService().build_template()
    assert filename.endswith(".xlsx")

    result = parse_welding_unit_price_excel(content)
    assert not result.errors
    assert len(result.rows) >= 1
    assert result.rows[0].item_name == "焊点"
    assert result.rows[0].unit == "个"


def test_export_includes_scope_audit_and_roundtrips() -> None:
    records = [
        SimpleNamespace(
            location="芜湖",
            year=2026,
            item_name="焊点",
            unit="个",
            unit_price=Decimal("0.0999"),
            remark="",
            audit_status="approved",
            created_by="张三",
            created_at=None,
            updated_by="李四",
            updated_at=None,
        ),
        SimpleNamespace(
            location="芜湖",
            year=2026,
            item_name="CO2焊缝",
            unit="CM",
            unit_price=Decimal("0.0700"),
            remark="样例",
            audit_status="pending",
            created_by="张三",
            created_at=None,
            updated_by=None,
            updated_at=None,
        ),
    ]
    service = BomWeldingUnitPriceExportService()
    service.repository = MagicMock()
    service.repository.list_for_export.return_value = records

    content, filename = service.export(
        MagicMock(),
        location="芜湖",
        year=2026,
        current_user=SYSTEM_USER,
    )
    assert "焊接单价" in filename
    assert "芜湖" in filename

    workbook = load_workbook(BytesIO(content), data_only=True)
    sheet = workbook.worksheets[0]
    assert sheet.title == "焊接单价"
    headers = [sheet.cell(1, col).value for col in range(1, 12)]
    assert headers == [
        "基地",
        "年份",
        "名称",
        "单位",
        "费用",
        "备注",
        "审核状态",
        "创建人",
        "创建时间",
        "更新人",
        "更新时间",
    ]
    assert sheet.cell(2, 1).value == "芜湖"
    assert sheet.cell(2, 3).value == "焊点"
    assert sheet.cell(2, 5).value == Decimal("0.0999") or float(sheet.cell(2, 5).value) == 0.0999
    assert sheet.cell(2, 7).value == "已通过"
    workbook.close()

    result = parse_welding_unit_price_excel(content)
    assert not result.errors
    assert [(row.item_name, row.unit, row.unit_price, row.remark) for row in result.rows] == [
        ("焊点", "个", Decimal("0.0999"), None),
        ("CO2焊缝", "CM", Decimal("0.0700"), "样例"),
    ]
