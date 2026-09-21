"""年度零件价导出与导入：核心列与模板一致；列表导出另含基地/年份/审核列。"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

from openpyxl import load_workbook

from app.common.audit import SYSTEM_USER
from app.parsers.bom.annual_part_price_parser import parse_annual_part_price_excel
from app.services.bom.annual_part_price.export.service import BomAnnualPartPriceExportService


def test_export_includes_scope_audit_and_roundtrips() -> None:
    records = [
        SimpleNamespace(
            location="芜湖",
            year=2026,
            part_number="N0111941",
            part_name="六角法兰面螺母",
            specification="",
            purchase_unit="pcs",
            purchase_unit_price=None,
            conversion_factor=Decimal("1"),
            unit="pcs",
            unit_price=Decimal("0.1200"),
            supply_type="标准件",
            supplier_name="",
            remark="",
            audit_status="approved",
            created_by="张三",
            created_at=None,
            updated_by="李四",
            updated_at=None,
        ),
        SimpleNamespace(
            location="芜湖",
            year=2026,
            part_number="551001133AA",
            part_name="板料样例",
            specification="1.2",
            purchase_unit="pcs",
            purchase_unit_price=Decimal("12.0000"),
            conversion_factor=Decimal("1"),
            unit="pcs",
            unit_price=Decimal("12.0000"),
            supply_type="沿用",
            supplier_name="某供应商",
            remark="样例",
            audit_status="pending",
            created_by="张三",
            created_at=None,
            updated_by=None,
            updated_at=None,
        ),
    ]
    service = BomAnnualPartPriceExportService()
    service.repository = MagicMock()
    service.repository.list_for_export.return_value = records

    content, filename = service.export(
        MagicMock(),
        location="芜湖",
        year=2026,
        current_user=SYSTEM_USER,
    )
    assert "年度零件价" in filename
    assert "芜湖" in filename

    workbook = load_workbook(BytesIO(content), data_only=True)
    sheet = workbook.worksheets[0]
    assert sheet.title == "年度零件价"
    headers = [sheet.cell(1, col).value for col in range(1, 19)]
    assert headers[:3] == ["基地", "年份", "零件号"]
    assert headers[9] == "单价"
    assert headers[10] == "原供货类型"
    assert headers[13:] == ["审核状态", "创建人", "创建时间", "更新人", "更新时间"]
    assert sheet.cell(2, 1).value == "芜湖"
    assert sheet.cell(2, 3).value == "N0111941"
    assert sheet.cell(2, 14).value == "已通过"
    workbook.close()

    result = parse_annual_part_price_excel(content)
    assert not result.errors
    assert len(result.rows) == 2
    assert result.rows[0].part_number == "N0111941"
    assert result.rows[0].unit_price == Decimal("0.1200")
    assert result.rows[0].supply_type == "标准件"
    assert result.rows[1].part_number == "551001133AA"
    assert result.rows[1].unit_price == Decimal("12.0000")
    assert result.rows[1].purchase_unit_price == Decimal("12.0000")
    assert result.rows[1].remark == "样例"
