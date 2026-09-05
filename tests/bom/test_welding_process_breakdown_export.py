"""焊接工艺拆解 L0 核对导出：短代码列。"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from app.schemas.vo.user_vo import UserVO
from app.services.bom.welding_process_breakdown.export.columns import COL, HEADERS, MAIN_COLS
from app.services.bom.welding_process_breakdown.export.values import part_short_code
from app.services.bom.welding_process_breakdown.export.workbook import L0CheckWorkbookBuilder


def test_part_short_code_last_four_digits() -> None:
    assert part_short_code("501042816AA") == "2816"
    assert part_short_code(None) == ""


def test_short_code_column_sits_after_part_number() -> None:
    assert HEADERS[4] == "零件号"
    assert HEADERS[5] == "短代码"
    assert MAIN_COLS == 10
    assert COL["短代码"] == 6


def test_workbook_writes_short_code_on_l0_row() -> None:
    line = SimpleNamespace(
        line_no=1,
        process_token="焊缝",
        matched_span="焊缝：3条 68mm",
        standard_process_name="CO2焊缝",
        standard_process_unit="CM",
        process_qty=Decimal("3"),
        process_qty_unit="条",
        process_qty_secondary=Decimal("68"),
        process_qty_secondary_unit="mm",
        map_status="ready",
    )
    node = SimpleNamespace(
        location="青岛",
        vehicle_model="T2AA",
        bom_version="V1",
        part_number="501042816AA",
        part_name="后副车架左前固定加强板总成",
        weld_info_raw="焊缝：3条 68mm",
        process_line_count=1,
        parse_status="ok",
        lines=[line],
    )
    content, row_count = L0CheckWorkbookBuilder().build(
        [node],  # type: ignore[list-item]
        filter_text="L0 · 基地 青岛",
        current_user=UserVO(id=1, username="tester"),
    )
    assert row_count == 1
    ws = load_workbook(BytesIO(content)).worksheets[0]
    assert ws.cell(row=5, column=COL["短代码"]).value == "短代码"
    assert ws.cell(row=6, column=COL["零件号"]).value == "501042816AA"
    assert ws.cell(row=6, column=COL["短代码"]).value == "2816"
    assert ws.freeze_panes == "G6"
