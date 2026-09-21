"""年度零件价导入模板：基地+项目外购叶子（other_leaf ∪ 非工艺库）。"""

from __future__ import annotations

from collections.abc import Generator
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base
from app.exceptions.base import BusinessException
from app.models.bom.node import BomNode, BomSnapshot  # noqa: F401
from app.models.masterdata.location import Location, LocationAlias  # noqa: F401
from app.models.stamping.part import StampingPart
from app.parsers.bom.annual_part_price_parser import (
    HEADER_CONVERSION_FACTOR,
    HEADER_PURCHASE_UNIT,
    HEADER_PURCHASE_UNIT_PRICE,
    HEADER_SPECIFICATION,
    HEADER_SUPPLY_TYPE,
    HEADER_UNIT,
    HEADER_UNIT_PRICE,
    parse_annual_part_price_excel,
)
from app.repositories.bom.bom.node import BomNodeRepository
from app.repositories.bom.bom.snapshot import BomSnapshotRepository
from app.repositories.masterdata.location.location_repository import LocationRepository
from app.schemas.dto.bom_dto import BomNodeWrite, BomSnapshotWrite
from app.schemas.dto.location_dto import LocationAliasInsert, LocationInsert
from app.services.bom.annual_part_price.importing.service import BomAnnualPartPriceImportService
from app.services.bom.annual_part_price.template_source import (
    SUPPLY_TYPE_CARRYOVER,
    SUPPLY_TYPE_GLUE,
    SUPPLY_TYPE_STANDARD,
)


@pytest.fixture()
def db() -> Generator[Session, None, None]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    session = factory()
    repo = LocationRepository()
    repo.create(
        session,
        LocationInsert(
            location="芜湖",
            is_enabled=True,
            sort_order=1,
            created_by="system",
            updated_by="system",
        ),
    )
    repo.create_alias(session, LocationAliasInsert(alias="芜湖基地", location="芜湖"))
    session.flush()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def _create_snapshot(
    db: Session,
    *,
    location: str,
    vehicle_model: str,
    bom_version: str,
) -> int:
    snapshot = BomSnapshotRepository().create(
        db,
        BomSnapshotWrite(
            location=location,
            vehicle_model=vehicle_model,
            bom_version=bom_version,
            source_filename=f"{vehicle_model}.xlsx",
            created_by="tester",
            updated_by="tester",
        ),
    )
    return snapshot.id


def _add_process_library_part(
    db: Session,
    *,
    location: str,
    vehicle_model: str,
    part_number: str,
) -> None:
    db.add(
        StampingPart(
            location=location,
            vehicle_model=vehicle_model,
            version="V1",
            part_number=part_number,
            part_name=part_number,
            excel_filename="process.xlsx",
        )
    )


def _add_stamping_tree(db: Session, *, snapshot_id: int) -> None:
    node_repo = BomNodeRepository()
    l0 = node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot_id,
            level=0,
            part_number="L0-STAMP",
            node_role="supply",
            created_by="tester",
            updated_by="tester",
            is_supply_unit=True,
            excel_filename="t.xlsx",
            excel_row=8,
        ),
    )
    node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot_id,
            parent_id=l0.id,
            level=1,
            part_number="L1-STAMP",
            node_role="make",
            created_by="tester",
            updated_by="tester",
            part_name="加强板",
            thickness_mm=Decimal("0.8000"),
            thickness_text="T=0.8",
            excel_filename="t.xlsx",
            excel_row=9,
        ),
    )


def _add_other_leaf(
    db: Session,
    *,
    snapshot_id: int,
    l0_part_number: str,
    part_number: str,
    part_name: str,
    excel_row: int,
    thickness_text: str | None = None,
) -> None:
    node_repo = BomNodeRepository()
    l0 = node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot_id,
            level=0,
            part_number=l0_part_number,
            node_role="supply",
            created_by="tester",
            updated_by="tester",
            is_supply_unit=True,
            excel_filename="t.xlsx",
            excel_row=excel_row,
        ),
    )
    node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot_id,
            parent_id=l0.id,
            level=1,
            part_number=part_number,
            part_name=part_name,
            node_role="buy",
            created_by="tester",
            updated_by="tester",
            qty_unit="pcs",
            thickness_text=thickness_text,
            excel_filename="t.xlsx",
            excel_row=excel_row + 1,
        ),
    )


def _add_subassembly_with_l2(
    db: Session,
    *,
    snapshot_id: int,
) -> None:
    """L1 分总成 + L2 板料 + L2 螺母。"""
    node_repo = BomNodeRepository()
    l0 = node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot_id,
            level=0,
            part_number="L0-HINGE",
            node_role="supply",
            created_by="tester",
            updated_by="tester",
            is_supply_unit=True,
            excel_filename="t.xlsx",
            excel_row=40,
        ),
    )
    l1 = node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot_id,
            parent_id=l0.id,
            level=1,
            part_number="551001137AA",
            part_name="左后门上铰链安装板总成",
            node_role="make",
            created_by="tester",
            updated_by="tester",
            excel_filename="t.xlsx",
            excel_row=41,
        ),
    )
    node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot_id,
            parent_id=l1.id,
            level=2,
            part_number="551001133AA",
            part_name="左后门上铰链安装板",
            node_role="make",
            created_by="tester",
            updated_by="tester",
            thickness_mm=Decimal("1.2000"),
            thickness_text="T=1.2",
            excel_filename="t.xlsx",
            excel_row=42,
        ),
    )
    node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot_id,
            parent_id=l1.id,
            level=2,
            part_number="QR36410",
            part_name="焊接方螺母",
            node_role="buy",
            created_by="tester",
            updated_by="tester",
            qty_unit="pcs",
            excel_filename="t.xlsx",
            excel_row=43,
        ),
    )


def _header_col(sheet, title: str) -> int:
    for col in range(1, (sheet.max_column or 1) + 1):
        if sheet.cell(1, col).value == title:
            return col
    raise AssertionError(f"missing header {title}")


def _row_map(sheet, row_idx: int) -> dict[str, object]:
    return {
        str(sheet.cell(1, col).value): sheet.cell(row_idx, col).value
        for col in range(1, (sheet.max_column or 1) + 1)
        if sheet.cell(1, col).value
    }


def _download_template(db: Session, *, vehicle_model: str = "T13J"):
    return BomAnnualPartPriceImportService().build_template(
        db, location="芜湖", vehicle_model=vehicle_model, year=2026
    )


def test_template_rejected_without_project(db: Session) -> None:
    with pytest.raises(BusinessException, match="请选择项目"):
        BomAnnualPartPriceImportService().build_template(
            db, location="芜湖", vehicle_model="  ", year=2026
        )


def test_template_rejected_without_bom(db: Session) -> None:
    with pytest.raises(BusinessException, match="请先上传该基地该项目的材料清单"):
        _download_template(db)


def test_template_rejected_when_all_leaves_in_process_library(db: Session) -> None:
    snapshot_id = _create_snapshot(db, location="芜湖", vehicle_model="T13J", bom_version="V3")
    _add_stamping_tree(db, snapshot_id=snapshot_id)
    _add_process_library_part(db, location="芜湖", vehicle_model="T13J", part_number="L1-STAMP")
    db.commit()
    with pytest.raises(BusinessException, match="没有需填年度价的外购叶子"):
        _download_template(db)


def test_template_unions_project_versions_and_dedupes(db: Session) -> None:
    first_id = _create_snapshot(db, location="芜湖", vehicle_model="T13J", bom_version="V3")
    _add_other_leaf(
        db,
        snapshot_id=first_id,
        l0_part_number="L0-A",
        part_number="QR36410",
        part_name="焊接方螺母",
        excel_row=8,
        thickness_text="M6",
    )
    second_id = _create_snapshot(db, location="芜湖", vehicle_model="T13J", bom_version="V4")
    _add_other_leaf(
        db,
        snapshot_id=second_id,
        l0_part_number="L0-B",
        part_number="QR36410",
        part_name="焊接方螺母-重复",
        excel_row=20,
    )
    _add_other_leaf(
        db,
        snapshot_id=second_id,
        l0_part_number="L0-C",
        part_number="M14-4102311",
        part_name="结构胶",
        excel_row=30,
    )
    other_project = _create_snapshot(db, location="芜湖", vehicle_model="T28", bom_version="V1")
    _add_other_leaf(
        db,
        snapshot_id=other_project,
        l0_part_number="L0-D",
        part_number="T28-ONLY",
        part_name="他项目件",
        excel_row=50,
    )
    db.commit()

    content, filename = _download_template(db)
    assert filename == "年度零件价导入模板_芜湖_T13J_2026.xlsx"

    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook.sheetnames[0] == "年度零件价"
    assert workbook.sheetnames[1] == "填写说明"
    sheet = workbook.worksheets[0]
    numbers = [str(row[0]) for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]]
    assert numbers == ["QR36410", "M14-4102311"]
    assert "T28-ONLY" not in numbers
    row2 = _row_map(sheet, 2)
    row3 = _row_map(sheet, 3)
    assert row2["零件名称"] == "焊接方螺母"
    assert row2[HEADER_SPECIFICATION] == "M6"
    assert row2[HEADER_PURCHASE_UNIT] == "pcs"
    assert row2[HEADER_PURCHASE_UNIT_PRICE] in (None, "")
    assert row2[HEADER_CONVERSION_FACTOR] in (1, Decimal("1"), Decimal("1.00000000"))
    assert row2[HEADER_UNIT] == "pcs"
    assert row2[HEADER_UNIT_PRICE] == Decimal("0.0000") or row2[HEADER_UNIT_PRICE] == 0
    assert row2[HEADER_SUPPLY_TYPE] == SUPPLY_TYPE_STANDARD
    assert row3[HEADER_SUPPLY_TYPE] == SUPPLY_TYPE_GLUE
    notes_values = [workbook.worksheets[1].cell(row=idx, column=1).value for idx in range(3, 10)]
    notes_text = "\n".join(str(item) for item in notes_values if item)
    assert "核价单价 = 采购单价 ÷ 转换系数" in notes_text
    assert "须大于 0" in notes_text
    assert "不要填 1/12" in notes_text
    assert "不要把采购包装单位" in notes_text

    parsed = parse_annual_part_price_excel(content)
    assert parsed.errors == []
    assert [row.part_number for row in parsed.rows] == ["QR36410", "M14-4102311"]
    assert parsed.rows[0].supply_type == SUPPLY_TYPE_STANDARD
    assert parsed.rows[1].supply_type == SUPPLY_TYPE_GLUE
    assert parsed.rows[0].unit_price == Decimal("0.0000")
    assert parsed.rows[0].conversion_factor == 1
    assert parsed.rows[0].purchase_unit_price is None
    assert parsed.rows[0].specification == "M6"


def test_template_includes_l2_not_in_process_library(db: Session) -> None:
    snapshot_id = _create_snapshot(db, location="芜湖", vehicle_model="T13J", bom_version="D/09")
    _add_subassembly_with_l2(db, snapshot_id=snapshot_id)
    db.commit()

    content, _filename = _download_template(db)
    workbook = load_workbook(BytesIO(content), data_only=True)
    sheet = workbook.active
    numbers = [str(row[0]) for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]]
    assert numbers == ["551001133AA", "QR36410"]
    assert "551001137AA" not in numbers
    by_number = {}
    supply_col = _header_col(sheet, HEADER_SUPPLY_TYPE)
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        by_number[str(row[0])] = str(row[supply_col - 1])
    assert by_number["551001133AA"] == SUPPLY_TYPE_CARRYOVER
    assert by_number["QR36410"] == SUPPLY_TYPE_STANDARD


def test_template_excludes_l2_that_is_in_process_library(db: Session) -> None:
    snapshot_id = _create_snapshot(db, location="芜湖", vehicle_model="T13J", bom_version="D/09")
    _add_subassembly_with_l2(db, snapshot_id=snapshot_id)
    _add_process_library_part(db, location="芜湖", vehicle_model="T13J", part_number="551001133AA")
    db.commit()

    content, _filename = _download_template(db)
    workbook = load_workbook(BytesIO(content), data_only=True)
    sheet = workbook.active
    numbers = [str(row[0]) for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]]
    assert numbers == ["QR36410"]
    assert "551001133AA" not in numbers
    assert "551001137AA" not in numbers
