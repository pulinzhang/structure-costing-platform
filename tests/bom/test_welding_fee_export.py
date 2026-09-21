"""焊接费导出：零件短代码、表头汇总数与表尾 SUM。"""

from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import MagicMock

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.services.bom.welding_fee.export.columns import (
    COL,
    COL_COUNT,
    HEADERS,
    MAIN_COLS,
    WeldingSheetLayout,
)
from app.services.bom.welding_fee.export.service import BomWeldingFeeExportService
from app.services.bom.welding_fee.export.values import (
    _count_label,
    part_short_code,
)


def test_short_code_column_after_part_number() -> None:
    assert HEADERS[6] == "零件号"
    assert HEADERS[7] == "短代码"
    assert MAIN_COLS == 19
    assert COL_COUNT == 26
    assert COL["短代码"] == 8
    assert COL["行费"] == 25


def test_compact_layout_drops_scope_columns() -> None:
    layout = WeldingSheetLayout(omit_scope=True)
    assert "年度" not in layout.headers
    assert "基地" not in layout.headers
    assert layout.headers[0] == "序号"
    assert layout.headers[1] == "层级"
    assert layout.col["零件号"] == 3
    assert layout.main_cols == 15
    assert layout.col_count == 22


def test_part_short_code_by_level() -> None:
    assert part_short_code(part_number="501029961AA", level=0) == "9961"
    assert part_short_code(part_number="501005123AA-1", level=1) == "5123AA-1"
    assert part_short_code(part_number="501005123AA", level=1) == "5123"
    assert part_short_code(part_number=None, level=0) == ""


def _fee_item(**overrides: object) -> SimpleNamespace:
    data: dict[str, object] = {
        "line_no": 1,
        "id": 1,
        "standard_process_name": "点焊",
        "process_token": None,
        "process_qty": Decimal("2"),
        "process_qty_unit": "点",
        "unit_price": Decimal("0.1000"),
        "unit": "元/点",
        "line_fee": Decimal("0.2000"),
        "match_status": "matched",
    }
    data.update(overrides)
    return SimpleNamespace(**data)


def _fee_record(**overrides: object) -> SimpleNamespace:
    data: dict[str, object] = {
        "year": 2026,
        "location": "青岛",
        "vehicle_model": "T28",
        "bom_version": "D_00",
        "level": 0,
        "part_number": "501006875AA",
        "part_name": "左前门内板",
        "weld_info_raw": "",
        "process_line_count": 1,
        "matched_line_count": 1,
        "missing_price_count": 0,
        "total_welding_fee": Decimal("1.0000"),
        "other_fee": Decimal("0.1000"),
        "final_welding_fee": Decimal("1.1000"),
        "remark": "",
        "calc_status": "ok",
        "audit_status": "pending",
        "source_breakdown_node_id": None,
        "items": [_fee_item()],
    }
    data.update(overrides)
    return SimpleNamespace(**data)


def _write_sheet(
    records: list[SimpleNamespace],
    *,
    omit_scope: bool = False,
) -> object:
    service = BomWeldingFeeExportService()
    service.fee_repo.list_for_export = MagicMock(return_value=records)
    service.breakdown_repo.map_weld_info_by_ids = MagicMock(return_value={})
    workbook = Workbook()
    service.write_to_workbook(
        workbook,
        MagicMock(),
        year=2026,
        location="青岛",
        omit_scope_columns=omit_scope,
    )
    return workbook.worksheets[0]


def test_count_label_includes_money_totals() -> None:
    label = _count_label(
        [
            _fee_record(),
            _fee_record(
                calc_status="partial",
                total_welding_fee=Decimal("2.0000"),
                other_fee=Decimal("-0.0500"),
                final_welding_fee=Decimal("1.9500"),
                items=[_fee_item(id=2, line_fee=Decimal("2.0000"))],
            ),
        ]
    )
    assert "共 2 条（核算正常 1）" in label
    assert "本表合计" in label
    assert "系统计算 3.0000" in label
    assert "其它费用 0.0500" in label
    assert "确认单价 3.0500" in label
    assert "行费 2.2000" in label


def test_export_header_and_footer_totals() -> None:
    sheet = _write_sheet(
        [
            _fee_record(
                items=[
                    _fee_item(line_fee=Decimal("0.2000")),
                    _fee_item(id=2, line_no=2, line_fee=Decimal("0.8000")),
                ]
            )
        ]
    )
    banner = str(sheet.cell(row=2, column=1).value)
    assert "本表合计" in banner
    assert "共 1 条（核算正常 1）" in banner
    assert "系统计算 1.0000" in banner
    assert "其它费用 0.1000" in banner
    assert "确认单价 1.1000" in banner
    assert "行费 1.0000" in banner
    note = str(sheet.cell(row=3, column=1).value)
    assert "表头与末行为本表合计" in note
    assert "末行公式不进筛选" in note

    system_col = COL["系统计算"]
    other_col = COL["其它费用"]
    final_col = COL["确认单价"]
    line_col = COL["行费"]
    last_col = get_column_letter(COL_COUNT)
    assert sheet.cell(row=8, column=1).value == "合计"
    assert sheet.cell(row=8, column=system_col).value == "=SUM(N6:N7)"
    assert sheet.cell(row=8, column=other_col).value == "=SUM(O6:O7)"
    assert sheet.cell(row=8, column=final_col).value == "=SUM(P6:P7)"
    assert sheet.cell(row=8, column=line_col).value == "=SUM(Y6:Y7)"
    assert sheet.auto_filter.ref == f"A5:{last_col}7"


def test_export_empty_footer_is_zero_and_excluded_from_filter() -> None:
    sheet = _write_sheet([])
    last_col = get_column_letter(COL_COUNT)
    assert sheet.cell(row=6, column=1).value == "合计"
    assert sheet.cell(row=6, column=COL["系统计算"]).value == Decimal("0.0000")
    assert sheet.cell(row=6, column=COL["行费"]).value == Decimal("0.0000")
    assert sheet.auto_filter.ref == f"A5:{last_col}5"


def test_compact_export_footer_uses_shifted_columns() -> None:
    layout = WeldingSheetLayout(omit_scope=True)
    sheet = _write_sheet([_fee_record()], omit_scope=True)
    system_letter = get_column_letter(layout.col["系统计算"])
    line_letter = get_column_letter(layout.col["行费"])
    last_col = get_column_letter(layout.col_count)
    assert sheet.cell(row=6, column=1).value != "合计"
    assert sheet.cell(row=7, column=1).value == "合计"
    assert sheet.cell(row=7, column=layout.col["系统计算"]).value == (
        f"=SUM({system_letter}6:{system_letter}6)"
    )
    assert sheet.cell(row=7, column=layout.col["行费"]).value == (
        f"=SUM({line_letter}6:{line_letter}6)"
    )
    assert sheet.auto_filter.ref == f"A5:{last_col}6"
