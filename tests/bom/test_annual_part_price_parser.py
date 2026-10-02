"""年度零件价：采购单价 ÷ 转换系数 → 核价单价。"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook

from app.common.audit import SYSTEM_USER
from app.config import settings
from app.exceptions.base import BusinessException
from app.parsers.bom.annual_part_price_parser import (
    HEADER_CONVERSION_FACTOR,
    HEADER_PART_NUMBER,
    HEADER_PURCHASE_UNIT,
    HEADER_PURCHASE_UNIT_PRICE,
    HEADER_SUPPLY_TYPE,
    HEADER_UNIT,
    HEADER_UNIT_PRICE,
    format_conversion_factor,
    parse_annual_part_price_excel,
    parse_conversion_factor,
    resolve_annual_unit_price,
)
from app.services.bom.annual_part_price.importing.service import (
    BomAnnualPartPriceImportService,
)
from app.services.bom.annual_part_price.template_builder import (
    AnnualPartPriceTemplateRow,
    render_import_template,
)


def test_resolve_same_unit_keeps_unit_price() -> None:
    assert resolve_annual_unit_price(
        purchase_unit_price=None,
        conversion_factor=parse_conversion_factor(1),
        unit_price=Decimal("2.1300"),
    ) == Decimal("2.1300")


def test_resolve_dozen_divides() -> None:
    assert resolve_annual_unit_price(
        purchase_unit_price=Decimal("120.0000"),
        conversion_factor=parse_conversion_factor(12),
        unit_price=Decimal("0"),
    ) == Decimal("10.0000")


def test_resolve_bottle_to_mm() -> None:
    assert resolve_annual_unit_price(
        purchase_unit_price=Decimal("90.0000"),
        conversion_factor=parse_conversion_factor(30000),
        unit_price=None,
    ) == Decimal("0.0030")


def test_parse_conversion_factor_rejects_fraction() -> None:
    with pytest.raises(ValueError, match="分数"):
        parse_conversion_factor("1/12")


def test_format_conversion_factor_does_not_pad() -> None:
    assert format_conversion_factor(Decimal("1")) == "1"
    assert format_conversion_factor(Decimal("12")) == "12"
    assert format_conversion_factor(Decimal("25.4")) == "25.4"
    assert format_conversion_factor(Decimal("10000")) == "10000"
    assert format_conversion_factor(parse_conversion_factor("25.400000")) == "25.4"


def test_resolve_inch_divides() -> None:
    assert resolve_annual_unit_price(
        purchase_unit_price=Decimal("25.4000"),
        conversion_factor=parse_conversion_factor("25.4"),
        unit_price=Decimal("0"),
    ) == Decimal("1.0000")


def test_resolve_rejects_mismatch() -> None:
    try:
        resolve_annual_unit_price(
            purchase_unit_price=Decimal("120.0000"),
            conversion_factor=parse_conversion_factor(12),
            unit_price=Decimal("9.0000"),
        )
    except ValueError as exc:
        assert "不一致" in str(exc)
    else:
        raise AssertionError("expected mismatch")


def _xlsx_with_rows(headers: list[str], rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    for col, title in enumerate(headers, start=1):
        sheet.cell(1, col, title)
    for row_idx, values in enumerate(rows, start=2):
        for col, value in enumerate(values, start=1):
            sheet.cell(row_idx, col, value)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parser_old_seven_columns_still_works() -> None:
    content = _xlsx_with_rows(
        [
            HEADER_PART_NUMBER,
            "零件名称",
            HEADER_UNIT,
            HEADER_UNIT_PRICE,
            HEADER_SUPPLY_TYPE,
            "供应商",
            "备注",
        ],
        [["P1", "螺母", "pcs", Decimal("2.1300"), "标准件", "", ""]],
    )
    parsed = parse_annual_part_price_excel(content)
    assert parsed.errors == []
    assert parsed.rows[0].unit_price == Decimal("2.1300")
    assert parsed.rows[0].conversion_factor == 1
    assert parsed.rows[0].purchase_unit_price is None


def test_parser_purchase_price_overrides_zero_unit_price() -> None:
    content = _xlsx_with_rows(
        [
            HEADER_PART_NUMBER,
            HEADER_PURCHASE_UNIT,
            HEADER_PURCHASE_UNIT_PRICE,
            HEADER_CONVERSION_FACTOR,
            HEADER_UNIT,
            HEADER_UNIT_PRICE,
            HEADER_SUPPLY_TYPE,
        ],
        [["P-DOZEN", "打", Decimal("120"), 12, "pcs", Decimal("0"), "标准件"]],
    )
    parsed = parse_annual_part_price_excel(content)
    assert parsed.errors == []
    assert parsed.rows[0].unit_price == Decimal("10.0000")
    assert parsed.rows[0].purchase_unit == "打"
    assert parsed.rows[0].conversion_factor == Decimal("12")


def test_parser_inch_factor() -> None:
    content = _xlsx_with_rows(
        [
            HEADER_PART_NUMBER,
            HEADER_PURCHASE_UNIT_PRICE,
            HEADER_CONVERSION_FACTOR,
            HEADER_UNIT_PRICE,
            HEADER_SUPPLY_TYPE,
        ],
        [["P-INCH", Decimal("25.4"), Decimal("25.4"), Decimal("0"), "沿用"]],
    )
    parsed = parse_annual_part_price_excel(content)
    assert parsed.errors == []
    assert parsed.rows[0].conversion_factor == Decimal("25.400000")
    assert parsed.rows[0].unit_price == Decimal("1.0000")


def test_parser_rejects_fraction_factor() -> None:
    content = _xlsx_with_rows(
        [
            HEADER_PART_NUMBER,
            HEADER_PURCHASE_UNIT_PRICE,
            HEADER_CONVERSION_FACTOR,
            HEADER_UNIT_PRICE,
            HEADER_SUPPLY_TYPE,
        ],
        [["P-DOZEN", Decimal("120"), "1/12", Decimal("0"), "标准件"]],
    )
    parsed = parse_annual_part_price_excel(content)
    assert parsed.rows == []
    assert any("分数" in err for err in parsed.errors)


def test_blank_template_roundtrip() -> None:
    content = render_import_template(
        [
            AnnualPartPriceTemplateRow(
                part_number="QR36410",
                part_name="焊接方螺母",
                specification="M6",
                purchase_unit="pcs",
                unit="pcs",
                supply_type="标准件",
            )
        ]
    )
    parsed = parse_annual_part_price_excel(content)
    assert parsed.errors == []
    assert len(parsed.rows) == 1
    row = parsed.rows[0]
    assert row.part_number == "QR36410"
    assert row.specification == "M6"
    assert row.purchase_unit == "pcs"
    assert row.purchase_unit_price is None
    assert row.conversion_factor == 1
    assert row.unit_price == Decimal("0.0000")


@patch("app.services.bom.annual_part_price.importing.service.LocationService")
def test_import_rejects_workbook_when_any_row_invalid(
    mock_location_cls: MagicMock,
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "UPLOAD_DIR", tmp_path)
    mock_location_cls.return_value.require_enabled.return_value = "青岛"
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.append(["零件号", "单价", "原供货类型"])
    sheet.append(["P001", Decimal("1.5000"), "标准件"])
    sheet.append(["", Decimal("2.0000"), "标准件"])
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(BusinessException, match="整本未导入") as raised:
        BomAnnualPartPriceImportService().import_excel(
            db=None,  # type: ignore[arg-type]
            file_content=buffer.getvalue(),
            filename="mixed.xlsx",
            location="青岛",
            year=2025,
            current_user=SYSTEM_USER,
        )
    message = str(raised.value)
    assert "「零件号」是空的" in message
    assert "请填写零件号" in message
    assert "\n· " in message
