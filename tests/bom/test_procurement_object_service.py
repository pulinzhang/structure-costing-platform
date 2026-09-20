"""采购对象清单 Service 测试（内存 SQLite）。"""

from __future__ import annotations

from collections.abc import Generator
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.audit import SYSTEM_USER
from app.database import Base
from app.exceptions.base import BusinessException
from app.models.bom.node import BomNode, BomSnapshot
from app.models.bom.procurement_object import (  # noqa: F401
    BomProcurementObject,
    BomProcurementObjectList,
)
from app.models.bom.procurement_object_generate_log import (  # noqa: F401
    BomProcurementObjectGenerateLog,
)
from app.models.masterdata.location import Location, LocationAlias  # noqa: F401
from app.models.masterdata.supplier import Supplier, SupplierAlias  # noqa: F401
from app.models.masterdata.vehicle_model import VehicleModel
from app.repositories.bom.bom.node import BomNodeRepository
from app.repositories.bom.bom.snapshot import BomSnapshotRepository
from app.repositories.masterdata.location.location_repository import LocationRepository
from app.schemas.dto.bom_dto import BomNodeWrite, BomSnapshotWrite
from app.schemas.dto.bom_procurement_object_dto import (
    BomProcurementObjectBatchApproveDTO,
    BomProcurementObjectBatchDeleteDTO,
    BomProcurementObjectBatchReverseAuditDTO,
    BomProcurementObjectConfirmDTO,
    BomProcurementObjectGenerateDTO,
    BomProcurementObjectGenerateLogQuery,
    BomProcurementObjectPrepQuery,
    BomProcurementObjectQuery,
    BomProcurementObjectSupplierDTO,
)
from app.schemas.dto.location_dto import LocationAliasInsert, LocationInsert
from app.schemas.dto.supplier_dto import SupplierCreateDTO
from app.services.bom.procurement_object.export import BomProcurementObjectExportService
from app.services.bom.procurement_object.insert_service import BomProcurementObjectInsertService
from app.services.bom.procurement_object.service import BomProcurementObjectService
from app.services.masterdata.supplier import SupplierService


@pytest.fixture()
def db() -> Generator[Session, None, None]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    session = factory()
    repo = LocationRepository()
    repo.create(
        session,
        LocationInsert(
            location="芜湖",
            is_enabled=True,
            sort_order=1,
            created_by="system",
            updated_by="system",
        ),
    )
    repo.create_alias(session, LocationAliasInsert(alias="芜湖基地", location="芜湖"))
    session.add(
        VehicleModel(
            location="芜湖",
            vehicle_model="T13J",
            effective_year=2020,
            effective_month=1,
            current_supplied_units=0,
            created_by="system",
            updated_by="system",
        )
    )
    session.flush()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def _seed_tree(db: Session, *, with_review_rows: bool = False) -> None:
    snapshot = BomSnapshotRepository().create(
        db,
        BomSnapshotWrite(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            source_filename="t13j.xlsx",
            created_by="tester",
            updated_by="tester",
        ),
    )
    node_repo = BomNodeRepository()
    l0 = node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot.id,
            level=0,
            part_number="L0-A",
            node_role="supply",
            created_by="tester",
            updated_by="tester",
            sort_no=0,
            is_supply_unit=True,
            excel_filename="t13j.xlsx",
            excel_row=8,
            weld_info="焊点:4",
        ),
    )
    node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot.id,
            parent_id=l0.id,
            level=1,
            part_number="L1-STAMP",
            node_role="make",
            created_by="tester",
            updated_by="tester",
            sort_no=1,
            part_name="加强板",
            internal_number="ACA674001",
            thickness_mm=Decimal("0.8000"),
            thickness_text="T=0.8",
            excel_filename="t13j.xlsx",
            excel_row=9,
        ),
    )
    if with_review_rows:
        node_repo.create(
            db,
            BomNodeWrite(
                snapshot_id=snapshot.id,
                parent_id=l0.id,
                level=1,
                part_number="B11-4102613",
                node_role="buy",
                created_by="tester",
                updated_by="tester",
                part_name="点焊胶",
                excel_filename="t13j.xlsx",
                excel_row=10,
            ),
        )
        l0_out = node_repo.create(
            db,
            BomNodeWrite(
                snapshot_id=snapshot.id,
                level=0,
                part_number="L0-OUT",
                node_role="supply",
                created_by="tester",
                updated_by="tester",
                is_supply_unit=True,
                remark="凸焊总成供货",
                excel_filename="t13j.xlsx",
                excel_row=20,
            ),
        )
        node_repo.create(
            db,
            BomNodeWrite(
                snapshot_id=snapshot.id,
                parent_id=l0_out.id,
                level=1,
                part_number="L1-OUT-CHILD",
                node_role="make",
                created_by="tester",
                updated_by="tester",
                part_name="加强板",
                internal_number="ACA674099",
                thickness_mm=Decimal("0.8000"),
                thickness_text="T=0.8",
                excel_filename="t13j.xlsx",
                excel_row=21,
            ),
        )
    db.commit()


def _scope() -> BomProcurementObjectGenerateDTO:
    return BomProcurementObjectGenerateDTO(
        location="芜湖",
        vehicle_model="T13J",
        bom_version="V3",
    )


def test_generate_and_regenerate_keeps_confirm(db: Session) -> None:
    _seed_tree(db)
    service = BomProcurementObjectService()
    first = service.generate(db, _scope(), current_user=SYSTEM_USER)
    assert first.object_count == 1
    items, total, _filters = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3", size=50),
    )
    assert total == 1
    assert items[0].internal_number == "ACA674001"
    confirmed = service.confirm(
        db,
        items[0].id,
        BomProcurementObjectConfirmDTO(confirm_status="outsourced", confirm_remark="外发"),
        current_user=SYSTEM_USER,
    )
    assert confirmed.confirm_status == "outsourced"
    second = service.generate(db, _scope(), current_user=SYSTEM_USER)
    assert second.object_count == 1
    items2, _, _ = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3", size=50),
    )
    assert items2[0].confirm_status == "outsourced"
    logs, log_total = service.list_generate_logs(db, BomProcurementObjectGenerateLogQuery(size=20))
    assert log_total == 2
    assert logs[0].status == "success"
    assert logs[0].object_count == 1
    assert "采购对象" in logs[0].detail_text
    assert "退化供货件亚类" in logs[0].detail_text
    assert "已按毛坯供货" not in logs[0].detail_text


def test_generate_blocked_when_approved(db: Session) -> None:
    _seed_tree(db)
    service = BomProcurementObjectService()
    service.generate(db, _scope(), current_user=SYSTEM_USER)
    service.batch_approve(
        db,
        BomProcurementObjectBatchApproveDTO(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            audit_remark="通过",
        ),
        current_user=SYSTEM_USER,
    )
    with pytest.raises(BusinessException) as exc:
        service.generate(db, _scope(), current_user=SYSTEM_USER)
    assert exc.value.code == 409
    logs, log_total = service.list_generate_logs(db, BomProcurementObjectGenerateLogQuery(size=20))
    assert log_total == 2
    assert logs[0].status == "failed"
    assert "已审核" in logs[0].detail_text


def test_batch_approve_empty_remark_rejected_by_dto() -> None:
    with pytest.raises(ValidationError):
        BomProcurementObjectBatchApproveDTO(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            audit_remark="",
        )


def test_batch_delete_includes_other_leaf_and_anomaly(db: Session) -> None:
    _seed_tree(db, with_review_rows=True)
    service = BomProcurementObjectService()
    generated = service.generate(db, _scope(), current_user=SYSTEM_USER)
    assert generated.object_count == 2
    items, total, _filters = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3", size=50),
    )
    assert total == 4
    sections = {item.section for item in items}
    assert "other_leaf" in sections
    assert "anomaly" in sections
    # 无版本且无已审 → 清该基地+项目全部清单头
    result = service.batch_delete(
        db,
        BomProcurementObjectBatchDeleteDTO(
            location="芜湖",
            vehicle_model="T13J",
        ),
        current_user=SYSTEM_USER,
    )
    assert result.deleted_count == 4
    assert result.skipped_approved_count == 0
    remaining, remaining_total, _ = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", size=50),
    )
    assert remaining_total == 0
    assert remaining == []
    assert (
        service.repository.list_lists_by_location_vehicle(db, location="芜湖", vehicle_model="T13J")
        == []
    )


def test_batch_delete_clears_review_rows_while_skipping_approved(db: Session) -> None:
    _seed_tree(db, with_review_rows=True)
    service = BomProcurementObjectService()
    service.generate(db, _scope(), current_user=SYSTEM_USER)
    service.batch_approve(
        db,
        BomProcurementObjectBatchApproveDTO(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            audit_remark="通过",
        ),
        current_user=SYSTEM_USER,
    )
    result = service.batch_delete(
        db,
        BomProcurementObjectBatchDeleteDTO(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
        ),
        current_user=SYSTEM_USER,
    )
    assert result.deleted_count == 2
    assert result.skipped_approved_count == 2
    remaining, remaining_total, _ = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3", size=50),
    )
    assert remaining_total == 2
    assert {item.section for item in remaining} == {"object"}
    assert all(item.audit_status == "approved" for item in remaining)
    header = service.repository.get_list_by_scope(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    assert header is not None
    assert (
        header.object_l0_degenerate_count
        + header.object_l1_subassembly_count
        + header.object_l1_stamping_count
        == 2
    )
    assert header.other_leaf_count == 0
    assert header.anomaly_count == 0
    assert header.qa_passed is False
    assert header.qa_error_text is not None
    assert "部分删除" in header.qa_error_text


def test_engineering_change_rollback_delete_keeps_header_when_rows_remain(db: Session) -> None:
    """整包 generate 后的清单：回滚只删部分行时头保留，且不改 QA。"""
    _seed_tree(db, with_review_rows=True)
    service = BomProcurementObjectService()
    result = service.generate(db, _scope(), current_user=SYSTEM_USER)
    assert result.header.qa_passed is True
    items, _, _ = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3", size=50),
    )
    object_items = [item for item in items if item.section == "object"]
    assert len(object_items) >= 2
    target_id = object_items[0].id
    before = service.repository.get_list_by_scope(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    assert before is not None
    before_qa = before.qa_passed
    before_error = before.qa_error_text
    before_stamping = before.object_l1_stamping_count

    service.delete_object_for_engineering_change_rollback(db, target_id)
    db.commit()

    after = service.repository.get_list_by_scope(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    assert after is not None
    assert after.qa_passed is before_qa
    assert after.qa_error_text == before_error
    assert after.object_l1_stamping_count == before_stamping
    assert service.repository.get_object_by_id(db, target_id) is None


def test_engineering_change_rollback_deletes_empty_header(db: Session) -> None:
    """EC 建空头插行后全部回滚：头须删掉（不依赖 procurement_list 效果链接）。"""
    _seed_tree(db)
    snapshot = db.query(BomSnapshot).one()
    service = BomProcurementObjectService()
    inserted_ids = BomProcurementObjectInsertService().insert_for_new_tree(
        db,
        snapshot_id=snapshot.id,
        new_l0_id=db.query(BomNode).filter(BomNode.level == 0).one().id,
        old_l0_id=None,
        current_user=SYSTEM_USER,
    )
    assert inserted_ids
    header = service.repository.get_list_by_snapshot_id(db, snapshot.id)
    assert header is not None
    assert header.qa_passed is True

    for object_id in inserted_ids:
        service.delete_object_for_engineering_change_rollback(db, object_id)
    db.commit()

    assert service.repository.get_list_by_snapshot_id(db, snapshot.id) is None


def test_confirm_skipped_when_approved(db: Session) -> None:
    _seed_tree(db)
    service = BomProcurementObjectService()
    service.generate(db, _scope(), current_user=SYSTEM_USER)
    items, _, _ = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3", size=50),
    )
    service.batch_approve(
        db,
        BomProcurementObjectBatchApproveDTO(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            audit_remark="通过",
        ),
        current_user=SYSTEM_USER,
    )
    vo = service.confirm(
        db,
        items[0].id,
        BomProcurementObjectConfirmDTO(confirm_status="in_house"),
        current_user=SYSTEM_USER,
    )
    assert vo.confirm_status != "in_house"
    assert vo.audit_status == "approved"


def test_confirm_allowed_when_qa_failed(db: Session) -> None:
    _seed_tree(db)
    service = BomProcurementObjectService()
    service.generate(db, _scope(), current_user=SYSTEM_USER)
    header = service.repository.get_list_by_scope(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    assert header is not None
    header.qa_passed = False
    header.qa_error_text = "结构缺口"
    db.commit()
    items, _, _ = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3", size=50),
    )
    vo = service.confirm(
        db,
        items[0].id,
        BomProcurementObjectConfirmDTO(confirm_status="outsourced", confirm_remark="外发"),
        current_user=SYSTEM_USER,
    )
    assert vo.confirm_status == "outsourced"


def test_batch_reverse_audit_required_remark() -> None:
    with pytest.raises(ValidationError):
        BomProcurementObjectBatchReverseAuditDTO(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            reverse_audit_remark="",
        )


def test_export_workbook_has_styled_sheets_without_cover(db: Session) -> None:
    _seed_tree(db)
    BomProcurementObjectService().generate(db, _scope(), current_user=SYSTEM_USER)
    content, filename = BomProcurementObjectExportService().export_excel(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3"),
        SYSTEM_USER,
    )
    assert filename.startswith("采购对象清单_芜湖_T13J_V3_")
    assert filename.endswith(".xlsx")
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["采购对象", "标准件及胶等", "异常"]
    object_sheet = workbook["采购对象"]
    assert object_sheet["A1"].value == "采购对象清单 · 采购对象"
    assert object_sheet["A4"].value == "对着原件"
    assert object_sheet["A5"].value == "Excel行号"
    assert object_sheet["C5"].value == "零件号"
    assert object_sheet["D5"].value == "短码"
    assert object_sheet["E5"].value == "内部号"
    assert object_sheet["F5"].value == "零件名称"
    assert object_sheet["I5"].value == "对象种类"
    assert object_sheet["K5"].value == "供应商"
    assert object_sheet["M5"].value == "材料"
    assert object_sheet["M4"].value == "佐证"
    assert object_sheet.freeze_panes == "I6"
    assert object_sheet["A6"].value == 9
    assert object_sheet["C6"].value == "L1-STAMP"
    assert object_sheet["D6"].value is None
    assert object_sheet["E6"].value == "ACA674001"
    assert object_sheet["U6"].value == "芜湖"
    assert object_sheet["X5"].value == "生失效"
    assert object_sheet["X6"].value == "2020年01月 – 未结束"
    assert object_sheet["C6"].font.color.rgb in {"0037474F", "37474F"}
    assert object_sheet["D6"].font.color.rgb in {"001F4E79", "1F4E79"}
    assert object_sheet["E6"].font.color.rgb in {"00808080", "808080"}
    assert object_sheet["U6"].font.color.rgb in {"00455A64", "455A64"}
    assert object_sheet["A1"].fill.fgColor.rgb in {"001A237E", "1A237E"}
    workbook.close()


def test_export_merges_same_l0_assembly_cells(db: Session) -> None:
    _seed_tree(db)
    snapshot = db.query(BomSnapshot).filter(BomSnapshot.vehicle_model == "T13J").one()
    l0 = db.query(BomNode).filter(BomNode.part_number == "L0-A").one()
    BomNodeRepository().create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot.id,
            parent_id=l0.id,
            level=1,
            part_number="L1-STAMP-2",
            node_role="make",
            created_by="tester",
            updated_by="tester",
            part_name="加强板2",
            internal_number="ACA674002",
            thickness_mm=Decimal("0.8000"),
            thickness_text="T=0.8",
            excel_filename="t13j.xlsx",
            excel_row=10,
        ),
    )
    db.commit()
    BomProcurementObjectService().generate(db, _scope(), current_user=SYSTEM_USER)
    content, _filename = BomProcurementObjectExportService().export_excel(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3"),
        SYSTEM_USER,
    )
    workbook = load_workbook(BytesIO(content))
    object_sheet = workbook["采购对象"]
    merged = {str(item) for item in object_sheet.merged_cells.ranges}
    assert "G6:G7" in merged
    assert "H6:H7" in merged
    assert "T6:T7" in merged
    assert object_sheet["G6"].value == "L0-A"
    assert object_sheet["T6"].value == "t13j.xlsx"
    workbook.close()


def test_list_default_includes_expired_and_reads_l0_period(db: Session) -> None:
    _seed_tree(db)
    service = BomProcurementObjectService()
    service.generate(db, _scope(), current_user=SYSTEM_USER)
    l0 = db.query(BomNode).filter(BomNode.part_number == "L0-A", BomNode.level == 0).one()
    l0.expiry_year = 2020
    l0.expiry_month = 6
    for child in db.query(BomNode).filter(BomNode.parent_id == l0.id).all():
        child.expiry_year = 2020
        child.expiry_month = 6
    db.commit()

    all_items, all_total, _ = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3", size=50),
    )
    assert all_total == 1
    assert all_items[0].effective_year == 2020
    assert all_items[0].effective_month == 1
    assert all_items[0].expiry_year == 2020
    assert all_items[0].expiry_month == 6

    active_items, active_total, _ = service.paginate(
        db,
        BomProcurementObjectQuery(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            include_expired=False,
            size=50,
        ),
    )
    assert active_total == 0
    assert active_items == []


def test_list_period_reads_l0_even_when_source_node_differs(db: Session) -> None:
    """防回归：VO/过期筛只认树根；故意把 L1 四列改成别的也不能当真相。"""
    _seed_tree(db)
    service = BomProcurementObjectService()
    service.generate(db, _scope(), current_user=SYSTEM_USER)
    l0 = db.query(BomNode).filter(BomNode.part_number == "L0-A", BomNode.level == 0).one()
    l1 = db.query(BomNode).filter(BomNode.part_number == "L1-STAMP").one()
    l0.effective_year = 2020
    l0.effective_month = 1
    l0.expiry_year = 2021
    l0.expiry_month = 6
    l1.effective_year = 2090
    l1.effective_month = 1
    l1.expiry_year = 2099
    l1.expiry_month = 12
    db.commit()

    items, total, _ = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3", size=50),
    )
    assert total == 1
    assert items[0].part_number == "L1-STAMP"
    assert items[0].effective_year == 2020
    assert items[0].effective_month == 1
    assert items[0].expiry_year == 2021
    assert items[0].expiry_month == 6


def test_insert_pair_by_sort_no_not_letter_stem() -> None:
    """§14.3：换号抄确认按树上同一 sort_no；禁止字母版本同茎猜旧号。"""
    from types import SimpleNamespace

    from app.services.bom.procurement_object.insert_service import _new_to_old_part_numbers

    def node(
        *,
        node_id: int,
        level: int,
        parent_id: int | None,
        part_number: str,
        sort_no: int,
    ) -> SimpleNamespace:
        return SimpleNamespace(
            id=node_id,
            level=level,
            parent_id=parent_id,
            part_number=part_number,
            sort_no=sort_no,
        )

    old_tree = [
        node(node_id=1, level=0, parent_id=None, part_number="3086AA", sort_no=0),
        node(node_id=2, level=1, parent_id=1, part_number="PART-AA", sort_no=1),
        node(node_id=3, level=1, parent_id=1, part_number="OTHER-1", sort_no=2),
    ]
    # 同茎 PART-AB 故意放在 sort_no=2；若按字母同茎会错配到 PART-AA
    new_tree = [
        node(node_id=11, level=0, parent_id=None, part_number="3086AB", sort_no=0),
        node(node_id=12, level=1, parent_id=11, part_number="OTHER-2", sort_no=1),
        node(node_id=13, level=1, parent_id=11, part_number="PART-AB", sort_no=2),
    ]
    mapping = _new_to_old_part_numbers(old_tree, new_tree)  # type: ignore[arg-type]
    assert mapping["3086AB"] == "3086AA"
    assert mapping["OTHER-2"] == "PART-AA"
    assert mapping["PART-AB"] == "OTHER-1"
    assert mapping["PART-AB"] != "PART-AA"


def test_insert_for_new_tree_copies_confirm_not_audit(db: Session) -> None:
    """工程变更插行：按 sort_no 抄确认与厂家；新行待审；旧已审行仍在。"""
    _seed_tree(db)
    service = BomProcurementObjectService()
    service.generate(db, _scope(), current_user=SYSTEM_USER)
    SupplierService().create(
        db,
        SupplierCreateDTO(
            abbreviation="常庆",
            full_name="安庆常庆汽车部件有限公司",
            erp_code="CQ001",
        ),
        current_user=SYSTEM_USER,
    )
    items, _, _ = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3", size=50),
    )
    old_item = items[0]
    service.confirm(
        db,
        old_item.id,
        BomProcurementObjectConfirmDTO(confirm_status="outsourced", confirm_remark="外发旧树"),
        current_user=SYSTEM_USER,
    )
    service.update_supplier(
        db,
        old_item.id,
        BomProcurementObjectSupplierDTO(supplier_name="安庆常庆汽车部件有限公司"),
        current_user=SYSTEM_USER,
    )
    service.batch_approve(
        db,
        BomProcurementObjectBatchApproveDTO(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            audit_remark="旧树已审",
        ),
        current_user=SYSTEM_USER,
    )

    old_l0 = db.query(BomNode).filter(BomNode.part_number == "L0-A", BomNode.level == 0).one()
    snapshot_id = old_l0.snapshot_id
    node_repo = BomNodeRepository()
    new_l0 = node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot_id,
            level=0,
            part_number="L0-B",
            node_role="supply",
            created_by="tester",
            updated_by="tester",
            sort_no=0,
            seq_no=2,
            is_supply_unit=True,
            excel_filename="t13j.xlsx",
            excel_row=30,
            effective_year=2025,
            effective_month=7,
        ),
    )
    node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot_id,
            parent_id=new_l0.id,
            level=1,
            part_number="L1-STAMP-NEW",
            node_role="make",
            created_by="tester",
            updated_by="tester",
            sort_no=1,
            part_name="加强板新",
            internal_number="ACA674002",
            thickness_mm=Decimal("0.8000"),
            thickness_text="T=0.8",
            excel_filename="t13j.xlsx",
            excel_row=31,
            effective_year=2025,
            effective_month=7,
        ),
    )
    db.flush()

    inserted_ids = BomProcurementObjectInsertService().insert_for_new_tree(
        db,
        snapshot_id=snapshot_id,
        new_l0_id=new_l0.id,
        old_l0_id=old_l0.id,
        current_user=SYSTEM_USER,
    )
    db.commit()
    assert len(inserted_ids) == 1

    all_items, total, _ = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3", size=50),
    )
    assert total == 2
    by_number = {item.part_number: item for item in all_items}
    old_row = by_number["L1-STAMP"]
    new_row = by_number["L1-STAMP-NEW"]
    assert old_row.audit_status == "approved"
    assert old_row.confirm_status == "outsourced"
    assert new_row.id == inserted_ids[0]
    assert new_row.confirm_status == "outsourced"
    assert new_row.confirm_remark == "外发旧树"
    assert new_row.supplier_name == "安庆常庆汽车部件有限公司"
    assert new_row.audit_status == "pending"
    assert new_row.l0_part_number == "L0-B"


def test_has_supplier_filter(db: Session) -> None:
    _seed_tree(db)
    service = BomProcurementObjectService()
    service.generate(db, _scope(), current_user=SYSTEM_USER)
    SupplierService().create(
        db,
        SupplierCreateDTO(
            abbreviation="常庆",
            full_name="安庆常庆汽车部件有限公司",
            erp_code="CQ001",
        ),
        current_user=SYSTEM_USER,
    )
    items, _, _ = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3", size=50),
    )
    service.confirm(
        db,
        items[0].id,
        BomProcurementObjectConfirmDTO(confirm_status="outsourced"),
        current_user=SYSTEM_USER,
    )
    service.update_supplier(
        db,
        items[0].id,
        BomProcurementObjectSupplierDTO(supplier_name="安庆常庆汽车部件有限公司"),
        current_user=SYSTEM_USER,
    )

    with_supplier, with_total, _ = service.paginate(
        db,
        BomProcurementObjectQuery(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            has_supplier=True,
            size=50,
        ),
    )
    assert with_total == 1
    assert with_supplier[0].supplier_name == "安庆常庆汽车部件有限公司"

    without_supplier, without_total, _ = service.paginate(
        db,
        BomProcurementObjectQuery(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            has_supplier=False,
            size=50,
        ),
    )
    assert without_total == 0
    assert without_supplier == []


def test_list_expands_location_alias_export_uses_exact(db: Session) -> None:
    """列表 expand 别名命中规范名行；导出按请求 location 精确匹配。"""
    _seed_tree(db)
    service = BomProcurementObjectService()
    service.generate(db, _scope(), current_user=SYSTEM_USER)

    listed, total, _ = service.paginate(
        db,
        BomProcurementObjectQuery(
            location="芜湖基地",
            vehicle_model="T13J",
            bom_version="V3",
            size=50,
        ),
    )
    assert total == 1
    assert listed[0].location == "芜湖"

    content, filename = BomProcurementObjectExportService().export_excel(
        db,
        BomProcurementObjectQuery(
            location="芜湖基地",
            vehicle_model="T13J",
            bom_version="V3",
        ),
        SYSTEM_USER,
    )
    assert "芜湖基地" in filename
    workbook = load_workbook(BytesIO(content))
    object_sheet = workbook["采购对象"]
    # 精确匹配无行：表头下无数据（数据从第 6 行起）
    assert object_sheet["C6"].value is None
    workbook.close()

    content_ok, _ = BomProcurementObjectExportService().export_excel(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3"),
        SYSTEM_USER,
    )
    workbook_ok = load_workbook(BytesIO(content_ok))
    assert workbook_ok["采购对象"]["C6"].value == "L1-STAMP"
    workbook_ok.close()


def test_batch_delete_ignores_include_expired_filter(db: Session) -> None:
    """批量按整份清单打；列表「不含已过期」看不到的行仍会被删。"""
    _seed_tree(db)
    service = BomProcurementObjectService()
    service.generate(db, _scope(), current_user=SYSTEM_USER)
    l0 = db.query(BomNode).filter(BomNode.part_number == "L0-A", BomNode.level == 0).one()
    l0.expiry_year = 2020
    l0.expiry_month = 6
    for child in db.query(BomNode).filter(BomNode.parent_id == l0.id).all():
        child.expiry_year = 2020
        child.expiry_month = 6
    db.commit()

    _, active_total, _ = service.paginate(
        db,
        BomProcurementObjectQuery(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            include_expired=False,
            size=50,
        ),
    )
    assert active_total == 0

    result = service.batch_delete(
        db,
        BomProcurementObjectBatchDeleteDTO(
            location="芜湖",
            vehicle_model="T13J",
        ),
        current_user=SYSTEM_USER,
    )
    assert result.deleted_count == 1
    assert (
        service.repository.list_lists_by_location_vehicle(db, location="芜湖", vehicle_model="T13J")
        == []
    )


def test_generate_normalizes_vehicle_model_case(db: Session) -> None:
    """§13：干净库精确匹配前 normalize；t13j 与 T13J 同一代号。"""
    _seed_tree(db)
    service = BomProcurementObjectService()
    result = service.generate(
        db,
        BomProcurementObjectGenerateDTO(
            location="芜湖",
            vehicle_model="t13j",
            bom_version="V3",
        ),
        current_user=SYSTEM_USER,
    )
    assert result.object_count == 1
    prep = service.prep_stats(
        db,
        BomProcurementObjectPrepQuery(
            location="芜湖",
            vehicle_model="t13j",
            bom_version="V3",
        ),
    )
    assert prep.object_count == 1
    assert prep.node_count == 2
    assert prep.message == ""


def test_generate_rejects_vehicle_model_that_normalizes_empty(db: Session) -> None:
    """§13：normalize 后空代号 → 项目代号无效（禁止靠 expand 猜）。"""
    _seed_tree(db)
    with pytest.raises(BusinessException, match="项目代号无效"):
        BomProcurementObjectService().generate(
            db,
            BomProcurementObjectGenerateDTO(
                location="芜湖",
                vehicle_model="青岛",
                bom_version="V3",
            ),
            current_user=SYSTEM_USER,
        )


def test_list_and_summary_by_location_vehicle_without_version(db: Session) -> None:
    """列表/summary 取件 = 基地+项目；版本可选。"""
    _seed_tree(db)
    service = BomProcurementObjectService()
    service.generate(db, _scope(), current_user=SYSTEM_USER)
    items, total, filters = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", size=50),
    )
    assert total >= 1
    assert items
    assert filters["summary"]["object_count"] >= 1
    summary = service.get_summary(db, location="芜湖", vehicle_model="T13J")
    assert summary.qa_passed is True


def test_optional_bom_version_filters_climbed_l0_only(db: Session) -> None:
    """有版本时只认爬到的 L0.bom_version，不是清单头键。"""
    _seed_tree(db)
    service = BomProcurementObjectService()
    service.generate(db, _scope(), current_user=SYSTEM_USER)
    header = service.repository.get_list_by_scope(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    assert header is not None
    snapshot_id = header.snapshot_id
    l0 = (
        db.query(BomNode)
        .filter(
            BomNode.snapshot_id == snapshot_id,
            BomNode.level == 0,
            BomNode.part_number == "L0-A",
        )
        .one()
    )
    l0.bom_version = "EC-202505"
    db.commit()

    ec_items, ec_total, _ = service.paginate(
        db,
        BomProcurementObjectQuery(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="EC-202505",
            size=50,
        ),
    )
    assert ec_total >= 1
    assert all(item.l0_part_number == "L0-A" for item in ec_items)

    _, sop_total, _ = service.paginate(
        db,
        BomProcurementObjectQuery(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            size=50,
        ),
    )
    assert sop_total == 0
    assert header.bom_version == "V3"
