"""采购对象供应商窄表模板 / 上传。"""

from __future__ import annotations

from collections.abc import Generator
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.audit import SYSTEM_USER
from app.config import settings
from app.database import Base
from app.exceptions.base import BusinessException
from app.models.bom.node import BomNode, BomSnapshot  # noqa: F401
from app.models.bom.procurement_object import (  # noqa: F401
    BomProcurementObject,
    BomProcurementObjectList,
)
from app.models.bom.procurement_object_generate_log import (  # noqa: F401
    BomProcurementObjectGenerateLog,
)
from app.models.file_trace.upload_file import UploadFile  # noqa: F401
from app.models.masterdata.location import Location, LocationAlias  # noqa: F401
from app.models.masterdata.supplier import Supplier, SupplierAlias  # noqa: F401
from app.models.masterdata.vehicle_model import VehicleModel
from app.parsers.bom.procurement_object_supplier_parser import (
    HEADER_INTERNAL_NUMBER,
    HEADER_SUPPLIER,
)
from app.repositories.bom.bom.node import BomNodeRepository
from app.repositories.bom.bom.snapshot import BomSnapshotRepository
from app.repositories.masterdata.location.location_repository import LocationRepository
from app.schemas.dto.bom_dto import BomNodeWrite, BomSnapshotWrite
from app.schemas.dto.bom_procurement_object_dto import (
    BomProcurementObjectBatchApproveDTO,
    BomProcurementObjectConfirmDTO,
    BomProcurementObjectGenerateDTO,
    BomProcurementObjectQuery,
)
from app.schemas.dto.location_dto import LocationAliasInsert, LocationInsert
from app.schemas.dto.supplier_dto import SupplierCreateDTO, SupplierUpdateDTO
from app.services.bom.procurement_object.service import BomProcurementObjectService
from app.services.masterdata.supplier.supplier_service import SupplierService

FULL_NAME = "安庆常庆汽车部件有限公司"


@pytest.fixture()
def db(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Generator[Session, None, None]:
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(settings, "UPLOAD_DIR", upload_dir)

    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _enable_fk(dbapi_connection: object, _connection_record: object) -> None:
        cursor = dbapi_connection.cursor()  # type: ignore[attr-defined]
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    session = factory()
    repo = LocationRepository()
    repo.create(
        session,
        LocationInsert(
            location="芜湖",
            is_enabled=True,
            sort_order=1,
            created_by="system",
            updated_by="system",
        ),
    )
    repo.create_alias(session, LocationAliasInsert(alias="芜湖基地", location="芜湖"))
    session.add(
        VehicleModel(
            location="芜湖",
            vehicle_model="T13J",
            effective_year=2020,
            effective_month=1,
            current_supplied_units=0,
            created_by="system",
            updated_by="system",
        )
    )
    session.flush()
    SupplierService().create(
        session,
        SupplierCreateDTO(
            abbreviation="常庆",
            full_name=FULL_NAME,
            erp_code="CQ001",
            aliases=["常庆汽配"],
        ),
        current_user=SYSTEM_USER,
    )
    session.flush()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def _seed_tree(db: Session) -> None:
    snapshot = BomSnapshotRepository().create(
        db,
        BomSnapshotWrite(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            source_filename="t13j.xlsx",
            created_by="tester",
            updated_by="tester",
        ),
    )
    node_repo = BomNodeRepository()
    l0 = node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot.id,
            level=0,
            part_number="L0-A",
            node_role="supply",
            created_by="tester",
            updated_by="tester",
            is_supply_unit=True,
            excel_filename="t13j.xlsx",
            excel_row=8,
            weld_info="焊点:4",
        ),
    )
    node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot.id,
            parent_id=l0.id,
            level=1,
            part_number="L1-STAMP",
            node_role="make",
            created_by="tester",
            updated_by="tester",
            part_name="加强板",
            internal_number="ACA674001",
            thickness_mm=Decimal("0.8000"),
            thickness_text="T=0.8",
            excel_filename="t13j.xlsx",
            excel_row=9,
        ),
    )
    db.commit()


def _seed_shared_part_tree(db: Session) -> None:
    _seed_tree(db)
    snapshot = BomSnapshotRepository().get_by_unique_key(db, "芜湖", "T13J", "V3")
    assert snapshot is not None
    node_repo = BomNodeRepository()
    l0_b = node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot.id,
            level=0,
            part_number="L0-B",
            node_role="supply",
            created_by="tester",
            updated_by="tester",
            is_supply_unit=True,
            excel_filename="t13j.xlsx",
            excel_row=18,
            weld_info="焊点:2",
        ),
    )
    node_repo.create(
        db,
        BomNodeWrite(
            snapshot_id=snapshot.id,
            parent_id=l0_b.id,
            level=1,
            part_number="L1-STAMP",
            node_role="make",
            created_by="tester",
            updated_by="tester",
            part_name="加强板",
            internal_number="ACA674001",
            thickness_mm=Decimal("0.8000"),
            thickness_text="T=0.8",
            excel_filename="t13j.xlsx",
            excel_row=19,
        ),
    )
    db.commit()


def _generate(db: Session) -> BomProcurementObjectService:
    _seed_tree(db)
    service = BomProcurementObjectService()
    result = service.generate(
        db,
        BomProcurementObjectGenerateDTO(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
        ),
        current_user=SYSTEM_USER,
    )
    assert result.header.qa_passed is True
    assert result.object_count == 1
    return service


def _generate_shared_part(db: Session) -> BomProcurementObjectService:
    _seed_shared_part_tree(db)
    service = BomProcurementObjectService()
    result = service.generate(
        db,
        BomProcurementObjectGenerateDTO(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
        ),
        current_user=SYSTEM_USER,
    )
    assert result.header.qa_passed is True
    assert result.object_count == 2
    return service


def _object_rows(db: Session, service: BomProcurementObjectService):
    items, _, _ = service.paginate(
        db,
        BomProcurementObjectQuery(location="芜湖", vehicle_model="T13J", bom_version="V3", size=50),
    )
    return items


def _object_row(db: Session, service: BomProcurementObjectService):
    return _object_rows(db, service)[0]


def _set_supplier_cell(file_content: bytes, value: str | None) -> bytes:
    workbook = load_workbook(BytesIO(file_content))
    sheet = workbook.worksheets[0]
    header_row = 1
    supplier_col = None
    for col in range(1, (sheet.max_column or 1) + 1):
        if sheet.cell(header_row, col).value == HEADER_SUPPLIER:
            supplier_col = col
            break
    assert supplier_col is not None
    sheet.cell(2, supplier_col).value = value
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _drop_data_rows(file_content: bytes) -> bytes:
    workbook = load_workbook(BytesIO(file_content))
    sheet = workbook.worksheets[0]
    if sheet.max_row and sheet.max_row >= 2:
        sheet.delete_rows(2, sheet.max_row - 1)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _upload(service: BomProcurementObjectService, db: Session, content: bytes):
    return service.upload_suppliers(
        db,
        location="芜湖",
        vehicle_model="T13J",
        bom_version="V3",
        filename="po-supplier.xlsx",
        file_content=content,
        current_user=SYSTEM_USER,
    )


def test_template_sheet_order_and_existing_supplier(db: Session) -> None:
    service = _generate(db)
    content, filename = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    assert "T13J" in filename
    workbook = load_workbook(BytesIO(content))
    assert [sheet.title for sheet in workbook.worksheets] == ["采购对象", "供应商", "填写说明"]
    object_sheet = workbook.worksheets[0]
    assert object_sheet["C1"].value == HEADER_INTERNAL_NUMBER
    assert object_sheet["C2"].value == "ACA674001"
    assert object_sheet["G1"].value == HEADER_SUPPLIER
    assert object_sheet["G2"].value is None
    assert object_sheet["A1"].value == "零件号"
    assert workbook.worksheets[1]["B2"].value == FULL_NAME
    workbook.close()

    filled = _set_supplier_cell(content, "常庆")
    result = _upload(service, db, filled)
    assert result.updated_count == 1
    assert result.cleared_count == 0
    assert _object_row(db, service).supplier_name == FULL_NAME

    content2, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    workbook2 = load_workbook(BytesIO(content2))
    assert workbook2.worksheets[0]["G2"].value == FULL_NAME
    workbook2.close()


def test_missing_row_does_not_clear(db: Session) -> None:
    service = _generate(db)
    content, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    _upload(service, db, _set_supplier_cell(content, FULL_NAME))
    header_only, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    result = _upload(service, db, _drop_data_rows(header_only))
    assert result.updated_count == 0
    assert result.cleared_count == 0
    assert _object_row(db, service).supplier_name == FULL_NAME


def test_empty_cell_clears_supplier(db: Session) -> None:
    service = _generate(db)
    content, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    _upload(service, db, _set_supplier_cell(content, FULL_NAME))
    content2, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    result = _upload(service, db, _set_supplier_cell(content2, None))
    assert result.cleared_count == 1
    assert _object_row(db, service).supplier_name is None


def test_qa_failed_upload_still_updates_pending(db: Session) -> None:
    service = _generate(db)
    header = service.repository.get_list_by_scope(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    assert header is not None
    header.qa_passed = False
    header.qa_error_text = "清单已被部分删除，须重新生成后才能审核"
    db.commit()
    content, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    result = _upload(service, db, _set_supplier_cell(content, "常庆"))
    assert result.updated_count == 1
    assert _object_row(db, service).supplier_name == FULL_NAME


def test_approved_change_skipped(db: Session) -> None:
    service = _generate(db)
    content, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    _upload(service, db, _set_supplier_cell(content, FULL_NAME))
    service.batch_approve(
        db,
        BomProcurementObjectBatchApproveDTO(
            location="芜湖",
            vehicle_model="T13J",
            bom_version="V3",
            audit_remark="通过",
        ),
        current_user=SYSTEM_USER,
    )
    content2, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    result = _upload(service, db, _set_supplier_cell(content2, None))
    assert result.cleared_count == 0
    assert result.unchanged_count == 1
    assert _object_row(db, service).supplier_name == FULL_NAME


def test_in_house_nonempty_rejected(db: Session) -> None:
    service = _generate(db)
    row = _object_row(db, service)
    service.confirm(
        db,
        row.id,
        BomProcurementObjectConfirmDTO(confirm_status="in_house", confirm_remark="自制"),
        current_user=SYSTEM_USER,
    )
    content, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    with pytest.raises(BusinessException, match="自制禁止"):
        _upload(service, db, _set_supplier_cell(content, "常庆"))
    assert _object_row(db, service).supplier_name is None


def test_unchanged_disabled_full_name_skips_gate(db: Session) -> None:
    service = _generate(db)
    content, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    _upload(service, db, _set_supplier_cell(content, FULL_NAME))
    SupplierService().update(
        db,
        FULL_NAME,
        SupplierUpdateDTO(is_enabled=False),
        current_user=SYSTEM_USER,
    )
    content2, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    result = _upload(service, db, content2)
    assert result.unchanged_count == 1
    assert result.updated_count == 0
    assert _object_row(db, service).supplier_name == FULL_NAME


def test_same_part_number_fills_once_updates_all_rows(db: Session) -> None:
    service = _generate_shared_part(db)
    content, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.worksheets[0]
    assert sheet.max_row == 2
    assert sheet["A2"].value == "L1-STAMP"
    workbook.close()
    result = _upload(service, db, _set_supplier_cell(content, "常庆"))
    assert result.updated_count == 2
    matched = [row for row in _object_rows(db, service) if row.part_number == "L1-STAMP"]
    assert len(matched) == 2
    assert {row.supplier_name for row in matched} == {FULL_NAME}


def test_out_of_scope_part_number_skipped_with_detail(db: Session) -> None:
    service = _generate(db)
    content, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.worksheets[0]
    sheet.cell(2, 1).value = "NOT-IN-SCOPE"
    for col in range(2, (sheet.max_column or 1) + 1):
        if sheet.cell(1, col).value == HEADER_SUPPLIER:
            sheet.cell(2, col).value = "常庆"
            break
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    result = _upload(service, db, buffer.getvalue())
    assert result.updated_count == 0
    assert result.skipped_count == 1
    assert result.skipped_details
    assert "NOT-IN-SCOPE" in result.skipped_details[0]
    assert "第2行" in result.skipped_details[0]
    assert _object_row(db, service).supplier_name is None


def test_duplicate_part_number_in_file_rejected(db: Session) -> None:
    service = _generate(db)
    content, _ = service.build_supplier_template(
        db, location="芜湖", vehicle_model="T13J", bom_version="V3"
    )
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.worksheets[0]
    for col in range(1, (sheet.max_column or 1) + 1):
        sheet.cell(3, col).value = sheet.cell(2, col).value
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    with pytest.raises(BusinessException, match="零件号重复"):
        _upload(service, db, buffer.getvalue())
