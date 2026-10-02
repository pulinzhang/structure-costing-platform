"""一级件出现在总成之前时，整本拒绝，不补 UNKNOWN-ASM。"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook

from app.exceptions.base import BusinessException
from app.parsers.bom.bom_parser import parse_bom_excel


def _workbook(data_rows: list[dict[str, object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.cell(1, 3, "总成零件号")
    sheet.cell(1, 9, "一级零件号")
    for offset, row in enumerate(data_rows, start=2):
        if "seq" in row:
            sheet.cell(offset, 1, row["seq"])
        if row.get("asm"):
            sheet.cell(offset, 3, row["asm"])
        if row.get("l1"):
            sheet.cell(offset, 9, row["l1"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_l1_before_any_assembly_rejects_whole_workbook() -> None:
    content = _workbook([{"seq": 1, "l1": "502000001AA"}])
    with pytest.raises(BusinessException, match="整本未导入") as raised:
        parse_bom_excel(content)
    message = str(raised.value)
    assert "第 2 行" in message
    assert "还没有总成" in message
    assert "UNKNOWN-ASM" not in message
    assert "未识别总成" not in message


def test_orphan_l1_rejects_even_when_a_later_row_has_assembly() -> None:
    content = _workbook(
        [
            {"seq": 1, "l1": "502000001AA"},
            {"seq": 2, "l1": "502000002AA"},
            {"seq": 3, "asm": "501000001AA", "l1": "502000003AA"},
        ]
    )
    with pytest.raises(BusinessException, match="整本未导入") as raised:
        parse_bom_excel(content)
    message = str(raised.value)
    assert "第 2 行" in message
    assert "第 3 行" in message
    assert "UNKNOWN-ASM" not in message


def test_l1_after_assembly_still_parses() -> None:
    content = _workbook(
        [
            {"seq": 1, "asm": "501000001AA", "l1": "502000001AA"},
            {"l1": "502000002AA"},
        ]
    )
    parsed = parse_bom_excel(content)
    part_numbers = [node.part_number for node in parsed.nodes]
    assert part_numbers == ["501000001AA", "502000001AA", "502000002AA"]
    assert parsed.nodes[1].parent_index == 0
    assert parsed.nodes[2].parent_index == 0
    assert all(not number.startswith("UNKNOWN-ASM") for number in part_numbers)
